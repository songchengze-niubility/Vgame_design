# -*- coding: utf-8 -*-
"""Sync monster-level recommended power to the authoritative player power curve."""

from __future__ import annotations

import argparse
import json
import math
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.comments import Comment


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_MONSTER = ROOT / "numerical-workflows/v2/framework/01_源表/战斗公式.xlsx"
DEFAULT_PLAYER = ROOT / "numerical-workflows/v2/framework/01_源表/养成属性分配.xlsx"
PROGRESS_KEYS = ("E", "G", "H", "I", "J", "K", "L")


def excel_round(value: float) -> int:
    return math.floor(value + 0.5)


def load_player_model(path: Path) -> dict:
    wb = load_workbook(path, data_only=True)
    ws = wb["玩家战力曲线"]
    module = wb["模块规划"]
    baseline = wb["_曲线基线"]

    stage_rows = []
    for row in ws.iter_rows(min_row=12, max_row=ws.max_row, values_only=True):
        stage, monster_level, team_power = row[1], row[41], row[46]
        if all(isinstance(value, (int, float)) for value in (monster_level, team_power)):
            stage_rows.append((int(monster_level), int(team_power), stage))
    if not stage_rows:
        raise RuntimeError("玩家战力曲线 AP/AU 没有可用的已计算值")
    if any(a[0] > b[0] for a, b in zip(stage_rows, stage_rows[1:])):
        raise RuntimeError("玩家战力曲线 AP（怪物等级）不是升序")

    unique_mapping = {}
    for monster_level, team_power, stage in stage_rows:
        current = unique_mapping.get(monster_level)
        if current is None or team_power >= current[0]:
            unique_mapping[monster_level] = (team_power, stage)

    parameters = {
        "base_crit": baseline["L3"].value,
        "base_crit_damage": baseline["L4"].value,
        "defense_percent_max": ws["O9"].value,
        "attack_percent_max": ws["T9"].value,
        "crit_growth": ws["AH9"].value,
        "crit_damage_growth": ws["AI9"].value,
        "damage_bonus_growth": ws["AJ9"].value,
        "hp_power": ws["AS3"].value,
        "attack_power": ws["AS4"].value,
        "defense_power": ws["AS5"].value,
        "crit_power": ws["AS6"].value,
        "crit_damage_power": ws["AS7"].value,
        "damage_bonus_power": ws["AS8"].value,
        "energy_power": ws["AS9"].value,
        "cooldown_power": ws["AS10"].value,
        "max_energy": ws["AU3"].value,
        "max_cooldown": ws["AU4"].value,
        "weight_h": ws["H10"].value,
        "weight_i": ws["I10"].value,
        "weight_j": ws["J10"].value,
        "weight_k": ws["K10"].value,
        "weight_l": ws["L10"].value,
        "hp_initial": module["D10"].value,
        "hp_level": module["D11"].value,
        "hp_star": module["D13"].value,
        "hp_containment": module["D15"].value,
        "hp_mapping": module["D17"].value,
        "hp_ticket": module["D16"].value,
        "hp_rv": module["D18"].value,
        "hp_star_percent": module["G14"].value,
        "hp_mapping_percent": module["J17"].value,
        "attack_initial": module["E10"].value,
        "attack_level": module["E11"].value,
        "attack_star": module["E13"].value,
        "attack_containment": module["E15"].value,
        "attack_mapping": module["E17"].value,
        "attack_ticket": module["E16"].value,
        "attack_rv": module["E18"].value,
        "defense_initial": module["F10"].value,
        "defense_level": module["F11"].value,
        "defense_star": module["F13"].value,
        "defense_containment": module["F15"].value,
        "defense_mapping": module["F17"].value,
        "defense_ticket": module["F16"].value,
        "defense_rv": module["F18"].value,
    }
    invalid = [name for name, value in parameters.items() if not isinstance(value, (int, float))]
    if invalid:
        raise RuntimeError(f"玩家战力参数不是有效数值：{', '.join(invalid)}")

    result = {
        "unique_mapping": sorted(
            (level, power, stage) for level, (power, stage) in unique_mapping.items()
        ),
        "end_progress": {key: float(ws[f"{key}4"].value) for key in PROGRESS_KEYS},
        "end_power": int(stage_rows[-1][1]),
        "end_stage": stage_rows[-1][2],
        "curve_max_level": int(stage_rows[-1][0]),
        "parameters": {name: float(value) for name, value in parameters.items()},
    }
    wb.close()
    return result


def calculate_team_power(progress: dict[str, float], p: dict[str, float]) -> int:
    e, g, h, i, j, k, l = (progress[key] for key in PROGRESS_KEYS)
    weighted = (
        p["weight_h"] * h
        + p["weight_i"] * i
        + p["weight_j"] * j
        + p["weight_k"] * k
        + p["weight_l"] * l
    )
    weight_sum = sum(p[f"weight_{key}"] for key in "hijkl")
    if weight_sum == 0:
        raise RuntimeError("玩家战力曲线二级属性权重合计为0")

    raw_hp = (
        (e * p["hp_initial"] + g * p["hp_level"] + h * p["hp_star"])
        * (1 + p["hp_star_percent"] * h)
        + i * p["hp_containment"]
        + j * p["hp_mapping"]
        + k * p["hp_ticket"]
        + l * p["hp_rv"]
    ) * (1 + p["hp_mapping_percent"] * j)
    raw_attack = e * p["attack_initial"] + g * p["attack_level"] + h * p["attack_star"]
    extra_attack = (
        i * p["attack_containment"]
        + j * p["attack_mapping"]
        + k * p["attack_ticket"]
        + l * p["attack_rv"]
    )
    attack = raw_attack * (1 + p["attack_percent_max"] * weighted / weight_sum) + extra_attack
    raw_defense = e * p["defense_initial"] + g * p["defense_level"] + h * p["defense_star"]
    extra_defense = (
        i * p["defense_containment"]
        + j * p["defense_mapping"]
        + k * p["defense_ticket"]
        + l * p["defense_rv"]
    )
    defense = raw_defense * (1 + p["defense_percent_max"] * weighted / weight_sum) + extra_defense
    crit = min(1, p["base_crit"] + p["crit_growth"] * weighted)
    crit_damage = p["base_crit_damage"] + p["crit_damage_growth"] * weighted
    damage_bonus = 1 + p["damage_bonus_growth"] * weighted
    base_power = raw_hp * p["hp_power"] + attack * p["attack_power"] + defense * p["defense_power"]
    secondary = (
        1
        + crit * p["crit_power"]
        + crit_damage * p["crit_damage_power"]
        + (damage_bonus - 1) * p["damage_bonus_power"]
        + weighted * p["max_energy"] * p["energy_power"]
        + weighted * p["max_cooldown"] * p["cooldown_power"]
    )
    return excel_round(excel_round(base_power * secondary) * 5)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", type=Path, default=DEFAULT_MONSTER)
    parser.add_argument("--player-workbook", type=Path, default=DEFAULT_PLAYER)
    args = parser.parse_args()

    print("读取玩家战力曲线…", flush=True)
    model = load_player_model(args.player_workbook)
    p = model["parameters"]
    replicated_end = calculate_team_power(model["end_progress"], p)
    if replicated_end != model["end_power"]:
        raise RuntimeError(f"玩家战力终点复现失败：公式复现={replicated_end}，源表AU={model['end_power']}")
    full_progress = {key: 1.0 for key in PROGRESS_KEYS}
    full_power = calculate_team_power(full_progress, p)

    print("生成统一推荐战力映射…", flush=True)
    wb = load_workbook(args.workbook, data_only=False)
    ws = wb["怪物标准模型"]
    for sheet_name in ("_推荐战力映射", "_推荐战力参数"):
        if sheet_name in wb.sheetnames:
            wb.remove(wb[sheet_name])
    map_ws = wb.create_sheet("_推荐战力映射")

    input_rows = {
        "curve_max_level": model["curve_max_level"],
        **{f"end_{key.lower()}": model["end_progress"][key] for key in PROGRESS_KEYS},
        **p,
    }
    map_ws["X1"] = "参数"
    map_ws["Y1"] = "值"
    map_ws["Z1"] = "来源"
    param_rows = {}
    for row, (name, value) in enumerate(input_rows.items(), 2):
        map_ws.cell(row, 24).value = name
        map_ws.cell(row, 25).value = value
        map_ws.cell(row, 26).value = "养成属性分配.xlsx/玩家战力曲线及模块规划"
        param_rows[name] = row
    ref = {name: f"$Y${row}" for name, row in param_rows.items()}

    map_ws["U1"] = "玩家曲线怪物等级"
    map_ws["V1"] = "玩家曲线队伍战力"
    map_ws["W1"] = "玩家进度"
    for row, (level, power, stage) in enumerate(model["unique_mapping"], 2):
        map_ws.cell(row, 21).value = level
        map_ws.cell(row, 22).value = power
        map_ws.cell(row, 23).value = stage
    mapping_end_row = 1 + len(model["unique_mapping"])

    headers = [
        "怪物等级", "统一推荐战力", "口径", "延伸进度", "基础进度", "角色等级进度",
        "星盘/升星进度", "收容物进度", "映射装备进度", "旅券进度", "房车进度",
        "玩家生命", "玩家攻击", "玩家防御", "二级属性权重", "单人基础战力",
        "二级属性倍率", "单人战力", "队伍战力",
    ]
    for col, header in enumerate(headers, 1):
        map_ws.cell(1, col).value = header

    weight_sum = "+".join(ref[f"weight_{key}"] for key in "hijkl")
    progress_columns = {"E": 5, "G": 6, "H": 7, "I": 8, "J": 9, "K": 10, "L": 11}
    for row, level in enumerate(range(1, 501), 2):
        map_ws.cell(row, 1).value = level
        map_ws.cell(row, 4).value = f"=MAX(0,(A{row}-{ref['curve_max_level']})/(500-{ref['curve_max_level']}))"
        for key, col in progress_columns.items():
            map_ws.cell(row, col).value = f"={ref[f'end_{key.lower()}']}+(1-{ref[f'end_{key.lower()}']})*$D{row}"
        weighted = (
            f"{ref['weight_h']}*G{row}+{ref['weight_i']}*H{row}+{ref['weight_j']}*I{row}+"
            f"{ref['weight_k']}*J{row}+{ref['weight_l']}*K{row}"
        )
        map_ws.cell(row, 15).value = f"={weighted}"
        map_ws.cell(row, 12).value = (
            f"=((E{row}*{ref['hp_initial']}+F{row}*{ref['hp_level']}+G{row}*{ref['hp_star']})"
            f"*(1+{ref['hp_star_percent']}*G{row})+H{row}*{ref['hp_containment']}+I{row}*{ref['hp_mapping']}+"
            f"J{row}*{ref['hp_ticket']}+K{row}*{ref['hp_rv']})*(1+{ref['hp_mapping_percent']}*I{row})"
        )
        map_ws.cell(row, 13).value = (
            f"=(E{row}*{ref['attack_initial']}+F{row}*{ref['attack_level']}+G{row}*{ref['attack_star']})"
            f"*(1+{ref['attack_percent_max']}*O{row}/({weight_sum}))+H{row}*{ref['attack_containment']}+"
            f"I{row}*{ref['attack_mapping']}+J{row}*{ref['attack_ticket']}+K{row}*{ref['attack_rv']}"
        )
        map_ws.cell(row, 14).value = (
            f"=(E{row}*{ref['defense_initial']}+F{row}*{ref['defense_level']}+G{row}*{ref['defense_star']})"
            f"*(1+{ref['defense_percent_max']}*O{row}/({weight_sum}))+H{row}*{ref['defense_containment']}+"
            f"I{row}*{ref['defense_mapping']}+J{row}*{ref['defense_ticket']}+K{row}*{ref['defense_rv']}"
        )
        map_ws.cell(row, 16).value = f"=L{row}*{ref['hp_power']}+M{row}*{ref['attack_power']}+N{row}*{ref['defense_power']}"
        map_ws.cell(row, 17).value = (
            f"=1+MIN(1,{ref['base_crit']}+{ref['crit_growth']}*O{row})*{ref['crit_power']}+"
            f"({ref['base_crit_damage']}+{ref['crit_damage_growth']}*O{row})*{ref['crit_damage_power']}+"
            f"({ref['damage_bonus_growth']}*O{row})*{ref['damage_bonus_power']}+"
            f"O{row}*{ref['max_energy']}*{ref['energy_power']}+O{row}*{ref['max_cooldown']}*{ref['cooldown_power']}"
        )
        map_ws.cell(row, 18).value = f"=ROUND(P{row}*Q{row},0)"
        map_ws.cell(row, 19).value = f"=ROUND(R{row}*5,0)"
        map_ws.cell(row, 2).value = (
            f"=IF(A{row}<={ref['curve_max_level']},"
            f"XLOOKUP(A{row},$U$2:$U${mapping_end_row},$V$2:$V${mapping_end_row},\"\",1),S{row})"
        )
        map_ws.cell(row, 3).value = f'=IF(A{row}<={ref["curve_max_level"]},"玩家曲线直接映射","同公式延伸至100%养成")'
    map_ws.sheet_state = "hidden"

    value_wb = load_workbook(args.workbook, data_only=True, read_only=True)
    value_ws = value_wb["怪物标准模型"]
    levels = [value_ws[f"B{row}"].value for row in range(12, 512)]
    value_wb.close()
    if levels != list(range(1, 501)):
        raise RuntimeError("怪物标准模型等级必须连续为1~500")
    ws["BD9"] = "推荐战力"
    ws["BD10"] = "玩家曲线统一口径"
    ws["BD11"] = "基准推荐战力"
    ws["BD9"]._style = copy(ws["AZ9"]._style)
    ws["BD10"]._style = copy(ws["AZ9"]._style)
    ws["BD11"]._style = copy(ws["BC11"]._style)
    ws["BD9"].alignment = copy(ws["AZ9"].alignment)
    ws["BD10"].alignment = copy(ws["AZ9"].alignment)
    ws["BD11"].alignment = copy(ws["BC11"].alignment)
    ws["BD11"].comment = Comment(
        "唯一权威：养成属性分配.xlsx/玩家战力曲线。当前曲线覆盖范围内按AP怪物等级与AU预估队伍战力直接映射；"
        "超过当前曲线终点后，从源表配置的终点养成进度继续使用同一套玩家战力公式延伸，至Lv500达到全部养成100%。"
        "不使用怪物EHP反推，不另建怪物属性战力公式。",
        "Codex",
    )
    ws.column_dimensions["BD"].width = 20
    for row in range(12, 512):
        ws.cell(row, 56)._style = copy(ws[f"BC{row}"]._style)
        ws.cell(row, 56).number_format = "#,##0"
        ws.cell(row, 56).value = f"='_推荐战力映射'!B{row - 10}"

    wb.calculation.calcMode = "auto"
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    print("保存工作簿…", flush=True)
    wb.save(args.workbook)

    print(json.dumps({
        "workbook": str(args.workbook),
        "sheet": "怪物标准模型",
        "column": "BD",
        "levels": 500,
        "authority": "养成属性分配.xlsx/玩家战力曲线",
        "curve_last_stage": model["end_stage"],
        "curve_max_monster_level": model["curve_max_level"],
        "curve_end_power": model["end_power"],
        "replicated_curve_end_power": replicated_end,
        "full_progress_power": full_power,
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
