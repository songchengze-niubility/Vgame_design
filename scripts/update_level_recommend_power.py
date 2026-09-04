# -*- coding: utf-8 -*-
"""Fill UIlevel.Power from the unified monster-level recommended-power curve."""

from __future__ import annotations

import argparse
import json
from bisect import bisect_left
from pathlib import Path

from openpyxl import load_workbook

from update_monster_recommend_power import calculate_team_power, load_player_model


DESIGN_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_PLAYER = (
    DESIGN_ROOT / "numerical-workflows/v2/framework/01_源表/养成属性分配.xlsx"
)
import vgame_paths

DEFAULT_LEVEL = vgame_paths.config_datas() / "level" / "level.xlsx"
DEFAULT_UILEVEL = vgame_paths.config_datas() / "level" / "UIlevel.xlsx"
DEFAULT_DELIVERY = (
    DESIGN_ROOT / "numerical-workflows/v2/framework/06_最终配置/全玩法难度/UIlevel.xlsx"
)
PROGRESS_KEYS = ("E", "G", "H", "I", "J", "K", "L")
SIGNIFICANT_DIGITS = 3


def round_to_significant_int(value: int, digits: int = SIGNIFICANT_DIGITS) -> int:
    """Round a positive integer to the requested significant digits, half up."""
    if value <= 0 or digits <= 0:
        return value
    places_to_drop = max(0, len(str(value)) - digits)
    if places_to_drop == 0:
        return value
    factor = 10**places_to_drop
    return ((value + factor // 2) // factor) * factor


def build_power_map(player_workbook: Path) -> tuple[dict[int, int], dict]:
    model = load_player_model(player_workbook)
    curve = model["unique_mapping"]
    curve_levels = [level for level, _, _ in curve]
    curve_powers = [power for _, power, _ in curve]
    max_level = model["curve_max_level"]
    end_progress = model["end_progress"]
    parameters = model["parameters"]

    result: dict[int, int] = {}
    for level in range(1, 501):
        if level <= max_level:
            index = bisect_left(curve_levels, level)
            if index >= len(curve_levels):
                raise RuntimeError(f"玩家战力曲线缺少怪物等级 {level} 的映射锚点")
            result[level] = round_to_significant_int(int(curve_powers[index]))
            continue

        extension = (level - max_level) / (500 - max_level)
        progress = {
            key: end_progress[key] + (1.0 - end_progress[key]) * extension
            for key in PROGRESS_KEYS
        }
        result[level] = round_to_significant_int(
            calculate_team_power(progress, parameters)
        )

    return result, model


def read_headers(ws) -> dict[str, int]:
    return {cell.value: cell.column for cell in ws[1] if cell.value}


def fill_workbook(
    level_workbook: Path,
    uilevel_workbook: Path,
    output: Path,
    power_map: dict[int, int],
) -> dict:
    level_wb = load_workbook(level_workbook, read_only=True, data_only=True)
    level_ws = level_wb["level"]
    level_headers = read_headers(level_ws)
    required_level = {"LevelId", "Grade"}
    if missing := required_level - level_headers.keys():
        raise RuntimeError(f"level.xlsx 缺少字段：{', '.join(sorted(missing))}")

    grades: dict[int, int | None] = {}
    duplicate_level_ids: set[int] = set()
    conflicting_grade_ids: set[int] = set()
    for row in level_ws.iter_rows(min_row=5, values_only=True):
        level_id = row[level_headers["LevelId"] - 1]
        if not isinstance(level_id, (int, float)):
            continue
        level_id = int(level_id)
        grade = row[level_headers["Grade"] - 1]
        grade = int(grade) if isinstance(grade, (int, float)) else None
        if level_id in grades:
            duplicate_level_ids.add(level_id)
            current = grades[level_id]
            if current is not None and grade is not None and current != grade:
                conflicting_grade_ids.add(level_id)
            elif current is None and grade is not None:
                grades[level_id] = grade
        else:
            grades[level_id] = grade
    level_wb.close()
    if conflicting_grade_ids:
        sample = ", ".join(map(str, sorted(conflicting_grade_ids)[:10]))
        raise RuntimeError(f"level.xlsx 重复 LevelId 的 Grade 冲突：{sample}")

    wb = load_workbook(uilevel_workbook, data_only=False)
    ws = wb["level"]
    headers = read_headers(ws)
    required_ui = {"LevelId", "Power", "LevelType"}
    if missing := required_ui - headers.keys():
        raise RuntimeError(f"UIlevel.xlsx 缺少字段：{', '.join(sorted(missing))}")

    power_col = headers["Power"]
    filled = 0
    cleared = 0
    missing_grade: list[dict] = []
    out_of_range: list[dict] = []
    seen_ui_ids: set[int] = set()
    duplicate_ui_ids: set[int] = set()

    for row_index in range(5, ws.max_row + 1):
        level_id = ws.cell(row_index, headers["LevelId"]).value
        if not isinstance(level_id, (int, float)):
            continue
        level_id = int(level_id)
        if level_id in seen_ui_ids:
            duplicate_ui_ids.add(level_id)
        seen_ui_ids.add(level_id)

        grade = grades.get(level_id)
        if isinstance(grade, int) and grade in power_map:
            ws.cell(row_index, power_col).value = power_map[grade]
            filled += 1
            continue

        if ws.cell(row_index, power_col).value not in (None, ""):
            cleared += 1
        ws.cell(row_index, power_col).value = None
        item = {
            "LevelId": level_id,
            "LevelType": ws.cell(row_index, headers["LevelType"]).value,
            "Grade": grade,
        }
        if isinstance(grade, int):
            out_of_range.append(item)
        else:
            missing_grade.append(item)

    if duplicate_ui_ids:
        sample = ", ".join(map(str, sorted(duplicate_ui_ids)[:10]))
        raise RuntimeError(f"UIlevel.xlsx 存在重复 LevelId：{sample}")

    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    wb.close()
    return {
        "output": str(output),
        "filled": filled,
        "cleared_stale_values": cleared,
        "missing_or_dynamic_grade": len(missing_grade),
        "grade_out_of_range": len(out_of_range),
        "duplicate_level_ids_resolved": sorted(duplicate_level_ids),
        "missing_or_dynamic_examples": missing_grade[:20],
        "grade_out_of_range_examples": out_of_range[:20],
    }


def verify(output: Path, level_workbook: Path, power_map: dict[int, int]) -> dict:
    level_wb = load_workbook(level_workbook, read_only=True, data_only=True)
    level_ws = level_wb["level"]
    level_headers = read_headers(level_ws)
    grades = {
        int(row[level_headers["LevelId"] - 1]): int(row[level_headers["Grade"] - 1])
        for row in level_ws.iter_rows(min_row=5, values_only=True)
        if isinstance(row[level_headers["LevelId"] - 1], (int, float))
        and isinstance(row[level_headers["Grade"] - 1], (int, float))
    }
    level_wb.close()

    wb = load_workbook(output, read_only=True, data_only=True)
    ws = wb["level"]
    headers = read_headers(ws)
    errors = []
    populated = 0
    for row in ws.iter_rows(min_row=5, values_only=True):
        level_id = row[headers["LevelId"] - 1]
        if not isinstance(level_id, (int, float)):
            continue
        level_id = int(level_id)
        actual = row[headers["Power"] - 1]
        grade = grades.get(level_id)
        expected = power_map.get(grade) if grade is not None else None
        if actual not in (None, ""):
            populated += 1
        if actual != expected:
            errors.append({"LevelId": level_id, "Grade": grade, "expected": expected, "actual": actual})
    wb.close()
    return {"populated": populated, "errors": len(errors), "error_examples": errors[:20]}


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--player-workbook", type=Path, default=DEFAULT_PLAYER)
    parser.add_argument("--level-workbook", type=Path, default=DEFAULT_LEVEL)
    parser.add_argument("--uilevel-workbook", type=Path, default=DEFAULT_UILEVEL)
    parser.add_argument("--output", type=Path, default=DEFAULT_UILEVEL)
    parser.add_argument("--delivery-output", type=Path, default=DEFAULT_DELIVERY)
    args = parser.parse_args()

    power_map, model = build_power_map(args.player_workbook)
    result = fill_workbook(args.level_workbook, args.uilevel_workbook, args.output, power_map)
    result["verification"] = verify(args.output, args.level_workbook, power_map)

    if args.delivery_output.resolve() != args.output.resolve():
        delivery = fill_workbook(
            args.level_workbook, args.output, args.delivery_output, power_map
        )
        delivery["verification"] = verify(args.delivery_output, args.level_workbook, power_map)
        result["delivery"] = delivery

    result.update(
        {
            "authority": "养成属性分配.xlsx/玩家战力曲线",
            "mapping": "UIlevel.LevelId -> level.Grade -> unified recommended power",
            "rounding": f"round half up to {SIGNIFICANT_DIGITS} significant digits",
            "curve_max_monster_level": model["curve_max_level"],
            "power_at_level_1": power_map[1],
            "power_at_level_500": power_map[500],
        }
    )
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
