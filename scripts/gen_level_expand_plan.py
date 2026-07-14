# -*- coding: utf-8 -*-
"""把「主线关卡扩展」Design Doc md 转成 openspec 风格策划案 xlsx。
3 个 sheet：修订记录与概述 / 系统功能正文 / 数据埋点方案。
通用化渲染：按 md 的 ## / ### 结构遍历，表格/bullets/段落自动识别；埋点从 md 解析。"""
import re
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DESIGN_DIR = os.path.join(REPO_ROOT, 'design')
OUT_DIR = os.path.join(os.path.dirname(REPO_ROOT), 'Output')
os.makedirs(OUT_DIR, exist_ok=True)

MD_NAME = '主线关卡扩展-支线·小游戏·引导·节点演出Design Doc.md'
OUT_NAME = '主线关卡扩展策划案.xlsx'

TITLE_FONT = Font(bold=True, size=14, color='1F3864')
SECTION_FILL = PatternFill('solid', start_color='305496')
SECTION_FONT = Font(bold=True, size=12, color='FFFFFF')
SUB_FILL = PatternFill('solid', start_color='D9E1F2')
SUB_FONT = Font(bold=True, size=11, color='1F3864')
LABEL_FONT = Font(bold=True, size=11, color='1F3864')
HDR_FILL = PatternFill('solid', start_color='8EA9DB')
HDR_FONT = Font(bold=True, color='FFFFFF')
WRAP = Alignment(wrap_text=True, vertical='top')
TOP = Alignment(vertical='top', horizontal='left')
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
THIN = Side(style='thin', color='BFBFBF')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
NCOL = 7


def clean(t):
    return t.replace('**', '').replace('`', '')


def split_sections(md):
    sections, cur_title, cur_body = [], None, []
    for line in md.splitlines():
        if line.startswith('## ') and not line.startswith('### '):
            if cur_title is not None:
                sections.append((cur_title, cur_body))
            cur_title, cur_body = line[3:].strip(), []
        elif cur_title is not None:
            cur_body.append(line)
    if cur_title is not None:
        sections.append((cur_title, cur_body))
    return sections


def split_subsections(body):
    subs, cur, cur_lines = [], None, []
    for line in body:
        if line.startswith('### '):
            if cur is not None:
                subs.append((cur, cur_lines))
            cur, cur_lines = line[4:].strip(), []
        elif cur is not None:
            cur_lines.append(line)
        elif line.strip():
            if not subs or subs[-1][0] != '':
                subs.append(('', []))
            subs[-1][1].append(line)
    if cur is not None:
        subs.append((cur, cur_lines))
    return subs


def parse_table(lines):
    rows = []
    for ln in lines:
        s = ln.strip()
        if not s.startswith('|') or re.match(r'^\|[\s:|-]+\|$', s):
            continue
        rows.append([clean(c.strip()) for c in s.strip('|').split('|')])
    return rows


def parse_bullets(lines):
    items = []
    for ln in lines:
        if not ln.strip():
            continue
        m = re.match(r'^(\s*)- \*\*(.+?)\*\*[：:](.*)$', ln)
        if m:
            items.append([m.group(2).strip(), m.group(3).strip(), len(m.group(1))])
            continue
        m2 = re.match(r'^(\s*)-\s+(.*)$', ln)
        if m2:
            indent, text = len(m2.group(1)), m2.group(2).strip()
            if items and indent > items[-1][2]:
                items[-1][1] = (items[-1][1] + '\n' + text) if items[-1][1] else text
            else:
                items.append(['', text, indent])
        elif items and ln.strip():
            items[-1][1] = (items[-1][1] + '\n' + ln.strip()) if items[-1][1] else ln.strip()
    return [[f, clean(v)] for f, v, _ in items]


def subsection_text(lines):
    out, in_code = [], False
    for ln in lines:
        s = ln.rstrip()
        if s.strip().startswith('```'):
            in_code = not in_code
            continue
        if not s.strip():
            continue
        s = re.sub(r'^-\s+', '', s.strip())
        out.append(s)
    return '\n'.join(out)


def has_table(lines):
    return bool(parse_table(lines))


def meta_dict(sections):
    d = {}
    for r in parse_table(sections.get('元信息', [])):
        if len(r) >= 2:
            d[r[0]] = r[1]
    return d


def style_row(ws, row, ncol=NCOL):
    for ci in range(1, ncol + 1):
        ws.cell(row=row, column=ci).border = BORDER


def merge_row(ws, row, text, font, fill=None, align=None, ncol=NCOL):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncol)
    c = ws.cell(row=row, column=1, value=text)
    c.font = font
    if fill:
        c.fill = fill
    c.alignment = align or Alignment(horizontal='left', vertical='center', wrap_text=True)
    style_row(ws, row, ncol)


def render_content(ws, r, lines, ncol=NCOL):
    """渲染一个子节的内容：优先表格，其次bullets，最后段落文本。返回下一行号。"""
    rows = parse_table(lines)
    if rows:
        hdr = rows[0]
        for ci, h in enumerate(hdr, 2):
            c = ws.cell(row=r, column=ci, value=h)
            c.fill = HDR_FILL
            c.font = HDR_FONT
            c.alignment = CENTER
        style_row(ws, r)
        r += 1
        for row in rows[1:]:
            if not any(row):
                continue
            for ci, val in enumerate(row, 2):
                c = ws.cell(row=r, column=ci, value=val)
                c.alignment = WRAP
            style_row(ws, r)
            ws.row_dimensions[r].height = max(20, min(80, max((str(v).count('\n') + 1) * 16 for v in row if v) + 4))
            r += 1
        return r + 1
    bullets = parse_bullets(lines)
    if bullets:
        for f, v in bullets:
            ws.cell(row=r, column=2, value=f).alignment = TOP
            ws.cell(row=r, column=2).font = Font(bold=True)
            ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=ncol)
            ws.cell(row=r, column=3, value=v).alignment = WRAP
            style_row(ws, r)
            ws.row_dimensions[r].height = max(20, min(120, v.count('\n') * 16 + 22))
            r += 1
        return r + 1
    txt = subsection_text(lines)
    if txt:
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=ncol)
        ws.cell(row=r, column=2, value=txt).alignment = WRAP
        style_row(ws, r)
        ws.row_dimensions[r].height = max(28, min(140, txt.count('\n') * 18 + 28))
        r += 2
    return r


def write_overview_sheet(wb, doc_title, meta, bg, target):
    ws = wb.create_sheet('修订记录与概述')
    ws.column_dimensions['A'].width = 16
    for col in 'BCDEFG':
        ws.column_dimensions[col].width = 18
    r = 1
    merge_row(ws, r, f'{doc_title} 策划案', TITLE_FONT, align=Alignment(horizontal='center', vertical='center'))
    ws.row_dimensions[r].height = 28
    r += 2
    merge_row(ws, r, '修订记录', LABEL_FONT)
    r += 1
    hdr = ['版本号', '修订人', '日期', '简介', '修订内容索引']
    for ci, h in enumerate(hdr, 2):
        c = ws.cell(row=r, column=ci, value=h)
        c.fill = HDR_FILL
        c.font = HDR_FONT
        c.alignment = CENTER
    style_row(ws, r)
    r += 1
    rev = ['v1', meta.get('作者', '策划'), meta.get('创建日期', '2026-07-02'), '初版策划案', '--']
    for ci, v in enumerate(rev, 2):
        c = ws.cell(row=r, column=ci, value=v)
        c.alignment = WRAP
    style_row(ws, r)
    r += 2
    merge_row(ws, r, '文档属性', LABEL_FONT)
    r += 1
    for ci, h in enumerate(['项', '内容'], 2):
        c = ws.cell(row=r, column=ci, value=h)
        c.fill = HDR_FILL
        c.font = HDR_FONT
        c.alignment = CENTER
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=7)
    style_row(ws, r)
    r += 1
    for k, v in meta.items():
        ws.cell(row=r, column=2, value=k).alignment = TOP
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=7)
        ws.cell(row=r, column=3, value=v).alignment = WRAP
        style_row(ws, r)
        r += 1
    r += 1
    merge_row(ws, r, '系统概述', LABEL_FONT)
    r += 1
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=7)
    ws.cell(row=r, column=2, value='背景：' + bg).alignment = WRAP
    style_row(ws, r)
    r += 1
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=7)
    ws.cell(row=r, column=2, value='目标：' + target).alignment = WRAP
    style_row(ws, r)
    ws.freeze_panes = 'A2'


BODY_ORDER = [
    '背景与目标', '节点拓扑与类型扩展', '核心规则', '节点切换演出设计',
    '数值与产出框架', '配置表初版结构（落地点汇总）', '美术与演出参考',
    '风险评估', '工作假设',
]


def write_body_sheet(wb, sections):
    ws = wb.create_sheet('系统功能正文')
    ws.column_dimensions['A'].width = 8
    ws.column_dimensions['B'].width = 24
    for col in 'CDEFG':
        ws.column_dimensions[col].width = 20
    r = 1
    merge_row(ws, r, '系统功能正文', TITLE_FONT, align=Alignment(horizontal='center', vertical='center'))
    ws.row_dimensions[r].height = 26
    r += 2
    for idx, sec_key in enumerate(BODY_ORDER, 1):
        if sec_key not in sections:
            continue
        body = sections[sec_key]
        merge_row(ws, r, f'{idx}. {sec_key}', SECTION_FONT, SECTION_FILL,
                  align=Alignment(horizontal='left', vertical='center'))
        r += 2
        subs = split_subsections(body)
        if subs and any(s[0] for s in subs):
            for sub, lines in subs:
                if not sub and not any(lines):
                    continue
                if sub:
                    ws.cell(row=r, column=1, value='').alignment = TOP
                    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=NCOL)
                    c = ws.cell(row=r, column=2, value=sub)
                    c.font = SUB_FONT
                    c.fill = SUB_FILL
                    c.alignment = Alignment(vertical='center')
                    style_row(ws, r)
                    r += 1
                r = render_content(ws, r, lines)
        else:
            r = render_content(ws, r, body)
    ws.freeze_panes = 'A2'


def write_analytics_sheet(wb, sections):
    ws = wb.create_sheet('数据埋点方案')
    ws.column_dimensions['A'].width = 24
    ws.column_dimensions['B'].width = 30
    ws.column_dimensions['C'].width = 44
    ws.column_dimensions['D'].width = 30
    r = 1
    merge_row(ws, r, '数据埋点方案', TITLE_FONT, align=Alignment(horizontal='center', vertical='center'))
    ws.row_dimensions[r].height = 26
    r += 2
    body = sections.get('数据埋点方案', [])
    subs = split_subsections(body)
    for sub, lines in subs:
        if sub:
            merge_row(ws, r, sub, LABEL_FONT)
            r += 1
        rows = parse_table(lines)
        if rows:
            hdr = rows[0]
            for ci, h in enumerate(hdr, 1):
                c = ws.cell(row=r, column=ci, value=h)
                c.fill = HDR_FILL
                c.font = HDR_FONT
                c.alignment = CENTER
            for ci in range(len(hdr) + 1, 5):
                ws.cell(row=r, column=ci).border = BORDER
            r += 1
            for row in rows[1:]:
                if not any(row):
                    continue
                for ci, val in enumerate(row, 1):
                    c = ws.cell(row=r, column=ci, value=val)
                    c.alignment = WRAP
                for ci in range(len(row) + 1, 5):
                    ws.cell(row=r, column=ci).border = BORDER
                ws.row_dimensions[r].height = 32
                r += 1
            r += 1
    ws.freeze_panes = 'A2'


def build():
    md_path = os.path.join(DESIGN_DIR, MD_NAME)
    md = open(md_path, encoding='utf-8').read()
    sections = dict(split_sections(md))
    meta = meta_dict(sections)
    doc_title = meta.get('功能名称', '主线关卡扩展')
    bg, target = '', ''
    for sub, lines in split_subsections(sections.get('背景与目标', [])):
        t = subsection_text(lines)
        if sub == '背景':
            bg = t
        elif sub == '目标':
            target = t
    wb = Workbook()
    wb.remove(wb.active)
    write_overview_sheet(wb, doc_title, meta, bg, target)
    write_body_sheet(wb, sections)
    write_analytics_sheet(wb, sections)
    out = os.path.join(OUT_DIR, OUT_NAME)
    wb.save(out)
    print(f'-> {os.path.basename(out)}')


if __name__ == '__main__':
    build()
