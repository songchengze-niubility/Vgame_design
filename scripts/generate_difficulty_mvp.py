#!/usr/bin/env python3
"""Generate the read-only Vgame difficulty MVP workbook.

The generator reads formal level configs plus the current player/monster numerical
models. It never writes back to the Vgame SVN workspace.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


FONT_NAME = "Microsoft YaHei"
NAVY = "17365D"
BLUE = "1F4E78"
LIGHT_BLUE = "D9EAF7"
LIGHT_GREEN = "E2F0D9"
LIGHT_YELLOW = "FFF2CC"
LIGHT_RED = "FCE4D6"
GRAY = "E7E6E6"
WHITE = "FFFFFF"
INPUT_BLUE = "0000FF"
FORMULA_BLACK = "000000"
LINK_GREEN = "008000"
SOURCE_RED = "C00000"
THIN_GRAY = Side(style="thin", color="B7B7B7")


@dataclass
class Sources:
    config_root: Path
    battle_workbook: Path
    mainlevel: Path
    level: Path
    boss: Path


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--config-root", type=Path)
    parser.add_argument("--battle-workbook", type=Path)
    parser.add_argument("--output", type=Path)
    return parser.parse_args()


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def find_battle_workbook(explicit: Path | None) -> Path:
    candidates: list[Path] = []
    if explicit:
        candidates.append(explicit)
    root = os.environ.get("VGAME_SOURCE_DOCS_ROOT")
    if root:
        candidates.append(Path(root) / "Vgame基本战斗框架.xlsx")
    candidates.append(Path("D:/数值文档/数值文档/Vgame基本战斗框架.xlsx"))
    for candidate in candidates:
        if candidate.is_file():
            return candidate.resolve()
    raise FileNotFoundError("未找到 Vgame基本战斗框架.xlsx，请使用 --battle-workbook 指定")


def resolve_sources(args: argparse.Namespace) -> Sources:
    config_root = args.config_root or Path(
        os.environ.get("VGAME_CONFIG_DATAS", "D:/Vgame/Config/GameConfig/Datas")
    )
    config_root = config_root.resolve()
    sources = Sources(
        config_root=config_root,
        battle_workbook=find_battle_workbook(args.battle_workbook),
        mainlevel=config_root / "level" / "MainLevel.xlsx",
        level=config_root / "level" / "level.xlsx",
        boss=config_root / "BossDungeon" / "Boss_Dungeon.xlsx",
    )
    for path in (sources.mainlevel, sources.level, sources.boss, sources.battle_workbook):
        if not path.is_file():
            raise FileNotFoundError(path)
    return sources


def load_level_map(path: Path) -> dict[int, dict[str, Any]]:
    wb = load_workbook(path, data_only=True, keep_links=False)
    ws = wb["level"]
    result: dict[int, dict[str, Any]] = {}
    for row in range(5, ws.max_row + 1):
        level_id = ws.cell(row, 2).value
        if not level_id:
            continue
        result[int(level_id)] = {
            "active": ws.cell(row, 1).value != "##",
            "number": ws.cell(row, 7).value,
            "grade": ws.cell(row, 22).value,
            "hp": ws.cell(row, 23).value,
            "attack": ws.cell(row, 24).value,
        }
    return result


def load_mainline(path: Path, levels: dict[int, dict[str, Any]]) -> list[dict[str, Any]]:
    wb = load_workbook(path, data_only=True, keep_links=False)
    ws = wb["MainLevel"]
    result: list[dict[str, Any]] = []
    for row in range(5, ws.max_row + 1):
        level_id = ws.cell(row, 2).value
        if not level_id or ws.cell(row, 8).value != "LevelType_Main":
            continue
        current = levels.get(int(level_id), {})
        result.append(
            {
                "sequence": len(result) + 1,
                "active": ws.cell(row, 1).value != "##" and current.get("active", False),
                "level_id": int(level_id),
                "number": current.get("number", ""),
                "chapter_id": ws.cell(row, 3).value,
                "is_boss": bool(ws.cell(row, 7).value),
                "grade": current.get("grade"),
                "hp": current.get("hp") if current.get("hp") not in (None, "") else 1,
                "attack": current.get("attack") if current.get("attack") not in (None, "") else 1,
            }
        )
    return result


def load_boss(path: Path, levels: dict[int, dict[str, Any]]) -> list[dict[str, Any]]:
    wb = load_workbook(path, data_only=True, keep_links=False)
    ws = wb.active
    result: list[dict[str, Any]] = []
    for row in range(5, ws.max_row + 1):
        boss_id = ws.cell(row, 2).value
        if not boss_id or ws.cell(row, 1).value == "##":
            continue
        level_id = int(ws.cell(row, 7).value)
        current = levels.get(level_id, {})
        result.append(
            {
                "id": int(boss_id),
                "name": ws.cell(row, 4).value,
                "layer": int(ws.cell(row, 6).value),
                "level_id": level_id,
                "grade": current.get("grade"),
                "hp": current.get("hp") if current.get("hp") not in (None, "") else 1,
                "attack": current.get("attack") if current.get("attack") not in (None, "") else 1,
            }
        )
    return result


def load_models(path: Path) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    wb = load_workbook(path, data_only=True, read_only=True, keep_links=False)
    player_ws = wb["玩家战力曲线"]
    players: list[dict[str, Any]] = []
    for row in player_ws.iter_rows(
        min_row=12, max_row=116, min_col=1, max_col=37, values_only=True
    ):
        players.append(
            {
                "player_level": row[0],
                "sequence": int(row[1]),
                "chapter": row[2],
                "stage": row[3],
                "defense": row[14],
                "reduction": row[15],
                "attack": row[19],
                "dps": row[36],
            }
        )
    monster_ws = wb["怪物标准模型"]
    monsters: list[dict[str, Any]] = []
    for row in monster_ws.iter_rows(
        min_row=12, max_row=511, min_col=1, max_col=51, values_only=True
    ):
        monsters.append(
            {
                "level": int(row[1]),
                "boss_ref_hp_60": row[50],
                "boss_hp": row[37],
                "boss_attack": row[38],
                "boss_defense": row[39],
                "mob_hp": row[40],
                "mob_attack": row[41],
            }
        )
    return players, monsters


def validate_inputs(
    mainline: list[dict[str, Any]],
    bosses: list[dict[str, Any]],
    players: list[dict[str, Any]],
    monsters: list[dict[str, Any]],
) -> None:
    checks = {
        "主线计划节点": (len(mainline), 105),
        "Boss挑战档位": (len(bosses), 70),
        "玩家成长快照": (len(players), 105),
        "怪物等级模型": (len(monsters), 500),
    }
    failures = [f"{name}={actual}, expected={expected}" for name, (actual, expected) in checks.items() if actual != expected]
    if failures:
        raise ValueError("输入覆盖检查失败: " + "; ".join(failures))
    if any(row["dps"] in (None, 0) for row in players):
        raise ValueError("玩家成长快照存在空DPS")
    if any(row["boss_ref_hp_60"] in (None, 0) for row in monsters):
        raise ValueError("怪物模型存在空的Boss 60秒基准生命")


def style_title(ws, title: str, subtitle: str, end_col: int) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_col)
    ws["A1"] = title
    ws["A1"].font = Font(name=FONT_NAME, size=16, bold=True, color=WHITE)
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 28
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=end_col)
    ws["A2"] = subtitle
    ws["A2"].font = Font(name=FONT_NAME, size=10, color="666666")
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[2].height = 32


def style_table(ws, header_row: int, last_row: int, last_col: int) -> None:
    for cell in ws[header_row]:
        if cell.column > last_col:
            break
        cell.font = Font(name=FONT_NAME, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=THIN_GRAY)
    for row in ws.iter_rows(min_row=header_row + 1, max_row=last_row, min_col=1, max_col=last_col):
        for cell in row:
            cell.font = Font(name=FONT_NAME, size=9, color=FORMULA_BLACK)
            cell.alignment = Alignment(vertical="center", wrap_text=False)
            cell.border = Border(bottom=THIN_GRAY)
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(last_col)}{last_row}"
    ws.freeze_panes = f"A{header_row + 1}"


def set_widths(ws, widths: dict[int, float]) -> None:
    for column, width in widths.items():
        ws.column_dimensions[get_column_letter(column)].width = width


def mark_input(cell, note: str | None = None) -> None:
    cell.font = Font(name=FONT_NAME, color=INPUT_BLUE)
    cell.fill = PatternFill("solid", fgColor=LIGHT_YELLOW)
    if note:
        cell.comment = Comment(note, "Codex")


def mark_source(cell) -> None:
    cell.font = Font(name=FONT_NAME, color=SOURCE_RED)


def mark_formula(cell, linked: bool = False) -> None:
    cell.font = Font(name=FONT_NAME, color=LINK_GREEN if linked else FORMULA_BLACK)


def nearest_level(target: float, monsters: list[dict[str, Any]]) -> int:
    return min(monsters, key=lambda row: abs(float(row["boss_ref_hp_60"]) - target))["level"]


def build_workbook(
    sources: Sources,
    mainline: list[dict[str, Any]],
    bosses: list[dict[str, Any]],
    players: list[dict[str, Any]],
    monsters: list[dict[str, Any]],
) -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"

    guide = wb.create_sheet("使用说明")
    style_title(guide, "Vgame 怪物难度生成 MVP", "只读生成建议，不写回正式配置；当前覆盖主线105个计划节点与Boss挑战70档。", 8)
    guide_rows = [
        ("使用步骤", "1. 修改“参数”蓝色单元格；2. 查看主线/Boss建议；3. 有实机反馈时填写手动覆盖列；4. 审核“配置建议”后再人工落表。"),
        ("主线算法", "P50玩家DPS × 目标参考Boss TTK → 目标生命；在500级模型的AY列（60秒Boss基准生命）中搜索最近等级。"),
        ("Boss算法", "每层绑定一个主线玩家快照和目标TTK；默认值校准到当前正式Boss等级，保证MVP初始结果安全。"),
        ("不使用的数据", "生成公式不读取旧玩家战力曲线AO怪物等级列；它只读取玩家DPS/属性快照。"),
        ("当前边界", "没有线上通过率；攻击系数暂不自动反推，建议值默认保持当前值（空白按1）。生存模型后续接实机反馈再升级。"),
        ("覆盖说明", "MainLevel包含105个计划节点，其中6个当前被##注释；工作簿仍完整展示，并在检查列标记。"),
        ("回写边界", "本工具不会修改 D:/Vgame 下任何Excel；“配置建议”只是人工评审清单。"),
    ]
    for idx, (label, value) in enumerate(guide_rows, 4):
        guide.cell(idx, 1, label).font = Font(name=FONT_NAME, bold=True, color=NAVY)
        guide.cell(idx, 2, value).font = Font(name=FONT_NAME)
        guide.cell(idx, 2).alignment = Alignment(wrap_text=True, vertical="top")
        guide.merge_cells(start_row=idx, start_column=2, end_row=idx, end_column=8)
        guide.row_dimensions[idx].height = 38
    set_widths(guide, {1: 18, 2: 24, 3: 16, 4: 16, 5: 16, 6: 16, 7: 16, 8: 16})

    params = wb.create_sheet("参数")
    style_title(params, "可调参数", "蓝字黄底为人工输入；修改后Excel会自动重算。", 7)
    param_rows = [
        ("玩家P25 DPS倍率", 0.85, "无线上分布时的工作假设"),
        ("玩家P50 DPS倍率", 1.00, "生成主基准"),
        ("玩家P75 DPS倍率", 1.15, "验证养成较好玩家"),
        ("主线起始参考Boss TTK(秒)", 35.0, "当前战力框架起点"),
        ("主线结束参考Boss TTK(秒)", 55.0, "当前战力框架终点"),
        ("主线Boss关TTK倍率", 1.00, "大于1会形成Boss等级尖峰"),
        ("建议H下限", 0.85, "超出则优先调整怪物等级"),
        ("建议H上限", 1.15, "超出则优先调整怪物等级"),
        ("相邻主线最大等级跳变", 10, "仅做异常提醒"),
        ("主线计划关卡数", 105, "覆盖校验"),
    ]
    params.append(["参数", "值", "说明"])
    for label, value, note in param_rows:
        params.append([label, value, note])
        mark_input(params.cell(params.max_row, 2), note)
    style_table(params, 3, 3 + len(param_rows), 3)
    set_widths(params, {1: 34, 2: 16, 3: 48})
    params["B4"].number_format = "0.00x"
    params["B5"].number_format = "0.00x"
    params["B6"].number_format = "0.00x"
    params["B7"].number_format = "0.0"
    params["B8"].number_format = "0.0"
    params["B9"].number_format = "0.00x"
    params["B10"].number_format = "0.00"
    params["B11"].number_format = "0.00"

    current_grade_by_layer: dict[int, int] = {}
    for boss in bosses:
        current_grade_by_layer.setdefault(boss["layer"], int(boss["grade"]))
    anchors = [45, 50, 55, 60, 65, 75, 85, 95, 100, 105]
    ref_hp_by_level = {row["level"]: float(row["boss_ref_hp_60"]) for row in monsters}
    dps_by_sequence = {row["sequence"]: float(row["dps"]) for row in players}
    boss_start = params.max_row + 3
    params.cell(boss_start, 1, "Boss层")
    params.cell(boss_start, 2, "锚定主线序号")
    params.cell(boss_start, 3, "目标TTK(秒)")
    params.cell(boss_start, 4, "当前基准等级")
    params.cell(boss_start, 5, "初始化说明")
    for layer in range(1, 11):
        anchor = anchors[layer - 1]
        grade = current_grade_by_layer[layer]
        ttk = round(ref_hp_by_level[grade] / dps_by_sequence[anchor], 1)
        row = boss_start + layer
        params.cell(row, 1, layer)
        params.cell(row, 2, anchor)
        params.cell(row, 3, ttk)
        params.cell(row, 4, grade)
        params.cell(row, 5, "以当前Boss等级反算初始TTK；实机反馈后改B/C列")
        mark_input(params.cell(row, 2), "该Boss层使用哪个主线玩家成长快照")
        mark_input(params.cell(row, 3), "该Boss层希望P50玩家打多久")
        mark_source(params.cell(row, 4))
    style_table(params, boss_start, boss_start + 10, 5)
    set_widths(params, {4: 18, 5: 52})

    player_sheet = wb.create_sheet("玩家快照")
    style_title(player_sheet, "玩家成长快照", "P50来自当前玩家战力曲线；P25/P75为参数倍率。未读取旧怪物等级AO列。", 12)
    player_headers = ["序号", "玩家等级", "章节", "关卡", "P50 DPS", "P25 DPS", "P75 DPS", "最终攻击", "最终防御", "减伤率", "来源", "置信度"]
    player_sheet.append(player_headers)
    for idx, player in enumerate(players, 4):
        player_sheet.cell(idx, 1, player["sequence"])
        player_sheet.cell(idx, 2, player["player_level"])
        player_sheet.cell(idx, 3, player["chapter"])
        player_sheet.cell(idx, 4, player["stage"])
        player_sheet.cell(idx, 5, player["dps"])
        player_sheet.cell(idx, 6, f"=E{idx}*参数!$B$4")
        player_sheet.cell(idx, 7, f"=E{idx}*参数!$B$6")
        player_sheet.cell(idx, 8, player["attack"])
        player_sheet.cell(idx, 9, player["defense"])
        player_sheet.cell(idx, 10, player["reduction"])
        player_sheet.cell(idx, 11, "Vgame基本战斗框架.xlsx/玩家战力曲线")
        player_sheet.cell(idx, 12, "中：在研期模型快照")
        for col in (1, 2, 3, 4, 5, 8, 9, 10, 11, 12):
            mark_source(player_sheet.cell(idx, col))
        mark_formula(player_sheet.cell(idx, 6), linked=True)
        mark_formula(player_sheet.cell(idx, 7), linked=True)
        for col in (5, 6, 7, 8, 9):
            player_sheet.cell(idx, col).number_format = "#,##0"
        player_sheet.cell(idx, 10).number_format = "0.0%"
    style_table(player_sheet, 3, 3 + len(players), len(player_headers))
    set_widths(player_sheet, {1: 8, 2: 10, 3: 8, 4: 8, 5: 14, 6: 14, 7: 14, 8: 14, 9: 14, 10: 11, 11: 40, 12: 20})

    monster_sheet = wb.create_sheet("怪物模型")
    style_title(monster_sheet, "500级怪物标准模型", "等级搜索使用Boss 60秒基准生命（原模型AY列）；其他字段用于审阅。", 8)
    monster_headers = ["怪物等级", "Boss 60秒基准生命", "Boss实际生命", "Boss攻击", "Boss防御", "小怪生命", "小怪攻击", "来源"]
    monster_sheet.append(monster_headers)
    for idx, monster in enumerate(monsters, 4):
        values = [monster["level"], monster["boss_ref_hp_60"], monster["boss_hp"], monster["boss_attack"], monster["boss_defense"], monster["mob_hp"], monster["mob_attack"], "Vgame基本战斗框架.xlsx/怪物标准模型"]
        for col, value in enumerate(values, 1):
            monster_sheet.cell(idx, col, value)
            mark_source(monster_sheet.cell(idx, col))
        for col in range(2, 8):
            monster_sheet.cell(idx, col).number_format = "#,##0"
    style_table(monster_sheet, 3, 3 + len(monsters), len(monster_headers))
    set_widths(monster_sheet, {1: 12, 2: 20, 3: 16, 4: 14, 5: 14, 6: 16, 7: 14, 8: 42})

    main_sheet = wb.create_sheet("主线105")
    style_title(main_sheet, "主线105关难度建议", "自动等级由P50 DPS与目标参考Boss TTK反推；手动覆盖列用于实机反馈。", 28)
    main_headers = [
        "序号", "配置启用", "LevelId", "关卡号", "ChapterId", "Boss关", "玩家等级", "P25 DPS", "P50 DPS", "P75 DPS",
        "目标TTK", "目标Boss生命", "下界位置", "上界位置", "自动等级", "等级基准生命", "自动H", "当前等级", "当前H", "当前A",
        "手动等级", "手动H", "最终等级", "最终H", "建议A", "当前预测TTK", "建议预测TTK", "等级差", "检查", "调整原因",
    ]
    main_sheet.append(main_headers)
    monster_level_range = "'怪物模型'!$A$4:$A$503"
    monster_hp_range = "'怪物模型'!$B$4:$B$503"
    for idx, entry in enumerate(mainline, 4):
        seq = entry["sequence"]
        snapshot_row = 3 + seq
        values = [seq, entry["active"], entry["level_id"], entry["number"], entry["chapter_id"], entry["is_boss"], None, None, None, None, None, None, None, None, None, None, None, entry["grade"], entry["hp"], entry["attack"], None, None, None, None, None, None, None, None, None, None]
        for col, value in enumerate(values, 1):
            main_sheet.cell(idx, col, value)
        formulas = {
            7: f"='玩家快照'!B{snapshot_row}",
            8: f"='玩家快照'!F{snapshot_row}",
            9: f"='玩家快照'!E{snapshot_row}",
            10: f"='玩家快照'!G{snapshot_row}",
            11: f"=(参数!$B$7+(参数!$B$8-参数!$B$7)*(A{idx}-1)/(参数!$B$13-1))*IF(F{idx},参数!$B$9,1)",
            12: f"=I{idx}*K{idx}",
            13: f"=MATCH(L{idx},{monster_hp_range},1)",
            14: f"=MIN(M{idx}+1,ROWS({monster_hp_range}))",
            15: f"=IF(ABS(INDEX({monster_hp_range},M{idx})-L{idx})<=ABS(INDEX({monster_hp_range},N{idx})-L{idx}),INDEX({monster_level_range},M{idx}),INDEX({monster_level_range},N{idx}))",
            16: f"=INDEX({monster_hp_range},MATCH(O{idx},{monster_level_range},0))",
            17: f"=L{idx}/P{idx}",
            23: f"=IF(U{idx}=\"\",O{idx},U{idx})",
            24: f"=IF(V{idx}=\"\",Q{idx},V{idx})",
            25: f"=T{idx}",
            26: f"=INDEX({monster_hp_range},MATCH(R{idx},{monster_level_range},0))*S{idx}/I{idx}",
            27: f"=INDEX({monster_hp_range},MATCH(W{idx},{monster_level_range},0))*X{idx}/I{idx}",
            28: f"=W{idx}-R{idx}",
            29: f"=IF(NOT(B{idx}),\"配置已注释\",IF(OR(X{idx}<参数!$B$10,X{idx}>参数!$B$11),\"H超范围\",IF(A{idx}=1,\"OK\",IF(ABS(W{idx}-W{idx-1})>参数!$B$12,\"等级跳变\",\"OK\"))))",
        }
        for col, formula in formulas.items():
            main_sheet.cell(idx, col, formula)
            mark_formula(main_sheet.cell(idx, col), linked=col in (7, 8, 9, 10, 23, 24, 25))
        mark_input(main_sheet.cell(idx, 21), "实机反馈后填写；留空使用自动等级")
        mark_input(main_sheet.cell(idx, 22), "实机反馈后填写；留空使用自动H")
        mark_input(main_sheet.cell(idx, 30), "记录调整原因")
        for col in (1, 2, 3, 4, 5, 6, 18, 19, 20):
            mark_source(main_sheet.cell(idx, col))
        for col in (8, 9, 10, 12, 16):
            main_sheet.cell(idx, col).number_format = "#,##0"
        for col in (11, 17, 19, 20, 22, 24, 25, 26, 27):
            main_sheet.cell(idx, col).number_format = "0.00"
    style_table(main_sheet, 3, 3 + len(mainline), len(main_headers))
    main_sheet.column_dimensions["M"].hidden = True
    main_sheet.column_dimensions["N"].hidden = True
    main_sheet.conditional_formatting.add(f"AC4:AC{3+len(mainline)}", FormulaRule(formula=["$AC4<>\"OK\""], fill=PatternFill("solid", fgColor=LIGHT_RED)))
    main_sheet.conditional_formatting.add(f"AB4:AB{3+len(mainline)}", CellIsRule(operator="notEqual", formula=["0"], fill=PatternFill("solid", fgColor=LIGHT_YELLOW)))
    set_widths(main_sheet, {1: 8, 2: 10, 3: 12, 4: 10, 5: 12, 6: 9, 7: 10, 8: 13, 9: 13, 10: 13, 11: 11, 12: 16, 15: 11, 16: 17, 17: 10, 18: 11, 19: 10, 20: 10, 21: 12, 22: 10, 23: 11, 24: 10, 25: 10, 26: 14, 27: 14, 28: 10, 29: 14, 30: 28})

    boss_factor = wb.create_sheet("Boss修正")
    style_title(boss_factor, "Boss个体修正", "默认1.00；机制导致有效输出窗口变化时，用此系数修正目标生命。", 4)
    boss_factor.append(["Boss名称", "难度修正", "说明", "来源"])
    boss_names = list(dict.fromkeys(row["name"] for row in bosses))
    for idx, name in enumerate(boss_names, 4):
        boss_factor.cell(idx, 1, name)
        boss_factor.cell(idx, 2, 1.0)
        boss_factor.cell(idx, 3, "实机反馈后调整；例如机制压缩输出窗口可降低")
        boss_factor.cell(idx, 4, "人工输入")
        mark_source(boss_factor.cell(idx, 1))
        mark_input(boss_factor.cell(idx, 2), "Boss个体机制修正")
        mark_input(boss_factor.cell(idx, 3))
        boss_factor.cell(idx, 2).number_format = "0.00x"
    style_table(boss_factor, 3, 3 + len(boss_names), 4)
    set_widths(boss_factor, {1: 20, 2: 14, 3: 50, 4: 16})

    boss_sheet = wb.create_sheet("Boss70")
    style_title(boss_sheet, "Boss挑战70档难度建议", "默认目标校准到当前十档等级；修改参数中的锚定进度/目标TTK后自动重算。", 27)
    boss_headers = [
        "配置Id", "Boss名称", "层", "LevelId", "Boss修正", "锚定主线序号", "P50 DPS", "目标TTK", "目标Boss生命", "下界位置", "上界位置",
        "自动等级", "等级基准生命", "自动H", "当前等级", "当前H", "当前A", "手动等级", "手动H", "最终等级", "最终H", "建议A",
        "当前预测TTK", "建议预测TTK", "等级差", "检查", "调整原因",
    ]
    boss_sheet.append(boss_headers)
    boss_param_first = boss_start + 1
    boss_param_last = boss_start + 10
    for idx, entry in enumerate(bosses, 4):
        values = [entry["id"], entry["name"], entry["layer"], entry["level_id"], None, None, None, None, None, None, None, None, None, None, entry["grade"], entry["hp"], entry["attack"], None, None, None, None, None, None, None, None, None, None]
        for col, value in enumerate(values, 1):
            boss_sheet.cell(idx, col, value)
        formulas = {
            5: f"=INDEX('Boss修正'!$B$4:$B${3+len(boss_names)},MATCH(B{idx},'Boss修正'!$A$4:$A${3+len(boss_names)},0))",
            6: f"=INDEX(参数!$B${boss_param_first}:$B${boss_param_last},C{idx})",
            7: f"=INDEX('玩家快照'!$E$4:$E$108,F{idx})",
            8: f"=INDEX(参数!$C${boss_param_first}:$C${boss_param_last},C{idx})",
            9: f"=G{idx}*H{idx}*E{idx}",
            10: f"=MATCH(I{idx},{monster_hp_range},1)",
            11: f"=MIN(J{idx}+1,ROWS({monster_hp_range}))",
            12: f"=IF(ABS(INDEX({monster_hp_range},J{idx})-I{idx})<=ABS(INDEX({monster_hp_range},K{idx})-I{idx}),INDEX({monster_level_range},J{idx}),INDEX({monster_level_range},K{idx}))",
            13: f"=INDEX({monster_hp_range},MATCH(L{idx},{monster_level_range},0))",
            14: f"=I{idx}/M{idx}",
            20: f"=IF(R{idx}=\"\",L{idx},R{idx})",
            21: f"=IF(S{idx}=\"\",N{idx},S{idx})",
            22: f"=Q{idx}",
            23: f"=INDEX({monster_hp_range},MATCH(O{idx},{monster_level_range},0))*P{idx}/G{idx}",
            24: f"=INDEX({monster_hp_range},MATCH(T{idx},{monster_level_range},0))*U{idx}/G{idx}",
            25: f"=T{idx}-O{idx}",
            26: f"=IF(OR(U{idx}<参数!$B$10,U{idx}>参数!$B$11),\"H超范围\",\"OK\")",
        }
        for col, formula in formulas.items():
            boss_sheet.cell(idx, col, formula)
            mark_formula(boss_sheet.cell(idx, col), linked=col in (5, 6, 7, 8, 20, 21, 22))
        mark_input(boss_sheet.cell(idx, 18), "实机反馈后填写；留空使用自动等级")
        mark_input(boss_sheet.cell(idx, 19), "实机反馈后填写；留空使用自动H")
        mark_input(boss_sheet.cell(idx, 27), "记录调整原因")
        for col in (1, 2, 3, 4, 15, 16, 17):
            mark_source(boss_sheet.cell(idx, col))
        for col in (7, 9, 13):
            boss_sheet.cell(idx, col).number_format = "#,##0"
        for col in (5, 8, 14, 16, 17, 19, 21, 22, 23, 24):
            boss_sheet.cell(idx, col).number_format = "0.00"
    style_table(boss_sheet, 3, 3 + len(bosses), len(boss_headers))
    boss_sheet.column_dimensions["J"].hidden = True
    boss_sheet.column_dimensions["K"].hidden = True
    boss_sheet.conditional_formatting.add(f"Z4:Z{3+len(bosses)}", FormulaRule(formula=["$Z4<>\"OK\""], fill=PatternFill("solid", fgColor=LIGHT_RED)))
    set_widths(boss_sheet, {1: 10, 2: 18, 3: 8, 4: 12, 5: 11, 6: 15, 7: 14, 8: 11, 9: 16, 12: 11, 13: 17, 14: 10, 15: 11, 16: 10, 17: 10, 18: 12, 19: 10, 20: 11, 21: 10, 22: 10, 23: 14, 24: 14, 25: 10, 26: 14, 27: 28})

    config_sheet = wb.create_sheet("配置建议")
    style_title(config_sheet, "配置建议清单", "供人工评审与落表；本工具不会自动修改正式配置。", 12)
    config_headers = ["玩法", "LevelId", "关卡/层", "当前Grade", "建议Grade", "当前H", "建议H", "当前A", "建议A", "差值", "检查", "人工确认"]
    config_sheet.append(config_headers)
    row_idx = 4
    for source_name, source_sheet, count, label_col, level_col, current_grade_col, final_grade_col, current_h_col, final_h_col, current_a_col, final_a_col, delta_col, check_col in [
        ("主线", "主线105", len(mainline), "D", "C", "R", "W", "S", "X", "T", "Y", "AB", "AC"),
        ("Boss挑战", "Boss70", len(bosses), "C", "D", "O", "T", "P", "U", "Q", "V", "Y", "Z"),
    ]:
        for offset in range(count):
            source_row = 4 + offset
            config_sheet.cell(row_idx, 1, source_name)
            config_sheet.cell(row_idx, 2, f"='{source_sheet}'!{level_col}{source_row}")
            config_sheet.cell(row_idx, 3, f"='{source_sheet}'!{label_col}{source_row}")
            config_sheet.cell(row_idx, 4, f"='{source_sheet}'!{current_grade_col}{source_row}")
            config_sheet.cell(row_idx, 5, f"='{source_sheet}'!{final_grade_col}{source_row}")
            config_sheet.cell(row_idx, 6, f"='{source_sheet}'!{current_h_col}{source_row}")
            config_sheet.cell(row_idx, 7, f"='{source_sheet}'!{final_h_col}{source_row}")
            config_sheet.cell(row_idx, 8, f"='{source_sheet}'!{current_a_col}{source_row}")
            config_sheet.cell(row_idx, 9, f"='{source_sheet}'!{final_a_col}{source_row}")
            config_sheet.cell(row_idx, 10, f"='{source_sheet}'!{delta_col}{source_row}")
            config_sheet.cell(row_idx, 11, f"='{source_sheet}'!{check_col}{source_row}")
            config_sheet.cell(row_idx, 12, "")
            for col in range(2, 12):
                mark_formula(config_sheet.cell(row_idx, col), linked=True)
            mark_input(config_sheet.cell(row_idx, 12), "人工审核后填写：通过/调整/暂缓")
            for col in (6, 7, 8, 9):
                config_sheet.cell(row_idx, col).number_format = "0.00"
            row_idx += 1
    style_table(config_sheet, 3, row_idx - 1, len(config_headers))
    config_sheet.conditional_formatting.add(f"J4:J{row_idx-1}", CellIsRule(operator="notEqual", formula=["0"], fill=PatternFill("solid", fgColor=LIGHT_YELLOW)))
    set_widths(config_sheet, {1: 12, 2: 13, 3: 12, 4: 13, 5: 13, 6: 10, 7: 10, 8: 10, 9: 10, 10: 10, 11: 16, 12: 18})

    checks = wb.create_sheet("检查")
    style_title(checks, "覆盖与风险检查", "公式通过不等于体验通过；最终仍需实机验证和人工确认。", 8)
    checks.append(["检查项", "结果", "期望", "状态", "说明"])
    check_rows = [
        ("主线计划节点", "=COUNTA('主线105'!$A$4:$A$108)", 105, "含6个注释节点"),
        ("主线启用节点", "=COUNTIF('主线105'!$B$4:$B$108,TRUE)", 99, "当前正式配置状态"),
        ("主线注释节点", "=COUNTIF('主线105'!$B$4:$B$108,FALSE)", 6, "仍生成建议但不应直接启用"),
        ("Boss挑战档位", "=COUNTA('Boss70'!$A$4:$A$73)", 70, "7个Boss×10层"),
        ("玩家快照", "=COUNTA('玩家快照'!$A$4:$A$108)", 105, "仅使用属性/DPS，不用旧怪物等级"),
        ("怪物等级模型", "=COUNTA('怪物模型'!$A$4:$A$503)", 500, "等级1~500"),
        ("主线异常", "=COUNTIF('主线105'!$AC$4:$AC$108,\"<>OK\")", 6, "默认应只有6个配置注释提醒"),
        ("Boss异常", "=COUNTIF('Boss70'!$Z$4:$Z$73,\"<>OK\")", 0, "默认安全基线"),
        ("人工覆盖数", "=COUNT('主线105'!$U$4:$U$108)+COUNT('Boss70'!$R$4:$R$73)", 0, "收到实机反馈后允许增加"),
    ]
    for row_idx, (name, formula, expected, note) in enumerate(check_rows, 4):
        checks.cell(row_idx, 1, name)
        checks.cell(row_idx, 2, formula)
        checks.cell(row_idx, 3, expected)
        checks.cell(row_idx, 4, f"=IF(B{row_idx}=C{row_idx},\"OK\",\"CHECK\")")
        checks.cell(row_idx, 5, note)
        mark_formula(checks.cell(row_idx, 2), linked=True)
        mark_formula(checks.cell(row_idx, 4))
    style_table(checks, 3, 3 + len(check_rows), 5)
    checks.conditional_formatting.add(f"D4:D{3+len(check_rows)}", FormulaRule(formula=["$D4<>\"OK\""], fill=PatternFill("solid", fgColor=LIGHT_RED)))
    set_widths(checks, {1: 24, 2: 16, 3: 12, 4: 12, 5: 48})

    chart_sheet = wb.create_sheet("曲线图")
    style_title(chart_sheet, "难度曲线对比", "当前Grade与MVP最终建议Grade；手动覆盖会同步更新建议曲线。", 12)
    chart = LineChart()
    chart.title = "主线105关：当前与建议怪物等级"
    chart.y_axis.title = "怪物等级"
    chart.x_axis.title = "主线序号"
    data = Reference(main_sheet, min_col=18, max_col=23, min_row=3, max_row=108)
    cats = Reference(main_sheet, min_col=1, min_row=4, max_row=108)
    chart.add_data(data, titles_from_data=True, from_rows=False)
    chart.set_categories(cats)
    keep = [0, 5]
    chart.series = [chart.series[i] for i in keep]
    chart.height = 12
    chart.width = 25
    chart_sheet.add_chart(chart, "A4")
    chart_sheet["A30"] = "注：建议曲线根据玩家DPS与目标TTK生成；当前没有线上通过率，实机反馈优先级更高。"
    chart_sheet["A30"].font = Font(name=FONT_NAME, color="666666")

    source_sheet = wb.create_sheet("来源追踪")
    style_title(source_sheet, "来源追踪", "源文件只读；哈希用于确认重新生成时是否换了版本。", 5)
    source_sheet.append(["来源", "路径", "SHA256", "用途", "写入行为"])
    source_rows = [
        ("MainLevel", sources.mainlevel, "105个计划主线节点/启用状态", "只读"),
        ("level", sources.level, "当前Grade/H/A", "只读"),
        ("Boss_Dungeon", sources.boss, "70档Boss挑战LevelId", "只读"),
        ("基本战斗框架", sources.battle_workbook, "玩家DPS快照与500级怪物模型", "只读"),
    ]
    for idx, (name, path, purpose, behavior) in enumerate(source_rows, 4):
        source_sheet.cell(idx, 1, name)
        source_sheet.cell(idx, 2, str(path))
        source_sheet.cell(idx, 3, sha256(path))
        source_sheet.cell(idx, 4, purpose)
        source_sheet.cell(idx, 5, behavior)
        for col in range(1, 6):
            mark_source(source_sheet.cell(idx, col))
    style_table(source_sheet, 3, 3 + len(source_rows), 5)
    set_widths(source_sheet, {1: 20, 2: 75, 3: 68, 4: 38, 5: 12})

    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None and cell.font.name != FONT_NAME:
                    cell.font = Font(name=FONT_NAME, size=cell.font.sz or 10, bold=cell.font.bold, italic=cell.font.italic, color=cell.font.color)

    return wb


def main() -> int:
    args = parse_args()
    sources = resolve_sources(args)
    output = args.output or Path("output/全玩法智能难度生成工作流/全玩法智能难度MVP.xlsx")
    output = output.resolve()
    output.parent.mkdir(parents=True, exist_ok=True)

    levels = load_level_map(sources.level)
    mainline = load_mainline(sources.mainlevel, levels)
    bosses = load_boss(sources.boss, levels)
    players, monsters = load_models(sources.battle_workbook)
    validate_inputs(mainline, bosses, players, monsters)

    wb = build_workbook(sources, mainline, bosses, players, monsters)
    wb.save(output)

    report = {
        "status": "generated",
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "output": str(output),
        "read_only_sources": True,
        "counts": {
            "mainline_planned": len(mainline),
            "mainline_active": sum(1 for row in mainline if row["active"]),
            "mainline_commented": sum(1 for row in mainline if not row["active"]),
            "boss_challenge": len(bosses),
            "player_snapshots": len(players),
            "monster_levels": len(monsters),
        },
        "source_hashes": {
            "MainLevel.xlsx": sha256(sources.mainlevel),
            "level.xlsx": sha256(sources.level),
            "Boss_Dungeon.xlsx": sha256(sources.boss),
            "Vgame基本战斗框架.xlsx": sha256(sources.battle_workbook),
        },
        "guardrails": [
            "不读取玩家战力曲线AO怪物等级作为生成输入",
            "不写回Vgame正式配置",
            "攻击系数在无生存模型时保持当前值",
        ],
    }
    report_path = output.with_suffix(".build.json")
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
