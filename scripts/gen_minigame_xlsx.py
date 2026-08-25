# -*- coding: utf-8 -*-
"""把 design/小游戏-*.md 批量转成 openspec 风格策划案 xlsx。
3 个 sheet：修订记录与概述 / 系统功能正文 / 数据埋点方案。"""
import re
import os
import glob
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DESIGN_DIR = os.path.join(REPO_ROOT, 'design')
OUT_DIR = os.path.join(os.path.dirname(REPO_ROOT), 'Output')
os.makedirs(OUT_DIR, exist_ok=True)

TITLE_FONT = Font(bold=True, size=14, color='1F3864')
SECTION_FILL = PatternFill('solid', start_color='305496')
SECTION_FONT = Font(bold=True, size=12, color='FFFFFF')
SUB_FILL = PatternFill('solid', start_color='D9E1F2')
SUB_FONT = Font(bold=True, size=11, color='1F3864')
LABEL_FONT = Font(bold=True, size=11, color='1F3864')
HDR_FILL = PatternFill('solid', start_color='8EA9DB')
HDR_FONT = Font(bold=True, color='FFFFFF')
WRAP = Alignment(wrap_text=True, vertical='top')
TOP = Alignment(vertical='top', horizontal='left')
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
THIN = Side(style='thin', color='BFBFBF')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
NCOL = 7

GAME_EVENTS = {
    '齿轮啮合': [
        ('gear_place_wrong', '放入尺寸不符的齿轮', 'game_id, slot_idx, gear_size, error_type', '玩法ID/空位序号/齿轮尺寸/错误类型', '分析错放原因分布与空位难度'),
        ('gear_broken_place', '放入损坏齿轮', 'game_id, slot_idx, is_removed', '玩法ID/空位序号/是否取出', '分析干扰项辨识率'),
        ('gear_chain_rotate', '传动链接通开始转动', 'game_id, chain_length, time_used, is_correct_dir', '玩法ID/齿轮数/用时/转向是否正确', '分析通关效率与转向判断'),
    ],
    '元能切割': [
        ('fruit_slice', '切中水果', 'game_id, fruit_type, combo_count, slice_count', '玩法ID/水果类型/combo数/本刀切中数', '分析水果偏好与combo节奏'),
        ('danger_slice', '切到危险物', 'game_id, danger_type(腐蚀/薯片), combo_broken', '玩法ID/危险类型/是否断combo', '分析误切分布与惩罚影响'),
        ('slowmo_trigger', '触发慢动作', 'game_id, remaining_time, score_gain', '玩法ID/剩余时间/慢动作期间得分', '分析慢动作高分窗口利用'),
    ],
    '薯片接食物': [
        ('food_catch', '接住好食物', 'game_id, food_type, combo_count', '玩法ID/食物类型/combo数', '分析接物偏好与连击节奏'),
        ('bad_food_catch', '接到坏食物', 'game_id, food_type, combo_broken', '玩法ID/坏食物类型/是否断combo', '分析误接分布'),
        ('doge_sleep', '薯片困倦触发', 'game_id, missed_good, missed_bad', '玩法ID/困倦期漏好食物/漏坏食物', '分析困倦机制影响'),
    ],
    '沙尘躲避': [
        ('obstacle_hit', '被障碍物撞击', 'game_id, obstacle_type, shield_left, hit_x, hit_y', '玩法ID/障碍类型/剩余护盾/撞击坐标', '分析撞击热点与障碍难度'),
        ('energy_collect', '收集能量碎片', 'game_id, collect_count, score', '玩法ID/收集数/得分', '分析风险收益取舍'),
    ],
    '信号解码': [
        ('code_submit', '提交解码判定', 'game_id, attempt_no, is_correct, correct_count', '玩法ID/第几次尝试/是否正确/正确符号数', '分析推理难度与尝试分布'),
        ('hint_use', '使用提示', 'game_id, hint_symbol_idx, grade_capped', '玩法ID/提示符号位/是否降档', '分析卡关卡口与提示依赖'),
    ],
    '货物整理': [
        ('cargo_place', '放入货物', 'game_id, cargo_id, is_valid, cell_x, cell_y', '玩法ID/货物ID/是否合法/格子坐标', '分析放置策略与失败热点'),
        ('cargo_rotate', '旋转货物', 'game_id, cargo_id, rotate_count', '玩法ID/货物ID/旋转次数', '分析旋转操作频率'),
        ('constraint_violate', '约束违反', 'game_id, cargo_id, constraint_type(易碎/重物/温控)', '玩法ID/货物ID/约束类型', '分析约束理解难度'),
    ],
    '元能管线拼接': [
        ('pipe_rotate', '旋转管道', 'game_id, cell_pos, rotate_count', '玩法ID/格子坐标/旋转次数', '分析操作路径与效率'),
        ('pipe_corrode', '管道腐蚀发生', 'game_id, cell_pos, is_on_path', '玩法ID/格子坐标/是否在规划路径上', '分析腐蚀对规划的影响'),
        ('energy_flow_reach', '能量流到达断口', 'game_id, cell_pos, time_used', '玩法ID/断口坐标/用时', '分析流动压力节奏'),
    ],
    '星图连线': [
        ('star_connect', '连通星点', 'game_id, star_idx, path_length, is_correct', '玩法ID/星点序号/当前路径长度/是否在正确路径', '分析连线策略'),
        ('star_retry', '连线断开重试', 'game_id, retry_count, fail_reason(重复/断开)', '玩法ID/重试次数/失败原因', '分析试错频率'),
        ('meteor_blocked', '被流星阻挡', 'game_id, star_idx, meteor_phase', '玩法ID/星点序号/流星相位', '分析时序干扰影响'),
    ],
    '异星钓鱼': [
        ('fish_cast', '抛竿入水', 'game_id, water_id, cast_x, cast_y', '玩法ID/水域ID/落点坐标', '分析抛竿落点偏好'),
        ('fish_bite', '浮标咬钩事件', 'game_id, wait_time, fish_rarity_hint', '玩法ID/等待时长/鱼稀有度预警', '分析等待节奏与稀有鱼识别'),
        ('fish_hook_strike', '起钩判定', 'game_id, strike_timing(早/准/晚/超时), fish_rarity', '玩法ID/起钩时机/鱼稀有度', '分析起钩时机掌握与脱钩分布'),
        ('tension_result', '张力控制结果', 'game_id, fish_rarity, duration, is_broken, is_perfect, max_tension', '玩法ID/鱼稀有度/遛鱼时长/是否断线/是否完美/最高张力', '分析遛鱼难度曲线与断线点'),
        ('fish_catch', '成功钓起鱼', 'game_id, fish_id, fish_rarity, score_gain, is_perfect, water_id', '玩法ID/鱼ID/稀有度/得分/是否完美/水域ID', '分析鱼种产出分布与完美收杆率'),
        ('fish_encyclopedia_unlock', '图鉴解锁', 'game_id, fish_id, unlock_count_total', '玩法ID/鱼ID/已解锁图鉴总数', '分析图鉴收集进度与动力'),
    ],
    '房车烹饪': [
        ('qte_hit', 'QTE命中', 'game_id, action_type(切/炒/煮/颠), is_perfect, combo_count', '玩法ID/动作类型/是否完美/当前combo', '分析各操作类型命中率与节奏'),
        ('qte_miss', 'QTE未命中', 'game_id, action_type, miss_reason(早/晚/未输入), combo_broken', '玩法ID/动作类型/失误原因/是否断combo', '分析失误热点与谱面难度'),
        ('combo_peak', '连击峰值', 'game_id, peak_combo, multiplier_reached', '玩法ID/峰值combo/达到倍率', '分析连击维持能力与倍率分布'),
        ('dish_phase_change', '菜品阶段切换', 'game_id, dish_id, from_phase, to_phase, combo_kept', '玩法ID/菜品ID/源阶段/目标阶段/是否保住combo', '分析阶段衔接流畅度'),
        ('perfect_serve', '完美出锅', 'game_id, dish_id, bonus_score', '玩法ID/菜品ID/出锅奖励分', '分析出锅判定命中率与奖励获取'),
        ('dish_complete', '菜品完成', 'game_id, dish_id, total_score, hit_rate, max_combo', '玩法ID/菜品ID/总分/命中率/最大combo', '分析菜品谱面完成质量与难度体感'),
    ],
}

COMMON_EVENTS = [
    ('minigame_enter', '进入小游戏', 'user_id, timestamp, game_id, level_no, is_replay, chapter_id', '用户ID/时间/玩法ID/关卡等级/是否复访/章节ID', '分析进入率、复访率、章节分布'),
    ('minigame_skip', '跳过小游戏', 'user_id, timestamp, game_id, chapter_id', '用户ID/时间/玩法ID/章节ID', '分析跳过率、哪些玩法/章节跳过最多'),
    ('minigame_complete', '通关', 'user_id, timestamp, game_id, level_no, time_used, score', '用户ID/时间/玩法ID/等级/用时/得分', '分析通关时长、分数分布、难度曲线'),
    ('minigame_fail', '失败或超时', 'user_id, timestamp, game_id, level_no, fail_reason(timeout/life_zero/give_up)', '用户ID/时间/玩法ID/等级/失败原因', '分析失败原因、卡点等级'),
    ('minigame_grade', '评价结算', 'user_id, timestamp, game_id, level_no, grade(S/A/B/C), is_best, reward_group_id', '用户ID/时间/玩法ID/等级/评价/是否破纪录/奖励组ID', '分析评价分布、奖励发放、破纪录率'),
]


def clean(t):
    return t.replace('**', '').replace('`', '')


def split_sections(md):
    sections, cur_title, cur_body = [], None, []
    for line in md.splitlines():
        if line.startswith('## ') and not line.startswith('### '):
            if cur_title is not None:
                sections.append((cur_title, cur_body))
            cur_title, cur_body = line[3:].strip(), []
        elif cur_title is not None:
            cur_body.append(line)
    if cur_title is not None:
        sections.append((cur_title, cur_body))
    return sections


def split_subsections(body):
    subs, cur, cur_lines = [], None, []
    for line in body:
        if line.startswith('### '):
            if cur is not None:
                subs.append((cur, cur_lines))
            cur, cur_lines = line[4:].strip(), []
        elif cur is not None:
            cur_lines.append(line)
        elif line.strip():
            if not subs or subs[-1][0] != '':
                subs.append(('', []))
            subs[-1][1].append(line)
    if cur is not None:
        subs.append((cur, cur_lines))
    return subs


def parse_table(lines):
    rows = []
    for ln in lines:
        s = ln.strip()
        if not s.startswith('|') or re.match(r'^\|[\s:|-]+\|$', s):
            continue
        rows.append([clean(c.strip()) for c in s.strip('|').split('|')])
    return rows


def parse_bullets(lines):
    items = []
    for ln in lines:
        if not ln.strip():
            continue
        m = re.match(r'^(\s*)- \*\*(.+?)\*\*[：:](.*)$', ln)
        if m:
            items.append([m.group(2).strip(), m.group(3).strip(), len(m.group(1))])
            continue
        m2 = re.match(r'^(\s*)-\s+(.*)$', ln)
        if m2:
            indent, text = len(m2.group(1)), m2.group(2).strip()
            if items and indent > items[-1][2]:
                items[-1][1] = (items[-1][1] + '\n' + text) if items[-1][1] else text
            else:
                items.append(['', text, indent])
        elif items and ln.strip():
            items[-1][1] = (items[-1][1] + '\n' + ln.strip()) if items[-1][1] else ln.strip()
    return [[f, clean(v)] for f, v, _ in items]


def subsection_text(lines):
    out = []
    for ln in lines:
        s = ln.strip()
        if not s:
            continue
        s = re.sub(r'^-\s+', '', s)
        out.append(s)
    return '\n'.join(out)


def meta_dict(sections):
    d = {}
    for r in parse_table(sections.get('元信息', [])):
        if len(r) >= 2:
            d[r[0]] = r[1]
    return d


def style_row(ws, row, ncol=NCOL):
    for ci in range(1, ncol + 1):
        ws.cell(row=row, column=ci).border = BORDER


def merge_row(ws, row, text, font, fill=None, align=None, ncol=NCOL):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncol)
    c = ws.cell(row=row, column=1, value=text)
    c.font = font
    if fill:
        c.fill = fill
    c.alignment = align or Alignment(horizontal='left', vertical='center', wrap_text=True)
    style_row(ws, row, ncol)


def write_overview_sheet(wb, game_name, meta, bg, target):
    ws = wb.create_sheet('修订记录与概述')
    ws.column_dimensions['A'].width = 16
    for col in 'BCDEFG':
        ws.column_dimensions[col].width = 18
    r = 1
    merge_row(ws, r, f'{game_name} 策划案', TITLE_FONT, align=Alignment(horizontal='center', vertical='center'))
    ws.row_dimensions[r].height = 28
    r += 2
    merge_row(ws, r, '修订记录', LABEL_FONT)
    r += 1
    hdr = ['版本号', '修订人', '日期', '简介', '修订内容索引']
    for ci, h in enumerate(hdr, 2):
        c = ws.cell(row=r, column=ci, value=h)
        c.fill = HDR_FILL
        c.font = HDR_FONT
        c.alignment = CENTER
    style_row(ws, r)
    r += 1
    rev = ['v1', meta.get('作者', '策划'), meta.get('创建日期', '2026-07-01'), '初版策划案', '--']
    for ci, v in enumerate(rev, 2):
        c = ws.cell(row=r, column=ci, value=v)
        c.alignment = WRAP
    style_row(ws, r)
    r += 2
    merge_row(ws, r, '文档属性', LABEL_FONT)
    r += 1
    for ci, h in enumerate(['项', '内容'], 2):
        c = ws.cell(row=r, column=ci, value=h)
        c.fill = HDR_FILL
        c.font = HDR_FONT
        c.alignment = CENTER
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=7)
    style_row(ws, r)
    r += 1
    for k, v in meta.items():
        ws.cell(row=r, column=2, value=k).alignment = TOP
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=7)
        ws.cell(row=r, column=3, value=v).alignment = WRAP
        style_row(ws, r)
        r += 1
    r += 1
    merge_row(ws, r, '系统概述', LABEL_FONT)
    r += 1
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=7)
    ws.cell(row=r, column=2, value='背景：' + bg).alignment = WRAP
    style_row(ws, r)
    r += 1
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=7)
    ws.cell(row=r, column=2, value='目标：' + target).alignment = WRAP
    style_row(ws, r)
    ws.freeze_panes = 'A2'


SECTION_MAP = [
    ('背景与目标', '1. 系统概述', [
        ('背景', '1.1', '背景'),
        ('目标', '1.2', '目标'),
        ('设计意图与体验目标', '1.3', '设计意图与体验目标'),
        ('非目标', '1.4', '非目标'),
    ]),
    ('核心规则', '2. 核心规则', None),
    ('关卡递进设计', '3. 关卡递进设计', None),
    ('数值框架', '4. 数值框架', None),
    ('美术风格参考', '5. 美术风格参考', None),
    ('玩法参考', '6. 玩法参考', None),
    ('风险评估', '7. 风险评估', None),
    ('工作假设', '8. 工作假设', None),
]


def write_body_sheet(wb, sections):
    ws = wb.create_sheet('系统功能正文')
    ws.column_dimensions['A'].width = 8
    ws.column_dimensions['B'].width = 22
    for col in 'CDEFG':
        ws.column_dimensions[col].width = 20
    r = 1
    merge_row(ws, r, '系统功能正文', TITLE_FONT, align=Alignment(horizontal='center', vertical='center'))
    ws.row_dimensions[r].height = 26
    r += 2

    for sec_key, sec_title, subs_def in SECTION_MAP:
        if sec_key not in sections:
            continue
        body = sections[sec_key]
        r += 0
        merge_row(ws, r, sec_title, SECTION_FONT, SECTION_FILL, align=Alignment(horizontal='left', vertical='center'))
        r += 2

        if sec_key == '背景与目标':
            smap = {s: (n, t) for s, n, t in subs_def}
            for sub, lines in split_subsections(body):
                if sub in smap:
                    num, title = smap[sub]
                    ws.cell(row=r, column=1, value=num).alignment = TOP
                    ws.cell(row=r, column=1).font = SUB_FONT
                    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=NCOL)
                    c = ws.cell(row=r, column=2, value=title)
                    c.font = SUB_FONT
                    c.fill = SUB_FILL
                    c.alignment = Alignment(vertical='center')
                    style_row(ws, r)
                    r += 1
                    txt = subsection_text(lines)
                    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=NCOL)
                    ws.cell(row=r, column=2, value=txt).alignment = WRAP
                    style_row(ws, r)
                    ws.row_dimensions[r].height = max(30, min(120, txt.count('\n') * 18 + 30))
                    r += 2
        elif sec_key == '核心规则':
            for sub, lines in split_subsections(body):
                ws.cell(row=r, column=1, value='').alignment = TOP
                ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=NCOL)
                c = ws.cell(row=r, column=2, value=sub)
                c.font = SUB_FONT
                c.fill = SUB_FILL
                c.alignment = Alignment(vertical='center')
                style_row(ws, r)
                r += 1
                for f, v in parse_bullets(lines):
                    ws.cell(row=r, column=2, value=f).alignment = TOP
                    ws.cell(row=r, column=2).font = Font(bold=True)
                    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=NCOL)
                    ws.cell(row=r, column=3, value=v).alignment = WRAP
                    style_row(ws, r)
                    ws.row_dimensions[r].height = max(20, min(100, v.count('\n') * 16 + 20))
                    r += 1
                r += 1
        else:
            rows = parse_table(body)
            if rows:
                hdr = rows[0]
                for ci, h in enumerate(hdr, 2):
                    c = ws.cell(row=r, column=ci, value=h)
                    c.fill = HDR_FILL
                    c.font = HDR_FONT
                    c.alignment = CENTER
                style_row(ws, r)
                r += 1
                for row in rows[1:]:
                    if not any(row):
                        continue
                    for ci, val in enumerate(row, 2):
                        c = ws.cell(row=r, column=ci, value=val)
                        c.alignment = WRAP
                    style_row(ws, r)
                    ws.row_dimensions[r].height = max(20, min(80, max((str(v).count('\n') + 1) * 16 for v in row if v) + 4))
                    r += 1
                r += 1
            else:
                bullets = parse_bullets(body)
                if bullets:
                    hdr = ['要素', '说明']
                    for ci, h in enumerate(hdr, 2):
                        c = ws.cell(row=r, column=ci, value=h)
                        c.fill = HDR_FILL
                        c.font = HDR_FONT
                        c.alignment = CENTER
                    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=NCOL)
                    style_row(ws, r)
                    r += 1
                    for f, v in bullets:
                        ws.cell(row=r, column=2, value=f).alignment = TOP
                        ws.cell(row=r, column=2).font = Font(bold=True)
                        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=NCOL)
                        ws.cell(row=r, column=3, value=v).alignment = WRAP
                        style_row(ws, r)
                        ws.row_dimensions[r].height = max(20, min(120, v.count('\n') * 16 + 22))
                        r += 1
                    r += 1
    ws.freeze_panes = 'A2'


def write_analytics_sheet(wb, game_key):
    ws = wb.create_sheet('数据埋点方案')
    ws.column_dimensions['A'].width = 22
    ws.column_dimensions['B'].width = 30
    ws.column_dimensions['C'].width = 40
    ws.column_dimensions['D'].width = 40
    ws.column_dimensions['E'].width = 28
    r = 1
    merge_row(ws, r, '数据埋点方案', TITLE_FONT, align=Alignment(horizontal='center', vertical='center'))
    ws.row_dimensions[r].height = 26
    r += 2
    merge_row(ws, r, '一、通用埋点（所有小游戏共享）', LABEL_FONT)
    r += 1
    hdr = ['埋点事件名', '触发时机', '上报字段', '字段说明', '分析用途']
    for ci, h in enumerate(hdr, 1):
        c = ws.cell(row=r, column=ci, value=h)
        c.fill = HDR_FILL
        c.font = HDR_FONT
        c.alignment = CENTER
    style_row(ws, r)
    r += 1
    for ev in COMMON_EVENTS:
        for ci, v in enumerate(ev, 1):
            ws.cell(row=r, column=ci, value=v).alignment = WRAP
        style_row(ws, r)
        ws.row_dimensions[r].height = 32
        r += 1
    r += 1
    merge_row(ws, r, '二、本玩法特有埋点', LABEL_FONT)
    r += 1
    for ci, h in enumerate(hdr, 1):
        c = ws.cell(row=r, column=ci, value=h)
        c.fill = HDR_FILL
        c.font = HDR_FONT
        c.alignment = CENTER
    style_row(ws, r)
    r += 1
    for ev in GAME_EVENTS.get(game_key, []):
        for ci, v in enumerate(ev, 1):
            ws.cell(row=r, column=ci, value=v).alignment = WRAP
        style_row(ws, r)
        ws.row_dimensions[r].height = 32
        r += 1
    ws.freeze_panes = 'A2'


def build(md_path):
    md = open(md_path, encoding='utf-8').read()
    sections = dict(split_sections(md))
    meta = meta_dict(sections)
    game_name = meta.get('功能名称', os.path.basename(md_path))
    bg, target = '', ''
    for sub, lines in split_subsections(sections.get('背景与目标', [])):
        t = subsection_text(lines)
        if sub == '背景':
            bg = t
        elif sub == '目标':
            target = t
    wb = Workbook()
    wb.remove(wb.active)
    write_overview_sheet(wb, game_name, meta, bg, target)
    write_body_sheet(wb, sections)
    game_key = next((k for k in GAME_EVENTS if k in game_name), '')
    write_analytics_sheet(wb, game_key)
    base = os.path.splitext(os.path.basename(md_path))[0].replace('Design Doc', '').strip()
    out = os.path.join(OUT_DIR, f'{base}.xlsx')
    wb.save(out)
    return out


def main():
    files = sorted(glob.glob(os.path.join(DESIGN_DIR, '小游戏-*Design Doc.md')))
    print(f'found {len(files)} minigame design docs')
    for f in files:
        out = build(f)
        print(f'  -> {os.path.basename(out)}')


if __name__ == '__main__':
    main()
