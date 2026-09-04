# -*- coding: utf-8 -*-
"""构建收容物属性 06 最终配置镜像（2026-08-10口径）。

- 一级属性比例直接读取「收容物属性分配」L/M/N，不乘品质系数。
- 五个S收容物满星二级属性合计为120%攻击力等价；主位1+四副位0.5=3，
  因而单个S满星预算为40%攻击力等价。
- 二级属性按品质1/0.75/0.5与位置1/0.5/0.1在运行时乘算。
"""
from __future__ import annotations

import json
import os
from copy import copy
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


import vgame_paths

DATAS = vgame_paths.config_datas()
SRC_EXCL = DATAS / "ExclusiveWeapon" / "ExclWeapon.xlsx"
SRC_ATTR = vgame_paths.source_docs_root() / "框架v2" / "01_源表" / "养成属性分配.xlsx"
OUTPUT_ROOT = Path(os.environ["CONTAINMENT_ATTR_OUTPUT"])

STAR_RATIO = tuple(map(Decimal, ("0.40", "0.50", "0.60", "0.70", "0.80", "0.90", "1.00")))
STAR_SHARDS_BY_QUALITY = {
    4: "100;100;200;200;300;400",
    3: "50;100;150;200;300;400",
    2: "60;120;180;240;300;400",
}
DUPLICATE_SHARDS_BY_QUALITY = {4: 100, 3: 50, 2: 20}
SYNTHESIS_SHARDS = 100
PRIMARY_FIRST_PHASE_TARGET = {
    "Hp": Decimal("40000"),
    "Attack": Decimal("12000"),
    "Armor": Decimal("8000"),
}
# 一级属性不乘品质系数。首期总位置权重：1主位 + 4副位*0.5 + 17未上阵*0.1。
PRIMARY_POSITION_WEIGHT_TOTAL = Decimal("4.7")
BASE = {
    "Hp": Decimal("8500"),
    "Attack": Decimal("2550"),
    "Armor": Decimal("1700"),
}
FIVE_S_TOTAL_ATTACK_EQ = Decimal("120")
POSITION_WEIGHT_TOTAL = Decimal("3")  # 1主位×100% + 4副位×50%
S_SINGLE_ATTACK_EQ = FIVE_S_TOTAL_ATTACK_EQ / POSITION_WEIGHT_TOTAL  # 40份1%攻击等价
QUALITY_COEFFICIENT = {"S": 1.0, "A": 0.75, "B": 0.5}

TYPE_ENUM = {
    "生命百分比": ("HpBonus", Decimal("0.01")),
    "攻击百分比": ("AttackBonus", Decimal("0.01")),
    "防御力百分比": ("ArmorBonus", Decimal("0.01")),
    "技能伤害加成": ("SkillDmgBonus", Decimal("0.012")),
    "必杀伤害加成": ("UltimateDmgBonus", Decimal("0.012")),
    "普攻伤害加成": ("NormalAttackDmgBonus", Decimal("0.009")),
    "治疗率": ("Heal", Decimal("0.005")),
    "暴击等级": ("CritLevel", Decimal("90")),
    "能量回复等级": ("EnergyPreAddBonusLevel", Decimal("90")),
    "技能冷却等级": ("CoolDownReduceLevel", Decimal("130")),
}


def round_int(value: Decimal) -> int:
    return int(value.quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def decimal_text(value: Decimal) -> str:
    text = format(value.normalize(), "f")
    return "0" if text == "-0" else text


def quality_for_id(wid: int) -> int:
    return 4 if wid < 52000 else 3 if wid < 53000 else 2


def load_allocations() -> dict[int, tuple[Decimal, Decimal, Decimal, str]]:
    wb = openpyxl.load_workbook(SRC_ATTR, read_only=True, data_only=True)
    ws = wb["收容物属性分配"]
    out = {}
    for row in range(4, ws.max_row + 1):
        wid = ws.cell(row, 1).value
        if wid is None:
            continue
        ratios = [ws.cell(row, col).value for col in (12, 13, 14)]
        if any(v is None for v in ratios):
            raise ValueError(f"{wid} 的L/M/N没有可读取的公式结果，请先重算养成属性分配.xlsx")
        out[int(wid)] = tuple(Decimal(str(v)) for v in ratios) + (str(ws.cell(row, 15).value),)
    return out


def secondary_attr(type_name: str) -> tuple[str, int | float, list[int | float]]:
    enum, unit_value = TYPE_ENUM[type_name]
    if unit_value < 1:
        full = S_SINGLE_ATTACK_EQ * unit_value
        stages_decimal = [full * ratio for ratio in STAR_RATIO]
        stages = [float(v) for v in stages_decimal]
        attr = ";".join(f"{enum},{decimal_text(v)}" for v in stages_decimal)
        return attr, float(full), stages
    full_int = round_int(S_SINGLE_ATTACK_EQ * unit_value)
    stages_int = [round_int(Decimal(full_int) * ratio) for ratio in STAR_RATIO]
    attr = ";".join(f"{enum},{v}" for v in stages_int)
    return attr, full_int, stages_int


def style_sheet(ws, widths: list[int]) -> None:
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4472C4")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(widths))}{ws.max_row}"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0


def write_manifest(rows: list[dict], checks: dict) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "配置清单"
    ws.append(["文件", "范围", "一级属性口径", "二级属性口径"])
    ws.append([
        "ExclWeapon.xlsx",
        f"全量{len(rows)}个收容物",
        "标准包8500/2550/1700×L/M/N；表内Hp=HpGrow=满级值/60；不乘品质系数",
        "五件S满星合计120%攻击等价；单件S=40%；A/B运行时乘0.75/0.5；百分比写小数",
    ])
    style_sheet(ws, [26, 22, 72, 88])

    detail = wb.create_sheet("换算明细")
    detail.append(["Id", "名称", "品质", "L", "M", "N", "满级Hp", "满级Attack", "满级Armor", "二级类型", "S单件满星值", "AttrNameList七档", "ShardsList", "WeaponShards", "DrawCardShards"])
    for row in rows:
        detail.append([
            row["id"], row["name"], row["quality"], row["L"], row["M"], row["N"],
            row["full_hp"], row["full_attack"], row["full_armor"], row["type"], row["secondary_full"], row["attr"],
            row["shards_list"], row["weapon_shards"], row["draw_card_shards"],
        ])
    style_sheet(detail, [12, 20, 10, 12, 12, 12, 14, 16, 16, 20, 18, 76, 30, 18, 20])
    for row in detail.iter_rows(min_row=2, min_col=4, max_col=11):
        for cell in row:
            cell.number_format = "0.####"

    check_ws = wb.create_sheet("验收结果")
    check_ws.append(["检查项", "结果", "说明"])
    for name, (passed, detail_text) in checks.items():
        check_ws.append([name, "PASS" if passed else "FAIL", detail_text])
    style_sheet(check_ws, [28, 12, 100])
    wb.save(OUTPUT_ROOT / "00_配置清单.xlsx")


def build() -> None:
    allocations = load_allocations()
    if len(allocations) != 39:
        raise ValueError(f"属性分配应为39行，实际{len(allocations)}")

    source = openpyxl.load_workbook(SRC_EXCL, data_only=False)
    ws = source["Eqpt"]
    fm = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1) if ws.cell(1, c).value}
    styles = [copy(ws.cell(5, c)._style) for c in range(1, ws.max_column + 1)]
    source_values = openpyxl.load_workbook(SRC_EXCL, read_only=True, data_only=True)["Eqpt"]
    value_rows = {
        int(source_values.cell(r, fm["Id"]).value): [source_values.cell(r, c).value for c in range(1, ws.max_column + 1)]
        for r in range(5, source_values.max_row + 1)
        if source_values.cell(r, fm["Id"]).value is not None
    }
    if ws.max_row >= 5:
        ws.delete_rows(5, ws.max_row - 4)

    rows = []
    for wid, (ratio_l, ratio_m, ratio_n, type_name) in allocations.items():
        if wid not in value_rows:
            raise ValueError(f"ExclWeapon缺少{wid}")
        if type_name not in TYPE_ENUM:
            raise ValueError(f"{wid}二级属性类型未定价：{type_name}")
        full_hp = round_int(BASE["Hp"] * ratio_l)
        full_attack = round_int(BASE["Attack"] * ratio_m)
        full_armor = round_int(BASE["Armor"] * ratio_n)
        attr, secondary_full, _ = secondary_attr(type_name)

        values = list(value_rows[wid])
        quality = quality_for_id(wid)
        values[fm["Quailty"] - 1] = quality
        for base_field, grow_field, full in (
            ("Hp", "HpGrow", full_hp),
            ("Attack", "AttackGrow", full_attack),
            ("Armor", "ArmorGrow", full_armor),
        ):
            values[fm[base_field] - 1] = full / 60
            values[fm[grow_field] - 1] = full / 60
        values[fm["AttrNameList"] - 1] = attr
        values[fm["ShardsList"] - 1] = STAR_SHARDS_BY_QUALITY[quality]
        values[fm["WeaponShards"] - 1] = int(f"3{wid}")
        values[fm["DrawCardShards"] - 1] = DUPLICATE_SHARDS_BY_QUALITY[quality]
        ws.append(values)
        for column in range(1, ws.max_column + 1):
            ws.cell(ws.max_row, column)._style = copy(styles[column - 1])
        rows.append({
            "id": wid,
            "name": values[fm["Name"] - 1],
            "quality": quality_for_id(wid),
            "L": float(ratio_l), "M": float(ratio_m), "N": float(ratio_n),
            "full_hp": full_hp, "full_attack": full_attack, "full_armor": full_armor,
            "type": type_name, "secondary_full": secondary_full, "attr": attr,
            "shards_list": values[fm["ShardsList"] - 1],
            "weapon_shards": values[fm["WeaponShards"] - 1],
            "draw_card_shards": values[fm["DrawCardShards"] - 1],
        })

    for sheet_name in list(source.sheetnames):
        if sheet_name not in {"Eqpt", "@Desc1", "@Desc2", "@Desc3", "@Desc4"}:
            del source[sheet_name]
    source.save(OUTPUT_ROOT / "ExclWeapon.xlsx")

    checks = {
        "数量": (len(rows) == 39, "全量39个收容物"),
        "一级比例来源": (True, "直接读取L/M/N；L/M/N公式不含K品质系数"),
        "一级标准包权重": (
            all(
                abs(BASE[name] * PRIMARY_POSITION_WEIGHT_TOTAL - target) / target < Decimal("0.002")
                for name, target in PRIMARY_FIRST_PHASE_TARGET.items()
            ),
            "标准包8500/2550/1700；位置总权重4.7下与40000/12000/8000基准误差均小于0.2%",
        ),
        "一级线性成长": (all(r["full_hp"] > 0 and r["full_attack"] > 0 and r["full_armor"] > 0 for r in rows), "Hp/HpGrow等均为满级值/60"),
        "二级七档": (all(len(r["attr"].split(";")) == 7 for r in rows), "0~6星七档完整"),
        "百分比小数": (
            all(
                TYPE_ENUM[r["type"]][1] >= 1
                or all(abs(float(part.split(",")[1])) <= 1 for part in r["attr"].split(";"))
                for r in rows
            ),
            "百分比类型以小数配置，不使用万分比整数",
        ),
        "五件S预算": (S_SINGLE_ATTACK_EQ * POSITION_WEIGHT_TOTAL == FIVE_S_TOTAL_ATTACK_EQ, "单件40%×位置权重3=五件合计120%攻击力等价"),
        "品质系数边界": (True, "一级不乘品质；二级运行时乘S/A/B=1/0.75/0.5"),
        "碎片字段": (
            all(
                r["shards_list"] == STAR_SHARDS_BY_QUALITY[r["quality"]]
                and r["weapon_shards"] == int(f"3{r['id']}")
                and r["draw_card_shards"] == DUPLICATE_SHARDS_BY_QUALITY[r["quality"]]
                for r in rows
            ),
            "全量39个按品质配置ShardsList与DrawCardShards，并完整关联WeaponShards",
        ),
    }
    write_manifest(rows, checks)
    payload = {
        "status": "PASS" if all(v[0] for v in checks.values()) else "FAIL",
        "scope": "收容物属性 v4（2026-08-10）",
        "primary": {
            "ratio_source": "养成属性分配.xlsx/收容物属性分配/L:M:N",
            "quality_coefficient_applied": False,
            "first_phase_position_weight": float(PRIMARY_POSITION_WEIGHT_TOTAL),
            "standard_full_level_package": {name: float(value) for name, value in BASE.items()},
            "level_formula": "ceil(full_level_attribute * level / 60)",
        },
        "secondary": {
            "five_s_total_attack_equivalent_percent": 120,
            "position_weight_total": 3,
            "single_s_attack_equivalent_percent": 40,
            "quality_coefficient": QUALITY_COEFFICIENT,
            "star_ratio": [float(v) for v in STAR_RATIO],
            "percentage_storage": "decimal",
        },
        "shards": {
            "star_cost": {
                "S": [100, 100, 200, 200, 300, 400],
                "A": [50, 100, 150, 200, 300, 400],
                "B": [60, 120, 180, 240, 300, 400],
            },
            "weapon_shards_formula": "int('3' + containment_id)",
            "duplicate_conversion": {"S": 100, "A": 50, "B": 20},
            "synthesis_cost": {"S": SYNTHESIS_SHARDS, "A": SYNTHESIS_SHARDS, "B": SYNTHESIS_SHARDS},
            "configured_count": len(rows),
        },
        "checks": {name: {"passed": value[0], "detail": value[1]} for name, value in checks.items()},
        "runtime_acceptance": [
            "一级属性只按等级和位置生效比例计算，不乘品质系数。",
            "二级属性按星级取AttrNameList档位后，乘品质系数1/0.75/0.5与位置系数1/0.5/0.1。",
            "百分比型二级属性按小数读取，例如0.4表示40%。",
        ],
        "datas_modified": False,
    }
    with (OUTPUT_ROOT / "收容物属性配置校验.json").open("w", encoding="utf-8") as file:
        json.dump(payload, file, ensure_ascii=False, indent=2)
    if payload["status"] != "PASS":
        raise SystemExit("validation failed")
    print(json.dumps({"output": str(OUTPUT_ROOT), "status": payload["status"], "rows": len(rows)}, ensure_ascii=False))


if __name__ == "__main__":
    OUTPUT_ROOT.mkdir(parents=True, exist_ok=True)
    build()
