#!/usr/bin/env python3
"""Generate a traceable 105-stage P50 core-growth player snapshot.

This generator intentionally does not depend on the deprecated
Vgame基本战斗框架.xlsx. It reads formal Vgame configuration in read-only mode,
combines it with the pacing bindings and an explicit scenario policy, then emits
JSON and XLSX artifacts for difficulty calibration.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import os
import re
from datetime import datetime
from pathlib import Path
from statistics import median
from typing import Any, Iterable

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


import vgame_paths

DEFAULT_CONFIG_ROOT = vgame_paths.config_datas()
DEFAULT_SOURCE_DOCS_ROOT = vgame_paths.source_docs_root()
DEFAULT_PACING_MD = DEFAULT_SOURCE_DOCS_ROOT / "框架v2" / "03_沉淀MD" / "参数基线" / "节奏参数基线.md"
DEFAULT_WORKFLOW_DIR = Path(__file__).resolve().parents[1] / "output" / "全玩法智能难度生成工作流"
DEFAULT_BINDINGS_MD = DEFAULT_WORKFLOW_DIR / "03-specialist-results" / "standard-progression-nodes.md"
DEFAULT_POLICY = DEFAULT_WORKFLOW_DIR / "03-specialist-results" / "player-snapshot-policy.v1.json"
DEFAULT_OUTPUT_DIR = DEFAULT_WORKFLOW_DIR


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--config-root", type=Path, default=DEFAULT_CONFIG_ROOT)
    parser.add_argument("--pacing-md", type=Path, default=DEFAULT_PACING_MD)
    parser.add_argument("--bindings-md", type=Path, default=DEFAULT_BINDINGS_MD)
    parser.add_argument("--policy", type=Path, default=DEFAULT_POLICY)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    return parser.parse_args()


def read_json_fence(path: Path) -> dict[str, Any]:
    text = path.read_text(encoding="utf-8")
    match = re.search(r"```json\s*(\{.*?\})\s*```", text, re.S)
    if not match:
        raise ValueError(f"No JSON fence found: {path}")
    return json.loads(match.group(1))


def load_table(path: Path) -> tuple[list[str | None], list[tuple[Any, ...]]]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if len(rows) < 5:
        raise ValueError(f"Unexpected table layout: {path}")
    return list(rows[0]), rows[4:]


def row_dict(headers: list[str | None], row: tuple[Any, ...]) -> dict[str, Any]:
    return {
        str(headers[i]): row[i]
        for i in range(1, min(len(headers), len(row)))
        if headers[i] not in (None, "")
    }


def index_table(path: Path, key: str) -> dict[int, dict[str, Any]]:
    headers, rows = load_table(path)
    result: dict[int, dict[str, Any]] = {}
    for row in rows:
        data = row_dict(headers, row)
        value = data.get(key)
        if value is None:
            continue
        data["__row_marker__"] = row[0]
        result[int(value)] = data
    return result


def parse_int_list(value: Any) -> list[int]:
    if value in (None, ""):
        return []
    if isinstance(value, (int, float)):
        return [int(value)]
    return [int(part.strip()) for part in str(value).replace(",", ";").split(";") if part.strip()]


def round_half_up(value: float) -> int:
    return int(math.floor(value + 0.5))


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def parse_pacing_days(pacing_md: Path) -> dict[int, int]:
    payload = read_json_fence(pacing_md)
    sheet = next(item for item in payload["sheets"] if item["name"] == "主角等级_输入")
    rows = sheet["rows"]
    header_index = next(i for i, row in enumerate(rows) if row and row[0] == "等级")
    headers = rows[header_index]
    level_index = headers.index("等级")
    day_index = headers.index("对应天数")
    return {
        int(row[level_index]): int(row[day_index])
        for row in rows[header_index + 1 :]
        if row and row[level_index] is not None and row[day_index] is not None
    }


def build_stage_bindings(bindings: dict[str, Any], pacing_days: dict[int, int]) -> list[dict[str, Any]]:
    formation_unlocks = bindings["formation_unlocks"]
    result: list[dict[str, Any]] = []
    for chapter in bindings["chapter_bindings"]:
        seq_start = int(chapter["seq_start"])
        seq_end = int(chapter["seq_end"])
        level_start = int(chapter["level_start"])
        level_end = int(chapter["level_end"])
        count = seq_end - seq_start + 1
        for index, seq in enumerate(range(seq_start, seq_end + 1), start=1):
            fraction = 0 if count <= 1 else (index - 1) / (count - 1)
            pacing_level = round_half_up(level_start + (level_end - level_start) * fraction)
            chapter_no = int(chapter["chapter"])
            level_id = 100000 + chapter_no * 100 + index
            slot_cap = sum(1 for item in formation_unlocks if int(item["effective_seq"]) <= seq)
            result.append(
                {
                    "seq": seq,
                    "level_id": level_id,
                    "stage": f"{chapter_no}-{index}",
                    "chapter": chapter_no,
                    "chapter_stage": index,
                    "pacing_level": pacing_level,
                    "day": int(pacing_days[pacing_level]),
                    "slot_cap": slot_cap,
                }
            )
    expected = int(bindings["compile"]["current_stage_count"])
    if len(result) != expected:
        raise ValueError(f"Stage binding count {len(result)} != expected {expected}")
    return result


def load_globals(path: Path) -> dict[str, Any]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    result: dict[str, Any] = {}
    for row in ws.iter_rows(min_row=4, values_only=True):
        if row[0] is not None:
            result[str(row[0])] = row[3]
    wb.close()
    return result


def load_drop_pools(path: Path, item_to_hero: dict[int, int]) -> dict[int, list[int]]:
    headers, rows = load_table(path)
    pools: dict[int, list[int]] = {}
    for row in rows:
        data = row_dict(headers, row)
        if data.get("DropId") is None or data.get("DropItem_id") is None:
            continue
        drop_id = int(data["DropId"])
        item_id = int(data["DropItem_id"])
        hero_id = item_to_hero.get(item_id)
        if hero_id is not None:
            pools.setdefault(drop_id, []).append(hero_id)
    return pools


def ascend_for_level(level: int, ascend_rows: Iterable[dict[str, Any]]) -> dict[str, Any]:
    ordered = sorted(ascend_rows, key=lambda item: int(item["AscendId"]))
    for item in ordered:
        if level <= int(item["LvLimit"]):
            return item
    return ordered[-1]


def numeric(data: dict[str, Any], key: str, default: float = 0.0) -> float:
    value = data.get(key)
    return default if value in (None, "") else float(value)


def calculate_exact_stats(
    hero_id: int,
    level: int,
    hero_attributes: dict[int, dict[str, Any]],
    ascend_rows: list[dict[str, Any]],
    armor_reduce_x: float,
) -> dict[str, Any]:
    current = hero_attributes.get(hero_id * 1000 + level)
    base = hero_attributes.get(hero_id * 1000 + 1)
    if current is None or base is None:
        raise KeyError(f"Missing HeroAttribute: hero={hero_id}, level={level}")
    ascend = ascend_for_level(level, ascend_rows)
    attr_up = numeric(ascend, "AttrUp", 1.0)
    increment = attr_up - 1.0

    hp = round_half_up(numeric(current, "Hp") + numeric(base, "Hp") * increment)
    attack = round_half_up(numeric(current, "Attack") + numeric(base, "Attack") * increment)
    armor = round_half_up(numeric(current, "Armor") + numeric(base, "Armor") * increment)
    crit = numeric(current, "Crit")
    crit_dmg = numeric(current, "CritDmgBonus", 2.0)
    dmg_bonus = numeric(current, "DmgBonus")
    dmg_immune = min(numeric(current, "DmgImmune"), 0.95)
    attack_index = attack * (1 + crit * max(0.0, crit_dmg - 1.0)) * (1 + dmg_bonus)
    ehp_proxy = hp * (1 + armor / armor_reduce_x) / max(0.05, 1 - dmg_immune)
    return {
        "level": level,
        "ascend_id": int(ascend["AscendId"]),
        "attr_up": attr_up,
        "hp": hp,
        "attack": attack,
        "armor": armor,
        "crit": crit,
        "crit_dmg_bonus": crit_dmg,
        "dmg_bonus": dmg_bonus,
        "dmg_immune": dmg_immune,
        "attack_index": round(attack_index, 2),
        "ehp_proxy": round(ehp_proxy, 2),
    }


def calculate_pool_median_stats(
    hero_ids: list[int],
    level: int,
    hero_attributes: dict[int, dict[str, Any]],
    ascend_rows: list[dict[str, Any]],
    armor_reduce_x: float,
) -> dict[str, Any]:
    values = [
        calculate_exact_stats(hero_id, level, hero_attributes, ascend_rows, armor_reduce_x)
        for hero_id in hero_ids
    ]
    result: dict[str, Any] = {
        "level": level,
        "ascend_id": int(median(item["ascend_id"] for item in values)),
        "attr_up": float(median(item["attr_up"] for item in values)),
    }
    for key in ("hp", "attack", "armor", "crit", "crit_dmg_bonus", "dmg_bonus", "dmg_immune"):
        result[key] = median(item[key] for item in values)
    result["hp"] = round_half_up(float(result["hp"]))
    result["attack"] = round_half_up(float(result["attack"]))
    result["armor"] = round_half_up(float(result["armor"]))
    result["attack_index"] = round(
        result["attack"]
        * (1 + result["crit"] * max(0.0, result["crit_dmg_bonus"] - 1.0))
        * (1 + result["dmg_bonus"]),
        2,
    )
    result["ehp_proxy"] = round(
        result["hp"]
        * (1 + result["armor"] / armor_reduce_x)
        / max(0.05, 1 - result["dmg_immune"]),
        2,
    )
    return result


def build_snapshots(
    stages: list[dict[str, Any]],
    policy: dict[str, Any],
    heroes: dict[int, dict[str, Any]],
    hero_attributes: dict[int, dict[str, Any]],
    ascend_rows: list[dict[str, Any]],
    drop_pools: dict[int, list[int]],
    formal_levels: dict[int, dict[str, Any]],
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    armor_reduce_x = float(policy["metrics"]["armor_reduce_x"])
    snapshots: list[dict[str, Any]] = []
    details: list[dict[str, Any]] = []
    for stage in stages:
        fielded: list[dict[str, Any]] = []
        for member in sorted(policy["members"], key=lambda item: int(item["position"])):
            if int(member["position"]) > stage["slot_cap"] or int(member["acquisition_seq"]) > stage["seq"]:
                continue
            hero_level = max(1, stage["pacing_level"] - int(member["level_offset"]))
            kind = member["kind"]
            if kind == "exact_hero":
                hero_id = int(member["hero_id"])
                stats = calculate_exact_stats(hero_id, hero_level, hero_attributes, ascend_rows, armor_reduce_x)
                identity = heroes[hero_id]
                display_name = str(identity["Name"])
                hero_class = str(identity["HeroClass"])
                quality = int(identity["Quality"])
                pool_members = [hero_id]
                evidence = "formal_exact"
            elif kind == "drop_pool_median":
                pool_id = int(member["drop_pool_id"])
                pool_members = sorted(set(drop_pools.get(pool_id, [])))
                if not pool_members:
                    raise KeyError(f"No hero members in Drop_draw pool {pool_id}")
                stats = calculate_pool_median_stats(
                    pool_members, hero_level, hero_attributes, ascend_rows, armor_reduce_x
                )
                display_name = f"掉落包{pool_id}属性中位模板"
                hero_class = "Mixed"
                quality = int(median(int(heroes[hero_id]["Quality"]) for hero_id in pool_members))
                evidence = "p50_pool_median"
            else:
                raise ValueError(f"Unsupported member kind: {kind}")

            detail = {
                "seq": stage["seq"],
                "level_id": stage["level_id"],
                "stage": stage["stage"],
                "position": int(member["position"]),
                "member_id": member["member_id"],
                "display_name": display_name,
                "kind": kind,
                "hero_id": member.get("hero_id"),
                "pool_hero_ids": pool_members,
                "quality": quality,
                "hero_class": hero_class,
                "acquisition_seq": int(member["acquisition_seq"]),
                "level_offset": int(member["level_offset"]),
                "availability_mode": member.get("availability_mode", "owned"),
                "source": member["source"],
                "evidence": evidence,
                **stats,
            }
            details.append(detail)
            fielded.append(detail)

        formal = formal_levels.get(stage["level_id"])
        formal_state = "missing"
        if formal is not None:
            formal_state = "commented" if formal.get("__row_marker__") == "##" else "active"
        snapshots.append(
            {
                **stage,
                "formal_level_state": formal_state,
                "actual_team_size": len(fielded),
                "slot_gap": stage["slot_cap"] - len(fielded),
                "team_members": [item["member_id"] for item in fielded],
                "team_member_names": [item["display_name"] for item in fielded],
                "hero_levels": [item["level"] for item in fielded],
                "team_hp": sum(item["hp"] for item in fielded),
                "team_attack": sum(item["attack"] for item in fielded),
                "team_armor": sum(item["armor"] for item in fielded),
                "team_attack_index": round(sum(item["attack_index"] for item in fielded), 2),
                "team_ehp_proxy": round(sum(item["ehp_proxy"] for item in fielded), 2),
                "team_dps": None,
                "dps_status": "calibration_pending",
                "resource_feasibility": policy["assumptions"]["resource_feasibility"],
                "snapshot_confidence": "medium" if all(item["evidence"] == "formal_exact" for item in fielded) else "low",
                "publish_ready": False,
            }
        )
    return snapshots, details


NAVY = "17365D"
BLUE = "2F75B5"
LIGHT_BLUE = "D9EAF7"
LIGHT_ORANGE = "FCE4D6"
LIGHT_GREEN = "E2F0D9"
LIGHT_GRAY = "E7E6E6"
WHITE = "FFFFFF"
THIN_GRAY = Side(style="thin", color="D9E1F2")


def style_sheet(ws, freeze: str | None = None, auto_filter: bool = False) -> None:
    ws.sheet_view.showGridLines = False
    if freeze:
        ws.freeze_panes = freeze
    if auto_filter and ws.max_row >= 2:
        ws.auto_filter.ref = ws.dimensions
    for row in ws.iter_rows():
        for cell in row:
            cell.font = Font(name="Microsoft YaHei", size=10)
            cell.alignment = Alignment(vertical="center", wrap_text=False)
            cell.border = Border(bottom=THIN_GRAY)


def apply_header(ws, row: int = 1) -> None:
    for cell in ws[row]:
        if cell.value is not None:
            cell.fill = PatternFill("solid", fgColor=BLUE)
            cell.font = Font(name="Microsoft YaHei", size=10, bold=True, color=WHITE)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 32


def autosize(ws, max_width: int = 42) -> None:
    for column in range(1, ws.max_column + 1):
        width = 0
        for row in range(1, min(ws.max_row, 300) + 1):
            value = ws.cell(row, column).value
            if value is not None:
                width = max(width, len(str(value)))
        ws.column_dimensions[get_column_letter(column)].width = min(max(width + 2, 10), max_width)


def write_xlsx(
    output_path: Path,
    metadata: dict[str, Any],
    snapshots: list[dict[str, Any]],
    details: list[dict[str, Any]],
    policy: dict[str, Any],
    bindings: dict[str, Any],
    sources: list[dict[str, Any]],
    checks: list[dict[str, Any]],
) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "说明"
    ws.merge_cells("A1:F1")
    ws["A1"] = "105关 P50 玩家核心养成快照 MVP"
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A1"].font = Font(name="Microsoft YaHei", size=16, bold=True, color=WHITE)
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 34
    intro = [
        ("模型ID", metadata["model_id"]),
        ("状态", metadata["status"]),
        ("生成时间", metadata["generated_at"]),
        ("有效范围", "主线序号1～105，节奏Lv1～39 / Day1～8"),
        ("已计入", "阵位解锁、可确认角色来源、角色等级目标、基础HP/攻击/防御、突破属性增量"),
        ("未计入", "星级、装备、徽章、能效、技能倍率/循环、阵容协同"),
        ("攻击指数", policy["metrics"]["attack_index"]),
        ("EHP代理", policy["metrics"]["ehp_proxy"]),
        ("DPS", "留空；标准战斗校准前不伪造DPS"),
        ("用法", "先看105关快照的实际队伍人数、攻击指数和EHP跳点；试玩后只调策略参数，不改生成器逻辑。"),
    ]
    for row_index, (key, value) in enumerate(intro, start=3):
        ws.cell(row_index, 1, key)
        ws.cell(row_index, 2, value)
        ws.merge_cells(start_row=row_index, start_column=2, end_row=row_index, end_column=6)
        ws.cell(row_index, 1).fill = PatternFill("solid", fgColor=LIGHT_BLUE)
        ws.cell(row_index, 1).font = Font(name="Microsoft YaHei", bold=True)
        ws.cell(row_index, 2).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[row_index].height = 34 if row_index >= 7 else 24
    style_sheet(ws)
    ws.column_dimensions["A"].width = 18
    for col in "BCDEF":
        ws.column_dimensions[col].width = 18

    ws = wb.create_sheet("105关快照")
    headers = [
        "序号", "LevelId", "关卡", "章节", "正式关卡行", "节奏等级", "天数", "阵位上限", "实际队伍人数", "缺编",
        "队伍成员", "角色等级", "总HP", "总攻击", "总防御", "攻击指数", "EHP代理", "DPS", "DPS状态",
        "资源可达性", "快照信心", "可发布性"
    ]
    ws.append(headers)
    for item in snapshots:
        ws.append([
            item["seq"], item["level_id"], item["stage"], item["chapter"], item["formal_level_state"],
            item["pacing_level"], item["day"], item["slot_cap"], item["actual_team_size"], item["slot_gap"],
            " / ".join(item["team_member_names"]), "/".join(map(str, item["hero_levels"])), item["team_hp"],
            item["team_attack"], item["team_armor"], item["team_attack_index"], item["team_ehp_proxy"],
            item["team_dps"], item["dps_status"], item["resource_feasibility"], item["snapshot_confidence"],
            "是" if item["publish_ready"] else "否",
        ])
    style_sheet(ws, freeze="A2", auto_filter=True)
    apply_header(ws)
    for row in range(2, ws.max_row + 1):
        if ws.cell(row, 10).value:
            ws.cell(row, 9).fill = PatternFill("solid", fgColor=LIGHT_ORANGE)
            ws.cell(row, 10).fill = PatternFill("solid", fgColor=LIGHT_ORANGE)
        if ws.cell(row, 1).value in {1, 15, 30, 45, 60, 75, 90, 105}:
            for cell in ws[row]:
                cell.fill = PatternFill("solid", fgColor=LIGHT_GREEN)
    autosize(ws, 34)

    ws = wb.create_sheet("角色明细")
    headers = [
        "序号", "LevelId", "关卡", "位置", "成员ID", "显示名", "类型", "角色ID", "池角色ID",
        "品质", "职责", "获得序号", "可用方式", "等级差", "角色等级", "突破阶段", "突破系数", "HP", "攻击", "防御",
        "暴击", "暴伤", "攻击指数", "EHP代理", "证据级别", "来源"
    ]
    ws.append(headers)
    for item in details:
        ws.append([
            item["seq"], item["level_id"], item["stage"], item["position"], item["member_id"], item["display_name"],
            item["kind"], item["hero_id"], ";".join(map(str, item["pool_hero_ids"])), item["quality"], item["hero_class"],
            item["acquisition_seq"], item["availability_mode"], item["level_offset"], item["level"], item["ascend_id"], item["attr_up"], item["hp"],
            item["attack"], item["armor"], item["crit"], item["crit_dmg_bonus"], item["attack_index"], item["ehp_proxy"],
            item["evidence"], item["source"],
        ])
    style_sheet(ws, freeze="A2", auto_filter=True)
    apply_header(ws)
    autosize(ws, 38)

    ws = wb.create_sheet("策略参数")
    ws.append(["位置", "成员ID", "类型", "角色/掉落池", "获得序号", "可用方式", "等级差", "来源"])
    for item in sorted(policy["members"], key=lambda value: int(value["position"])):
        target = item.get("hero_id", f"DropPool {item.get('drop_pool_id')}")
        ws.append([item["position"], item["member_id"], item["kind"], target, item["acquisition_seq"], item.get("availability_mode", "owned"), item["level_offset"], item["source"]])
    start = ws.max_row + 3
    ws.cell(start, 1, "其他假设")
    ws.cell(start, 1).fill = PatternFill("solid", fgColor=NAVY)
    ws.cell(start, 1).font = Font(name="Microsoft YaHei", bold=True, color=WHITE)
    for index, (key, value) in enumerate(policy["assumptions"].items(), start=start + 1):
        ws.cell(index, 1, key)
        ws.cell(index, 2, json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else value)
    style_sheet(ws, freeze="A2")
    apply_header(ws)
    autosize(ws, 54)

    ws = wb.create_sheet("阵位解锁")
    ws.append(["阵位", "SystemOpenId", "UnlockId", "条件类型", "条件参数", "首个生效序号"])
    for item in bindings["formation_unlocks"]:
        ws.append([item["slot"], item["system_open_id"], item["unlock_id"], item["unlock_type"], item["unlock_param"], item["effective_seq"]])
    style_sheet(ws, freeze="A2", auto_filter=True)
    apply_header(ws)
    autosize(ws)

    ws = wb.create_sheet("来源追踪")
    ws.append(["来源", "路径", "SHA256"])
    for item in sources:
        ws.append([item["name"], item["path"], item["sha256"]])
    style_sheet(ws, freeze="A2", auto_filter=True)
    apply_header(ws)
    autosize(ws, 70)

    ws = wb.create_sheet("校验")
    ws.append(["检查项", "状态", "结果"])
    for item in checks:
        ws.append([item["check"], item["status"], item["result"]])
    style_sheet(ws, freeze="A2", auto_filter=True)
    apply_header(ws)
    for row in range(2, ws.max_row + 1):
        fill = LIGHT_GREEN if ws.cell(row, 2).value == "PASS" else LIGHT_ORANGE
        ws.cell(row, 2).fill = PatternFill("solid", fgColor=fill)
    autosize(ws, 70)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def main() -> int:
    args = parse_args()
    config_root = args.config_root.resolve()
    required_paths = {
        "生成器": Path(__file__).resolve(),
        "节奏参数基线": args.pacing_md.resolve(),
        "主线绑定": args.bindings_md.resolve(),
        "快照策略": args.policy.resolve(),
        "globals_conf": config_root / "globals_conf" / "globals_conf.xlsx",
        "Hero": config_root / "Hero" / "Hero.xlsx",
        "HeroAttribute": config_root / "Hero" / "HeroAttribute.xlsx",
        "HeroAscendLv": config_root / "Hero" / "HeroAscendLv.xlsx",
        "Items": config_root / "items" / "Items.xlsx",
        "SystemOpen": config_root / "system_open" / "SystemOpen.xlsx",
        "Unlock": config_root / "unlock" / "Unlock.xlsx",
        "DrawCard": config_root / "Activity" / "DrawCard" / "DrawCard.xlsx",
        "Drop_draw": config_root / "Drop" / "Drop_draw.xlsx",
        "Drop_mainlevel": config_root / "Drop" / "Drop_mainlevel.xlsx",
        "Level": config_root / "level" / "level.xlsx",
    }
    missing = [str(path) for path in required_paths.values() if not path.exists()]
    if missing:
        raise FileNotFoundError("Missing inputs: \n" + "\n".join(missing))

    pacing_days = parse_pacing_days(required_paths["节奏参数基线"])
    bindings = read_json_fence(required_paths["主线绑定"])
    policy = json.loads(required_paths["快照策略"].read_text(encoding="utf-8"))
    stages = build_stage_bindings(bindings, pacing_days)

    globals_conf = load_globals(required_paths["globals_conf"])
    heroes = index_table(required_paths["Hero"], "HeroId")
    hero_attributes = index_table(required_paths["HeroAttribute"], "AttributeID")
    ascend_rows = list(index_table(required_paths["HeroAscendLv"], "Id").values())
    items = index_table(required_paths["Items"], "Id")
    system_open = index_table(required_paths["SystemOpen"], "Id")
    unlock = index_table(required_paths["Unlock"], "Id")
    draw_card = index_table(required_paths["DrawCard"], "Id")
    formal_levels = index_table(required_paths["Level"], "LevelId")
    item_to_hero = {
        int(hero["HeroItem"]): hero_id
        for hero_id, hero in heroes.items()
        if hero.get("HeroItem") is not None
    }
    drop_pools = load_drop_pools(required_paths["Drop_draw"], item_to_hero)
    mainline_drop_pools = load_drop_pools(required_paths["Drop_mainlevel"], item_to_hero)

    if parse_int_list(globals_conf.get("InitialHero")) != [1001]:
        raise ValueError(f"Unexpected InitialHero: {globals_conf.get('InitialHero')}")
    if 311002 not in items or int(item_to_hero.get(311002, 0)) != 1002:
        raise ValueError("Forced guide item 311002 does not map to hero 1002")
    if 311040 not in items or int(item_to_hero.get(311040, 0)) != 1040:
        raise ValueError("Mainline item 311040 does not map to hero 1040 (小寒)")
    if 1040 not in mainline_drop_pools.get(10003, []):
        raise ValueError("Drop_mainlevel.DropId=10003 no longer contains hero item 311040 (小寒)")
    if int(system_open[17]["UnlockId"]) != 4017 or int(unlock[4017]["UnLockParams"]) != 100101:
        raise ValueError("Recruitment unlock chain changed: SystemOpen.17 -> Unlock.4017 -> Level.100101 expected")
    if int(draw_card[2]["ForcedGuideReplacement"]) != 311002:
        raise ValueError("DrawCard.2.ForcedGuideReplacement changed from 311002")
    formal_armor_reduce_x = float(globals_conf.get("CombatArmorReduceX"))
    if formal_armor_reduce_x != float(policy["metrics"]["armor_reduce_x"]):
        raise ValueError(
            f"Armor reduce parameter mismatch: globals={formal_armor_reduce_x}, policy={policy['metrics']['armor_reduce_x']}"
        )

    level_id_to_seq = {item["level_id"]: item["seq"] for item in stages}
    for item in bindings["formation_unlocks"]:
        slot = int(item["slot"])
        sys_row = system_open[int(item["system_open_id"])]
        if int(sys_row["UnlockId"]) != int(item["unlock_id"]):
            raise ValueError(f"Formation slot {slot} SystemOpen -> Unlock mismatch")
        unlock_row = unlock[int(item["unlock_id"])]
        if str(unlock_row["UnLockType"]) != str(item["unlock_type"]):
            raise ValueError(f"Formation slot {slot} unlock type mismatch")
        if int(unlock_row["UnLockParams"]) != int(item["unlock_param"]):
            raise ValueError(f"Formation slot {slot} unlock param mismatch")
        if str(item["unlock_type"]) == "UnLock_UUT_Level":
            expected_effective_seq = level_id_to_seq[int(item["unlock_param"])] + 1
            if expected_effective_seq != int(item["effective_seq"]):
                raise ValueError(f"Formation slot {slot} effective sequence mismatch")

    ten_pull_seq = int(policy["assumptions"]["newbie_ten_pull_completion_seq"])
    modeled_pool_members = [item for item in policy["members"] if item["kind"] == "drop_pool_median"]
    if any(int(item["acquisition_seq"]) != ten_pull_seq for item in modeled_pool_members):
        raise ValueError("Pool-median acquisition_seq must match newbie_ten_pull_completion_seq in MVP v0.1")

    snapshots, details = build_snapshots(
        stages, policy, heroes, hero_attributes, ascend_rows, drop_pools, formal_levels
    )
    key_team_sizes = {1: 1, 3: 2, 7: 3, 10: 4, 20: 5}
    actual_key_team_sizes = {
        item["seq"]: item["actual_team_size"] for item in snapshots if item["seq"] in key_team_sizes
    }
    stage_0_3 = snapshots[2]
    checks = [
        {"check": "快照关卡数", "status": "PASS" if len(snapshots) == 105 else "FAIL", "result": len(snapshots)},
        {"check": "序号唯一", "status": "PASS" if len({item['seq'] for item in snapshots}) == 105 else "FAIL", "result": len({item['seq'] for item in snapshots})},
        {"check": "LevelId唯一", "status": "PASS" if len({item['level_id'] for item in snapshots}) == 105 else "FAIL", "result": len({item['level_id'] for item in snapshots})},
        {"check": "节奏终点", "status": "PASS" if snapshots[-1]["pacing_level"] == 39 and snapshots[-1]["day"] == 8 else "FAIL", "result": f"Lv{snapshots[-1]['pacing_level']} / Day{snapshots[-1]['day']}"},
        {"check": "阵容终点", "status": "PASS" if snapshots[-1]["actual_team_size"] == 5 else "FAIL", "result": snapshots[-1]["actual_team_size"]},
        {"check": "关键阵容节点", "status": "PASS" if actual_key_team_sizes == key_team_sizes else "FAIL", "result": str(actual_key_team_sizes)},
        {"check": "0-3小寒上阵", "status": "PASS" if "hero_1040" in stage_0_3["team_members"] else "FAIL", "result": " / ".join(stage_0_3["team_member_names"])},
        {"check": "DPS未伪造", "status": "PASS" if all(item["team_dps"] is None for item in snapshots) else "FAIL", "result": "105关均为空"},
        {"check": "正式属性引用", "status": "PASS" if details and all(item["hp"] > 0 and item["attack"] > 0 for item in details) else "FAIL", "result": f"{len(details)}条上阵角色明细"},
        {"check": "招募引用链", "status": "PASS", "result": "SystemOpen.17 -> Unlock.4017 -> Level.100101; DrawCard.2 -> Item.311002 -> Hero.1002"},
        {"check": "小寒引用链", "status": "PASS", "result": "0-3战中固定上阵；Drop_mainlevel.10003 -> Item.311040 -> Hero.1040"},
        {"check": "阵位引用链", "status": "PASS", "result": "5/5 SystemOpen -> Unlock -> effective_seq"},
        {"check": "防御减伤参数", "status": "PASS", "result": formal_armor_reduce_x},
        {"check": "可发布性门禁", "status": "PASS" if not any(item["publish_ready"] for item in snapshots) else "FAIL", "result": "待战斗校准"},
    ]
    failed = [item for item in checks if item["status"] != "PASS"]
    if failed:
        raise RuntimeError("Validation failed: " + json.dumps(failed, ensure_ascii=False))

    generated_at = datetime.now().astimezone().isoformat(timespec="seconds")
    metadata = {
        "schema": "vgame.player-snapshot-output.v1",
        "model_id": policy["model_id"],
        "status": "mvp_core_growth_ready_battle_calibration_pending",
        "generated_at": generated_at,
        "stage_count": len(snapshots),
        "pacing_level_end": snapshots[-1]["pacing_level"],
        "day_end": snapshots[-1]["day"],
        "publish_ready": False,
    }
    sources = [
        {"name": name, "path": str(path), "sha256": sha256(path)}
        for name, path in required_paths.items()
    ]
    payload = {
        **metadata,
        "policy": policy,
        "checks": checks,
        "sources": sources,
        "stages": snapshots,
        "details": details,
    }
    args.output_dir.mkdir(parents=True, exist_ok=True)
    json_path = args.output_dir / "P50玩家快照MVP.json"
    xlsx_path = args.output_dir / "P50玩家快照MVP.xlsx"
    json_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    write_xlsx(xlsx_path, metadata, snapshots, details, policy, bindings, sources, checks)
    print(json.dumps({"json": str(json_path), "xlsx": str(xlsx_path), "checks": checks}, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
