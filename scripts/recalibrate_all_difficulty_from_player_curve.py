# -*- coding: utf-8 -*-
from __future__ import annotations

import argparse
import hashlib
import json
import math
import statistics
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


import vgame_paths

DESIGN_ROOT = Path(__file__).resolve().parents[1]
SOURCE_ROOT = vgame_paths.source_docs_root() / "框架v2" / "01_源表"
CONFIG_ROOT = vgame_paths.config_datas()
OUTPUT_ROOT = DESIGN_ROOT / "output" / "全玩法智能难度生成工作流"

CURVE_PATH = SOURCE_ROOT / "养成属性分配.xlsx"
DIFFICULTY_PATH = SOURCE_ROOT / "全玩法难度填写表.xlsx"
LEVEL_PATH = CONFIG_ROOT / "level" / "level.xlsx"
ROGUE_PATH = CONFIG_ROOT / "Rogue" / "RoguelikeLevel.xlsx"

PLAY_SHEETS = [
    "主线", "困难主线", "角色经验", "金币", "收容经验",
    "旅券", "映射", "Boss挑战", "肉鸽", "大秘境",
]
RESOURCE_SHEETS = ["角色经验", "金币", "收容经验"]
EXCLUDED_SHEETS = ["爬塔", "PVPVE", "无限挑战"]
DEFENSE_CONSTANT = 24000.0


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def scalar(value, default=1.0) -> float:
    return float(default if value in (None, "") else value)


def excel_round(value: float) -> int:
    return int(math.floor(value + 0.5))


def header_map(ws, row: int = 1) -> dict[str, int]:
    return {
        str(ws.cell(row, col).value): col
        for col in range(1, min(ws.max_column, 512) + 1)
        if ws.cell(row, col).value is not None
    }


def load_curve() -> list[dict]:
    wb = load_workbook(CURVE_PATH, data_only=True, read_only=False)
    ws = wb["玩家战力曲线"]
    rows = []
    for row in range(12, ws.max_row + 1):
        stage = ws.cell(row, 2).value
        if stage is None:
            continue
        record = {
            "source_row": row,
            "stage": int(stage),
            "player_level": ws.cell(row, 1).value,
            "chapter": ws.cell(row, 3).value,
            "stage_no": ws.cell(row, 4).value,
            "team_count": ws.cell(row, 43).value,
            "attack": ws.cell(row, 21).value,
            "defense": ws.cell(row, 16).value,
            "dps": ws.cell(row, 38).value,
            "ehp": ws.cell(row, 40).value,
        }
        for field in ("attack", "defense", "dps", "ehp"):
            if not isinstance(record[field], (int, float)) or record[field] <= 0:
                raise AssertionError(f"玩家战力曲线 row {row} {field} 无有效计算值：{record[field]}")
        rows.append(record)
    wb.close()
    if not rows or len({row["stage"] for row in rows}) != len(rows):
        raise AssertionError("玩家战力曲线为空或关卡序号重复")
    return rows


def load_baseline(wb) -> dict[int, dict]:
    ws = wb["_玩家曲线"]
    rows = {}
    for row in range(2, ws.max_row + 1):
        stage = ws.cell(row, 1).value
        if stage is None:
            continue
        rows[int(stage)] = {
            "player_level": ws.cell(row, 2).value,
            "chapter": ws.cell(row, 3).value,
            "stage_no": ws.cell(row, 4).value,
            "team_count": ws.cell(row, 5).value,
            "attack": ws.cell(row, 6).value,
            "defense": ws.cell(row, 7).value,
            "dps": ws.cell(row, 8).value,
            "ehp": ws.cell(row, 9).value,
        }
    return rows


def load_monster_model(wb) -> dict[tuple[int, str], tuple[float, float, float]]:
    ws = wb["_怪物模型"]
    model = {}
    for row in range(2, ws.max_row + 1):
        level = ws.cell(row, 2).value
        monster_type = ws.cell(row, 3).value
        if level is None or monster_type is None:
            continue
        model[(int(level), str(monster_type))] = (
            float(ws.cell(row, 4).value),
            float(ws.cell(row, 5).value),
            float(ws.cell(row, 6).value),
        )
    if not model:
        raise AssertionError("_怪物模型为空")
    return model


def row_metrics(record: dict, level: int, player: dict, model: dict) -> tuple[float, float]:
    try:
        hp, attack, defense = model[(level, record["monster_type"])]
    except KeyError as exc:
        raise AssertionError(f"怪物模型缺少 Lv{level}|{record['monster_type']}") from exc
    hp *= record["hp_coeff"]
    attack *= record["attack_coeff"]
    defense *= record["armor_coeff"]
    reduction = min(0.75, defense / (defense + DEFENSE_CONSTANT))
    monster_ehp = hp / (1.0 - reduction) * record["ehp_coeff"]
    ttk = monster_ehp / float(player["dps"])
    survival = float(player["ehp"]) / (attack * record["pressure_coeff"])
    return ttk, survival


def objective(metrics: tuple[float, float], target: tuple[float, float]) -> float:
    ttk, survival = metrics
    old_ttk, old_survival = target
    return 0.6 * math.log(ttk / old_ttk) ** 2 + 0.4 * math.log(survival / old_survival) ** 2


def load_play_rows(wb_values) -> dict[str, list[dict]]:
    result = {}
    for sheet in PLAY_SHEETS:
        ws = wb_values[sheet]
        rows = []
        for row in range(5, ws.max_row + 1):
            level_ids = ws.cell(row, 5).value
            old_level = ws.cell(row, 10).value
            if level_ids in (None, "") or old_level in (None, ""):
                continue
            rows.append(
                {
                    "sheet": sheet,
                    "row": row,
                    "node": str(ws.cell(row, 4).value or ""),
                    "level_ids": [part.strip() for part in str(level_ids).split(";") if part.strip().isdigit()],
                    "monster_type": str(ws.cell(row, 6).value),
                    "stage": int(ws.cell(row, 9).value),
                    "old_level": int(old_level),
                    "hp_coeff": scalar(ws.cell(row, 11).value),
                    "attack_coeff": scalar(ws.cell(row, 12).value),
                    "armor_coeff": scalar(ws.cell(row, 13).value),
                    "ehp_coeff": scalar(ws.cell(row, 14).value),
                    "pressure_coeff": scalar(ws.cell(row, 15).value),
                }
            )
        result[sheet] = rows
    return result


def build_costs(rows: list[dict], old_curve: dict, new_curve: dict, model: dict, weaker: bool) -> list[list[float]]:
    costs = []
    for row in rows:
        old_player = old_curve[row["stage"]]
        new_player = new_curve[row["stage"]]
        target = row_metrics(row, row["old_level"], old_player, model)
        row["old_ttk"], row["old_survival"] = target
        upper = row["old_level"] if weaker else 500
        values = [math.inf] * 501
        for level in range(1, upper + 1):
            values[level] = objective(row_metrics(row, level, new_player, model), target)
        row["ideal_level"] = min(range(1, upper + 1), key=lambda level: values[level])
        costs.append(values)
    return costs


def solve_monotonic(costs: list[list[float]], minimum: int = 1) -> list[int]:
    if not costs:
        return []
    size = 501
    back = []
    previous = [math.inf] * size
    for level in range(max(1, minimum), size):
        previous[level] = costs[0][level]
    back.append([-1] * size)
    for row_cost in costs[1:]:
        prefix_value = math.inf
        prefix_level = -1
        choices = [-1] * size
        current = [math.inf] * size
        for level in range(1, size):
            if previous[level] < prefix_value:
                prefix_value = previous[level]
                prefix_level = level
            choices[level] = prefix_level
            if math.isfinite(row_cost[level]) and math.isfinite(prefix_value):
                current[level] = row_cost[level] + prefix_value
        previous = current
        back.append(choices)
    last = min(range(1, size), key=lambda level: previous[level])
    if not math.isfinite(previous[last]):
        raise AssertionError("单调难度优化无可行解")
    levels = [last]
    for index in range(len(costs) - 1, 0, -1):
        last = back[index][last]
        levels.append(last)
    return list(reversed(levels))


def solve_mapping_linear(rows: list[dict], costs: list[list[float]], weaker: bool) -> list[int]:
    count = len(rows)
    best = None
    start_limit = rows[0]["old_level"] if weaker else 500
    end_limit = rows[-1]["old_level"] if weaker else 500
    for start in range(1, start_limit + 1):
        for end in range(start, end_limit + 1):
            levels = [excel_round(start + index * (end - start) / (count - 1)) for index in range(count)]
            if weaker and any(level > row["old_level"] for level, row in zip(levels, rows)):
                continue
            score = sum(costs[index][level] for index, level in enumerate(levels))
            if best is None or score < best[0]:
                best = (score, levels)
    if best is None:
        raise AssertionError("映射装备严格线性优化无可行解")
    return best[1]


def assign_levels(play_rows: dict, old_curve: dict, new_curve: dict, model: dict, weaker: bool) -> None:
    costs_by_sheet = {
        sheet: build_costs(rows, old_curve, new_curve, model, weaker)
        for sheet, rows in play_rows.items()
    }

    for sheet in PLAY_SHEETS:
        if sheet in RESOURCE_SHEETS or sheet in ("困难主线", "映射"):
            continue
        levels = solve_monotonic(costs_by_sheet[sheet])
        for row, level in zip(play_rows[sheet], levels):
            row["new_level"] = level

    resource_rows = play_rows[RESOURCE_SHEETS[0]]
    combined = []
    for index in range(len(resource_rows)):
        combined.append([
            sum(costs_by_sheet[sheet][index][level] for sheet in RESOURCE_SHEETS)
            for level in range(501)
        ])
    resource_levels = solve_monotonic(combined)
    for sheet in RESOURCE_SHEETS:
        if len(play_rows[sheet]) != len(resource_levels):
            raise AssertionError("三个资源副本档位数量不一致")
        for row, level in zip(play_rows[sheet], resource_levels):
            row["new_level"] = level

    mapping_levels = solve_mapping_linear(play_rows["映射"], costs_by_sheet["映射"], weaker)
    for row, level in zip(play_rows["映射"], mapping_levels):
        row["new_level"] = level

    main_515 = next((row for row in play_rows["主线"] if row["node"] == "5-15"), None)
    if main_515 is None:
        raise AssertionError("主线未找到5-15，无法校验困难主线起点")
    old_gap = play_rows["困难主线"][0]["old_level"] - main_515["old_level"]
    hard_minimum = main_515["new_level"] + max(1, old_gap)
    hard_levels = solve_monotonic(costs_by_sheet["困难主线"], minimum=hard_minimum)
    for row, level in zip(play_rows["困难主线"], hard_levels):
        row["new_level"] = level

    for sheet, rows in play_rows.items():
        for row in rows:
            new_ttk, new_survival = row_metrics(row, row["new_level"], new_curve[row["stage"]], model)
            row["new_ttk"] = new_ttk
            row["new_survival"] = new_survival
            row["ttk_log_error"] = abs(math.log(new_ttk / row["old_ttk"]))
            row["survival_log_error"] = abs(math.log(new_survival / row["old_survival"]))


def build_level_writes(play_rows: dict) -> tuple[list[dict], dict[str, dict]]:
    desired = {}
    for sheet, rows in play_rows.items():
        for row in rows:
            for level_id in row["level_ids"]:
                if level_id in desired:
                    raise AssertionError(f"LevelId重复映射：{level_id}")
                desired[level_id] = {
                    "sheet": sheet,
                    "node": row["node"],
                    "grade": row["new_level"],
                }
    rogue_ids = sorted(int(level_id) for level_id, item in desired.items() if item["sheet"] == "肉鸽")
    if rogue_ids != list(range(1000001, 1000006)):
        raise AssertionError(f"肉鸽LevelId应为1000001~1000005，实际{rogue_ids}")

    wb = load_workbook(LEVEL_PATH, data_only=False, read_only=False)
    ws = wb[wb.sheetnames[0]]
    headers = header_map(ws)
    id_col = headers["LevelId"]
    grade_col = headers["Grade"]
    matched = set()
    writes = []
    for row in range(5, ws.max_row + 1):
        value = ws.cell(row, id_col).value
        if value is None:
            continue
        level_id = str(int(value)) if isinstance(value, (int, float)) else str(value).strip()
        item = desired.get(level_id)
        if item is None:
            continue
        grade_cell = ws.cell(row, grade_col)
        if grade_cell.data_type == "f":
            raise AssertionError(f"level.xlsx {grade_cell.coordinate} Grade为公式")
        matched.add(level_id)
        writes.append(
            {
                "sheet": ws.title,
                "row": row,
                "column": grade_col,
                "cell": grade_cell.coordinate,
                "level_id": level_id,
                "play": item["sheet"],
                "node": item["node"],
                "old_grade": grade_cell.value,
                "new_grade": item["grade"],
            }
        )
    wb.close()
    missing = sorted(set(desired) - matched, key=int)
    if missing:
        raise AssertionError(f"level.xlsx 未命中{len(missing)}个LevelId：{missing[:10]}")
    return writes, desired


def summarize(play_rows: dict) -> dict:
    result = {}
    for sheet, rows in play_rows.items():
        ttk_errors = [row["ttk_log_error"] for row in rows]
        survival_errors = [row["survival_log_error"] for row in rows]
        result[sheet] = {
            "count": len(rows),
            "old_range": [min(row["old_level"] for row in rows), max(row["old_level"] for row in rows)],
            "new_range": [min(row["new_level"] for row in rows), max(row["new_level"] for row in rows)],
            "changed": sum(row["old_level"] != row["new_level"] for row in rows),
            "delta_range": [
                min(row["new_level"] - row["old_level"] for row in rows),
                max(row["new_level"] - row["old_level"] for row in rows),
            ],
            "old_levels": [row["old_level"] for row in rows] if len(rows) <= 35 else None,
            "new_levels": [row["new_level"] for row in rows] if len(rows) <= 35 else None,
            "monotonic": all(rows[index]["new_level"] <= rows[index + 1]["new_level"] for index in range(len(rows) - 1)),
            "ttk_log_error_median_max": [statistics.median(ttk_errors), max(ttk_errors)],
            "survival_log_error_median_max": [statistics.median(survival_errors), max(survival_errors)],
            "first3": [
                {key: row[key] for key in ("row", "node", "stage", "old_level", "ideal_level", "new_level", "old_ttk", "new_ttk", "old_survival", "new_survival")}
                for row in rows[:3]
            ],
            "last3": [
                {key: row[key] for key in ("row", "node", "stage", "old_level", "ideal_level", "new_level", "old_ttk", "new_ttk", "old_survival", "new_survival")}
                for row in rows[-3:]
            ],
        }
    return result


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()

    for path in (CURVE_PATH, DIFFICULTY_PATH, LEVEL_PATH, ROGUE_PATH):
        if not path.exists():
            raise FileNotFoundError(path)

    curve_rows = load_curve()
    new_curve = {row["stage"]: row for row in curve_rows}
    difficulty_values = load_workbook(DIFFICULTY_PATH, data_only=True, read_only=False)
    difficulty_formulas = load_workbook(DIFFICULTY_PATH, data_only=False, read_only=False)
    old_curve = load_baseline(difficulty_values)
    model = load_monster_model(difficulty_values)
    play_rows = load_play_rows(difficulty_values)

    missing_stages = sorted({row["stage"] for rows in play_rows.values() for row in rows} - set(new_curve))
    if missing_stages:
        raise AssertionError(f"新玩家曲线缺少对照关卡：{missing_stages}")
    missing_old_stages = sorted(set(new_curve) - set(old_curve))
    if missing_old_stages:
        raise AssertionError(f"旧玩家曲线缺少关卡：{missing_old_stages}")

    ratios = []
    for stage, player in new_curve.items():
        old = old_curve[stage]
        ratios.append(
            {
                "stage": stage,
                "dps": float(player["dps"]) / float(old["dps"]),
                "ehp": float(player["ehp"]) / float(old["ehp"]),
            }
        )
    weaker = all(item["dps"] <= 1.0 + 1e-10 and item["ehp"] <= 1.0 + 1e-10 for item in ratios)
    assign_levels(play_rows, old_curve, new_curve, model, weaker)

    source_writes = []
    for sheet, rows in play_rows.items():
        for row in rows:
            if sheet == "映射" and row["row"] not in (5, 34):
                formula = difficulty_formulas[sheet].cell(row["row"], 10).value
                if not (isinstance(formula, str) and formula.startswith("=ROUND(")):
                    raise AssertionError(f"映射 {sheet}!J{row['row']} 线性公式被破坏")
                continue
            source_writes.append(
                {
                    "sheet": sheet,
                    "row": row["row"],
                    "column": 10,
                    "cell": f"J{row['row']}",
                    "old_level": row["old_level"],
                    "new_level": row["new_level"],
                }
            )

    level_writes, desired = build_level_writes(play_rows)
    if len(desired) != 526 or len(level_writes) != 526:
        raise AssertionError(f"level.xlsx 映射数量应为526，实际{len(desired)}/{len(level_writes)}")

    hard_first = play_rows["困难主线"][0]["new_level"]
    main_515 = next(row["new_level"] for row in play_rows["主线"] if row["node"] == "5-15")
    if hard_first <= main_515:
        raise AssertionError(f"困难主线首档Lv{hard_first}未高于主线5-15 Lv{main_515}")

    plan = {
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "source": {
            "curve_path": str(CURVE_PATH),
            "curve_sha256": sha256(CURVE_PATH),
            "difficulty_path": str(DIFFICULTY_PATH),
            "difficulty_sha256": sha256(DIFFICULTY_PATH),
            "level_path": str(LEVEL_PATH),
            "level_sha256": sha256(LEVEL_PATH),
            "roguelike_path_readonly": str(ROGUE_PATH),
            "roguelike_sha256": sha256(ROGUE_PATH),
        },
        "algorithm": {
            "target": "保持当前每节点预测TTK与预测存活时间",
            "objective": "0.6*ln(TTK新/TTK旧)^2 + 0.4*ln(存活新/存活旧)^2",
            "player_curve_is_globally_weaker": weaker,
            "constraints": [
                "怪物等级1-500；玩家整体变弱时新等级不高于旧值",
                "各玩法表内等级不下降",
                "角色经验/金币/收容经验共用同一等级序列",
                "映射装备1-30档由首末等级严格线性生成，保留J6:J33原公式",
                "困难主线首档严格高于主线5-15",
                "爬塔由TeamHero承载；PVPVE双方属性对算；无限挑战不在本轮",
                "肉鸽写level.xlsx的LevelId 1000001~1000005，不改RoguelikeLevel.xlsx",
            ],
        },
        "curve_change": {
            "count": len(curve_rows),
            "changed_count": sum(abs(item["dps"] - 1.0) > 1e-10 or abs(item["ehp"] - 1.0) > 1e-10 for item in ratios),
            "dps_ratio_min_max": [min(item["dps"] for item in ratios), max(item["dps"] for item in ratios)],
            "ehp_ratio_min_max": [min(item["ehp"] for item in ratios), max(item["ehp"] for item in ratios)],
            "samples": [item for item in ratios if item["stage"] in (1, 16, 54, 99, 146, 240, 279)],
        },
        "curve_rows": [
            {
                "target_row": index + 2,
                "values": [
                    row["stage"], row["player_level"], row["chapter"], row["stage_no"],
                    row["team_count"], row["attack"], row["defense"], row["dps"], row["ehp"],
                ],
            }
            for index, row in enumerate(curve_rows)
        ],
        "source_writes": source_writes,
        "level_writes": level_writes,
        "summary": summarize(play_rows),
        "checks": {
            "difficulty_source_write_cells": len(source_writes),
            "level_id_count": len(desired),
            "level_changed_count": sum(item["old_grade"] != item["new_grade"] for item in level_writes),
            "level_noop_count": sum(item["old_grade"] == item["new_grade"] for item in level_writes),
            "hard_first": hard_first,
            "main_5_15": main_515,
            "mapping_formula_cells_preserved": 28,
            "excluded_sheets": EXCLUDED_SHEETS,
        },
    }
    difficulty_values.close()
    difficulty_formulas.close()

    output = args.output or OUTPUT_ROOT / f"player-curve-resync-write-plan-{datetime.now():%Y%m%d-%H%M%S}.json"
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(plan, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"plan": str(output), "curve_change": plan["curve_change"], "summary": plan["summary"], "checks": plan["checks"]}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
