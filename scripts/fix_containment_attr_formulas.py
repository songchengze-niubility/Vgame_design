# -*- coding: utf-8 -*-
"""修正 insert_rows 后原 A/B 行公式行号偏移（+13）。"""
import openpyxl
import re

import vgame_paths

SRC = str(vgame_paths.source_docs_root() / "框架v2" / "01_源表" / "养成属性分配.xlsx")
wb = openpyxl.load_workbook(SRC, data_only=False)
ws = wb["收容物属性分配"]

fixed = 0
for r in range(17, 43):  # 原 A级15 + B级11 = 26 行，现位于 17~42
    old = r - 13
    for col in ("E", "F", "G"):
        cell = ws[f"{col}{r}"]
        v = cell.value
        if isinstance(v, str) and v.startswith("="):
            v = v.replace(f"$D{old}", f"$D{r}")
            cell.value = v
            fixed += 1
    for col in ("L", "M", "N"):
        cell = ws[f"{col}{r}"]
        v = cell.value
        if isinstance(v, str) and v.startswith("="):
            # 形如 =E4*H4*$K4
            colA = {"L": "E", "M": "F", "N": "G"}[col]
            colB = {"L": "H", "M": "I", "N": "J"}[col]
            cell.value = f"={colA}{r}*{colB}{r}*$K{r}"
            fixed += 1
    cell = ws[f"K{r}"]
    v = cell.value
    if isinstance(v, str) and v.startswith("="):
        v = v.replace(f"VLOOKUP(C{old},", f"VLOOKUP(C{r},")
        cell.value = v
        fixed += 1

wb.save(SRC)
print("修正公式数:", fixed)

# 验证
wb2 = openpyxl.load_workbook(SRC, data_only=False)
ws2 = wb2["收容物属性分配"]
for r in (17, 18, 31, 42):
    print(f"r{r} {ws2['B'+str(r)].value}: E={ws2['E'+str(r)].value}")
    print(f"     K={ws2['K'+str(r)].value} L={ws2['L'+str(r)].value}")
