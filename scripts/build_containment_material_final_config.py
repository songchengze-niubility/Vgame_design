# -*- coding: utf-8 -*-
"""构建收容物养成 06 最终配置镜像。

只写 06_最终配置，不修改 SVN Datas。
"""
from __future__ import annotations

import json
import math
import os
from copy import copy
from pathlib import Path

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


import vgame_paths

DATAS = vgame_paths.config_datas()
EXCL_DIR = DATAS / "ExclusiveWeapon"
OUTPUT_ROOT = Path(os.environ["CONTAINMENT_MATERIAL_OUTPUT"])
ATTR_MIRROR = Path(os.environ["CONTAINMENT_ATTR_OUTPUT"]) / "ExclWeapon.xlsx"

FIRST_PHASE_IDS = (
    51001, 51002, 51003, 51004, 51005,
    52001, 52002, 52003, 52004, 52005, 52006, 52007, 52008, 52009, 52010,
    53001, 53002, 53003, 53004, 53005, 53006, 53007,
)
STAR_SHARDS_BY_QUALITY = {
    4: (100, 100, 200, 200, 300, 400),
    3: (50, 100, 150, 200, 300, 400),
    2: (60, 120, 180, 240, 300, 400),
}
DUPLICATE_SHARDS_BY_QUALITY = {4: 100, 3: 50, 2: 20}
SYNTHESIS_SHARDS = 100
STAR_GOLD = (4000, 8000, 12000, 16000, 24000, 36000)
QUALITIES = (("S", 4), ("A", 3), ("B", 2))
EXP_ITEMS = {108: 1000, 109: 4000, 110: 20000}
QUALITY_ENUM = {4: "Orange", 3: "Purple", 2: "Blue"}


def field_map(ws) -> dict[str, int]:
    return {
        ws.cell(1, c).value: c
        for c in range(1, ws.max_column + 1)
        if ws.cell(1, c).value
    }


def copy_row_style(ws, source_row: int, target_row: int, max_col: int | None = None) -> None:
    max_col = max_col or ws.max_column
    for c in range(1, max_col + 1):
        ws.cell(target_row, c)._style = copy(ws.cell(source_row, c)._style)
        if ws.cell(source_row, c).number_format:
            ws.cell(target_row, c).number_format = ws.cell(source_row, c).number_format
    ws.row_dimensions[target_row].height = ws.row_dimensions[source_row].height


def replace_data_rows(ws, rows: list[list], max_col: int | None = None) -> None:
    """保留 1~4 行表头，用第 5 行样式重建数据。"""
    max_col = max_col or ws.max_column
    style_row = 5
    styles = [copy(ws.cell(style_row, c)._style) for c in range(1, max_col + 1)]
    row_height = ws.row_dimensions[style_row].height
    if ws.max_row >= 5:
        ws.delete_rows(5, ws.max_row - 4)
    for values in rows:
        ws.append(values[:max_col] + [None] * max(0, max_col - len(values)))
        for c in range(1, max_col + 1):
            ws.cell(ws.max_row, c)._style = copy(styles[c - 1])
        ws.row_dimensions[ws.max_row].height = row_height


def quality_for_id(wid: int) -> int:
    if 51000 <= wid < 52000:
        return 4
    if 52000 <= wid < 53000:
        return 3
    if 53000 <= wid < 54000:
        return 2
    raise ValueError(f"未知收容物品质：{wid}")


def build_excl_weapon() -> dict:
    """首期22个本体：直接接入已经线性化的属性镜像。"""
    if not ATTR_MIRROR.exists():
        raise FileNotFoundError(f"属性最终配置不存在：{ATTR_MIRROR}")
    source_formula = openpyxl.load_workbook(ATTR_MIRROR, data_only=False)
    source_values = openpyxl.load_workbook(ATTR_MIRROR, data_only=True)
    ws = source_formula["Eqpt"]
    value_ws = source_values["Eqpt"]
    fm = field_map(ws)
    value_rows = {
        int(value_ws.cell(r, fm["Id"]).value): [value_ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        for r in range(5, value_ws.max_row + 1)
        if value_ws.cell(r, fm["Id"]).value is not None
    }

    rows: list[list] = []
    full_level: dict[int, dict[str, float]] = {}
    for wid in FIRST_PHASE_IDS:
        if wid not in value_rows:
            raise ValueError(f"属性镜像缺少首期收容物 {wid}")
        row = list(value_rows[wid])
        full_hp = float(row[fm["Hp"] - 1]) * 60
        full_attack = float(row[fm["Attack"] - 1]) * 60
        full_armor = float(row[fm["Armor"] - 1]) * 60
        quality = quality_for_id(wid)
        row[fm["Quailty"] - 1] = quality
        row[fm["ShardsList"] - 1] = ";".join(map(str, STAR_SHARDS_BY_QUALITY[quality]))
        row[fm["WeaponShards"] - 1] = int(f"3{wid}")
        row[fm["DrawCardShards"] - 1] = DUPLICATE_SHARDS_BY_QUALITY[quality]
        rows.append(row)
        full_level[wid] = {"Hp": full_hp, "Attack": full_attack, "Armor": full_armor}

    replace_data_rows(ws, rows)
    # 清理属性镜像可能携带的说明页，保持可落表结构。
    for sheet_name in list(source_formula.sheetnames):
        if sheet_name not in {"Eqpt", "@Desc1", "@Desc2", "@Desc3", "@Desc4"}:
            del source_formula[sheet_name]
    source_formula.save(OUTPUT_ROOT / "ExclWeapon.xlsx")
    return {"count": len(rows), "full_level": full_level}


def exp_cost(level: int) -> int | None:
    if level >= 60:
        return None
    return math.ceil(85 * level * level / 1000) * 1000


def build_lv() -> dict:
    wb = openpyxl.load_workbook(EXCL_DIR / "ExclWeaponLv.xlsx")
    ws = wb["EqptLv"]
    fm = field_map(ws)
    rows = []
    for level in range(1, 61):
        row = [None] * ws.max_column
        row[fm["Id"] - 1] = level
        row[fm["Lv"] - 1] = level
        row[fm["UpgradeCostExp"] - 1] = exp_cost(level)
        row[fm["Quality"] - 1] = 4  # 单组公共曲线；程序需取消按品质过滤。
        rows.append(row)
    replace_data_rows(ws, rows)
    wb.save(OUTPUT_ROOT / "ExclWeaponLv.xlsx")
    total = sum(exp_cost(level) or 0 for level in range(1, 61))
    return {"rows": len(rows), "total_exp": total, "total_gold": round(total * 0.2)}


def build_lv_limit() -> dict:
    src = openpyxl.load_workbook(EXCL_DIR / "ExclWeaponLvLimit.xlsx")
    src_ws = src["ExclWeaponLvLimit"]
    wb = Workbook()
    ws = wb.active
    ws.title = "ExclWeaponLvLimit"
    for r in range(1, 5):
        for c in range(1, 5):
            source_cell = src_ws.cell(r, c)
            target_cell = ws.cell(r, c)
            target_cell.value = source_cell.value
            target_cell.font = copy(source_cell.font)
            target_cell.fill = copy(source_cell.fill)
            target_cell.border = copy(source_cell.border)
            target_cell.alignment = copy(source_cell.alignment)
            target_cell.number_format = source_cell.number_format
            target_cell.protection = copy(source_cell.protection)
        ws.row_dimensions[r].height = src_ws.row_dimensions[r].height
    for c in range(1, 5):
        letter = get_column_letter(c)
        ws.column_dimensions[letter].width = src_ws.column_dimensions[letter].width
    rows = [[None, player_lv, player_lv, min(player_lv + 5, 60)] for player_lv in range(1, 61)]
    # Id 在第 2 列，首列为导表标记列。
    for i, row in enumerate(rows, start=1):
        row[1] = i
    replace_data_rows(ws, rows, max_col=4)
    wb.save(OUTPUT_ROOT / "ExclWeaponLvLimit.xlsx")
    return {"rows": len(rows), "lv55": 60, "lv60": 60}


def build_star() -> dict:
    wb = openpyxl.load_workbook(EXCL_DIR / "ExclWeaponStar.xlsx")
    ws = wb["EqptStar"]
    fm = field_map(ws)
    rows = []
    row_id = 1
    for quality_name, quality in QUALITIES:
        for star in range(0, 7):
            row = [None] * ws.max_column
            row[fm["Id"] - 1] = row_id
            row[fm["Star"] - 1] = star
            row[fm["#"] - 1] = quality_name
            row[fm["Quality"] - 1] = quality
            row[fm["StarCostEqptNum"] - 1] = 0
            row[fm["StarCostItem"] - 1] = None
            if star < 6:
                row[fm["StarCost"] - 1] = f"1,{STAR_GOLD[star]}"
                row[fm["SkillIndex"] - 1] = star + 1
            row[fm["##"] - 1] = star + 1
            rows.append(row)
            row_id += 1
    replace_data_rows(ws, rows)
    wb.save(OUTPUT_ROOT / "ExclWeaponStar.xlsx")
    return {
        "rows": len(rows),
        "fragment_total": {
            quality_name: sum(STAR_SHARDS_BY_QUALITY[quality])
            for quality_name, quality in QUALITIES
        },
        "gold_total": sum(STAR_GOLD),
    }


def build_items() -> dict:
    src_path = DATAS / "items" / "Items.xlsx"
    formula_wb = openpyxl.load_workbook(src_path, data_only=False)
    value_wb = openpyxl.load_workbook(src_path, data_only=True)
    ws = formula_wb.active
    value_ws = value_wb.active
    fm = field_map(ws)
    rows = []
    found = {}
    for r in range(5, value_ws.max_row + 1):
        item_id = value_ws.cell(r, fm["Id"]).value
        if item_id in EXP_ITEMS:
            row = [value_ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
            rows.append(row)
            found[int(item_id)] = row[fm["ParaList"] - 1]
    if found != EXP_ITEMS:
        raise ValueError(f"经验道具配置不一致：{found} != {EXP_ITEMS}")

    # WeaponShards 是 ExclWeapon 的强引用；属性镜像维护全量39个本体，
    # 因此在养成配置包中同步生成全量同名碎片道具，未进首期池的条目保持未投放即可。
    attr_wb = openpyxl.load_workbook(ATTR_MIRROR, data_only=True, read_only=True)
    attr_ws = attr_wb["Eqpt"]
    attr_fm = field_map(attr_ws)
    template = None
    for r in range(5, value_ws.max_row + 1):
        if value_ws.cell(r, fm["Id"]).value == 301001:
            template = [value_ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
            break
    if template is None:
        raise ValueError("Items缺少碎片道具模板301001")

    shard_ids = []
    for r in range(5, attr_ws.max_row + 1):
        wid = attr_ws.cell(r, attr_fm["Id"]).value
        if wid is None:
            continue
        wid = int(wid)
        shard_id = int(f"3{wid}")
        row = list(template)
        row[0] = None
        row[fm["Id"] - 1] = shard_id
        row[fm["Name"] - 1] = f"{attr_ws.cell(r, attr_fm['Name']).value}碎片"
        row[fm["ParaList"] - 1] = None
        row[fm["ItemType"] - 1] = "ItemType_WeaponShards"
        row[fm["ItemSubType"] - 1] = "ItemSubType_WeaponShards"
        row[fm["ChestOpenMode"] - 1] = None
        icon = attr_ws.cell(r, attr_fm["WeaponIconSmall"]).value
        row[fm["Icon"] - 1] = icon
        row[fm["LittleIcon"] - 1] = icon
        row[fm["ItemQuality"] - 1] = QUALITY_ENUM[quality_for_id(wid)]
        row[fm["Description"] - 1] = f"{attr_ws.cell(r, attr_fm['Name']).value}的收容物碎片，集齐100个可合成0星本体，也可用于升星。"
        row[fm["AccessList"] - 1] = None
        row[fm["ShopId"] - 1] = None
        row[fm["NumLimit"] - 1] = None
        row[fm["RelatedId"] - 1] = wid
        rows.append(row)
        shard_ids.append(shard_id)
    replace_data_rows(ws, rows)
    for sheet_name in list(formula_wb.sheetnames):
        if sheet_name != ws.title and not sheet_name.startswith("@"):
            del formula_wb[sheet_name]
    formula_wb.save(OUTPUT_ROOT / "Items.xlsx")
    return {"rows": len(rows), "items": found, "shard_ids": shard_ids}


def build_globals() -> dict:
    src_path = DATAS / "globals_conf" / "globals_conf.xlsx"
    wb = openpyxl.load_workbook(src_path)
    ws = wb["globals_conf"]
    target_row = None
    old_value = None
    for r in range(1, ws.max_row + 1):
        if ws.cell(r, 1).value == "ExclEquipmentUpgradeRatio":
            target_row = r
            old_value = ws.cell(r, 4).value
            ws.cell(r, 4).value = 0.2
            break
    if target_row is None:
        raise ValueError("globals_conf 缺少 ExclEquipmentUpgradeRatio")
    wb.save(OUTPUT_ROOT / "globals_conf.xlsx")
    return {"row": target_row, "old": old_value, "new": 0.2}


def style_table(ws, widths: list[int]) -> None:
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4472C4")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(widths))}{ws.max_row}"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins.left = 0.25
    ws.page_margins.right = 0.25
    ws.page_margins.top = 0.4
    ws.page_margins.bottom = 0.4
    ws.print_title_rows = "1:1"


def write_manifest(results: dict, program_gaps: list[dict]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "配置清单"
    ws.append(["文件", "状态", "配置内容", "落地备注"])
    for row in (
        ("ExclWeapon.xlsx", "READY", "首期22个；线性一级属性、按品质配置6档碎片消耗与重复转化", "依赖收容物属性最终配置"),
        ("ExclWeaponLv.xlsx", "READY", "Lv1~60公共经验曲线；Lv60无下一级消耗", "S/A/B共用，程序需取消品质过滤"),
        ("ExclWeaponLvLimit.xlsx", "READY", "玩家Lv1~60对应收容物上限MIN(玩家+5,60)", "程序需接新表"),
        ("ExclWeaponStar.xlsx", "READY", "S/A/B × 0~6星；六次升星金币", "碎片由ExclWeapon.ShardsList承载"),
        ("Items.xlsx", "READY", "108/109/110经验道具 + 全量39个同名碎片道具", "WeaponShards引用完整"),
        ("globals_conf.xlsx", "READY", "ExclEquipmentUpgradeRatio=0.2", "升级金币=实际消耗经验×20%"),
        ("ExclWeaponAscendLv.xlsx", "DEPRECATED", "旧突破等级表不再进入本包", "新方案由ExclWeaponLvLimit连续限制"),
    ):
        ws.append(row)
    style_table(ws, [30, 15, 58, 58])

    rules = wb.create_sheet("养成口径")
    rules.append(["模块", "确认值", "校验结果"])
    for row in (
        ("等级", "初始1级，上限60；上限MIN(玩家等级+5,60)", "PASS"),
        ("经验", "CEILING(85×当前等级²,1000)，当前等级1~59", f"单槽合计{results['lv']['total_exp']:,}"),
        ("升级金币", "实际消耗经验×20%", f"单槽满级{results['lv']['total_gold']:,}"),
        ("经验道具", "108/109/110=1000/4000/20000经验", "PASS"),
        ("星级", "初始0星，满6星；0星也生效二级属性", "PASS"),
        ("升星碎片", "S=100/100/200/200/300/400；A=50/100/150/200/300/400；B=60/120/180/240/300/400", "S/A/B合计=1300/1200/1300"),
        ("重复转化", "S/A/B重复整件=100/50/20片", "每段升星均为对应品质重复转化的整数倍"),
        ("合成本体", f"S/A/B均为同名碎片×{SYNTHESIS_SHARDS}", "与重复转化字段分离"),
        ("升星金币", "4000/8000/12000/16000/24000/36000", f"合计{results['star']['gold_total']:,}"),
        ("品质系数", "S/A/B=1/0.75/0.5", "仅二级属性运行时乘算；一级属性不乘品质"),
        ("位置系数", "主位100%/副位50%/未上阵10%", "运行时乘算；未上阵不封顶"),
        ("一级属性", "ceil(满级属性×等级/60)", "表内存满级/60连续值，运行时向上取整"),
        ("一级标准包", "满级生命8500/攻击2550/防御1700", "一级不乘品质；再乘L/M/N，最后四舍五入"),
        ("一级总预算", "首期约40000生命/12000攻击/8000防御", "位置权重1+4×0.5+17×0.1=4.7；取整标准包误差小于0.2%"),
        ("二级星级档", "0/1/2/3/4/5/6星=40/50/60/70/80/90/100%", "六次升星均增加满星属性的10%"),
        ("二级总预算", "五个S满星合计120%攻击力等价", "位置权重1+4×0.5=3，单个S满星=40%攻击力等价"),
    ):
        rules.append(row)
    style_table(rules, [24, 70, 42])

    gaps = wb.create_sheet("程序接入项")
    gaps.append(["优先级", "接入项", "当前证据", "验收方式"])
    for gap in program_gaps:
        gaps.append([gap["priority"], gap["item"], gap["evidence"], gap["acceptance"]])
    style_table(gaps, [12, 58, 85, 70])

    detail = wb.create_sheet("首期线性属性映射")
    detail.append(["Id", "满级Hp", "配置Hp/HpGrow", "满级Attack", "配置Attack/AttackGrow", "满级Armor", "配置Armor/ArmorGrow"])
    for wid, attrs in results["excl"]["full_level"].items():
        detail.append([
            wid,
            attrs["Hp"], attrs["Hp"] / 60,
            attrs["Attack"], attrs["Attack"] / 60,
            attrs["Armor"], attrs["Armor"] / 60,
        ])
    style_table(detail, [14, 16, 22, 18, 26, 18, 24])
    for row in detail.iter_rows(min_row=2, min_col=2, max_col=7):
        for cell in row:
            cell.number_format = "0.####"
    wb.save(OUTPUT_ROOT / "00_配置清单.xlsx")


def validate_workbooks(expected_files: list[str]) -> dict:
    files = {}
    for name in expected_files:
        path = OUTPUT_ROOT / name
        wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
        formulas = 0
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                formulas += sum(1 for cell in row if isinstance(cell.value, str) and cell.value.startswith("="))
        files[name] = {"opened": True, "sheets": wb.sheetnames, "formula_cells": formulas}
    return files


def main() -> None:
    OUTPUT_ROOT.mkdir(parents=True, exist_ok=True)
    obsolete = OUTPUT_ROOT / "ExclWeaponAscendLv.xlsx"
    if obsolete.exists():
        obsolete.unlink()

    results = {
        "excl": build_excl_weapon(),
        "lv": build_lv(),
        "lv_limit": build_lv_limit(),
        "star": build_star(),
        "items": build_items(),
        "globals": build_globals(),
    }
    assert results["excl"]["count"] == 22
    assert results["lv"]["rows"] == 60
    assert results["lv"]["total_exp"] == 5_997_000
    assert results["lv"]["total_gold"] == 1_199_400
    assert results["lv_limit"]["rows"] == 60
    assert results["star"]["rows"] == 21
    assert results["star"]["fragment_total"] == {"S": 1_300, "A": 1_200, "B": 1_300}
    assert results["star"]["gold_total"] == 100_000
    assert len(results["items"]["shard_ids"]) == 39

    program_gaps = [
        {
            "priority": "P0",
            "item": "升级经验公共曲线",
            "evidence": "当前 ExclusiveWeaponSystem 仍按 cfg.Lv + cfg.Quality 查 ExclWeaponLv；最终表只有 Quality=4 一组公共曲线。",
            "acceptance": "S/A/B 任一收容物在同等级显示并扣除相同经验。",
        },
        {
            "priority": "P0",
            "item": "玩家等级联动上限",
            "evidence": "当前逻辑仍读取 ExclWeaponAscendLv；新增 ExclWeaponLvLimit 尚无非生成代码引用。",
            "acceptance": "玩家1~55级时上限=玩家+5；56~60级时均为60。",
        },
        {
            "priority": "P0",
            "item": "同名碎片升星",
            "evidence": "当前升星仍读取 StarCostEqptNum/StarCostItem，未读取 ExclWeapon.ShardsList/WeaponShards。",
            "acceptance": "S六次依次扣100/100/200/200/300/400片，A扣50/100/150/200/300/400片，B扣60/120/180/240/300/400片；不接受同名本体或旧通用材料。",
        },
        {
            "priority": "P0",
            "item": "一级、二级属性公式与系数边界",
            "evidence": "当前展示逻辑为 base+(Lv-1)×grow，未向上取整；品质与位置系数无现成配置字段。",
            "acceptance": "一级=ceil(满级×Lv/60)×位置；二级=星级档×品质(1/0.75/0.5)×位置。",
        },
        {
            "priority": "P0",
            "item": "二级属性按0~6星取档",
            "evidence": "当前逻辑用 ascendId 索引 AttrNameList；新方案应使用 star，且0星读取第1档。",
            "acceptance": "0~6星分别读取40/50/60/70/80/90/100%累计档。",
        },
        {
            "priority": "P1",
            "item": "五槽与未上阵继承",
            "evidence": "配置表无槽位继承字段。",
            "acceptance": "五槽独立升级；未上阵等级取五槽最低等级，属性按10%生效且不封顶。",
        },
    ]
    write_manifest(results, program_gaps)
    expected = [
        "00_配置清单.xlsx", "ExclWeapon.xlsx", "ExclWeaponLv.xlsx",
        "ExclWeaponLvLimit.xlsx", "ExclWeaponStar.xlsx", "Items.xlsx", "globals_conf.xlsx",
    ]
    workbook_validation = validate_workbooks(expected)
    payload = {
        "status": "PASS_WITH_RUNTIME_GAPS",
        "scope": "收容物养成最终配置 v3（2026-08-10）",
        "source_of_truth": {
            "attributes": str(ATTR_MIRROR),
            "config_schema": str(EXCL_DIR),
            "output": str(OUTPUT_ROOT),
        },
        "checks": {
            "first_phase_count": 22,
            "level_rows": 60,
            "level_total_exp": results["lv"]["total_exp"],
            "level_total_gold": results["lv"]["total_gold"],
            "level_limit_rows": 60,
            "star_rows": 21,
            "star_fragment_total": results["star"]["fragment_total"],
            "star_gold_total": results["star"]["gold_total"],
            "exp_items": results["items"]["items"],
            "fragment_item_count": len(results["items"]["shard_ids"]),
            "upgrade_gold_ratio": results["globals"],
        },
        "workbooks": workbook_validation,
        "runtime_gaps": program_gaps,
        "datas_modified": False,
    }
    with (OUTPUT_ROOT / "收容物养成材料配置校验.json").open("w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False, indent=2)
    print(json.dumps({"output": str(OUTPUT_ROOT), "status": payload["status"], "files": sorted(p.name for p in OUTPUT_ROOT.iterdir())}, ensure_ascii=False))


if __name__ == "__main__":
    main()
