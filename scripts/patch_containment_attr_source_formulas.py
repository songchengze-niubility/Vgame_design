# -*- coding: utf-8 -*-
"""修正收容物属性源表：L/M/N不含品质系数，K保留二级属性品质系数。"""
from pathlib import Path

import openpyxl


import vgame_paths

SOURCE = vgame_paths.source_docs_root() / "框架v2" / "01_源表" / "养成属性分配.xlsx"


def main() -> None:
    wb = openpyxl.load_workbook(SOURCE, data_only=False)
    role = wb["角色属性分配"]
    containment = wb["收容物属性分配"]

    quality = {"SSR": 1.0, "SR": 0.75, "R": 0.5}
    for row in range(19, 22):
        name = role.cell(row, 2).value
        if name in quality:
            role.cell(row, 3).value = quality[name]

    for row in range(4, containment.max_row + 1):
        if containment.cell(row, 1).value is None:
            continue
        containment.cell(row, 12).value = f"=E{row}*H{row}"
        containment.cell(row, 13).value = f"=F{row}*I{row}"
        containment.cell(row, 14).value = f"=G{row}*J{row}"

    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    wb.save(SOURCE)


if __name__ == "__main__":
    main()
