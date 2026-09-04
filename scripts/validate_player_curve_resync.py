# -*- coding: utf-8 -*-
from __future__ import annotations

import argparse
import hashlib
import json
import math
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.xml.functions import tostring


ERROR_VALUES = {"#REF!", "#VALUE!", "#DIV/0!", "#NAME?", "#N/A", "#NUM!", "#NULL!"}


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def chart_signature(ws) -> list[dict]:
    result = []
    for chart in ws._charts:
        anchor = chart.anchor
        if hasattr(anchor, "_from"):
            anchor_value = (
                anchor._from.col, anchor._from.row, anchor._from.colOff, anchor._from.rowOff,
                anchor.to.col, anchor.to.row, anchor.to.colOff, anchor.to.rowOff,
            )
        else:
            anchor_value = str(anchor)
        result.append(
            {
                "type": type(chart).__name__,
                "anchor": anchor_value,
                "xml": hashlib.sha256(tostring(chart.to_tree())).hexdigest(),
            }
        )
    return result


def dimension_signature(ws) -> dict:
    columns = {
        key: (value.width, value.hidden, value.outlineLevel, value.collapsed, value.style)
        for key, value in ws.column_dimensions.items()
    }
    rows = {
        str(key): (value.height, value.hidden, value.outlineLevel, value.collapsed, value.style)
        for key, value in ws.row_dimensions.items()
    }
    return {
        "max_row": ws.max_row,
        "max_column": ws.max_column,
        "freeze_panes": str(ws.freeze_panes),
        "merged": sorted(str(item) for item in ws.merged_cells.ranges),
        "columns": columns,
        "rows": rows,
        "charts": chart_signature(ws),
        "images": len(ws._images),
        "auto_filter": ws.auto_filter.ref,
    }


def compare_workbooks(before_path: Path, after_path: Path, allowed: dict[str, set[tuple[int, int]]]) -> dict:
    before = load_workbook(before_path, data_only=False, read_only=False)
    after = load_workbook(after_path, data_only=False, read_only=False)
    errors = []
    content_changes = []
    style_changes = []
    if before.sheetnames != after.sheetnames:
        errors.append(f"页签顺序变化：{before.sheetnames} -> {after.sheetnames}")
    for name in before.sheetnames:
        if name not in after.sheetnames:
            continue
        bws, aws = before[name], after[name]
        if bws.sheet_state != aws.sheet_state:
            errors.append(f"{name} 可见状态变化：{bws.sheet_state}->{aws.sheet_state}")
        if dimension_signature(bws) != dimension_signature(aws):
            errors.append(f"{name} 行列/合并/冻结/图表/图片结构变化")
        max_row = max(bws.max_row, aws.max_row)
        max_col = max(bws.max_column, aws.max_column)
        sheet_allowed = allowed.get(name, set())
        for row in range(1, max_row + 1):
            for col in range(1, max_col + 1):
                before_cell = bws.cell(row, col)
                after_cell = aws.cell(row, col)
                if before_cell.value != after_cell.value:
                    item = {
                        "sheet": name,
                        "cell": after_cell.coordinate,
                        "before": before_cell.value,
                        "after": after_cell.value,
                    }
                    if (row, col) not in sheet_allowed:
                        content_changes.append(item)
                before_style = tuple(before_cell._style) if before_cell._style is not None else (0,) * 9
                after_style = tuple(after_cell._style) if after_cell._style is not None else (0,) * 9
                if before_style != after_style:
                    style_changes.append({"sheet": name, "cell": after_cell.coordinate})
    before.close()
    after.close()
    if content_changes:
        errors.append(f"存在{len(content_changes)}个非授权内容变化")
    if style_changes:
        errors.append(f"存在{len(style_changes)}个单元格样式变化")
    return {
        "status": "PASS" if not errors else "FAIL",
        "errors": errors,
        "unexpected_content_changes": content_changes[:20],
        "style_changes": style_changes[:20],
    }


def scan_errors_and_links(path: Path) -> dict:
    formulas = load_workbook(path, data_only=False, read_only=False)
    values = load_workbook(path, data_only=True, read_only=False)
    errors = []
    external_formulas = []
    formula_count = 0
    for ws in formulas.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.data_type == "f":
                    formula_count += 1
                    if "[" in str(cell.value) and "]" in str(cell.value):
                        external_formulas.append(f"{ws.title}!{cell.coordinate}:{cell.value}")
    for ws in values.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.data_type == "e" or cell.value in ERROR_VALUES:
                    errors.append(f"{ws.title}!{cell.coordinate}:{cell.value}")
    external_links = len(getattr(formulas, "_external_links", []))
    formulas.close()
    values.close()
    return {
        "status": "PASS" if not errors and not external_formulas and external_links == 0 else "FAIL",
        "formula_count": formula_count,
        "formula_errors": errors[:50],
        "external_formula_count": len(external_formulas),
        "external_formulas": external_formulas[:20],
        "external_link_records": external_links,
    }


def validate_difficulty(plan: dict, path: Path) -> dict:
    formula_wb = load_workbook(path, data_only=False, read_only=False)
    value_wb = load_workbook(path, data_only=True, read_only=False)
    errors = []

    curve = value_wb["_玩家曲线"]
    for item in plan["curve_rows"]:
        row = item["target_row"]
        for index, expected in enumerate(item["values"], 1):
            actual = curve.cell(row, index).value
            if isinstance(expected, (int, float)) and isinstance(actual, (int, float)):
                if not math.isclose(float(expected), float(actual), rel_tol=1e-10, abs_tol=1e-7):
                    errors.append(f"_玩家曲线!{curve.cell(row, index).coordinate} {actual}!={expected}")
            elif expected != actual:
                errors.append(f"_玩家曲线!{curve.cell(row, index).coordinate} {actual}!={expected}")

    summary_actual = {}
    for sheet, summary in plan["summary"].items():
        ws = value_wb[sheet]
        rows = []
        for row in range(5, ws.max_row + 1):
            if ws.cell(row, 5).value not in (None, "") and ws.cell(row, 10).value not in (None, ""):
                rows.append(row)
        levels = [int(ws.cell(row, 10).value) for row in rows]
        summary_actual[sheet] = {
            "count": len(levels),
            "range": [min(levels), max(levels)],
            "levels": levels if len(levels) <= 35 else None,
            "monotonic": all(levels[index] <= levels[index + 1] for index in range(len(levels) - 1)),
        }
        if len(levels) != summary["count"] or [min(levels), max(levels)] != summary["new_range"]:
            errors.append(f"{sheet} 数量或范围不一致")
        expected_levels = summary["new_levels"]
        if expected_levels is not None and levels != expected_levels:
            errors.append(f"{sheet} 等级序列不一致")
        if not summary_actual[sheet]["monotonic"]:
            errors.append(f"{sheet} 等级非单调")

    resource_levels = [summary_actual[sheet]["levels"] for sheet in ("角色经验", "金币", "收容经验")]
    if not (resource_levels[0] == resource_levels[1] == resource_levels[2]):
        errors.append("三个资源副本等级序列不一致")

    mapping_formula_ws = formula_wb["映射"]
    for row in range(6, 34):
        formula = mapping_formula_ws.cell(row, 10).value
        if not (isinstance(formula, str) and formula.startswith("=ROUND($J$5+")):
            errors.append(f"映射!J{row} 严格线性公式异常：{formula}")

    hard_first = value_wb["困难主线"].cell(5, 10).value
    main_ws = value_wb["主线"]
    main_515 = next(main_ws.cell(row, 10).value for row in range(5, main_ws.max_row + 1) if main_ws.cell(row, 4).value == "5-15")
    if hard_first - main_515 != 2:
        errors.append(f"困难主线首档与主线5-15差值应为2，实际{hard_first}-{main_515}")

    formula_wb.close()
    value_wb.close()
    return {
        "status": "PASS" if not errors else "FAIL",
        "errors": errors[:50],
        "summary": summary_actual,
        "hard_first": hard_first,
        "main_5_15": main_515,
    }


def validate_level(plan: dict, path: Path) -> dict:
    wb = load_workbook(path, data_only=False, read_only=False)
    ws = wb[wb.sheetnames[0]]
    errors = []
    by_play = {}
    for item in plan["level_writes"]:
        actual = ws.cell(item["row"], item["column"]).value
        if actual != item["new_grade"]:
            errors.append(f"LevelId {item['level_id']} Grade {actual}!={item['new_grade']}")
        by_play.setdefault(item["play"], []).append(actual)
    wb.close()
    return {
        "status": "PASS" if not errors and len(plan["level_writes"]) == 526 else "FAIL",
        "checked": len(plan["level_writes"]),
        "errors": errors[:50],
        "by_play": {
            play: {"count": len(values), "range": [min(values), max(values)]}
            for play, values in sorted(by_play.items())
        },
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--plan", type=Path, required=True)
    parser.add_argument("--apply-report", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()

    plan = json.loads(args.plan.read_text(encoding="utf-8"))
    apply_report = json.loads(args.apply_report.read_text(encoding="utf-8"))
    difficulty_path = Path(plan["source"]["difficulty_path"])
    level_path = Path(plan["source"]["level_path"])
    rogue_path = Path(plan["source"]["roguelike_path_readonly"])
    backup_dir = Path(apply_report["backup_dir"])

    difficulty_allowed = {"_玩家曲线": {(row, col) for row in range(2, 281) for col in range(1, 10)}}
    for item in plan["source_writes"]:
        difficulty_allowed.setdefault(item["sheet"], set()).add((item["row"], item["column"]))
    level_allowed = {}
    for item in plan["level_writes"]:
        level_allowed.setdefault(item["sheet"], set()).add((item["row"], item["column"]))

    checks = {
        "difficulty_content_and_layout": compare_workbooks(
            backup_dir / "全玩法难度填写表.xlsx", difficulty_path, difficulty_allowed
        ),
        "level_content_and_layout": compare_workbooks(
            backup_dir / "level.xlsx", level_path, level_allowed
        ),
        "difficulty_formula_errors_and_links": scan_errors_and_links(difficulty_path),
        "level_formula_errors_and_links": scan_errors_and_links(level_path),
        "difficulty_values": validate_difficulty(plan, difficulty_path),
        "level_values": validate_level(plan, level_path),
        "roguelike_unchanged": {
            "status": "PASS" if sha256(rogue_path) == plan["source"]["roguelike_sha256"] else "FAIL",
            "sha256": sha256(rogue_path),
        },
    }
    status = "PASS" if all(item["status"] == "PASS" for item in checks.values()) else "FAIL"
    report = {
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "status": status,
        "plan": str(args.plan.resolve()),
        "apply_report": str(args.apply_report.resolve()),
        "checks": checks,
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    if status != "PASS":
        raise SystemExit(1)


if __name__ == "__main__":
    main()
