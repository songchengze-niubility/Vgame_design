# -*- coding: utf-8 -*-
"""在 01_源表/养成属性分配.xlsx 的「收容物属性分配」页签补 S 级 13 个收容物（CHG-005 首期+后续）。
S 级稀有度=SSR（品质系数1.0）；职业细分与专武比留空待策划填写；二级属性名来自 ExclWeapon.AttrNameList。
"""
import openpyxl
from copy import copy

import vgame_paths

SRC = str(vgame_paths.source_docs_root() / "框架v2" / "01_源表" / "养成属性分配.xlsx")

# 名称 / 职业大类(ExclWeapon.HeroClass) / 二级属性名(页签命名)
S_LEVEL = [
    (51001, "虚妄之镜",  "坦克",  "生命百分比"),
    (51002, "玉液葫芦",  "输出",  "技能伤害加成"),
    (51003, "海妖",      "输出",  "必杀伤害加成"),
    (51004, "回声",      "辅助",  "攻击百分比"),
    (51005, "亚瑟之手",  "坦克",  "攻击百分比"),
    (51006, "象",        "输出",  "技能伤害加成"),
    (51007, "狂乱之牙",  "输出",  "攻击百分比"),
    (51008, "陈旧的烟匣", "坦克", "生命百分比"),
    (51009, "断事官印章", "输出",  "攻击百分比"),
    (51010, "十字剑章",  "坦克",  "技能伤害加成"),
    (51011, "紫罗兰王冠", "辅助",  "治疗率"),
    (51012, "破烂玩偶",  "输出",  "攻击百分比"),
    (51013, "六芒圣楔",  "输出",  "普攻伤害加成"),
]

wb = openpyxl.load_workbook(SRC, data_only=False)
ws = wb["收容物属性分配"]
assert ws["A3"].value == "收容物id" and ws["B3"].value == "武将名称", "表头位置异常，停止写入"

# 捕获原第4行（A级延爆手环）样式
styles = [copy(ws.cell(4, c)._style) for c in range(1, 16)]

# 在第4行前插入 13 行
ws.insert_rows(4, len(S_LEVEL))

for i, (wid, name, job, sec) in enumerate(S_LEVEL):
    r = 4 + i
    ws.cell(r, 1).value = wid          # 收容物id
    ws.cell(r, 2).value = name         # 武将名称
    ws.cell(r, 3).value = "SSR"        # 稀有度（SSR=1.0）
    ws.cell(r, 4).value = None         # 武将职业：细分待策划填
    for c in range(5, 11):             # E~J 职业比/专武比：待填
        ws.cell(r, c).value = None
    ws.cell(r, 11).value = 1.0         # K 品质系数（SSR）
    for c in range(12, 15):            # L~N 最终比：公式待补
        ws.cell(r, c).value = None
    ws.cell(r, 15).value = sec         # O 二级属性名
    for c in range(1, 16):
        ws.cell(r, c)._style = copy(styles[c - 1])

wb.save(SRC)
print("已写入", len(S_LEVEL), "行 S 级")

# 验证：原 A/B 行公式行号是否被正确平移
wb2 = openpyxl.load_workbook(SRC, data_only=False)
ws2 = wb2["收容物属性分配"]
print("第4行(S级首行):", [ws2.cell(4, c).value for c in range(1, 16)])
print("第17行(原A级延爆手环) E4公式:", ws2["E17"].value)
print("第17行 L公式:", ws2["L17"].value)
print("第18行(原A级紊乱透镜) L公式:", ws2["L18"].value)
print("第31行(原B级防辐射贴纸) K公式:", ws2["K31"].value)
print("第42行(技能标签):", ws2["A42"].value)
