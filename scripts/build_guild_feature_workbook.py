from __future__ import annotations

import argparse
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation


NAVY = "1F4E78"
BLUE = "D9EAF7"
LIGHT_BLUE = "DDEBF7"
GREEN = "E2F0D9"
YELLOW = "FFF2CC"
ORANGE = "FCE4D6"
GRAY = "E7E6E6"
WHITE = "FFFFFF"
RED = "F4CCCC"
INPUT_BLUE = "0000FF"

THIN = Side(style="thin", color="B7B7B7")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def style_title(ws, cell_range: str, title: str, color: str = NAVY) -> None:
    ws.merge_cells(cell_range)
    cell = ws[cell_range.split(":")[0]]
    cell.value = title
    cell.fill = PatternFill("solid", fgColor=color)
    cell.font = Font(color=WHITE, bold=True, size=16)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[cell.row].height = 28


def style_header(ws, row: int, first_col: int, last_col: int, color: str = NAVY) -> None:
    for col in range(first_col, last_col + 1):
        cell = ws.cell(row, col)
        cell.fill = PatternFill("solid", fgColor=color)
        cell.font = Font(color=WHITE, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[row].height = 34


def style_table(ws, start_row: int, end_row: int, first_col: int, last_col: int) -> None:
    for row in range(start_row, end_row + 1):
        for col in range(first_col, last_col + 1):
            cell = ws.cell(row, col)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F7FBFF")


def style_parameter_block(ws, start_row: int, end_row: int) -> None:
    style_header(ws, start_row, 1, 3)
    for row in range(start_row + 1, end_row + 1):
        for col in range(1, 4):
            ws.cell(row, col).border = BORDER
            ws.cell(row, col).alignment = Alignment(vertical="center", wrap_text=True)
        ws.cell(row, 1).fill = PatternFill("solid", fgColor=LIGHT_BLUE)
        ws.cell(row, 1).font = Font(bold=True)


def input_style(cell) -> None:
    cell.fill = PatternFill("solid", fgColor=YELLOW)
    cell.font = Font(color=INPUT_BLUE, bold=True)
    cell.alignment = Alignment(horizontal="center")


def formula_style(cell) -> None:
    cell.fill = PatternFill("solid", fgColor=GREEN)
    cell.font = Font(color="006100", bold=True)
    cell.alignment = Alignment(horizontal="center")


def add_whole_validation(ws, cell_range: str, minimum: int, maximum: int) -> None:
    dv = DataValidation(type="whole", operator="between", formula1=str(minimum), formula2=str(maximum))
    dv.error = f"请输入 {minimum} 到 {maximum} 的整数"
    dv.errorTitle = "输入超出范围"
    dv.prompt = f"允许范围：{minimum}–{maximum}"
    dv.promptTitle = "参数输入"
    dv.showErrorMessage = True
    dv.showInputMessage = True
    ws.add_data_validation(dv)
    dv.add(cell_range)


def add_decimal_validation(ws, cell_range: str, minimum: float, maximum: float) -> None:
    dv = DataValidation(type="decimal", operator="between", formula1=str(minimum), formula2=str(maximum))
    dv.error = f"请输入 {minimum} 到 {maximum} 之间的数值"
    dv.errorTitle = "输入超出范围"
    dv.showErrorMessage = True
    ws.add_data_validation(dv)
    dv.add(cell_range)


def set_widths(ws, widths: dict[str, float]) -> None:
    for col, width in widths.items():
        ws.column_dimensions[col].width = width


def add_note_block(ws, start_row: int, title: str, notes: list[str], last_col: int) -> None:
    ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=last_col)
    title_cell = ws.cell(start_row, 1, title)
    title_cell.fill = PatternFill("solid", fgColor=GRAY)
    title_cell.font = Font(bold=True)
    title_cell.border = BORDER
    for offset, note in enumerate(notes, start=1):
        row = start_row + offset
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)
        cell = ws.cell(row, 1, f"• {note}")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = BORDER
        ws.row_dimensions[row].height = 28


def update_overview(ws) -> None:
    # 修复原表中与“贡献值：集体代币=1:1”相冲突的明显录入值。
    ws["G43"] = 100
    ws["G43"].comment = Comment("v0.1一致性修正：协作活跃贡献值为100，按1:1折算集体代币应为100。", "Codex")
    ws["H38"] = 900
    ws["H38"].comment = Comment("v0.1一致性修正：按当前任务次数汇总，任务贡献最大值为900；1000仍作为个人每日总贡献上限。", "Codex")

    ws["E4"] = "个人与公会资源分离；集体代币用于公会建设、公会补给和公会科技"
    ws["D22"] = "仅管理职位可消耗，用于公会建设、补给和公会科技"
    ws["D23"] = "通过任务、派遣、心愿捐赠获得；用于商店与个人科技"

    ws["B9"] = "新增三功能"
    ws["C9"] = "公会补给：管理层消耗集体代币向符合资格的全体成员发福利；个人/公会科技提供永久成长"
    ws["D9"] = "补给按符合资格成员数计费；两类科技分别消耗个人代币与集体代币"
    ws["E9"] = "详见“公会补给”“个人科技树”“公会科技树”页；数值为v0.1评审稿"
    for col in range(2, 6):
        source = ws.cell(8, col)
        target = ws.cell(9, col)
        target._style = copy(source._style)
        target.alignment = copy(source.alignment)
    ws.row_dimensions[9].height = 46


def build_supply_sheet(wb) -> None:
    name = "公会补给"
    if name in wb.sheetnames:
        del wb[name]
    ws = wb.create_sheet(name)
    ws.sheet_properties.tabColor = "5B9BD5"
    ws.freeze_panes = "A13"
    ws.auto_filter.ref = "A12:K15"
    style_title(ws, "A1:K1", "公会补给（v0.1建议值）")
    ws.merge_cells("A2:K2")
    ws["A2"] = "设计目标：让会长/副会长把集体活跃转化为全员福利，同时以人数计费避免大公会获得固定成本套利。"
    ws["A2"].alignment = Alignment(wrap_text=True)
    ws["A2"].font = Font(italic=True, color="666666")

    params = [
        ("参数", "建议值", "口径说明"),
        ("符合资格成员数", 30, "演算输入；实际激活时由服务器快照"),
        ("全公会每周激活上限", 2, "三个档位共用；每档每周最多1次"),
        ("两次激活最短间隔（小时）", 24, "避免短时间连续刷屏与误操作"),
        ("新成员保护期（小时）", 24, "入会满24小时且激活时仍在会，才进入发放名单"),
        ("奖励邮件有效期（小时）", 48, "逾期不补发"),
        ("活跃成员日集体代币预算", 1000, "沿用基础页：每日贡献上限1000，贡献与集体代币1:1"),
        ("公会周集体代币毛产出", "=B4*B9*7", "符合资格成员数 × 日预算 × 7"),
    ]
    for r, row in enumerate(params, start=3):
        for c, value in enumerate(row, start=1):
            ws.cell(r, c, value)
    style_parameter_block(ws, 3, 10)
    for ref in ["B4", "B5", "B6", "B7", "B8", "B9"]:
        input_style(ws[ref])
    formula_style(ws["B10"])
    ws["B10"].number_format = "#,##0"
    add_whole_validation(ws, "B4", 1, 53)
    add_whole_validation(ws, "B5", 1, 3)
    add_whole_validation(ws, "B6:B8", 1, 168)
    add_whole_validation(ws, "B9", 1, 100000)
    ws["B9"].comment = Comment("来源：基础页每日贡献上限1000，且贡献值与集体代币1:1。这里作为高活跃预算上界，不代表所有成员必然拉满。", "Codex")

    headers = ["补给ID", "档位", "公会等级", "操作权限", "人均集体代币成本", "计费人数", "本次总成本", "单档周限", "每名成员奖励", "发放方式", "占周毛产出"]
    for c, value in enumerate(headers, start=1):
        ws.cell(12, c, value)
    style_header(ws, 12, 1, 11)
    rows = [
        (1, "基础补给", 1, "会长/副会长", 300, "=$B$4", "=E13*F13", 1, "金币×50,000；角色经验高(10,000)×1；专武经验高(20,000)×1", "按资格快照发邮件", "=G13/$B$10"),
        (2, "进阶补给", 4, "会长/副会长", 600, "=$B$4", "=E14*F14", 1, "金币×100,000；角色经验高(10,000)×2；专武经验高(20,000)×2", "按资格快照发邮件", "=G14/$B$10"),
        (3, "豪华补给", 7, "会长/副会长", 1000, "=$B$4", "=E15*F15", 1, "金币×200,000；角色经验高(10,000)×3；专武经验高(20,000)×3", "按资格快照发邮件", "=G15/$B$10"),
    ]
    for r, row in enumerate(rows, start=13):
        for c, value in enumerate(row, start=1):
            ws.cell(r, c, value)
    style_table(ws, 13, 15, 1, 11)
    for row in range(13, 16):
        ws.cell(row, 5).number_format = "#,##0"
        ws.cell(row, 6).number_format = "0"
        ws.cell(row, 7).number_format = "#,##0"
        ws.cell(row, 11).number_format = "0.0%"
        formula_style(ws.cell(row, 7))
        formula_style(ws.cell(row, 11))
        ws.row_dimensions[row].height = 50

    add_note_block(
        ws,
        18,
        "结算与异常规则",
        [
            "资格快照：激活时已入会满24小时、仍为正式成员；激活后加入者不追发，激活后退出者仍保留该次邮件。",
            "扣费顺序：先生成资格快照并计算总成本，再原子扣除集体代币；余额不足则整次失败，不允许部分发放。",
            "周规则：每周一05:00重置；全公会共2次、同档最多1次；两次激活至少间隔24小时。",
            "成本折扣：若受公会科技影响，先计算人均成本折扣并向上取整，再乘计费人数。",
            "配置落地：现有GuildSupplies可复用商品/价格/周限骨架；还需支持MinGuildLv、MemberUnitCost、RewardDropId、GuildWeeklyLimit、MinJoinHours、MailExpireHours。",
        ],
        11,
    )
    set_widths(ws, {"A": 10, "B": 13, "C": 10, "D": 15, "E": 18, "F": 11, "G": 14, "H": 10, "I": 54, "J": 18, "K": 14})
    ws.print_area = "A1:K23"
    ws.sheet_view.showGridLines = False


def build_personal_tech_sheet(wb) -> None:
    name = "个人科技树"
    if name in wb.sheetnames:
        del wb[name]
    ws = wb.create_sheet(name)
    ws.sheet_properties.tabColor = "70AD47"
    ws.freeze_panes = "A12"
    ws.auto_filter.ref = "A11:Q21"
    style_title(ws, "A1:Q1", "个人科技树（v0.1建议值）", color="548235")
    ws.merge_cells("A2:Q2")
    ws["A2"] = "设计目标：消耗个人代币形成可控的长期个人成长；不直接提高个人代币产出，避免自我复利。"
    ws["A2"].alignment = Alignment(wrap_text=True)
    ws["A2"].font = Font(italic=True, color="666666")

    params = [
        ("参数", "建议值", "口径说明"),
        ("个人代币日毛产出预算", 1000, "任务/派遣/心愿合计的建模上界；上线前需按实测活跃分布校准"),
        ("投入科技比例", 0.5, "其余50%预留公会商店消费"),
        ("每日科技预算", "=B4*B5", "日毛产出 × 投入比例"),
        ("全树总成本", "=SUM(M12:M21)", "10节点全部升至3级"),
        ("预计毕业天数", "=ROUNDUP(B7/B6,0)", "未计研究队列等待，仅按代币预算"),
        ("研究队列数", 1, "同一玩家同时只能研究1项"),
        ("重置规则", "首版不开放", "避免频繁切换最优解；后续可增加限次重置"),
    ]
    for r, row in enumerate(params, start=3):
        for c, value in enumerate(row, start=1):
            ws.cell(r, c, value)
    style_parameter_block(ws, 3, 10)
    for ref in ["B4", "B5", "B9"]:
        input_style(ws[ref])
    for ref in ["B6", "B7", "B8"]:
        formula_style(ws[ref])
    ws["B5"].number_format = "0%"
    ws["B6"].number_format = "#,##0"
    ws["B7"].number_format = "#,##0"
    ws["B8"].number_format = "0"
    add_whole_validation(ws, "B4", 1, 100000)
    add_decimal_validation(ws, "B5", 0.1, 1)
    add_whole_validation(ws, "B9", 1, 2)

    headers = ["科技ID", "分支", "节点名称", "解锁条件", "前置节点", "效果类型", "Lv1累计效果", "Lv2累计效果", "Lv3累计效果", "Lv1成本", "Lv2成本", "Lv3成本", "总成本", "Lv1时长(h)", "Lv2时长(h)", "Lv3时长(h)", "生效范围/说明"]
    for c, value in enumerate(headers, start=1):
        ws.cell(11, c, value)
    style_header(ws, 11, 1, 17, color="548235")
    rows = [
        (201, "战术强化", "攻击校准", "加入公会", "无", "角色攻击%", 0.005, 0.010, 0.015, 500, 1000, 2000, None, 2, 4, 8, "PvE默认生效；PVPVE/竞技玩法待确认"),
        (202, "战术强化", "体能强化", "加入公会", "201 Lv1", "角色生命%", 0.008, 0.016, 0.024, 500, 1000, 2000, None, 2, 4, 8, "PvE默认生效；PVPVE/竞技玩法待确认"),
        (203, "战术强化", "防护校准", "加入公会", "201 Lv1", "角色防御%", 0.008, 0.016, 0.024, 500, 1000, 2000, None, 2, 4, 8, "PvE默认生效；PVPVE/竞技玩法待确认"),
        (204, "战术强化", "精准打击", "加入公会", "202、203 Lv2", "暴击率百分点", 0.002, 0.004, 0.006, 1000, 2000, 4000, None, 4, 8, 16, "加法叠加；建议竞技玩法暂不生效"),
        (205, "挂机效率", "自动采集", "加入公会", "无", "挂机金币产出%", 0.03, 0.06, 0.10, 500, 1000, 2000, None, 2, 4, 8, "仅放大基础挂机金币，不放大活动/邮件奖励"),
        (206, "挂机效率", "数据整理", "加入公会", "205 Lv1", "挂机角色经验%", 0.03, 0.06, 0.10, 500, 1000, 2000, None, 2, 4, 8, "仅放大基础挂机角色经验"),
        (207, "挂机效率", "装备维护", "加入公会", "205 Lv1", "挂机专武经验%", 0.03, 0.06, 0.10, 500, 1000, 2000, None, 2, 4, 8, "仅放大基础挂机专武经验"),
        (208, "挂机效率", "延时缓存", "加入公会", "206、207 Lv2", "挂机储存上限(h)", 1, 2, 4, 1000, 2000, 4000, None, 4, 8, 16, "加到原有挂机储存时长上限"),
        (209, "派遣效率", "流程优化", "加入公会", "无", "派遣时长%", -0.02, -0.04, -0.06, 500, 1000, 2000, None, 2, 4, 8, "结算时长向上取整到分钟"),
        (210, "派遣效率", "额外备货", "加入公会", "209 Lv2", "派遣额外奖励%", 0.02, 0.04, 0.06, 1000, 2000, 4000, None, 4, 8, 16, "仅放大额外材料奖励，不放大个人/集体代币"),
    ]
    for r, row in enumerate(rows, start=12):
        for c, value in enumerate(row, start=1):
            ws.cell(r, c, value)
        ws.cell(r, 13, f"=SUM(J{r}:L{r})")
    style_table(ws, 12, 21, 1, 17)
    for row in range(12, 22):
        for col in range(7, 10):
            ws.cell(row, col).number_format = "0.0%" if row != 19 else '0"h"'
        for col in range(10, 14):
            ws.cell(row, col).number_format = "#,##0"
        formula_style(ws.cell(row, 13))
        ws.row_dimensions[row].height = 42

    add_note_block(
        ws,
        24,
        "科技通用规则",
        [
            "效果值均为该等级累计值，不是每级增量；同类效果默认加法叠加，再进入对应系统的最终结算。",
            "个人科技永久绑定玩家；退出或更换公会不清除、不失效，避免退会造成已消费资产损失。",
            "研究开始时一次性扣除个人代币；取消研究不返还；研究完成后效果生效。",
            "PVPVE/竞技玩法是否应用战斗属性加成是上线前阻断项，当前工作稿按PvE默认生效处理。",
            "现有GuildTech可复用Personal路由骨架；建议UpgradeCost、ResearchHours按等级数组配置，并补齐EffectId/BuffId。",
        ],
        17,
    )
    set_widths(ws, {"A": 10, "B": 12, "C": 14, "D": 13, "E": 16, "F": 18, "G": 14, "H": 14, "I": 14, "J": 11, "K": 11, "L": 11, "M": 12, "N": 12, "O": 12, "P": 12, "Q": 42})
    ws.print_area = "A1:Q29"
    ws.sheet_view.showGridLines = False


def build_guild_tech_sheet(wb) -> None:
    name = "公会科技树"
    if name in wb.sheetnames:
        del wb[name]
    ws = wb.create_sheet(name)
    ws.sheet_properties.tabColor = "ED7D31"
    ws.freeze_panes = "A15"
    ws.auto_filter.ref = "A14:Q24"
    style_title(ws, "A1:Q1", "公会科技树（v0.1建议值）", color="C65911")
    ws.merge_cells("A2:Q2")
    ws["A2"] = "设计目标：由管理层消耗集体代币建立公会长期成长；在标准30人活跃公会中，全树约2个月完成。"
    ws["A2"].alignment = Alignment(wrap_text=True)
    ws["A2"].font = Font(italic=True, color="666666")

    params = [
        ("参数", "建议值", "口径说明"),
        ("活跃成员数", 30, "演算输入"),
        ("每名活跃成员日集体代币", 1000, "每日贡献上限1000，贡献与集体代币1:1"),
        ("周集体代币毛产出", "=B4*B5*7", "活跃成员 × 日产出 × 7"),
        ("计划补给人均成本", 600, "默认按进阶补给演算"),
        ("计划每周补给次数", 2, "公会补给全局周限"),
        ("每周补给预算", "=B4*B7*B8", "活跃成员 × 人均成本 × 次数"),
        ("每周可投科技预算", "=B6-B9", "暂未扣除其它公会建设支出"),
        ("全树总成本", "=SUM(M15:M24)", "10节点全部升至3级"),
        ("预计毕业周数", "=ROUNDUP(B11/B10,1)", "仅按代币预算，未计研究队列等待"),
        ("预计毕业天数", "=ROUNDUP(B11/B10*7,0)", "仅按代币预算"),
    ]
    for r, row in enumerate(params, start=3):
        for c, value in enumerate(row, start=1):
            ws.cell(r, c, value)
    style_parameter_block(ws, 3, 13)
    for ref in ["B4", "B5", "B7", "B8"]:
        input_style(ws[ref])
    for ref in ["B6", "B9", "B10", "B11", "B12", "B13"]:
        formula_style(ws[ref])
    for ref in ["B5", "B6", "B7", "B9", "B10", "B11"]:
        ws[ref].number_format = "#,##0"
    ws["B12"].number_format = "0.0"
    ws["B13"].number_format = "0"
    add_whole_validation(ws, "B4", 1, 53)
    add_whole_validation(ws, "B5", 1, 100000)
    add_whole_validation(ws, "B7", 1, 100000)
    add_whole_validation(ws, "B8", 0, 3)

    headers = ["科技ID", "分支", "节点名称", "解锁公会等级", "前置节点", "效果类型", "Lv1累计效果", "Lv2累计效果", "Lv3累计效果", "Lv1成本", "Lv2成本", "Lv3成本", "总成本", "Lv1时长(h)", "Lv2时长(h)", "Lv3时长(h)", "结算说明"]
    for c, value in enumerate(headers, start=1):
        ws.cell(14, c, value)
    style_header(ws, 14, 1, 17, color="C65911")
    rows = [
        (101, "运营研发", "研发提速", 1, "无", "公会科技研究时长%", -0.05, -0.10, -0.15, 10000, 20000, 40000, None, 6, 12, 24, "不影响自身当前等级；后续研究开始时读取"),
        (102, "运营研发", "任务协同", 2, "101 Lv1", "公会任务贡献值%", 0.02, 0.04, 0.06, 10000, 20000, 40000, None, 6, 12, 24, "仅放大贡献值，并按1:1继续生成公会经验/集体代币"),
        (103, "运营研发", "委托统筹", 3, "102 Lv1", "公会派遣贡献值%", 0.02, 0.04, 0.06, 10000, 20000, 40000, None, 6, 12, 24, "不放大派遣个人代币和额外奖励"),
        (104, "福利保障", "补给统筹", 4, "101 Lv2", "补给人均成本%", -0.03, -0.06, -0.10, 20000, 40000, 80000, None, 12, 24, 48, "折后人均成本向上取整，再乘资格人数"),
        (105, "福利保障", "补给扩容", 5, "104 Lv1", "补给基础材料数量%", 0.02, 0.04, 0.06, 20000, 40000, 80000, None, 12, 24, 48, "只放大金币/经验材料；不放大抽卡券、元晶等高价值资源"),
        (106, "成长协作", "经验积累", 5, "102 Lv2", "公会经验获取%", 0.03, 0.06, 0.10, 20000, 40000, 80000, None, 12, 24, 48, "独立放大公会经验，不额外生成集体代币"),
        (107, "成长协作", "心愿互助", 6, "103 Lv2", "心愿捐赠个人代币%", 0.02, 0.04, 0.06, 20000, 40000, 80000, None, 12, 24, 48, "单次奖励向下取整；受每日心愿捐赠次数限制"),
        (108, "商业发展", "商店议价", 7, "107 Lv1", "公会商店个人代币价格%", -0.02, -0.04, -0.06, 30000, 60000, 120000, None, 18, 36, 72, "折后价格向上取整；不影响其它商店"),
        (109, "商业发展", "拍卖扩容", 8, "106 Lv2", "每周拍卖掉落上限", 2, 4, 6, 30000, 60000, 120000, None, 18, 36, 72, "加到公会等级提供的周上限，仍受系统硬上限约束"),
        (110, "组织治理", "组织扩编", 9, "108、109 Lv2", "公会成员上限", 1, 2, 3, 40000, 80000, 160000, None, 24, 48, 96, "加到公会等级人数上限；当前建议硬上限53人"),
    ]
    for r, row in enumerate(rows, start=15):
        for c, value in enumerate(row, start=1):
            ws.cell(r, c, value)
        ws.cell(r, 13, f"=SUM(J{r}:L{r})")
    style_table(ws, 15, 24, 1, 17)
    for row in range(15, 25):
        if row <= 22:
            for col in range(7, 10):
                ws.cell(row, col).number_format = "0%"
        else:
            for col in range(7, 10):
                ws.cell(row, col).number_format = '0"点"'
        for col in range(10, 14):
            ws.cell(row, col).number_format = "#,##0"
        formula_style(ws.cell(row, 13))
        ws.row_dimensions[row].height = 46

    add_note_block(
        ws,
        27,
        "科技通用规则",
        [
            "仅会长/副会长可开始研究；同一公会只有1条研究队列；开始时一次性扣除集体代币。",
            "科技状态绑定公会；成员加入后立即获得已完成效果，退出后立即失去，避免跨公会叠加。",
            "效果值均为该等级累计值；同类效果默认加法叠加，再应用系统硬上限与取整规则。",
            "当前演算未扣除公会建设等其它集体代币支出；若建设同步上线，应先设周预算占比或会长消费提醒。",
            "现有GuildTech可复用Collective路由、前置节点、成本与研究时长字段；EffectId/BuffId需补齐并逐项验证消费链。",
        ],
        17,
    )
    ws.conditional_formatting.add("B10", CellIsRule(operator="lessThan", formula=["0"], fill=PatternFill("solid", fgColor=RED)))
    set_widths(ws, {"A": 10, "B": 12, "C": 14, "D": 14, "E": 16, "F": 21, "G": 14, "H": 14, "I": 14, "J": 12, "K": 12, "L": 12, "M": 13, "N": 12, "O": 12, "P": 12, "Q": 46})
    ws.print_area = "A1:Q32"
    ws.sheet_view.showGridLines = False


def main() -> None:
    parser = argparse.ArgumentParser(description="基于公会系统数值表生成三功能扩展版工作簿")
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()

    args.output.parent.mkdir(parents=True, exist_ok=True)
    wb = load_workbook(args.source)
    update_overview(wb["Sheet1"])
    build_supply_sheet(wb)
    build_personal_tech_sheet(wb)
    build_guild_tech_sheet(wb)
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    wb.active = 0
    wb.save(args.output)
    print(args.output)


if __name__ == "__main__":
    main()
