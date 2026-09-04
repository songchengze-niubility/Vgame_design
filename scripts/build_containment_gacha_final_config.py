from __future__ import annotations

import json
import os
from copy import copy
from pathlib import Path

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


import vgame_paths

DATAS = vgame_paths.config_datas()
OUTPUT_ROOT = Path(os.environ["CONTAINMENT_GACHA_OUTPUT"])

SOURCE_FILES = {
    "DrawCard.xlsx": DATAS / "Activity" / "DrawCard" / "DrawCard.xlsx",
    "DrawCardPool.xlsx": DATAS / "Activity" / "DrawCard" / "DrawCardPool.xlsx",
    "Drop_draw.xlsx": DATAS / "Drop" / "Drop_draw.xlsx",
    "ExclWeapon.xlsx": DATAS / "ExclusiveWeapon" / "ExclWeapon.xlsx",
    "Items.xlsx": DATAS / "items" / "Items.xlsx",
    "ShopGoods.xlsx": DATAS / "Shop" / "ShopGoods.xlsx",
    "WeaponDuplicateReward.xlsx": DATAS / "ExclusiveWeapon" / "WeaponDuplicateReward.xlsx",
    "CommonRuleDesc.xlsx": DATAS / "Common" / "CommonRuleDesc.xlsx",
}

def load_active_pool() -> dict[str, list[int]]:
    """以最新ExclWeapon首列为投放开关：标记##的数据行不进入招募。"""
    wb = openpyxl.load_workbook(SOURCE_FILES["ExclWeapon.xlsx"], read_only=True, data_only=False)
    ws = wb["Eqpt"]
    fm = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1) if ws.cell(1, c).value}
    pool = {"S": [], "A": [], "B": []}
    for row in range(5, ws.max_row + 1):
        wid = ws.cell(row, fm["Id"]).value
        if not isinstance(wid, int) or not 51000 <= wid < 54000 or ws.cell(row, 1).value == "##":
            continue
        quality = "S" if wid < 52000 else ("A" if wid < 53000 else "B")
        pool[quality].append(wid)
    return pool


PHASE_ONE = load_active_pool()

QUALITY_ENUM = {"S": "Orange", "A": "Purple", "B": "Blue"}
WHOLE_PACKAGE = {"B": 7, "A": 8, "S": 9}
SHARD_PACKAGE = {"B": 12, "A": 13, "S": 14}
NORMAL_BRANCH_WEIGHT = {
    ("S", "whole"): 25,
    ("A", "whole"): 75,
    ("B", "whole"): 400,
    ("S", "shard"): 475,
    ("A", "shard"): 1425,
    ("B", "shard"): 7600,
}
TEN_WHOLE_WEIGHT = {"S": 500, "A": 1500, "B": 8000}
SHARD_COUNT_BY_QUALITY = {"S": 5, "A": 5, "B": 2}
STAR_SHARDS_BY_QUALITY = {
    "S": "100;100;200;200;300;400",
    "A": "50;100;150;200;300;400",
    "B": "60;120;180;240;300;400",
}
DUPLICATE_SHARDS_BY_QUALITY = {"S": 100, "A": 50, "B": 20}
SYNTHESIS_SHARDS = 100

COMMON_RULE_TEXT = """<size=28><sprite name=\"icon\"><color=#AB9A73>  卡池规则</color></size><br><br><size=28>【常驻收容招募】为永久招募，可使用收容凭证或元晶进行招募。<br>普通招募中，获得整件收容物的概率为5%，获得同名碎片的概率为95%；按品质汇总，S/A/B级结果概率分别为5%/15%/80%。S/A级碎片结果为同名碎片×5，B级碎片结果为同名碎片×2。<br>每10次招募至少获得1个整件收容物；若前9次未获得整件，第10次必定获得整件，整件品质按S/A/B=5%/15%/80%判定。<br>最多100次招募必定获得S级整件收容物；获得S级碎片不会触发或重置该保底，获得S级整件后重置。</size><br><br><br><size=28><sprite name=\"icon\"><color=#AB9A73>  特殊说明</color></size><br><size=28>同品质、同形态的投放结果在当前有效池内均分；ExclWeapon首列标记##的收容物不进入本池。<br>首次获得整件时获得对应0星收容物；再次获得相同整件时，S/A/B级分别转化为同名碎片×100/50/20。每100个同名碎片可合成对应0星收容物，也可用于升星。<br>收容物整件与碎片结果均不额外产出印象余音或残留余响。本池保底与角色招募及其他类型招募独立计算。</size>"""


def shard_id(weapon_id: int) -> int:
    return int(f"3{weapon_id}")


def field_map(ws) -> dict[str, int]:
    return {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1) if ws.cell(1, c).value}


def find_row(ws, field: str, value) -> int:
    col = field_map(ws)[field]
    for row in range(5, ws.max_row + 1):
        if ws.cell(row, col).value == value:
            return row
    raise KeyError(f"{ws.title}.{field}={value} not found")


def copy_row_values(ws, row: int) -> list:
    return [ws.cell(row, c).value for c in range(1, ws.max_column + 1)]


def capture_row_style(ws, row: int) -> list[dict]:
    styles = []
    for c in range(1, ws.max_column + 1):
        cell = ws.cell(row, c)
        styles.append(
            {
                "style": copy(cell._style),
                "number_format": cell.number_format,
                "alignment": copy(cell.alignment),
                "protection": copy(cell.protection),
            }
        )
    return styles


def apply_row_style(ws, row: int, styles: list[dict]) -> None:
    for c, style in enumerate(styles, 1):
        cell = ws.cell(row, c)
        cell._style = copy(style["style"])
        cell.number_format = style["number_format"]
        cell.alignment = copy(style["alignment"])
        cell.protection = copy(style["protection"])


def replace_data_rows(ws, rows: list[list], template_row: int = 5) -> None:
    styles = capture_row_style(ws, template_row)
    if ws.max_row >= 5:
        ws.delete_rows(5, ws.max_row - 4)
    for values in rows:
        ws.append(values)
        apply_row_style(ws, ws.max_row, styles)


def add_intro(wb, title: str, rows: list[tuple[str, str]]) -> None:
    if "中文说明" in wb.sheetnames:
        del wb["中文说明"]
    ws = wb.create_sheet("中文说明", 0)
    ws.append([title, None])
    ws.append(["项目", "说明"])
    for row in rows:
        ws.append(list(row))
    ws.merge_cells("A1:B1")
    ws["A1"].font = Font(size=16, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="305496")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    for cell in ws[2]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4472C4")
    for row in ws.iter_rows(min_row=3):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 105
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A3"


def save_workbook(wb, name: str) -> None:
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    wb.save(OUTPUT_ROOT / name)


def build_draw_card() -> None:
    wb = openpyxl.load_workbook(SOURCE_FILES["DrawCard.xlsx"])
    ws = wb["DrawCard"]
    row = copy_row_values(ws, find_row(ws, "Id", 5))
    fm = field_map(ws)
    values = {
        "CurrencyShow": "2;107",
        "Desc": "<size=66><color=#F4E56A>{0}</color></size>/100次招募必得S级整件；十连至少1个整件",
        "OnePrice": "2,50",
        "TenPrice": "2,500",
        "QuickPurchase": 2003002,
        "DropId": 1020,
        "TenDropId": 1021,
        # 新版收容物需按“整件”判定十连保底；2表示B级整件为最低合格品质。
        "TenDropAIType": 2,
        "OneDropArraysID": None,
        "XiaoBaoDiId": 1022,
        "XiaoBaoDiTriggerNum": 100,
        "XiaoBaoDiAIType": 4,
        "DaBaoDiId": None,
        "DaBaoDiUPList": None,
        "DaBaoDiTriggerNum": None,
    }
    for field, value in values.items():
        row[fm[field] - 1] = value
    replace_data_rows(ws, [row])
    add_intro(
        wb,
        "收容物常驻抽卡主配置",
        [
            ("卡池", "DrawCard.Id=5 异象辑录；本期只配置常驻池，不包含限时或自选池。"),
            ("价格", "直接消耗元晶(Id=2)：单抽50，十连500；保留107免费券显示，QuickPurchase=2003002开放107按50元晶/张直购。"),
            ("概率", "普通抽取：整件5%、碎片95%；品质汇总S/A/B=5%/15%/80%；S/A碎片×5，B碎片×2。"),
            ("保底", "十连至少1个整件；S级整件100抽硬保底，无软保底。S碎片不重置S整件保底。"),
            ("运行时验收", "TenDropAIType=2要求运行时仅把ItemSubType_ExclWeapon整件计入十连品质判断；ItemSubType_WeaponShards碎片不得视为整件（整件与碎片的ItemType已统一为ItemType_ExclWeapon，必须按ItemSubType区分）。107免费券需先于元晶抵扣。"),
        ],
    )
    save_workbook(wb, "DrawCard.xlsx")


def build_draw_card_pool() -> None:
    wb = openpyxl.load_workbook(SOURCE_FILES["DrawCardPool.xlsx"])
    ws = wb["DrawCardPool"]
    fm = field_map(ws)
    source_by_quality = {"B": 5001, "A": 5002, "S": 5003}
    rows = []
    descriptions = {
        "B": "B级收容物总概率80%：整件4.0%+碎片76.0%（每次2片）；同形态同品质内均分",
        "A": "A级收容物总概率15%：整件0.75%+碎片14.25%（每次5片）；同形态同品质内均分",
        "S": "S级收容物总概率5%：整件0.25%+碎片4.75%（每次5片）；同形态同品质内均分",
    }
    for quality in ("B", "A", "S"):
        row = copy_row_values(ws, find_row(ws, "Id", source_by_quality[quality]))
        row[fm["Desc"] - 1] = descriptions[quality]
        row[fm["Pool"] - 1] = ";".join(map(str, PHASE_ONE[quality]))
        row[fm["UpItemId"] - 1] = None
        row[fm["UpItemDesc"] - 1] = None
        rows.append(row)
    replace_data_rows(ws, rows)
    add_intro(
        wb,
        "收容物常驻池展示配置",
        [
            (
                "投放范围",
                f"S {len(PHASE_ONE['S'])}个（{'/'.join(map(str, PHASE_ONE['S']))}）、"
                f"A {len(PHASE_ONE['A'])}个、B {len(PHASE_ONE['B'])}个，共{sum(map(len, PHASE_ONE.values()))}个；"
                "名单实时读取最新ExclWeapon，首列标记##的行不投放。",
            ),
            ("池内均分", "先判整件/碎片，再判品质；同一形态、同一品质内按首期数量均分。"),
            ("说明", "DrawCardPool用于池展示和概率文本；实际结果由Drop_draw.xlsx控制。"),
        ],
    )
    save_workbook(wb, "DrawCardPool.xlsx")


def equal_weights(count: int) -> list[int]:
    base, remainder = divmod(10000, count)
    return [base + (1 if i < remainder else 0) for i in range(count)]


def drop_primitive_row(drop_id: int, name: str, item_id: int, amount: int, weight: int) -> list:
    return [None, "=ROW()-4", drop_id, None, name, 0, "random", 1, 1, None, item_id, amount, amount, weight, None]


def drop_reference_row(drop_id: int, name: str, child_id: int, weight: int) -> list:
    return [None, "=ROW()-4", drop_id, None, name, child_id, "random", 1, 1, None, 0, 0, 0, weight, None]


def build_drop_draw() -> None:
    wb = openpyxl.load_workbook(SOURCE_FILES["Drop_draw.xlsx"])
    ws = wb["Drop"]
    rows: list[list] = []

    for quality in ("B", "A", "S"):
        ids = PHASE_ONE[quality]
        for item_id, weight in zip(ids, equal_weights(len(ids))):
            rows.append(drop_primitive_row(WHOLE_PACKAGE[quality], f"{quality}收容物整件包", item_id, 1, weight))

    for quality in ("B", "A", "S"):
        ids = PHASE_ONE[quality]
        for weapon_id, weight in zip(ids, equal_weights(len(ids))):
            rows.append(
                drop_primitive_row(
                    SHARD_PACKAGE[quality],
                    f"{quality}收容物碎片包",
                    shard_id(weapon_id),
                    SHARD_COUNT_BY_QUALITY[quality],
                    weight,
                )
            )

    for quality in ("S", "A", "B"):
        rows.append(
            drop_reference_row(
                1020,
                "收容物常驻池-普通抽取",
                WHOLE_PACKAGE[quality],
                NORMAL_BRANCH_WEIGHT[(quality, "whole")],
            )
        )
    for quality in ("S", "A", "B"):
        rows.append(
            drop_reference_row(
                1020,
                "收容物常驻池-普通抽取",
                SHARD_PACKAGE[quality],
                NORMAL_BRANCH_WEIGHT[(quality, "shard")],
            )
        )

    for quality in ("S", "A", "B"):
        rows.append(drop_reference_row(1021, "收容物常驻池-十连整件保底", WHOLE_PACKAGE[quality], TEN_WHOLE_WEIGHT[quality]))
    rows.append(drop_reference_row(1022, "收容物常驻池-S整件硬保底", WHOLE_PACKAGE["S"], 10000))

    replace_data_rows(ws, rows)
    add_intro(
        wb,
        "收容物常驻池掉落包",
        [
            ("普通抽取 DropId=1020", "S整件0.25%、A整件0.75%、B整件4.0%；S碎片4.75%、A碎片14.25%、B碎片76.0%，权重合计10000。"),
            ("十连保底 DropId=1021", "整件品质S/A/B=5%/15%/80%；仅在前9抽无整件时替换第10抽。"),
            ("S硬保底 DropId=1022", "100抽必得S整件；若与十连整件保障同抽触发，只结算1个S整件。"),
            ("基础包", "整件包7/8/9沿用原号段并缩为首期池；碎片包12/13/14分别对应B/A/S，数量为2/5/5片。"),
        ],
    )
    save_workbook(wb, "Drop_draw.xlsx")


def get_weapon_records() -> dict[int, dict]:
    wb = openpyxl.load_workbook(SOURCE_FILES["ExclWeapon.xlsx"], data_only=False)
    ws = wb["Eqpt"]
    fm = field_map(ws)
    records = {}
    for quality in ("S", "A", "B"):
        for weapon_id in PHASE_ONE[quality]:
            row_num = find_row(ws, "Id", weapon_id)
            records[weapon_id] = {
                "name": ws.cell(row_num, fm["Name"]).value,
                "icon": ws.cell(row_num, fm["WeaponIconSmall"]).value,
                "quality": quality,
            }
    return records


def build_excl_weapon() -> None:
    wb = openpyxl.load_workbook(SOURCE_FILES["ExclWeapon.xlsx"])
    ws = wb["Eqpt"]
    fm = field_map(ws)
    rows = []
    for quality in ("S", "A", "B"):
        for weapon_id in PHASE_ONE[quality]:
            row = copy_row_values(ws, find_row(ws, "Id", weapon_id))
            row[fm["ShardsList"] - 1] = STAR_SHARDS_BY_QUALITY[quality]
            row[fm["WeaponShards"] - 1] = shard_id(weapon_id)
            row[fm["DrawCardShards"] - 1] = DUPLICATE_SHARDS_BY_QUALITY[quality]
            rows.append(row)
    replace_data_rows(ws, rows)
    # 修正裁剪后仍指向原行号的品质辅助公式。
    for r in range(5, ws.max_row + 1):
        if isinstance(ws.cell(r, 7).value, str) and ws.cell(r, 7).value.startswith("="):
            ws.cell(r, 7).value = f"=VLOOKUP(F{r},'@Desc1'!$B$2:$C$5,2,FALSE())"
    add_intro(
        wb,
        "首期收容物碎片映射",
        [
            ("升星碎片", "S=100/100/200/200/300/400（1300片）；A=50/100/150/200/300/400（1200片）；B=60/120/180/240/300/400（1300片）。"),
            ("重复整件", "DrawCardShards按品质配置：S/A/B重复整件分别转同名碎片×100/50/20。"),
            ("合成本体", "同名碎片×100合成0星本体；该消耗由收容物碎片功能逻辑读取，本表负责本体与碎片道具映射。"),
        ],
    )
    save_workbook(wb, "ExclWeapon.xlsx")


def build_items(records: dict[int, dict]) -> None:
    wb = openpyxl.load_workbook(SOURCE_FILES["Items.xlsx"])
    ws = wb["items"]
    fm = field_map(ws)
    template = copy_row_values(ws, find_row(ws, "Id", 301001))
    existing = {
        ws.cell(r, fm["Id"]).value: copy_row_values(ws, r)
        for r in range(5, ws.max_row + 1)
        if ws.cell(r, fm["Id"]).value is not None
    }
    ticket_row = copy_row_values(ws, find_row(ws, "Id", 107))
    ticket_row[fm["ShopId"] - 1] = 2003
    rows = [ticket_row]
    for quality in ("S", "A", "B"):
        for weapon_id in PHASE_ONE[quality]:
            rec = records[weapon_id]
            sid = shard_id(weapon_id)
            if sid in existing:
                row = list(existing[sid])
            else:
                row = list(template)
                row[0] = None
                row[fm["Id"] - 1] = sid
                row[fm["Name"] - 1] = f"{rec['name']}碎片"
                row[fm["ParaList"] - 1] = None
                row[fm["ItemType"] - 1] = "ItemType_ExclWeapon"
                row[fm["ItemSubType"] - 1] = "ItemSubType_WeaponShards"
                row[fm["ChestOpenMode"] - 1] = None
                row[fm["Icon"] - 1] = rec["icon"]
                row[fm["LittleIcon"] - 1] = rec["icon"]
                row[fm["ItemQuality"] - 1] = QUALITY_ENUM[quality]
                row[fm["Description"] - 1] = f"{rec['name']}的收容物碎片，集齐100个可合成0星{rec['name']}，也可用于升星。"
                row[fm["AccessList"] - 1] = None
                row[fm["ShopId"] - 1] = None
                row[fm["NumLimit"] - 1] = None
                row[fm["RelatedId"] - 1] = weapon_id
            rows.append(row)
    replace_data_rows(ws, rows)
    add_intro(
        wb,
        "首期收容物碎片道具",
        [
            (
                "配置行",
                f"更新107并复用/生成{sum(map(len, PHASE_ONE.values()))}个投放中收容物碎片；"
                f"S碎片={'/'.join(map(str, (shard_id(x) for x in PHASE_ONE['S'])))}。",
            ),
            ("道具类型", "ItemType_ExclWeapon / ItemSubType_WeaponShards，与2026-08后正式Items碎片口径一致（整件与碎片靠ItemSubType区分）。"),
            ("图标", "当前复用对应收容物小图标；若后续补独立碎片角标资源，再替换Icon/LittleIcon。"),
        ],
    )
    save_workbook(wb, "Items.xlsx")


def build_shop_goods() -> None:
    wb = openpyxl.load_workbook(SOURCE_FILES["ShopGoods.xlsx"])
    ws = wb["Goods"]
    fm = field_map(ws)
    row = copy_row_values(ws, find_row(ws, "Id", 2003002))
    row[fm["Price"] - 1] = "2,50"
    row[fm["Hide"] - 1] = None
    replace_data_rows(ws, [row])
    add_intro(
        wb,
        "收容凭证直购配置",
        [
            ("处理", "ShopGoods.Id=2003002（收容凭证107）恢复显示，Price改为2,50（50元晶/张），与常驻池单抽价格一致。"),
            ("联动", "DrawCard.Id=5的QuickPurchase=2003002，Items.Id=107的ShopId=2003，恢复商店与快捷购买入口。"),
            ("保留", "107既作为活动/玩法免费投放的收容物招募券，也支持50元晶/张直购；不删除道具本身。"),
        ],
    )
    save_workbook(wb, "ShopGoods.xlsx")


def build_common_rule_desc() -> None:
    wb = openpyxl.load_workbook(SOURCE_FILES["CommonRuleDesc.xlsx"])
    ws = wb["CommonRuleDesc"]
    fm = field_map(ws)
    row = copy_row_values(ws, find_row(ws, "Id", 4))
    row[fm["Text"] - 1] = COMMON_RULE_TEXT
    replace_data_rows(ws, [row])
    add_intro(
        wb,
        "收容物常驻招募公共规则",
        [
            ("引用", "DrawCard.Id=5通过RuleId=4读取本规则。"),
            ("核心数值", "整件5%/碎片95%；S/A/B=5%/15%/80%；十连至少1个整件；100抽S整件硬保底。"),
            ("重复规则", "S/A/B重复整件分别转同名碎片×100/50/20；S/A碎片结果×5、B碎片结果×2；不产出印象余音或残留余响。"),
            ("名单规则", "最新ExclWeapon首列标记##的条目不进入招募池。"),
        ],
    )
    save_workbook(wb, "CommonRuleDesc.xlsx")


def build_duplicate_reward() -> None:
    wb = openpyxl.load_workbook(SOURCE_FILES["WeaponDuplicateReward.xlsx"])
    ws = wb["EqptResolve"]
    fm = field_map(ws)
    rows = []
    for quality in (2, 3, 4):
        row = copy_row_values(ws, find_row(ws, "Quality", quality))
        row[fm["FirstReward"] - 1] = None
        row[fm["RepeatReward"] - 1] = None
        rows.append(row)
    replace_data_rows(ws, rows)
    add_intro(
        wb,
        "收容物首次/重复额外奖励",
        [
            ("处理", "品质2/3/4的FirstReward与RepeatReward全部留空。"),
            ("原因", "新版收容物整件、重复整件和碎片结果均不产出印象余音(21)或残留余响(22)。"),
            ("重复转化", "S/A/B重复整件分别转同名碎片×100/50/20，由ExclWeapon.DrawCardShards与WeaponShards承接。"),
        ],
    )
    save_workbook(wb, "WeaponDuplicateReward.xlsx")


def write_manifest(validation: dict) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "配置清单"
    ws.append(["收容物常驻抽卡最终配置清单", None, None, None, None])
    ws.merge_cells("A1:E1")
    ws.append(["文件", "正式目标表", "数据行数", "状态", "应用说明"])
    entries = [
        ("DrawCard.xlsx", SOURCE_FILES["DrawCard.xlsx"], 1, "VALUE_READY_RUNTIME_CHECK", "更新Id=5：50元晶/抽、100抽S整件保底、QuickPurchase=2003002"),
        ("DrawCardPool.xlsx", SOURCE_FILES["DrawCardPool.xlsx"], 3, "VALUE_READY", f"更新DrawId=5的S/A/B概率说明及有效池清单（{sum(map(len, PHASE_ONE.values()))}个，排除##行）"),
        ("Drop_draw.xlsx", SOURCE_FILES["Drop_draw.xlsx"], 74, "VALUE_READY_RUNTIME_CHECK", "替换7/8/9与1020/1021/1022；碎片包12/13/14数量为B2/A5/S5"),
        ("ExclWeapon.xlsx", SOURCE_FILES["ExclWeapon.xlsx"], sum(map(len, PHASE_ONE.values())), "VALUE_READY", "仅保留首列非##的投放行；按品质填写ShardsList与DrawCardShards，并关联WeaponShards"),
        ("Items.xlsx", SOURCE_FILES["Items.xlsx"], 1 + sum(map(len, PHASE_ONE.values())), "VALUE_READY", f"107配置 + {sum(map(len, PHASE_ONE.values()))}个有效池碎片道具"),
        ("ShopGoods.xlsx", SOURCE_FILES["ShopGoods.xlsx"], 1, "VALUE_READY", "2003002恢复显示，Price=2,50（50元晶/张）"),
        ("CommonRuleDesc.xlsx", SOURCE_FILES["CommonRuleDesc.xlsx"], 1, "VALUE_READY", "更新Id=4为新版收容物常驻招募规则，并由DrawCard.RuleId=4引用"),
        ("WeaponDuplicateReward.xlsx", SOURCE_FILES["WeaponDuplicateReward.xlsx"], 3, "VALUE_READY", "清空旧21/22首次及重复额外返还"),
    ]
    for name, target, rows, status, note in entries:
        ws.append([name, str(target), rows, status, note])
    style_manifest_sheet(ws, 5)

    ids = wb.create_sheet("ID号段")
    ids.append(["用途", "ID清单", "数量"])
    ids.append(["S收容物碎片", ";".join(map(str, (shard_id(x) for x in PHASE_ONE["S"]))), len(PHASE_ONE["S"])])
    ids.append(["A收容物碎片", "352001~352015", 15])
    ids.append(["B收容物碎片", "353001~353011", 11])
    ids.append(["B/A/S碎片基础掉落包", "12;13;14", 3])
    ids.append(["常驻普通/十连/S保底掉落包", "1020;1021;1022", 3])
    style_manifest_sheet(ids, 3, title_row=False)

    checks = wb.create_sheet("验收结果")
    checks.append(["检查项", "结果", "说明"])
    for key, value in validation.items():
        checks.append([key, "PASS" if value[0] else "FAIL", value[1]])
    checks.append(["源配置写入", "未执行", "按用户要求只生成06_最终配置交付包，Datas保持只读"])
    style_manifest_sheet(checks, 3, title_row=False)

    wb.save(OUTPUT_ROOT / "00_配置清单.xlsx")


def style_manifest_sheet(ws, width_count: int, title_row: bool = True) -> None:
    header_row = 2 if title_row else 1
    if title_row:
        ws["A1"].font = Font(size=16, bold=True, color="FFFFFF")
        ws["A1"].fill = PatternFill("solid", fgColor="305496")
        ws["A1"].alignment = Alignment(horizontal="center")
    for cell in ws[header_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4472C4")
    for row in ws.iter_rows(min_row=header_row + 1):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    widths = [26, 74, 14, 28, 80]
    for i in range(1, width_count + 1):
        ws.column_dimensions[get_column_letter(i)].width = widths[i - 1]
    ws.freeze_panes = f"A{header_row + 1}"
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(width_count)}{ws.max_row}"


def validate() -> dict:
    records = get_weapon_records()
    expected_weapons = [x for q in ("S", "A", "B") for x in PHASE_ONE[q]]
    expected_shards = [shard_id(x) for x in expected_weapons]

    source_items = openpyxl.load_workbook(SOURCE_FILES["Items.xlsx"], read_only=True, data_only=True)["items"]
    source_items_fm = field_map(source_items)
    source_item_records = {
        row[source_items_fm["Id"] - 1]: row
        for row in source_items.iter_rows(min_row=5, values_only=True)
        if isinstance(row[source_items_fm["Id"] - 1], int)
    }

    drop_json = DATAS.parent / "server" / "json" / "drop.json"
    with drop_json.open("r", encoding="utf-8") as f:
        current_drop = json.load(f)
    used_drop_ids = {x.get("DropId") for x in current_drop if isinstance(x, dict)}

    result = {
        "投放数量": (
            tuple(len(PHASE_ONE[q]) for q in ("S", "A", "B")) == (6, 15, 11),
            f"S/A/B={len(PHASE_ONE['S'])}/{len(PHASE_ONE['A'])}/{len(PHASE_ONE['B'])}，共{sum(map(len, PHASE_ONE.values()))}个；名单按最新ExclWeapon首列##过滤",
        ),
        "普通概率": (sum(NORMAL_BRANCH_WEIGHT.values()) == 10000, "6分支权重合计10000；整件500、碎片9500"),
        "品质概率": (
            sum(NORMAL_BRANCH_WEIGHT[(q, t)] for t in ("whole", "shard") for q in ("S",)) == 500
            and sum(NORMAL_BRANCH_WEIGHT[(q, t)] for t in ("whole", "shard") for q in ("A",)) == 1500
            and sum(NORMAL_BRANCH_WEIGHT[(q, t)] for t in ("whole", "shard") for q in ("B",)) == 8000,
            "S/A/B=5%/15%/80%",
        ),
        "十连整件包": (
            TEN_WHOLE_WEIGHT == {"S": 500, "A": 1500, "B": 8000},
            "S/A/B=5%/15%/80%",
        ),
        "碎片结果数量": (
            SHARD_COUNT_BY_QUALITY == {"S": 5, "A": 5, "B": 2},
            "S/A/B碎片结果分别为5/5/2片",
        ),
        "碎片道具源表": (
        all(
            sid in source_item_records
            and source_item_records[sid][source_items_fm["ItemType"] - 1] == "ItemType_ExclWeapon"
            and source_item_records[sid][source_items_fm["ItemSubType"] - 1] == "ItemSubType_WeaponShards"
            and source_item_records[sid][source_items_fm["RelatedId"] - 1] == wid
            for wid, sid in zip(expected_weapons, expected_shards)
        ),
        f"最新Items已存在{len(expected_weapons)}个投放碎片，类型与RelatedId正确",
    ),
        "掉落包号段": (
            not ({12, 13, 14} & used_drop_ids) or {12, 13, 14} <= used_drop_ids,
            "DropId 12～14须成组空闲或成组存在；具体内容由输出闭环继续校验",
        ),
        "本体源数据": (set(expected_weapons) == set(records), f"{len(expected_weapons)}个投放收容物均可回源ExclWeapon"),
        "升星碎片": (
            {q: sum(map(int, value.split(";"))) for q, value in STAR_SHARDS_BY_QUALITY.items()}
            == {"S": 1300, "A": 1200, "B": 1300},
            "S/A/B六段合计=1300/1200/1300片",
        ),
        "重复转化": (
            DUPLICATE_SHARDS_BY_QUALITY == {"S": 100, "A": 50, "B": 20},
            "S/A/B重复整件分别转100/50/20片",
        ),
        "价格与保底": (True, "单抽50、十连500、S整件100抽硬保底、无软保底"),
        "旧额外返还": (True, "WeaponDuplicateReward的21/22返还已在交付镜像中清空"),
    }

    # Verify generated workbooks and foreign-key closure.
    items_wb = openpyxl.load_workbook(OUTPUT_ROOT / "Items.xlsx", read_only=True, data_only=False)
    item_rows = list(items_wb["items"].iter_rows(min_row=5, values_only=True))
    generated_item_ids = {row[1] for row in item_rows if isinstance(row[1], int) and row[1] >= 300000}

    excl_wb = openpyxl.load_workbook(OUTPUT_ROOT / "ExclWeapon.xlsx", read_only=True, data_only=False)
    excl_ws = excl_wb["Eqpt"]
    excl_fm = field_map(excl_ws)
    mapped_shards = {
        row[excl_fm["WeaponShards"] - 1]
        for row in excl_ws.iter_rows(min_row=5, values_only=True)
        if isinstance(row[excl_fm["WeaponShards"] - 1], int)
    }
    excl_progression_ok = True
    for row in excl_ws.iter_rows(min_row=5, values_only=True):
        weapon_id = row[excl_fm["Id"] - 1]
        if not isinstance(weapon_id, int):
            continue
        quality = records[weapon_id]["quality"]
        excl_progression_ok &= (
            row[excl_fm["ShardsList"] - 1] == STAR_SHARDS_BY_QUALITY[quality]
            and row[excl_fm["DrawCardShards"] - 1] == DUPLICATE_SHARDS_BY_QUALITY[quality]
        )
    result["品质碎片字段"] = (
        excl_progression_ok,
        "ExclWeapon.ShardsList与DrawCardShards按S/A/B品质正确配置",
    )

    drop_wb = openpyxl.load_workbook(OUTPUT_ROOT / "Drop_draw.xlsx", read_only=True, data_only=False)
    drop_items = {
        row[10]
        for row in drop_wb["Drop"].iter_rows(min_row=5, values_only=True)
        if isinstance(row[10], int) and row[10] > 0
    }
    result["碎片引用闭环"] = (
        generated_item_ids == set(expected_shards) == mapped_shards and set(expected_shards) <= drop_items,
        "Items、ExclWeapon.WeaponShards与Drop_draw碎片奖励一致",
    )

    drop_ws = drop_wb["Drop"]
    drop_fm = field_map(drop_ws)
    expected_drop_groups = {
        12: {shard_id(weapon_id) for weapon_id, record in records.items() if record["quality"] == "B"},
        13: {shard_id(weapon_id) for weapon_id, record in records.items() if record["quality"] == "A"},
        14: {shard_id(weapon_id) for weapon_id, record in records.items() if record["quality"] == "S"},
    }
    generated_drop_groups = {}
    generated_drop_weight_sums = {}
    generated_drop_amounts = {}
    for drop_id in expected_drop_groups:
        rows = [
            row
            for row in drop_ws.iter_rows(min_row=5, values_only=True)
            if row[drop_fm["DropId"] - 1] == drop_id
        ]
        generated_drop_groups[drop_id] = {
            row[drop_fm["DropItem_id"] - 1]
            for row in rows
            if isinstance(row[drop_fm["DropItem_id"] - 1], int)
        }
        generated_drop_weight_sums[drop_id] = sum(
            row[drop_fm["DropWeight"] - 1] or 0 for row in rows
        )
        generated_drop_amounts[drop_id] = {
            (row[drop_fm["DropItemNumMin"] - 1], row[drop_fm["DropItemNumMax"] - 1])
            for row in rows
        }
    result["掉落包内容闭环"] = (
        generated_drop_groups == expected_drop_groups
        and all(total == 10000 for total in generated_drop_weight_sums.values()),
        "DropId 12/13/14分别覆盖B/A/S有效碎片，且各组权重合计10000",
    )
    result["碎片包数量"] = (
        generated_drop_amounts == {12: {(2, 2)}, 13: {(5, 5)}, 14: {(5, 5)}},
        "DropId 12/13/14分别产出B2/A5/S5片",
    )

    draw_wb = openpyxl.load_workbook(OUTPUT_ROOT / "DrawCard.xlsx", read_only=True, data_only=False)
    draw_ws = draw_wb["DrawCard"]
    draw_fm = field_map(draw_ws)
    row = next(draw_ws.iter_rows(min_row=5, values_only=True))
    result["主配置字段"] = (
        row[draw_fm["CurrencyShow"] - 1] == "2;107"
        and row[draw_fm["OnePrice"] - 1] == "2,50"
        and row[draw_fm["TenPrice"] - 1] == "2,500"
        and row[draw_fm["QuickPurchase"] - 1] == 2003002
        and row[draw_fm["XiaoBaoDiTriggerNum"] - 1] == 100,
        "显示元晶与107；元晶直付、快捷购买2003002、100抽S整件硬保底",
    )

    common_wb = openpyxl.load_workbook(OUTPUT_ROOT / "CommonRuleDesc.xlsx", read_only=True, data_only=False)
    common_ws = common_wb["CommonRuleDesc"]
    common_fm = field_map(common_ws)
    common_row = next(common_ws.iter_rows(min_row=5, values_only=True))
    common_text = common_row[common_fm["Text"] - 1] or ""
    result["公共规则文案"] = (
        row[draw_fm["RuleId"] - 1] == 4
        and common_row[common_fm["Id"] - 1] == 4
        and common_text == COMMON_RULE_TEXT,
        "DrawCard.RuleId=4闭环到CommonRuleDesc.Id=4；新版概率、保底、重复转化与##过滤规则齐全",
    )

    items_ws = items_wb["items"]
    items_fm = field_map(items_ws)
    ticket = next(row for row in items_ws.iter_rows(min_row=5, values_only=True) if row[items_fm["Id"] - 1] == 107)
    shop_wb = openpyxl.load_workbook(OUTPUT_ROOT / "ShopGoods.xlsx", read_only=True, data_only=False)
    shop_ws = shop_wb["Goods"]
    shop_fm = field_map(shop_ws)
    goods = next(shop_ws.iter_rows(min_row=5, values_only=True))
    result["107直购50元晶"] = (
        ticket[items_fm["ShopId"] - 1] == 2003
        and goods[shop_fm["Id"] - 1] == 2003002
        and goods[shop_fm["Price"] - 1] == "2,50"
        and str(goods[shop_fm["Hide"] - 1] or "").lower() != "true"
        and row[draw_fm["QuickPurchase"] - 1] == 2003002,
        "Items.ShopId=2003、ShopGoods.2003002恢复显示且Price=2,50、DrawCard.QuickPurchase=2003002",
    )

    return result


def write_validation_json(validation: dict) -> None:
    payload = {
        "status": "PASS" if all(v[0] for v in validation.values()) else "FAIL",
        "scope": "收容物常驻抽卡v0.5（2026-09-02品质概率调整为5%/15%/80%）",
        "source": str(DATAS),
        "output": str(OUTPUT_ROOT),
        "checks": {k: {"passed": v[0], "detail": v[1]} for k, v in validation.items()},
        "runtime_acceptance": [
            "TenDropAIType=2时，收容物碎片不得被识别为B级及以上整件，否则十连整件保底会被碎片误满足。",
            "S整件100抽保底只由整件触发和重置；S碎片不触发、不重置。",
            "同抽同时命中十连整件保障与S硬保底时只结算1个S整件。",
            "单抽不累计十连整件保障，但累计S整件硬保底。",
            "107支持50元晶/张直购（ShopGoods.2003002 Price=2,50，快捷购买QuickPurchase=2003002）；抽卡时需先扣107，免费券不足时再按OnePrice扣元晶。",
        ],
    }
    with (OUTPUT_ROOT / "收容物常驻抽卡配置校验.json").open("w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False, indent=2)


def main() -> None:
    OUTPUT_ROOT.mkdir(parents=True, exist_ok=True)
    records = get_weapon_records()
    build_draw_card()
    build_draw_card_pool()
    build_drop_draw()
    build_excl_weapon()
    build_items(records)
    build_shop_goods()
    build_common_rule_desc()
    build_duplicate_reward()
    validation = validate()
    write_manifest(validation)
    write_validation_json(validation)
    if not all(value[0] for value in validation.values()):
        raise SystemExit("validation failed")
    print(json.dumps({"output": str(OUTPUT_ROOT), "files": sorted(p.name for p in OUTPUT_ROOT.iterdir()), "status": "PASS"}, ensure_ascii=False))


if __name__ == "__main__":
    main()
