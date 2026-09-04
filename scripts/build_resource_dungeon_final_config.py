from __future__ import annotations

import hashlib
import json
import shutil
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


import vgame_paths

SOURCE_ROOT = vgame_paths.config_datas()
FRAMEWORK_ROOT = vgame_paths.source_docs_root() / "框架v2"
OUTPUT_ROOT = FRAMEWORK_ROOT / "06_最终配置" / "资源副本"

SOURCE_FILES = [
    Path("__enums__.xlsx"),
    Path("globals_conf/globals_conf.xlsx"),
    Path("level/UIlevel.xlsx"),
    Path("level/level.xlsx"),
    Path("level/LevelType.xlsx"),
    Path("level/ChallengeDungeonSupport.xlsx"),
    Path("level/CoinLevel.xlsx"),
    Path("system_open/SystemOpen.xlsx"),
    Path("unlock/Unlock.xlsx"),
    Path("items/Items.xlsx"),
    Path("Drop/Drop_dungeon.xlsx"),
]

LEVEL_REQUIREMENTS = [1, 20, 30, 40, 50]

CATEGORIES: dict[str, dict[str, Any]] = {
    "角色经验": {
        "type": "LevelType_Exp",
        "level_base": 210000,
        "unlock_base": 1090,
        "system_id": 301,
        "repeat_drop_base": 21000010,
        "first_drop_base": 21001010,
        "repeat": [
            [(101, 10), (100, 20)],
            [(101, 15), (100, 10)],
            [(102, 4), (101, 2), (100, 2)],
            [(102, 5)],
            [(102, 6)],
        ],
        "first": [[(101, 15)], [(101, 45)], [(102, 15)], [(102, 20)], [(102, 30)]],
    },
    "金币": {
        "type": "LevelType_Gold",
        "level_base": 220000,
        "unlock_base": 1100,
        "system_id": 302,
        "repeat_drop_base": 21000000,
        "first_drop_base": 21001000,
        "repeat": [[(1, 30000)], [(1, 35000)], [(1, 45000)], [(1, 50000)], [(1, 60000)]],
        "first": [[(1, 30000)], [(1, 90000)], [(1, 140000)], [(1, 220000)], [(1, 300000)]],
    },
    "收容物经验": {
        "type": "LevelType_WeaponExp",
        "level_base": 221000,
        "unlock_base": 1110,
        "system_id": 303,
        "repeat_drop_base": 21000020,
        "first_drop_base": 21001020,
        "repeat": [
            [(109, 10), (108, 20)],
            [(109, 15), (108, 10)],
            [(110, 4), (109, 2), (108, 2)],
            [(110, 5)],
            [(110, 6)],
        ],
        "first": [[(109, 15)], [(109, 45)], [(110, 15)], [(110, 20)], [(110, 30)]],
    },
    "舱室材料": {
        "type": "LevelType_RvCabin",
        "level_base": 222000,
        "unlock_base": 1210,
        "system_id": 327,
        "repeat_drop_base": 21000030,
        "first_drop_base": 21001030,
        "repeat": [
            [(3028, 6)],
            [(3028, 7)],
            [(3029, 3)],
            [(3029, 3), (3028, 1)],
            [(3030, 1), (3029, 1)],
        ],
        "first": [
            [(3028, 6)],
            [(3028, 18)],
            [(3029, 9)],
            [(3029, 15)],
            [(3030, 6)],
        ],
    },
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def headers(ws) -> dict[str, int]:
    result: dict[str, int] = {}
    for column in range(1, ws.max_column + 1):
        value = ws.cell(1, column).value
        if value is not None and value not in result:
            result[str(value)] = column
    return result


def find_rows(ws, header: dict[str, int], field: str, values: set[Any]) -> dict[Any, list[int]]:
    rows: dict[Any, list[int]] = defaultdict(list)
    column = header[field]
    for row in range(5, ws.max_row + 1):
        value = ws.cell(row, column).value
        if value in values:
            rows[value].append(row)
    return rows


def reward_text(rewards: list[tuple[int, int]]) -> str:
    return ";".join(f"{item_id},{quantity}" for item_id, quantity in rewards)


def save_atomic(wb, path: Path) -> None:
    temp = path.with_name(f".{path.name}.tmp.xlsx")
    wb.save(temp)
    temp.replace(path)


def copy_sources() -> dict[str, str]:
    source_hashes: dict[str, str] = {}
    for relative in SOURCE_FILES:
        source = SOURCE_ROOT / relative
        if not source.exists():
            raise FileNotFoundError(source)
        destination = OUTPUT_ROOT / "Datas" / relative
        destination.parent.mkdir(parents=True, exist_ok=True)
        source_hashes[str(relative).replace("\\", "/")] = sha256(source)
        shutil.copy2(source, destination)
    return source_hashes


def modify_ui_level() -> None:
    path = OUTPUT_ROOT / "Datas/level/UIlevel.xlsx"
    wb = load_workbook(path, data_only=False)
    ws = wb["level"]
    h = headers(ws)
    all_ids = {
        category["level_base"] + floor
        for category in CATEGORIES.values()
        for floor in range(1, 10)
    }
    rows = find_rows(ws, h, "LevelId", all_ids)
    missing = sorted(level_id for level_id in all_ids if not rows[level_id])
    if missing:
        raise ValueError(f"UIlevel 缺少 LevelId: {missing}")

    for category_name, category in CATEGORIES.items():
        for floor in range(1, 10):
            level_id = category["level_base"] + floor
            row = rows[level_id][0]
            if floor <= 5:
                ws.cell(row, h["##var"]).value = None
                ws.cell(row, h["Name"]).value = f"第{floor}层"
                ws.cell(row, h["LevelType"]).value = category["type"]
                ws.cell(row, h["LevelRefreshType"]).value = "Unlimited"
                ws.cell(row, h["RefreshNum"]).value = None
                ws.cell(row, h["TierNum"]).value = floor
                ws.cell(row, h["Reward"]).value = reward_text(category["repeat"][floor - 1])
                ws.cell(row, h["FirstReward"]).value = reward_text(category["first"][floor - 1])
                ws.cell(row, h["FirstRewardDrop"]).value = category["first_drop_base"] + floor
                ws.cell(row, h["FallRewardDrop"]).value = category["repeat_drop_base"] + floor
                ws.cell(row, h["Vitality"]).value = None
                ws.cell(row, h["RecommendLevel"]).value = LEVEL_REQUIREMENTS[floor - 1]
            else:
                ws.cell(row, h["##var"]).value = "##"
    save_atomic(wb, path)


def modify_level() -> None:
    path = OUTPUT_ROOT / "Datas/level/level.xlsx"
    wb = load_workbook(path, data_only=False)
    ws = wb["level"]
    h = headers(ws)
    all_ids = {
        category["level_base"] + floor
        for category in CATEGORIES.values()
        for floor in range(1, 10)
    }
    rows = find_rows(ws, h, "LevelId", all_ids)
    missing = sorted(level_id for level_id in all_ids if not rows[level_id])
    if missing:
        raise ValueError(f"level 缺少 LevelId: {missing}")

    for category_name, category in CATEGORIES.items():
        for floor in range(1, 10):
            level_id = category["level_base"] + floor
            row = rows[level_id][0]
            if floor <= 5:
                ws.cell(row, h["##var"]).value = None
                ws.cell(row, h["LevelName"]).value = f"{category_name}副本·第{floor}层"
                ws.cell(row, h["FirstReward"]).value = category["first_drop_base"] + floor
                ws.cell(row, h["FallReward"]).value = category["repeat_drop_base"] + floor
                ws.cell(row, h["Vitality"]).value = 0
            else:
                ws.cell(row, h["##var"]).value = "##"
    save_atomic(wb, path)


def modify_level_type() -> None:
    path = OUTPUT_ROOT / "Datas/level/LevelType.xlsx"
    wb = load_workbook(path, data_only=False)
    ws = wb["LevelType"]
    h = headers(ws)
    target_types = {category["type"] for category in CATEGORIES.values()}
    rows = find_rows(ws, h, "Id", target_types)
    display_rewards = {
        "角色经验": "102,-1;101,-1;100,-1",
        "金币": "1,-1",
        "收容物经验": "110,-1;109,-1;108,-1",
        "舱室材料": "3030,-1;3029,-1;3028,-1",
    }
    for category_name, category in CATEGORIES.items():
        row_list = rows[category["type"]]
        if len(row_list) != 1:
            raise ValueError(f"LevelType {category['type']} 行数异常: {row_list}")
        row = row_list[0]
        ws.cell(row, h["Name"]).value = f"{category_name}副本"
        ws.cell(row, h["Reward"]).value = display_rewards[category_name]
        ws.cell(row, h["SystemId"]).value = category["system_id"]
        if category_name == "舱室材料":
            # 当前没有专用资源副本大图，先使用舱室材料自身图标，避免继续显示收容物经验资产。
            ws.cell(row, h["Icon"]).value = "UIRes/UIAtlas/Icon/Items/item_1073.png"
            ws.cell(row, h["NameIcon"]).value = "UIRes/UIAtlas/Icon/Items/item_1073.png"
    save_atomic(wb, path)


def modify_system_open() -> None:
    path = OUTPUT_ROOT / "Datas/system_open/SystemOpen.xlsx"
    wb = load_workbook(path, data_only=False)
    ws = wb["SystemOpen"]
    h = headers(ws)
    target_ids = {category["system_id"] for category in CATEGORIES.values()}
    rows = find_rows(ws, h, "Id", target_ids)
    for category_name, category in CATEGORIES.items():
        row_list = rows[category["system_id"]]
        if len(row_list) != 1:
            raise ValueError(f"SystemOpen {category['system_id']} 行数异常: {row_list}")
        ws.cell(row_list[0], h["Name"]).value = f"{category_name}副本"
    save_atomic(wb, path)


def modify_challenge_support() -> None:
    path = OUTPUT_ROOT / "Datas/level/ChallengeDungeonSupport.xlsx"
    wb = load_workbook(path, data_only=False)
    ws = wb[wb.sheetnames[0]]
    h = headers(ws)
    target_types = {category["type"] for category in CATEGORIES.values()}
    rows = find_rows(ws, h, "ID", target_types)
    icon_by_category = {
        "角色经验": "Assets/GameResources/UIRes/UIAtlas/Level/Resource/resource_rukou_jingyan.png",
        "金币": "Assets/GameResources/UIRes/UIAtlas/Level/Resource/resource_rukou_jinbi.png",
        "收容物经验": "Assets/GameResources/UIRes/UIAtlas/Level/Resource/resource_rukou_zhuanwujingyan.png",
        "舱室材料": "Assets/GameResources/UIRes/UIAtlas/Icon/Items/item_1073.png",
    }
    for sort, (category_name, category) in enumerate(CATEGORIES.items(), start=1):
        row_list = rows[category["type"]]
        if row_list:
            row = row_list[0]
        else:
            row = next(
                candidate
                for candidate in range(5, ws.max_row + 1)
                if ws.cell(candidate, h["ID"]).value is None
            )
            ws.cell(row, h["##var"]).value = None
            ws.cell(row, h["ID"]).value = category["type"]
        ws.cell(row, h["Name"]).value = category_name
        ws.cell(row, h["GroupParent"]).value = "LevelTypeGroup_Resource"
        ws.cell(row, h["Icon"]).value = icon_by_category[category_name]
        ws.cell(row, h["Description"]).value = f"获取{category_name}资源"
        ws.cell(row, h["Sort"]).value = sort
        ws.cell(row, h["SystemOpenId"]).value = category["system_id"]
        ws.cell(row, h["UIView"]).value = "UIResourceLevelDetailView"
    save_atomic(wb, path)


def modify_drop() -> list[dict[str, Any]]:
    path = OUTPUT_ROOT / "Datas/Drop/Drop_dungeon.xlsx"
    wb = load_workbook(path, data_only=False)
    ws = wb["Drop"]
    h = headers(ws)

    protected_before = {
        (row, field): ws.cell(row, h[field]).value
        for row in range(5, ws.max_row + 1)
        for field in ("ID", "DropId", "DropId2")
    }
    target_drop_ids = {
        base + floor
        for category in CATEGORIES.values()
        for base in (category["repeat_drop_base"], category["first_drop_base"])
        for floor in range(1, 6)
    }
    drop_rows = find_rows(ws, h, "DropId", target_drop_ids)
    missing_rows: list[dict[str, Any]] = []

    def write_package(drop_id: int, rewards: list[tuple[int, int]], category_name: str, floor: int, kind: str) -> None:
        rows = drop_rows[drop_id]
        if not rows:
            for item_id, quantity in rewards:
                missing_rows.append(
                    {
                        "奖励类型": kind,
                        "副本": category_name,
                        "层": floor,
                        "DropId": drop_id,
                        "道具ID": item_id,
                        "数量": quantity,
                        "原因": "DropId 无现有行，需人工插行并复制公式",
                    }
                )
            return
        for index, row in enumerate(rows):
            if index < len(rewards):
                item_id, quantity = rewards[index]
                ws.cell(row, h["DropType"]).value = "fixed_"
                ws.cell(row, h["DropMinNum"]).value = 1
                ws.cell(row, h["DropMaxNum"]).value = 1
                ws.cell(row, h["DropItem_id"]).value = item_id
                ws.cell(row, h["DropItemNumMin"]).value = quantity
                ws.cell(row, h["DropItemNumMax"]).value = quantity
                ws.cell(row, h["DropWeight"]).value = 10000
                ws.cell(row, h["ShowItemID"]).value = None
            else:
                # 仅写奖励纯数据列，不触碰公式列。多余旧行设为零次无效掉落。
                ws.cell(row, h["DropType"]).value = "drop_Invalid"
                ws.cell(row, h["DropMinNum"]).value = 0
                ws.cell(row, h["DropMaxNum"]).value = 0
                # 预览代码可能先收集道具ID再过滤掉落类型；替换成当前包道具，避免旧108/金币图标穿透。
                ws.cell(row, h["DropItem_id"]).value = rewards[0][0]
                ws.cell(row, h["DropItemNumMin"]).value = 0
                ws.cell(row, h["DropItemNumMax"]).value = 0
                ws.cell(row, h["DropWeight"]).value = 0
                ws.cell(row, h["ShowItemID"]).value = None
        if len(rewards) > len(rows):
            for item_id, quantity in rewards[len(rows) :]:
                missing_rows.append(
                    {
                        "奖励类型": kind,
                        "副本": category_name,
                        "层": floor,
                        "DropId": drop_id,
                        "道具ID": item_id,
                        "数量": quantity,
                        "原因": "现有奖励行不足；禁止程序自动插入 Drop_dungeon 公式行",
                    }
                )

    for category_name, category in CATEGORIES.items():
        for floor in range(1, 6):
            write_package(
                category["repeat_drop_base"] + floor,
                category["repeat"][floor - 1],
                category_name,
                floor,
                "重复奖励",
            )
            write_package(
                category["first_drop_base"] + floor,
                category["first"][floor - 1],
                category_name,
                floor,
                "首通奖励",
            )

    protected_after = {
        (row, field): ws.cell(row, h[field]).value
        for row in range(5, ws.max_row + 1)
        for field in ("ID", "DropId", "DropId2")
    }
    if protected_before != protected_after:
        raise AssertionError("Drop_dungeon 受保护列 ID/DropId/DropId2 被修改")

    save_atomic(wb, path)
    return missing_rows


def style_sheet(ws, freeze: str = "A2") -> None:
    ws.freeze_panes = freeze
    ws.auto_filter.ref = ws.dimensions
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for column in range(1, ws.max_column + 1):
        values = [str(ws.cell(row, column).value or "") for row in range(1, min(ws.max_row, 120) + 1)]
        width = min(max(max((len(value) for value in values), default=8) + 2, 10), 44)
        ws.column_dimensions[get_column_letter(column)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def add_table_sheet(wb: Workbook, title: str, columns: list[str], rows: list[list[Any]]) -> None:
    ws = wb.create_sheet(title)
    ws.append(columns)
    for row in rows:
        ws.append(row)
    style_sheet(ws)


def create_manifest(missing_drop_rows: list[dict[str, Any]], validation: dict[str, Any]) -> None:
    wb = Workbook()
    wb.remove(wb.active)

    file_rows = []
    modified = {
        "level/UIlevel.xlsx": "更新20层展示/奖励/等级；6-9层保持停用；清空储能",
        "level/level.xlsx": "20层无储能；同步重复Drop；6-9层保持停用",
        "level/LevelType.xlsx": "收容物改名；补舱室材料奖励预览",
        "level/ChallengeDungeonSupport.xlsx": "补舱室材料入口；四类排序",
        "system_open/SystemOpen.xlsx": "更新资源副本名称",
        "Drop/Drop_dungeon.xlsx": "按沉淀MD更新固定重复奖励及首通奖励现有行",
    }
    for relative in SOURCE_FILES:
        key = str(relative).replace("\\", "/")
        file_rows.append(
            [
                key,
                "已修改" if key in modified else "只读随包镜像",
                modified.get(key, "用于引用、枚举或运行时核对，未改内容"),
                str(SOURCE_ROOT / relative),
            ]
        )
    add_table_sheet(wb, "配置清单", ["相对路径", "状态", "本批用途", "SVN源文件"], file_rows)

    level_rows = []
    for category_name, category in CATEGORIES.items():
        for floor in range(1, 6):
            level_rows.append(
                [
                    category_name,
                    floor,
                    category["level_base"] + floor,
                    category["type"],
                    LEVEL_REQUIREMENTS[floor - 1],
                    category["unlock_base"] + floor - 1,
                    0,
                    "每类共享3次/日（程序接入，未写成每层3次）",
                    category["repeat_drop_base"] + floor,
                    category["first_drop_base"] + floor,
                ]
            )
    add_table_sheet(
        wb,
        "目标关卡20行",
        ["副本", "层", "LevelId", "LevelType", "玩家等级要求", "等级UnlockId", "储能", "每日次数", "重复DropId", "首通DropId"],
        level_rows,
    )

    reward_rows = []
    missing_keys = {(r["奖励类型"], r["副本"], r["层"], r["道具ID"], r["数量"]) for r in missing_drop_rows}
    for category_name, category in CATEGORIES.items():
        for floor in range(1, 6):
            for kind, base_name, rewards in (
                ("重复奖励", "repeat_drop_base", category["repeat"][floor - 1]),
                ("首通奖励", "first_drop_base", category["first"][floor - 1]),
            ):
                drop_id = category[base_name] + floor
                for item_id, quantity in rewards:
                    status = "需人工新增Drop行" if (kind, category_name, floor, item_id, quantity) in missing_keys else "已写入镜像"
                    reward_rows.append([category_name, floor, kind, drop_id, item_id, quantity, status])
    add_table_sheet(wb, "奖励目标", ["副本", "层", "类型", "DropId", "道具ID", "数量", "写入状态"], reward_rows)

    gap_rows = [
        ["P0", "每类共享每日3次", "现有 UIlevel.RefreshNum 的客户端/服务端实现为按 LevelId 计数；若给每层填3会变成每类15次", "镜像保留 Unlimited/空次数；程序新增按LevelType共享的日次数状态与结算原子扣减", "BLOCKED_RUNTIME"],
        ["CLOSED", "舱室材料3-5层首通单物品包", "用户最终确认只保留一种材料：3029×9、3029×15、3030×6", "现有Drop行可完整承载，不再需要人工插行", "VERIFIED"],
        ["P1", "旧三星奖励", "UIlevel.ThreeStarRewardDrop=10000002 仍保留，是否取消玩家经验/旧附赠未确认", "确认后统一清空或保留", "PENDING_DESIGN"],
        ["P1", "金币局内产出", "CoinLevel 局内金币与固定Drop的最终叠加关系尚未做服务端联调", "验证结算总额等于目标固定重复奖励", "PENDING_RUNTIME_TEST"],
        ["P1", "首通扣次数与失败规则", "首通是否同时发重复奖励、失败/退出/断线是否扣次数尚未最终确认", "按设计建议补验收用例并锁定服务端规则", "PENDING_DESIGN"],
        ["P1", "扫荡/多倍", "旧4倍结算移除方向已写入设计，但扫荡、批量结算规则未最终确认", "确认后再更新globals_conf与UI行为", "PENDING_DESIGN"],
        ["P2", "舱室材料入口图标", "现有配置无专用资源副本大图，镜像暂使用舱室材料item_1073图标，已移除旧收容物经验图片引用", "UI资产提供后替换Icon/NameIcon", "PLACEHOLDER_ASSET"],
    ]
    add_table_sheet(wb, "程序与人工接入项", ["优先级", "事项", "证据/原因", "落地动作", "状态"], gap_rows)

    manual_rows = [[r[key] for key in ("奖励类型", "副本", "层", "DropId", "道具ID", "数量", "原因")] for r in missing_drop_rows]
    add_table_sheet(wb, "Drop人工新增行", ["奖励类型", "副本", "层", "DropId", "道具ID", "数量", "原因"], manual_rows)

    summary_rows = [["生成时间", validation["generated_at"]], ["状态", validation["status"]]]
    summary_rows.extend([[check["name"], check["result"]] for check in validation["checks"]])
    add_table_sheet(wb, "校验摘要", ["项目", "结果"], summary_rows)

    path = OUTPUT_ROOT / "00_资源副本配置清单.xlsx"
    save_atomic(wb, path)


def create_manual_drop_workbook(missing_drop_rows: list[dict[str, Any]]) -> None:
    directory = OUTPUT_ROOT / "人工操作"
    directory.mkdir(parents=True, exist_ok=True)
    legacy_path = directory / "Drop_dungeon_人工新增3行.xlsx"
    if legacy_path.exists() and legacy_path.parent.resolve() == directory.resolve():
        legacy_path.unlink()
    wb = Workbook()
    ws = wb.active
    ws.title = "人工操作结论"
    ws.append(["操作顺序", "DropId", "奖励类型", "副本", "层", "DropType", "DropMinNum", "DropMaxNum", "DropItem_id", "DropItemNumMin", "DropItemNumMax", "DropWeight", "说明"])
    if not missing_drop_rows:
        ws.append([None, None, None, "舱室材料", "3-5", None, None, None, None, None, None, None, "无需新增Drop行；现有21001033/34/35分别配置3029×9、3029×15、3030×6"])
    for index, row in enumerate(missing_drop_rows, start=1):
        ws.append(
            [
                index,
                row["DropId"],
                row["奖励类型"],
                row["副本"],
                row["层"],
                "fixed_",
                1,
                1,
                row["道具ID"],
                row["数量"],
                row["数量"],
                10000,
                "在同DropId现有行旁人工插入；ID列复制相邻公式；DropId填本行；DropId2保持0",
            ]
        )
    style_sheet(ws)

    instructions = wb.create_sheet("操作说明")
    instructions.append(["步骤", "说明"])
    instructions.append([1, "本版舱室材料第3-5层首通均为单物品包，不需要人工插入Drop行。"])
    instructions.append([2, "21001033=3029×9；21001034=3029×15；21001035=3030×6。"])
    instructions.append([3, "若后续再次改为多物品包，再按Drop公式安全规则人工插行。"])
    style_sheet(instructions)
    save_atomic(wb, directory / "Drop_dungeon_无需新增行.xlsx")


def validate(source_hashes: dict[str, str], missing_drop_rows: list[dict[str, Any]]) -> dict[str, Any]:
    checks: list[dict[str, str]] = []

    def record(name: str, ok: bool, detail: str) -> None:
        checks.append({"name": name, "result": "PASS" if ok else "FAIL", "detail": detail})

    ui_path = OUTPUT_ROOT / "Datas/level/UIlevel.xlsx"
    wb = load_workbook(ui_path, read_only=False, data_only=False)
    ws = wb["level"]
    h = headers(ws)
    ids = {
        category["level_base"] + floor
        for category in CATEGORIES.values()
        for floor in range(1, 10)
    }
    rows = find_rows(ws, h, "LevelId", ids)
    active = []
    inactive = []
    active_details_ok = True
    for category in CATEGORIES.values():
        for floor in range(1, 10):
            row = rows[category["level_base"] + floor][0]
            flag = ws.cell(row, h["##var"]).value
            (active if flag != "##" else inactive).append(category["level_base"] + floor)
            if floor <= 5:
                active_details_ok &= ws.cell(row, h["RecommendLevel"]).value == LEVEL_REQUIREMENTS[floor - 1]
                active_details_ok &= ws.cell(row, h["Vitality"]).value in (None, "")
                active_details_ok &= ws.cell(row, h["LevelRefreshType"]).value == "Unlimited"
                active_details_ok &= ws.cell(row, h["RefreshNum"]).value in (None, "")
    record("四类各5层启用", len(active) == 20 and len(inactive) == 16, f"启用{len(active)}行，停用{len(inactive)}行")
    record("等级与储能字段", active_details_ok, "等级要求1/20/30/40/50；UI储能为空；未误填每层3次")
    wb.close()

    level_path = OUTPUT_ROOT / "Datas/level/level.xlsx"
    level_wb = load_workbook(level_path, read_only=False, data_only=False)
    level_ws = level_wb["level"]
    level_h = headers(level_ws)
    level_rows = find_rows(level_ws, level_h, "LevelId", ids)
    first_reward_sync_ok = all(
        level_ws.cell(level_rows[category["level_base"] + floor][0], level_h["FirstReward"]).value
        == category["first_drop_base"] + floor
        for category in CATEGORIES.values()
        for floor in range(1, 6)
    )
    record("底层首通奖励引用", first_reward_sync_ok, "level.FirstReward与UIlevel.FirstRewardDrop同步到四类20个首通包")
    level_wb.close()

    drop_path = OUTPUT_ROOT / "Datas/Drop/Drop_dungeon.xlsx"
    source_drop_path = SOURCE_ROOT / "Drop/Drop_dungeon.xlsx"
    source_wb = load_workbook(source_drop_path, read_only=False, data_only=False)
    output_wb = load_workbook(drop_path, read_only=False, data_only=False)
    source_ws = source_wb["Drop"]
    output_ws = output_wb["Drop"]
    sh = headers(source_ws)
    oh = headers(output_ws)
    protected_ok = all(
        source_ws.cell(row, sh[field]).value == output_ws.cell(row, oh[field]).value
        for row in range(5, source_ws.max_row + 1)
        for field in ("ID", "DropId", "DropId2")
    )
    record("Drop公式保护列", protected_ok, "ID/DropId/DropId2 与SVN源完全一致；未插入或删除行")
    source_wb.close()
    output_wb.close()

    expected_missing: list[tuple[int, int, int]] = []
    actual_missing = [(r["DropId"], r["道具ID"], r["数量"]) for r in missing_drop_rows]
    record("人工新增Drop行清单", actual_missing == expected_missing, f"需人工新增{len(actual_missing)}行：{actual_missing}")

    source_unchanged = all(sha256(SOURCE_ROOT / relative) == source_hashes[str(relative).replace("\\", "/")] for relative in SOURCE_FILES)
    record("SVN源文件未修改", source_unchanged, "生成前后11个SVN源文件SHA256一致")

    formula_errors = []
    for relative in SOURCE_FILES:
        output = OUTPUT_ROOT / "Datas" / relative
        wb_scan = load_workbook(output, read_only=True, data_only=False)
        for ws_scan in wb_scan.worksheets:
            for row in ws_scan.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and cell.value in {"#VALUE!", "#DIV/0!", "#REF!", "#NAME?", "#NULL!", "#NUM!", "#N/A"}:
                        formula_errors.append(f"{relative}:{ws_scan.title}!{cell.coordinate}={cell.value}")
        wb_scan.close()
    record("公式错误扫描", not formula_errors, "未发现Excel错误值" if not formula_errors else "; ".join(formula_errors[:10]))

    return {
        "generated_at": datetime.now().astimezone().isoformat(timespec="seconds"),
        "status": "REVIEW_MIRROR_GENERATED_NOT_READY_FOR_SVN",
        "scope": "06_最终配置评审镜像；非正式Datas",
        "checks": checks,
        "blockers": [
            "每类共享每日3次需要程序按LevelType实现，当前镜像未误配为每层3次。",
        ],
        "pending": [
            "旧ThreeStarRewardDrop=10000002是否取消",
            "金币CoinLevel与固定Drop结算总额联调",
            "首通扣次数、失败/退出/断线、刷新、扫荡和批量结算规则",
            "舱室材料入口专用UI图标",
        ],
    }


def create_readme(validation: dict[str, Any]) -> None:
    lines = [
        "# 资源副本 06 最终配置评审镜像",
        "",
        f"> 生成时间：{validation['generated_at']}",
        f"> 状态：`{validation['status']}`",
        "> 本目录不是正式 Datas，SVN 源文件未修改。",
        "",
        "## 已落入",
        "",
        "- 金币、角色经验、收容物经验、舱室材料四类副本，各启用前5层；旧6至9层保持停用。",
        "- 五层玩家等级要求为1、20、30、40、50；沿用既有Unlock与LevelId骨架。",
        "- UIlevel储能已清空，level.Vitality保持0。",
        "- 重复奖励已按沉淀MD写入固定Drop；首通奖励已写入现有Drop行。",
        "- 原专武经验入口在展示层改名为收容物经验副本；底层旧枚举名暂保留兼容。",
        "- 舱室材料入口已补入ChallengeDungeonSupport空行，当前图标为临时复用。",
        "- 舱室材料第3至5层首通采用单物品包：3029×9、3029×15、3030×6；无需人工新增Drop行。",
        "",
        "## 阻塞项",
        "",
        "1. 每类共享每日3次不能用现有每LevelId的RefreshNum直接表达；当前保持Unlimited，等待程序按LevelType接入。",
        "",
        "## 交付文件",
        "",
        "- `00_资源副本配置清单.xlsx`：20层目标、奖励目标、接入项、校验摘要。",
        "- `Datas/`：按正式目录结构保存的评审镜像。",
        "- `人工操作/Drop_dungeon_无需新增行.xlsx`：本版Drop人工操作结论。",
        "- `资源副本配置校验.json`：机器校验结果。",
        "",
    ]
    (OUTPUT_ROOT / "README.md").write_text("\n".join(lines), encoding="utf-8")


def main() -> None:
    OUTPUT_ROOT.mkdir(parents=True, exist_ok=True)
    source_hashes = copy_sources()
    modify_ui_level()
    modify_level()
    modify_level_type()
    modify_system_open()
    modify_challenge_support()
    missing_drop_rows = modify_drop()
    validation = validate(source_hashes, missing_drop_rows)
    create_manifest(missing_drop_rows, validation)
    create_manual_drop_workbook(missing_drop_rows)
    create_readme(validation)
    (OUTPUT_ROOT / "资源副本配置校验.json").write_text(
        json.dumps(validation, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    print(json.dumps(validation, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
