# -*- coding: utf-8 -*-
"""
gen_hero_level.py — 主角等级 生成器

从节奏参数基线 MD 读取输入数据（每级成长等），计算派生列，输出完整等级表。
不再依赖旧源表——v2 参数基线 MD 就是唯一参数源。

派生公式：
  升级需要经验(Lv1) = 首级经验(参数)
  升级需要经验(Lv N) = 上一级升级需要经验 + 本级每级成长
  所需总经验(Lv N) = Σ 升级需要经验(1..N)

用法: python gen_hero_level.py
"""
import sys, io, os
import warnings
warnings.filterwarnings('ignore')
import openpyxl
from openpyxl.styles import Font, PatternFill

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

HERE = os.path.dirname(os.path.abspath(__file__))
PACE_DIR = os.path.dirname(HERE)
V2_DIR = os.path.dirname(PACE_DIR)
NUM_DIR = os.path.dirname(V2_DIR)
sys.path.insert(0, V2_DIR)
from _shared.md_parameter_reader import load_parameter_workbook, parameter_md_path

PACE_PARAM = parameter_md_path('节奏')
OUT = os.path.join(PACE_DIR, '生成_主角等级单源.xlsx')
REPORT = os.path.join(PACE_DIR, '_reports', 'P5_主角等级单源_生成与对照.md')

# 旧源表（仅用于对照验证，可选）
SRC_FILE = os.path.join(NUM_DIR, 'Vgame新手引导以及开启节奏.xlsx')


def read_params():
    """从节奏参数基线 MD 读取主角等级参数和输入数据。"""
    wb = load_parameter_workbook(PACE_PARAM, data_only=True)

    # 主角等级参数页
    ws_p = wb['主角等级参数']
    params = {}
    for r in range(1, ws_p.max_row + 1):
        k = ws_p.cell(r, 1).value
        v = ws_p.cell(r, 2).value
        if k:
            params[k] = v

    # 主角等级_输入页
    ws_i = wb['主角等级_输入']
    levels = []
    for r in range(5, ws_i.max_row + 1):
        lv = ws_i.cell(r, 1).value
        if lv is None:
            break
        levels.append({
            '等级': int(lv),
            '每级成长': ws_i.cell(r, 2).value or 0,
            '升级返还体力': ws_i.cell(r, 3).value,
            '对应主线章节': ws_i.cell(r, 4).value,
            '对应天数': ws_i.cell(r, 5).value,
        })

    wb.close()
    return params, levels


def generate(params, levels):
    """计算派生列。"""
    first_exp = params.get('主角首级升级经验', 80)
    cap = int(params.get('主角等级上限', 60))

    results = []
    prev_need = 0
    cum_total = 0

    for i, lv_data in enumerate(levels[:cap]):
        grow = lv_data['每级成长']
        if i == 0:
            need = first_exp
        else:
            need = prev_need + grow

        cum_total += need
        prev_need = need

        results.append({
            '等级': lv_data['等级'],
            '每级成长': grow,
            '升级需要经验': need,
            '升级返还体力': lv_data['升级返还体力'],
            '所需总经验': cum_total,
            '对应主线章节': lv_data['对应主线章节'],
            '对应天数': lv_data['对应天数'],
        })

    return results


def verify_against_source(results):
    """可选：与旧源表对照验证（证明搬迁无误）。"""
    if not os.path.exists(SRC_FILE):
        return None, None

    try:
        wb = openpyxl.load_workbook(SRC_FILE, data_only=True)
        ws = wb['主角等级']
    except Exception:
        return None, None

    max_err_need = 0
    max_err_total = 0
    for i, res in enumerate(results):
        r = 5 + i
        src_need = ws.cell(r, 4).value
        src_total = ws.cell(r, 6).value
        if src_need is not None:
            max_err_need = max(max_err_need, abs(res['升级需要经验'] - src_need))
        if src_total is not None:
            max_err_total = max(max_err_total, abs(res['所需总经验'] - src_total))

    wb.close()
    return max_err_need, max_err_total


def write_output(results):
    """写生成表。"""
    wb = openpyxl.Workbook()
    bold = Font(bold=True)
    hf = PatternFill('solid', fgColor='D9E1F2')
    der = PatternFill('solid', fgColor='E2EFDA')

    ws = wb.active
    ws.title = '主角等级_单源'
    ws.cell(1, 1, '主角等级（唯一权威源）— 由 gen_hero_level.py 从参数基线 MD 生成')
    ws.cell(1, 1).font = Font(bold=True, size=12)
    ws.cell(2, 1, '输入列从 03_沉淀MD/参数基线/节奏参数基线.md [主角等级_输入] 读取; 派生列(绿色)由公式计算')

    headers = ['等级', '每级成长', '升级需要经验', '升级返还体力', '所需总经验', '对应主线章节', '对应天数']
    for c, h in enumerate(headers, 1):
        ws.cell(4, c, h)
        ws.cell(4, c).font = bold
        ws.cell(4, c).fill = hf

    for res in results:
        r = ws.max_row + 1
        ws.cell(r, 1, res['等级'])
        ws.cell(r, 2, res['每级成长'])
        ws.cell(r, 3, res['升级需要经验'])
        ws.cell(r, 3).fill = der
        ws.cell(r, 4, res['升级返还体力'])
        ws.cell(r, 5, res['所需总经验'])
        ws.cell(r, 5).fill = der
        ws.cell(r, 6, res['对应主线章节'])
        ws.cell(r, 7, res['对应天数'])

    ws.freeze_panes = 'A5'
    for w, col in zip([8, 12, 14, 14, 14, 14, 12], 'ABCDEFG'):
        ws.column_dimensions[col].width = w

    wb.save(OUT)


def write_report(results, err_need, err_total):
    """写报告。"""
    chapter_filled = sum(1 for r in results if r['对应主线章节'] not in (None, ''))
    day_filled = sum(1 for r in results if r['对应天数'] not in (None, ''))
    missing_days = [r['等级'] for r in results if r['对应天数'] in (None, '')]
    md = [
        '# P5 · 主角等级 生成与对照', '',
        '> 从 `03_沉淀MD/参数基线/节奏参数基线.md` 读取输入数据(每级成长、章节/天数锚点等)，计算派生列(升级需要经验/所需总经验)。', '',
        '## 派生公式', '',
        '```',
        '升级需要经验(Lv1) = 首级经验(%d)' % results[0]['升级需要经验'],
        '升级需要经验(Lv N) = 上一级 + 本级每级成长',
        '所需总经验(Lv N) = 累计',
        '```', '',
        '## 关键数据', '',
        '| 指标 | 值 |', '|---|---|',
        '| 等级上限 | %d |' % len(results),
        '| 首级经验 | %d |' % results[0]['升级需要经验'],
        '| 末级升级需要 | %d |' % results[-1]['升级需要经验'],
        '| 末级总经验 | %d |' % results[-1]['所需总经验'],
        '| 对应主线章节覆盖 | %d/%d |' % (chapter_filled, len(results)),
        '| 对应天数覆盖 | %d/%d |' % (day_filled, len(results)),
        '',
    ]

    if missing_days:
        md += [
            '## 待补日程锚点', '',
            '以下等级在当前旧源天数曲线中没有明确达到天数，保留为空，等待玩家进度模拟或新版日程参数：',
            '',
            ', '.join('Lv%d' % lv for lv in missing_days),
            '',
        ]

    if err_need is not None:
        if err_need < 1e-6 and err_total < 1e-6:
            md.append('## 对照旧源表')
            md.append('')
            md.append('→ ✅ 源内部一致, 派生列可由输入生成')
        else:
            md.append('## 对照旧源表')
            md.append('')
            md.append('→ 🔴 存在偏差: 升级需要最大误差 %.3g, 总经验最大误差 %.3g' % (err_need, err_total))
    else:
        md.append('> 旧源表不存在或无法读取，跳过对照。')

    os.makedirs(os.path.dirname(REPORT), exist_ok=True)
    with open(REPORT, 'w', encoding='utf-8') as f:
        f.write('\n'.join(md))


def main():
    print('=== 主角等级 生成 ===\n')

    params, levels = read_params()
    print('参数: 等级上限=%s, 首级经验=%s, 输入数据=%d级' % (
        params.get('主角等级上限'), params.get('主角首级升级经验'), len(levels)))

    results = generate(params, levels)
    print('生成: %d级, 末级升级需要=%d, 末级总经验=%d' % (
        len(results), results[-1]['升级需要经验'], results[-1]['所需总经验']))
    chapter_filled = sum(1 for r in results if r['对应主线章节'] not in (None, ''))
    day_filled = sum(1 for r in results if r['对应天数'] not in (None, ''))
    print('锚点: 对应主线章节=%d/%d, 对应天数=%d/%d' % (
        chapter_filled, len(results), day_filled, len(results)))

    err_need, err_total = verify_against_source(results)
    if err_need is not None:
        print('\n对照旧源表: 升级需要最大误差=%.3g, 总经验最大误差=%.3g' % (err_need, err_total))
        if err_need < 1e-6 and err_total < 1e-6:
            print('→ ✅ 源内部一致, 派生列可由输入生成')
        else:
            print('→ 🔴 存在偏差')

    write_output(results)
    print('\n[生成表] %s' % OUT)

    write_report(results, err_need, err_total)
    print('[报告]   %s' % REPORT)


if __name__ == '__main__':
    main()
