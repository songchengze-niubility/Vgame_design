# -*- coding: utf-8 -*-
"""
gen_difficulty_inf.py — 无限挑战难度线 生成器 (框架重搭 P-难度线)

读取节奏参数基线 MD / 难度线参数 (a,b,c,怪物等级上限,最大大波数) →
生成 关卡大波数 → 怪物等级 曲线，并与现有 战斗框架/无限挑战难度线 对照。

公式 (已逆向验证):
  怪物等级(n) = round( b + c×(n-1)/((n-1)+a) )，封顶 怪物等级上限

只读源表。生成表为独立新文件。
用法: python gen_difficulty_inf.py [--check]
"""
import sys, io, os, argparse
import warnings
warnings.filterwarnings('ignore')
import openpyxl
from openpyxl.styles import Font, PatternFill

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

HERE = os.path.dirname(os.path.abspath(__file__))
PACE_DIR = os.path.dirname(HERE)  # 节奏/
V2_DIR = os.path.dirname(PACE_DIR)  # 框架v2/
NUM_DIR = os.path.dirname(V2_DIR)  # 数值文档/数值文档/
sys.path.insert(0, V2_DIR)
from _shared.md_parameter_reader import load_parameter_workbook, parameter_md_path

PACE_PARAM = parameter_md_path('节奏')

OUT     = os.path.join(PACE_DIR, '生成_无限挑战难度线.xlsx')
REPORT  = os.path.join(PACE_DIR, '_reports', 'P难度线_无限挑战_生成与对照.md')
SRC     = os.path.join(NUM_DIR, 'Vgame基本战斗框架.xlsx')
SRC_SHEET = '无限挑战难度线'
DATA_ROW0 = 11   # n=1 所在行 (col2=大波数, col3=怪物等级)

def read_params():
    wb = load_parameter_workbook(PACE_PARAM, data_only=True)
    ws = wb['难度线参数']
    p = {}
    for r in range(1, ws.max_row + 1):
        k = ws.cell(r, 1).value
        if k in ('a', 'b', 'c', '怪物等级上限', '最大大波数'):
            p[k] = ws.cell(r, 2).value
    wb.close()
    return p

def monster_level(n, a, b, c, cap):
    val = b + c * (n - 1) / ((n - 1) + a)
    return min(round(val), int(cap))

def main():
    ap = argparse.ArgumentParser(); ap.add_argument('--check', action='store_true')
    args = ap.parse_args()
    p = read_params()
    a, b, c, cap, nmax = p['a'], p['b'], p['c'], p['怪物等级上限'], int(p['最大大波数'])
    levels = [(n, monster_level(n, a, b, c, cap)) for n in range(1, nmax + 1)]

    # 对照
    wb = openpyxl.load_workbook(SRC, data_only=True); ws = wb[SRC_SHEET]
    checks = fails = 0; first = []
    for i, (n, lv) in enumerate(levels):
        r = DATA_ROW0 + i
        act = ws.cell(r, 3).value
        if act is None:
            continue
        checks += 1
        if int(act) != lv:
            fails += 1
            if len(first) < 15: first.append((n, lv, act))
    wb.close()

    md=[]
    def o(s=''):
        print(s); md.append(s)
    o('# 难度线 · 无限挑战 生成与对照'); o()
    o('> 怪物等级(n)=round(b+c×(n-1)/((n-1)+a))，封顶%d。a=%g b=%g c=%g。源表只读。' % (cap, a, b, c)); o()
    o('- 对照波数 %d，不一致 %d' % (checks, fails)); o()
    if not fails:
        o('✅ **逐格复现**：%d 波怪物等级全部一致。' % checks)
    else:
        o('🔴 不一致样本：'); o('| n | 生成 | 源表 |'); o('|---:|---:|---:|')
        for (n, lv, act) in first: o('| %d | %d | %s |' % (n, lv, act))
    o()
    o('- 样本：n1=%d, n3=%d, n5=%d, n35=%d, n%d=%d' %
      (levels[0][1], levels[2][1], levels[4][1], levels[34][1], nmax, levels[-1][1]))

    if not args.check:
        owb = openpyxl.Workbook(); bold=Font(bold=True); hf=PatternFill('solid',fgColor='D9E1F2')
        ows = owb.active; ows.title='无限挑战难度线'
        ows.append(['【自动生成·勿手改】源:03_沉淀MD/参数基线/节奏参数基线.md/难度线参数  怪物等级=round(b+c×(n-1)/((n-1)+a)) 封顶%d' % cap])
        ows.append([]); ows.append(['关卡大波数', '怪物等级'])
        ows.cell(3,1).font=bold; ows.cell(3,2).font=bold; ows.cell(3,1).fill=hf; ows.cell(3,2).fill=hf
        for (n, lv) in levels: ows.append([n, lv])
        owb.save(OUT)

    os.makedirs(os.path.dirname(REPORT), exist_ok=True)
    with open(REPORT, 'w', encoding='utf-8') as f: f.write('\n'.join(md))
    print('\n[报告] %s' % REPORT)
    if not args.check: print('[生成表] %s' % OUT)

if __name__ == '__main__':
    main()
