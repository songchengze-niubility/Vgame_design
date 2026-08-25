# -*- coding: utf-8 -*-
"""
gen_difficulty_tower.py — 爬塔难度线 · 天数→层数预计 生成器

读取节奏参数基线 MD / 难度线参数(爬塔段) → 生成天数→层数预计曲线，
与现有 战斗框架/爬塔难度线 对照。

公式 (已逆向验证, 两段式):
  d <= 保底天数: 层数 = 首层 + (保底层-首层)×(1-r^(d-1))/(1-r^(保底天-1))
  d >  保底天数: 层数 = 保底层 + (总层-保底层)×(d-保底天)/(满天-保底天)
  取整 = round

范围:
  1. 天数→层数预计(原始+取整)
  2. 层数→难度映射(层→天数桶→怪物等级+层型/奖励系数)
只读源表。生成表为独立新文件。
用法: python gen_difficulty_tower.py [--check]
"""
import sys, io, os, argparse
import warnings
warnings.filterwarnings('ignore')
import openpyxl
from openpyxl.styles import Font, PatternFill

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

HERE = os.path.dirname(os.path.abspath(__file__))
PACE_DIR = os.path.dirname(HERE)  # 节奏/
V2_DIR = os.path.dirname(PACE_DIR)  # 框架v2/
NUM_DIR = os.path.dirname(V2_DIR)  # 数值文档/数值文档/
sys.path.insert(0, V2_DIR)
from _shared.md_parameter_reader import load_parameter_workbook, parameter_md_path

PACE_PARAM = parameter_md_path('节奏')

OUT     = os.path.join(PACE_DIR, '生成_爬塔难度线_天数层数.xlsx')
REPORT  = os.path.join(PACE_DIR, '_reports', 'P难度线_爬塔天数层数_生成与对照.md')
SRC     = os.path.join(NUM_DIR, 'Vgame基本战斗框架.xlsx')
SRC_SHEET = '爬塔难度线'
DAY_ROW0 = 20   # 天数=1 所在行 (c5=天数, c6=层数预计, c7=取整)

KEYS = ['爬塔第一天层数','爬塔保底层数','爬塔打到保底天数','爬塔总层数','爬塔打满天数','爬塔衰减比']

def read_params():
    wb = load_parameter_workbook(PACE_PARAM, data_only=True); ws = wb['难度线参数']
    p = {}
    for r in range(1, ws.max_row+1):
        k = ws.cell(r,1).value
        if k in KEYS: p[k] = ws.cell(r,2).value
    wb.close()
    return p

def floors(d, p):
    first=p['爬塔第一天层数']; base=p['爬塔保底层数']; baseD=int(p['爬塔打到保底天数'])
    mx=p['爬塔总层数']; mxD=int(p['爬塔打满天数']); r=p['爬塔衰减比']
    if d <= baseD:
        return first + (base-first)*(1 - r**(d-1))/(1 - r**(baseD-1))
    return base + (mx-base)*(d-baseD)/(mxD-baseD)


def earliest_day_for_floor(floor, day_rows):
    for d, raw, rnd in day_rows:
        if rnd >= floor:
            return d
    return ''


def read_floor_mapping(day_rows):
    wb = openpyxl.load_workbook(SRC, data_only=True)
    ws = wb[SRC_SHEET]
    type_name = {0: '小层', 1: '中层', 2: '大层'}
    rows = []
    for r in range(20, 120):
        floor = ws.cell(r, 12).value
        monster_lv = ws.cell(r, 22).value
        if floor is None or monster_lv is None:
            continue
        floor = int(floor)
        rows.append({
            '层数': floor,
            '层型': type_name.get(ws.cell(r, 11).value, ''),
            '预计最早天数': earliest_day_for_floor(floor, day_rows),
            '对应主角等级': ws.cell(r, 14).value,
            '怪物等级': monster_lv,
            '奖励系数': ws.cell(r, 19).value,
            '来源行': r,
        })
    wb.close()
    return rows

def main():
    ap=argparse.ArgumentParser(); ap.add_argument('--check',action='store_true'); args=ap.parse_args()
    p = read_params(); mxD = int(p['爬塔打满天数'])
    rows = [(d, floors(d,p), round(floors(d,p))) for d in range(1, mxD+1)]
    floor_rows = read_floor_mapping(rows)

    wb = openpyxl.load_workbook(SRC, data_only=True); ws = wb[SRC_SHEET]
    raw_e = int_e = 0; cnt = 0
    for i,(d,raw,rnd) in enumerate(rows):
        r = DAY_ROW0 + i
        a_raw = ws.cell(r,6).value; a_int = ws.cell(r,7).value
        if a_raw is not None: raw_e = max(raw_e, abs(raw-a_raw)); cnt += 1
        if a_int is not None: int_e = max(int_e, abs(rnd-a_int))
    wb.close()

    if not args.check:
        owb = openpyxl.Workbook(); bold=Font(bold=True); hf=PatternFill('solid',fgColor='D9E1F2')
        ows = owb.active; ows.title='爬塔天数层数'
        ows.append(['【自动生成·勿手改】两段式: 几何爬升到保底, 之后线性到满'])
        ows.append([]); ows.append(['天数','层数预计','取整'])
        for c in range(1,4): ows.cell(3,c).font=bold; ows.cell(3,c).fill=hf
        for (d,raw,rnd) in rows: ows.append([d, round(raw,6), rnd])

        mws = owb.create_sheet('爬塔层数难度')
        headers = ['层数','层型','预计最早天数','对应主角等级','怪物等级','奖励系数','来源页','口径']
        mws.append(['【自动生成·勿手改】源: Vgame基本战斗框架.xlsx/爬塔难度线；预计最早天数来自天数→层数曲线'])
        mws.append([])
        mws.append(headers)
        for c in range(1, len(headers)+1):
            mws.cell(3,c).font=bold; mws.cell(3,c).fill=hf
        for row in floor_rows:
            mws.append([
                row['层数'], row['层型'], row['预计最早天数'], row['对应主角等级'],
                row['怪物等级'], row['奖励系数'],
                'Vgame基本战斗框架.xlsx/爬塔难度线:%d' % row['来源行'],
                '层数→天数桶→等级+层型/奖励系数',
            ])
        for wsx in (ows, mws):
            wsx.freeze_panes = 'A4'
            for col in wsx.columns:
                wsx.column_dimensions[col[0].column_letter].width = min(
                    max(len(str(cell.value)) if cell.value is not None else 0 for cell in col) + 2, 36)
        owb.save(OUT)

    md=[]
    def o(s=''):
        print(s); md.append(s)
    o('# 爬塔难度线 · 天数→层数预计 生成与对照'); o()
    o('> 两段式(几何爬升→线性)。第一天%g层, 保底%g@%g天, 满%g@%g天, 衰减%g。源表只读。' %
      (p['爬塔第一天层数'],p['爬塔保底层数'],p['爬塔打到保底天数'],p['爬塔总层数'],p['爬塔打满天数'],p['爬塔衰减比'])); o()
    o('- 对照天数 %d, 层数预计最大误差 %.2e, 取整不一致 %s' % (cnt, raw_e, '0' if int_e==0 else ('%g'%int_e))); o()
    if raw_e < 1e-9 and int_e == 0:
        o('✅ **天数→层数曲线逐格复现**(原始+取整)。')
        o('- 层数→难度映射已补: %d 层，包含预计最早天数、层型、主角等级、怪物等级和奖励系数。' % len(floor_rows))
    else:
        o('🔴 存在偏差，需复查。')
    os.makedirs(os.path.dirname(REPORT),exist_ok=True)
    with open(REPORT,'w',encoding='utf-8') as f: f.write('\n'.join(md))
    print('\n[报告] %s' % REPORT)
    if not args.check: print('[生成表] %s' % OUT)

if __name__ == '__main__':
    main()
