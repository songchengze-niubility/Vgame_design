# -*- coding: utf-8 -*-
"""
gen_onboarding_pace.py — 开启节奏与玩家日程总览

从旧源 `Vgame新手引导以及开启节奏.xlsx` 抽取:
  - 开启节奏和引导
  - 新手流程
  - 主角等级天数曲线
  - 养成进度预估

输出 `节奏/生成_开启节奏总览.xlsx`，作为 v2 节奏层的人读入口。
旧源只读；正式 SystemOpen/Unlock/LevelId 链暂不在本脚本内落地。
"""
import io
import os
import sys
import warnings
from collections import defaultdict

warnings.filterwarnings("ignore")

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

HERE = os.path.dirname(os.path.abspath(__file__))
PACE_DIR = os.path.dirname(HERE)
V2_DIR = os.path.dirname(PACE_DIR)
NUM_DIR = os.path.dirname(V2_DIR)

SRC_FILE = os.path.join(NUM_DIR, "Vgame新手引导以及开启节奏.xlsx")
DIFFICULTY_ALL = os.path.join(PACE_DIR, "全玩法难度线.xlsx")
OUT = os.path.join(PACE_DIR, "生成_开启节奏总览.xlsx")
REPORT = os.path.join(PACE_DIR, "_reports", "P节奏_开启与日程总览.md")


def blank(v):
    return "" if v is None else v


def as_int(v):
    if isinstance(v, (int, float)):
        return int(v)
    return v


def trigger_of(row):
    lv = row["角色等级"]
    ch = row["章节"]
    stage = row["主线关卡"]
    day = row["天数"]
    if lv not in ("", None):
        return "主角等级%d" % int(lv) if isinstance(lv, (int, float)) else "主角等级%s" % lv
    if ch not in ("", None) and stage not in ("", None):
        return "主线%s-%s" % (ch, stage)
    if day not in ("", None):
        return "第%s天" % day
    return "待确认"


def node_type(row):
    has_growth = row["养成线"] not in ("", None)
    has_play = row["玩法线"] not in ("", None)
    has_guide = row["引导"] not in ("", None, "")
    if has_growth and has_play:
        return "养成+玩法"
    if has_play:
        return "玩法开启"
    if has_growth:
        return "养成开启"
    if has_guide:
        return "引导节点"
    return "进度节点"


def read_open_rows(wb):
    ws = wb["开启节奏和引导"]
    rows = []
    for r in range(3, ws.max_row + 1):
        vals = {
            "源行": r,
            "开启顺序": blank(ws.cell(r, 2).value),
            "养成线": blank(ws.cell(r, 3).value),
            "玩法线": blank(ws.cell(r, 4).value),
            "引导": blank(ws.cell(r, 5).value),
            "额外奖励": blank(ws.cell(r, 6).value),
            "角色等级": blank(ws.cell(r, 7).value),
            "章节": blank(ws.cell(r, 8).value),
            "主线关卡": blank(ws.cell(r, 9).value),
            "预计游戏时间(min)": blank(ws.cell(r, 10).value),
            "天数": blank(ws.cell(r, 11).value),
        }
        if not any(vals[k] not in ("", None) for k in vals if k != "源行"):
            continue
        vals["节点类型"] = node_type(vals)
        vals["开启锚点"] = trigger_of(vals)
        vals["来源"] = "Vgame新手引导以及开启节奏.xlsx/开启节奏和引导:%d" % r
        rows.append(vals)
    return rows


def read_flow_rows(wb):
    ws = wb["新手流程"]
    rows = []
    for r in range(1, ws.max_row + 1):
        step = ws.cell(r, 1).value
        text = ws.cell(r, 2).value
        if step is None and text is None:
            continue
        if isinstance(step, (int, float)):
            rows.append([int(step), blank(text), r])
        elif text not in (None, "") and r > 20:
            rows.append(["补充", blank(text), r])
    return rows


def read_day_curve(wb):
    ws = wb["主角等级"]
    rows = []
    for r in range(21, ws.max_row + 1):
        day = ws.cell(r, 14).value
        lv = ws.cell(r, 15).value
        daily_exp = ws.cell(r, 17).value
        if not isinstance(day, (int, float)) or not isinstance(lv, (int, float)):
            continue
        if daily_exp in (None, ""):
            continue
        rows.append({
            "天数": int(day),
            "主角等级预计": int(lv),
            "所需经验": ws.cell(r, 16).value,
            "当天经验": daily_exp,
            "来源行": r,
        })
    return rows


def read_growth_rows(wb):
    ws = wb["养成进度预估"]
    headers = [blank(ws.cell(1, c).value) for c in range(1, ws.max_column + 1)]
    rows = []
    by_lv = {}
    for r in range(2, ws.max_row + 1):
        lv = ws.cell(r, 1).value
        if not isinstance(lv, (int, float)):
            continue
        vals = [blank(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)]
        rows.append(vals)
        by_lv[int(lv)] = dict(zip(headers, vals))
    return headers, rows, by_lv


def read_difficulty_gaps():
    if not os.path.exists(DIFFICULTY_ALL):
        return []
    wb = openpyxl.load_workbook(DIFFICULTY_ALL, data_only=True)
    ws = wb["覆盖状态"]
    rows = []
    for r in range(2, ws.max_row + 1):
        play = ws.cell(r, 1).value
        total = ws.cell(r, 3).value or 0
        monster = ws.cell(r, 4).value or 0
        status = ws.cell(r, 5).value
        if play and total and monster < total:
            rows.append([play, total, monster, total - monster, status])
    wb.close()
    return rows


def build_day_milestones(open_rows, day_curve, growth_by_lv):
    by_day = defaultdict(list)
    for row in open_rows:
        day = row["天数"]
        if day not in ("", None):
            by_day[str(day)].append(row)

    out = []
    for day_row in day_curve:
        day_key = str(day_row["天数"])
        lv = day_row["主角等级预计"]
        growth = growth_by_lv.get(lv, {})
        nodes = by_day.get(day_key, [])
        out.append([
            day_row["天数"],
            lv,
            day_row["所需经验"],
            day_row["当天经验"],
            "；".join(filter(None, [blank(n["养成线"]) for n in nodes])),
            "；".join(filter(None, [blank(n["玩法线"]) for n in nodes])),
            "；".join(filter(None, [blank(n["引导"]) for n in nodes])),
            blank(growth.get("代表主线章节")),
            blank(growth.get("代表主线关卡")),
            blank(growth.get("代表关卡ID")),
            blank(growth.get("怪物等级")),
            blank(growth.get("角色平均等级")),
            blank(growth.get("专武平均等级")),
            blank(growth.get("旅券平均等级")),
            blank(growth.get("映射装备平均等级")),
            "主角等级:%d; 养成预估按同等级精确匹配" % day_row["来源行"],
        ])

    existing = {str(r[0]) for r in out}
    for day_key, nodes in by_day.items():
        if day_key in existing:
            continue
        out.append([
            day_key, "", "", "",
            "；".join(filter(None, [blank(n["养成线"]) for n in nodes])),
            "；".join(filter(None, [blank(n["玩法线"]) for n in nodes])),
            "；".join(filter(None, [blank(n["引导"]) for n in nodes])),
            "", "", "", "", "", "", "", "",
            "开启节奏有该天数，但主角等级日程表未覆盖",
        ])
    return out


def append_table(ws, headers, rows, header_row=1):
    hf = PatternFill("solid", fgColor="D9E1F2")
    bold = Font(bold=True)
    for c, h in enumerate(headers, 1):
        ws.cell(header_row, c, h)
        ws.cell(header_row, c).font = bold
        ws.cell(header_row, c).fill = hf
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "A%d" % (header_row + 1)
    for col in ws.columns:
        letter = col[0].column_letter
        width = min(max(len(str(cell.value)) if cell.value is not None else 0 for cell in col) + 2, 48)
        ws.column_dimensions[letter].width = max(width, 10)
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def write_output(open_rows, flow_rows, day_curve, growth_headers, growth_rows, milestones, gaps):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet("阅读导航")
    append_table(ws, ["项目", "内容"], [
        ["产物", OUT],
        ["旧源", SRC_FILE],
        ["用途", "节奏层人读入口：开启节点、新手流程、逐日等级、养成预估、缺口追踪"],
        ["开启节点数", len(open_rows)],
        ["新手流程步骤", len(flow_rows)],
        ["主角日程天数", len(day_curve)],
    ])

    ws = wb.create_sheet("开启节奏")
    headers = ["源行", "节点类型", "开启锚点", "养成线", "玩法线", "引导", "额外奖励",
               "角色等级", "章节", "主线关卡", "预计游戏时间(min)", "天数", "来源"]
    append_table(ws, headers, [[row[h] for h in headers] for row in open_rows])

    ws = wb.create_sheet("主角日程")
    append_table(ws, [
        "天数", "主角等级预计", "所需经验", "当天经验", "当日养成开启", "当日玩法开启",
        "当日引导", "代表主线章节", "代表主线关卡", "代表关卡ID", "怪物等级",
        "角色平均等级", "专武平均等级", "旅券平均等级", "映射装备平均等级", "口径",
    ], milestones)

    ws = wb.create_sheet("新手流程")
    append_table(ws, ["步骤", "内容", "源行"], flow_rows)

    ws = wb.create_sheet("养成进度预估")
    append_table(ws, growth_headers, growth_rows)

    ws = wb.create_sheet("缺口追踪")
    max_day_level = max((r["主角等级预计"] for r in day_curve), default=0)
    rows = [
        ["正式开放链", "SystemOpen/Unlock", "本表来自策划旧源，不等同正式配置", "回正式配置 SystemOpen.xlsx + Unlock.xlsx 校验", "P1"],
        ["爬塔", "层数→难度映射", "已由 `生成_爬塔难度线_天数层数.xlsx` 补一张映射页；正式 TowerTrial 配置仍未接", "后续接 TowerTrial 正式表", "P1"],
        ["PVPVE", "战斗怪物等级/系数", "当前只有竞技场分层模型", "回 LevelType_PVPVE/UIlevel/Arena/level 补战斗源", "P0"],
        ["BOSS挑战", "战斗怪物等级/系数", "当前只有奖励难度档", "回 BossDungeon/Boss_Dungeon.xlsx 补战斗源", "P0"],
    ]
    if max_day_level < 60:
        rows.insert(0, [
            "主角等级日程",
            "Lv%d-Lv60" % (max_day_level + 1),
            "旧源有效逐日表未覆盖",
            "等待玩家进度模拟或新版日程参数",
            "P1",
        ])
    for play, total, monster, missing, status in gaps:
        rows.append(["全玩法难度线", play, "缺怪物等级 %s/%s 行" % (missing, total), status, "自动统计"])
    append_table(ws, ["模块", "对象", "缺口", "下一步", "优先级"], rows)

    wb.save(OUT)


def write_report(open_rows, flow_rows, day_curve, milestones, gaps):
    play_nodes = sum(1 for r in open_rows if r["玩法线"] not in ("", None))
    growth_nodes = sum(1 for r in open_rows if r["养成线"] not in ("", None))
    guide_nodes = sum(1 for r in open_rows if r["引导"] not in ("", None, ""))
    max_day = max((r["天数"] for r in day_curve), default=0)
    max_level = max((r["主角等级预计"] for r in day_curve), default=0)
    md = [
        "# 节奏 · 开启与日程总览",
        "",
        "> 从 `Vgame新手引导以及开启节奏.xlsx` 迁移开启节奏、新手流程、主角日程和养成进度预估，输出 `节奏/生成_开启节奏总览.xlsx`。",
        "",
        "## 覆盖",
        "",
        "| 项目 | 数量 |",
        "|---|---:|",
        f"| 开启节点 | {len(open_rows)} |",
        f"| 玩法开启节点 | {play_nodes} |",
        f"| 养成开启节点 | {growth_nodes} |",
        f"| 引导节点 | {guide_nodes} |",
        f"| 新手流程步骤 | {len(flow_rows)} |",
        f"| 主角日程天数 | {len(day_curve)} |",
        f"| 旧源逐日表最高覆盖 | Day{max_day} / Lv{max_level} |",
        f"| 日程总览行 | {len(milestones)} |",
        "",
        "## 仍需正式配置接入",
        "",
        "- `SystemOpen/Unlock`：本表保留策划旧源口径，不能替代正式配置开放条件。",
        "- `PVPVE` 与 `BOSS挑战`：仍缺战斗怪物等级/系数源，当前只记录分层/奖励难度。",
    ]
    if max_level < 60:
        md.append("- `Lv%d-Lv60`：旧源有效逐日等级表只覆盖到 Day%d/Lv%d，后段日程需要玩家进度模拟或新版参数。" % (max_level + 1, max_day, max_level))
    if gaps:
        md += ["", "## 全玩法难度线缺口统计", "", "| 玩法 | 总行数 | 含怪物等级 | 缺口 | 状态 |", "|---|---:|---:|---:|---|"]
        for play, total, monster, missing, status in gaps:
            md.append(f"| {play} | {total} | {monster} | {missing} | {status} |")
    md += ["", "开启节奏总览生成完成。"]

    os.makedirs(os.path.dirname(REPORT), exist_ok=True)
    with open(REPORT, "w", encoding="utf-8") as f:
        f.write("\n".join(md))
    for line in md:
        print(line)
    print("\n[生成表] %s" % OUT)
    print("[报告]   %s" % REPORT)


def main():
    src_wb = openpyxl.load_workbook(SRC_FILE, data_only=True)
    open_rows = read_open_rows(src_wb)
    flow_rows = read_flow_rows(src_wb)
    day_curve = read_day_curve(src_wb)
    growth_headers, growth_rows, growth_by_lv = read_growth_rows(src_wb)
    src_wb.close()

    milestones = build_day_milestones(open_rows, day_curve, growth_by_lv)
    gaps = read_difficulty_gaps()
    write_output(open_rows, flow_rows, day_curve, growth_headers, growth_rows, milestones, gaps)
    write_report(open_rows, flow_rows, day_curve, milestones, gaps)


if __name__ == "__main__":
    main()
