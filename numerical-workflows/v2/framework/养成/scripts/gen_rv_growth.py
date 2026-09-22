# -*- coding: utf-8 -*-
"""房车养成正式数值文档生成器。

房车整备保持账号全局额外属性口径，参数仍来自养成参数基线 MD。
出行指南自 2026-08-24 起改为每职业 66 个多级节点（Lv2/Lv3/Lv5/Lv10）；
2026-09-04 CHG-053 将百分比三维与二级属性的高级指南材料(1156)拆分到
Lv2 与 Lv3。本生成器直接读取 `06_最终配置/房车` 评审镜像
（RvGuideNode/RvGuideNodeLv/RvStage/RvGuide）重建多级消耗明细，
并按 `03_沉淀MD/养成/房车养成系统数值设计文档-2026-07-03.md`
的验收标准做全量断言，防止旧单级口径回流。
"""
import io
import os
import re
import sys
from collections import defaultdict

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

HERE = os.path.dirname(os.path.abspath(__file__))
GROWTH_DIR = os.path.dirname(HERE)
V2_DIR = os.path.dirname(GROWTH_DIR)
sys.path.insert(0, V2_DIR)
from _shared.md_parameter_reader import load_parameter_workbook, parameter_md_path

SHARED_PARAM = parameter_md_path("共享")
GROWTH_PARAM = parameter_md_path("养成")
OUT = os.path.join(GROWTH_DIR, "生成_房车.xlsx")
REPORT = os.path.join(GROWTH_DIR, "_reports", "P房车养成_生成与对照.md")
MIRROR_DIR = os.path.join(V2_DIR, "06_最终配置", "房车")

# 指南/职业/职业代码/职业基础材料 ItemId；HeroJob 取正式 RvGuide。
GUIDE_CAREERS = {
    1: ("行程指南", "旅行家", "Explorer", 1075),
    2: ("安行指南", "战士", "Soldier", 1076),
    3: ("探路指南", "佣兵", "Contractor", 1077),
    4: ("见闻指南", "学者", "Technologist", 1078),
    5: ("共鸣指南", "元能者", "Arcanist", 1079),
}
STAGE_NAMES = {1: "基础", 2: "中级", 3: "高级", 4: "精通"}
# 名称取正式 Items.xlsx（2026-09-04 回读）；ItemId 才是稳定键。
ITEM_NAMES = {
    1: "金币",
    1074: "房车零件",
    1075: "行程便签",
    1076: "安途清单",
    1077: "探路标记",
    1078: "见闻手札",
    1079: "调律灯片",
    1156: "高阶课程笔记",
    1157: "精英课程笔记",
}
# 沉淀MD/参数基线沿用旧别名，此处保留映射便于人读对照。
GUIDE_ALIAS = {1075: "行程便签", 1076: "安行便签", 1077: "探路便签", 1078: "见闻便签", 1079: "共鸣便签"}
ATTR_NAMES = {
    "Hp": "生命",
    "Attack": "攻击",
    "Armor": "防御",
    "HpBonus": "生命%",
    "AttackBonus": "攻击%",
    "ArmorBonus": "防御%",
    "EnergyPreAddBonusLevel": "能量回复等级",
    "CoolDownReduceLevel": "冷却缩减等级",
    "CritLevel": "暴击等级",
    "CritDmgBonusLevel": "暴伤等级",
    "DmgBonus": "伤害加成",
    "DmgImmune": "伤害减免",
}
FIXED_ATTR_CODES = {"Hp", "Attack", "Armor"}
PERCENT_ATTRS = {"生命%", "攻击%", "防御%", "伤害加成", "伤害减免"}

# 验收常量：来自房车养成系统数值设计文档「生成与验收」。
EXPECTED_TOTAL_NODES = 330
EXPECTED_NODES_PER_CAREER = 66
EXPECTED_LV_ROWS = 1800
EXPECTED_LVMAX_DIST = {10: 15, 5: 31, 3: 15, 2: 5}
EXPECTED_LV_ROWS_PER_CAREER = 360
EXPECTED_CAREER_TOTALS = {"career_material": 5000, 1156: 500, 1157: 50, "gold": 1364500}
EXPECTED_STAGE_GOLD = {"基础": 8900, "中级": 54000, "高级": 291000, "精通": 1010600}
EXPECTED_STAGE_MATERIAL = {"基础": 200, "中级": 600, "高级": 1500, "精通": 2700}
EXPECTED_MAX_ATTR_TOTALS = {
    "生命": 10000,
    "攻击": 3000,
    "防御": 2000,
    "生命%": 0.5,
    "攻击%": 0.5,
    "防御%": 0.5,
    "伤害加成": 0.1,
    "伤害减免": 0.1,
}


def clean(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def to_num(value, default=0):
    if value is None or value == "":
        return default
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return value
    text = str(value).strip()
    try:
        num = float(text)
    except ValueError:
        return default
    return int(num) if num.is_integer() else num


def find_row(ws, title):
    for row in range(1, ws.max_row + 1):
        if ws.cell(row, 1).value == title:
            return row
    raise RuntimeError("房车参数缺少区块: %s" % title)


def read_table(ws, title):
    title_row = find_row(ws, title)
    header_row = title_row + 1
    headers = [clean(c.value) for c in ws[header_row]]
    last_col = 0
    for idx, header in enumerate(headers, start=1):
        if header != "":
            last_col = idx
    headers = headers[:last_col]
    rows = []
    row_idx = header_row + 1
    while row_idx <= ws.max_row:
        first = ws.cell(row_idx, 1).value
        if first is None or first == "":
            break
        if isinstance(first, str) and first.endswith("参数") and row_idx > header_row + 1:
            break
        values = [clean(ws.cell(row_idx, col).value) for col in range(1, last_col + 1)]
        rows.append({headers[i]: values[i] if i < len(values) else "" for i in range(len(headers))})
        row_idx += 1
    return rows


def read_key_values(ws, title):
    rows = read_table(ws, title)
    return {row["参数"]: row["值"] for row in rows if row.get("参数")}


def read_shared():
    wb = load_parameter_workbook(SHARED_PARAM, data_only=True, read_only=True)
    sp = wb["标准人"]
    std = {sp.cell(r, 1).value: to_num(sp.cell(r, 2).value) for r in range(2, sp.max_row + 1) if sp.cell(r, 1).value}
    fp = wb["投放系数"]
    rv_coef = None
    for r in range(2, fp.max_row + 1):
        if fp.cell(r, 1).value == "房车整备":
            rv_coef = to_num(fp.cell(r, 2).value)
            break
    wb.close()
    if rv_coef is None:
        raise RuntimeError("共享参数基线 MD 缺少投放系数: 房车整备")
    return std, rv_coef


def read_params():
    wb = load_parameter_workbook(GROWTH_PARAM, data_only=True, read_only=True)
    if "房车参数" not in wb.sheetnames:
        raise RuntimeError("养成参数基线 MD 缺少房车参数页")
    ws = wb["房车参数"]
    general = read_key_values(ws, "房车整备参数")
    cost_bands = read_table(ws, "房车强化消耗分段")
    results = read_table(ws, "房车强化结果")
    careers = read_table(ws, "出行指南职业")
    wb.close()
    return general, cost_bands, results, careers


def read_luban_rows(path, sheet):
    """读取 Luban 格式评审镜像表。

    首个非空行为 ##var 列名（##var 标记常落在 A 列占位，真实数据列右移一列）；
    以 ## 开头的行是类型/分组/中文说明，全部跳过；全空行跳过。
    """

    if not os.path.exists(path):
        raise FileNotFoundError(path)
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb[sheet]
    header = None
    marker_idx = -1
    result = []
    for values in ws.iter_rows(values_only=True):
        cells = [clean(v) for v in values]
        non_empty = [c for c in cells if c != ""]
        if not non_empty:
            continue
        if header is None:
            if str(non_empty[0]) != "##var":
                raise RuntimeError("镜像表首行不是 ##var: %s" % path)
            header = cells
            marker_idx = cells.index("##var")
            continue
        if str(non_empty[0]).startswith("#"):
            continue
        row = {}
        for idx, name in enumerate(header):
            if idx == marker_idx or name in (None, ""):
                continue
            row[str(name)] = cells[idx] if idx < len(cells) else ""
        result.append(row)
    wb.close()
    return result


def parse_cost_items(text):
    pairs = []
    for part in str(text or "").split(";"):
        part = part.strip()
        if not part:
            continue
        match = re.fullmatch(r"(\d+),(\d+)", part)
        if not match:
            raise RuntimeError("CostItemList 格式异常: %s" % part)
        pairs.append((int(match.group(1)), int(match.group(2))))
    return pairs


def attr_curve(std, coef, level_cap, base_ratio):
    full = {
        "生命": std["生命"] * coef,
        "攻击": std["攻击"] * coef,
        "防御": std["防御"] * coef,
    }
    rows = []
    denom = max(level_cap - 1, 1)
    for lv in range(1, level_cap + 1):
        progress = (lv - 1) / denom
        ratio = base_ratio + (1 - base_ratio) * progress
        rows.append({
            "目标等级": lv,
            "属性进度": progress,
            "生命": full["生命"] * ratio,
            "攻击": full["攻击"] * ratio,
            "防御": full["防御"] * ratio,
        })
    return rows, full


def cost_for_level(cost_bands, lv):
    if lv <= 1:
        return "", ""
    for band in cost_bands:
        lo = int(to_num(band["目标等级下限"]))
        hi = int(to_num(band["目标等级上限"]))
        if lo <= lv <= hi:
            delta = lv - lo
            part = to_num(band["零件基数"]) + to_num(band["零件成长"]) * delta
            gold = to_num(band["金币基数"]) + to_num(band["金币成长"]) * delta
            return part, gold
    return "", ""


def load_guide_mirror():
    nodes = read_luban_rows(os.path.join(MIRROR_DIR, "RvGuideNode.xlsx"), "RvGuideNode")
    levels = read_luban_rows(os.path.join(MIRROR_DIR, "RvGuideNodeLv.xlsx"), "RvGuideNodeLv")
    stages = read_luban_rows(os.path.join(MIRROR_DIR, "RvStage.xlsx"), "RvStage")
    guides = read_luban_rows(os.path.join(MIRROR_DIR, "RvGuide.xlsx"), "RvGuide")
    if len(nodes) != EXPECTED_TOTAL_NODES:
        raise RuntimeError("RvGuideNode 镜像应为 %d 节点，实际 %d" % (EXPECTED_TOTAL_NODES, len(nodes)))
    if len(levels) != EXPECTED_LV_ROWS:
        raise RuntimeError("RvGuideNodeLv 镜像应为 %d 行，实际 %d" % (EXPECTED_LV_ROWS, len(levels)))
    stage_guide = {}
    stage_name = {}
    for row in stages:
        stage_guide[int(row["Id"])] = int(row["GuideId"])
        stage_name[int(row["Id"])] = str(row["Name"]).replace("阶段", "")
    guide_jobs = {int(row["Id"]): str(row["HeroJob"]) for row in guides}
    if set(stage_guide.values()) != set(GUIDE_CAREERS):
        raise RuntimeError("RvStage 指南集合异常: %s" % sorted(set(stage_guide.values())))
    for guide_no, (_name, _career, job, _item) in GUIDE_CAREERS.items():
        if guide_jobs.get(guide_no) != job:
            raise RuntimeError("RvGuide HeroJob 与生成器职业代码不一致: %s" % guide_no)
    return nodes, levels, stage_guide, stage_name


def validate_levels(levels_raw):
    """断言逐级行结构：TargetLv 连续、Value 严格递增、1156 拆分符合 CHG-053。"""

    by_node = defaultdict(list)
    for row in levels_raw:
        by_node[int(row["NodeId"])].append(row)
    for node_id, rows in by_node.items():
        rows.sort(key=lambda r: int(r["TargetLv"]))
        target_lvs = [int(r["TargetLv"]) for r in rows]
        if target_lvs != list(range(1, len(rows) + 1)):
            raise RuntimeError("节点 %s TargetLv 不连续: %s" % (node_id, target_lvs))
        values = [to_num(r["Value"]) for r in rows]
        if any(b <= a for a, b in zip(values, values[1:])):
            raise RuntimeError("节点 %s Value 未严格递增: %s" % (node_id, values))
        attr_code = str(rows[0]["Attribute"])
        if attr_code in FIXED_ATTR_CODES:
            continue
        # CHG-053：百分比三维与二级属性的 1156 仅允许出现在 Lv2/Lv3。
        for r in rows:
            lv = int(r["TargetLv"])
            has_1156 = any(iid == 1156 for iid, _q in parse_cost_items(r["CostItemList"]))
            if has_1156 and lv not in (2, 3):
                raise RuntimeError("节点 %s Lv%s 不应消耗1156（CHG-053拆分仅落Lv2/Lv3）" % (node_id, lv))
    return by_node


def build_guide_tables(nodes, levels, stage_guide, stage_name):
    node_by_id = {}
    for row in nodes:
        node_id = int(row["Id"])
        node_by_id[node_id] = {
            "NodeId": node_id,
            "指南编号": stage_guide[int(row["StageId"])],
            "阶段": stage_name[int(row["StageId"])],
            "节点名": str(row["Name"]),
            "等级上限": int(row["NodeLvMax"]),
        }
    attr_of_node = {}
    level_rows = []
    for row in levels:
        node_id = int(row["NodeId"])
        node = node_by_id[node_id]
        attr_code = str(row["Attribute"])
        attr_of_node[node_id] = attr_code
        guide_no = node["指南编号"]
        career_item = GUIDE_CAREERS[guide_no][3]
        costs = parse_cost_items(row["CostItemList"])
        for item_id, _qty in costs:
            if item_id not in (career_item, 1156, 1157):
                raise RuntimeError("节点 %s 出现非预期材料 %s" % (node_id, item_id))
        level_rows.append({
            "NodeId": node_id,
            "指南": GUIDE_CAREERS[guide_no][0],
            "职业": GUIDE_CAREERS[guide_no][1],
            "阶段": node["阶段"],
            "节点名": node["节点名"],
            "目标等级": int(row["TargetLv"]),
            "属性": ATTR_NAMES.get(attr_code, attr_code),
            "累计Value": to_num(row["Value"]),
            "职业基础材料": sum(q for iid, q in costs if iid == career_item),
            "高级指南材料": sum(q for iid, q in costs if iid == 1156),
            "超高级指南材料": sum(q for iid, q in costs if iid == 1157),
            "金币": to_num(row["CostCurrencyNum"]),
        })
    levels_by_node = defaultdict(list)
    for row in level_rows:
        levels_by_node[row["NodeId"]].append(row)
    for node_id, node in node_by_id.items():
        node["属性"] = ATTR_NAMES.get(attr_of_node[node_id], attr_of_node[node_id])
        rows = levels_by_node[node_id]
        node["升级记录数"] = len(rows)
        node["满级Value"] = max(r["累计Value"] for r in rows)
        node["职业基础材料"] = sum(r["职业基础材料"] for r in rows)
        node["高级指南材料"] = sum(r["高级指南材料"] for r in rows)
        node["超高级指南材料"] = sum(r["超高级指南材料"] for r in rows)
        node["金币"] = sum(r["金币"] for r in rows)

    career_rows = []
    stage_rows = {}
    career_agg = defaultdict(lambda: defaultdict(float))
    stage_agg = defaultdict(lambda: defaultdict(float))
    for guide_no, (guide_name, career, _job, career_item) in GUIDE_CAREERS.items():
        nodes_of_career = [n for n in node_by_id.values() if n["指南编号"] == guide_no]
        if len(nodes_of_career) != EXPECTED_NODES_PER_CAREER:
            raise RuntimeError("%s 节点数应为 %d，实际 %d" % (guide_name, EXPECTED_NODES_PER_CAREER, len(nodes_of_career)))
        dist = defaultdict(int)
        for node in nodes_of_career:
            dist[node["等级上限"]] += 1
        if dict(dist) != EXPECTED_LVMAX_DIST:
            raise RuntimeError("%s 节点等级上限分布异常: %s" % (guide_name, dict(dist)))
        for node in nodes_of_career:
            stage_key = (guide_name, node["阶段"])
            stage_agg[stage_key]["节点数"] += 1
            stage_agg[stage_key]["升级记录"] += node["升级记录数"]
            stage_agg[stage_key]["职业基础材料"] += node["职业基础材料"]
            stage_agg[stage_key]["高级指南材料"] += node["高级指南材料"]
            stage_agg[stage_key]["超高级指南材料"] += node["超高级指南材料"]
            stage_agg[stage_key]["金币"] += node["金币"]
            career_agg[guide_name]["职业基础材料"] += node["职业基础材料"]
            career_agg[guide_name]["高级指南材料"] += node["高级指南材料"]
            career_agg[guide_name]["超高级指南材料"] += node["超高级指南材料"]
            career_agg[guide_name]["金币"] += node["金币"]
            career_agg[guide_name]["升级记录"] += node["升级记录数"]
        totals = career_agg[guide_name]
        if totals["职业基础材料"] != EXPECTED_CAREER_TOTALS["career_material"]:
            raise RuntimeError("%s 职业基础材料总量异常: %s" % (guide_name, totals["职业基础材料"]))
        if totals["高级指南材料"] != EXPECTED_CAREER_TOTALS[1156]:
            raise RuntimeError("%s 高级指南材料(1156)总量异常: %s" % (guide_name, totals["高级指南材料"]))
        if totals["超高级指南材料"] != EXPECTED_CAREER_TOTALS[1157]:
            raise RuntimeError("%s 超高级指南材料(1157)总量异常: %s" % (guide_name, totals["超高级指南材料"]))
        if totals["金币"] != EXPECTED_CAREER_TOTALS["gold"]:
            raise RuntimeError("%s 金币总量异常: %s" % (guide_name, totals["金币"]))
        if totals["升级记录"] != EXPECTED_LV_ROWS_PER_CAREER:
            raise RuntimeError("%s 升级记录应为 %d，实际 %s" % (guide_name, EXPECTED_LV_ROWS_PER_CAREER, totals["升级记录"]))
        attr_totals = defaultdict(float)
        for node in nodes_of_career:
            attr_totals[node["属性"]] += node["满级Value"]
        for attr, expected in EXPECTED_MAX_ATTR_TOTALS.items():
            if abs(attr_totals[attr] - expected) > 1e-6:
                raise RuntimeError("%s %s 满级总量异常: %s" % (guide_name, attr, attr_totals[attr]))
        for attr, total in sorted(attr_totals.items()):
            career_rows.append({
                "指南": guide_name,
                "职业": career,
                "属性": attr,
                "满级属性增量": int(total) if float(total).is_integer() else total,
                "是否百分比": "是" if attr in PERCENT_ATTRS else "否",
            })
    for (guide_name, stage), agg in stage_agg.items():
        if int(agg["金币"]) != EXPECTED_STAGE_GOLD[stage]:
            raise RuntimeError("%s %s 阶段金币异常: %s" % (guide_name, stage, agg["金币"]))
        if int(agg["职业基础材料"]) != EXPECTED_STAGE_MATERIAL[stage]:
            raise RuntimeError("%s %s 阶段职业基础材料异常: %s" % (guide_name, stage, agg["职业基础材料"]))
        stage_rows[(guide_name, stage)] = {
            "指南": guide_name,
            "阶段": stage,
            "节点数": int(agg["节点数"]),
            "升级记录": int(agg["升级记录"]),
            "职业基础材料": int(agg["职业基础材料"]),
            "高级指南材料(1156)": int(agg["高级指南材料"]),
            "超高级指南材料(1157)": int(agg["超高级指南材料"]),
            "金币": int(agg["金币"]),
        }
    stage_order = {"基础": 0, "中级": 1, "高级": 2, "精通": 3}
    stage_list = [stage_rows[key] for key in sorted(stage_rows, key=lambda k: (k[0], stage_order[k[1]]))]
    node_list = [node_by_id[nid] for nid in sorted(node_by_id)]
    return node_list, level_rows, stage_list, career_rows


def append_sheet(wb, title, headers, rows):
    ws = wb.create_sheet(title)
    ws.append(headers)
    fill = PatternFill("solid", fgColor="D9E1F2")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in rows:
        if isinstance(row, dict):
            ws.append([row.get(h, "") for h in headers])
        else:
            ws.append(row)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    widths = {
        "指南": 14, "职业": 12, "阶段": 10, "节点名": 18, "属性": 20,
        "进度/条件": 46, "口径": 30, "备注": 26,
    }
    for idx, header in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = widths.get(header, 14)
    return ws


def fmt(value):
    if isinstance(value, float):
        return round(value, 6)
    return value


def main():
    std, coef = read_shared()
    general, cost_bands, results, baseline_careers = read_params()
    level_cap = int(to_num(general.get("等级上限"), 60))
    base_ratio = float(to_num(general.get("基础占比"), 0.1))

    nodes_raw, levels_raw, stage_guide, stage_name = load_guide_mirror()
    validate_levels(levels_raw)
    node_list, level_rows, stage_list, career_rows = build_guide_tables(nodes_raw, levels_raw, stage_guide, stage_name)

    rv_level_rows, full = attr_curve(std, coef, level_cap, base_ratio)
    for row in rv_level_rows:
        part, gold = cost_for_level(cost_bands, row["目标等级"])
        row["房车零件"] = part
        row["金币"] = gold

    economy_rows = []

    summary_rows = [
        ["标准人生命", std["生命"]],
        ["标准人攻击", std["攻击"]],
        ["标准人防御", std["防御"]],
        ["房车整备投放系数", coef],
        ["房车满级生命", full["生命"]],
        ["房车满级攻击", full["攻击"]],
        ["房车满级防御", full["防御"]],
        ["房车等级上限", level_cap],
        ["出行指南职业数", len(GUIDE_CAREERS)],
        ["出行指南节点数", len(node_list)],
        ["单职业节点数", EXPECTED_NODES_PER_CAREER],
        ["节点等级上限分布(每职业)", "Lv10×15；Lv5×31；Lv3×15；Lv2×5"],
        ["出行指南升级记录数", len(level_rows)],
        ["单职业升级记录数", EXPECTED_LV_ROWS_PER_CAREER],
        ["节点解锁规则", "节点Lv1即可解锁后续；升级消耗见「出行指南节点等级」"],
        ["高级指南材料(1156)口径", "CHG-053：百分比三维与二级属性拆分落Lv2/Lv3，单节点合计2/4/11/21"],
        ["出行指南数据来源", "06_最终配置/房车/RvGuideNode.xlsx + RvGuideNodeLv.xlsx（CHG-053评审镜像）"],
        ["房车整备数据来源", "03_沉淀MD/参数基线/养成参数基线.md 房车参数"],
    ]

    career_display = []
    for guide_no, (guide_name, career, job, item_id) in GUIDE_CAREERS.items():
        career_display.append({
            "指南": guide_name,
            "职业": career,
            "职业Code": job,
            "专属材料": ITEM_NAMES[item_id],
            "专属ItemId": item_id,
            "说明": "统一均衡养成路线（仅材料不同）；CHG-053后每职业66个多级节点",
        })

    wb = openpyxl.Workbook()
    del wb[wb.sheetnames[0]]
    append_sheet(wb, "口径总览", ["项目", "值"], summary_rows)
    append_sheet(
        wb,
        "房车整备曲线",
        ["目标等级", "属性进度", "生命", "攻击", "防御", "房车零件", "金币"],
        [{k: fmt(v) for k, v in row.items()} for row in rv_level_rows],
    )
    append_sheet(wb, "强化结果", list(results[0].keys()), results)
    append_sheet(wb, "出行指南职业", list(career_display[0].keys()), career_display)
    append_sheet(wb, "出行指南阶段", list(stage_list[0].keys()), stage_list)
    append_sheet(wb, "出行指南节点", list(node_list[0].keys()), [{k: fmt(v) for k, v in row.items()} for row in node_list])
    append_sheet(
        wb,
        "出行指南节点等级",
        ["NodeId", "指南", "职业", "阶段", "节点名", "目标等级", "属性", "累计Value", "职业基础材料", "高级指南材料", "超高级指南材料", "金币"],
        [{k: fmt(v) for k, v in row.items()} for row in level_rows],
    )
    append_sheet(wb, "职业满级属性汇总", ["指南", "职业", "属性", "满级属性增量", "是否百分比"], career_rows)
    wb.save(OUT)

    os.makedirs(os.path.dirname(REPORT), exist_ok=True)
    lines = [
        "# 房车养成 · 生成与对照",
        "",
        "> 房车整备数值口径来自共享/养成参数基线；出行指南多级节点口径来自",
        "> `06_最终配置/房车` 评审镜像（RvGuideNode + RvGuideNodeLv），并按",
        "> `03_沉淀MD/养成/房车养成系统数值设计文档-2026-07-03.md` 验收标准全量断言。",
        "",
        "## 全局升级",
        "",
        "- 房车整备属性分类：账号全局额外属性。",
        "- 投放系数：%s 个标准人。" % coef,
        "- 满级属性：生命 %.0f / 攻击 %.0f / 防御 %.0f。" % (full["生命"], full["攻击"], full["防御"]),
        "- 等级上限：%s；Lv1 基础占比：%s。" % (level_cap, base_ratio),
        "",
        "## 出行指南（多级节点）",
        "",
        "- 职业数：%d；总节点：%d；单职业节点：%d。" % (len(GUIDE_CAREERS), len(node_list), EXPECTED_NODES_PER_CAREER),
        "- 单职业节点等级上限分布：Lv10×15、Lv5×31、Lv3×15、Lv2×5；升级记录 360/职业，五职业合计 %d。" % len(level_rows),
        "- 每职业满级属性：固定生命10,000/攻击3,000/防御2,000；生命%/攻击%/防御%各50%；增伤/减伤各10%。",
        "- 每职业消耗：职业基础材料5,000、高级指南材料(1156)×500、超高级指南材料(1157)×50、金币1,364,500。",
        "- CHG-053：百分比三维与二级属性的1156拆分落Lv2/Lv3（Lv2 1/2/5/10，Lv3 1/2/6/11），单节点合计2/4/11/21不变；生成器已断言1156不出现在这些节点的Lv1/Lv4+。",
        "- 阶段金币预算：8,900/54,000/291,000/1,010,600（合计1,364,500）；阶段职业基础材料 200/600/1,500/2,700。",
        "- 消耗校验总量（验收断言值）：房车零件15,457；五职业职业基础材料各5,000；高级指南材料(1156)五职业2,500；超高级指南材料(1157)五职业250；金币合计8,991,650（整备2,169,150 + 指南6,822,500）。",
        "- 2026-09-04起本表为QA交付口径：不再输出「资源消耗汇总」「经济消耗明细」账本页；经济总账改读「房车整备曲线」与「出行指南节点等级」结构化页签。",
        "",
        "## 旧口径切换记录",
        "",
        "- 2026-08-24 前生成链为 310 个单级节点；2026-09-04 起出行指南改读 06 评审镜像多级节点，防止旧单级口径回流。",
        "- 养成参数基线 `房车参数` 中的旧「出行指南节点」单级块不再作为出行指南输入；房车整备参数、消耗分段与强化结果仍从基线读取。",
        "",
        "## 产物",
        "",
        "- `养成/生成_房车.xlsx`",
        "- `养成/_reports/P房车养成_生成与对照.md`",
    ]
    with open(REPORT, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))

    print("房车养成生成完成：整备系数 %.2f，等级 %d，职业 %d，节点 %d，升级记录 %d" % (
        coef, level_cap, len(GUIDE_CAREERS), len(node_list), len(level_rows)))
    print("[生成表] %s" % OUT)
    print("[报告] %s" % REPORT)


if __name__ == "__main__":
    main()
