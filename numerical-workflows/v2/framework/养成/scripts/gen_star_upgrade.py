# -*- coding: utf-8 -*-
"""
gen_star_upgrade.py — 角色升星 属性派生验证 生成器

读取共享参数基线 MD + 养成参数基线 MD → 计算升星的基础属性%投放 + 偏向二级属性投放，与源表逐格对照。

派生公式：
  最终基础属性%(生/攻/防) = 投放系数页「角色升星」系数 (直接等于，走%不走绝对值)
  偏向普通二级属性 = 属性节点数 × 每节点二级属性投放
  各节点分配: 全属性% + 偏向二级（星3/星5各一次）

对照目标：角色升星sheet(R12) + 养成属性规划(R040)

只读源表。生成表为独立新文件。
用法: python gen_star_upgrade.py
"""
import sys, io, os
import warnings
warnings.filterwarnings('ignore')
import openpyxl
from openpyxl.styles import Font, PatternFill

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

HERE = os.path.dirname(os.path.abspath(__file__))
GROWTH_DIR = os.path.dirname(HERE)  # 养成/
V2_DIR = os.path.dirname(GROWTH_DIR)  # 框架v2/
NUM_DIR = os.path.dirname(V2_DIR)  # 数值文档/数值文档/
sys.path.insert(0, V2_DIR)
from _shared.md_parameter_reader import load_parameter_workbook, parameter_md_path

SHARED_PARAM = parameter_md_path('共享')
GROWTH_PARAM = parameter_md_path('养成')
OUT     = os.path.join(GROWTH_DIR, '生成_角色升星.xlsx')
REPORT  = os.path.join(GROWTH_DIR, '_reports', 'P养成模块_角色升星_生成与对照.md')
SRC     = os.path.join(NUM_DIR, 'Vgame养成数值.xlsx')


def almost(a, b, tol=0.001):
    if a is None and b is None:
        return True
    try:
        return abs(float(a) - float(b)) <= tol
    except (TypeError, ValueError):
        return str(a) == str(b)


def read_params():
    """从共享参数读取投放系数，从养成参数读取升星相关参数。"""
    # 投放系数 → 从共享参数读
    wb_s = load_parameter_workbook(SHARED_PARAM, data_only=True)
    fp = wb_s['投放系数']
    star_pct = None
    for r in range(2, fp.max_row + 1):
        if fp.cell(r, 1).value == '角色升星':
            star_pct = fp.cell(r, 2).value
            break
    wb_s.close()

    # 养成模块参数页 - 升星区 → 从养成参数读
    wb_g = load_parameter_workbook(GROWTH_PARAM, data_only=True)
    gp = wb_g['养成模块参数']
    gm = {}
    for r in range(1, gp.max_row + 1):
        k = gp.cell(r, 1).value
        v = gp.cell(r, 2).value
        if k and str(k).startswith('升星'):
            gm[k] = v
    wb_g.close()

    return star_pct, gm


def read_source():
    """从源表读取升星的手搓值用于对照。"""
    wb = openpyxl.load_workbook(SRC, data_only=True)
    src = {}

    # 角色升星 sheet
    ws = wb['角色升星']
    # R12: 最终基础属性% (B/C/D = 生命/攻击/防御)
    src['基础属性%_生命'] = ws.cell(12, 2).value
    src['基础属性%_攻击'] = ws.cell(12, 3).value
    src['基础属性%_防御'] = ws.cell(12, 4).value
    # R12 F: 偏向值(普通二级)
    src['偏向_普通二级'] = ws.cell(12, 6).value
    # R12 J: 偏向值(稀有)
    src['偏向_稀有'] = ws.cell(12, 10).value

    # 星3(R18): C6=全属性%, C7=二级属性
    src['星3_全属性%'] = ws.cell(18, 6).value
    src['星3_二级属性'] = ws.cell(18, 7).value
    # 星5(R20): C6=全属性%, C7=二级属性
    src['星5_全属性%'] = ws.cell(20, 6).value
    src['星5_二级属性'] = ws.cell(20, 7).value

    # 养成属性规划 R040
    ws2 = wb['养成属性规划']
    src['规划_基础属性%_生命'] = ws2.cell(40, 7).value
    src['规划_基础属性%_攻击'] = ws2.cell(40, 8).value
    src['规划_基础属性%_防御'] = ws2.cell(40, 9).value
    src['规划_偏向二级'] = ws2.cell(40, 16).value

    wb.close()
    return src


def generate(star_pct, gm):
    """计算升星派生值。"""
    result = {}

    # 基础属性% = 投放系数(直接)
    result['基础属性%_生命'] = star_pct
    result['基础属性%_攻击'] = star_pct
    result['基础属性%_防御'] = star_pct

    # 属性节点参数
    node_count = int(gm.get('升星属性节点数', 2))
    per_node_pct = float(gm.get('升星每节点全属性%', 0.15))
    per_node_sec = float(gm.get('升星每节点二级属性', 0.3))

    # 偏向普通二级 = 节点数 × 每节点二级
    result['偏向_普通二级'] = node_count * per_node_sec
    result['偏向_稀有'] = 0  # 当前未投放

    # 各节点投放明细
    result['星3_全属性%'] = per_node_pct
    result['星3_二级属性'] = per_node_sec
    result['星5_全属性%'] = per_node_pct
    result['星5_二级属性'] = per_node_sec

    # 养成属性规划行的值(应与上面一致)
    result['规划_基础属性%_生命'] = star_pct
    result['规划_基础属性%_攻击'] = star_pct
    result['规划_基础属性%_防御'] = star_pct
    result['规划_偏向二级'] = result['偏向_普通二级']

    return result


def compare(gen, src):
    """逐格对照。"""
    keys = [
        '基础属性%_生命', '基础属性%_攻击', '基础属性%_防御',
        '偏向_普通二级', '偏向_稀有',
        '星3_全属性%', '星3_二级属性',
        '星5_全属性%', '星5_二级属性',
        '规划_基础属性%_生命', '规划_基础属性%_攻击', '规划_基础属性%_防御',
        '规划_偏向二级',
    ]
    checks = []
    for k in keys:
        gv = gen.get(k)
        sv = src.get(k)
        match = almost(gv, sv)
        checks.append((k, gv, sv, match))
    return checks


def write_output(gen, checks, star_pct, gm):
    """写生成表。"""
    wb = openpyxl.Workbook()
    bold = Font(bold=True)
    hf = PatternFill('solid', fgColor='D9E1F2')
    okf = PatternFill('solid', fgColor='E2EFDA')
    title_font = Font(bold=True, size=13)

    ws = wb.active
    ws.title = '角色升星'

    ws.cell(1, 1, '角色升星 — 各星级投放总览')
    ws.cell(1, 1).font = title_font
    ws.cell(2, 1, '由 gen_star_upgrade.py 自动生成，勿手改。主要付费养成，解锁天赋+投放属性%。')

    # === 星级投放总表 ===
    r = 4
    headers = ['星级', '投放内容', '生命%', '攻击%', '防御%', '偏向二级属性']
    for c, h in enumerate(headers, 1):
        ws.cell(r, c, h); ws.cell(r, c).font = bold; ws.cell(r, c).fill = hf

    node_pct = float(gm.get('升星每节点全属性%', 0.15))
    node_sec = float(gm.get('升星每节点二级属性', 0.3))

    # 6个星级
    star_data = [
        (1, '天赋1', 0, 0, 0, 0),
        (2, '天赋2', 0, 0, 0, 0),
        (3, '全属性%+普通二级属性', node_pct, node_pct, node_pct, node_sec),
        (4, '天赋3', 0, 0, 0, 0),
        (5, '全属性%+普通二级属性', node_pct, node_pct, node_pct, node_sec),
        (6, '天赋4', 0, 0, 0, 0),
    ]
    for (star, desc, hp, atk, df, sec) in star_data:
        r += 1
        ws.cell(r, 1, star)
        ws.cell(r, 2, desc)
        ws.cell(r, 3, hp if hp else '—')
        ws.cell(r, 4, atk if atk else '—')
        ws.cell(r, 5, df if df else '—')
        ws.cell(r, 6, sec if sec else '—')

    # === 合计 ===
    r += 1
    ws.cell(r, 1, '合计(显式属性)'); ws.cell(r, 1).font = bold
    ws.cell(r, 2, '')
    total_pct = node_pct * 2  # 星3+星5
    total_sec = node_sec * 2
    ws.cell(r, 3, total_pct); ws.cell(r, 3).font = bold
    ws.cell(r, 4, total_pct); ws.cell(r, 4).font = bold
    ws.cell(r, 5, total_pct); ws.cell(r, 5).font = bold
    ws.cell(r, 6, total_sec); ws.cell(r, 6).font = bold

    r += 1
    ws.cell(r, 1, '投放系数(含天赋折算)')
    ws.cell(r, 2, ''); ws.cell(r, 3, star_pct)
    ws.cell(r, 4, '← 含4个天赋的战力贡献折算')

    # === 说明 ===
    r += 2
    ws.cell(r, 1, '说明'); ws.cell(r, 1).font = bold
    r += 1
    ws.cell(r, 1, '• 最终基础属性% = 投放系数(0.9)，含天赋带来的战中提升折算')
    r += 1
    ws.cell(r, 1, '• 星3/星5各投放一次属性节点：全属性%.15 + 偏向普通二级0.3')
    r += 1
    ws.cell(r, 1, '• 天赋节点(星1/2/4/6)为设计内容，不由数值公式生成')
    r += 1
    ws.cell(r, 1, '• 偏向稀有属性当前=0（未投放）')

    # 列宽
    ws.column_dimensions['A'].width = 8
    ws.column_dimensions['B'].width = 26
    ws.column_dimensions['C'].width = 10
    ws.column_dimensions['D'].width = 10
    ws.column_dimensions['E'].width = 10
    ws.column_dimensions['F'].width = 16

    # === 各角色升星属性选择 ===
    ws_h = wb.create_sheet('各角色升星属性选择')
    ws_h.append(['各角色升星偏向属性 -- 星3/星5各投放一种二级属性，每次0.3(攻击%等价)'])
    ws_h.append(['另外星3/星5还各给全属性%+0.15 (生命/攻击/防御各+15%)'])
    ws_h.append([])
    hdrs = ['角色', '品质', '升星属性1(3星)', '预算', '升星属性2(5星)', '预算']
    ws_h.append(hdrs)
    for c in range(1, len(hdrs)+1): ws_h.cell(4,c).font = bold; ws_h.cell(4,c).fill = hf
    try:
        wb_p = load_parameter_workbook(GROWTH_PARAM, data_only=True)
        ws_wj = wb_p['武将系数']
        sec = None
        for r in range(1, ws_wj.max_row+1):
            v = ws_wj.cell(r,1).value
            if v == '角色系数列表': sec = 'hero'; continue
            if v == '专武系数列表': break
            if sec == 'hero' and v and v != '武将名称':
                ws_h.append([ws_wj.cell(r,1).value, ws_wj.cell(r,2).value,
                             ws_wj.cell(r,17).value, 0.3,
                             ws_wj.cell(r,18).value, 0.3])
        wb_p.close()
    except Exception:
        pass
    for w,col in zip([12,8,18,8,18,8],'ABCDEF'): ws_h.column_dimensions[col].width = w

    wb.save(OUT)


def write_report(checks):
    """写对照报告。"""
    total = len(checks)
    passed = sum(1 for c in checks if c[3])

    lines = [
        '# P养成模块 · 角色升星 生成与对照',
        '',
        '> `gen_star_upgrade.py` 从共享参数+养成参数读取升星参数，计算基础属性%投放 + 偏向二级，与源表逐格对照。',
        '',
        '## 结论',
        '',
    ]
    if passed == total:
        lines.append('✅ **完全复现**：生成结果与现有手搓表逐格一致（%d/%d 格）。' % (passed, total))
    else:
        lines.append('⚠️ **存在偏差**：%d/%d 格一致，%d 格不一致。' % (passed, total, total - passed))

    lines += [
        '',
        '## 派生公式',
        '',
        '```',
        '基础属性%(生/攻/防) = 投放系数页「角色升星」= 0.9',
        '偏向普通二级 = 属性节点数(2) × 每节点二级投放(0.3) = 0.6',
        '各节点: 全属性%(0.15) + 二级属性(0.3)  [星3 + 星5 各一次]',
        '```',
        '',
        '## 逐格对照',
        '',
        '| 字段 | 生成值 | 源表值 | 一致 |',
        '|---|---|---|---|',
    ]
    for (key, gv, sv, match) in checks:
        lines.append('| %s | %s | %s | %s |' % (key, gv, sv, '✓' if match else '**✗**'))

    os.makedirs(os.path.dirname(REPORT), exist_ok=True)
    with open(REPORT, 'w', encoding='utf-8') as f:
        f.write('\n'.join(lines))


def main():
    print('=== 角色升星 属性派生验证 ===\n')

    star_pct, gm = read_params()
    print('参数读取: 升星基础属性%%=%s, 节点数=%s, 每节点全属性%%=%s, 每节点二级=%s' % (
        star_pct, gm.get('升星属性节点数'), gm.get('升星每节点全属性%'), gm.get('升星每节点二级属性')))

    gen = generate(star_pct, gm)
    print('生成: 基础属性%%=%.1f, 偏向普通二级=%.1f, 偏向稀有=%s' % (
        gen['基础属性%_生命'], gen['偏向_普通二级'], gen['偏向_稀有']))

    src = read_source()
    checks = compare(gen, src)

    total = len(checks)
    passed = sum(1 for c in checks if c[3])
    print('\n对照: %d/%d 格一致' % (passed, total))
    for (key, gv, sv, match) in checks:
        if not match:
            print('  ✗ %s: 生成=%s, 源=%s' % (key, gv, sv))

    write_output(gen, checks, star_pct, gm)
    print('\n[生成表] %s' % OUT)

    write_report(checks)
    print('[报告]   %s' % REPORT)

    if passed == total:
        print('\n✅ 完全复现：%d/%d 格逐格一致。' % (passed, total))
    else:
        print('\n⚠️ 存在偏差：%d/%d 格一致。' % (passed, total))


if __name__ == '__main__':
    main()
