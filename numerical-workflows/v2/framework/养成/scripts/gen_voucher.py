# -*- coding: utf-8 -*-
"""
gen_voucher.py — 旅券 主/副属性曲线 生成器 (养成系数6)

读取共享参数基线 MD（标准人/投放系数）+ 养成参数基线 MD（旅券参数）→ 生成旅券主/副属性传说基准曲线
及 配置页-旅券 各品质最终输出，与现有 养成数值 逐格对照。

链路:
  基础属性投放 = 标准人 × 旅券系数(投放系数页=6)  → 生命/攻击/防御 (30000/9000/6000)
  二级投放(4.2) → 主二级 = 权重占比×二级投放/主二级数量;  副二级 = 权重占比×二级投放/部位/词条
  传说主属性满级:  非直投 = 主二级数值 × 二级价值 / 二级基准 ('等级'类 ROUND0); 直投 = 对应投放
            基础 = 满级 × 主基础占比 ; 成长 = 满级 × (1-主基础占比)/上限
  传说副属性满级:  非直投 = 单词条数值 × 二级价值 / 二级基准 ; 直投 = 主同名满级/5
            1级  = 满级 × 副基础占比 ; 成长 = 满级 × (1-副基础占比)/上限 × 词条数 (直投 ROUND0)
  配置页主属性: 传说 = ROUND(旅券基础/成长, %类3位 否则 -1); 其余品质 = 传说 × 品质系数 (不再取整)
  配置页副属性: 传说 = ROUND(旅券副1级, %类3位 否则 0); 其余品质 = ROUND(传说 × 品质系数, 同精度)

只读源表。生成表为独立新文件。
用法: python gen_voucher.py [--check]
"""
import sys, io, os, argparse
from decimal import Decimal, ROUND_HALF_UP
import warnings
warnings.filterwarnings('ignore')
import openpyxl
from openpyxl.styles import Font, PatternFill

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

HERE = os.path.dirname(os.path.abspath(__file__))
GROWTH_DIR = os.path.dirname(HERE)  # 养成/
V2_DIR = os.path.dirname(GROWTH_DIR)  # 框架v2/
NUM_DIR = os.path.dirname(V2_DIR)  # 数值文档/数值文档/
sys.path.insert(0, V2_DIR)
from _shared.md_parameter_reader import load_parameter_workbook, parameter_md_path

SHARED_PARAM = parameter_md_path('共享')
GROWTH_PARAM = parameter_md_path('养成')
OUT     = os.path.join(GROWTH_DIR, '生成_旅券属性.xlsx')
REPORT  = os.path.join(GROWTH_DIR, '_reports', 'P养成模块_旅券属性_生成与对照.md')
SRC     = os.path.join(NUM_DIR, 'Vgame养成数值.xlsx')


def D(x):
    return Decimal(repr(x)) if not isinstance(x, Decimal) else x


def almost(a, b, tol=Decimal('0.0005')):
    if a is None or b is None:
        return a == b
    try:
        return abs(D(float(a)) - D(float(b))) <= tol
    except Exception:
        return a == b


def is_pct(name):
    return ('百分比' in name) or ('加成' in name) or (name == '治疗率')


def normalize_attack_name(v):
    if not isinstance(v, str):
        return v
    return (v.replace('物理攻击百分比', '攻击百分比')
             .replace('魔法攻击百分比', '攻击百分比')
             .replace('物理攻击加成', '攻击加成')
             .replace('魔法攻击加成', '攻击加成')
             .replace('物理攻击力', '攻击力')
             .replace('魔法攻击力', '攻击力')
             .replace('物理攻击%', '攻击%')
             .replace('魔法攻击%', '攻击%')
             .replace('物理攻击', '攻击')
             .replace('魔法攻击', '攻击')
             .replace('物攻', '攻击')
             .replace('魔攻', '攻击'))


def is_level(name):
    return name.endswith('等级')


def xround(v, prec):
    return D(v).quantize(Decimal(1).scaleb(-prec), rounding=ROUND_HALF_UP)


def read_params():
    # 标准人、投放系数 → 从共享参数读
    wb_s = load_parameter_workbook(SHARED_PARAM, data_only=True)
    sp = wb_s['标准人']
    std = {}
    for r in range(2, sp.max_row + 1):
        if sp.cell(r, 1).value:
            std[sp.cell(r, 1).value] = sp.cell(r, 2).value
    fp = wb_s['投放系数']
    voucher_coef = None
    for r in range(2, fp.max_row + 1):
        if fp.cell(r, 1).value == '旅券':
            voucher_coef = fp.cell(r, 2).value
    wb_s.close()

    # 旅券参数 → 从养成参数读
    wb_g = load_parameter_workbook(GROWTH_PARAM, data_only=True)
    ws = wb_g['旅券参数']
    scal = {}; quality = []; zhu = []; fu = []
    mode = None
    for r in range(1, ws.max_row + 1):
        a = ws.cell(r, 1).value
        if a == '参数': mode = 's'; continue
        if a == '品质值': mode = 'q'; continue
        if a == '主属性': mode = 'z'; continue
        if a == '副属性': mode = 'f'; continue
        if a is None: continue
        if mode == 's':
            scal[a] = ws.cell(r, 2).value
        elif mode == 'q':
            quality.append((int(a), ws.cell(r, 2).value, D(ws.cell(r, 3).value)))
        elif mode == 'z':
            zhu.append((normalize_attack_name(a), ws.cell(r, 2).value, normalize_attack_name(ws.cell(r, 3).value)))
        elif mode == 'f':
            fu.append((normalize_attack_name(a), ws.cell(r, 2).value, normalize_attack_name(ws.cell(r, 3).value)))
    wb_g.close()
    return std, voucher_coef, scal, quality, zhu, fu


def gen(std, vcoef, scal, zhu, fu):
    invest = {k: D(std[k]) * D(vcoef) for k in ('生命', '攻击', '防御')}
    direct_map = {'攻击': invest['攻击'], '防御': invest['防御'], '生命': invest['生命']}
    second = D(scal['二级属性投放'])
    base_price = D(scal['二级基准价格'])
    wz = D(scal['主属性权重']); wf = D(scal['副属性权重'])
    cap = D(scal['规划上限等级'])
    zhu_base = D(scal['主属性基础占比']); zhu_cnt = D(scal['主属性二级数量'])
    fu_base = D(scal['副属性基础占比'])
    parts = D(scal['部位数']); entries = D(scal['每部位词条数'])

    zhu_second = wz / (wz + wf) * second / zhu_cnt          # 主属性二级数值 C74
    fu_second = wf / (wz + wf) * second / parts / entries   # 单词条数值 C120
    zhu_grow_rate = (1 - zhu_base) / cap                    # C72
    fu_grow_rate = (1 - fu_base) / cap * entries            # C122

    # 主属性: name -> (满级C, 基础E, 成长F)
    zhu_out = {}
    for (name, price, direct) in zhu:
        if direct:
            sat = direct_map[direct]
        else:
            sat = zhu_second * D(price) / base_price
            if is_level(name):
                sat = xround(sat, 0)
        base = sat * zhu_base
        grow = sat * zhu_grow_rate
        zhu_out[name] = (sat, base, grow)

    # 副属性: name -> (满级C, 1级E, 成长F)
    fu_out = {}
    for (name, price, direct) in fu:
        if direct:
            sat = zhu_out[name][0] / 5          # 主同名满级/5
        else:
            sat = fu_second * D(price) / base_price
        lv1 = sat * fu_base
        grow = sat * fu_grow_rate
        if direct:
            lv1 = xround(lv1, 0); grow = xround(grow, 0)
        fu_out[name] = (sat, lv1, grow)

    return zhu_out, fu_out


# ---------- 对照: 旅券 sheet ----------
def parity_voucher_sheet(zhu_out, fu_out):
    wb = openpyxl.load_workbook(SRC, data_only=True); ws = wb['旅券']
    fails = []; checks = 0
    # 主属性 rows 94-111 : B名 C满级 E基础 F成长
    for r in range(94, 112):
        name = normalize_attack_name(ws.cell(r, 2).value)
        if name not in zhu_out: continue
        sat, base, grow = zhu_out[name]
        for val, col, lbl in [(sat, 3, '满级'), (base, 5, '基础'), (grow, 6, '成长')]:
            act = ws.cell(r, col).value
            checks += 1
            if not almost(val, act):
                fails.append(('旅券主 R%d %s %s' % (r, name, lbl), float(val), act))
    # 副属性 rows 126-135 : B名 C满级 E1级 F成长
    for r in range(126, 136):
        name = normalize_attack_name(ws.cell(r, 2).value)
        if name not in fu_out: continue
        sat, lv1, grow = fu_out[name]
        for val, col, lbl in [(sat, 3, '满级'), (lv1, 5, '1级'), (grow, 6, '成长')]:
            act = ws.cell(r, col).value
            checks += 1
            if not almost(val, act):
                fails.append(('旅券副 R%d %s %s' % (r, name, lbl), float(val), act))
    wb.close()
    return checks, fails


# ---------- 对照: 配置页-旅券 ----------
def parity_config(zhu_out, fu_out, quality):
    wb = openpyxl.load_workbook(SRC, data_only=True); ws = wb['配置页-旅券']
    qmap = {q: coef for (q, n, coef) in quality}   # 品质值->系数
    fails = []; checks = 0

    # 主属性: row3 表头每属性占两列(基础,成长); 行 4..7 品质 2..5
    # 找到每个属性的起始列(name在偶数列, 重复名列=成长)
    cols = []  # (col_base, name)
    c = 2
    maxc = ws.max_column
    while c <= maxc:
        nm = normalize_attack_name(ws.cell(3, c).value)
        if nm is None: c += 1; continue
        cols.append((c, nm)); c += 2
    qrows = {ws.cell(r, 1).value: r for r in range(4, 8)}  # 品质值->行
    for (cb, name) in cols:
        if name not in zhu_out: continue
        prec = 3 if is_pct(name) else -1
        _, base_e, grow_f = zhu_out[name]
        legend_base = xround(base_e, prec)   # 传说(品质5)
        legend_grow = xround(grow_f, prec)
        for q, row in qrows.items():
            coef = qmap[q]
            if q == 5:
                exp_b, exp_g = legend_base, legend_grow
            else:
                exp_b = legend_base * coef
                exp_g = legend_grow * coef
            ab = ws.cell(row, cb).value
            ag = ws.cell(row, cb + 1).value
            for val, act, lbl in [(exp_b, ab, '基础'), (exp_g, ag, '成长')]:
                checks += 1
                if not almost(val, act):
                    fails.append(('配置主 品质%d %s %s' % (q, name, lbl), float(val), act))

    # 副属性: row12 表头(6属性单列), 行13..16 品质2..5
    fu_cols = []
    for c in range(2, ws.max_column + 1):
        nm = normalize_attack_name(ws.cell(12, c).value)
        if nm and nm in fu_out:
            fu_cols.append((c, nm))
    qrows2 = {ws.cell(r, 1).value: r for r in range(13, 17)}
    for (c, name) in fu_cols:
        prec = 3 if is_pct(name) else 0
        _, lv1, _ = fu_out[name]
        legend = xround(lv1, prec)
        for q, row in qrows2.items():
            coef = qmap[q]
            exp = legend if q == 5 else xround(legend * coef, prec)
            act = ws.cell(row, c).value
            checks += 1
            if not almost(exp, act):
                fails.append(('配置副 品质%d %s' % (q, name), float(exp), act))
    wb.close()
    return checks, fails


def write_out(zhu_out, fu_out, quality):
    wb = openpyxl.Workbook(); bold = Font(bold=True); hf = PatternFill('solid', fgColor='D9E1F2')
    title_font = Font(bold=True, size=13)

    # Sheet 1: 总览
    ws0 = wb.active; ws0.title = '旅券总览'
    ws0.cell(1, 1, '旅券 — 养成总览'); ws0.cell(1, 1).font = title_font
    ws0.cell(2, 1, '由 gen_voucher.py 自动生成。6件套装，4+2套装效果，投放基础属性+二级属性。')

    r = 4
    ws0.cell(r, 1, '体系结构'); ws0.cell(r, 1).font = bold
    r += 1
    for c, h in enumerate(['项目', '值', '说明'], 1):
        ws0.cell(r, c, h); ws0.cell(r, c).font = bold; ws0.cell(r, c).fill = hf
    info = [
        ('套装件数', 6, '4+2套装，2件套=属性，4件套=效果'),
        ('部位', 6, '每部位2条词条'),
        ('主属性数', len(zhu_out), '二级属性+基础属性(直投)，攻击已统一'),
        ('副属性数', len(fu_out), '暴击/暴伤/攻防生%/攻防生绝对值，攻击已统一'),
        ('品质', '普通/稀有/史诗/传说', '系数 0.4/0.6/0.8/1'),
        ('等级上限', 5, ''),
        ('基础属性投放', '生命30000/攻击9000/防御6000', '标准人×投放系数6'),
        ('二级属性投放', 4.2, '主副按权重32:68分配'),
        ('套装被动提升', '30%', '6件总战斗提升期望'),
    ]
    for (item, val, desc) in info:
        r += 1
        ws0.cell(r, 1, item); ws0.cell(r, 2, val); ws0.cell(r, 3, desc)

    r += 2
    ws0.cell(r, 1, '品质系数'); ws0.cell(r, 1).font = bold
    r += 1
    for c, h in enumerate(['品质', '系数', '对应品质值'], 1):
        ws0.cell(r, c, h); ws0.cell(r, c).font = bold; ws0.cell(r, c).fill = hf
    for (qv, qn, qc) in quality:
        r += 1
        ws0.cell(r, 1, qn); ws0.cell(r, 2, qc); ws0.cell(r, 3, qv)

    ws0.column_dimensions['A'].width = 16
    ws0.column_dimensions['B'].width = 30
    ws0.column_dimensions['C'].width = 36

    # 部位主属性分配
    r += 2
    ws0.cell(r, 1, '各部位主属性池'); ws0.cell(r, 1).font = bold
    r += 1
    for c, h in enumerate(['部位', '可选主属性(权重)'], 1):
        ws0.cell(r, c, h); ws0.cell(r, c).font = bold; ws0.cell(r, c).fill = hf
    pos_data = [
        ('部位1', '攻击(200)'),
        ('部位2', '生命(100)'),
        ('部位3', '防御(100)'),
        ('部位4', '攻击%(400), 生命%(200), 防御%(200), 必杀回复等级(100)'),
        ('部位5', '攻击%(400), 生命%(200), 防御%(200)'),
        ('部位6', '攻击%(400), 生命%(200), 防御%(200), 暴击等级(100), 暴伤等级(100)'),
    ]
    for (pos, attrs) in pos_data:
        r += 1
        ws0.cell(r, 1, pos); ws0.cell(r, 2, attrs)

    # Sheet 2: 主属性
    ws = wb.create_sheet('主属性(传说基准)')
    ws.append(['主属性 传说基准曲线 — 等级上限%d级' % 5])
    ws.append(['公式: 面板属性 = 基础 + (等级-1) × 成长; 其余品质 = 传说 × 品质系数'])
    ws.append([])
    ws.append(['主属性', '满级(Lv5)', '基础(Lv1)', '成长/级'])
    for c in range(1, 5): ws.cell(4, c).font = bold; ws.cell(4, c).fill = hf
    for name, (sat, base, grow) in zhu_out.items():
        ws.append([name, float(sat), float(base), float(grow)])
    for w, col in zip([20, 14, 14, 14], 'ABCD'):
        ws.column_dimensions[col].width = w

    # Sheet 3: 副属性
    ws2 = wb.create_sheet('副属性(传说基准)')
    ws2.append(['副属性 传说基准曲线 — 等级上限%d级' % 5])
    ws2.append(['公式: 面板属性 = 1级 + (等级-1) × 成长; 其余品质 = 传说 × 品质系数'])
    ws2.append([])
    ws2.append(['副属性', '满级(Lv5)', '1级', '成长/级'])
    for c in range(1, 5): ws2.cell(4, c).font = bold; ws2.cell(4, c).fill = hf
    for name, (sat, lv1, grow) in fu_out.items():
        ws2.append([name, float(sat), float(lv1), float(grow)])
    for w, col in zip([20, 14, 14, 14], 'ABCD'):
        ws2.column_dimensions[col].width = w

    wb.save(OUT)


def main():
    ap = argparse.ArgumentParser(); ap.add_argument('--check', action='store_true'); args = ap.parse_args()
    std, vcoef, scal, quality, zhu, fu = read_params()
    zhu_out, fu_out = gen(std, vcoef, scal, zhu, fu)
    vc, vf = parity_voucher_sheet(zhu_out, fu_out)
    cc, cf = parity_config(zhu_out, fu_out, quality)
    if not args.check:
        write_out(zhu_out, fu_out, quality)

    md = []
    def o(s=''):
        print(s); md.append(s)
    o('# 养成模块 · 旅券 主/副属性 生成与对照'); o()
    o('> 主/副二级分配×属性定价→传说基准曲线; 配置页其余品质=传说×品质系数。源表只读。'); o()
    o('## 旅券 sheet 传说基准 (主 行94-111 / 副 行126-135)')
    o('- 对照 %d 格，不一致 %d' % (vc, len(vf)))
    if vf:
        o('| 位置 | 生成 | 源 |'); o('|---|--:|--:|')
        for (loc, v, a) in vf[:25]:
            o('| %s | %s | %s |' % (loc, v, a))
    else:
        o('- ✅ **旅券基准曲线逐格复现** (主%d属性 + 副%d属性 × 满级/基础(1级)/成长)' % (len(zhu_out), len(fu_out)))
    o()
    o('## 配置页-旅券 各品质最终输出')
    o('- 对照 %d 格，不一致 %d' % (cc, len(cf)))
    if cf:
        o('| 位置 | 生成 | 源 |'); o('|---|--:|--:|')
        for (loc, v, a) in cf[:25]:
            o('| %s | %s | %s |' % (loc, v, a))
    else:
        o('- ✅ **配置页最终输出逐格复现** (主%d属性×4品质×基础/成长 + 副6属性×4品质)' % len(zhu_out))
    o()
    o('**合计对照 %d 格，不一致 %d**' % (vc + cc, len(vf) + len(cf)))
    os.makedirs(os.path.dirname(REPORT), exist_ok=True)
    with open(REPORT, 'w', encoding='utf-8') as f:
        f.write('\n'.join(md))
    print('\n[报告] %s' % REPORT)
    if not args.check:
        print('[生成表] %s' % OUT)


if __name__ == '__main__':
    main()
