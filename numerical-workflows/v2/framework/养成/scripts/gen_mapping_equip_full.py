# -*- coding: utf-8 -*-
"""从 01_源表与 03_沉淀MD 重建 19 页映射装备数值总表。"""

from __future__ import annotations

import argparse
import hashlib
import os
import re
import sys
from datetime import date
from pathlib import Path

import openpyxl
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


HERE = Path(__file__).resolve().parent
GROWTH_DIR = HERE.parent
V2_DIR = GROWTH_DIR.parent
SOURCE_DIR = V2_DIR / "01_源表"
MD_DIR = V2_DIR / "03_沉淀MD"

GROWTH_MD = MD_DIR / "养成" / "映射装备养成数值沉淀-2026-07-07.md"
ECON_MD = MD_DIR / "经济" / "21_养成_映射装备经济消耗沉淀-2026-07-09.md"
DUNGEON_MD = MD_DIR / "经济" / "05_玩法_映射副本经济沉淀-2026-07-08.md"
TOWER_MD = MD_DIR / "经济" / "08_玩法_爬塔经济沉淀-2026-07-09.md"
SCENARIO_MD = MD_DIR / "养成" / "映射装备数值总表生成参数-2026-07-10.md"

SOURCE_FILES = [
    SOURCE_DIR / "属性定价.xlsx",
    SOURCE_DIR / "养成属性分配.xlsx",
    SOURCE_DIR / "道具定价.xlsx",
    SOURCE_DIR / "经济口径.xlsx",
]
MD_FILES = [GROWTH_MD, ECON_MD, DUNGEON_MD, TOWER_MD, SCENARIO_MD]

OUT = GROWTH_DIR / "生成_映射装备数值总表.xlsx"
REPORT = GROWTH_DIR / "_reports" / "P映射装备数值总表_生成与对照.md"
LEGACY_FINAL = V2_DIR / "04_数值文档" / "大模块" / "映射装备数值总表_数值文档.xlsx"

SHEET_ORDER = [
    "00_口径", "01_循环总览", "02_品质等级", "03_主属性基础", "04_主属性计算",
    "05_副属性基础", "06_副属性计算", "07_副本宝箱", "08_副本材料", "09_日产验算",
    "10_爬塔挂机", "11_锻造台经验", "12_合成消耗", "13_合成概率", "14_打造消耗",
    "15_打造概率", "16_分解回收", "17_边界待确认", "18_产销汇总",
]

BLUE = "0000FF"
BLACK = "000000"
WHITE = "FFFFFF"
HEADER = "1F4E78"
SUBHEADER = "D9EAF7"
SECTION = "D9E1F2"
INPUT_FILL = "FFF2CC"
THIN = Side(style="thin", color="D9E2F3")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
FONT_NAME = "Microsoft YaHei"


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def split_md_row(line: str) -> list[str]:
    body = line.strip().strip("|")
    return [part.replace("\\|", "|").strip() for part in re.split(r"(?<!\\)\|", body)]


def is_separator(line: str) -> bool:
    if not line.strip().startswith("|"):
        return False
    cells = split_md_row(line)
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", c.replace(" ", "")) for c in cells)


def parse_tables(path: Path) -> list[tuple[list[str], list[list[str]], int]]:
    lines = path.read_text(encoding="utf-8-sig").splitlines()
    tables = []
    i = 0
    while i + 1 < len(lines):
        if lines[i].lstrip().startswith("|") and is_separator(lines[i + 1]):
            headers = split_md_row(lines[i])
            rows = []
            j = i + 2
            while j < len(lines) and lines[j].lstrip().startswith("|"):
                row = split_md_row(lines[j])
                row += [""] * (len(headers) - len(row))
                rows.append(row[: len(headers)])
                j += 1
            tables.append((headers, rows, i + 1))
            i = j
        else:
            i += 1
    return tables


def find_table(tables, headers: list[str], occurrence: int = 0) -> list[dict[str, str]]:
    matches = []
    for actual, rows, _line in tables:
        if actual == headers:
            matches.append([dict(zip(actual, row)) for row in rows])
    if occurrence >= len(matches):
        raise KeyError("未找到 Markdown 表: %s occurrence=%d" % (headers, occurrence))
    return matches[occurrence]


def compatible_table(tables, headers: list[str], replacements: dict[str, str] | None = None, optional: bool = False) -> list[dict[str, str]]:
    replacements = replacements or {}
    for actual, rows, _line in tables:
        canonical = [replacements.get(header, header) for header in actual]
        if all(header in canonical for header in headers):
            return [{key: value for key, value in zip(canonical, row)} for row in rows]
    if optional:
        return []
    raise KeyError("未找到当前版本 Markdown 表: %s" % headers)


def clean(value) -> str:
    return str(value).strip().strip("`").replace("，", ",")


def number(value):
    text = clean(value).replace(",", "")
    if text in ("", "-", "—"):
        return None
    if text.endswith("%"):
        return float(text[:-1]) / 100
    if re.fullmatch(r"-?\d+", text):
        return int(text)
    if re.fullmatch(r"-?\d+(?:\.\d+)?", text):
        return float(text)
    raise ValueError("不是数值: %r" % value)


def percent_literal(value: float) -> str:
    """把 0.08 还原为 Excel 可读的 8%，避免把沉淀口径降成小数常量。"""
    percent = value * 100
    return ("%g" % percent) + "%"


def probs(value: str, count: int) -> list[float]:
    values = [number(part) or 0 for part in re.split(r"\s*/\s*", clean(value))]
    if len(values) != count:
        raise ValueError("概率列数量错误: %s -> %s" % (value, values))
    return values


def ints(value: str) -> list[int]:
    return [int(v) for v in re.findall(r"\d+", clean(value))]


def table_data() -> dict:
    gt = parse_tables(GROWTH_MD)
    et = parse_tables(ECON_MD)
    dt = parse_tables(DUNGEON_MD)
    tt = parse_tables(TOWER_MD)
    st = parse_tables(SCENARIO_MD)
    data = {
        "qualities": find_table(gt, ["品质顺序", "品质"]),
        "quality_coef": find_table(gt, ["品质", "主属性品质系数", "副属性品质系数"]),
        "parts": find_table(gt, ["部位", "主属性", "口径"]),
        "main_level": find_table(gt, ["等级", "主属性等级倍率"]),
        "main_target": compatible_table(gt, ["主属性", "目标值", "说明"]),
        "sub_level": find_table(gt, ["等级", "副属性等级倍率"]),
        "sub_target": compatible_table(gt, ["副属性", "60级红装单条目标值", "换算说明"], {"60级红装单条上限": "60级红装单条目标值"}),
        "forge_rules": find_table(et, ["项", "最新口径"], 0),
        "equip_exp": find_table(et, ["装备等级", "绿", "蓝", "紫", "金", "红"], 0),
        "synthesis_tier": find_table(et, ["合成档", "角色等级要求", "产物等级池", "等级权重", "投入装备等级下限", "基础金币"]),
        "quality_gold_synthesis": find_table(et, ["品质", "绿", "蓝", "紫", "金", "红"], 0),
        "forge_tier": compatible_table(et, ["打造档", "打造档位等级", "可打造装备等级", "金币(暂沿用旧档)", "紫主材每种副材", "金主材每种副材", "红主材每种副材"], {"成品等级概率池": "可打造装备等级"}),
        "decompose_base": find_table(et, ["装备等级", "5", "10", "15", "20", "25", "30", "35", "40", "45", "50", "55", "60"]),
        "quality_gold_decompose": find_table(et, ["品质", "绿", "蓝", "紫", "金", "红"], 1),
        "dungeon_box": find_table(dt, ["层", "消耗", "固定宝箱", "额外宝箱概率", "期望宝箱", "宝箱档位概率(绿/蓝/紫/金/红)", "装备等级池"]),
        "dungeon_material": find_table(dt, ["层", "主材等级", "主材掉率", "主材品质权重(紫/金/红)", "副材总量", "单种副材期望"]),
        "daily_equipment": find_table(dt, ["刷层", "装备/日", "绿装/日", "蓝装/日", "紫装/日", "金装/日", "红装/日", "全分解金币上限/日"]),
        "daily_material": find_table(dt, ["阶段", "代表刷层", "次数/日", "主材等级", "主材总/日", "紫主材/日", "金主材/日", "红主材/日", "副材总/日", "单种副材/日"]),
        "daily_material_class": find_table(dt, ["代表刷层", "武器主材/日", "防具主材/日", "首饰主材/日"]),
        "tower_gold": compatible_table(tt, ["最高层", "金币/日", "房车经验/日"], {"房车零件/日": "房车经验/日"}),
        "tower_material": compatible_table(tt, ["最高层", "主材阶位", "主材总/日", "紫/金/红概率", "紫期望/日", "金期望/日", "红期望/日", "副材总/日", "单种副材/日"], optional=True),
        "scenario": find_table(st, ["阶段", "副本代表层", "副本次数/日", "爬塔金币最高层", "爬塔材料区间", "合成档", "打造档位等级", "备注"]),
    }
    econ_text = ECON_MD.read_text(encoding="utf-8-sig")
    m = re.search(r"按 10 取整\(([-\d.]+) \+ ([-\d.]+)n \+ ([-\d.]+)n\^2\)", econ_text)
    if m:
        data["forge_exp_formula"] = tuple(float(v) for v in m.groups())
        data["forge_exp_values"] = []
    else:
        curve = compatible_table(et, ["等级段", "本段经验", "总经验占比", "逐级 `CostExp`（从段首等级依次升级）"])
        values = []
        for row in curve:
            values.extend(int(value.replace(",", "")) for value in re.findall(r"\d+(?:,\d{3})?", row["逐级 `CostExp`（从段首等级依次升级）"]))
        if len(values) != 59 or sum(values) != 94500:
            raise ValueError("锻造台离散经验曲线必须为59级且合计94500")
        data["forge_exp_formula"] = None
        data["forge_exp_values"] = values
    formulas = {}
    for name in ("绿合紫", "蓝合金", "紫合金", "紫合红", "金合红", "紫主材出金", "金主材出红"):
        m = re.search(re.escape(name) + r" = ([\d.]+)% \+ ([\d.]+)%", econ_text)
        if not m:
            raise ValueError("未解析到概率公式: %s" % name)
        formulas[name] = tuple(float(v) / 100 for v in m.groups())
    data["prob_formulas"] = formulas
    return data


def source_evidence(data: dict) -> dict:
    for path in SOURCE_FILES + MD_FILES:
        if not path.exists():
            raise FileNotFoundError(path)

    wb = load_workbook(SOURCE_DIR / "养成属性分配.xlsx", data_only=True, read_only=True)
    ws = wb["模块规划"]
    mapping_row = None
    for row in ws.iter_rows(values_only=True):
        if len(row) > 1 and row[1] == "映射装备":
            mapping_row = row
            break
    wb.close()
    if mapping_row is None:
        raise ValueError("养成属性分配缺少映射装备模块")
    allocation = {"生命": mapping_row[3], "攻击": mapping_row[4], "防御": mapping_row[5]}

    targets = {}
    for row in data["main_target"]:
        targets[row["主属性"]] = targets.get(row["主属性"], 0) + number(row["目标值"])
    for name in ("生命", "攻击", "防御"):
        expected = allocation[name] * 0.75
        if abs(targets[name] - expected) > 1e-6:
            raise ValueError("%s目标值未通过01源表校验: MD=%s, 01*75%%=%s" % (name, targets[name], expected))

    wb = load_workbook(SOURCE_DIR / "属性定价.xlsx", data_only=True, read_only=True)
    ws = wb.active
    attr_prices = {ws.cell(r, 1).value: ws.cell(r, 2).value for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value}
    wb.close()
    for name in ("攻击百分比", "防御百分比", "生命百分比"):
        if name not in attr_prices:
            raise ValueError("属性定价缺少: %s" % name)

    wb = load_workbook(SOURCE_DIR / "道具定价.xlsx", data_only=True, read_only=True)
    ws = wb.active
    gold = None
    for r in range(1, ws.max_row + 1):
        if ws.cell(r, 3).value == "金币":
            gold = {"id": ws.cell(r, 4).value, "rmb": ws.cell(r, 5).value, "diamond": ws.cell(r, 6).value}
            break
    wb.close()
    if not gold or gold["id"] != 1:
        raise ValueError("道具定价未确认金币 ItemId=1")

    wb = load_workbook(SOURCE_DIR / "经济口径.xlsx", data_only=True, read_only=True)
    ws = wb["经济口径说明"]
    ruling = None
    for r in range(1, ws.max_row + 1):
        if ws.cell(r, 1).value == "映射装备裁决":
            ruling = ws.cell(r, 2).value
            break
    wb.close()
    if not ruling or "升级、打孔、鉴定已删除" not in ruling:
        raise ValueError("经济口径未确认映射装备废弃边界")

    return {
        "allocation": allocation,
        "mapping_coefficient": mapping_row[2],
        "attr_price_count": len(attr_prices),
        "gold": gold,
        "economy_ruling": ruling,
        "hashes": {os.fspath(path.relative_to(V2_DIR)): sha256(path) for path in SOURCE_FILES + MD_FILES},
    }


def add_rows(ws, rows) -> None:
    for row in rows:
        ws.append(list(row))


def style_sheet(ws, header_rows=(1,), source_note="") -> None:
    for row in ws.iter_rows():
        for cell in row:
            cell.font = Font(name=FONT_NAME, size=10, color=BLACK)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if cell.value is not None:
                cell.border = BORDER
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    cell.font = Font(name=FONT_NAME, size=10, color=BLACK)
                elif cell.row not in header_rows:
                    cell.font = Font(name=FONT_NAME, size=10, color=BLUE)
    for row_idx in header_rows:
        for cell in ws[row_idx]:
            if cell.value is not None:
                cell.fill = PatternFill("solid", fgColor=HEADER)
                cell.font = Font(name=FONT_NAME, size=10, bold=True, color=WHITE)
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                if source_note:
                    cell.comment = Comment(source_note, "Codex")
    ws.freeze_panes = "A2"
    if ws.max_row > 1 and ws.max_column > 1:
        ws.auto_filter.ref = ws.dimensions
    for col in range(1, ws.max_column + 1):
        values = [str(ws.cell(r, col).value or "") for r in range(1, min(ws.max_row, 80) + 1)]
        width = min(max(max((len(v) for v in values), default=8) + 2, 10), 36)
        ws.column_dimensions[get_column_letter(col)].width = width


def build_workbook(data: dict, evidence: dict) -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet("00_口径")
    rows = [
        ["项", "值", "备注"],
        ["生成日期", date.today().isoformat(), ""],
        ["生成物", os.fspath(OUT), "运行时不读取04终稿"],
        ["口径", "01_源表提供全局事实校验；03_沉淀MD提供系统参数与公式", ""],
        ["公式说明", "派生项使用 Excel 公式；蓝字为源值/输入，黑字为公式", ""],
    ]
    for path in SOURCE_FILES:
        rel = path.relative_to(V2_DIR)
        rows.append(["源表", os.fspath(rel), evidence["hashes"][os.fspath(rel)][:16]])
    for path in MD_FILES:
        rel = path.relative_to(V2_DIR)
        rows.append(["源沉淀MD", os.fspath(rel), evidence["hashes"][os.fspath(rel)][:16]])
    rows += [
        ["01校验", "映射装备投放系数=%s；生命/攻击/防御=%s" % (evidence["mapping_coefficient"], evidence["allocation"]), "三项60红目标=01投放的75%"],
        ["金币锚点", "金币(%s)，钻石定价=%s" % (evidence["gold"]["id"], evidence["gold"]["diamond"]), "来自01_源表/道具定价.xlsx"],
        ["经济边界", evidence["economy_ruling"], "来自01_源表/经济口径.xlsx"],
        ["新增汇总", "18_产销汇总：按前/中/后期汇总日产出、打造/合成承载和金币压力", "阶段假设维护在生成参数MD"],
    ]
    add_rows(ws, rows)
    style_sheet(ws, source_note="来源链与文件哈希")

    ws = wb.create_sheet("01_循环总览")
    add_rows(ws, [
        ["环节", "投入/对象", "系统", "产出/去向", "备注"],
        ["副本入场", "映射之匙(25)x1", "映射装备副本", "每日/周期由钥匙投放控制", "进入消耗，不归养成消耗"],
        ["装备产出", "宝箱", "副本通关掉落", "装备进入穿戴/合成/分解", "红宝箱必出红装"],
        ["打造材料产出", "主材/副材", "映射装备副本常驻", "打造指定等级/部位装备", "爬塔新版不再投放装备材料；主材决定打造档位和品质池"],
        ["装备消耗", "9件同品质装备 + 金币", "合成", "随机部位，品质受锻造台影响", "合成给锻造台经验"],
        ["材料消耗", "主材x1 + 2种副材 + 金币", "打造", "指定档内等级/部位，品质受锻造台影响", "打造不给锻造台经验"],
        ["回收", "非目标装备", "分解", "金币 + 锻造台经验；红装额外高阶零件(23)", "洗练本版不做"],
        ["沉淀", "高阶零件(23)", "红装分解", "后续版本装备高级材料储备", "当前无消耗"],
    ])
    style_sheet(ws, source_note=ECON_MD.name + "；" + DUNGEON_MD.name)

    quality_order = {r["品质"]: number(r["品质顺序"]) for r in data["qualities"]}
    quality_coef = {r["品质"]: (number(r["主属性品质系数"]), number(r["副属性品质系数"])) for r in data["quality_coef"]}
    syn_gold = {k: number(v) for k, v in data["quality_gold_synthesis"][0].items() if k != "品质"}
    dec_gold = {k: number(v) for k, v in data["quality_gold_decompose"][0].items() if k != "品质"}
    qualities = [r["品质"] for r in data["qualities"]]

    ws = wb.create_sheet("02_品质等级")
    add_rows(ws, [["品质", "品质顺序", "主属性品质系数", "副属性品质系数", "合成金币系数", "分解金币系数"]] + [
        [q, quality_order[q], quality_coef[q][0], quality_coef[q][1], syn_gold[q], dec_gold[q]] for q in qualities
    ])
    style_sheet(ws, source_note=GROWTH_MD.name + "；" + ECON_MD.name)

    main_levels = [(number(r["等级"]), number(r["主属性等级倍率"])) for r in data["main_level"]]
    main_targets = [
        (
            r["主属性"],
            number(r["目标值"]),
            "比例" if "%" in r["目标值"] else ("等级" if r["主属性"].endswith("等级") else "数值"),
        )
        for r in data["main_target"]
    ]
    ws = wb.create_sheet("03_主属性基础")
    ws.append(["部位", "主属性", "口径", "", "主属性", "60级红装目标值", "单位", "", "等级", "主属性等级倍率"])
    for i in range(max(len(data["parts"]), len(main_targets), len(main_levels))):
        part = data["parts"][i] if i < len(data["parts"]) else None
        target = main_targets[i] if i < len(main_targets) else None
        level = main_levels[i] if i < len(main_levels) else None
        ws.append([
            part["部位"] if part else "", part["主属性"] if part else "", part["口径"] if part else "", "",
            target[0] if target else "", target[1] if target else "", target[2] if target else "", "",
            level[0] if level else "", level[1] if level else "",
        ])
    style_sheet(ws, source_note=GROWTH_MD.name)
    for row in range(2, ws.max_row + 1):
        if ws.cell(row, 7).value == "比例":
            ws.cell(row, 6).number_format = "0.00%"

    ws = wb.create_sheet("04_主属性计算")
    ws.append(["部位", "主属性", "品质", "等级", "60红目标", "主属性品质系数", "主属性等级倍率", "主属性值"])
    for part in data["parts"]:
        for quality in qualities:
            for level, _ratio in main_levels:
                r = ws.max_row + 1
                ws.append([
                    part["部位"],
                    "=INDEX('03_主属性基础'!$B$2:$B$9,MATCH(A%d,'03_主属性基础'!$A$2:$A$9,0))" % r,
                    quality, level,
                    "=INDEX('03_主属性基础'!$F$2:$F$9,MATCH(B%d,'03_主属性基础'!$E$2:$E$9,0))" % r,
                    "=INDEX('02_品质等级'!$C$2:$C$6,MATCH(C%d,'02_品质等级'!$A$2:$A$6,0))" % r,
                    "=INDEX('03_主属性基础'!$J$2:$J$13,MATCH(D%d,'03_主属性基础'!$I$2:$I$13,0))" % r,
                    "=E%d*F%d*G%d" % (r, r, r),
                ])
    style_sheet(ws, source_note=GROWTH_MD.name)

    sub_levels = [(number(r["等级"]), number(r["副属性等级倍率"])) for r in data["sub_level"]]
    sub_targets = [
        (
            r["副属性"],
            number(r["60级红装单条目标值"]),
            "比例" if "%" in r["60级红装单条目标值"] else ("等级" if r["副属性"].endswith("等级") else "数值"),
            r["换算说明"].replace(",", ""),
        )
        for r in data["sub_target"]
    ]
    ws = wb.create_sheet("05_副属性基础")
    ws.append(["副属性", "60级红装单条目标值", "单位", "换算说明", "", "等级", "副属性等级倍率"])
    for i in range(max(len(sub_targets), len(sub_levels))):
        target = sub_targets[i] if i < len(sub_targets) else None
        level = sub_levels[i] if i < len(sub_levels) else None
        ws.append([
            target[0] if target else "", target[1] if target else "", target[2] if target else "", target[3] if target else "", "",
            level[0] if level else "", level[1] if level else "",
        ])
    style_sheet(ws, source_note=GROWTH_MD.name + "；01_源表/属性定价.xlsx")
    for row in range(2, ws.max_row + 1):
        if ws.cell(row, 3).value == "比例":
            ws.cell(row, 2).number_format = "0.00%"

    ws = wb.create_sheet("06_副属性计算")
    ws.append(["副属性", "品质", "等级", "60红单条目标", "副属性品质系数", "副属性等级倍率", "副属性值"])
    for attr, _target, _unit, _note in sub_targets:
        for quality in qualities:
            for level, _ratio in sub_levels:
                r = ws.max_row + 1
                ws.append([
                    attr, quality, level,
                    "=INDEX('05_副属性基础'!$B$2:$B$13,MATCH(A%d,'05_副属性基础'!$A$2:$A$13,0))" % r,
                    "=INDEX('02_品质等级'!$D$2:$D$6,MATCH(B%d,'02_品质等级'!$A$2:$A$6,0))" % r,
                    "=INDEX('05_副属性基础'!$G$2:$G$13,MATCH(C%d,'05_副属性基础'!$F$2:$F$13,0))" % r,
                    "=D%d*E%d*F%d" % (r, r, r),
                ])
    style_sheet(ws, source_note=GROWTH_MD.name)

    ws = wb.create_sheet("07_副本宝箱")
    ws.append(["层", "消耗", "固定宝箱", "额外宝箱概率", "期望宝箱", "绿宝箱概率", "蓝宝箱概率", "紫宝箱概率", "金宝箱概率", "红宝箱概率", "装备等级池"])
    for row in data["dungeon_box"]:
        p = probs(row["宝箱档位概率(绿/蓝/紫/金/红)"], 5)
        ws.append([number(row["层"]), row["消耗"], number(row["固定宝箱"]), number(row["额外宝箱概率"]), number(row["期望宝箱"]), *p, row["装备等级池"]])
    style_sheet(ws, source_note=DUNGEON_MD.name)
    for row in range(2, ws.max_row + 1):
        for col in range(4, 11):
            ws.cell(row, col).number_format = "0.00%" if col != 5 else "0.00"

    forge_rows = data["forge_tier"]
    roman = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII"]
    def forge_for_level(level: int):
        for idx, row in enumerate(forge_rows):
            if level in ints(row["可打造装备等级"]):
                return roman[idx], number(row["打造档位等级"]), row["可打造装备等级"]
        raise KeyError("主材等级没有打造档: %s" % level)

    ws = wb.create_sheet("08_副本材料")
    ws.append(["层", "主材阶位", "打造档位等级", "可打造装备等级", "主材掉率", "紫主材概率", "金主材概率", "红主材概率", "副材总量/次", "单种副材期望/次"])
    for row in data["dungeon_material"]:
        level = number(row["主材等级"])
        tier, build_level, craft_levels = forge_for_level(level)
        p = probs(row["主材品质权重(紫/金/红)"], 3)
        r = ws.max_row + 1
        ws.append([number(row["层"]), tier, build_level, craft_levels, number(row["主材掉率"]), *p, number(row["副材总量"]), "=I%d/4" % r])
    style_sheet(ws, source_note=DUNGEON_MD.name + "；" + ECON_MD.name)
    for row in range(2, ws.max_row + 1):
        for col in range(5, 9):
            ws.cell(row, col).number_format = "0.00%"

    ws = wb.create_sheet("09_日产验算")
    ws.append(["类型", "阶段/刷层", "代表层", "次数/日", "装备/日", "绿装/日", "蓝装/日", "紫装/日", "金装/日", "红装/日", "全分解金币上限/日", "打造档位等级", "主材总/日", "紫主材/日", "金主材/日", "红主材/日", "副材总/日", "单种副材/日", "武器主材/日", "防具主材/日", "首饰主材/日"])
    default_runs = number(data["daily_material"][0]["次数/日"])
    for row in data["daily_equipment"]:
        ws.append(["装备验算", "刷%s层" % row["刷层"], number(row["刷层"]), default_runs] + [number(row[h]) for h in ["装备/日", "绿装/日", "蓝装/日", "紫装/日", "金装/日", "红装/日", "全分解金币上限/日"]] + [""] * 10)
    for row in data["daily_material"]:
        r = ws.max_row + 1
        layer = number(row["代表刷层"])
        ws.append(["材料验算", row["阶段"], layer, number(row["次数/日"]), "", "", "", "", "", "", "",
                   "=INDEX('08_副本材料'!$C$2:$C$31,MATCH(C%d,'08_副本材料'!$A$2:$A$31,0))" % r,
                   "=D%d*INDEX('08_副本材料'!$E$2:$E$31,MATCH(C%d,'08_副本材料'!$A$2:$A$31,0))" % (r, r),
                   "=M%d*INDEX('08_副本材料'!$F$2:$F$31,MATCH(C%d,'08_副本材料'!$A$2:$A$31,0))" % (r, r),
                   "=M%d*INDEX('08_副本材料'!$G$2:$G$31,MATCH(C%d,'08_副本材料'!$A$2:$A$31,0))" % (r, r),
                   "=M%d*INDEX('08_副本材料'!$H$2:$H$31,MATCH(C%d,'08_副本材料'!$A$2:$A$31,0))" % (r, r),
                   "=D%d*INDEX('08_副本材料'!$I$2:$I$31,MATCH(C%d,'08_副本材料'!$A$2:$A$31,0))" % (r, r),
                   "=Q%d/4" % r, "=M%d*20%%" % r, "=M%d*55%%" % r, "=M%d*25%%" % r])
    style_sheet(ws, source_note=DUNGEON_MD.name)

    ws = wb.create_sheet("10_爬塔挂机")
    ws.append(["最高层", "金币/日", "房车经验/日", "", "最高层材料区间", "主材阶位", "打造档位等级", "主材总/日", "紫概率", "金概率", "红概率", "紫期望/日", "金期望/日", "红期望/日", "副材总/日", "单种副材/日"])
    gold_rows = data["tower_gold"]
    mat_rows = data["tower_material"]
    tier_to_build = {roman[i]: number(row["打造档位等级"]) for i, row in enumerate(forge_rows)}
    for i in range(max(len(gold_rows), len(mat_rows))):
        gold = gold_rows[i] if i < len(gold_rows) else None
        mat = mat_rows[i] if i < len(mat_rows) else None
        tier = mat["主材阶位"] if mat else ""
        p = probs(mat["紫/金/红概率"], 3) if mat else ["", "", ""]
        ws.append([
            gold["最高层"] if gold else "", number(gold["金币/日"]) if gold else "", number(gold["房车经验/日"]) if gold else "", "",
            mat["最高层"] if mat else "", tier, tier_to_build.get(tier, "-" if tier == "-" else ""), number(mat["主材总/日"]) if mat else "",
            *p,
            number(mat["紫期望/日"]) if mat else "", number(mat["金期望/日"]) if mat else "", number(mat["红期望/日"]) if mat else "",
            number(mat["副材总/日"]) if mat else "", number(mat["单种副材/日"]) if mat else "",
        ])
    style_sheet(ws, source_note=TOWER_MD.name)
    for row in range(2, ws.max_row + 1):
        for col in range(9, 12):
            if isinstance(ws.cell(row, col).value, (int, float)):
                ws.cell(row, col).number_format = "0.00%"

    ws = wb.create_sheet("11_锻造台经验")
    ws.append(["锻造台Lv", "升下一级经验", "累计经验", "绿装经验", "蓝装经验", "紫装经验", "金装经验", "红装经验", "备注", "", "装备等级段", "绿", "蓝", "紫", "金", "红"])
    formula = data["forge_exp_formula"]
    values = data["forge_exp_values"]
    for lv in range(1, 61):
        r = ws.max_row + 1
        cost = values[lv - 1] if values and lv <= 59 else ("" if values else '=IF(A%d=60,"",ROUND((%g+%g*A%d+%g*A%d^2)/10,0)*10)' % (r, formula[0], formula[1], r, formula[2], r))
        ws.append([lv, cost, 0 if lv == 1 else "=C%d+B%d" % (r - 1, r - 1), "", "", "", "", "", "离散曲线来自沉淀MD" if values else "升级公式来自沉淀MD", None, None, None, None, None, None, None])
    for idx, row in enumerate(data["equip_exp"], 2):
        for col, key in enumerate(["装备等级", "绿", "蓝", "紫", "金", "红"], 11):
            ws.cell(idx, col, row[key].replace(" ", "") if key == "装备等级" else number(row[key]))
    style_sheet(ws, source_note=ECON_MD.name)

    ws = wb.create_sheet("12_合成消耗")
    ws.append(["合成档", "角色等级要求", "产物等级池", "等级权重", "投入装备等级下限", "基础金币", "绿金币", "蓝金币", "紫金币", "金金币", "红金币"])
    for row in data["synthesis_tier"]:
        r = ws.max_row + 1
        ws.append([number(row["合成档"]), number(row["角色等级要求"]), row["产物等级池"], row["等级权重"], number(row["投入装备等级下限"]), number(row["基础金币"])] + [
            "=$F%d*INDEX('02_品质等级'!$E$2:$E$6,MATCH(LEFT(%s$1,1),'02_品质等级'!$A$2:$A$6,0))" % (r, get_column_letter(col)) for col in range(7, 12)
        ])
    style_sheet(ws, source_note=ECON_MD.name)

    ws = wb.create_sheet("13_合成概率")
    ws.append(["锻造台Lv", "绿合蓝", "绿合紫", "蓝合紫", "蓝合金", "紫合紫", "紫合金", "紫合红", "金合金", "金合红", "红合红"])
    pf = data["prob_formulas"]
    for lv in range(1, 61):
        r = ws.max_row + 1
        ws.append([lv, "=1-C%d" % r, "=%s+%s*(A%d-1)/59" % (*map(percent_literal, pf["绿合紫"]), r), "=1-E%d" % r,
                   "=%s+%s*(A%d-1)/59" % (*map(percent_literal, pf["蓝合金"]), r), "=1-G%d-H%d" % (r, r),
                   "=%s+%s*(A%d-1)/59" % (*map(percent_literal, pf["紫合金"]), r), "=%s+%s*(A%d-1)/59" % (*map(percent_literal, pf["紫合红"]), r),
                   "=1-J%d" % r, "=%s+%s*(A%d-1)/59" % (*map(percent_literal, pf["金合红"]), r), 1])
    style_sheet(ws, source_note=ECON_MD.name)
    for row in range(2, 62):
        for col in range(2, 12):
            ws.cell(row, col).number_format = "0.00%"

    ws = wb.create_sheet("14_打造消耗")
    ws.append(["打造档", "主材阶位", "打造档位等级", "可打造装备等级", "金币(暂沿用旧档)", "紫主材每种副材", "金主材每种副材", "红主材每种副材", "武器消耗副材", "防具消耗副材", "首饰消耗副材", "备注"])
    for idx, row in enumerate(forge_rows):
        ws.append([number(row["打造档"]), roman[idx], number(row["打造档位等级"]), row["可打造装备等级"], number(row["金币(暂沿用旧档)"]), number(row["紫主材每种副材"]), number(row["金主材每种副材"]), number(row["红主材每种副材"]), "副材1 + 副材2", "副材2 + 副材3", "副材3 + 副材4", "每次主材 x1；副材数量为每种数量"])
    style_sheet(ws, source_note=ECON_MD.name)

    ws = wb.create_sheet("15_打造概率")
    ws.append(["锻造台Lv", "紫主材出紫", "紫主材出金", "金主材出金", "金主材出红", "红主材出红"])
    for lv in range(1, 61):
        r = ws.max_row + 1
        ws.append([lv, "=1-C%d" % r, "=%s+%s*(A%d-1)/59" % (*map(percent_literal, pf["紫主材出金"]), r), "=1-E%d" % r, "=%s+%s*(A%d-1)/59" % (*map(percent_literal, pf["金主材出红"]), r), 1])
    style_sheet(ws, source_note=ECON_MD.name)
    for row in range(2, 62):
        for col in range(2, 7):
            ws.cell(row, col).number_format = "0.00%"

    decompose = data["decompose_base"][0]
    base_gold = {int(level): number(decompose[str(level)]) for level in range(5, 61, 5)}
    exp_by_level = {}
    for row in data["equip_exp"]:
        levels = ints(row["装备等级"])
        for level in levels:
            exp_by_level[level] = {q: number(row[q]) for q in qualities}
    ws = wb.create_sheet("16_分解回收")
    ws.append(["装备等级", "基础金币", "绿金币", "蓝金币", "紫金币", "金金币", "红金币", "绿经验", "蓝经验", "紫经验", "金经验", "红经验", "红装额外产出"])
    for level in range(5, 61, 5):
        r = ws.max_row + 1
        formulas = ["=$B%d*INDEX('02_品质等级'!$F$2:$F$6,MATCH(LEFT(%s$1,1),'02_品质等级'!$A$2:$A$6,0))" % (r, get_column_letter(col)) for col in range(3, 8)]
        ws.append([level, base_gold[level], *formulas, *[exp_by_level[level][q] for q in qualities], "高阶零件(23)x1"])
    style_sheet(ws, source_note=ECON_MD.name)

    ws = wb.create_sheet("17_边界待确认")
    add_rows(ws, [
        ["事项", "状态", "说明"],
        ["红装洗练", "本版本先不做", "不登记任何洗练材料或金币消耗"],
        ["升级系统", "已废弃", "旧 SpecialEquipmentLvUp 不计入当前版本"],
        ["打孔/鉴定", "已废弃", "旧精密打孔器/鉴定卷轴不计入当前版本"],
        ["装备高级材料", "高阶零件(23)", "仅红装分解产出；当前作为后续材料沉淀"],
        ["主材/副材 ItemId", "待规划", "主材为大类 x 品质 x 等级；副材为4种"],
        ["映射之匙(25)", "副本入场消耗", "钥匙来源/购买/上限不在本表展开"],
        ["首通旧胚料", "待重配", "05副本沉淀中旧 I-VIII 胚料首通表仅留档，不纳入新版主材/副材最终口径"],
        ["爬塔房车经验", "待确认", "需确认是新道具，还是房车零件改名/经验化"],
        ["已确认：打造8档", "打造只分8个档位：10/20/30/40/45/50/55/60；15/25/35仅作为档内可打造装备等级，不单独拆主材档。", "2026-07-09 用户反馈修正"],
    ])
    style_sheet(ws, source_note=ECON_MD.name + "；" + DUNGEON_MD.name + "；" + TOWER_MD.name)

    build_summary_sheet(wb, data["scenario"])

    if wb.sheetnames != SHEET_ORDER:
        raise AssertionError("页签顺序错误: %s" % wb.sheetnames)
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    return wb


def build_summary_sheet(wb: Workbook, scenarios: list[dict[str, str]]) -> None:
    ws = wb.create_sheet("18_产销汇总")
    ws.append(["映射装备总产销汇总（前 / 中 / 后期）"])
    ws.append([])
    ws.append(["阶段假设（蓝字可改）"])
    ws.append(["阶段", "副本代表层", "副本次数/日", "爬塔金币最高层", "爬塔材料区间", "合成档", "打造档位等级", "备注"])
    for row in scenarios:
        ws.append([row["阶段"], number(row["副本代表层"]), number(row["副本次数/日"]), row["爬塔金币最高层"], row["爬塔材料区间"], number(row["合成档"]), number(row["打造档位等级"]), row["备注"]])
    ws.append([])
    ws.append(["一、日产出汇总"])
    ws.append(["指标", "单位/口径", "前期", "中期", "后期", "说明"])
    labels = [
        ("映射之匙消耗/日", "次/日", "按副本次数折算"),
        ("爬塔金币/日", "金币/日", "爬塔挂机金币；非映射装备专属，但会进入金币总账"),
        ("装备总产出/日", "件/日", "副本宝箱开装期望"),
        ("绿装/日", "件/日", ""), ("蓝装/日", "件/日", ""), ("紫装/日", "件/日", ""),
        ("金装/日", "件/日", ""), ("红装/日", "件/日", "红宝箱必出红装已计入"),
        ("副本主材总/日", "个/日", "副本独立主材判定"),
        ("副本紫主材/日", "个/日", ""), ("副本金主材/日", "个/日", ""), ("副本红主材/日", "个/日", ""),
        ("副本副材总/日", "个/日", "副材1-4等概率"),
        ("爬塔主材总/日", "个/日", "爬塔挂机材料"), ("爬塔紫主材/日", "个/日", ""),
        ("爬塔金主材/日", "个/日", ""), ("爬塔红主材/日", "个/日", ""), ("爬塔副材总/日", "个/日", "副材1-4等概率"),
        ("主材合计/日", "个/日", "副本 + 爬塔"), ("紫主材合计/日", "个/日", ""),
        ("金主材合计/日", "个/日", ""), ("红主材合计/日", "个/日", ""), ("副材合计/日", "个/日", "副本 + 爬塔"),
        ("单种副材合计/日", "个/日", "副材1/2/3/4等概率"),
    ]
    for label, unit, note in labels:
        ws.append([label, unit, "", "", "", note])
    for col, stage_row in zip("CDE", (5, 6, 7)):
        ws[f"{col}11"] = f"=$C${stage_row}"
        ws[f"{col}12"] = f"=INDEX('10_爬塔挂机'!$B$2:$B$11,MATCH($D${stage_row},'10_爬塔挂机'!$A$2:$A$11,0))"
        for row, source_col in zip(range(13, 19), "EFGHIJ"):
            ws[f"{col}{row}"] = f'=INDEX(\'09_日产验算\'!${source_col}$2:${source_col}$8,MATCH("刷"&$B${stage_row}&"层",\'09_日产验算\'!$B$2:$B$8,0))'
        ws[f"{col}19"] = f"=$C${stage_row}*INDEX('08_副本材料'!$E$2:$E$31,MATCH($B${stage_row},'08_副本材料'!$A$2:$A$31,0))"
        for row, source_col in zip(range(20, 23), "FGH"):
            ws[f"{col}{row}"] = f"={col}19*INDEX('08_副本材料'!${source_col}$2:${source_col}$31,MATCH($B${stage_row},'08_副本材料'!$A$2:$A$31,0))"
        ws[f"{col}23"] = f"=$C${stage_row}*INDEX('08_副本材料'!$I$2:$I$31,MATCH($B${stage_row},'08_副本材料'!$A$2:$A$31,0))"
        for row, source_col in zip(range(24, 29), "HLMNO"):
            ws[f"{col}{row}"] = 0 if not any(wb['10_爬塔挂机'].cell(index, 5).value for index in range(2, wb['10_爬塔挂机'].max_row + 1)) else f"=INDEX('10_爬塔挂机'!${source_col}$2:${source_col}$17,MATCH($E${stage_row},'10_爬塔挂机'!$E$2:$E$17,0))"
        for row, formula in {
            29: f"={col}19+{col}24", 30: f"={col}20+{col}25", 31: f"={col}21+{col}26", 32: f"={col}22+{col}27",
            33: f"={col}23+{col}28", 34: f"={col}33/4",
        }.items():
            ws[f"{col}{row}"] = formula

    ws.append([])
    ws.append(["二、打造消耗承载（按每日材料全投入，三种品质不同时发生）"])
    ws.append(["指标", "单位/口径", "前期", "中期", "后期", "说明"])
    for row in [
        ["代表打造金币/次", "金币/次", "", "", "", "取阶段假设中的打造档位等级"],
        ["紫主材每种副材消耗", "个/次", "", "", "", "每次消耗2种副材，各消耗该数值"],
        ["金主材每种副材消耗", "个/次", "", "", "", ""],
        ["红主材每种副材消耗", "个/次", "", "", "", ""],
        ["紫主材可支撑打造次数/日", "次/日", "", "", "", "若全投入紫主材打造"],
        ["金主材可支撑打造次数/日", "次/日", "", "", "", "若全投入金主材打造"],
        ["红主材可支撑打造次数/日", "次/日", "", "", "", "若全投入红主材打造"],
        ["紫打造金币消耗/日", "金币/日", "", "", "", "按可支撑次数折算"],
        ["金打造金币消耗/日", "金币/日", "", "", "", "按可支撑次数折算"],
        ["红打造金币消耗/日", "金币/日", "", "", "", "按可支撑次数折算"],
    ]:
        ws.append(row)
    for col, stage_row in zip("CDE", (5, 6, 7)):
        for row, source_col in zip(range(38, 42), "EFGH"):
            ws[f"{col}{row}"] = f"=INDEX('14_打造消耗'!${source_col}$2:${source_col}$9,MATCH($G${stage_row},'14_打造消耗'!$C$2:$C$9,0))"
        ws[f"{col}42"] = f"=MIN({col}30,{col}33/(2*{col}39))"
        ws[f"{col}43"] = f"=MIN({col}31,{col}33/(2*{col}40))"
        ws[f"{col}44"] = f"=MIN({col}32,{col}33/(2*{col}41))"
        for row, source_row in zip(range(45, 48), range(42, 45)):
            ws[f"{col}{row}"] = f"={col}{source_row}*{col}38"

    ws.append([])
    ws.append(["三、合成消耗承载（按掉落装备全用于同品质合成的潜在上限）"])
    ws.append(["指标", "单位/口径", "前期", "中期", "后期", "说明"])
    for q in ["绿", "蓝", "紫", "金", "红"]:
        ws.append([q + "装合成金币/次", "金币/次", "", "", "", "代表合成档" if q == "绿" else ""])
    for q in ["绿", "蓝", "紫", "金", "红"]:
        ws.append([q + "装潜在合成次数/日", "次/日", "", "", "", "9件同品质合1次" if q == "绿" else ""])
    ws.append(["全品质潜在合成次数/日", "次/日", "", "", "", "各品质分别除以9后汇总"])
    ws.append(["全品质潜在合成金币消耗/日", "金币/日", "", "", "", "潜在上限，不代表实际每日必消耗"])
    for col, stage_row in zip("CDE", (5, 6, 7)):
        for row, source_col in zip(range(51, 56), "GHIJK"):
            ws[f"{col}{row}"] = f"=INDEX('12_合成消耗'!${source_col}$2:${source_col}$9,MATCH($F${stage_row},'12_合成消耗'!$A$2:$A$9,0))"
        for row, source_row in zip(range(56, 61), range(14, 19)):
            ws[f"{col}{row}"] = f"={col}{source_row}/9"
        ws[f"{col}61"] = f"=SUM({col}56:{col}60)"
        ws[f"{col}62"] = "+".join(f"={col}56*{col}51" if i == 0 else f"{col}{56+i}*{col}{51+i}" for i in range(5))

    ws.append([])
    ws.append(["四、回收与金币压力参考"])
    ws.append(["指标", "单位/口径", "前期", "中期", "后期", "说明"])
    for row in [
        ["全分解金币回收上限/日", "金币/日", "", "", "", "若当日装备全部分解"],
        ["全分解锻造台经验上限/日", "经验/日", "", "", "", "按阶段打造档位等级匹配分解经验"],
        ["高阶零件(23)上限/日", "个/日", "", "", "", "按红装全部分解、每件红装x1估算"],
        ["金币供给参考/日", "金币/日", "", "", "", "爬塔金币 + 全分解金币上限"],
        ["扣除潜在合成后金币余量", "金币/日", "", "", "", "仅用于压力参考"],
        ["扣除紫打造后金币余量", "金币/日", "", "", "", "若每日材料全投紫打造"],
        ["扣除金打造后金币余量", "金币/日", "", "", "", "若每日材料全投金打造"],
        ["扣除红打造后金币余量", "金币/日", "", "", "", "若每日材料全投红打造"],
    ]:
        ws.append(row)
    for col, stage_row in zip("CDE", (5, 6, 7)):
        ws[f"{col}66"] = f'=INDEX(\'09_日产验算\'!$K$2:$K$8,MATCH("刷"&$B${stage_row}&"层",\'09_日产验算\'!$B$2:$B$8,0))'
        exp_terms = []
        for source_row, source_col in zip(range(14, 19), "HIJKL"):
            exp_terms.append(f"{col}{source_row}*INDEX('16_分解回收'!${source_col}$2:${source_col}$13,MATCH($G${stage_row},'16_分解回收'!$A$2:$A$13,0))")
        ws[f"{col}67"] = "=" + "+".join(exp_terms)
        ws[f"{col}68"] = f"={col}18"
        ws[f"{col}69"] = f"={col}12+{col}66"
        ws[f"{col}70"] = f"={col}69-{col}62"
        ws[f"{col}71"] = f"={col}69-{col}45"
        ws[f"{col}72"] = f"={col}69-{col}46"
        ws[f"{col}73"] = f"={col}69-{col}47"

    for merged in ("A1:H1", "A3:H3", "A9:H9", "A36:H36", "A49:H49", "A64:H64"):
        ws.merge_cells(merged)
    style_sheet(ws, header_rows=(4, 10, 37, 50, 65), source_note=SCENARIO_MD.name)
    for row in (1, 3, 9, 36, 49, 64):
        ws.cell(row, 1).fill = PatternFill("solid", fgColor=SECTION)
        ws.cell(row, 1).font = Font(name=FONT_NAME, size=11, bold=True, color=BLACK)
    for row in range(5, 8):
        for col in range(1, 9):
            ws.cell(row, col).fill = PatternFill("solid", fgColor=INPUT_FILL)
            ws.cell(row, col).font = Font(name=FONT_NAME, size=10, color=BLUE)
    ws.freeze_panes = "C11"


def workbook_stats(path: Path) -> dict:
    wb = load_workbook(path, data_only=False, read_only=True, keep_links=False)
    result = {}
    for ws in wb.worksheets:
        formulas = 0
        nonempty = 0
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None:
                    nonempty += 1
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formulas += 1
        result[ws.title] = (ws.max_row, ws.max_column, nonempty, formulas)
    wb.close()
    return result


def compare_workbooks(new_path: Path, old_path: Path) -> dict:
    new = load_workbook(new_path, data_only=False, read_only=True, keep_links=False)
    old = load_workbook(old_path, data_only=False, read_only=True, keep_links=False)
    diffs = 0
    checked = 0
    by_sheet = {}
    for name in SHEET_ORDER:
        if name not in new.sheetnames or name not in old.sheetnames:
            by_sheet[name] = "缺页"
            continue
        nw, ow = new[name], old[name]
        sd = 0
        if name != "00_口径":
            new_rows = list(nw.iter_rows(values_only=True))
            old_rows = list(ow.iter_rows(values_only=True))
            max_row = max(len(new_rows), len(old_rows))
            max_col = max(nw.max_column, ow.max_column)
            for r in range(max_row):
                nr = new_rows[r] if r < len(new_rows) else ()
                orow = old_rows[r] if r < len(old_rows) else ()
                for c in range(max_col):
                    checked += 1
                    nv = nr[c] if c < len(nr) else None
                    ov = orow[c] if c < len(orow) else None
                    blanks_equal = nv in (None, "") and ov in (None, "")
                    numbers_equal = (
                        isinstance(nv, (int, float))
                        and not isinstance(nv, bool)
                        and isinstance(ov, (int, float))
                        and not isinstance(ov, bool)
                        and abs(float(nv) - float(ov)) <= 1e-12
                    )
                    if not (nv == ov or blanks_equal or numbers_equal):
                        diffs += 1
                        sd += 1
        by_sheet[name] = sd
    new.close()
    old.close()
    return {"checked": checked, "diffs": diffs, "by_sheet": by_sheet}


def validate_output(path: Path) -> tuple[int, int]:
    wb = load_workbook(path, data_only=False, read_only=False, keep_links=True)
    if wb.sheetnames != SHEET_ORDER:
        raise AssertionError("输出缺页或顺序错误")
    formulas = 0
    errors = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                if isinstance(value, str) and value.startswith("="):
                    formulas += 1
                    if "[" in value and "]" in value:
                        errors.append("%s!%s 外链公式" % (ws.title, cell.coordinate))
                if value in ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A"):
                    errors.append("%s!%s=%s" % (ws.title, cell.coordinate, value))
    discrete_curve_cells = sum(isinstance(wb["11_锻造台经验"].cell(row, 2).value, (int, float)) for row in range(2, 61))
    tower_material_retired = not any(wb["10_爬塔挂机"].cell(row, 5).value for row in range(2, wb["10_爬塔挂机"].max_row + 1))
    wb.close()
    if errors:
        raise RuntimeError("；".join(errors[:20]))
    minimum_formulas = 6000 - discrete_curve_cells - (15 if tower_material_retired else 0)
    if formulas < minimum_formulas:
        raise AssertionError("公式数量异常: %d" % formulas)
    return len(SHEET_ORDER), formulas


def write_report(stats: dict, evidence: dict, compare: dict | None) -> None:
    REPORT.parent.mkdir(parents=True, exist_ok=True)
    lines = [
        "# 映射装备数值总表生成与验证",
        "",
        "> 从 `01_源表` 与 `03_沉淀MD` 重建 19 页工作簿；运行时不读取 `04_数值文档`。",
        "",
        "## 结果",
        "",
        "- 生成物：`养成/生成_映射装备数值总表.xlsx`",
        "- 页签：19/19",
        "- 公式：%d" % sum(v[3] for v in stats.values()),
        "- 01 校验：映射装备投放系数 `%s`，金币 ItemId `%s`，经济边界已确认。" % (evidence["mapping_coefficient"], evidence["gold"]["id"]),
        "",
        "## 页签覆盖",
        "",
        "| 页签 | 行 | 列 | 非空 | 公式 |",
        "|---|---:|---:|---:|---:|",
    ]
    for name in SHEET_ORDER:
        r, c, n, f = stats[name]
        lines.append("| %s | %d | %d | %d | %d |" % (name, r, c, n, f))
    lines += ["", "## 运行时输入", ""]
    for rel, digest in evidence["hashes"].items():
        lines.append("- `%s` — `%s`" % (rel.replace("\\", "/"), digest[:16]))
    if compare:
        lines += [
            "", "## 旧终稿只读对照", "",
            "- 对照单元格：%d" % compare["checked"],
            "- 值/公式差异：%d（`00_口径` 因新来源链不计）" % compare["diffs"],
            "- 该对照仅用于开发验证，不是运行时输入。",
        ]
    REPORT.write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", default=os.fspath(OUT))
    parser.add_argument("--compare-final", action="store_true", help="开发时只读对照旧04终稿")
    args = parser.parse_args()
    output = Path(args.output)

    data = table_data()
    evidence = source_evidence(data)
    wb = build_workbook(data, evidence)
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    wb.close()
    sheets, formulas = validate_output(output)
    stats = workbook_stats(output)
    compare = compare_workbooks(output, LEGACY_FINAL) if args.compare_final and LEGACY_FINAL.exists() else None
    write_report(stats, evidence, compare)
    print("映射装备数值总表生成完成：页签 %d，公式 %d" % (sheets, formulas))
    print("[生成表] %s" % output)
    print("[报告] %s" % REPORT)
    if compare:
        print("[旧终稿对照] 单元格 %d，差异 %d" % (compare["checked"], compare["diffs"]))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
