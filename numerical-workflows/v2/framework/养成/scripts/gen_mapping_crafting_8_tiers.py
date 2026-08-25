# -*- coding: utf-8 -*-
"""从养成参数基线 MD 生成映射装备合成/制作 8 档工作簿。

运行时权威输入仅为：
``03_沉淀MD/参数基线/养成参数基线.md`` 中的三个机器参数页：

- 映射装备_合成8档
- 映射装备_制作8档
- 映射装备_制作材料

本脚本不读取 ``04_数值文档``，也不把任何既有 Excel 当作模板。

用法：

    python gen_mapping_crafting_8_tiers.py
    python gen_mapping_crafting_8_tiers.py --check
"""

from __future__ import annotations

import argparse
import io
import os
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


HERE = Path(__file__).resolve().parent
GROWTH_DIR = HERE.parent
V2_DIR = GROWTH_DIR.parent
sys.path.insert(0, os.fspath(V2_DIR))

from _shared.md_parameter_reader import load_parameter_workbook, parameter_md_path


PARAMETER_SCOPE = "养成"
PARAMETER_MD = Path(parameter_md_path(PARAMETER_SCOPE))
OUTPUT = GROWTH_DIR / "生成_映射装备合成制作8档.xlsx"

SOURCE_TO_OUTPUT = (
    ("映射装备_合成8档", "合成8档"),
    ("映射装备_制作8档", "制作8档"),
    ("映射装备_制作材料", "制作材料"),
)

FONT_NAME = "Microsoft YaHei"
NAVY = "1F4E78"
LIGHT_BLUE = "BDD7EE"
SECTION_BLUE = "D9EAF7"
NOTE_YELLOW = "FFF2CC"
INPUT_BLUE = "0000FF"
BLACK = "000000"
WHITE = "FFFFFF"
GRID = Side(style="thin", color="A6A6A6")
THIN_BORDER = Border(left=GRID, right=GRID, top=GRID, bottom=GRID)


def copy_parameter_sheet(source_ws, target_ws) -> None:
    """逐格复制 MD 内存工作簿，公式文本保持不变。"""

    for row in source_ws.iter_rows():
        for cell in row:
            if cell.value is not None:
                target_ws.cell(cell.row, cell.column, cell.value)


def style_cell(cell, *, bold=False, color=BLACK, fill=None, size=10, center=False) -> None:
    cell.font = Font(name=FONT_NAME, size=size, bold=bold, color=color)
    if fill:
        cell.fill = PatternFill("solid", fgColor=fill)
    cell.border = THIN_BORDER
    cell.alignment = Alignment(
        horizontal="center" if center else None,
        vertical="center",
        wrap_text=True,
    )


def style_rect(ws, min_row, max_row, min_col, max_col, *, input_cells=True) -> None:
    for row in ws.iter_rows(
        min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col
    ):
        for cell in row:
            is_formula = isinstance(cell.value, str) and cell.value.startswith("=")
            color = BLACK if is_formula or not input_cells else INPUT_BLUE
            style_cell(cell, color=color)


def style_title(ws, max_col: int) -> None:
    for col in range(1, max_col + 1):
        cell = ws.cell(1, col)
        style_cell(cell, bold=True, color=WHITE, fill=NAVY, size=14)
    ws.row_dimensions[1].height = 28


def style_header(ws, row: int, max_col: int) -> None:
    for col in range(1, max_col + 1):
        style_cell(ws.cell(row, col), bold=True, fill=LIGHT_BLUE, center=True)


def style_section(ws, row: int, max_col: int, *, yellow=False) -> None:
    for col in range(1, max_col + 1):
        style_cell(
            ws.cell(row, col),
            bold=True,
            fill=NOTE_YELLOW if yellow else SECTION_BLUE,
        )


def add_source_comment(ws, parameter_sheet: str) -> None:
    ws["A1"].comment = Comment(
        "参数权威：03_沉淀MD/参数基线/养成参数基线.md\n"
        "稳定来源：养成/%s!A1\n"
        "本工作簿由 gen_mapping_crafting_8_tiers.py 独立生成，不读取 04_数值文档。"
        % parameter_sheet,
        "Vgame 数值框架",
    )


def build_overview(ws) -> None:
    rows = [
        ["口径总览", None],
        [None, None],
        ["项", "内容"],
        ["档位", "8 档：10 / 20 / 30 / 40 / 45 / 50 / 55 / 60"],
        ["合成", "9 件同品质装备 -> 1 件同品质随机部位装备；产物等级按档内等级池随机"],
        ["制作", "指定部位、档内等级可选、品质随机；制作结果等级不随机"],
        ["材料", "8 档 * 3 部位大类 = 24 种制作材料"],
        ["正式配置", "未修改真实配置；后续需要确认表结构承接等级池和可制作等级"],
        ["参数权威", "03_沉淀MD/参数基线/养成参数基线.md"],
        [
            "参数页",
            "养成/映射装备_合成8档；养成/映射装备_制作8档；养成/映射装备_制作材料",
        ],
        ["生成器", "养成/scripts/gen_mapping_crafting_8_tiers.py"],
        ["运行边界", "仅读取上述 MD 参数页；不读取或复制 04_数值文档，也不依赖旧参数 Excel"],
        ["维护规则", "先同步 MD 参数/公式，再运行本生成器；经对照评审后才升级正式终稿"],
    ]
    for row in rows:
        ws.append(row)
    style_title(ws, 2)
    style_header(ws, 3, 2)
    style_rect(ws, 4, len(rows), 1, 2, input_cells=False)
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 88
    ws.freeze_panes = "A3"
    for row in range(3, len(rows) + 1):
        ws.row_dimensions[row].height = 28 if row >= 5 else 22


def style_compose(ws) -> None:
    style_title(ws, 12)
    style_header(ws, 3, 12)
    style_rect(ws, 4, 11, 1, 12)
    style_section(ws, 13, 12, yellow=True)
    widths = [12, 14, 16, 16, 18, 12, 10, 10, 10, 10, 10, 30]
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    for row in range(3, 14):
        ws.row_dimensions[row].height = 24
    ws.row_dimensions[13].height = 46
    ws.freeze_panes = "A3"
    ws.auto_filter.ref = "A3:L11"
    for row in range(4, 12):
        for col in range(6, 12):
            ws.cell(row, col).number_format = "#,##0.##"


def style_craft(ws) -> None:
    style_title(ws, 9)
    style_header(ws, 3, 9)
    style_rect(ws, 4, 11, 1, 9)
    style_section(ws, 13, 9)
    style_header(ws, 14, 5)
    style_rect(ws, 15, 15, 1, 5)
    widths = [12, 14, 16, 14, 18, 18, 18, 12, 22]
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    for row in range(3, 16):
        ws.row_dimensions[row].height = 24
    ws.freeze_panes = "A3"
    ws.auto_filter.ref = "A3:I11"
    for col in range(1, 6):
        ws.cell(15, col).number_format = "0%"


def style_material(ws) -> None:
    style_title(ws, 6)
    style_header(ws, 3, 6)
    style_rect(ws, 4, 11, 1, 6)
    style_section(ws, 13, 6)
    style_rect(ws, 14, 21, 1, 2)
    style_section(ws, 23, 6)
    style_rect(ws, 24, 26, 1, 3)
    widths = [14, 34, 22, 22, 18, 20]
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    for row in range(3, 27):
        ws.row_dimensions[row].height = 24
    ws.freeze_panes = "A3"
    ws.auto_filter.ref = "A3:F11"
    for row in range(24, 27):
        ws.cell(row, 2).number_format = "0%"


def build_workbook() -> Workbook:
    source = load_parameter_workbook(PARAMETER_SCOPE, data_only=False)
    missing = [name for name, _ in SOURCE_TO_OUTPUT if name not in source.sheetnames]
    if missing:
        raise RuntimeError("养成参数基线 MD 缺少参数页：%s" % "、".join(missing))

    wb = Workbook()
    overview = wb.active
    overview.title = "口径总览"
    build_overview(overview)

    for source_name, target_name in SOURCE_TO_OUTPUT:
        ws = wb.create_sheet(target_name)
        copy_parameter_sheet(source[source_name], ws)
        if target_name == "制作材料":
            ws["A1"] = "制作材料"
        add_source_comment(ws, source_name)
        if target_name == "合成8档":
            style_compose(ws)
        elif target_name == "制作8档":
            style_craft(ws)
        else:
            style_material(ws)
    source.close()

    wb.properties.creator = "Vgame 数值框架"
    wb.properties.title = "映射装备合成制作 8 档"
    wb.properties.subject = "由养成参数基线 MD 独立生成"
    wb.calculation.calcMode = "auto"
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    return wb


def parse_percent_weights(text: str) -> float:
    total = 0.0
    for token in text.replace("/", " ").split():
        if token.endswith("%"):
            total += float(token[:-1]) / 100.0
    return total


def validate(path: Path) -> list[str]:
    problems: list[str] = []
    if not path.exists():
        return ["生成文件不存在：%s" % path]

    wb = load_workbook(path, data_only=False, read_only=False)
    expected = ["口径总览", "合成8档", "制作8档", "制作材料"]
    if wb.sheetnames != expected:
        problems.append("页签顺序错误：%s" % wb.sheetnames)

    if all(name in wb.sheetnames for name in expected):
        compose = wb["合成8档"]
        craft = wb["制作8档"]
        material = wb["制作材料"]

        compose_tiers = [compose.cell(row, 1).value for row in range(4, 12)]
        craft_tiers = [craft.cell(row, 1).value for row in range(4, 12)]
        material_tiers = [material.cell(row, 1).value for row in range(4, 12)]
        if compose_tiers != list(range(1, 9)):
            problems.append("合成档不完整：%s" % compose_tiers)
        if craft_tiers != list(range(1, 9)):
            problems.append("制作档不完整：%s" % craft_tiers)
        if material_tiers != ["I", "II", "III", "IV", "V", "VI", "VII", "VIII"]:
            problems.append("制作材料档不完整：%s" % material_tiers)

        formulas = [
            cell.value
            for row in compose.iter_rows(min_row=4, max_row=11, min_col=7, max_col=11)
            for cell in row
            if isinstance(cell.value, str) and cell.value.startswith("=")
        ]
        if len(formulas) != 40:
            problems.append("合成金币公式应为 40 个，实际 %d" % len(formulas))
        for row in range(4, 12):
            expected_formulas = [
                "=F%d*0.5" % row,
                "=F%d*0.7" % row,
                "=F%d*1" % row,
                "=F%d*1.4" % row,
                "=F%d*2" % row,
            ]
            actual = [compose.cell(row, col).value for col in range(7, 12)]
            if actual != expected_formulas:
                problems.append("合成8档第 %d 行公式异常：%s" % (row, actual))

        quality_total = sum(float(craft.cell(15, col).value or 0) for col in range(1, 6))
        if abs(quality_total - 1.0) > 1e-9:
            problems.append("制作品质权重合计不是 100%%：%s" % quality_total)
        material_total = sum(float(material.cell(row, 2).value or 0) for row in range(24, 27))
        if abs(material_total - 1.0) > 1e-9:
            problems.append("材料类型权重合计不是 100%%：%s" % material_total)
        for row in range(14, 22):
            total = parse_percent_weights(str(material.cell(row, 2).value or ""))
            if abs(total - 1.0) > 1e-9:
                problems.append("制作材料第 %d 行档位权重不是 100%%：%s" % (row, total))
    wb.close()
    return problems


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--check", action="store_true", help="只校验已有生成表")
    args = parser.parse_args()

    if not args.check:
        wb = build_workbook()
        wb.save(OUTPUT)
        wb.close()

    problems = validate(OUTPUT)
    if problems:
        print("[FAIL] %s" % OUTPUT)
        for problem in problems:
            print("  - %s" % problem)
        return 1

    print("[OK] 参数源：%s" % PARAMETER_MD)
    print("[OK] 生成表：%s" % OUTPUT)
    print("[OK] 页签=4；合成档=8；制作档=8；材料档=8；公式=40；权重校验通过")
    return 0


if __name__ == "__main__":
    if hasattr(sys.stdout, "buffer"):
        sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
    raise SystemExit(main())
