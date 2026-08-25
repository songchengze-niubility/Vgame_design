# -*- coding: utf-8 -*-
"""
gen_power_curve.py — 玩家战力曲线 完整链路 生成器

共享参数基线 MD（标准人/投放系数）+ 养成参数基线 MD（升级曲线参数/战力曲线参数）+ 战力曲线表内系数(行3-10) →
生成 5条旧模块进度 + 完整战力链，并追加房车整备扩展列，直到最终输出:
  总单位时间伤害(玩家DPS) / 预计bossTTK / boss生命。
与现有 战斗框架/玩家战力曲线 逐列逐格对照；房车为 v2 新增扩展列，不参与旧源表对照。

链路(源公式逆向, 已核对; 含源表两处保留怪癖):
  进度_m(i)=结束×(1-0.97^(i-起始))/(1-0.97^(总-起始)); i=关卡id(累计序号)
  裸体防御 L=1×初始防御+角色进度×升级防御+星盘进度×星盘防御
  额外防御 M=专武进度×F41+映射进度×F42+旅券进度×F43   (源表投放列序)
  防御% N=1+3.6×(权重·进度); 算后防御 O=L×N+M; 减伤 P=O/(O+24000)
  裸体攻击 Q=初始攻击+角色进度×升级攻击  (源表此处空引用→星盘不计,如实复现)
  额外攻击 R=同M列序; 攻击% S; 算后攻击 T=Q×S+R
  系数倍率 U/V/W(随星盘进度); 战斗内总倍率 AC=U+V+W
  攻速/技能/必杀倍率 AD/AE/AF; 暴击 AG; 暴伤 AH; 增伤 AI (随权重·进度)
  养成总倍率 AJ=(1+AG×(AH-1))×AD×AF×AI
  总单位时间伤害 AK=AJ×AC×算后攻击×队伍人数进度×(1-减伤)
  预计bossTTK AL(插值); boss生命 AN=AL×AK
  房车整备：按账号全局额外属性追加「含房车」扩展列，旧列仍保持源表复现。

只读源表。生成表为独立新文件。
用法: python gen_power_curve.py [--check]
"""
import sys, io, os, argparse
import warnings
warnings.filterwarnings('ignore')
import openpyxl
from openpyxl.styles import Font, PatternFill

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

HERE = os.path.dirname(os.path.abspath(__file__))
GROWTH_DIR = os.path.dirname(HERE)  # 养成/
V2_DIR = os.path.dirname(GROWTH_DIR)  # 框架v2/
NUM_DIR = os.path.dirname(V2_DIR)  # 数值文档/数值文档/
sys.path.insert(0, V2_DIR)
from _shared.md_parameter_reader import load_parameter_workbook, parameter_md_path

SHARED_PARAM = parameter_md_path('共享')
GROWTH_PARAM = parameter_md_path('养成')
COMBAT_PARAM = parameter_md_path('战斗')
OUT     = os.path.join(GROWTH_DIR, '生成_玩家战力曲线.xlsx')
REPORT  = os.path.join(GROWTH_DIR, '_reports', 'P战力曲线_完整链路_生成与对照.md')
SRC     = os.path.join(NUM_DIR, 'Vgame基本战斗框架.xlsx')
SRC_SHEET = '玩家战力曲线'
ROW0 = 12  # i=1 行

def read_params():
    # 标准人、投放系数 → 从共享参数读
    wb_s = load_parameter_workbook(SHARED_PARAM, data_only=True)
    sp = wb_s['标准人']; std = {sp.cell(r,1).value: sp.cell(r,2).value for r in range(2,5)}
    fp = wb_s['投放系数']; coef = {}
    for r in range(2, fp.max_row+1):
        if fp.cell(r,1).value: coef[fp.cell(r,1).value] = (fp.cell(r,2).value, fp.cell(r,3).value)
    wb_s.close()

    # 升级曲线参数、战力曲线参数 → 从养成参数读
    wb_g = load_parameter_workbook(GROWTH_PARAM, data_only=True)
    cp = wb_g['升级曲线参数']; pr = {cp.cell(r,1).value: cp.cell(r,2).value for r in range(2,cp.max_row+1) if cp.cell(r,1).value}
    pc = wb_g['战力曲线参数']; N=None; pcurve=None; mods={}
    for r in range(1, pc.max_row+1):
        k=pc.cell(r,1).value
        if k=='总关卡': N=int(pc.cell(r,2).value)
        elif k=='函数参数p': pcurve=float(pc.cell(r,2).value)
        elif k in ('角色等级','星盘','专武','映射装备','旅券','房车整备'):
            mods[k]=(int(pc.cell(r,2).value), float(pc.cell(r,3).value))
    wb_g.close()

    # 战斗模型参数 → 从战斗参数读
    wb_s2 = load_parameter_workbook(COMBAT_PARAM, data_only=True)
    mp=wb_s2['战斗模型参数']; Kd=None
    for r in range(2,mp.max_row+1):
        if mp.cell(r,1).value=='防御常数K': Kd=float(mp.cell(r,2).value)
    wb_s2.close()

    lvl = int(pr['等级上限'])*float(pr['单级成长'])
    def c(name): return float(coef[name][0])
    sysm = {'初始':1,'升级':lvl,'星盘':c('角色星盘'),'专武':c('专武'),'旅券':c('旅券'),'映射':c('映射装备')}
    if '房车整备' in coef:
        sysm['房车'] = c('房车整备')
    return std, sysm, N, pcurve, mods, Kd

def main():
    ap=argparse.ArgumentParser(); ap.add_argument('--check',action='store_true'); args=ap.parse_args()
    std, sysm, N, pcurve, mods, Kd = read_params()
    ATK, DEF = std['攻击'], std['防御']
    F = {k: DEF*sysm[k] for k in sysm}
    E = {k: ATK*sysm[k] for k in sysm}

    wb = openpyxl.load_workbook(SRC, data_only=True); ws = wb[SRC_SHEET]
    def R(r,c): return ws.cell(r,c).value
    H10,I10,J10,K10 = R(10,8),R(10,9),R(10,10),R(10,11)
    U8,U9=R(8,21),R(9,21); V8,V9=R(8,22),R(9,22); W8,W9=R(8,23),R(9,23)
    N9=R(9,14); S9=R(9,19); AD9=R(9,30); AE9=R(9,31); AF9=R(9,32); AG9=R(9,33); AH9=R(9,34); AI9=R(9,35)
    F11,F12=0.05,2.0
    AL3,AL4,AL5,AL8=R(3,38),R(4,38),R(5,38),R(8,38)
    team=[R(ROW0+i,6) for i in range(N)]

    def prog(name,i):
        st,e = mods[name]
        return 0.0 if i<st else e*(1-pcurve**(i-st))/(1-pcurve**(N-st))

    cols={'L':12,'M':13,'N':14,'O':15,'P':16,'Q':17,'R':18,'S':19,'T':20,
          'AC':29,'AD':30,'AE':31,'AF':32,'AG':33,'AH':34,'AI':35,'AJ':36,'AK':37,'AL':38,'AN':40}
    prog_cols={'角色等级':7,'星盘':8,'专武':9,'映射装备':10,'旅券':11}
    err={k:0.0 for k in list(cols)+list(prog_cols)}
    out_rows=[]
    for i in range(1,N+1):
        g=prog('角色等级',i); h=prog('星盘',i); ii=prog('专武',i); j=prog('映射装备',i); k=prog('旅券',i)
        rv=prog('房车整备',i) if '房车整备' in mods and '房车' in sysm else 0.0
        Ft=team[i-1]
        L=1*F['初始']+g*F['升级']+h*F['星盘']
        M=ii*F['专武']+j*F['旅券']+k*F['映射']
        w3=H10*h+I10*ii+J10*j; w4=w3+K10*k
        Nv=1+N9*w3; O=L*Nv+M; P=O/(O+Kd)
        Q=1*E['初始']+g*E['升级']
        Rr=ii*E['专武']+j*E['旅券']+k*E['映射']
        Sv=1+S9*w3; T=Q*Sv+Rr
        U=U8+(U9-U8)*h; V=0.0 if i<14 else V8+(V9-V8)*h; W=W8+(W9-W8)*h
        AC=U+V+W; AD=1+AD9*w3; AE=1+AE9*w4; AF=1+AF9*w3
        AG=F11+AG9*w4; AH=F12+AH9*w4; AI=1+AI9*w4
        AJ=(1+AG*(AH-1))*AD*AF*AI
        AK=AJ*AC*T*Ft*(1-P)
        AL=0.0 if i<AL8 else AL3+(AL4-AL3)*(AL5**(i-AL8)-1)/(AL5**(N-AL8)-1)
        AN=AL*AK
        M_rv=M+rv*F.get('房车',0)
        O_rv=L*Nv+M_rv
        P_rv=O_rv/(O_rv+Kd)
        Rr_rv=Rr+rv*E.get('房车',0)
        T_rv=Q*Sv+Rr_rv
        AK_rv=AJ*AC*T_rv*Ft*(1-P_rv)
        AN_rv=AL*AK_rv
        calc={'L':L,'M':M,'N':Nv,'O':O,'P':P,'Q':Q,'R':Rr,'S':Sv,'T':T,'AC':AC,'AD':AD,'AE':AE,'AF':AF,'AG':AG,'AH':AH,'AI':AI,'AJ':AJ,'AK':AK,'AL':AL,'AN':AN}
        pcalc={'角色等级':g,'星盘':h,'专武':ii,'映射装备':j,'旅券':k}
        for nm,c in cols.items():
            a=R(ROW0+i-1,c)
            if a is not None: err[nm]=max(err[nm],abs(calc[nm]-a)/max(abs(a),1e-9))
        for nm,c in prog_cols.items():
            a=R(ROW0+i-1,c)
            if a is not None: err[nm]=max(err[nm],abs(pcalc[nm]-a)/max(abs(a),1e-9))
        out_rows.append((i,g,h,ii,j,k,rv,L,M,Nv,O,P,Q,Rr,Sv,T,AC,AJ,AK,AL,AN,M_rv,O_rv,P_rv,Rr_rv,T_rv,AK_rv,AN_rv))
    wb.close()

    npass=sum(1 for v in err.values() if v<1e-5); ntot=len(err)
    if not args.check:
        owb=openpyxl.Workbook(); bold=Font(bold=True); hf=PatternFill('solid',fgColor='D9E1F2'); ff=PatternFill('solid',fgColor='FFF2CC')
        o=owb.active; o.title='玩家战力曲线'
        o.append(['【自动生成·勿手改】完整链路至最终输出; 黄列=玩家DPS/bossTTK/boss生命'])
        o.append([])
        head=['关卡id','角色进度','星盘进度','专武进度','映射进度','旅券进度','房车进度','裸体防御','额外防御','防御%','算后防御','减伤率','裸体攻击','额外攻击','攻击%','算后攻击','战斗内总倍率','养成总倍率','总单位时间伤害','预计bossTTK','boss生命','额外防御_含房车','算后防御_含房车','减伤率_含房车','额外攻击_含房车','算后攻击_含房车','总单位时间伤害_含房车','boss生命_含房车']
        o.append(head)
        for c in range(1,len(head)+1): o.cell(3,c).font=bold; o.cell(3,c).fill=hf
        for c in (19,20,21,27,28): o.cell(3,c).fill=ff
        for row in out_rows: o.append([round(x,6) if isinstance(x,float) else x for x in row])
        owb.save(OUT)
        old=os.path.join(GROWTH_DIR,'生成_玩家战力曲线_进度.xlsx')
        if os.path.exists(old): os.remove(old)

    md=[]
    def w(s=''):
        print(s); md.append(s)
    w('# 玩家战力曲线 · 完整链路 生成与对照'); w()
    w('> 参数→进度+完整战力链→**最终输出(总单位时间伤害/预计bossTTK/boss生命)**。源表只读。'); w()
    w('- 对照列 %d (5进度+20链路)，**复现通过 %d / %d**。房车整备为新增扩展列，不参与旧源表对照。' % (ntot, npass, ntot)); w()
    w('| 列 | 最大相对误差 | 复现 |'); w('|---|---:|:--:|')
    order=['角色等级','星盘','专武','映射装备','旅券','L','M','N','O','P','Q','R','S','T','AC','AD','AE','AF','AG','AH','AI','AJ','AK','AL','AN']
    name={'AK':'AK总单位时间伤害','AL':'AL预计bossTTK','AN':'AN boss生命'}
    for nm in order:
        w('| %s | %.2e | %s |' % (name.get(nm,nm), err[nm], '✅' if err[nm]<1e-5 else '🔴'))
    w()
    if npass==ntot: w('✅ **全链路逐格复现**，含最终输出(总单位时间伤害/bossTTK/boss生命)。')
    else: w('🔴 有列未通过，需复查。')
    os.makedirs(os.path.dirname(REPORT),exist_ok=True)
    open(REPORT,'w',encoding='utf-8').write('\n'.join(md))
    print('\n[报告] %s'%REPORT)
    if not args.check: print('[生成表] %s'%OUT)

if __name__=='__main__':
    main()
