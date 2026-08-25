# -*- coding: utf-8 -*-
"""读取和写回 ``03_沉淀MD/参数基线`` 中的机器可读参数块。

MD 是参数权威源。旧 ``参数.xlsx`` 已在完成逐格迁移与 MD-only 回归后删除，
正式生成器只读取这里的参数块。
"""

from __future__ import annotations

import json
import os
from copy import deepcopy
from datetime import date, datetime, time
from pathlib import Path

from openpyxl import Workbook


V2_DIR = Path(__file__).resolve().parent.parent
PARAM_DIR = V2_DIR / "03_沉淀MD" / "参数基线"
BLOCK_BEGIN = "<!-- VGAME_PARAMETER_JSON_BEGIN -->"
BLOCK_END = "<!-- VGAME_PARAMETER_JSON_END -->"

PARAMETER_FILES = {
    "共享": PARAM_DIR / "共享参数基线.md",
    "养成": PARAM_DIR / "养成参数基线.md",
    "战斗": PARAM_DIR / "战斗参数基线.md",
    "经济": PARAM_DIR / "经济参数基线.md",
    "节奏": PARAM_DIR / "节奏参数基线.md",
}


def parameter_md_path(scope: str) -> str:
    try:
        return os.fspath(PARAMETER_FILES[scope])
    except KeyError as exc:
        raise KeyError("未知参数域: %s" % scope) from exc


def _resolve(scope_or_path: str | os.PathLike[str]) -> Path:
    value = os.fspath(scope_or_path)
    if value in PARAMETER_FILES:
        return PARAMETER_FILES[value]
    path = Path(value)
    if path.suffix.lower() == ".md":
        return path
    raise ValueError("参数源必须是参数域名称或 .md 路径: %s" % value)


def _decode_cell(value, data_only: bool):
    if not isinstance(value, dict):
        return value
    if "formula" in value:
        return value.get("value") if data_only else value.get("formula")
    kind = value.get("type")
    raw = value.get("value")
    if kind == "datetime":
        return datetime.fromisoformat(raw)
    if kind == "date":
        return date.fromisoformat(raw)
    if kind == "time":
        return time.fromisoformat(raw)
    return raw


def _encode_cell(value):
    if isinstance(value, datetime):
        return {"type": "datetime", "value": value.isoformat()}
    if isinstance(value, date):
        return {"type": "date", "value": value.isoformat()}
    if isinstance(value, time):
        return {"type": "time", "value": value.isoformat()}
    return value


def read_parameter_payload(scope_or_path: str | os.PathLike[str]) -> dict:
    path = _resolve(scope_or_path)
    text = path.read_text(encoding="utf-8")
    try:
        start = text.index(BLOCK_BEGIN) + len(BLOCK_BEGIN)
        end = text.index(BLOCK_END, start)
    except ValueError as exc:
        raise ValueError("参数 MD 缺少机器参数块: %s" % path) from exc
    block = text[start:end].strip()
    if block.startswith("```json"):
        block = block[len("```json") :].lstrip("\r\n")
    if block.endswith("```"):
        block = block[:-3].rstrip("\r\n")
    payload = json.loads(block)
    if payload.get("schema") != "vgame-parameter-workbook/v1":
        raise ValueError("不支持的参数 schema: %s" % payload.get("schema"))
    return payload


def load_parameter_workbook(
    scope_or_path: str | os.PathLike[str], data_only: bool = True, **_kwargs
) -> Workbook:
    """把 MD 参数块还原为内存工作簿，兼容现有读取逻辑。"""

    payload = read_parameter_payload(scope_or_path)
    wb = Workbook()
    wb.remove(wb.active)
    for sheet in payload["sheets"]:
        ws = wb.create_sheet(sheet["name"])
        for row_idx, row in enumerate(sheet.get("rows", []), 1):
            for col_idx, raw in enumerate(row, 1):
                value = _decode_cell(raw, data_only=data_only)
                if value is not None:
                    ws.cell(row_idx, col_idx, value)
    wb._vgame_parameter_scope = payload.get("scope")
    wb._vgame_parameter_path = os.fspath(_resolve(scope_or_path))
    return wb


def _existing_formula_values(payload: dict) -> dict[tuple[str, int, int], dict]:
    result = {}
    for sheet in payload.get("sheets", []):
        for row_idx, row in enumerate(sheet.get("rows", []), 1):
            for col_idx, value in enumerate(row, 1):
                if isinstance(value, dict) and "formula" in value:
                    result[(sheet["name"], row_idx, col_idx)] = value
    return result


def workbook_to_payload(wb: Workbook, base_payload: dict) -> dict:
    payload = deepcopy(base_payload)
    formula_values = _existing_formula_values(base_payload)
    sheets = []
    for ws in wb.worksheets:
        max_row = max((cell.row for row in ws.iter_rows() for cell in row if cell.value is not None), default=0)
        max_col = max((cell.column for row in ws.iter_rows() for cell in row if cell.value is not None), default=0)
        rows = []
        for row_idx in range(1, max_row + 1):
            row = []
            for col_idx in range(1, max_col + 1):
                value = ws.cell(row_idx, col_idx).value
                if isinstance(value, str) and value.startswith("="):
                    old = formula_values.get((ws.title, row_idx, col_idx), {})
                    cached = old.get("value") if old.get("formula") == value else None
                    row.append({"formula": value, "value": cached})
                else:
                    row.append(_encode_cell(value))
            rows.append(row)
        sheets.append({"name": ws.title, "rows": rows})
    payload["sheets"] = sheets
    payload["updated_at"] = datetime.now().astimezone().isoformat(timespec="seconds")
    return payload


def _json_block(payload: dict) -> str:
    body = json.dumps(payload, ensure_ascii=False, indent=2, allow_nan=False)
    return "%s\n```json\n%s\n```\n%s" % (BLOCK_BEGIN, body, BLOCK_END)


def save_parameter_workbook(
    scope_or_path: str | os.PathLike[str], wb: Workbook
) -> Path:
    """只更新 MD 内机器参数块，保留块外的人读说明。"""

    path = _resolve(scope_or_path)
    old_text = path.read_text(encoding="utf-8")
    old_payload = read_parameter_payload(path)
    payload = workbook_to_payload(wb, old_payload)
    start = old_text.index(BLOCK_BEGIN)
    end = old_text.index(BLOCK_END, start) + len(BLOCK_END)
    path.write_text(old_text[:start] + _json_block(payload) + old_text[end:], encoding="utf-8")
    return path


def render_parameter_md(payload: dict) -> str:
    rows = [
        "# %s参数基线" % payload["scope"],
        "",
        "> 本文件是该参数域的唯一参数权威。旧参数工作簿已在完成逐格迁移与回归后删除。",
        "> 参数单元格稳定 ID 为 `%s/<Sheet>!<Cell>`。正式数值文档修改必须同步本参数块或对应系统沉淀 MD。" % payload["scope"],
        "",
        "## 编辑规则",
        "",
        "- 只编辑下方 JSON 参数块中的值；不要新增第二份参数副本。",
        "- `formula` 保存旧公式语义，`value` 是生成器实际读取的确定值。公式变化时必须同步确认并更新 `value`。",
        "- `01_源表` 和正式配置只提供全局事实或证据，不替代本文件中的单系统参数结论。",
        "",
        "## 参数页索引",
        "",
        "| 页签 | 行数 | 列数 | 非空单元格 | 公式单元格 |",
        "|---|---:|---:|---:|---:|",
    ]
    for sheet in payload["sheets"]:
        table = sheet.get("rows", [])
        nrows = len(table)
        ncols = max((len(row) for row in table), default=0)
        nonempty = formulas = 0
        for row in table:
            for value in row:
                if value is not None:
                    nonempty += 1
                if isinstance(value, dict) and "formula" in value:
                    formulas += 1
        rows.append("| %s | %d | %d | %d | %d |" % (sheet["name"], nrows, ncols, nonempty, formulas))
    rows.extend(["", "## 机器可读参数块", "", _json_block(payload), ""])
    return "\n".join(rows)
