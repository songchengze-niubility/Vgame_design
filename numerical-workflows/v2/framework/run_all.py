# -*- coding: utf-8 -*-
"""框架 v2 权威链全量重算入口。

门禁顺序：01_源表 → 03_沉淀MD 一致性、MD 参数完整性、生成器运行、
本次生成物存在性。旧外部 Excel 的历史逐格对照只保留在各自报告中，不再
覆盖当前源表/MD 的权威结论。
"""

from __future__ import annotations

import io
import os
import subprocess
import sys
import time
from dataclasses import dataclass
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill


if hasattr(sys.stdout, "buffer"):
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

V2_DIR = Path(__file__).resolve().parent
SUMMARY = V2_DIR / "汇总总览.xlsx"
MASTER = V2_DIR / "_reports" / "00_总览_全量校验.md"


@dataclass(frozen=True)
class Job:
    name: str
    script: str
    outputs: tuple[str, ...] = ()
    args: tuple[str, ...] = ()


INPUT_GATES = (
    Job("G1/源表→MD", "02_AI规则/sync_global_sources_to_md.py"),
    Job("G2/MD参数完整性", "02_AI规则/check_md_parameters.py"),
)

# 只列入当前权威链和正式交付所需任务。旧的 gen_mapping_equip.py（主属性2
# 迁移切片）和 gen_combat_late.py（旧外部表后期对照）不属于新 04 发布链。
JOBS = (
    Job("战斗/知识页镜像", "战斗/scripts/sync_battle_knowledge_pages.py", ("战斗/镜像_战斗知识工具页.xlsx",)),
    Job("经济/知识页镜像", "经济/scripts/sync_economy_knowledge_pages.py", ("经济/镜像_经济知识工具页.xlsx",)),
    Job("养成/总览与角色升级", "养成/scripts/gen_growth_budget.py", ("养成/生成_养成总览.xlsx", "养成/生成_角色升级.xlsx")),
    Job("养成/房车", "养成/scripts/gen_rv_growth.py", ("养成/生成_房车.xlsx",)),
    Job("养成/专武+炮台", "养成/scripts/gen_growth_modules.py", ("养成/生成_专武.xlsx", "养成/生成_炮台.xlsx")),
    Job("养成/角色星盘", "养成/scripts/gen_constellation.py", ("养成/生成_角色星盘.xlsx",)),
    Job("养成/角色升星", "养成/scripts/gen_star_upgrade.py", ("养成/生成_角色升星.xlsx",)),
    Job("养成/映射装备合成制作8档", "养成/scripts/gen_mapping_crafting_8_tiers.py", ("养成/生成_映射装备合成制作8档.xlsx",)),
    Job("养成/映射装备19页总表", "养成/scripts/gen_mapping_equip_full.py", ("养成/生成_映射装备数值总表.xlsx",)),
    Job("养成/旅券", "养成/scripts/gen_voucher.py", ("养成/生成_旅券属性.xlsx",)),
    Job("养成/战力曲线", "养成/scripts/gen_power_curve.py", ("养成/生成_玩家战力曲线.xlsx",)),
    Job("战斗/生存骨架验证", "战斗/scripts/gen_combat_model.py"),
    Job("战斗/怪物模型", "战斗/scripts/gen_monster_model.py", ("战斗/生成_怪物标准模型.xlsx",)),
    Job("战斗/技能系数验证", "战斗/scripts/gen_skill_coefficients.py"),
    Job("经济/升级成本", "经济/scripts/gen_levelup_cost.py", ("经济/生成_角色升级成本.xlsx",)),
    Job("经济/资源产销总账", "经济/scripts/gen_economy_ledger.py", ("经济/全资源产销总表.xlsx", "经济/模块产销表/01_玩法_主线.xlsx", "经济/模块产销表/22_外循环_章节任务.xlsx")),
    Job("经济/资源养成阅读版", "经济/scripts/gen_economy_readable_views.py", ("经济/资源养成_阅读版.xlsx",)),
    Job("经济/爬塔QA奖励表", "经济/scripts/gen_tower_reward_qa_table.py", ("经济/生成_爬塔首通挂机奖励.xlsx",)),
    Job("节奏/主角等级", "节奏/scripts/gen_hero_level.py", ("节奏/生成_主角等级单源.xlsx",)),
    Job("节奏/无限挑战", "节奏/scripts/gen_difficulty_inf.py", ("节奏/生成_无限挑战难度线.xlsx",)),
    Job("节奏/爬塔难度", "节奏/scripts/gen_difficulty_tower.py", ("节奏/生成_爬塔难度线_天数层数.xlsx",)),
    Job("节奏/全玩法难度线", "节奏/scripts/gen_difficulty_all.py", ("节奏/全玩法难度线.xlsx",)),
    Job("节奏/开启节奏总览", "节奏/scripts/gen_onboarding_pace.py", ("节奏/生成_开启节奏总览.xlsx",)),
)


def decode_output(raw: bytes) -> str:
    for encoding in ("utf-8", "gb18030"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            pass
    return raw.decode("utf-8", "replace")


def last_line(text: str) -> str:
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    return lines[-1] if lines else ""


def run_job(job: Job, *, require_fresh: bool) -> tuple[str, str, str]:
    script = V2_DIR / job.script
    if not script.exists():
        return "FAIL", "脚本不存在", os.fspath(script)
    started_ns = time.time_ns()
    env = os.environ.copy()
    env.setdefault("PYTHONIOENCODING", "utf-8")
    try:
        process = subprocess.run(
            [sys.executable, os.fspath(script), *job.args],
            cwd=os.fspath(V2_DIR),
            capture_output=True,
            timeout=600,
            env=env,
        )
    except Exception as exc:
        return "FAIL", "运行异常", str(exc)
    stdout = decode_output(process.stdout)
    stderr = decode_output(process.stderr)
    detail = last_line(stdout) or last_line(stderr) or "完成"
    if process.returncode != 0:
        return "FAIL", "退出码 %d" % process.returncode, detail

    missing = []
    stale = []
    for relative in job.outputs:
        path = V2_DIR / relative
        if not path.exists() or path.stat().st_size <= 0:
            missing.append(relative)
        elif require_fresh and path.stat().st_mtime_ns < started_ns - 2_000_000_000:
            stale.append(relative)
    if missing:
        return "FAIL", "缺少生成物", "、".join(missing)
    if stale:
        return "FAIL", "生成物未被本次任务刷新", "、".join(stale)
    return "PASS", "通过", detail


def build_summary(gates, results) -> None:
    wb = openpyxl.Workbook()
    bold = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="D9E1F2")
    ok_fill = PatternFill("solid", fgColor="E2EFDA")
    bad_fill = PatternFill("solid", fgColor="FCE4E4")
    ws = wb.active
    ws.title = "校验总览"
    ws.append(["Vgame 数值框架 v2 · 权威链全量重算"])
    ws.cell(1, 1).font = Font(bold=True, size=13)
    ws.append(["权威链：01_源表 + 03_沉淀MD → 本次生成物 → 新 04"])
    ws.append([])
    ws.append(["阶段", "任务", "状态", "说明"])
    for col in range(1, 5):
        ws.cell(4, col).font = bold
        ws.cell(4, col).fill = header_fill
    for phase, rows in (("输入门禁", gates), ("生成", results)):
        for name, status, label, detail in rows:
            ws.append([phase, name, status, "%s；%s" % (label, detail)])
            ws.cell(ws.max_row, 3).fill = ok_fill if status == "PASS" else bad_fill
    ws.append([])
    passed = sum(row[1] == "PASS" for row in results)
    ws.append(["合计", "%d 个生成任务" % len(results), "%d/%d 通过" % (passed, len(results)), "输入门禁 %d/%d" % (sum(row[1] == "PASS" for row in gates), len(gates))])
    ws.cell(ws.max_row, 1).font = bold
    for width, column in zip((14, 28, 12, 80), "ABCD"):
        ws.column_dimensions[column].width = width
    wb.save(SUMMARY)


def write_report(gates, results) -> None:
    MASTER.parent.mkdir(parents=True, exist_ok=True)
    passed = sum(row[1] == "PASS" for row in results)
    lines = [
        "# 框架 v2 权威链全量校验",
        "",
        "> `01_源表 + 03_沉淀MD → 本次生成物`；旧外部 Excel 对照仅作迁移记录。",
        "",
        "## 输入门禁",
        "",
        "| 门禁 | 状态 | 说明 |",
        "|---|---|---|",
    ]
    for name, status, label, detail in gates:
        lines.append("| %s | %s | %s；%s |" % (name, status, label, detail.replace("|", "\\|")))
    lines += ["", "## 生成结果（%d/%d）" % (passed, len(results)), "", "| 任务 | 状态 | 说明 |", "|---|---|---|"]
    for name, status, label, detail in results:
        lines.append("| %s | %s | %s；%s |" % (name, status, label, detail.replace("|", "\\|")))
    lines += [
        "",
        "## 非发布旧诊断",
        "",
        "- `养成/scripts/gen_mapping_equip.py`：旧主属性 2 迁移切片，已从新 04 下线。",
        "- `战斗/scripts/gen_combat_late.py`：旧外部战斗表后期段对照，不是新 04 生成链。",
    ]
    MASTER.write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> int:
    print("=== 框架 v2：权威链全量重算 ===")
    gates = []
    for job in INPUT_GATES:
        status, label, detail = run_job(job, require_fresh=False)
        gates.append((job.name, status, label, detail))
        print("  [%s] %s — %s" % (status, job.name, detail[:100]))
    if any(row[1] != "PASS" for row in gates):
        build_summary(gates, [])
        write_report(gates, [])
        print("[FAIL] 输入门禁未通过，未启动生成")
        return 1

    results = []
    for job in JOBS:
        status, label, detail = run_job(job, require_fresh=True)
        results.append((job.name, status, label, detail))
        print("  [%s] %s — %s" % (status, job.name, detail[:100]))
    build_summary(gates, results)
    write_report(gates, results)

    passed = sum(row[1] == "PASS" for row in results)
    print("=== %d/%d 个生成任务通过；输入门禁 %d/%d ===" % (passed, len(results), sum(row[1] == "PASS" for row in gates), len(gates)))
    print("[汇总Excel] %s" % SUMMARY)
    print("[报告] %s" % MASTER)
    return 0 if passed == len(results) else 1


if __name__ == "__main__":
    raise SystemExit(main())
