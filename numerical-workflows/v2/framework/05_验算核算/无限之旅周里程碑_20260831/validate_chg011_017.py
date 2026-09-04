# -*- coding: utf-8 -*-
"""CHG-011~017 独立复读校验：重读06候选包产物，与沉淀MD口径逐项核对。"""
import json, datetime, os
import openpyxl

BASE = r"D:/Vgame_design/numerical-workflows/v2/framework"
PKG = BASE + "/06_最终配置/无限之旅"
BATCH = BASE + "/05_验算核算/无限之旅周里程碑_20260831"

PACKAGES = {
    1000: [(2, 40), (1, 10000), (100, 20)],
    2000: [(2, 60), (25, 2), (108, 20)],
    3000: [(106, 1), (1, 15000), (6, 5000)],
    4000: [(2, 80), (25, 2), (3028, 20)],
    5000: [(106, 1), (1, 20000), (1074, 20)],
    6000: [(2, 100), (25, 2), (100, 30)],
    7000: [(2, 60), (1, 25000), (108, 30)],
    8000: [(106, 1), (25, 2), (6, 7500)],
    9000: [(2, 60), (1, 30000), (3027, 15)],
    10000: [(2, 60), (1, 35000), (1074, 30)],
    11000: [(2, 60), (25, 1), (101, 10)],
    12000: [(106, 2), (1, 40000), (109, 10)],
    13000: [(2, 60), (1, 45000), (6, 10000)],
    14000: [(2, 60), (25, 1), (3027, 20)],
    15000: [(2, 60), (1, 55000), (1074, 40)],
}
CAPS = {1: 6000, 2: 9000, 3: 12000, 4: 15000}
UNLOCK = {1: "默认开放", 2: "通关难度1＋35级", 3: "通关难度2＋45级", 4: "通关难度3＋55级"}
TOTALS = {
    1: [280, 2, 0, 6, 45000, 50, 20, 5000, 0, 0, 20, 20, 0],
    2: [400, 3, 0, 8, 100000, 50, 50, 12500, 0, 0, 20, 20, 15],
    3: [520, 5, 0, 9, 175000, 50, 50, 12500, 10, 10, 50, 20, 15],
    4: [700, 5, 0, 10, 275000, 50, 50, 22500, 10, 10, 90, 20, 35],
}

checks, fails = [], []
def ck(name, ok, detail=""):
    checks.append({"check": name, "result": "PASS" if ok else "FAIL", "detail": detail})
    if not ok:
        fails.append(f"{name}: {detail}")

def is_white(font):
    c = getattr(font, "color", None)
    return c is not None and getattr(c, "rgb", None) in ("FFFFFFFF", "00FFFFFF")

# ---- NR ----
wb = openpyxl.load_workbook(PKG + "/NR_HighestScoreReward.xlsx")
ws = wb["WeekScoreReward"]
rows = [[ws.cell(r, c).value for c in range(2, 6)] for r in range(5, ws.max_row + 1)]
ck("NR行数=42", ws.max_row == 46 and len(rows) == 42, f"max_row={ws.max_row}")
ck("NR Id=1..42连续", [r[0] for r in rows] == list(range(1, 43)))
ck("NR (难度,分数)唯一且%1000=0", len({(r[1], r[2]) for r in rows}) == 42 and all(r[2] % 1000 == 0 for r in rows))
exp_nr = []
for d in (1, 2, 3, 4):
    for s in range(1000, CAPS[d] + 1, 1000):
        exp_nr.append((d, s, 23000000 + s // 1000))
ck("NR内容=共享阶梯42行", [tuple(r[1:]) for r in rows] == exp_nr)
caps_seen = {d: max(r[2] for r in rows if r[1] == d) for d in (1, 2, 3, 4)}
ck("NR各档上限=6/9/12/15千分", caps_seen == CAPS, str(caps_seen))
ck("NR数据区无白字", all(not is_white(ws.cell(r, c).font) for r in range(5, 47) for c in range(2, 6)))

# ---- Drop ----
wb = openpyxl.load_workbook(PKG + "/Drop_dungeon.xlsx")
ws = wb["Drop"]
drows = [[ws.cell(r, c).value for c in range(2, 16)] for r in range(5, ws.max_row + 1)]
ck("Drop行数=45", ws.max_row == 49 and len(drows) == 45, f"max_row={ws.max_row}")
ck("Drop B列公式=ROW()+200000-4", all(r[0] == "=ROW()+200000-4" for r in drows))
by_pkg = {}
for r in drows:
    by_pkg.setdefault(r[1], []).append(r)
ck("Drop包数=15且ID=23000001~15", sorted(by_pkg) == [23000000 + i for i in range(1, 16)])
ok_pkg = True
for did, rs in by_pkg.items():
    score = did - 23000000
    exp = PACKAGES[score * 1000]
    got = [(r[9], r[10]) for r in rs]
    if got != exp or len(rs) != 3:
        ok_pkg = False
    if not all(r[5] == "fixed_" and r[6] == 1 and r[7] == 1 and r[12] == 10000 and r[4] == 0 for r in rs):
        ok_pkg = False
ck("Drop 15包内容/类型/权重与MD口径一致", ok_pkg)
ck("Drop数据区无白字", all(not is_white(ws.cell(r, c).font) for r in range(5, 50) for c in range(2, 16)))
nr_dropids = {r[3] for r in rows}
ck("NR引用DropId全部存在于Drop", nr_dropids <= set(by_pkg))
keys = sum(q for s, items in PACKAGES.items() for it, q in items if it == 25)
ck("周映射之匙=10把(单档满领)", keys == 10, str(keys))

# ---- 清单 ----
wb = openpyxl.load_workbook(PKG + "/00_配置清单.xlsx")
ws = wb["周奖励预览"]
pv = [[ws.cell(r, c).value for c in range(1, 15)] for r in range(2, ws.max_row + 1)]
ck("预览行数=42", len(pv) == 42, str(len(pv)))
ok_pv = True
for i, p in enumerate(pv):
    d, score, did = p[1], p[3], p[4]
    if did != 23000000 + score // 1000 or p[2] != UNLOCK[d] or p[0] != i + 1:
        ok_pv = False; break
    got = [(p[5], p[7]), (p[8], p[10]), (p[11], p[13])]
    names = [p[6], p[9], p[12]]
    if got != PACKAGES[score] or not all(names):
        ok_pv = False; break
ck("预览42行与阶梯/解锁条件一致", ok_pv)

ws = wb["周奖励汇总"]
recalc = []
for d in (1, 2, 3, 4):
    acc = [0] * 13
    for p in pv:
        if p[1] != d:
            continue
        its = [(p[5], p[7]), (p[8], p[10]), (p[11], p[13])]
        names = [p[6], p[9], p[12]]
        for (it, q), nm in zip(its, names):
            col = {2: 0, 106: 1, 107: 2, 25: 3, 1: 4, 100: 5, 108: 6, 6: 7, 101: 8, 109: 9, 1074: 10, 3028: 11, 3027: 12}.get(it)
            if col is None:
                acc = None; break
            acc[col] += q
        if acc is None:
            break
    recalc.append(acc)
ck("汇总静态值与预览重算一致", recalc == [TOTALS[d] for d in (1, 2, 3, 4)], f"recalc={recalc}")
ck("汇总收容凭证列全为0", all(r[2] == 0 for r in recalc))

status = "PASS_STATIC_WITH_RUNTIME_GATES" if not fails else "FAIL"
report = {
    "batch": "INFINITE-WEEKLY-20260831-CHG017",
    "generated_at": datetime.datetime.now().isoformat(),
    "status": status,
    "scope": "06候选包全量重铺复读校验：NR_HighestScoreReward/Drop_dungeon/00_配置清单；SVN只读未改",
    "design_basis": "03_沉淀MD/经济/06_玩法_无限挑战经济沉淀 CHG-011～CHG-017",
    "checks": checks,
    "weekly_totals_per_difficulty": {str(d): dict(zip(
        ["元晶", "旅券", "收容凭证", "映射之匙", "金币", "低阶记录盘", "泛黄档案", "沿途随笔",
         "中阶记录盘", "录像带", "房车零件", "低级舱室建材3028", "天赋树箱3027"], TOTALS[d])) for d in (1, 2, 3, 4)},
    "runtime_gates": [
        "同分节点跨档复用DropId（23000001~15）需程序确认按Difficult+Score领取无副作用。",
        "每周只领一档、周重置、各档分数隔离仍由服务端实现并待联调验收。",
        "周上限6/9/12/15千分由各难度最高Score行承载，服务端计分口径（本周最高单局）沿用待验证。",
        "正式合SVN时仍须人工预建行并下拉正式表ID公式；本候选不写死正式表ID。",
        "本轮未执行Luban与运行时验证；SVN未改。",
    ],
    "backups": BATCH + "/backup",
}
json.dump(report, open(BATCH + "/published-validation.json", "w", encoding="utf-8"), ensure_ascii=False, indent=2)

chk = json.load(open(PKG + "/无限之旅配置校验.json", encoding="utf-8"))
chk["status"] = status
chk["current_validation"] = "../../05_验算核算/无限之旅周里程碑_20260831/published-validation.json"
chk["validated_at"] = report["generated_at"]
json.dump(chk, open(PKG + "/无限之旅配置校验.json", "w", encoding="utf-8"), ensure_ascii=False, indent=2)

print(json.dumps({"status": status, "fails": fails, "checks": len(checks)}, ensure_ascii=False))
