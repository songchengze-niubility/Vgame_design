# -*- coding: utf-8 -*-
"""严格验算按源表与沉淀 MD 重建的正式数值文档。

默认验算 ``04_数值文档``，并生成本次运行的 Excel/JSON 快照报告。
任何前置检查、清单、数量、文件结构、外链或显式错误值异常都会返回非零。
"""

from __future__ import annotations

import argparse
import importlib.util
import json
import os
import re
import subprocess
import sys
import tempfile
import zipfile
from datetime import datetime
from pathlib import Path
from xml.etree import ElementTree

import openpyxl
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


V2_DIR = Path(__file__).resolve().parents[2]
AI_DIR = V2_DIR / "02_AI规则"
AUDIT_DIR = V2_DIR / "05_验算核算"
DEFAULT_TARGET = V2_DIR / "04_数值文档"
DEFAULT_XLSX_REPORT = AUDIT_DIR / "动态验算总表.xlsx"
DEFAULT_JSON_REPORT = AUDIT_DIR / "动态验算总表.json"

ERROR_VALUES = {
    "#NULL!",
    "#DIV/0!",
    "#VALUE!",
    "#REF!",
    "#NAME?",
    "#NUM!",
    "#N/A",
    "#GETTING_DATA",
    "#SPILL!",
    "#CALC!",
    "#FIELD!",
    "#BLOCKED!",
    "#UNKNOWN!",
    "#CONNECT!",
    "#BUSY!",
}
EXTERNAL_FORMULA_RE = re.compile(
    r"\[(?:\d+|[^\[\]]+\.(?:xlsx?|xlsm|xlsb|xltx|xltm|csv|ods))\]"
    r"[^!\r\n]{0,128}!",
    re.IGNORECASE,
)
MD_LINE_RE = re.compile(
    r"^\[(?P<scope>.+?)\]\s+OK\s+sheets=(?P<sheets>\d+)\s+"
    r"nonempty=(?P<nonempty>\d+)\s+formulas=(?P<formulas>\d+)$"
)

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
PASS_FILL = PatternFill("solid", fgColor="E2F0D9")
FAIL_FILL = PatternFill("solid", fgColor="FCE4D6")
INFO_FILL = PatternFill("solid", fgColor="FFF2CC")
WHITE_BOLD = Font(name="Microsoft YaHei", size=10, bold=True, color="FFFFFF")
BODY_FONT = Font(name="Microsoft YaHei", size=10, color="000000")


def now_iso() -> str:
    return datetime.now().astimezone().isoformat(timespec="seconds")


def compact(value, limit: int = 30000) -> str:
    if value is None:
        return ""
    if isinstance(value, (dict, list, tuple, set)):
        text = json.dumps(value, ensure_ascii=False, default=str)
    else:
        text = str(value)
    return text if len(text) <= limit else text[: limit - 3] + "..."


def run_python(script: Path, *arguments: str, timeout: int = 300) -> dict:
    command = [sys.executable, str(script), *arguments]
    env = os.environ.copy()
    env["PYTHONUTF8"] = "1"
    env["PYTHONIOENCODING"] = "utf-8"
    started_at = now_iso()
    try:
        completed = subprocess.run(
            command,
            cwd=V2_DIR,
            env=env,
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            timeout=timeout,
            check=False,
        )
        return {
            "command": command,
            "started_at": started_at,
            "finished_at": now_iso(),
            "returncode": completed.returncode,
            "stdout": completed.stdout,
            "stderr": completed.stderr,
            "timed_out": False,
        }
    except subprocess.TimeoutExpired as exc:
        return {
            "command": command,
            "started_at": started_at,
            "finished_at": now_iso(),
            "returncode": 124,
            "stdout": compact(exc.stdout),
            "stderr": compact(exc.stderr) or f"检查超时（{timeout} 秒）",
            "timed_out": True,
        }


def normalize_relative(value) -> str:
    path = Path(value)
    if path.is_absolute() or ".." in path.parts:
        raise ValueError(f"期望清单含不安全路径: {value!r}")
    normalized = path.as_posix().lstrip("./")
    if not normalized or Path(normalized).suffix.lower() != ".xlsx":
        raise ValueError(f"期望清单含非 xlsx 路径: {value!r}")
    return normalized


def load_expected_relatives() -> tuple[list[str], str | None]:
    stage_script = AI_DIR / "stage_final_documents_from_md.py"
    try:
        spec = importlib.util.spec_from_file_location("stage_final_documents_from_md", stage_script)
        if spec is None or spec.loader is None:
            raise ImportError(f"无法加载模块: {stage_script}")
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        public_api = getattr(module, "expected_relatives", None)
        if not callable(public_api):
            raise AttributeError("stage_final_documents_from_md.py 缺少 public expected_relatives()")
        raw = public_api()
        if not isinstance(raw, (list, tuple, set)):
            raise TypeError("expected_relatives() 必须返回 list/tuple/set")
        expected = [normalize_relative(item) for item in raw]
        duplicates = sorted({item for item in expected if expected.count(item) > 1})
        if duplicates:
            raise ValueError("期望清单含重复路径: " + ", ".join(duplicates))
        return sorted(expected), None
    except Exception as exc:  # 清单加载失败本身就是门禁失败，仍需输出报告。
        return [], f"{type(exc).__name__}: {exc}"


def explicit_error(cell) -> str | None:
    value = cell.value
    if value is None:
        return None
    normalized = str(value).strip().upper()
    if cell.data_type == "e" or normalized in ERROR_VALUES:
        return normalized
    return None


def defined_name_formulas(workbook) -> list[tuple[str, str]]:
    container = workbook.defined_names
    if callable(getattr(container, "values", None)):
        names = list(container.values())
    else:
        names = list(getattr(container, "definedName", []))
    result = []
    for item in names:
        formula = getattr(item, "attr_text", None)
        if isinstance(formula, str):
            result.append((str(getattr(item, "name", "")), formula))
    return result


def package_external_evidence(path: Path) -> tuple[list[str], list[str], list[str]]:
    external_links: list[str] = []
    external_relations: list[str] = []
    package_errors: list[str] = []
    try:
        with zipfile.ZipFile(path) as archive:
            corrupt_member = archive.testzip()
            if corrupt_member:
                package_errors.append(f"ZIP CRC 失败: {corrupt_member}")
            names = archive.namelist()
            external_links.extend(
                f"package:{name}" for name in names if name.startswith("xl/externalLinks/")
            )
            for name in names:
                if not name.endswith(".rels"):
                    continue
                try:
                    root = ElementTree.fromstring(archive.read(name))
                except Exception as exc:
                    package_errors.append(f"关系 XML 无法解析 {name}: {type(exc).__name__}: {exc}")
                    continue
                for relationship in root:
                    if str(relationship.attrib.get("TargetMode", "")).lower() != "external":
                        continue
                    rel_type = relationship.attrib.get("Type", "").rsplit("/", 1)[-1]
                    target = relationship.attrib.get("Target", "")
                    external_relations.append(f"{name}:{rel_type}:{target}")
    except (zipfile.BadZipFile, OSError) as exc:
        package_errors.append(f"不是有效 xlsx/ZIP: {type(exc).__name__}: {exc}")
    return sorted(set(external_links)), sorted(set(external_relations)), package_errors


def inspect_workbook(path: Path, relative: str, expected: bool) -> dict:
    external_links, external_relations, issues = package_external_evidence(path)
    external_formulas: list[str] = []
    explicit_errors: list[str] = []
    formula_coordinates: dict[str, set[str]] = {}
    sheet_count = 0
    formula_count = 0
    blank_cache_count = 0
    formula_wb = None
    cached_wb = None

    try:
        formula_wb = load_workbook(path, data_only=False, read_only=True, keep_links=True)
        sheet_count = len(formula_wb.sheetnames)
        if sheet_count < 1:
            issues.append("工作簿没有工作表")
        workbook_external_links = list(getattr(formula_wb, "_external_links", []) or [])
        if workbook_external_links:
            external_links.append(f"openpyxl:externalLinks={len(workbook_external_links)}")

        for name, formula in defined_name_formulas(formula_wb):
            if EXTERNAL_FORMULA_RE.search(formula):
                external_formulas.append(f"defined-name:{name}:{formula}")

        for worksheet in formula_wb.worksheets:
            coordinates: set[str] = set()
            for row in worksheet.iter_rows():
                for cell in row:
                    error = explicit_error(cell)
                    if error:
                        explicit_errors.append(f"formula-view:{worksheet.title}!{cell.coordinate}:{error}")
                    value = cell.value
                    if cell.data_type == "f" or (isinstance(value, str) and value.startswith("=")):
                        coordinates.add(cell.coordinate)
                        formula_count += 1
                        if EXTERNAL_FORMULA_RE.search(str(value)):
                            external_formulas.append(f"cell:{worksheet.title}!{cell.coordinate}:{value}")
            formula_coordinates[worksheet.title] = coordinates
    except Exception as exc:
        issues.append(f"公式视图无法打开: {type(exc).__name__}: {exc}")
    finally:
        if formula_wb is not None:
            formula_wb.close()

    try:
        cached_wb = load_workbook(path, data_only=True, read_only=True, keep_links=False)
        if sheet_count and cached_wb.sheetnames != list(formula_coordinates):
            issues.append("公式视图与缓存视图工作表清单不一致")
        for worksheet in cached_wb.worksheets:
            formula_cells = formula_coordinates.get(worksheet.title, set())
            seen_formula_cells: set[str] = set()
            for row in worksheet.iter_rows():
                for cell in row:
                    # read_only 模式会为未落盘的空格返回 EmptyCell；它没有
                    # coordinate 属性，也不属于任何公式缓存坐标。
                    coordinate = getattr(cell, "coordinate", None)
                    error = explicit_error(cell)
                    if error:
                        explicit_errors.append(
                            f"cache-view:{worksheet.title}!{coordinate or '(empty)'}:{error}"
                        )
                    if coordinate and coordinate in formula_cells:
                        seen_formula_cells.add(coordinate)
                        if cell.value is None:
                            blank_cache_count += 1
            blank_cache_count += len(formula_cells - seen_formula_cells)
    except Exception as exc:
        issues.append(f"缓存视图无法打开: {type(exc).__name__}: {exc}")
    finally:
        if cached_wb is not None:
            cached_wb.close()

    external_links = sorted(set(external_links))
    external_relations = sorted(set(external_relations))
    external_formulas = sorted(set(external_formulas))
    explicit_errors = sorted(set(explicit_errors))
    if not expected:
        issues.append("路径不在期望清单")
    if external_links:
        issues.append(f"存在 externalLinks: {len(external_links)}")
    if external_relations:
        issues.append(f"存在外部关系: {len(external_relations)}")
    if external_formulas:
        issues.append(f"公式含 [book] 外链: {len(external_formulas)}")
    if explicit_errors:
        issues.append(f"存在显式错误值: {len(explicit_errors)}")

    evidence = {
        "external_links": external_links,
        "external_relations": external_relations,
        "external_formulas": external_formulas,
        "explicit_errors": explicit_errors,
    }
    return {
        "path": relative,
        "expected": expected,
        "exists": True,
        "status": "PASS" if not issues else "FAIL",
        "sheet_count": sheet_count,
        "formula_count": formula_count,
        "blank_or_missing_formula_cache_count": blank_cache_count,
        "explicit_error_count": len(explicit_errors),
        "external_link_count": len(external_links),
        "external_relation_count": len(external_relations),
        "external_formula_count": len(external_formulas),
        "issues": issues,
        "evidence": evidence,
    }


def parse_md_integrity(command_result: dict) -> list[dict]:
    rows = []
    for line in command_result.get("stdout", "").splitlines():
        match = MD_LINE_RE.match(line.strip())
        if not match:
            continue
        rows.append(
            {
                "scope": match.group("scope"),
                "status": "PASS",
                "sheets": int(match.group("sheets")),
                "nonempty": int(match.group("nonempty")),
                "formulas": int(match.group("formulas")),
                "detail": "MD 参数块可解析且内存工作簿结构一致",
            }
        )
    if command_result["returncode"] != 0 or not rows:
        rows.append(
            {
                "scope": "命令级检查",
                "status": "FAIL",
                "sheets": "",
                "nonempty": "",
                "formulas": "",
                "detail": compact(command_result.get("stderr") or command_result.get("stdout")),
            }
        )
    return rows


def add_sheet(workbook: Workbook, name: str, headers: list[str], rows: list[list]) -> None:
    worksheet = workbook.create_sheet(name)
    worksheet.append(headers)
    for row in rows:
        worksheet.append([compact(value) for value in row])
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.row_dimensions[1].height = 24

    for cell in worksheet[1]:
        cell.fill = HEADER_FILL
        cell.font = WHITE_BOLD
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = BODY_FONT
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if cell.value == "PASS":
                cell.fill = PASS_FILL
            elif cell.value == "FAIL":
                cell.fill = FAIL_FILL
            elif cell.value == "INFO":
                cell.fill = INFO_FILL

    for column in worksheet.columns:
        letter = column[0].column_letter
        max_length = max(len(compact(cell.value, 200)) for cell in column)
        worksheet.column_dimensions[letter].width = min(max(max_length + 2, 10), 60)


def write_xlsx(report: dict, path: Path) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.creator = "Vgame 数值文档动态验算"
    workbook.properties.title = "动态验算总表"

    summary_rows = [
        [item["gate"], item["status"], item["count"], item["detail"]]
        for item in report["run_summary"]
    ]
    add_sheet(workbook, "RunSummary", ["门禁", "状态", "数量", "说明"], summary_rows)

    differences = report["source_md_diff"].get("differences", [])
    source_rows = [
        [
            item.get("parameter_id", ""),
            "FAIL",
            item.get("owner", ""),
            item.get("old", ""),
            item.get("new", ""),
            item.get("source", ""),
            item.get("transform", ""),
        ]
        for item in differences
    ]
    if not source_rows:
        source_command = report.get("source_md_command", {})
        if source_command.get("returncode", 0) != 0 or report["source_md_diff"].get("report_error"):
            source_rows = [[
                "检查失败",
                "FAIL",
                "",
                "",
                "",
                "",
                report["source_md_diff"].get("report_error")
                or compact(source_command.get("stderr") or source_command.get("stdout")),
            ]]
        else:
            source_rows = [["无差异", "PASS", "", "", "", "", ""]]
    add_sheet(
        workbook,
        "SourceMdDiff",
        ["参数ID", "状态", "Owner", "MD当前值", "源表/派生应为", "来源", "变换"],
        source_rows,
    )

    md_rows = [
        [
            item.get("scope", ""),
            item.get("status", ""),
            item.get("sheets", ""),
            item.get("nonempty", ""),
            item.get("formulas", ""),
            item.get("detail", ""),
        ]
        for item in report["md_integrity"]
    ]
    add_sheet(
        workbook,
        "MdIntegrity",
        ["参数域", "状态", "工作表数", "非空单元格", "公式数", "详情"],
        md_rows,
    )

    artifact_rows = []
    for item in report["artifact_check"]:
        artifact_rows.append(
            [
                item.get("path", ""),
                item.get("status", ""),
                item.get("expected", False),
                item.get("exists", False),
                item.get("sheet_count", ""),
                item.get("formula_count", ""),
                item.get("blank_or_missing_formula_cache_count", ""),
                item.get("explicit_error_count", ""),
                item.get("external_link_count", ""),
                item.get("external_relation_count", ""),
                item.get("external_formula_count", ""),
                compact(item.get("issues", [])),
                compact(item.get("evidence", {})),
            ]
        )
    add_sheet(
        workbook,
        "ArtifactCheck",
        [
            "相对路径",
            "状态",
            "在期望清单",
            "文件存在",
            "工作表数",
            "公式数",
            "空/缺公式缓存（信息）",
            "显式错误值",
            "externalLinks",
            "外部关系",
            "[book]公式外链",
            "问题",
            "证据",
        ],
        artifact_rows,
    )

    unresolved_rows = [
        [
            item.get("category", ""),
            item.get("id", ""),
            item.get("status", "FAIL"),
            item.get("owner", ""),
            item.get("old", ""),
            item.get("new", ""),
            item.get("source", ""),
            item.get("reason", ""),
        ]
        for item in report["unresolved_overrides"]
    ] or [["无", "无未解析覆盖/漂移", "PASS", "", "", "", "", ""]]
    add_sheet(
        workbook,
        "UnresolvedOverrides",
        ["类别", "ID", "状态", "Owner", "当前值", "应为", "来源", "未解析原因"],
        unresolved_rows,
    )

    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(path.stem + ".tmp" + path.suffix)
    workbook.save(temporary)
    workbook.close()
    temporary.replace(path)


def write_json(report: dict, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(path.name + ".tmp")
    temporary.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    temporary.replace(path)


def resolve_path(value: str, default: Path) -> Path:
    path = Path(value) if value else default
    if not path.is_absolute():
        path = V2_DIR / path
    return path.resolve()


def validate(args) -> tuple[dict, bool]:
    started_at = now_iso()
    target = resolve_path(args.target, DEFAULT_TARGET)

    descriptor, temporary_name = tempfile.mkstemp(prefix="vgame-source-md-", suffix=".json")
    os.close(descriptor)
    temporary_source_report = Path(temporary_name)
    try:
        source_command = run_python(
            AI_DIR / "sync_global_sources_to_md.py",
            "--report",
            str(temporary_source_report),
        )
        try:
            source_report = json.loads(temporary_source_report.read_text(encoding="utf-8"))
        except Exception as exc:
            source_report = {
                "schema": "vgame-source-md-sync/v1",
                "difference_count": None,
                "differences": [],
                "report_error": f"{type(exc).__name__}: {exc}",
            }
    finally:
        temporary_source_report.unlink(missing_ok=True)

    md_command = run_python(AI_DIR / "check_md_parameters.py")
    md_integrity = parse_md_integrity(md_command)

    expected, manifest_error = load_expected_relatives()
    expected_set = set(expected)
    actual_paths: list[Path] = []
    target_error = None
    if not target.is_dir():
        target_error = f"目标目录不存在: {target}"
    else:
        actual_paths = sorted(
            (path for path in target.rglob("*") if path.is_file() and path.suffix.lower() == ".xlsx"),
            key=lambda item: item.relative_to(target).as_posix(),
        )
    actual_relatives = [path.relative_to(target).as_posix() for path in actual_paths]
    actual_set = set(actual_relatives)
    missing = sorted(expected_set - actual_set)
    extra = sorted(actual_set - expected_set)

    artifact_check = [
        inspect_workbook(path, relative, relative in expected_set)
        for path, relative in zip(actual_paths, actual_relatives)
    ]
    for relative in missing:
        artifact_check.append(
            {
                "path": relative,
                "expected": True,
                "exists": False,
                "status": "FAIL",
                "sheet_count": 0,
                "formula_count": 0,
                "blank_or_missing_formula_cache_count": 0,
                "explicit_error_count": 0,
                "external_link_count": 0,
                "external_relation_count": 0,
                "external_formula_count": 0,
                "issues": ["期望文件缺失"],
                "evidence": {},
            }
        )
    if target_error:
        artifact_check.append(
            {
                "path": "<target>",
                "expected": False,
                "exists": False,
                "status": "FAIL",
                "sheet_count": 0,
                "formula_count": 0,
                "blank_or_missing_formula_cache_count": 0,
                "explicit_error_count": 0,
                "external_link_count": 0,
                "external_relation_count": 0,
                "external_formula_count": 0,
                "issues": [target_error],
                "evidence": {},
            }
        )
    artifact_check.sort(key=lambda item: item["path"])

    source_differences = source_report.get("differences", []) or []
    unresolved = [
        {
            "category": "source_md_drift",
            "id": item.get("parameter_id", ""),
            "status": "FAIL",
            "owner": item.get("owner", ""),
            "old": item.get("old", ""),
            "new": item.get("new", ""),
            "source": item.get("source", ""),
            "reason": "source_owned/derived 参数与沉淀 MD 不一致，未解析",
        }
        for item in source_differences
    ]
    if source_command["returncode"] != 0 and not source_differences:
        unresolved.append(
            {
                "category": "source_md_check_failure",
                "id": "sync_global_sources_to_md.py",
                "status": "FAIL",
                "owner": "",
                "old": "",
                "new": "",
                "source": "01_源表 / 参数 MD",
                "reason": compact(source_command.get("stderr") or source_command.get("stdout")),
            }
        )

    manifest_pass = manifest_error is None and len(expected) == args.expected_count
    count_pass = target_error is None and len(actual_relatives) == args.expected_count
    path_pass = manifest_error is None and not missing and not extra
    artifact_pass = target_error is None and all(
        item["status"] == "PASS" for item in artifact_check
    )
    source_pass = (
        source_command["returncode"] == 0
        and source_report.get("difference_count") == 0
        and not source_report.get("report_error")
    )
    md_pass_rows = [item for item in md_integrity if item.get("status") == "PASS"]
    md_pass = (
        md_command["returncode"] == 0
        and len(md_pass_rows) == 5
        and all(item.get("status") == "PASS" for item in md_integrity)
    )
    overall_pass = all((source_pass, md_pass, manifest_pass, count_pass, path_pass, artifact_pass))

    run_summary = [
        {
            "gate": "G1 SourceMdSync",
            "status": "PASS" if source_pass else "FAIL",
            "count": source_report.get("difference_count", ""),
            "detail": "01_源表 与 source_owned/derived MD 一致"
            if source_pass
            else compact(source_command.get("stderr") or source_command.get("stdout")),
        },
        {
            "gate": "G2 MdIntegrity",
            "status": "PASS" if md_pass else "FAIL",
            "count": len([item for item in md_integrity if item.get("status") == "PASS"]),
            "detail": "check_md_parameters.py exit=0"
            if md_pass
            else compact(md_command.get("stderr") or md_command.get("stdout")),
        },
        {
            "gate": "G3 ExpectedManifest",
            "status": "PASS" if manifest_pass else "FAIL",
            "count": len(expected),
            "detail": "expected_relatives() 清单有效且数量符合预期"
            if manifest_pass
            else (manifest_error or f"期望清单={len(expected)}，命令行预期={args.expected_count}"),
        },
        {
            "gate": "G4 ArtifactCount",
            "status": "PASS" if count_pass else "FAIL",
            "count": len(actual_relatives),
            "detail": f"目标应恰有 {args.expected_count} 个 xlsx"
            + (f"；{target_error}" if target_error else ""),
        },
        {
            "gate": "G5 ArtifactPathSet",
            "status": "PASS" if path_pass else "FAIL",
            "count": len(missing) + len(extra),
            "detail": compact({"missing": missing, "extra": extra}),
        },
        {
            "gate": "G6 ArtifactContent",
            "status": "PASS" if artifact_pass else "FAIL",
            "count": len([item for item in artifact_check if item["status"] == "FAIL"]),
            "detail": "逐文件可打开、有工作表、无外链、无显式错误值；空公式缓存仅记录不阻断",
        },
        {
            "gate": "FINAL",
            "status": "PASS" if overall_pass else "FAIL",
            "count": len([item for item in (source_pass, md_pass, manifest_pass, count_pass, path_pass, artifact_pass) if not item]),
            "detail": f"target={target}",
        },
    ]

    report = {
        "schema": "vgame-rebuild-validation/v1",
        "generated_at": now_iso(),
        "started_at": started_at,
        "target": str(target),
        "expected_count": args.expected_count,
        "actual_count": len(actual_relatives),
        "overall_status": "PASS" if overall_pass else "FAIL",
        "run_summary": run_summary,
        "source_md_diff": source_report,
        "source_md_command": source_command,
        "md_integrity": md_integrity,
        "md_integrity_command": md_command,
        "expected_relatives": expected,
        "actual_relatives": actual_relatives,
        "missing_relatives": missing,
        "extra_relatives": extra,
        "artifact_check": artifact_check,
        "unresolved_overrides": unresolved,
    }
    return report, overall_pass


def main() -> int:
    parser = argparse.ArgumentParser(description="严格动态验算重建后的数值文档")
    parser.add_argument("--target", default=str(DEFAULT_TARGET), help="待验算目录，默认 04_数值文档")
    parser.add_argument("--expected-count", type=int, default=44, help="预期 xlsx 总数，默认 44")
    parser.add_argument("--xlsx-report", default=str(DEFAULT_XLSX_REPORT), help="Excel 报告路径")
    parser.add_argument("--json-report", default=str(DEFAULT_JSON_REPORT), help="JSON 报告路径")
    args = parser.parse_args()
    if args.expected_count < 0:
        parser.error("--expected-count 必须大于等于 0")

    xlsx_report = resolve_path(args.xlsx_report, DEFAULT_XLSX_REPORT)
    json_report = resolve_path(args.json_report, DEFAULT_JSON_REPORT)
    try:
        report, passed = validate(args)
    except Exception as exc:
        report = {
            "schema": "vgame-rebuild-validation/v1",
            "generated_at": now_iso(),
            "target": str(resolve_path(args.target, DEFAULT_TARGET)),
            "expected_count": args.expected_count,
            "actual_count": None,
            "overall_status": "FAIL",
            "run_summary": [
                {
                    "gate": "UNHANDLED",
                    "status": "FAIL",
                    "count": 1,
                    "detail": f"{type(exc).__name__}: {exc}",
                }
            ],
            "source_md_diff": {"differences": []},
            "md_integrity": [
                {
                    "scope": "未完成",
                    "status": "FAIL",
                    "sheets": "",
                    "nonempty": "",
                    "formulas": "",
                    "detail": f"{type(exc).__name__}: {exc}",
                }
            ],
            "artifact_check": [],
            "unresolved_overrides": [
                {
                    "category": "unhandled_exception",
                    "id": type(exc).__name__,
                    "status": "FAIL",
                    "reason": str(exc),
                }
            ],
        }
        passed = False

    report["xlsx_report"] = str(xlsx_report)
    report["json_report"] = str(json_report)
    try:
        write_xlsx(report, xlsx_report)
        write_json(report, json_report)
    except Exception as exc:
        print(f"[FAIL] 报告写入失败: {type(exc).__name__}: {exc}", file=sys.stderr)
        return 2

    print(f"[{report['overall_status']}] target={report['target']}")
    print(f"[REPORT] {xlsx_report}")
    print(f"[REPORT] {json_report}")
    return 0 if passed else 1


if __name__ == "__main__":
    raise SystemExit(main())
