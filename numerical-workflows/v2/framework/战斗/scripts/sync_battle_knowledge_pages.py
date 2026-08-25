# -*- coding: utf-8 -*-
"""
sync_battle_knowledge_pages.py — 战斗知识/工具页选择性镜像 (P0)

只把不能直接参数量化替代的知识页、规则页和工具页纳入 v2：
  - 战斗属性
  - 战斗公式
  - 标准模型（技能系数公式依赖页，保留上下文）
  - 技能系数

怪物标准模型、玩家战力曲线、各类难度线等派生大表继续走参数 → 生成 → 对照。
Boss技能系数与草稿页不进入 P0 镜像。
"""
import io
import os
import sys
import warnings

warnings.filterwarnings("ignore")

import openpyxl

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

HERE = os.path.dirname(os.path.abspath(__file__))
COMBAT_DIR = os.path.dirname(HERE)
V2_DIR = os.path.dirname(COMBAT_DIR)
NUM_DIR = os.path.dirname(V2_DIR)

SRC = os.path.join(NUM_DIR, "Vgame基本战斗框架.xlsx")
OUT = os.path.join(COMBAT_DIR, "镜像_战斗知识工具页.xlsx")
REPORT = os.path.join(COMBAT_DIR, "_reports", "P0_战斗知识工具页_1比1镜像.md")

KEEP_SHEETS = [
    ("战斗属性", "知识/规则表", "战斗属性定义、含义和基础口径，不能由参数表替代。"),
    ("战斗公式", "知识/规则表", "战斗结算公式与常量口径，必须保留原始上下文。"),
    ("标准模型", "公式依赖页", "技能系数公式引用该页；结构化模型仍由生成器负责。"),
    ("技能系数", "工具/计算器表", "角色技能效果定价器，先保留原表，再做结构化校验。"),
]

EXCLUDED_SHEETS = [
    "Boss技能系数",
    "草稿-技能系数公式",
    "怪物标准模型",
    "怪物标准模型-可视版",
    "怪物标准模型趋势",
    "玩家战力曲线",
    "参考-养成属性规划（从养成属性复制）",
    "主角等级（从新手引导复制） ",
    "玩家等级对应关卡（从新手节奏复制）",
    "爬塔难度线",
    "玩法难度线",
    "无限挑战难度线",
]


def sheet_stats(ws):
    nonempty = 0
    formulas = 0
    for row in ws.iter_rows():
        for cell in row:
            value = cell.value
            if value not in (None, ""):
                nonempty += 1
                if is_formula(value):
                    formulas += 1
    return {
        "rows": ws.max_row,
        "cols": ws.max_column,
        "nonempty": nonempty,
        "formulas": formulas,
        "merged": len(ws.merged_cells.ranges),
    }


def is_formula(value):
    return (
        isinstance(value, str) and value.startswith("=")
    ) or (
        hasattr(value, "text") and hasattr(value, "ref")
    )


def comparable_value(value):
    if hasattr(value, "text") and hasattr(value, "ref"):
        return ("ArrayFormula", value.ref, value.text)
    return value


def compare_sheet(src_ws, out_ws):
    max_row = max(src_ws.max_row, out_ws.max_row)
    max_col = max(src_ws.max_column, out_ws.max_column)
    checked = 0
    fails = []
    for r in range(1, max_row + 1):
        for c in range(1, max_col + 1):
            checked += 1
            src_value = comparable_value(src_ws.cell(r, c).value)
            out_value = comparable_value(out_ws.cell(r, c).value)
            if src_value != out_value:
                fails.append((src_ws.cell(r, c).coordinate, src_value, out_value))
                if len(fails) >= 20:
                    return checked, fails
    return checked, fails


def write_report(rows, excluded, total_checked, total_fails):
    os.makedirs(os.path.dirname(REPORT), exist_ok=True)
    lines = [
        "# P0 · 战斗知识/工具页 1:1 镜像",
        "",
        "> 纠偏记录：这里不再做 `Vgame基本战斗框架.xlsx` 整本镜像，只镜像不能直接参数量化的知识页、规则页和工具页。",
        "",
        "## 结论",
        "",
        f"- 源表: `{SRC}`",
        f"- 产物: `{OUT}`",
        f"- 保留页签: {len(KEEP_SHEETS)} 个。",
        f"- 内容/公式逐格校验: {total_checked - total_fails}/{total_checked} 通过。",
        "- `Boss技能系数` 和 `草稿-技能系数公式` 不进入本镜像；Boss 技能系数如需使用，走结构化生成/校验链路。",
        "- 派生大表继续走正常生成链路，不通过镜像替代。",
        "",
        "## 保留页签",
        "",
        "| 页签 | 类型 | 保留原因 | 行×列 | 非空 | 公式 | 合并格 | 校验 |",
        "|---|---|---|---:|---:|---:|---:|---|",
    ]
    for row in rows:
        stats = row["stats"]
        lines.append(
            "| {sheet} | {kind} | {reason} | {rows}×{cols} | {nonempty} | {formulas} | {merged} | {status} |".format(
                sheet=row["sheet"],
                kind=row["kind"],
                reason=row["reason"],
                rows=stats["rows"],
                cols=stats["cols"],
                nonempty=stats["nonempty"],
                formulas=stats["formulas"],
                merged=stats["merged"],
                status=row["status"],
            )
        )

    lines += [
        "",
        "## 不镜像页签",
        "",
        "这些页签不进入 P0 镜像；其中派生模型、趋势视图、外部复制源或难度曲线继续由对应生成器/参数链路维护：",
        "",
    ]
    for sheet in excluded:
        lines.append(f"- `{sheet}`")

    if total_fails:
        lines += ["", "## 不一致样例", ""]
        for row in rows:
            for cell, src_value, out_value in row["fails"]:
                lines.append(f"- `{row['sheet']}`!`{cell}`: 源=`{src_value}` / 镜像=`{out_value}`")

    with open(REPORT, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))


def main():
    if not os.path.exists(SRC):
        raise FileNotFoundError(SRC)

    src_wb = openpyxl.load_workbook(SRC, data_only=False)
    keep_names = [item[0] for item in KEEP_SHEETS]
    missing = [name for name in keep_names if name not in src_wb.sheetnames]
    if missing:
        src_wb.close()
        raise RuntimeError("源表缺少必要页签: %s" % ", ".join(missing))

    mirror_wb = openpyxl.load_workbook(SRC, data_only=False)
    for sheet in list(mirror_wb.sheetnames):
        if sheet not in keep_names:
            del mirror_wb[sheet]
    mirror_wb.active = 0
    mirror_wb.save(OUT)
    mirror_wb.close()

    out_wb = openpyxl.load_workbook(OUT, data_only=False)
    rows = []
    total_checked = 0
    total_fails = 0
    for sheet, kind, reason in KEEP_SHEETS:
        stats = sheet_stats(src_wb[sheet])
        checked, fails = compare_sheet(src_wb[sheet], out_wb[sheet])
        total_checked += checked
        total_fails += len(fails)
        rows.append({
            "sheet": sheet,
            "kind": kind,
            "reason": reason,
            "stats": stats,
            "fails": fails,
            "status": "✅ 内容/公式一致" if not fails else "🔴 不一致",
        })

    excluded = [name for name in EXCLUDED_SHEETS if name in src_wb.sheetnames]
    write_report(rows, excluded, total_checked, total_fails)
    out_wb.close()
    src_wb.close()

    if total_fails:
        print(f"不一致: {total_fails}")
    else:
        print(f"选择页内容/公式逐格复现: {len(KEEP_SHEETS)}/{len(KEEP_SHEETS)}")
    print(f"[镜像] {OUT}")
    print(f"[报告] {REPORT}")


if __name__ == "__main__":
    main()
