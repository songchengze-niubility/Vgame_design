# -*- coding: utf-8 -*-
"""
gen_skill_coefficients.py — 技能系数 / Boss技能系数 生成器 (P3c)

把 Vgame基本战斗框架.xlsx 中的「技能系数」「Boss技能系数」纳入 v2:
  - 角色技能系数: 按技能定价器公式重算关键输出。
  - Boss技能系数: 按参数综合系数完整重算最终技能系数。

源表只读，仅用于逐格/逐项对照。参数来源为战斗参数基线 MD 中的页签:
  - 技能系数参数
  - Boss技能系数参数
  - Boss技能输入

用法:
    python gen_skill_coefficients.py
    python gen_skill_coefficients.py --export-xlsx
"""
import argparse
import io
import math
import os
import sys
import warnings

warnings.filterwarnings("ignore")

import openpyxl
from openpyxl.styles import Font, PatternFill

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

HERE = os.path.dirname(os.path.abspath(__file__))
COMBAT_DIR = os.path.dirname(HERE)
V2_DIR = os.path.dirname(COMBAT_DIR)
NUM_DIR = os.path.dirname(V2_DIR)
sys.path.insert(0, V2_DIR)
from _shared.md_parameter_reader import load_parameter_workbook, parameter_md_path

COMBAT_PARAM = parameter_md_path("战斗")
OUT = os.path.join(COMBAT_DIR, "生成_技能系数.xlsx")
REPORT = os.path.join(COMBAT_DIR, "_reports", "P3c_技能系数_Boss技能系数_生成与对照.md")
SRC = os.path.join(NUM_DIR, "Vgame基本战斗框架.xlsx")


def excel_round(value, ndigits=0):
    scale = 10 ** ndigits
    if value >= 0:
        return math.floor(value * scale + 0.5) / scale
    return math.ceil(value * scale - 0.5) / scale


def mround(value, multiple):
    if multiple == 0:
        return 0
    return round(value / multiple) * multiple


def almost(a, b, tol=1e-8):
    if a in ("", None) and b in ("", None):
        return True
    try:
        a = float(a)
        b = float(b)
    except (TypeError, ValueError):
        return a == b
    return abs(a - b) <= max(tol, abs(b) * 1e-8)


def read_key_values(ws, key_col=1, value_col=2, start=1):
    values = {}
    for r in range(start, ws.max_row + 1):
        key = ws.cell(r, key_col).value
        if key is not None:
            values[key] = ws.cell(r, value_col).value
    return values


def read_params():
    wb = load_parameter_workbook(COMBAT_PARAM, data_only=True)
    if "技能系数参数" not in wb.sheetnames or "Boss技能系数参数" not in wb.sheetnames or "Boss技能输入" not in wb.sheetnames:
        wb.close()
        raise RuntimeError("战斗参数基线 MD 缺少 技能系数参数 / Boss技能系数参数 / Boss技能输入 页，请先同步参数块。")

    skill_ws = wb["技能系数参数"]
    skill = read_key_values(skill_ws)
    qualities = {}
    split_basic = []
    split_skill = []
    atk_hp = {}
    for r in range(2, skill_ws.max_row + 1):
        group = skill_ws.cell(r, 4).value
        key = skill_ws.cell(r, 5).value
        val = skill_ws.cell(r, 6).value
        if group == "品质系数" and key:
            qualities[key] = val
        elif group == "拆分系数_普攻必杀" and key is not None:
            split_basic.append((key, val))
        elif group == "拆分系数_技能" and key is not None:
            split_skill.append((key, val))
        elif group == "攻血比" and key:
            atk_hp[key] = val

    boss_param_ws = wb["Boss技能系数参数"]
    boss = read_key_values(boss_param_ws)
    tier = {}
    block = {}
    area = {}
    pierce = {}
    for r in range(2, boss_param_ws.max_row + 1):
        group = boss_param_ws.cell(r, 4).value
        key = boss_param_ws.cell(r, 5).value
        v1 = boss_param_ws.cell(r, 6).value
        v2 = boss_param_ws.cell(r, 7).value
        if group == "技能档位" and key:
            tier[key] = (v1, v2)
        elif group == "释放屏蔽" and key:
            block[key] = v1
        elif group == "范围类型" and key:
            area[key] = v1
        elif group == "穿透类型" and key:
            pierce[key] = v1

    inputs_ws = wb["Boss技能输入"]
    inputs = []
    headers = [inputs_ws.cell(1, c).value for c in range(1, inputs_ws.max_column + 1)]
    for r in range(2, inputs_ws.max_row + 1):
        row = {headers[c - 1]: inputs_ws.cell(r, c).value for c in range(1, inputs_ws.max_column + 1)}
        if row.get("Boss名") or row.get("技能名") or row.get("技能档位"):
            inputs.append(row)
    wb.close()
    return skill, qualities, split_basic, split_skill, atk_hp, boss, tier, block, area, pierce, inputs


def pow_ratio(a, b, exp):
    if a in (None, 0) or b in (None, 0):
        return 0
    return math.pow(float(a), float(exp)) / math.pow(float(b), float(exp))


def compute_skill_outputs(skill, qualities, split_basic, split_skill):
    q_name = skill["技能当前品质"]
    q = qualities[q_name]
    base = skill["技能系数"]
    outputs = []

    def add(name, value, source_cell, note=""):
        outputs.append((name, value, source_cell, note))

    add("单体伤害系数", skill["单体系数"] / skill["单体攻速"] * q * base, "C8")
    add("穿透子弹伤害系数", skill["穿透子弹系数"] / skill["穿透子弹攻速"] * q * base, "C14")
    add("跟踪子弹伤害系数", skill["跟踪子弹系数"] / skill["跟踪子弹攻速"] * q * base, "C20")
    add(
        "AOE子弹伤害系数",
        pow_ratio(skill["AOE标准范围半径"], skill["AOE实际范围半径"], skill["AOE衰减系数"])
        * skill["AOE系数"] * q * base / skill["AOE攻速"],
        "C30",
    )
    dot_duration = pow_ratio(skill["DOT持续时间"], skill["DOT标准持续时间"], skill["DOT持续衰减系数"]) / skill["DOT持续时间"] * skill["DOT标准持续时间"]
    dot_layers = pow_ratio(skill["DOT标准最大层数"], skill["DOT最大层数"], skill["DOT最大层数衰减系数"]) * skill["DOT最大层数"] / skill["DOT标准最大层数"]
    add("DOT持续系数", dot_duration, "D37")
    add("DOT最大层数系数", dot_layers, "D40")
    add("DOT单层系数", skill["DOT系数"] / skill["DOT持续时间"] * dot_duration * dot_layers * q * base, "C41")

    heal_people = pow_ratio(skill["回血人数"], skill["标准回血人数"], skill["回血人数衰减系数"]) * skill["标准回血人数"] / skill["回血人数"]
    add("回血人数系数", heal_people, "D49")
    add("最终回血系数", skill["回血系数"] / skill["回血人数"] * skill["标准回血人数"] * heal_people * q * base, "C50")

    shield_people = pow_ratio(skill["护盾人数"], skill["标准护盾人数"], skill["护盾人数衰减系数"]) * skill["标准护盾人数"] / skill["护盾人数"]
    shield_duration = min(max(skill["护盾持续时间"], skill["护盾持续下限"]), skill["护盾持续上限"])
    shield_time = pow_ratio(shield_duration, skill["标准护盾持续时间"], skill["护盾时间衰减系数"]) * skill["标准护盾持续时间"] / shield_duration
    add("护盾人数系数", shield_people, "D58")
    add("护盾最终持续时间", shield_duration, "C62")
    add("护盾时间衰减系数", shield_time, "D64")
    add("最终护盾系数", skill["护盾系数"] * q / skill["护盾人数"] * shield_people * shield_time * base, "C65")

    for prefix, cell_value, cell_people, cell_result in [
        ("眩晕", "控制系数_眩晕", "控制人数_眩晕", "C74"),
        ("冰冻", "控制系数_冰冻", "控制人数_冰冻", "C83"),
        ("沉默", "控制系数_沉默", "控制人数_沉默", "C92"),
    ]:
        std_seconds = skill[f"标准控制秒数_{prefix}"]
        people = skill[cell_people]
        std_people = skill[f"标准控制人数_{prefix}"]
        decay = skill[f"控制人数衰减系数_{prefix}"]
        people_coeff = pow_ratio(people, std_people, decay) * std_people / people
        add(f"{prefix}控制人数系数", people_coeff, {"眩晕": "D73", "冰冻": "D82", "沉默": "D91"}[prefix])
        add(f"{prefix}实际控制秒数", base * std_seconds * skill[cell_value] * people_coeff * q, cell_result)

    for weight, val in split_basic:
        add(f"跟踪拆分_普攻必杀_{weight}", mround(skill["跟踪拆分基准"] * val, 5), None, "展示字符串拆分")
    for weight, val in split_skill:
        add(f"跟踪拆分_技能_{weight}", mround(val * skill["跟踪拆分基准"], 5), None, "展示字符串拆分")
    return outputs


def row_sign(row_index):
    # 旧表模板行 16-18 用减法降低威胁，后续空白扩展行用加法。
    return -1 if row_index in (16, 17, 18) else 1


def compute_boss_row(row, row_index, boss, tier, block, area, pierce):
    stage = row.get("技能档位")
    if not stage:
        return None
    base, default_cd = tier[stage]
    sign = row_sign(row_index)

    block_coeff = block[row.get("释放屏蔽")]
    area_coeff = area[row.get("范围类型")]
    radius = row.get("范围半径/宽度") or 0
    std_radius = row.get("标准范围") or 0
    if radius > 0:
        range_coeff = max(0.85, 1 + sign * (radius - std_radius) * boss["范围大小每单位系数"])
    else:
        range_coeff = 1
    pierce_coeff = pierce[row.get("穿透类型")]
    bullet_coeff = 1 + sign * max((row.get("弹体数") or 1) - 1, 0) * boss["额外弹体系数"]
    wave_coeff = 1 + sign * max((row.get("波次数") or 1) - 1, 0) * boss["额外波次系数"]
    duration = row.get("持续时间(s)") or 0
    interval = row.get("命中间隔(s)") or 0
    if duration > 0:
        duration_coeff = 1 + max(duration / max(interval, 0.1) - 1, 0) * boss["持续命中系数"]
    else:
        duration_coeff = 1
    control_time = row.get("控制时长(s)") or 0
    control_coeff = 1 + sign * max(control_time, 0) * boss["控制每秒系数"]
    mech = row.get("机制修正") if row.get("机制修正") not in (None, "") else 1

    param_coeff = block_coeff * area_coeff * range_coeff * pierce_coeff * bullet_coeff * wave_coeff * duration_coeff * control_coeff * mech
    final_coeff = base * param_coeff
    hp_pct = final_coeff / boss["最终TTK(s)"] if boss["最终TTK(s)"] else ""
    per_sec = final_coeff / default_cd if default_cd else ""
    reverse_base = base / param_coeff if param_coeff else ""
    out = dict(row)
    out.update({
        "行号": row_index,
        "基准倍率": base,
        "CD(s)": default_cd,
        "屏蔽系数": block_coeff,
        "范围类型系数": area_coeff,
        "范围大小系数": range_coeff,
        "穿透系数": pierce_coeff,
        "弹体系数": bullet_coeff,
        "波次系数": wave_coeff,
        "持续系数": duration_coeff,
        "控制系数": control_coeff,
        "参数综合系数": param_coeff,
        "最终技能系数": final_coeff,
        "100%命中掉血占比": hp_pct,
        "每秒威胁": per_sec,
        "反推基础倍率": reverse_base,
    })
    return out


def compute_boss_outputs(boss, tier, block, area, pierce, inputs):
    outputs = []
    for i, row in enumerate(inputs, start=16):
        out = compute_boss_row(row, i, boss, tier, block, area, pierce)
        if out:
            outputs.append(out)
    return outputs


def source_values(sheet_name, cells):
    wb = openpyxl.load_workbook(SRC, data_only=True, read_only=True)
    ws = wb[sheet_name]
    values = {cell: ws[cell].value for cell in cells}
    wb.close()
    return values


def parity_skill(outputs):
    cells = [source for _, _, source, _ in outputs if isinstance(source, str)]
    src = source_values("技能系数", cells)
    fails = []
    for name, value, cell, _ in outputs:
        if not isinstance(cell, str):
            continue
        if not almost(value, src[cell]):
            fails.append((name, cell, value, src[cell]))
    return len(cells), fails


def parity_boss(outputs):
    cells = []
    for row in outputs:
        r = row["行号"]
        for col in ("D", "E", "G", "I", "L", "N", "P", "R", "U", "W", "Y", "Z"):
            cells.append(f"{col}{r}")
    src = source_values("Boss技能系数", cells)
    key_map = {
        "D": "基准倍率", "E": "CD(s)", "G": "屏蔽系数", "I": "范围类型系数",
        "L": "范围大小系数", "N": "穿透系数", "P": "弹体系数", "R": "波次系数",
        "U": "持续系数", "W": "控制系数", "Y": "参数综合系数", "Z": "最终技能系数",
    }
    fails = []
    checks = 0
    for row in outputs:
        r = row["行号"]
        for col, key in key_map.items():
            checks += 1
            cell = f"{col}{r}"
            if not almost(row[key], src[cell]):
                fails.append((f"{row.get('技能名')}:{key}", cell, row[key], src[cell]))
    return checks, fails


def write_out(skill_outputs, boss_outputs):
    wb = openpyxl.Workbook()
    bold = Font(bold=True)
    hf = PatternFill("solid", fgColor="D9E1F2")
    ff = PatternFill("solid", fgColor="FFF2CC")

    ws = wb.active
    ws.title = "技能系数"
    ws.append(["【自动生成·勿手改】角色技能效果定价器关键输出"])
    ws.append(["效果/中间项", "生成值", "源表单元格", "说明"])
    for c in range(1, 5):
        ws.cell(2, c).font = bold
        ws.cell(2, c).fill = hf
    for name, value, cell, note in skill_outputs:
        ws.append([name, value, cell or "", note])
    for col, width in zip("ABCD", [28, 16, 14, 28]):
        ws.column_dimensions[col].width = width

    ws = wb.create_sheet("Boss技能系数")
    headers = [
        "Boss名", "技能名", "技能档位", "基准倍率", "CD(s)", "释放屏蔽", "屏蔽系数",
        "范围类型", "范围类型系数", "范围半径/宽度", "标准范围", "范围大小系数",
        "穿透类型", "穿透系数", "弹体数", "弹体系数", "波次数", "波次系数",
        "持续时间(s)", "命中间隔(s)", "持续系数", "控制时长(s)", "控制系数",
        "机制修正", "参数综合系数", "最终技能系数", "100%命中掉血占比", "每秒威胁", "反推基础倍率"
    ]
    ws.append(["【自动生成·勿手改】Boss技能系数：黄色列为v2新增诊断指标"])
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        ws.cell(2, c).font = bold
        ws.cell(2, c).fill = hf
    for c in range(27, 30):
        ws.cell(2, c).fill = ff
    for row in boss_outputs:
        ws.append([row.get(h, "") for h in headers])
    for c in range(1, len(headers) + 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width = 14
    wb.save(OUT)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--check", action="store_true", help="兼容旧参数；默认即只验证/写报告")
    ap.add_argument("--export-xlsx", action="store_true", help="临时导出调试用 xlsx")
    args = ap.parse_args()

    skill, qualities, split_basic, split_skill, _atk_hp, boss, tier, block, area, pierce, inputs = read_params()
    skill_outputs = compute_skill_outputs(skill, qualities, split_basic, split_skill)
    boss_outputs = compute_boss_outputs(boss, tier, block, area, pierce, inputs)
    skill_checks, skill_fails = parity_skill(skill_outputs)
    boss_checks, boss_fails = parity_boss(boss_outputs)

    if args.export_xlsx:
        write_out(skill_outputs, boss_outputs)

    md = []

    def o(text=""):
        print(text)
        md.append(text)

    o("# P3c · 技能系数 / Boss技能系数 生成与对照")
    o()
    o("> 将角色技能效果定价器与 Boss 技能威胁系数纳入 v2。源表只读，参数进入 `03_沉淀MD/参数基线/战斗参数基线.md`。")
    o()
    o("## 结论")
    o()
    o(f"- 角色技能系数关键项对照: {skill_checks - len(skill_fails)}/{skill_checks} 通过。")
    o(f"- Boss技能系数公式项对照: {boss_checks - len(boss_fails)}/{boss_checks} 通过，覆盖 {len(boss_outputs)} 条实际技能输入。")
    o("- 技能系数表当前定位为效果定价器，不直接等于 runtime 技能配置；Boss技能系数用于 Boss 技能威胁预算和 TTK 压力拆分。")
    o("- 定位修正：本生成器只负责结构化校验；`技能系数` 原工具页已由 `镜像_战斗知识工具页.xlsx` / `P0_战斗知识工具页_1比1镜像.md` 保留，`Boss技能系数` 不进 P0 镜像。")
    o()
    if skill_fails or boss_fails:
        o("## 不一致")
        o()
        o("| 项 | 单元格 | 生成 | 源 |")
        o("|---|---|---:|---:|")
        for name, cell, value, src in (skill_fails + boss_fails)[:40]:
            o(f"| {name} | {cell} | {value} | {src} |")
        o()
    else:
        o("✅ **逐项复现通过**：关键公式输出与旧表缓存值一致。")
        o()
    o("## 关键口径")
    o()
    o("- Vgame 是 5 人小队养成驱动的跑酷战斗射击，技能系数必须服务于自动普攻/自动技能/玩家择机必杀的节奏，而不是传统站桩卡牌单回合倍率。")
    o("- 知识类/工具类页签遵循“先 1:1 镜像，再结构化校验”的迁移规则，保留原表上下文优先于拆成生成结果。")
    o("- 控制、护盾、治疗、DOT、召唤等机制不能只看倍率，后续角色强度验证仍要按 Boss/小怪/无限挑战/PVPVE 等场景拆开。")
    o("- Boss 场景默认控制价值不应直接抬高 Boss 伤害结论；当前 Boss 技能表里的控制系数只作为技能威胁预算输入。")
    o()
    o("## 产物")
    o()
    o(f"- 报告: `{REPORT}`")
    if args.export_xlsx:
        o(f"- 临时导出表: `{OUT}`")

    os.makedirs(os.path.dirname(REPORT), exist_ok=True)
    with open(REPORT, "w", encoding="utf-8") as f:
        f.write("\n".join(md))
    print(f"\n[报告] {REPORT}")
    if args.export_xlsx:
        print(f"[生成表] {OUT}")


if __name__ == "__main__":
    main()
