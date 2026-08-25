# -*- coding: utf-8 -*-
"""
gen_monster_model.py — 怪物标准模型 大表生成器 (P3 + P3b 完整链路)

读取战斗参数基线 MD（怪物模型参数 / 战斗模型参数）→
按曲率插值 + 完整派生链 → 生成 N 级大表，
**直到游戏真正使用的最终输出**(单位秒最终属性 AL:AP, 各怪种生命 AR:AZ)，
并与现有 战斗框架/怪物标准模型 逐列逐格对照。

链路 (源公式逆向, 已核对):
  插值列(几何, 攻击%/防御%线性): 裸体/额外/百分比/各系数倍率/各乘区/各修正系数
  算后X = 裸体X×(1+X%) + 额外X
  无必杀增伤倍率 X = (1+暴击率×(暴伤-1))×普攻乘区×增伤乘区×(普攻系数+技能系数)×战斗支持
  有必杀增伤倍率 Y = (1+暴击率×(暴伤-1))×(普攻+技能+必杀系数)×普攻乘区×必杀乘区×增伤乘区×战斗支持
  最终生命 Z = Y×算后生命 ; 最终攻击 = 算后攻击 ; 最终防御 = 算后防御
  减伤率 = 防御/(防御+K) ; eHP = Z/(1-减伤) ; 无必杀TTK = eHP/攻击/X ; TTK = eHP/攻击/Y
  单位秒最终生命(BOSS侧) AL = round(Z/TTK×人数系数)
  单位秒最终生命(小怪侧) AO = round(Z/无必杀TTK×人数系数×小怪生命修正)
  单位秒最终攻击 AM = round(算后攻击×攻击系数修正) ; AN=round(算后防御) ; AP=round(AM×小怪攻击修正)
  小怪各类型生命 = ROUNDUP(类型倍率×AO) ; BOSS各类型生命 = round(AL×类型倍率×BOSS生命修正)

只读源表。生成表为独立新文件。
用法: python gen_monster_model.py [--check]
"""
import sys, io, os, argparse, math
import warnings
warnings.filterwarnings('ignore')
import openpyxl
from openpyxl.styles import Font, PatternFill

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

HERE = os.path.dirname(os.path.abspath(__file__))
COMBAT_DIR = os.path.dirname(HERE)  # 战斗/
V2_DIR = os.path.dirname(COMBAT_DIR)  # 框架v2/
NUM_DIR = os.path.dirname(V2_DIR)  # 数值文档/数值文档/
sys.path.insert(0, V2_DIR)
from _shared.md_parameter_reader import load_parameter_workbook, parameter_md_path

COMBAT_PARAM = parameter_md_path('战斗')
OUT     = os.path.join(COMBAT_DIR, '生成_怪物标准模型.xlsx')
REPORT  = os.path.join(COMBAT_DIR, '_reports', 'P3_怪物标准模型_生成与对照.md')
SRC     = os.path.join(NUM_DIR, 'Vgame基本战斗框架.xlsx')
SRC_SHEET = '怪物标准模型'
DATA_ROW0 = 12   # 源表 L1 所在行

def eround(x):   # Excel ROUND 半进位(正数)
    return int(math.floor(x + 0.5))
def eroundup(x): # Excel ROUNDUP = ceil
    return int(math.ceil(round(x, 6)))

def read_params():
    # 怪物模型参数 / 战斗模型参数从战斗参数 MD 读
    wb_combat = load_parameter_workbook(COMBAT_PARAM, data_only=True)
    ws = wb_combat['怪物模型参数']
    M = int(ws.cell(2, 2).value)
    attrs = []; types = []; mode = None
    for r in range(5, ws.max_row + 1):
        a = ws.cell(r, 1).value
        if a == '属性': continue
        if a == '类型名': mode = 'types'; continue
        if a is None: continue
        if mode == 'types':
            types.append((a, ws.cell(r,2).value, ws.cell(r,3).value))  # 名, 倍率, 侧
        elif isinstance(ws.cell(r,2).value, (int, float)) and ws.cell(r,4).value is not None:
            attrs.append((a, int(ws.cell(r,2).value), ws.cell(r,3).value,
                          ws.cell(r,4).value, ws.cell(r,5).value, ws.cell(r,6).value))
    mp = wb_combat['战斗模型参数']
    K = None
    for r in range(2, mp.max_row+1):
        if mp.cell(r,1).value == '防御常数K': K = float(mp.cell(r,2).value)
    wb_combat.close()
    return M, attrs, types, K

def interp(lo, hi, cur, kind, N, M):
    if kind == '线性' or cur is None or abs(cur - 1.0) < 1e-12:
        f = (N - 1) / (M - 1)
    else:
        f = (cur ** (N - 1) - 1) / (cur ** (M - 1) - 1)
    return lo + (hi - lo) * f

def generate(M, attrs, types, K):
    s = {}
    for (nm, col, lo, hi, cur, kind) in attrs:
        s[nm] = [interp(lo, hi, cur, kind, N, M) for N in range(1, M + 1)]
    for base, pct, extra, out in [('裸体生命','生命%','额外生命','算后生命'),
                                  ('裸体攻击','攻击%','额外攻击','算后攻击'),
                                  ('裸体防御','防御%','额外防御','算后防御')]:
        s[out] = [s[base][i]*(1+s[pct][i]) + s[extra][i] for i in range(M)]
    # 完整派生链
    chain = {k: [] for k in ['X','Y','Z','AA','AB','AC','AD','AE','AF','AL','AM','AN','AO','AP']}
    typ = {t[0]: [] for t in types}
    for i in range(M):
        O=s['普攻系数倍率'][i]; P=s['技能系数倍率'][i]; Q=s['必杀系数倍率'][i]; R=s['战斗支持倍率'][i]
        S=s['暴击率'][i]; T=s['暴击伤害'][i]; U=s['普攻乘区'][i]; V=s['必杀乘区'][i]; W=s['增伤乘区'][i]
        L=s['算后生命'][i]; Mt=s['算后攻击'][i]; Nn=s['算后防御'][i]
        AG=s['人数系数'][i]; AH=s['攻击系数修正'][i]; AI=s['BOSS生命系数修正'][i]; AJ=s['小怪生命系数修正'][i]; AK=s['小怪攻击系数修正'][i]
        base = 1 + S*(T-1)
        X = base*U*W*(O+P)*R
        Y = base*(O+P+Q)*U*V*W*R
        Z = Y*L; AA = Mt; AB = Nn
        AC = AB/(AB+K); AD = Z/(1-AC); AE = AD/AA/X; AF = AD/AA/Y
        AL = eround(Z/AF*AG); AM = eround(AA*AH); AN = eround(AB)
        AO = eround(Z/AE*AG*AJ); AP = eround(AM*AK)
        for k,v in [('X',X),('Y',Y),('Z',Z),('AA',AA),('AB',AB),('AC',AC),('AD',AD),
                    ('AE',AE),('AF',AF),('AL',AL),('AM',AM),('AN',AN),('AO',AO),('AP',AP)]:
            chain[k].append(v)
        for (name, mult, side) in types:
            typ[name].append(eroundup(mult*AO) if side=='小怪' else eround(AL*mult*AI))
    return s, chain, typ

# 源列映射
INTER_COLS = None  # 用 attrs 自带源列
CHAIN_COL = {'X':24,'Y':25,'Z':26,'AA':27,'AB':28,'AC':29,'AD':30,'AE':31,'AF':32,
             'AL':38,'AM':39,'AN':40,'AO':41,'AP':42}
TYPE_COL = {'普通怪':44,'普通盾怪':45,'2级普通怪':46,'站桩怪':47,'炮灰怪':48,'精英怪':49,'精英盾怪':50,'BOSS':51,'装备副本BOSS':52}

def parity(M, attrs, s, chain, typ):
    wb = openpyxl.load_workbook(SRC, data_only=True); ws = wb[SRC_SHEET]
    results = []
    def check(nm, col, seq, is_int):
        ma = mr = 0.0
        for N in range(1, M+1):
            a = ws.cell(DATA_ROW0+N-1, col).value
            if a is None: continue
            e = abs(seq[N-1] - a); ma = max(ma, e); mr = max(mr, e/max(abs(a),1e-12))
        ok = (ma < (0.5 if is_int else 1e-9)) or (not is_int and mr < 1e-6)
        results.append((nm, col, ma, ok))
    for (nm, col, *_ ) in attrs: check(nm, col, s[nm], False)
    for nm,col in [('算后生命',12),('算后攻击',13),('算后防御',14)]: check(nm, col, s[nm], False)
    for k,col in CHAIN_COL.items():
        is_int = k in ('AL','AM','AN','AO','AP')
        check(k, col, chain[k], is_int)
    for name, seq in typ.items():
        check(name, TYPE_COL[name], seq, True)
    wb.close()
    return results

def write_out(M, attrs, s, chain, typ):
    interp_order = [nm for (nm,*_ ) in attrs] + ['算后生命','算后攻击','算后防御']
    chain_order = ['X无必杀增伤','Y有必杀增伤','Z最终生命','最终攻击','最终防御','减伤率','eHP','无必杀TTK','TTK',
                   'AL_BOSS单位秒生命','AM最终攻击','AN最终防御','AO_小怪单位秒生命','AP小怪攻击']
    chain_keys = ['X','Y','Z','AA','AB','AC','AD','AE','AF','AL','AM','AN','AO','AP']
    type_order = list(typ.keys())
    wb = openpyxl.Workbook(); bold=Font(bold=True)
    hf=PatternFill('solid',fgColor='D9E1F2'); ff=PatternFill('solid',fgColor='FFF2CC')
    ws = wb.active; ws.title = '怪物标准模型'
    ws.append(['【自动生成·勿手改】完整链路至最终输出; 黄列=游戏使用的最终值(单位秒最终属性/各怪种生命)'])
    ws.append([])
    header = ['等级'] + interp_order + chain_order + ['['+t+']' for t in type_order]
    ws.append(header)
    nfix = 1 + len(interp_order)
    for c in range(1, len(header)+1):
        ws.cell(3,c).font=bold; ws.cell(3,c).fill=hf
    # 标黄最终列(chain 的 AL.. 与 类型列)
    final_start = 1 + len(interp_order) + 9  # 等级+interp+前9条chain
    for c in range(final_start+1, len(header)+1):
        ws.cell(3,c).fill = ff
    for N in range(1, M+1):
        row = [N] + [round(s[nm][N-1],6) for nm in interp_order]
        row += [round(chain[k][N-1],6) for k in chain_keys]
        row += [typ[t][N-1] for t in type_order]
        ws.append(row)
    wb.save(OUT)

def main():
    ap = argparse.ArgumentParser(); ap.add_argument('--check', action='store_true')
    args = ap.parse_args()
    M, attrs, types, K = read_params()
    s, chain, typ = generate(M, attrs, types, K)
    results = parity(M, attrs, s, chain, typ)
    if not args.check:
        write_out(M, attrs, s, chain, typ)

    md=[]
    def o(t=''):
        print(t); md.append(t)
    o('# P3 · 怪物标准模型 完整链路 生成与对照'); o()
    o('> 参数→%d级插值→完整派生链→**最终输出(单位秒最终属性+各怪种生命,boss/小怪分开)**。源表只读。' % M); o()
    npass = sum(1 for *_, ok in results if ok)
    o('- 等级数 %d，对照列 %d，**复现通过 %d / %d**' % (M, len(results), npass, len(results))); o()
    fails = [r for r in results if not r[3]]
    o('| 列 | 源列 | 最大误差 | 复现 |'); o('|---|---:|---:|:--:|')
    # 只详列最终输出 + 失败项
    show = [r for r in results if (r[0] in CHAIN_COL or r[0] in TYPE_COL) ]
    from openpyxl.utils import get_column_letter
    for (nm, col, ma, ok) in show:
        o('| %s | %s | %.2e | %s |' % (nm, get_column_letter(col), ma, '✅' if ok else '🔴'))
    o()
    if not fails:
        o('✅ **全链路逐格复现**，含游戏使用的最终输出(AL:AP 单位秒最终属性 + AR:AZ 各怪种生命, boss/小怪分开)。')
    else:
        o('🔴 未通过：');
        for (nm,col,ma,ok) in fails: o('- %s(%s) 最大误差 %.2e' % (nm, get_column_letter(col), ma))
    os.makedirs(os.path.dirname(REPORT), exist_ok=True)
    with open(REPORT, 'w', encoding='utf-8') as f: f.write('\n'.join(md))
    print('\n[报告] %s' % REPORT)
    if not args.check: print('[生成表] %s' % OUT)

if __name__ == '__main__':
    main()
