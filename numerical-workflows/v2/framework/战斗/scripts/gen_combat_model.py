# -*- coding: utf-8 -*-
"""
gen_combat_model.py — 标准战斗模型 生存骨架 生成器 (框架重搭 P2)

读取共享、养成、战斗参数基线 MD（标准人 / 投放系数 / 升级曲线 / 战斗模型参数）→
  生成 前期 + 中期 标准人基础属性 → 减伤率 / eHP / ttk
落地 ruling②：投放系数以「养成属性规划」版为准 (升级4.8/升阶0/星盘2.3)。

公式 (已逆向验证):
  减伤率 = round(防御 / (防御 + K), 减伤率小数位)   K=24000
  eHP   = 总生命 / (1 - 减伤率)
  ttk   = eHP / 总攻击                              (怪物DPS系数=1)
  某档基础属性 = 标准人 × 该档投放系数总和

只读现有源表。默认只写校验报告；如需临时导出 xlsx，可加 --export-xlsx。

用法:
    python gen_combat_model.py                # 公式验证 + ruling②差异报告
    python gen_combat_model.py --export-xlsx  # 额外导出调试 xlsx
"""
import sys, io, os, argparse
import warnings
warnings.filterwarnings('ignore')
import openpyxl
from openpyxl.styles import Font, PatternFill

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

HERE = os.path.dirname(os.path.abspath(__file__))
COMBAT_DIR = os.path.dirname(HERE)  # 战斗/
V2_DIR = os.path.dirname(COMBAT_DIR)  # 框架v2/
NUM_DIR = os.path.dirname(V2_DIR)  # 数值文档/数值文档/
sys.path.insert(0, V2_DIR)
from _shared.md_parameter_reader import load_parameter_workbook, parameter_md_path

SHARED_PARAM = parameter_md_path('共享')
GROWTH_PARAM = parameter_md_path('养成')
COMBAT_PARAM = parameter_md_path('战斗')
OUT     = os.path.join(COMBAT_DIR, '生成_战斗模型.xlsx')
REPORT  = os.path.join(COMBAT_DIR, '_reports', 'P2_战斗模型_生成与验证.md')

# 旧表(标准模型)已知基准值，用于验证公式可逐格复现 (来自 战斗框架/标准模型)
LEGACY = {
    '前期': dict(coefsum=1.0,  hp=5000,   atk=1500,  dfs=1000,  reduce=0.04, ehp=5208,   ttk=3.472),
    '中期': dict(coefsum=27.9, hp=139500, atk=41850, dfs=27900, reduce=0.54, ehp=303261, ttk=7.246),
}
# 旧表使用的(过期)投放系数, 仅供公式验证
STALE_COEF = dict(初始=1, 升级=4.2, 升阶=1.5, 星盘=1.2, 专武=4, 旅券=6, 映射装备=10)

def almost(a, b, tol=0.5):
    try: return abs(float(a) - float(b)) <= tol
    except (TypeError, ValueError): return a == b

def read_params():
    # 标准人 / 投放系数 从共享参数读
    wb_shared = load_parameter_workbook(SHARED_PARAM, data_only=True)
    sp = wb_shared['标准人']
    std = {sp.cell(r,1).value: sp.cell(r,2).value for r in range(2,5)}
    fp = wb_shared['投放系数']
    coefs = {}
    for r in range(2, fp.max_row+1):
        name = fp.cell(r,1).value
        if not name: continue
        coefs[name] = (fp.cell(r,2).value, fp.cell(r,3).value)
    wb_shared.close()
    # 升级曲线参数 → 从养成参数 MD 读
    wb_growth = load_parameter_workbook(GROWTH_PARAM, data_only=True)
    cp = wb_growth['升级曲线参数']
    curve = {cp.cell(r,1).value: cp.cell(r,2).value for r in range(2,cp.max_row+1) if cp.cell(r,1).value}
    wb_growth.close()
    # 战斗模型参数 → 从战斗参数 MD 读
    wb_combat = load_parameter_workbook(COMBAT_PARAM, data_only=True)
    mp = wb_combat['战斗模型参数']
    model = {mp.cell(r,1).value: mp.cell(r,2).value for r in range(2,mp.max_row+1) if mp.cell(r,1).value}
    wb_combat.close()
    return std, curve, coefs, model

def coef_sum(coefs, curve):
    """当前参数(ruling②)下的投放系数总和。升级=派生=等级上限×单级成长。"""
    lvl = int(curve['等级上限']) * float(curve['单级成长'])
    total = 0.0
    for name,(val,typ) in coefs.items():
        if typ == '锚': total += float(val)
        elif typ == '派生' and '升级' in name: total += lvl
        elif typ == '旋钮': total += float(val)
        # 旋钮(%) 不计入基础属性总和
    return round(total, 6), lvl

def phase(std, coefsum, K, ndigit):
    hp  = round(std['生命'] * coefsum)
    atk = round(std['攻击'] * coefsum)
    dfs = round(std['防御'] * coefsum)
    reduce = round(dfs / (dfs + K), int(ndigit))
    ehp = round(hp / (1 - reduce))
    ttk = round(ehp / atk, 3)
    return dict(coefsum=coefsum, hp=hp, atk=atk, dfs=dfs, reduce=reduce, ehp=ehp, ttk=ttk)

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--check', action='store_true', help='兼容旧参数；默认即只验证/写报告')
    ap.add_argument('--export-xlsx', action='store_true', help='临时导出调试用 xlsx')
    args = ap.parse_args()
    std, curve, coefs, model = read_params()
    K = float(model['防御常数K']); ndigit = int(model['减伤率小数位'])

    total, lvl = coef_sum(coefs, curve)

    # ruling② 下的三档(本期只做前期/中期)
    v2 = {'前期': phase(std, 1.0, K, ndigit), '中期': phase(std, total, K, ndigit)}

    # 公式验证: 用旧(过期)系数应能复现旧表
    stale_total = sum(STALE_COEF.values())  # 27.9
    chk = {'前期': phase(std, 1.0, K, ndigit), '中期': phase(std, stale_total, K, ndigit)}

    md=[]
    def o(s=''):
        print(s); md.append(s)
    o('# P2 · 标准战斗模型(生存骨架) 生成与验证'); o()
    o('> 前期+中期基础属性 → 减伤率/eHP/ttk。落地 ruling②(投放系数用养成属性规划版)。源表只读。'); o()
    o('- K(防御常数)=%g, 减伤率小数位=%d, 怪物DPS系数=%g' % (K, ndigit, model.get('怪物DPS系数',1)))
    o('- ruling② 投放系数总和(中期) = %.1f (升级派生=%.2f)' % (total, lvl)); o()

    # 1) 公式验证
    o('## ① 公式验证：用旧系数(升级4.2/升阶1.5/星盘1.2, 总27.9)能否复现旧表'); o()
    o('| 档 | 项 | 计算值 | 旧表值 | 一致 |'); o('|---|---|---:|---:|:--:|')
    ok = True
    for ph in ['前期','中期']:
        L = LEGACY[ph]; C = chk[ph]
        for key,label in [('hp','生命'),('atk','攻击'),('dfs','防御'),('reduce','减伤率'),('ehp','eHP'),('ttk','ttk')]:
            tol = 0.01 if key in ('reduce','ttk') else 1.0
            good = almost(C[key], L[key], tol)
            ok = ok and good
            o('| %s | %s | %s | %s | %s |' % (ph,label,C[key],L[key],'✅' if good else '🔴'))
    o()
    o('**公式验证：%s**' % ('✅ 通过 — 模型公式可逐格复现旧表' if ok else '🔴 有偏差，需复查'))
    o()

    # 2) ruling② 输出 + 差异
    o('## ② ruling② 修正后的 v2 值（与旧表差异 = 裁决带来的修正）'); o()
    o('| 档 | 生命 | 攻击 | 防御 | 减伤率 | eHP | ttk |'); o('|---|---:|---:|---:|---:|---:|---:|')
    for ph in ['前期','中期']:
        v = v2[ph]
        o('| %s | %d | %d | %d | %.2f | %d | %.3f |' % (ph,v['hp'],v['atk'],v['dfs'],v['reduce'],v['ehp'],v['ttk']))
    o()
    o('差异(中期, v2 − 旧表)：')
    d = v2['中期']; L = LEGACY['中期']
    o('- 投放系数总和 %.1f → %.1f' % (L['coefsum'], d['coefsum']))
    o('- 生命 %d → %d (%+d)　攻击 %d → %d (%+d)　防御 %d → %d (%+d)' %
      (L['hp'],d['hp'],d['hp']-L['hp'], L['atk'],d['atk'],d['atk']-L['atk'], L['dfs'],d['dfs'],d['dfs']-L['dfs']))
    o('- ttk %.3f → %.3f (生存/击杀节奏基本不变)' % (L['ttk'], d['ttk']))
    o()

    if args.export_xlsx:
        wb = openpyxl.Workbook(); bold=Font(bold=True); hf=PatternFill('solid',fgColor='D9E1F2')
        ws=wb.active; ws.title='标准模型-生存骨架'
        ws.append(['【自动生成·勿手改】源:03_沉淀MD/参数基线（共享+养成+战斗） 落地ruling②(投放系数=养成属性规划版)'])
        ws.append(['公式: 减伤率=round(防御/(防御+%g),%d); eHP=生命/(1-减伤率); ttk=eHP/攻击' % (K,ndigit)])
        ws.append([])
        ws.append(['档','投放系数总和','生命','攻击','防御','减伤率','eHP','ttk(秒)'])
        for c in range(1,9): ws.cell(4,c).font=bold; ws.cell(4,c).fill=hf
        for ph in ['前期','中期']:
            v=v2[ph]; ws.append([ph,v['coefsum'],v['hp'],v['atk'],v['dfs'],v['reduce'],v['ehp'],v['ttk']])
        for w,col in zip([8,14,10,10,10,8,12,10],'ABCDEFGH'): ws.column_dimensions[col].width=w
        wb.save(OUT)

    os.makedirs(os.path.dirname(REPORT), exist_ok=True)
    with open(REPORT,'w',encoding='utf-8') as f: f.write('\n'.join(md))
    print('\n[报告] %s' % REPORT)
    if args.export_xlsx:
        print('[生成表] %s' % OUT)

if __name__ == '__main__':
    main()
