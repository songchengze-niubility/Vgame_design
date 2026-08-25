# -*- coding: utf-8 -*-
"""
sync_skill_coefficient_params.py — 从旧战斗框架同步技能系数参数页

只用于把 Vgame基本战斗框架.xlsx 中「技能系数 / Boss技能系数」的输入型内容
抽取到战斗参数基线 MD。旧源表只读。

默认只预演，显式确认后才写入 MD：
    python sync_skill_coefficient_params.py --write
"""
import argparse
import io
import os
import sys
import warnings

warnings.filterwarnings("ignore")

import openpyxl
from openpyxl.styles import Font, PatternFill

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

HERE = os.path.dirname(os.path.abspath(__file__))
COMBAT_DIR = os.path.dirname(HERE)
V2_DIR = os.path.dirname(COMBAT_DIR)
NUM_DIR = os.path.dirname(V2_DIR)
sys.path.insert(0, V2_DIR)
from _shared.md_parameter_reader import (
    load_parameter_workbook,
    parameter_md_path,
    save_parameter_workbook,
)

PARAM = parameter_md_path("战斗")
SRC = os.path.join(NUM_DIR, "Vgame基本战斗框架.xlsx")


def style_header(ws, max_col):
    bold = Font(bold=True)
    fill = PatternFill("solid", fgColor="D9E1F2")
    for c in range(1, max_col + 1):
        ws.cell(1, c).font = bold
        ws.cell(1, c).fill = fill


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--write", action="store_true", help="确认用旧战斗框架重建 MD 中的三个技能参数页")
    args = parser.parse_args()
    wb_src = openpyxl.load_workbook(SRC, data_only=True)
    ws = wb_src["技能系数"]
    bs = wb_src["Boss技能系数"]

    wb = load_parameter_workbook(PARAM, data_only=False)
    for name in ["技能系数参数", "Boss技能系数参数", "Boss技能输入"]:
        if name in wb.sheetnames:
            del wb[name]

    yf = PatternFill("solid", fgColor="FFF2CC")

    skill = wb.create_sheet("技能系数参数")
    skill.append(["参数", "值", "说明", "分组", "键", "值", "说明"])
    style_header(skill, 7)
    basic = [
        ("技能系数", ws["C2"].value, "全局定价预算"),
        ("技能当前品质", ws["C3"].value, "R/SR/SSR"),
        ("单体系数", ws["C5"].value, ""), ("单体攻速", ws["C7"].value, ""),
        ("穿透子弹系数", ws["C11"].value, ""), ("穿透子弹攻速", ws["C13"].value, ""),
        ("跟踪子弹系数", ws["C17"].value, ""), ("跟踪子弹攻速", ws["C19"].value, ""),
        ("跟踪拆分基准", ws["J20"].value, "展示用百分比分拆"),
        ("AOE系数", ws["C24"].value, ""), ("AOE标准范围半径", ws["C26"].value, ""),
        ("AOE实际范围半径", ws["C27"].value, ""), ("AOE衰减系数", ws["C28"].value, ""), ("AOE攻速", ws["C29"].value, ""),
        ("DOT系数", ws["C33"].value, ""), ("DOT持续时间", ws["C35"].value, ""), ("DOT标准持续时间", ws["C36"].value, ""),
        ("DOT持续衰减系数", ws["C37"].value, ""), ("DOT标准最大层数", ws["C38"].value, ""), ("DOT最大层数", ws["C39"].value, ""),
        ("DOT最大层数衰减系数", ws["C40"].value, ""),
        ("回血系数", ws["C45"].value, "来自标准模型I75"), ("回血人数", ws["C47"].value, ""), ("标准回血人数", ws["C48"].value, ""),
        ("回血人数衰减系数", ws["C49"].value, ""),
        ("护盾系数", ws["C54"].value, "回血系数*1.25"), ("护盾人数", ws["C56"].value, ""), ("标准护盾人数", ws["C57"].value, ""),
        ("护盾人数衰减系数", ws["C58"].value, ""), ("护盾持续时间", ws["C59"].value, ""), ("护盾持续上限", ws["C60"].value, ""),
        ("护盾持续下限", ws["C61"].value, ""), ("标准护盾持续时间", ws["C63"].value, ""), ("护盾时间衰减系数", ws["C64"].value, ""),
        ("控制系数_眩晕", ws["C68"].value, ""), ("标准控制秒数_眩晕", ws["C70"].value, ""), ("控制人数_眩晕", ws["C71"].value, ""),
        ("标准控制人数_眩晕", ws["C72"].value, ""), ("控制人数衰减系数_眩晕", ws["C73"].value, ""),
        ("控制系数_冰冻", ws["C77"].value, ""), ("标准控制秒数_冰冻", ws["C79"].value, ""), ("控制人数_冰冻", ws["C80"].value, ""),
        ("标准控制人数_冰冻", ws["C81"].value, ""), ("控制人数衰减系数_冰冻", ws["C82"].value, ""),
        ("控制系数_沉默", ws["C86"].value, ""), ("标准控制秒数_沉默", ws["C88"].value, ""), ("控制人数_沉默", ws["C89"].value, ""),
        ("标准控制人数_沉默", ws["C90"].value, ""), ("控制人数衰减系数_沉默", ws["C91"].value, ""),
    ]
    lookup = []
    for r in range(4, 7):
        lookup.append(("品质系数", ws.cell(r, 14).value, ws.cell(r, 15).value, ""))
    for c in range(7, 11):
        lookup.append(("拆分系数_普攻必杀", ws.cell(16, c).value, ws.cell(16, c).value, ""))
    for c in range(7, 10):
        lookup.append(("拆分系数_技能", ws.cell(18, c).value, ws.cell(18, c).value, ""))
    for r in range(4, 7):
        lookup.append(("攻血比", ws.cell(r, 18).value, ws.cell(r, 19).value, ""))
    for i in range(max(len(basic), len(lookup))):
        row = i + 2
        if i < len(basic):
            skill.cell(row, 1).value = basic[i][0]
            skill.cell(row, 2).value = basic[i][1]
            skill.cell(row, 3).value = basic[i][2]
            skill.cell(row, 2).fill = yf
        if i < len(lookup):
            skill.cell(row, 4).value = lookup[i][0]
            skill.cell(row, 5).value = lookup[i][1]
            skill.cell(row, 6).value = lookup[i][2]
            skill.cell(row, 7).value = lookup[i][3]
    for col, width in zip("ABCDEFG", [24, 14, 34, 22, 16, 14, 20]):
        skill.column_dimensions[col].width = width

    bp = wb.create_sheet("Boss技能系数参数")
    bp.append(["参数", "值", "说明", "分组", "键", "值1", "值2", "说明"])
    style_header(bp, 8)
    basic = []
    for r in range(4, 14):
        if bs.cell(r, 1).value:
            basic.append((bs.cell(r, 1).value, bs.cell(r, 2).value, bs.cell(r, 3).value))
    lookup = []
    for r in range(4, 7):
        lookup.append(("技能档位", bs.cell(r, 4).value, bs.cell(r, 5).value, bs.cell(r, 6).value, bs.cell(r, 7).value))
    for r in range(4, 8):
        lookup.append(("释放屏蔽", bs.cell(r, 9).value, bs.cell(r, 10).value, "", bs.cell(r, 11).value))
    for r in range(4, 10):
        lookup.append(("范围类型", bs.cell(r, 13).value, bs.cell(r, 14).value, "", bs.cell(r, 15).value))
    for r in range(4, 7):
        lookup.append(("穿透类型", bs.cell(r, 17).value, bs.cell(r, 18).value, "", bs.cell(r, 19).value))
    for i in range(max(len(basic), len(lookup))):
        row = i + 2
        if i < len(basic):
            bp.cell(row, 1).value = basic[i][0]
            bp.cell(row, 2).value = basic[i][1]
            bp.cell(row, 3).value = basic[i][2]
            bp.cell(row, 2).fill = yf
        if i < len(lookup):
            bp.cell(row, 4).value = lookup[i][0]
            bp.cell(row, 5).value = lookup[i][1]
            bp.cell(row, 6).value = lookup[i][2]
            bp.cell(row, 7).value = lookup[i][3]
            bp.cell(row, 8).value = lookup[i][4]
    for col, width in zip("ABCDEFGH", [22, 12, 45, 16, 16, 12, 12, 38]):
        bp.column_dimensions[col].width = width

    inp = wb.create_sheet("Boss技能输入")
    headers = [bs.cell(15, c).value for c in range(1, 25)]
    inp.append(headers)
    style_header(inp, len(headers))
    for r in range(16, 61):
        inp.append([bs.cell(r, c).value for c in range(1, 25)])
    for r in range(2, inp.max_row + 1):
        for c in [1, 2, 3, 6, 8, 10, 11, 13, 15, 17, 19, 20, 22, 24]:
            inp.cell(r, c).fill = yf
    for c in range(1, len(headers) + 1):
        inp.column_dimensions[openpyxl.utils.get_column_letter(c)].width = 14

    if args.write:
        save_parameter_workbook(PARAM, wb)
    wb.close()
    wb_src.close()
    if args.write:
        print("[参数MD] 已同步到", PARAM)
    else:
        print("[DRY-RUN] 已完成同步预演，未修改 MD；确认后使用 --write")


if __name__ == "__main__":
    main()
