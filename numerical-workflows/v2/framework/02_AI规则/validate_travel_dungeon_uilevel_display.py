#!/usr/bin/env python3
"""只读校验旅券副本 UIlevel 奖励展示镜像。"""
from __future__ import annotations

import json
import os
from pathlib import Path

from openpyxl import load_workbook

ROOT=Path(os.environ.get("VGAME_NUMERICAL_RUNTIME_ROOT", str(Path(__file__).resolve().parents[1])))
CONFIG=Path(os.environ.get("VGAME_CONFIG_DATAS", ""))
OUT=Path(os.environ.get("TRAVEL_DUNGEON_FINAL_OUT",str(ROOT/"06_最终配置"/"旅券副本")))
GEN=ROOT/"05_验算核算"/"travel-dungeon-uilevel-display-F01.json"
REPORT=ROOT/"05_验算核算"/"travel-dungeon-uilevel-display-F01-validation.json"
TYPES={"LevelType_Strategy_First","LevelType_Strategy_Second","LevelType_Strategy_Three","LevelType_Strategy_Four"}

def main():
    generation=json.loads(GEN.read_text(encoding="utf-8")); checks=[]
    def check(name,ok,actual,expected): checks.append({"name":name,"result":"PASS" if ok else "FAIL","actual":actual,"expected":expected})
    path=OUT/"UIlevel.xlsx"; check("旅券UIlevel存在",path.exists(),path.exists(),True)
    wb=load_workbook(path,read_only=True,data_only=False); check("页签顺序",wb.sheetnames[:2]==["level","中文说明"],wb.sheetnames[:2],["level","中文说明"])
    ws=wb["level"]; headers=next(ws.iter_rows(min_row=1,max_row=1,values_only=True)); cols={field:headers.index(field) for field in ("LevelId","LevelType","LevelFPDrop","LevelPPDrop","Reward","FirstReward","FirstRewardDrop","FallRewardDrop")}
    rows=[row for row in ws.iter_rows(min_row=5,values_only=True) if row[cols["LevelId"]] is not None]
    check("旅券关卡数",len(rows)==40,len(rows),40)
    check("旅券关卡类型",all(row[cols["LevelType"]] in TYPES for row in rows),sorted({row[cols["LevelType"]] for row in rows}),sorted(TYPES))
    fuw=load_workbook(CONFIG/"level/UIlevel.xlsx",read_only=True,data_only=True); fus=fuw["level"]
    formal_headers=next(fus.iter_rows(min_row=1,max_row=1,values_only=True)); formal_cols={field:formal_headers.index(field) for field in ("LevelId","LevelFPDrop","LevelPPDrop")}
    formal_display_by_level={row[formal_cols["LevelId"]]:(row[formal_cols["LevelFPDrop"]],row[formal_cols["LevelPPDrop"]]) for row in fus.iter_rows(min_row=5,values_only=True) if row[formal_cols["LevelId"]] is not None}
    display_mismatches=[(row[cols["LevelId"]],(row[cols["LevelFPDrop"]],row[cols["LevelPPDrop"]]),formal_display_by_level.get(row[cols["LevelId"]])) for row in rows if (row[cols["LevelFPDrop"]],row[cols["LevelPPDrop"]])!=formal_display_by_level.get(row[cols["LevelId"]])]
    check("LevelFPDrop和LevelPPDrop保持正式源表原值",not display_mismatches,display_mismatches,[])
    fuw.close()
    preview_ids=set(); preview_quantity_errors=[]
    for row in rows:
        for field in ("Reward","FirstReward"):
            for part in str(row[cols[field]] or "").split(";"):
                if not part:
                    continue
                item_id,quantity=(int(x) for x in part.split(",")); preview_ids.add(item_id)
                if quantity < 0 and quantity != -1:
                    preview_quantity_errors.append((row[cols["LevelId"]],field,item_id,quantity))
    check("UI无数量展示统一为负一",not preview_quantity_errors,preview_quantity_errors,[])
    errors=sum(1 for row in ws.iter_rows(values_only=True) for value in row if isinstance(value,str) and value.startswith(("#REF!","#VALUE!","#DIV/0!","#NAME?","#N/A")))
    check("公式错误",errors==0,errors,0); wb.close()

    dw=load_workbook(CONFIG/"Drop/Drop_dungeon.xlsx",read_only=True,data_only=True); ds=dw["Drop"]
    dh=next(ds.iter_rows(min_row=1,max_row=1,values_only=True)); didx=dh.index("DropId")
    drop_ids={row[didx] for row in ds.iter_rows(min_row=5,values_only=True) if row[didx] is not None}; dw.close()
    real_drop_ids={row[cols["FirstRewardDrop"]] for row in rows}|{row[cols["FallRewardDrop"]] for row in rows}
    check("旅券真实Drop引用闭环",real_drop_ids<=drop_ids,sorted(real_drop_ids-drop_ids),[])
    iw=load_workbook(CONFIG/"items/Items.xlsx",read_only=True,data_only=True); isheet=iw["items"]
    ih=next(isheet.iter_rows(min_row=1,max_row=1,values_only=True)); item_col=ih.index("Id")
    formal_item_ids={row[item_col] for row in isheet.iter_rows(min_row=5,values_only=True) if isinstance(row[item_col],int)}; iw.close()
    check("旅券UI预览Items正式闭环",preview_ids<=formal_item_ids,sorted(preview_ids-formal_item_ids),[])
    check("正式Datas零写入",generation["formal_datas_unchanged"] is True,generation["formal_datas_unchanged"],True)
    result={"summary":{"checks":len(checks),"pass":sum(c["result"]=="PASS" for c in checks),"fail":sum(c["result"]=="FAIL" for c in checks)},"checks":checks}
    REPORT.write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding="utf-8")
    print(json.dumps(result["summary"],ensure_ascii=False))
    raise SystemExit(1 if result["summary"]["fail"] else 0)

if __name__=="__main__": main()
