"""Independent read-only semantic checks; never authors XLSX."""
import argparse,hashlib,importlib,json,re,sys
from collections import defaultdict
from fractions import Fraction
from pathlib import Path
from openpyxl import load_workbook
HERE=Path(__file__).resolve().parent;ROOT=HERE.parent.parent
parser=argparse.ArgumentParser();parser.add_argument('--published',action='store_true');args=parser.parse_args()
meta=json.loads((ROOT/'05_验算核算/mapping-equip-first-clear-CHG022.json' if args.published else HERE/'repair.json').read_text(encoding='utf-8'))
source=meta['source'];out=Path(meta['out'] if args.published else meta['staging']);backup=Path(meta['backup'])
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
def read(p,sheet=None):
 w=load_workbook(p,read_only=True,data_only=False);s=w[sheet] if sheet else w.worksheets[0]
 rs=[list(r) for r in s.values];w.close();h=rs[0] if rs[0][0]=='##var' else rs[3]
 return rs,{str(v).split('\n')[-1]:i for i,v in enumerate(h) if v}
def trim(r):
 r=list(r)
 while r and r[-1] is None:r.pop()
 return r
def mapping(rs,h,key):
 d=defaultdict(list)
 for r in rs[4:]:
  if len(r)>h[key] and r[h[key]] is not None:d[r[h[key]]].append(r)
 return d
checks=[]
def check(n,ok):checks.append({'name':n,'result':'PASS' if ok else 'FAIL'})

for f in meta['files'][:-1]:
 old,_=read(backup/f);new,_=read(out/f)
 for p in source['patches']:
  if p['file']==f:old[p['row']][p['col']]=p['new']
 if f=='SpecialEquipmentGenerate.xlsx':old.append(source['newGenerate']['values'])
 check(f+'仅指定业务单元格及新增4375行变化',[trim(r) for r in old]==[trim(r) for r in new])

dr,dh=read(out/'Drop_dungeon.xlsx');drops=mapping(dr,dh,'DropId')
ir,ih=read(out/'Items.xlsx');items=mapping(ir,ih,'Id')
gr,gh=read(out/'SpecialEquipmentGenerate.xlsx');gens=mapping(gr,gh,'Id')
ur,uh=read(out/'UIlevel.xlsx');uis=mapping(ur,uh,'LevelId')
check('第三层实际包固定发4375和25各1',len(drops[20000012])==2 and {r[dh['DropItem_id']] for r in drops[20000012]}=={25,4375} and all(r[dh['DropItemNumMin']]==r[dh['DropItemNumMax']]==1 and r[dh['DropType']]=='fixed_' and r[dh['DropId2']]==0 for r in drops[20000012]))
item=items[4375][0];gen=gens[4375][0]
check('4375唯一且确为实际开箱道具',len(items[4375])==len(gens[4375])==1 and item[ih['ItemSubType']]=='ItemSubType_SpecialEquip_Box' and item[ih['ParaList']]==4375)
check('4375外观蓝色而产物为金色',item[ih['ItemQuality']]=='Blue' and gen[gh['QualityList']]=='Orange,1000')
check('4375仅10级且固定一件',gen[gh['GradeList']]=='10,1000' and gen[gh['GenerateMin']]==gen[gh['GenerateMax']]==1)
check('4375文案精确且非占位',item[ih['Description']]=='开启后固定获得1件10级金色映射装备。' and '实际' in item[ih['##']] and 'UI_PREVIEW_ONLY' not in str(item))
check('4375只由第三层首通发放',{k for k,rs in drops.items() if any(r[dh['DropItem_id']]==4375 for r in rs)}=={20000012})
check('4375没有固定装备ID覆写',gen[gh['SpecialEquipmentId']] is None)
check('4375部位保留8部位',gen[gh['PartList']]=='1;2;3;4;5;6;7;8')
md=ROOT/'03_沉淀MD/经济/05_玩法_映射副本经济沉淀-2026-07-08.md'
colors={'绿':'Green','蓝':'Blue','紫':'Purple','金':'Orange','红':'Red'}
qualityOrder=list(colors.values());rules={}
for line in md.read_text(encoding='utf-8-sig').splitlines():
 m=re.match(r'^\| (\d+) \| 装备副本钥匙 x1 \| ([绿蓝紫金红])宝箱 x1',line)
 if m:rules[int(m[1])]=(colors[m[2]],line.split('|')[3].strip())
check('文档30层与4375例外完整',len(rules)==30 and '4375' in rules[3][1])
formal=Path('D:/Vgame/Config/GameConfig/Datas')
fir,fih=read(formal/'items/Items.xlsx');formalIds={r[fih['Id']] for r in fir[4:] if r[0]!='##'}
knownItems=set(items)|formalIds
pair=lambda text:[(int(p.split(',')[0]),int(p.split(',')[1])) for p in (text or '').split(';') if p]
dist=lambda text:{s.split(',')[0]:Fraction(s.split(',')[1]) for s in text.split(';') if s}
expectedTotals=defaultdict(Fraction)
for l,(q,archive) in rules.items():
 root=20000009+l;rs=drops[root];box=4375 if l==3 else 4201+(l-1)*5+qualityOrder.index(q)
 actual=[(r[dh['DropItem_id']],r[dh['DropItemNumMin']]) for r in rs]
 check(f'L{l}实际首通符合MD且无子池',(25,1) in actual and (box,1) in actual and all(r[dh['DropId2']]==0 and r[dh['DropItemNumMin']]==r[dh['DropItemNumMax']]==1 for r in rs))
 bi=items[box][0];bg=gens[bi[ih['ParaList']]][0]
 levels={int(x) for x,w in dist(bg[gh['GradeList']]).items() if w>0}
 qi=dist(bg[gh['QualityList']]);qs=sum(qi.values())
 for quality,w in qi.items():expectedTotals[quality]+=w/qs
 preview=items[4372+l][0]
 check(f'L{l}首通预览描述等级与开箱一致',{int(n) for n in re.findall(r'(\d+)级',preview[ih['Description']])}==levels)
 for variant in range(1,6):
  id=250000+l*10+variant;u=uis[id][0];fp=pair(u[uh['FirstReward']]);normal=pair(u[uh['Reward']])
  translated=[(box if itemid==4372+l else itemid,num) for itemid,num in fp]
  check(f'UI{id}首通预览完整对应实际必得项',sorted(translated)==sorted(actual))
  check(f'UI{id}首通常规联合展示引用有效',all(i in knownItems for i,n in fp+normal) and all(n==-1 for i,n in normal))
  check(f'UI{id}常规预览不含专用首通箱',4375 not in {i for i,n in normal})
check('首通开装品质期望更新为1.6/5.5/9.9/9.84/3.16',dict(expectedTotals)==dict(zip(qualityOrder,map(Fraction,['1.6','5.5','9.9','9.84','3.16']))))
for f in meta['files']:
 a=load_workbook(backup/f,data_only=False);b=load_workbook(out/f,data_only=False)
 formulas=lambda w:{(s.title,c.coordinate):c.value for s in w for row in s for c in row if c.data_type=='f'}
 check(f+'所有公式不变',formulas(a)==formulas(b));a.close();b.close()
 b=load_workbook(out/f,data_only=True);check(f+'无单元格错误',not any(c.data_type=='e' for s in b for row in s for c in row));b.close()
for p,h in source['formalHashes'].items():check('SVN未改 '+Path(p).name,sha(Path(p))==h)
for f,h in source['candidateHashes'].items():
 if f not in meta['files']:check('其他候选未改 '+f,sha(Path(meta['out'])/f)==h)
check('已确认MD未被并发修改',sha(md)==source['mdHash'])
sys.path.insert(0,str(HERE.parent));generator=importlib.import_module('generate_mapping_equip_full_package')
model={'layers':[{'layer':3,'first_archive':rules[3][1],'craft_tier':10,'grade_probs':[(5,Fraction(3,4)),(10,Fraction(1,4))]}]}
rows,_=generator.first_clear_rows(model)
check('生成器真实Drop使用4375',any(r[1]==20000012 and r[9]==4375 and r[10]==r[11]==1 for r in rows))
check('生成器实际箱身份正确',generator.first_special_box_item_row()[2:5]==[4375,'ItemType_Other','ItemSubType_SpecialEquip_Box'])
check('生成器开箱固定10级金装',generator.first_special_box_generate_row()[2:4]==['10,1000','Orange,1000'])
check('生成器五变体均匙1箱1',all(generator.first_preview_items(model['layers'][0],v)=='25,1;4375,1' for v in range(1,6)))
check('生成器包含专用开箱记录',any(r[0]==4375 for r in generator.generate_rows(model)))
full=[]
for l,(q,archive) in rules.items():
 normalGen=gens[4201+(l-1)*5][0]
 mats=[r[dh['DropItem_id']] for r in drops[20000009+l] if 1080<=r[dh['DropItem_id']]<=1151]
 tier=generator.CRAFT_TIERS[(mats[0]-1080)//9] if mats else 10
 gradeWeights=dist(normalGen[gh['GradeList']]);total=sum(gradeWeights.values())
 full.append({'layer':l,'first_archive':archive,'craft_tier':tier,'grade_probs':[(int(k),v/total) for k,v in gradeWeights.items()]})
fullModel={'layers':full}
generatedItems=generator.box_item_rows(fullModel);generatedOpening=generator.generate_rows(fullModel)
check('全量纯函数生成151个实际箱且4375唯一',len(generatedItems)==151 and sum(r[0]==4375 for r in generatedItems)==1)
check('全量纯函数生成151个开箱记录且4375唯一',len(generatedOpening)==151 and sum(r[0]==4375 for r in generatedOpening)==1)
generatedFirst,_=generator.first_clear_rows(fullModel)
check('全量首通75行且仅第三层发4375',len(generatedFirst)==75 and [r[1] for r in generatedFirst if r[9]==4375]==[20000012])
formalUI,fuh=read(formal/'level/UIlevel.xlsx')
fr=[r for r in formalUI[4:] if r[fuh['LevelType']]=='LevelType_SpecialEquipment']
missing=sorted({i for r in fr for i,n in pair(r[fuh['Reward']]) if i not in formalIds})
oldkeys=sum((25,1) not in pair(r[fuh['FirstReward']]) for r in fr)
summary={'checks':len(checks),'pass':sum(c['result']=='PASS' for c in checks),'fail':sum(c['result']=='FAIL' for c in checks)}
report=dict(revision='CHG-022',output=str(out),backup=str(backup),summary=summary,checks=checks,files={f:sha(out/f) for f in meta['files']},
 formalUnchangedGaps={'missingRewardItems':missing,'oldKeyPreviewRows':oldkeys,'notPublished':True},firstEquipmentExpectation={k:str(v) for k,v in expectedTotals.items()})
(ROOT/'05_验算核算/mapping-equip-first-clear-CHG022-validation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(summary));print('Formal gaps preserved: '+str(len(missing))+' missing items, '+str(oldkeys)+' old key previews; no SVN writes.')
for c in checks:
 if c['result']=='FAIL':print(c)
raise SystemExit(bool(summary['fail']))
