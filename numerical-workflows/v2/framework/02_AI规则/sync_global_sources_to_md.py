# -*- coding: utf-8 -*-
"""把 01_源表拥有的公共事实同步到参数基线 MD。

默认只检查；``--write`` 才写入 MD。最终生成器只消费同步后的 MD，
避免同一参数在源表和 MD 之间静默择值。
"""

from __future__ import annotations

import argparse
import json
import math
import sys
from datetime import datetime
from pathlib import Path

import openpyxl


V2_DIR = Path(__file__).resolve().parent.parent
SOURCE_DIR = V2_DIR / "01_源表"
REPORT = V2_DIR / "05_验算核算" / "source-md-sync-report.json"
sys.path.insert(0, str(V2_DIR))

from _shared.md_parameter_reader import load_parameter_workbook, save_parameter_workbook


SCOPES = ("共享", "养成", "战斗", "节奏")
ALIASES = {
    "生命值百分比": "生命百分比",
    "防御力百分比": "防御百分比",
    "攻击力百分比": "攻击百分比",
}


def canonical(value):
    return ALIASES.get(value, value)


def same(left, right):
    if isinstance(left, (int, float)) and isinstance(right, (int, float)):
        return math.isclose(float(left), float(right), rel_tol=1e-10, abs_tol=1e-10)
    return canonical(left) == canonical(right)


class Sync:
    def __init__(self):
        self.books = {scope: load_parameter_workbook(scope, data_only=False) for scope in SCOPES}
        self.values = {scope: load_parameter_workbook(scope, data_only=True) for scope in SCOPES}
        self.changes = []
        self.changed_scopes = set()

    def set(self, scope, sheet, row, col, value, source, transform="direct"):
        old = self.values[scope][sheet].cell(row, col).value
        value = canonical(value)
        if (old == value if transform == "terminology_alias" else same(old, value)):
            return
        cell = self.books[scope][sheet].cell(row, col)
        self.changes.append({
            "parameter_id": f"{scope}/{sheet}!{cell.coordinate}",
            "old": old,
            "new": value,
            "owner": "source_owned" if transform == "direct" else "derived",
            "source": source,
            "transform": transform,
        })
        cell.value = value
        self.changed_scopes.add(scope)

    def close(self):
        for wb in list(self.books.values()) + list(self.values.values()):
            wb.close()


def source_book(name):
    return openpyxl.load_workbook(SOURCE_DIR / name, data_only=True)


def sync_shared(sync: Sync):
    growth = source_book("养成属性分配.xlsx")
    plan = growth["模块规划"]
    for row, key in zip((2, 3, 4), ("生命", "攻击", "防御")):
        sync.set("共享", "标准人", row, 1, key, "养成属性分配.xlsx/模块规划!C2:E3")
        sync.set("共享", "标准人", row, 2, plan.cell(3, row + 1).value,
                 "养成属性分配.xlsx/模块规划!C3:E3")

    source_rows = {}
    for row in range(9, plan.max_row + 1):
        name = plan.cell(row, 2).value
        if name:
            source_rows[str(name)] = plan.cell(row, 3).value
    # 升星在源表中是百分比投放，入口位于 G 列，不是基础属性投放 C 列。
    source_rows["角色升星"] = next(
        (plan.cell(row, 7).value for row in range(9, plan.max_row + 1)
         if plan.cell(row, 2).value == "角色升星"),
        None,
    )
    name_map = {
        "角色初始": "初始",
        "角色升级": "角色升级",
        "角色升阶": "角色升阶",
        "角色星盘": "角色星盘",
        "角色升星": "角色升星",
        "专武": "专武",
        "旅券": "旅券",
        "映射装备": "映射装备",
        "房车整备": "房车升级",
    }
    target = sync.books["共享"]["投放系数"]
    for row in range(2, target.max_row + 1):
        target_name = target.cell(row, 1).value
        source_name = name_map.get(target_name)
        if source_name in source_rows:
            sync.set("共享", "投放系数", row, 2, source_rows[source_name],
                     f"养成属性分配.xlsx/模块规划[{source_name}]")

    pricing = source_book("属性定价.xlsx")
    src = pricing["属性定价"]
    dst = sync.books["共享"]["属性定价"]
    rows = max(src.max_row, dst.max_row)
    cols = max(src.max_column, dst.max_column)
    for row in range(1, rows + 1):
        for col in range(1, cols + 1):
            value = src.cell(row, col).value if row <= src.max_row and col <= src.max_column else None
            sync.set("共享", "属性定价", row, col, value,
                     f"属性定价.xlsx/属性定价!{src.cell(row, col).coordinate}")
    pricing.close()
    growth.close()


def section_rows(ws, begin, end=None):
    active = False
    result = []
    for row in range(1, ws.max_row + 1):
        value = ws.cell(row, 1).value
        if value == begin:
            active = True
            continue
        if active and end and value == end:
            break
        if active:
            result.append(row)
    return result


def sync_growth(sync: Sync):
    source = source_book("养成属性分配.xlsx")
    chars = source["角色属性分配"]
    weapons = source["专武属性分配"] if "专武属性分配" in source.sheetnames else None
    target = sync.books["养成"]["武将系数"]

    role_rows = {target.cell(row, 1).value: row for row in range(6, 15)}
    for row in range(3, 12):
        name = chars.cell(row, 2).value
        if name not in role_rows:
            continue
        dst = role_rows[name]
        for offset, col in enumerate((2, 3, 4, 5), 1):
            sync.set("养成", "武将系数", dst, offset, chars.cell(row, col).value,
                     f"养成属性分配.xlsx/角色属性分配!B{row}:E{row}")

    quality_rows = {target.cell(row, 1).value: row for row in range(18, 21)}
    for row in range(14, 17):
        quality = chars.cell(row, 2).value
        if quality in quality_rows:
            sync.set("养成", "武将系数", quality_rows[quality], 2, chars.cell(row, 3).value,
                     f"养成属性分配.xlsx/角色属性分配!B{row}:C{row}")

    # 统一攻击后，伤害类型不再二次缩放基础攻击。
    for row in range(24, 27):
        if target.cell(row, 1).value in ("物理", "魔法", "混合"):
            sync.set("养成", "武将系数", row, 2, 1,
                     "01_源表统一攻击口径", "normalize_single_attack")
            sync.set("养成", "武将系数", row, 3, 1,
                     "01_源表统一攻击口径", "normalize_single_attack")

    hero_rows = {}
    for row in section_rows(target, "角色系数列表", "专武系数列表"):
        name = target.cell(row, 1).value
        if name and target.cell(row, 2).value in ("SSR", "SR", "R"):
            hero_rows[name] = row
    for row in range(27, chars.max_row + 1):
        name = chars.cell(row, 2).value
        if name not in hero_rows:
            continue
        dst = hero_rows[name]
        mappings = {
            2: 3, 3: 4,
            5: 6, 6: 7, 7: 8,
            8: 9, 9: 10, 10: 10, 11: 11,
            13: 13,
            14: 17, 15: 18, 16: 19,
            17: 28, 18: 29,
        }
        for dst_col, src_col in mappings.items():
            sync.set("养成", "武将系数", dst, dst_col, chars.cell(row, src_col).value,
                     f"养成属性分配.xlsx/角色属性分配!{chars.cell(row, src_col).coordinate}")
        total = sum(float(chars.cell(row, col).value or 0) for col in (9, 10, 11))
        sync.set("养成", "武将系数", dst, 12, round(total, 10),
                 f"养成属性分配.xlsx/角色属性分配!I{row}:K{row}", "sum_single_attack")

    weapon_rows = {}
    for row in section_rows(target, "专武系数列表"):
        name = target.cell(row, 1).value
        if name and target.cell(row, 14).value == "通用":
            weapon_rows[name] = row
    for row in (range(4, weapons.max_row + 1) if weapons is not None else ()):
        if weapons.cell(row, 3).value not in ("SSR", "SR", "R"):
            continue
        name = weapons.cell(row, 2).value
        if name not in weapon_rows:
            continue
        dst = weapon_rows[name]
        mappings = {
            2: 3, 3: 4,
            5: 5, 6: 6, 7: 7,
            8: 8, 9: 9, 10: 9, 11: 10,
            13: 11,
        }
        for dst_col, src_col in mappings.items():
            sync.set("养成", "武将系数", dst, dst_col, weapons.cell(row, src_col).value,
                     f"养成属性分配.xlsx/专武属性分配!{weapons.cell(row, src_col).coordinate}")
        total = sum(float(weapons.cell(row, col).value or 0) for col in (8, 9, 10))
        sync.set("养成", "武将系数", dst, 12, round(total, 10),
                 f"养成属性分配.xlsx/专武属性分配!H{row}:J{row}", "sum_single_attack")

    # 没有专武源行的专属武器也按统一三属性规则重算合计。
    for row in section_rows(target, "专武系数列表"):
        if target.cell(row, 2).value not in ("SSR", "SR", "R"):
            continue
        total = sum(float(sync.books["养成"]["武将系数"].cell(row, col).value or 0)
                    for col in (8, 9, 11))
        sync.set("养成", "武将系数", row, 12, round(total, 10),
                 "统一攻击三属性派生", "sum_single_attack")

    sync.set("养成", "武将系数", 29, 10, "兼容攻击比（不计入统一攻击）",
             "01_源表统一攻击口径", "schema_label")
    sync.set("养成", "武将系数", 29, 12, "统一三属性调整合计",
             "01_源表统一攻击口径", "schema_label")
    sync.set("养成", "武将系数", 70, 10, "兼容攻击比（不计入统一攻击）",
             "01_源表统一攻击口径", "schema_label")
    sync.set("养成", "武将系数", 70, 12, "统一三属性调整合计",
             "01_源表统一攻击口径", "schema_label")

    # 全参数域统一百分比术语；源表旧写法仅作为别名。
    for sheet_name in ("武将系数", "专武突破"):
        ws = sync.books["养成"][sheet_name]
        values = sync.values["养成"][sheet_name]
        for row in range(1, ws.max_row + 1):
            for col in range(1, ws.max_column + 1):
                old = values.cell(row, col).value
                new = canonical(old)
                if old != new:
                    sync.set("养成", sheet_name, row, col, new,
                             "属性定价.xlsx/属性定价", "terminology_alias")
    source.close()


def sync_battle(sync: Sync):
    source = source_book("战斗公式.xlsx")
    model = source["怪物标准模型"]
    target = sync.books["战斗"]["怪物模型参数"]
    column_map = {
        "裸体生命": 3, "裸体攻击": 4, "裸体防御": 5,
        "额外生命": 6, "额外攻击": 7, "额外防御": 8,
        "生命%": 9, "攻击%": 10, "防御%": 11,
        "普攻系数倍率": 15, "技能系数倍率": 16, "必杀系数倍率": 17,
        "战斗支持倍率": 18, "暴击率": 19, "暴击伤害": 20,
        "普攻乘区": 21, "必杀乘区": 22, "增伤乘区": 24,
        "人数系数": 34, "攻击系数修正": 35, "BOSS生命系数修正": 36,
        "小怪生命系数修正": 37, "小怪攻击系数修正": 38,
    }
    sync.set("战斗", "怪物模型参数", 2, 2, model["B3"].value,
             "战斗公式.xlsx/怪物标准模型!B3")
    for row in range(5, 28):
        name = target.cell(row, 1).value
        if name not in column_map:
            continue
        col = column_map[name]
        sync.set("战斗", "怪物模型参数", row, 2, col,
                 f"战斗公式.xlsx/怪物标准模型!{openpyxl.utils.get_column_letter(col)}4", "source_column")
        for dst_col, src_row in ((3, 5), (4, 6), (5, 7)):
            sync.set("战斗", "怪物模型参数", row, dst_col, model.cell(src_row, col).value,
                     f"战斗公式.xlsx/怪物标准模型!{model.cell(src_row, col).coordinate}")

    for index, row in enumerate(range(32, 41), 45):
        sync.set("战斗", "怪物模型参数", row, 1, model.cell(7, index).value,
                 f"战斗公式.xlsx/怪物标准模型!{model.cell(7, index).coordinate}")
        sync.set("战斗", "怪物模型参数", row, 2, model.cell(10, index).value,
                 f"战斗公式.xlsx/怪物标准模型!{model.cell(10, index).coordinate}")

    delivery = source["预计投放属性数值"]
    sync.set("战斗", "战斗后期参数", 4, 2, delivery["B12"].value,
             "战斗公式.xlsx/预计投放属性数值!B12", "derived_source_total")
    sync.set("战斗", "战斗后期参数", 5, 2, delivery["C12"].value,
             "战斗公式.xlsx/预计投放属性数值!C12", "derived_source_total")
    source.close()


def sync_pace(sync: Sync):
    source = source_book("节奏口径.xlsx")
    levels = source["角色等级"]
    source_levels = {}
    for row in range(5, levels.max_row + 1):
        level = levels.cell(row, 2).value
        if isinstance(level, (int, float)):
            source_levels[int(level)] = (levels.cell(row, 3).value, levels.cell(row, 5).value, row)
    day_curve = []
    for row in range(5, levels.max_row + 1):
        day = levels.cell(row, 16).value
        actual_level = levels.cell(row, 18).value
        if isinstance(day, (int, float)) and isinstance(actual_level, (int, float)):
            day_curve.append((int(day), int(actual_level), row))
    day_curve.sort()

    target = sync.books["节奏"]["主角等级_输入"]
    for row in range(5, target.max_row + 1):
        level = target.cell(row, 1).value
        if not isinstance(level, (int, float)):
            continue
        level = int(level)
        if level in source_levels:
            growth, stamina, source_row = source_levels[level]
            sync.set("节奏", "主角等级_输入", row, 2, growth,
                     f"节奏口径.xlsx/角色等级!C{source_row}")
            sync.set("节奏", "主角等级_输入", row, 3, stamina,
                     f"节奏口径.xlsx/角色等级!E{source_row}")
        if day_curve:
            day = next((day for day, reached, _ in day_curve if reached >= level), None)
            sync.set("节奏", "主角等级_输入", row, 5, day,
                     "节奏口径.xlsx/角色等级!P:R", "earliest_day_reaching_level")
    source.close()


def build_report(sync: Sync, mode: str):
    return {
        "schema": "vgame-source-md-sync/v1",
        "generated_at": datetime.now().astimezone().isoformat(timespec="seconds"),
        "mode": mode,
        "canonical_terms": ALIASES,
        "difference_count": len(sync.changes),
        "changed_scopes": sorted(sync.changed_scopes),
        "differences": sync.changes,
    }


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--write", action="store_true", help="把 source_owned/derived 差异写入参数 MD")
    parser.add_argument("--report", default=str(REPORT))
    args = parser.parse_args()

    sync = Sync()
    try:
        sync_shared(sync)
        sync_growth(sync)
        sync_battle(sync)
        sync_pace(sync)
        report = build_report(sync, "write" if args.write else "check")
        report_path = Path(args.report)
        report_path.parent.mkdir(parents=True, exist_ok=True)
        report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        if args.write:
            for scope in sorted(sync.changed_scopes):
                save_parameter_workbook(scope, sync.books[scope])
            print(f"[WRITE] synchronized={len(sync.changes)} scopes={','.join(sorted(sync.changed_scopes))}")
            print(f"[REPORT] {report_path}")
            return 0
        if sync.changes:
            print(f"[DRIFT] differences={len(sync.changes)}")
            for item in sync.changes[:20]:
                print(f"  {item['parameter_id']}: {item['old']!r} -> {item['new']!r} ({item['source']})")
            print(f"[REPORT] {report_path}")
            return 1
        print("[OK] 01_源表与 source_owned/derived MD 参数一致")
        print(f"[REPORT] {report_path}")
        return 0
    finally:
        sync.close()


if __name__ == "__main__":
    raise SystemExit(main())
