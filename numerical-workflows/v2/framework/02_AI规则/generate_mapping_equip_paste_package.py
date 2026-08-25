#!/usr/bin/env python3
"""生成映射装备正式配置粘贴包 P01。

本脚本只生成独立 Excel 值包，绝不修改 Vgame Datas。复制区严格按正式表字段
顺序排列；概率权重采用项目现有千分制并以最大余数法保证每行合计 1000。
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import sys
from datetime import datetime
from decimal import Decimal, ROUND_HALF_UP
from fractions import Fraction
from pathlib import Path
from typing import Iterable, Sequence

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


PACKAGE_ID = "MAPEQUIP-20260710-P01"
PACKAGE_VERSION = "1.0"
GENERATION_MARKER = "mapping-equip-paste-package/v1"
WEIGHT_SCALE = 1000

FRAME_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUTPUT = FRAME_ROOT / "04_数值文档" / "配置粘贴包" / "映射装备正式配置粘贴包_P01.xlsx"
DEFAULT_REPORT = FRAME_ROOT / "05_验算核算" / "mapping-equip-paste-package-P01-generation.json"
DEFAULT_CONFIG_ROOT = Path(os.environ.get("VGAME_CONFIG_DATAS", ""))

TARGET_FILES = [
    "SpecialEquipment/Portable/PortableExp.xlsx",
    "SpecialEquipment/SpecialEquipmentSubAttrCoef.xlsx",
    "SpecialEquipment/Portable/PortableLvUp.xlsx",
    "SpecialEquipment/Portable/PortableCompoundProb.xlsx",
    "SpecialEquipment/Portable/PortableCraftProb.xlsx",
    "SpecialEquipment/Portable/PortableEquipCompound.xlsx",
    "SpecialEquipment/SpecialEquipmentQuality.xlsx",
]

SOURCE_MDS = [
    "03_沉淀MD/养成/映射装备养成数值沉淀-2026-07-07.md",
    "03_沉淀MD/经济/21_养成_映射装备经济消耗沉淀-2026-07-09.md",
    "03_沉淀MD/养成/映射装备数值总表生成参数-2026-07-10.md",
]

QUALITIES = ["Green", "Blue", "Purple", "Orange", "Red"]
QUALITY_INDEX = {quality: index for index, quality in enumerate(QUALITIES)}

NAVY = "17365D"
BLUE = "4472C4"
LIGHT_BLUE = "D9EAF7"
PALE_GREEN = "E2F0D9"
PALE_YELLOW = "FFF2CC"
PALE_RED = "F4CCCC"
PALE_GRAY = "E7E6E6"
WHITE = "FFFFFF"
BLACK = "000000"
INPUT_BLUE = "0070C0"
PASS_GREEN = "70AD47"
WARN_ORANGE = "ED7D31"
THIN_GRAY = Side(style="thin", color="B7B7B7")
GRID_BORDER = Border(left=THIN_GRAY, right=THIN_GRAY, top=THIN_GRAY, bottom=THIN_GRAY)


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def source_records(config_root: Path) -> list[dict]:
    records: list[dict] = []
    for relative_path in TARGET_FILES:
        path = config_root / relative_path
        if not path.is_file():
            raise FileNotFoundError(path)
        records.append(
            {
                "relative_path": relative_path,
                "bytes": path.stat().st_size,
                "sha256": sha256_file(path),
            }
        )
    return records


def aggregate_records(records: Iterable[dict]) -> str:
    digest = hashlib.sha256()
    for record in sorted(records, key=lambda item: item["relative_path"]):
        digest.update(
            (
                f"{record['relative_path']}\0{record['bytes']}\0{record['sha256']}\n"
            ).encode("utf-8")
        )
    return digest.hexdigest()


def round_to_ten(value: Decimal) -> int:
    return int((value / Decimal(10)).quantize(Decimal("1"), rounding=ROUND_HALF_UP) * 10)


FORGE_LEVEL_COSTS = (
    10, 20, 30, 40, 50, 60, 70, 100, 120,
    130, 140, 160, 180, 190, 210, 220, 240, 250, 280,
    350, 400, 450, 500, 550, 650, 750, 850, 950, 1050,
    1130, 1190, 1250, 1310, 1370, 1430, 1490, 1550, 1610, 1670,
    2000, 2100, 2200, 2300, 2400, 2500, 2600, 2700, 2800, 2900,
    3800, 4000, 4200, 4400, 4600, 4800, 5000, 5200, 5400, 5600,
)


def forge_level_cost(target_level: int) -> int:
    if not 2 <= target_level <= 60:
        raise ValueError(f"锻造台目标等级超出范围: {target_level}")
    return FORGE_LEVEL_COSTS[target_level - 2]


def largest_remainder(probabilities: dict[str, Fraction]) -> dict[str, int]:
    """按品质顺序稳定分配余数，返回合计严格为 WEIGHT_SCALE 的整数权重。"""

    if sum(probabilities.values(), Fraction(0, 1)) != 1:
        raise ValueError(f"概率合计不为1: {probabilities}")
    raw = {quality: probability * WEIGHT_SCALE for quality, probability in probabilities.items()}
    weights = {quality: value.numerator // value.denominator for quality, value in raw.items()}
    remainder = WEIGHT_SCALE - sum(weights.values())
    order = sorted(
        raw,
        key=lambda quality: (
            -(raw[quality] - weights[quality]),
            QUALITY_INDEX[quality],
        ),
    )
    for quality in order[:remainder]:
        weights[quality] += 1
    if sum(weights.values()) != WEIGHT_SCALE or any(value < 0 for value in weights.values()):
        raise AssertionError(weights)
    return weights


def serialize_weights(weights: dict[str, int]) -> str:
    return ";".join(
        f"{quality},{weights[quality]}"
        for quality in QUALITIES
        if weights.get(quality, 0) > 0
    )


def portable_exp_rows() -> list[list]:
    levels = list(range(5, 61, 5))
    values = {
        5: [1, 2, 5, 12, 30],
        10: [1, 2, 5, 12, 30],
        15: [2, 4, 10, 24, 60],
        20: [2, 4, 10, 24, 60],
        25: [3, 6, 15, 36, 90],
        30: [3, 6, 15, 36, 90],
        35: [4, 8, 20, 48, 120],
        40: [4, 8, 20, 48, 120],
        45: [5, 10, 25, 60, 150],
        50: [5, 10, 25, 60, 150],
        55: [6, 12, 30, 72, 180],
        60: [6, 12, 30, 72, 180],
    }
    rows: list[list] = []
    row_id = 1
    for level in levels:
        for quality, exp in zip(QUALITIES, values[level]):
            rows.append([row_id, level, quality, exp])
            row_id += 1
    return rows


def sub_attr_rows() -> list[list]:
    # 为避免既有 Id 2~11 引用漂移，保留原 Id，仅追加 Id12=Lv10。
    return [
        [1, 5, 0.200],
        [2, 15, 0.345],
        [3, 20, 0.418],
        [4, 25, 0.491],
        [5, 30, 0.564],
        [6, 35, 0.636],
        [7, 40, 0.709],
        [8, 45, 0.782],
        [9, 50, 0.855],
        [10, 55, 0.927],
        [11, 60, 1.000],
        [12, 10, 0.273],
    ]


def portable_level_rows() -> list[list]:
    return [[target_level - 1, target_level, forge_level_cost(target_level)] for target_level in range(2, 61)]


def compound_probabilities(level: int, quality: str) -> dict[str, Fraction]:
    progress = Fraction(level - 1, 59)
    if quality == "Green":
        high = Fraction(8, 100) + Fraction(17, 100) * progress
        return {"Blue": 1 - high, "Purple": high}
    if quality == "Blue":
        high = Fraction(5, 100) + Fraction(15, 100) * progress
        return {"Purple": 1 - high, "Orange": high}
    if quality == "Purple":
        orange = Fraction(30, 100) + Fraction(30, 100) * progress
        red = Fraction(1, 100) + Fraction(2, 100) * progress
        return {"Purple": 1 - orange - red, "Orange": orange, "Red": red}
    if quality == "Orange":
        red = Fraction(20, 100) + Fraction(20, 100) * progress
        return {"Orange": 1 - red, "Red": red}
    if quality == "Red":
        return {"Red": Fraction(1, 1)}
    raise ValueError(quality)


def craft_probabilities(level: int, quality: str) -> dict[str, Fraction]:
    progress = Fraction(level - 1, 59)
    if quality == "Purple":
        orange = Fraction(10, 100) + Fraction(10, 100) * progress
        return {"Purple": 1 - orange, "Orange": orange}
    if quality == "Orange":
        red = Fraction(2, 100) + Fraction(3, 100) * progress
        return {"Orange": 1 - red, "Red": red}
    if quality == "Red":
        return {"Red": Fraction(1, 1)}
    raise ValueError(quality)


def probability_rows(kind: str) -> tuple[list[list], list[list[int]]]:
    rows: list[list] = []
    audit: list[list[int]] = []
    input_qualities = QUALITIES if kind == "compound" else ["Purple", "Orange", "Red"]
    probability_builder = compound_probabilities if kind == "compound" else craft_probabilities
    row_id = 1
    for level in range(1, 61):
        for quality in input_qualities:
            probabilities = probability_builder(level, quality)
            weights = largest_remainder(probabilities)
            full_weights = {item: weights.get(item, 0) for item in QUALITIES}
            rows.append([row_id, level, quality, serialize_weights(full_weights)])
            audit.append([full_weights[item] for item in QUALITIES])
            row_id += 1
    return rows, audit


def compound_cost_rows() -> list[list]:
    matrix = {
        1: [100, 140, 200, 280, 400],
        2: [350, 490, 700, 980, 1400],
        3: [1000, 1400, 2000, 2800, 4000],
        4: [2500, 3500, 5000, 7000, 10000],
        5: [3800, 5320, 7600, 10640, 15200],
        6: [5500, 7700, 11000, 15400, 22000],
        7: [7500, 10500, 15000, 21000, 30000],
        8: [10000, 14000, 20000, 28000, 40000],
    }
    rows: list[list] = []
    row_id = 1
    for range_id in range(1, 9):
        for cost in matrix[range_id]:
            rows.append([row_id, cost])
            row_id += 1
    return rows


def quality_patch_rows() -> list[list]:
    return [
        [1, "Green", None, 0.20, 0.30],
        [2, "Blue", None, 0.30, 0.45],
        [3, "Purple", None, 0.45, 0.60],
        [4, "Orange", None, 0.70, 0.80],
        [5, "Red", "23,1", 1.00, 1.00],
    ]


def apply_global_sheet_style(ws) -> None:
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A5"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.outlinePr.summaryRight = True


def style_title(ws, title: str, max_col: int) -> None:
    max_col = max(max_col, 6)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max_col)
    cell = ws.cell(1, 1, title)
    cell.font = Font(name="微软雅黑", size=16, bold=True, color=WHITE)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 28


def set_info_row(ws, row: int, label: str, value: str, max_col: int, fill: str = LIGHT_BLUE) -> None:
    max_col = max(max_col, 6)
    ws.cell(row, 1, label)
    ws.cell(row, 1).font = Font(name="微软雅黑", bold=True, color=BLACK)
    ws.cell(row, 1).fill = PatternFill("solid", fgColor=fill)
    ws.cell(row, 1).alignment = Alignment(vertical="center")
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=max_col)
    ws.cell(row, 2, value)
    ws.cell(row, 2).font = Font(name="微软雅黑", color=BLACK)
    ws.cell(row, 2).fill = PatternFill("solid", fgColor=fill)
    ws.cell(row, 2).alignment = Alignment(vertical="center", wrap_text=True)
    for column in range(1, max_col + 1):
        ws.cell(row, column).border = GRID_BORDER
    ws.row_dimensions[row].height = 32 if row == 3 else 22


def style_header(ws, headers: Sequence[str | None], row: int = 4) -> None:
    for column, header in enumerate(headers, start=1):
        cell = ws.cell(row, column, header)
        if header is None:
            cell.fill = PatternFill("solid", fgColor=WHITE)
            continue
        cell.font = Font(name="微软雅黑", bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=BLUE if column <= 5 else "7F8C8D")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = GRID_BORDER
    ws.row_dimensions[row].height = 30


def write_rows(ws, rows: Sequence[Sequence], *, start_row: int, start_col: int, fill: str) -> None:
    for row_offset, values in enumerate(rows):
        row_number = start_row + row_offset
        for column_offset, value in enumerate(values):
            cell = ws.cell(row_number, start_col + column_offset, value)
            cell.font = Font(name="微软雅黑", size=10, color=INPUT_BLUE)
            cell.fill = PatternFill("solid", fgColor=fill)
            cell.alignment = Alignment(vertical="center", wrap_text=False)
            cell.border = GRID_BORDER


def add_data_sheet(
    wb: Workbook,
    *,
    name: str,
    title: str,
    target: str,
    instruction: str,
    headers: Sequence[str],
    rows: Sequence[Sequence],
    status: str,
    copy_columns: Sequence[int] | None = None,
) -> object:
    ws = wb.create_sheet(name)
    apply_global_sheet_style(ws)
    max_col = max(len(headers), 6)
    style_title(ws, title, max_col)
    set_info_row(ws, 2, "正式目标", target, max_col)
    set_info_row(ws, 3, "复制操作", instruction, max_col, PALE_GREEN if status == "READY" else PALE_YELLOW)
    style_header(ws, headers)
    write_rows(ws, rows, start_row=5, start_col=1, fill=PALE_GREEN if status == "READY" else PALE_YELLOW)
    if copy_columns:
        copy_set = set(copy_columns)
        for row_number in range(5, 5 + len(rows)):
            for column in range(1, len(headers) + 1):
                if column not in copy_set:
                    ws.cell(row_number, column).fill = PatternFill("solid", fgColor=PALE_GRAY)
                    ws.cell(row_number, column).font = Font(name="微软雅黑", size=10, color="666666")
    ws.auto_filter.ref = f"A4:{get_column_letter(len(headers))}{4 + len(rows)}"
    for column, header in enumerate(headers, start=1):
        width = 14
        if "QualityList" in header:
            width = 42
        elif header in {"Quality", "ExtraItemList"}:
            width = 18
        ws.column_dimensions[get_column_letter(column)].width = width
    ws.sheet_properties.tabColor = PASS_GREEN if status == "READY" else WARN_ORANGE
    return ws


def add_probability_sheet(
    wb: Workbook,
    *,
    name: str,
    title: str,
    target: str,
    instruction: str,
    level_header: str,
    rows: Sequence[Sequence],
    audits: Sequence[Sequence[int]],
) -> None:
    ws = wb.create_sheet(name)
    apply_global_sheet_style(ws)
    headers = ["Id", level_header, "Quality", "QualityList", None, "Green", "Blue", "Purple", "Orange", "Red", "权重和", "校验"]
    style_title(ws, title, len(headers))
    set_info_row(ws, 2, "正式目标", target, len(headers))
    set_info_row(
        ws,
        3,
        "复制操作",
        instruction + "；F:L 为隐藏验算列，不属于粘贴区。",
        len(headers),
        PALE_YELLOW,
    )
    style_header(ws, headers)
    write_rows(ws, rows, start_row=5, start_col=1, fill=PALE_YELLOW)
    for index, values in enumerate(audits, start=5):
        for offset, value in enumerate(values, start=6):
            cell = ws.cell(index, offset, value)
            cell.font = Font(name="微软雅黑", color="666666")
            cell.fill = PatternFill("solid", fgColor=PALE_GRAY)
            cell.border = GRID_BORDER
        ws.cell(index, 11, f"=SUM(F{index}:J{index})")
        ws.cell(index, 12, f'=IF(K{index}={WEIGHT_SCALE},"PASS","FAIL")')
        for column in (11, 12):
            ws.cell(index, column).font = Font(name="微软雅黑", color=BLACK)
            ws.cell(index, column).fill = PatternFill("solid", fgColor=PALE_GRAY)
            ws.cell(index, column).border = GRID_BORDER
    ws.auto_filter.ref = f"A4:D{4 + len(rows)}"
    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 44
    ws.column_dimensions["E"].width = 3
    for column in "FGHIJKL":
        ws.column_dimensions[column].width = 11
        ws.column_dimensions[column].hidden = True
    ws.sheet_properties.tabColor = WARN_ORANGE


def add_readme(wb: Workbook) -> None:
    ws = wb.create_sheet("00_说明")
    ws.sheet_view.showGridLines = False
    style_title(ws, "映射装备正式配置粘贴包 P01", 9)
    ws["A2"] = "交付结论"
    ws["B2"] = "本文件提供按正式字段顺序排好的纯值数据；未修改 Vgame Datas，未执行 Luban，未代表上线批准。"
    ws.merge_cells("B2:I2")
    for cell in ws[2]:
        cell.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
        cell.border = GRID_BORDER
        cell.font = Font(name="微软雅黑", bold=cell.column == 1)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[2].height = 34

    instructions = [
        "先备份正式工作簿；只复制本包标注的数据区域，不复制表头。",
        "在正式表中从 B5 或指定目标列开始‘仅粘贴值’；正式表 A 列是 Luban 控制列，禁止覆盖。",
        "绿色为可直接粘贴；黄色为数据完整但有运行时/字段语义门禁，确认后再粘贴。",
        "概率采用项目现有总权重 1000 约定，最大余数法量化；每项误差小于 0.1 个百分点，每行严格合计 1000。",
        "PortableCraftProb.xlsx 的实际 Sheet 名仍是 PortableCompoundProb，这是现有正式表命名，不是本包笔误。",
        "粘贴后仍需运行配置导出、引用检查和业务冒烟；本包不把‘能粘贴’等同于‘可上线’。",
        "本包只锁定7个正式目标文件；Items、Drop、UILevel、Unlock等跨表条件项落地前必须刷新全局映射快照。",
    ]
    ws["A4"] = "操作规则"
    ws["A4"].font = Font(name="微软雅黑", bold=True, color=WHITE)
    ws["A4"].fill = PatternFill("solid", fgColor=BLUE)
    ws.merge_cells("A4:I4")
    for row_number, instruction in enumerate(instructions, start=5):
        ws.cell(row_number, 1, f"{row_number - 4}.")
        ws.cell(row_number, 2, instruction)
        ws.merge_cells(start_row=row_number, start_column=2, end_row=row_number, end_column=9)
        for column in range(1, 10):
            cell = ws.cell(row_number, column)
            cell.border = GRID_BORDER
            cell.fill = PatternFill("solid", fgColor=WHITE)
            cell.font = Font(name="微软雅黑")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[row_number].height = 28

    headers = ["状态", "本包Sheet", "正式工作簿", "实际Sheet", "复制范围", "粘贴范围", "行数", "粘贴方式", "说明"]
    start = 12
    for column, header in enumerate(headers, start=1):
        cell = ws.cell(start, column, header)
        cell.font = Font(name="微软雅黑", bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = GRID_BORDER
    mappings = [
        ["READY", "01_PortableExp", "PortableExp.xlsx", "PortableExp", "A5:D64", "B5:E64", 60, "整表值替换", "12等级×5品质"],
        ["READY", "02_SubAttrCoef", "SpecialEquipmentSubAttrCoef.xlsx", "SpecialEquipmentSubAttrCoef", "A5:C16", "B5:D16", 12, "整表值替换", "保留Id1~11；Id12=Lv10"],
        ["READY", "03_PortableLvUp", "PortableLvUp.xlsx", "PortableLvUp", "A5:C63", "B5:D63", 59, "整表值替换", "TargetLv 2~60；总经验94500"],
        ["CONDITIONAL", "04_CompoundProb", "PortableCompoundProb.xlsx", "PortableCompoundProb", "A5:D304", "B5:E304", 300, "整表值替换", "总权重1000；需Portable运行时冒烟"],
        ["CONDITIONAL", "05_CraftProb", "PortableCraftProb.xlsx", "PortableCompoundProb", "A5:D184", "B5:E184", 180, "整表值替换", "正式Sheet误名；需Portable运行时冒烟"],
        ["CONDITIONAL", "10_CompoundCost", "PortableEquipCompound.xlsx", "PortableEquipCompound", "B5:B44", "F5:F44", 40, "单列值补丁", "只粘CostCurrency；不覆盖EquipResult"],
        ["CONDITIONAL", "11_QualityPatch", "SpecialEquipmentQuality.xlsx", "SpecialEquipment", "C5:C9", "E5:E9", 5, "单列值补丁", "空白代表清空；旧运行时退场后再贴"],
        ["CONDITIONAL", "11_QualityPatch", "SpecialEquipmentQuality.xlsx", "SpecialEquipment", "D5:E9", "M5:N9", 5, "两列值补丁", "只改两项品质系数"],
    ]
    for row_number, row in enumerate(mappings, start=start + 1):
        fill = PALE_GREEN if row[0] == "READY" else PALE_YELLOW
        for column, value in enumerate(row, start=1):
            cell = ws.cell(row_number, column, value)
            cell.font = Font(name="微软雅黑", color=INPUT_BLUE if column in (5, 6, 7) else BLACK)
            cell.fill = PatternFill("solid", fgColor=fill)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = GRID_BORDER
        ws.row_dimensions[row_number].height = 34
    ws.freeze_panes = "A13"
    widths = [15, 22, 34, 30, 15, 15, 10, 16, 42]
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.sheet_properties.tabColor = NAVY


def add_blockers(wb: Workbook) -> None:
    ws = wb.create_sheet("90_阻断清单")
    ws.sheet_view.showGridLines = False
    style_title(ws, "尚不能伪装成正式完整数据的配置项", 7)
    set_info_row(
        ws,
        2,
        "原则",
        "缺少字段语义、引用ID或运行时消费者证据时，只列阻断，不编造可粘贴值。",
        7,
        PALE_RED,
    )
    headers = ["状态", "正式表/模块", "阻断字段", "已知数据", "缺失决策", "解除条件", "本包处理"]
    style_header(ws, headers)
    rows = [
        ["BLOCKED", "PortableRange", "LevelMin/LevelMax/UnlockId", "目标8档；现正式表仅2行", "区间语义与8个UnlockId", "确认运行时取值语义并分配引用", "不提供假值"],
        ["BLOCKED", "PortableEquipCompound", "EquipResult", "8×5金币矩阵已知", "8档生成结果ID；Range3~8闭环", "先补SpecialEquipmentGenerate与PortableRange", "仅给CostCurrency单列"],
        ["BLOCKED", "PortableEquipCraft", "Tier/Range/Quality/材料ItemId", "打造档位与概率设计已知", "当前Schema无法表达8档打造；72主材+4副材ID未分配", "扩Schema并完成Items/产出链", "不提供完整表"],
        ["BLOCKED", "SpecialEquipmentQuality", "Basic/旧升级打孔鉴定字段", "品质系数、红装额外道具已知", "旧Lua仍读取旧字段；分解金币缺等级维度", "旧链退场并新增等级承载", "仅给条件字段补丁"],
        ["BLOCKED", "SpecialEquipmentMain/Effect/Generate", "曲线/生成池/不重复规则", "MD有数值目标", "现有字段承载与生成ID不完整", "Schema与运行时实现闭环", "不提供假值"],
        ["BLOCKED", "Items/Drop/UIlevel/Unlock", "ItemId/DropId/UnlockId", "来源与数量设计部分已知", "正式ID及跨表引用未分配", "完成ID规划与引用核验", "不提供假值"],
        ["RUNTIME_GATE", "Portable 7表", "客户端/服务端消费者", "生成表与UI百分比证据存在", "未找到非Generated的完整消费链", "导出后做服务器+客户端冒烟", "概率表标条件粘贴"],
    ]
    for row_number, row in enumerate(rows, start=5):
        for column, value in enumerate(row, start=1):
            cell = ws.cell(row_number, column, value)
            cell.font = Font(name="微软雅黑")
            cell.fill = PatternFill("solid", fgColor=PALE_RED if row[0] == "BLOCKED" else PALE_YELLOW)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = GRID_BORDER
        ws.row_dimensions[row_number].height = 58
    widths = [16, 30, 34, 34, 46, 42, 28]
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = "A5"
    ws.sheet_properties.tabColor = "C00000"


def add_validation_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("99_验算")
    ws.sheet_view.showGridLines = False
    style_title(ws, "粘贴包内置验算", 5)
    set_info_row(ws, 2, "说明", "本页公式只做包内复核，不属于任何复制区域；最终以独立Python门禁报告为准。", 5)
    headers = ["验算项", "实算", "期望", "结果", "口径"]
    style_header(ws, headers)
    checks = [
        ["PortableExp行数", "=COUNTA('01_PortableExp'!A5:A64)", 60, None, "12等级×5品质"],
        ["PortableExp最小经验", "=MIN('01_PortableExp'!D5:D64)", 1, None, "MD矩阵"],
        ["PortableExp最大经验", "=MAX('01_PortableExp'!D5:D64)", 180, None, "MD矩阵"],
        ["SubAttrCoef行数", "=COUNTA('02_SubAttrCoef'!A5:A16)", 12, None, "保留旧Id并追加Lv10"],
        ["PortableLvUp行数", "=COUNTA('03_PortableLvUp'!A5:A63)", 59, None, "TargetLv2~60"],
        ["PortableLvUp总经验", "=SUM('03_PortableLvUp'!C5:C63)", 94500, None, "升级曲线总和"],
        ["CompoundProb行数", "=COUNTA('04_CompoundProb'!A5:A304)", 300, None, "60级×5品质"],
        ["CompoundProb权重失败数", "=COUNTIF('04_CompoundProb'!L5:L304,\"FAIL\")", 0, None, "每行总权重1000"],
        ["CraftProb行数", "=COUNTA('05_CraftProb'!A5:A184)", 180, None, "60级×3品质"],
        ["CraftProb权重失败数", "=COUNTIF('05_CraftProb'!L5:L184,\"FAIL\")", 0, None, "每行总权重1000"],
        ["CompoundCost行数", "=COUNTA('10_CompoundCost'!A5:A44)", 40, None, "8档×5品质"],
        ["CompoundCost总金币", "=SUM('10_CompoundCost'!B5:B44)", 344400, None, "金币矩阵"],
        ["QualityPatch行数", "=COUNTA('11_QualityPatch'!A5:A9)", 5, None, "五品质"],
    ]
    for row_number, row in enumerate(checks, start=5):
        row[3] = f'=IF(B{row_number}=C{row_number},"PASS","FAIL")'
        for column, value in enumerate(row, start=1):
            cell = ws.cell(row_number, column, value)
            cell.font = Font(name="微软雅黑", color=BLACK)
            cell.fill = PatternFill("solid", fgColor=WHITE)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = GRID_BORDER
        ws.cell(row_number, 2).number_format = "0.000"
        ws.cell(row_number, 3).number_format = "0.000"
    ws.conditional_formatting.add(
        f"D5:D{4 + len(checks)}",
        CellIsRule(operator="equal", formula=['"PASS"'], fill=PatternFill("solid", fgColor=PALE_GREEN)),
    )
    widths = [32, 18, 18, 14, 42]
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = "A5"
    ws.sheet_properties.tabColor = PASS_GREEN


def add_meta(
    wb: Workbook,
    *,
    generated_at: str,
    config_root: Path,
    records: list[dict],
    source_snapshot_hash: str,
) -> None:
    ws = wb.create_sheet("__meta__")
    metadata = [
        ("package_id", PACKAGE_ID),
        ("package_version", PACKAGE_VERSION),
        ("generation_marker", GENERATION_MARKER),
        ("generated_at", generated_at),
        ("official_config_write", False),
        ("official_config_root", str(config_root)),
        ("official_target_hashes_json", json.dumps(records, ensure_ascii=False, separators=(",", ":"))),
        ("source_snapshot_hash", source_snapshot_hash),
        ("weight_scale", WEIGHT_SCALE),
        ("probability_rounding", "largest_remainder_quality_order_v1"),
        ("source_md_json", json.dumps(SOURCE_MDS, ensure_ascii=False, separators=(",", ":"))),
        ("direct_ready_sheets", "01_PortableExp;02_SubAttrCoef;03_PortableLvUp"),
        ("conditional_sheets", "04_CompoundProb;05_CraftProb;10_CompoundCost;11_QualityPatch"),
        ("luban_executed", False),
    ]
    ws.append(["key", "value"])
    for key, value in metadata:
        ws.append([key, value])
    ws.sheet_state = "veryHidden"


def build_workbook(output_path: Path, config_root: Path, report_path: Path) -> dict:
    if not config_root.is_dir():
        raise NotADirectoryError(config_root)
    generated_at = datetime.now().astimezone().isoformat(timespec="seconds")
    records = source_records(config_root)
    source_snapshot_hash = aggregate_records(records)

    exp_rows = portable_exp_rows()
    sub_rows = sub_attr_rows()
    level_rows = portable_level_rows()
    compound_rows, compound_audit = probability_rows("compound")
    craft_rows, craft_audit = probability_rows("craft")
    cost_rows = compound_cost_rows()
    quality_rows = quality_patch_rows()

    if len(exp_rows) != 60 or len(sub_rows) != 12 or len(level_rows) != 59:
        raise AssertionError("直接粘贴数据行数异常")
    if len(compound_rows) != 300 or len(craft_rows) != 180:
        raise AssertionError("概率表数据行数异常")
    if sum(row[2] for row in level_rows) != 94500:
        raise AssertionError("锻造台总经验异常")
    if sum(row[1] for row in cost_rows) != 344400:
        raise AssertionError("合成金币总和异常")

    wb = Workbook()
    wb.remove(wb.active)
    add_readme(wb)
    add_data_sheet(
        wb,
        name="01_PortableExp",
        title="PortableExp：装备提供锻造台经验（可直接粘贴）",
        target="SpecialEquipment/Portable/PortableExp.xlsx / PortableExp / B5:E64",
        instruction="复制 A5:D64 → 正式表 B5:E64；仅粘贴值",
        headers=["Id", "EquipLv", "Quality", "Exp"],
        rows=exp_rows,
        status="READY",
    )
    sub_ws = add_data_sheet(
        wb,
        name="02_SubAttrCoef",
        title="SpecialEquipmentSubAttrCoef：副属性等级倍率（可直接粘贴）",
        target="SpecialEquipment/SpecialEquipmentSubAttrCoef.xlsx / SpecialEquipmentSubAttrCoef / B5:D16",
        instruction="复制 A5:C16 → 正式表 B5:D16；仅粘贴值；Id12=Lv10用于避免既有Id漂移",
        headers=["Id", "EquipLv", "SubAttrCoef"],
        rows=sub_rows,
        status="READY",
    )
    for row_number in range(5, 17):
        sub_ws.cell(row_number, 3).number_format = "0.000"
    add_data_sheet(
        wb,
        name="03_PortableLvUp",
        title="PortableLvUp：锻造台升级经验（可直接粘贴）",
        target="SpecialEquipment/Portable/PortableLvUp.xlsx / PortableLvUp / B5:D63",
        instruction="复制 A5:C63 → 正式表 B5:D63；仅粘贴值",
        headers=["Id", "TargetLv", "CostExp"],
        rows=level_rows,
        status="READY",
    )
    add_probability_sheet(
        wb,
        name="04_CompoundProb",
        title="PortableCompoundProb：合成品质概率（条件粘贴）",
        target="SpecialEquipment/Portable/PortableCompoundProb.xlsx / PortableCompoundProb / B5:E304",
        instruction="确认Portable运行时后，复制 A5:D304 → 正式表 B5:E304；仅粘贴值",
        level_header="Grade",
        rows=compound_rows,
        audits=compound_audit,
    )
    add_probability_sheet(
        wb,
        name="05_CraftProb",
        title="PortableCraftProb：打造品质概率（条件粘贴）",
        target="SpecialEquipment/Portable/PortableCraftProb.xlsx / 实际Sheet PortableCompoundProb / B5:E184",
        instruction="确认Portable运行时后，复制 A5:D184 → 正式表 B5:E184；仅粘贴值",
        level_header="PortableLv",
        rows=craft_rows,
        audits=craft_audit,
    )
    add_data_sheet(
        wb,
        name="10_CompoundCost",
        title="PortableEquipCompound：合成金币单列补丁（条件粘贴）",
        target="SpecialEquipment/Portable/PortableEquipCompound.xlsx / PortableEquipCompound / F5:F44",
        instruction="Range与EquipResult闭环后，只复制 B5:B44 → 正式表 F5:F44；A列Id仅用于对行",
        headers=["Id", "CostCurrency"],
        rows=cost_rows,
        status="CONDITIONAL",
        copy_columns=[2],
    )
    quality_ws = add_data_sheet(
        wb,
        name="11_QualityPatch",
        title="SpecialEquipmentQuality：字段补丁（条件粘贴）",
        target="SpecialEquipment/SpecialEquipmentQuality.xlsx / SpecialEquipment / E列、M:N列",
        instruction="旧运行时退场后：复制 C5:C9 → E5:E9；复制 D5:E9 → M5:N9；空白代表清空",
        headers=["Id", "Quality", "ExtraItemList", "GradeCoefficient", "QualityCoefficient"],
        rows=quality_rows,
        status="CONDITIONAL",
        copy_columns=[3, 4, 5],
    )
    for row_number in range(5, 10):
        quality_ws.cell(row_number, 4).number_format = "0.00"
        quality_ws.cell(row_number, 5).number_format = "0.00"
    add_blockers(wb)
    add_validation_sheet(wb)
    add_meta(
        wb,
        generated_at=generated_at,
        config_root=config_root,
        records=records,
        source_snapshot_hash=source_snapshot_hash,
    )

    try:
        wb.calculation.calcMode = "auto"
        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True
    except Exception:
        pass
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    wb.close()

    report = {
        "package_id": PACKAGE_ID,
        "package_version": PACKAGE_VERSION,
        "generated_at": generated_at,
        "workbook": str(output_path),
        "workbook_sha256_before_excel_recalc": sha256_file(output_path),
        "official_config_root": str(config_root),
        "official_target_hashes": records,
        "source_snapshot_hash": source_snapshot_hash,
        "official_config_write": False,
        "luban_executed": False,
        "weight_scale": WEIGHT_SCALE,
        "counts": {
            "direct_ready_rows": len(exp_rows) + len(sub_rows) + len(level_rows),
            "probability_conditional_rows": len(compound_rows) + len(craft_rows),
            "conditional_patch_rows": len(cost_rows) + len(quality_rows),
            "portable_exp": len(exp_rows),
            "sub_attr_coef": len(sub_rows),
            "portable_lv_up": len(level_rows),
            "compound_prob": len(compound_rows),
            "craft_prob": len(craft_rows),
            "compound_cost": len(cost_rows),
            "quality_patch": len(quality_rows),
        },
    }
    report_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    return report


def main() -> int:
    parser = argparse.ArgumentParser(description="生成映射装备正式配置粘贴包 P01（不写正式Datas）")
    parser.add_argument("--output", default=str(DEFAULT_OUTPUT))
    parser.add_argument("--config-root", default=str(DEFAULT_CONFIG_ROOT))
    parser.add_argument("--report", default=str(DEFAULT_REPORT))
    args = parser.parse_args()
    try:
        report = build_workbook(Path(args.output).resolve(), Path(args.config_root).resolve(), Path(args.report).resolve())
    except Exception as exc:
        print(f"[FAIL] {type(exc).__name__}: {exc}", file=sys.stderr)
        return 1
    print("[OK] 映射装备正式配置粘贴包 P01 已生成；正式 Datas 未写入")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
