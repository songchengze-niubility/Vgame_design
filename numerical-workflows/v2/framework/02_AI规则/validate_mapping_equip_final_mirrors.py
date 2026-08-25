#!/usr/bin/env python3
"""只读校验 06_最终配置。"""
import hashlib, json, os
from pathlib import Path
from openpyxl import load_workbook

ROOT=Path(os.environ.get("VGAME_NUMERICAL_RUNTIME_ROOT", str(Path(__file__).resolve().parents[1])))
OUT=Path(os.environ.get("MAPPING_EQUIP_FINAL_OUT", str(ROOT/"06_最终配置"/"映射装备")))
CONFIG=Path(os.environ.get("VGAME_CONFIG_DATAS", ""))
GEN=ROOT/"05_验算核算"/"mapping-equip-final-mirrors-F01.json"
REPORT=ROOT/"05_验算核算"/"mapping-equip-final-mirrors-F01-validation.json"

MAIN_TARGET_BY_PART={
    "SEPart_Weapon":("Attack",11250),"SEPart_Head":("Hp",18750),"SEPart_Armour":("Hp",18750),
    "SEPart_Glove":("Armor",2500),"SEPart_Belt":("Armor",2500),"SEPart_Shoe":("Armor",2500),
    "SEPart_Device":("CritDmgBonusLevel",8800),"SEPart_Decoration":("CritLevel",7200),
}
SUB_VALUE_RANGE={
    "AttackBonus":(.329,.0987,.001),"ArmorBonus":(.329,.0987,.001),"HpBonus":(.329,.0987,.001),
    "Attack":(3701,1110,1),"Armor":(2468,740,1),"Hp":(12338,3701,1),
    "CritLevel":(2961,888,1),"CritDmgBonusLevel":(3619,1086,1),
    "EnergyPreAddBonusLevel":(2961,888,1),"CoolDownReduceLevel":(4277,1283,1),
}
ORDINARY_EFFECT_IDS={offset+slot for offset in (0,100,200,300) for slot in range(1,13)}
REMOVED_DAMAGE_ATTRS={"PhysicsDmgBonus","MagicDmgBonus"}
QUALITY_EXPECTED={
    "Green":(2,2,.2,.5),"Blue":(2,3,.3,.6),"Purple":(4,4,.45,.7),
    "Orange":(4,4,.7,.85),"Red":(4,4,1.0,1.0),
}
SUB_LEVEL_EXPECTED={5:.400,10:.455,15:.509,20:.564,25:.618,30:.673,35:.727,40:.782,45:.836,50:.891,55:.945,60:1.000}

def main():
    generation=json.loads(GEN.read_text(encoding="utf-8"))
    checks=[]
    def check(name,ok,actual,expected): checks.append({"name":name,"result":"PASS" if ok else "FAIL","actual":actual,"expected":expected})
    files=sorted(p for p in OUT.glob("*.xlsx") if not p.name.startswith("~$"))
    check("文件总数",len(files)==generation["files"],len(files),generation["files"])
    expected={r["file"]:r for r in generation["records"]}
    for fn,r in expected.items():
        p=OUT/fn; check(f"存在:{fn}",p.exists(),p.exists(),True)
        if not p.exists(): continue
        wb=load_workbook(p,read_only=True,data_only=False)
        check(f"说明页:{fn}","中文说明" in wb.sheetnames,wb.sheetnames,"含中文说明")
        check(f"配置页:{fn}",r["sheet"] in wb.sheetnames,wb.sheetnames,r["sheet"])
        check(f"页签顺序:{fn}",wb.sheetnames[:2]==[r["sheet"],"中文说明"],wb.sheetnames[:2],[r["sheet"],"中文说明"])
        if r["sheet"] in wb.sheetnames:
            ws=wb[r["sheet"]]
            actual=sum(1 for row in ws.iter_rows(min_row=5,values_only=True) if any(v is not None for v in row))
            check(f"行数:{fn}",actual==r["rows"],actual,r["rows"])
            errors=sum(1 for row in ws.iter_rows(values_only=True) for v in row if isinstance(v,str) and v.startswith(("#REF!","#VALUE!","#DIV/0!","#NAME?","#N/A")))
            check(f"错误值:{fn}",errors==0,errors,0)
        wb.close()
    # 映射装备属性原子配置与核心引用闭环
    sw=load_workbook(OUT/"SpecialEquipment.xlsx",read_only=True,data_only=True); ss=sw["SpecialEquipment"]
    equip_headers=next(ss.iter_rows(min_row=1,max_row=1,values_only=True)); equip_col={v:i for i,v in enumerate(equip_headers) if v is not None}
    equips=[list(row) for row in ss.iter_rows(min_row=5,values_only=True) if row[equip_col["Id"]] is not None]
    mw=load_workbook(OUT/"SpecialEquipmentMain.xlsx",read_only=True,data_only=True); ms=mw["SpecialEquipment"]
    main_headers=next(ms.iter_rows(min_row=1,max_row=1,values_only=True)); main_col={v:i for i,v in enumerate(main_headers) if v is not None}
    main_rows=[row for row in ms.iter_rows(min_row=5,values_only=True) if row[main_col["Id"]] is not None]
    mains={row[main_col["Id"]] for row in main_rows}
    dw=load_workbook(OUT/"SpecialEquipmentEffectDrop.xlsx",read_only=True,data_only=True); ds=dw["SpecialEquipment"]
    pool_headers=next(ds.iter_rows(min_row=1,max_row=1,values_only=True)); pool_col={v:i for i,v in enumerate(pool_headers) if v is not None}
    pools={row[pool_col["Id"]]:row[pool_col["EffectList"]] for row in ds.iter_rows(min_row=5,values_only=True) if row[pool_col["Id"]] is not None}
    ew=load_workbook(OUT/"SpecialEquipmentEffect.xlsx",read_only=True,data_only=True); es=ew["SpecialEquipment"]
    effect_headers=next(es.iter_rows(min_row=1,max_row=1,values_only=True)); effect_col={v:i for i,v in enumerate(effect_headers) if v is not None}
    effect_rows=[row for row in es.iter_rows(min_row=5,values_only=True) if row[effect_col["Id"]] is not None]
    effect_by_id={row[effect_col["Id"]]:row for row in effect_rows}; effects=set(effect_by_id)
    qw=load_workbook(OUT/"SpecialEquipmentQuality.xlsx",read_only=True,data_only=True); qs=qw["SpecialEquipment"]
    quality_headers=next(qs.iter_rows(min_row=1,max_row=1,values_only=True)); quality_col={v:i for i,v in enumerate(quality_headers) if v is not None}
    quality_rows=[row for row in qs.iter_rows(min_row=5,values_only=True) if row[quality_col["Id"]] is not None]
    scw=load_workbook(OUT/"SpecialEquipmentSubAttrCoef.xlsx",read_only=True,data_only=True); scs=scw["SpecialEquipmentSubAttrCoef"]
    sub_headers=next(scs.iter_rows(min_row=1,max_row=1,values_only=True)); sub_col={v:i for i,v in enumerate(sub_headers) if v is not None}
    sub_rows=[row for row in scs.iter_rows(min_row=5,values_only=True) if row[sub_col["Id"]] is not None]
    iw=load_workbook(OUT/"Items.xlsx",read_only=True,data_only=True)
    output_item_rows=[row for row in iw["items"].iter_rows(min_row=5,values_only=True) if row[0] is not None]
    output_item_ids={row[0] for row in output_item_rows}
    output_items_by_name={row[1]:row[0] for row in output_item_rows if isinstance(row[1],str)}
    fiw=load_workbook(CONFIG/"items/Items.xlsx",read_only=True,data_only=True)
    formal_item_ids={row[1] for row in fiw["items"].iter_rows(min_row=5,values_only=True) if row[1] is not None}
    formal_items_by_name={row[2]:row[1] for row in fiw["items"].iter_rows(min_row=5,values_only=True) if row[1] is not None and isinstance(row[2],str)}
    fsw=load_workbook(CONFIG/"SpecialEquipment/SpecialEquipment.xlsx",read_only=True,data_only=True); fss=fsw["SpecialEquipment"]
    formal_headers=next(fss.iter_rows(min_row=1,max_row=1,values_only=True)); formal_equip_col={v:i for i,v in enumerate(formal_headers) if v is not None}
    formal_equips={row[formal_equip_col["Id"]]:row for row in fss.iter_rows(min_row=5,values_only=True) if row[formal_equip_col["Id"]] is not None}
    equip_ids={r[equip_col["Id"]] for r in equips}; main_refs={int(str(r[equip_col["Attribute1List"]]).split(',')[0]) for r in equips}; pool_refs={r[equip_col["Effect"]] for r in equips}
    effect_refs={int(pair.split(',')[0]) for value in pools.values() for pair in str(value).split(';') if pair}
    check("正式装备190条",len(equips)==190,len(equips),190)
    check("正式装备ID全集保留",equip_ids==set(formal_equips),sorted(equip_ids^set(formal_equips)),[])
    preserved_fields=[field for field in formal_equip_col if field!="Attribute2List"]
    changed_unexpected=[]
    for row in equips:
        equip_id=row[equip_col["Id"]]; formal_row=formal_equips[equip_id]
        if any(row[equip_col[field]]!=formal_row[formal_equip_col[field]] for field in preserved_fields): changed_unexpected.append(equip_id)
    check("正式非主属性字段逐格保留",not changed_unexpected,changed_unexpected,[])
    check("装备Items正式闭环",equip_ids<=formal_item_ids,len(equip_ids-formal_item_ids),0)
    material_names={name for name,item_id in output_items_by_name.items() if not 4351<=item_id<=4372 and (name.endswith("主材") or name.startswith("映射装备副材"))}
    # 正式 Items 尚未写入本模块新材料时，以最终镜像为待落表新增；若正式表已有同名材料，则 ID 必须一致。
    material_name_mismatches=[(name,output_items_by_name[name],formal_items_by_name[name]) for name in sorted(material_names) if name in formal_items_by_name and output_items_by_name[name]!=formal_items_by_name[name]]
    pending_material_names=[name for name in sorted(material_names) if name not in formal_items_by_name]
    check("材料Items名称和ID可落表",len(material_names)==76 and not material_name_mismatches,{"count":len(material_names),"mismatches":material_name_mismatches,"pending_formal_insert":pending_material_names},{"count":76,"mismatches":[],"pending_formal_insert":"允许新增材料由最终镜像落表"})
    preview_item_ids=set(range(4351,4373))
    check("UI预览占位Items闭环",preview_item_ids<=output_item_ids and preview_item_ids<=formal_item_ids,
          {"missing_output":sorted(preview_item_ids-output_item_ids),"missing_formal":sorted(preview_item_ids-formal_item_ids)},
          {"missing_output":[],"missing_formal":[]})
    check("主属性引用闭环",main_refs<=mains,sorted(main_refs-mains),[])
    formal_main_refs={int(str(r[formal_equip_col["Attribute1List"]]).split(',')[0]) for r in formal_equips.values()}
    check("Main_190条独立参数",len(main_rows)==190 and len(mains)==190,(len(main_rows),len(mains)),(190,190))
    check("Main沿用正式独立ID",mains==formal_main_refs,sorted(mains^formal_main_refs),[])
    expected_main_attrs={v[0] for v in MAIN_TARGET_BY_PART.values()}
    check("Main属性集合",{r[main_col["Attribute"]] for r in main_rows}==expected_main_attrs,sorted({r[main_col["Attribute"]] for r in main_rows}),sorted(expected_main_attrs))
    main_row_by_id={r[main_col["Id"]]:r for r in main_rows}
    lv5_errors=[]; lv60_errors=[]
    main_mapping_errors=[]
    for equip in equips:
        main_id=int(str(equip[equip_col["Attribute1List"]]).split(',')[0]); r=main_row_by_id[main_id]
        attr,target=MAIN_TARGET_BY_PART[equip[equip_col["Part"]]]
        if r[main_col["Attribute"]]!=attr: main_mapping_errors.append((equip[equip_col["Id"]],equip[equip_col["Part"]],r[main_col["Attribute"]],attr))
        value,grow,param=r[main_col["Value"]],r[main_col["Grow"]],r[main_col["Param"]]
        lv5_errors.append(abs(value*grow+param-target*0.05))
        lv60_errors.append(abs(value*(grow**12)+param-target))
    check("Main_Lv5锚点",max(lv5_errors,default=999)<1e-5,max(lv5_errors,default=None),"<1e-5")
    check("Main_Lv60锚点",max(lv60_errors,default=999)<1e-5,max(lv60_errors,default=None),"<1e-5")
    check("副属性池引用闭环",pool_refs<=set(pools),sorted(pool_refs-set(pools)),[])
    check("副属性Effect引用闭环",effect_refs<=effects,sorted(effect_refs-effects),[])
    check("每件装备保留独立Main引用",all(r[equip_col["Attribute1List"]]==formal_equips[r[equip_col["Id"]]][formal_equip_col["Attribute1List"]] for r in equips),[(r[equip_col["Id"]],r[equip_col["Attribute1List"]],formal_equips[r[equip_col["Id"]]][formal_equip_col["Attribute1List"]]) for r in equips if r[equip_col["Attribute1List"]]!=formal_equips[r[equip_col["Id"]]][formal_equip_col["Attribute1List"]]],"保留正式Attribute1List")
    check("独立Main按部位配置属性",not main_mapping_errors,main_mapping_errors,[])
    check("单主属性规则",all(r[equip_col["Attribute1List"]] and r[equip_col["Attribute2List"]] is None for r in equips),sum(not(r[equip_col["Attribute1List"]] and r[equip_col["Attribute2List"]] is None) for r in equips),0)

    effect_range_errors=[]
    for effect_id in sorted(ORDINARY_EFFECT_IDS):
        row=effect_by_id.get(effect_id)
        if row is None: effect_range_errors.append((effect_id,"missing")); continue
        attr=row[effect_col["Attribute"]]; expected_range=SUB_VALUE_RANGE.get(attr)
        actual_range=(row[effect_col["ValueMax"]],row[effect_col["ValueMin"]],row[effect_col["RandomMin"]])
        if expected_range is None or any(abs(float(a)-float(e))>1e-9 for a,e in zip(actual_range,expected_range)):
            effect_range_errors.append((effect_id,attr,actual_range,expected_range))
    check("十类副属性随机区间",not effect_range_errors,effect_range_errors,[])
    pool_semantic_errors=[]
    for pool_id in sorted(pool_refs):
        ids=[int(token.split(',')[0]) for token in str(pools[pool_id]).split(';') if token]
        attrs=[effect_by_id[i][effect_col["Attribute"]] for i in ids if effect_by_id[i][effect_col["Attribute"]] is not None]
        if any(attr in REMOVED_DAMAGE_ATTRS or attr not in SUB_VALUE_RANGE for attr in attrs) or len(attrs)!=len(set(attrs)):
            pool_semantic_errors.append((pool_id,ids,attrs))
    check("在用副属性池十类且语义不重复",not pool_semantic_errors,pool_semantic_errors,[])
    generic_pool_attrs={pid:{effect_by_id[int(token.split(',')[0])][effect_col["Attribute"]] for token in str(pools[pid]).split(';') if token} for pid in (101,201,301)}
    check("通用品质池各含十类词条",all(attrs==set(SUB_VALUE_RANGE) for attrs in generic_pool_attrs.values()),{k:sorted(v) for k,v in generic_pool_attrs.items()},sorted(SUB_VALUE_RANGE))
    red_pool_refs={row[equip_col["Effect"]] for row in equips if row[equip_col["Quality"]]=="Red"}
    red_pool_errors=[]
    for pool_id in sorted(red_pool_refs):
        ids=[int(token.split(',')[0]) for token in str(pools[pool_id]).split(';') if token]
        attrs=[effect_by_id[i][effect_col["Attribute"]] for i in ids]
        if len(ids)!=10 or set(attrs)!=set(SUB_VALUE_RANGE) or any(attr is None for attr in attrs):
            red_pool_errors.append((pool_id,ids,attrs))
    check("红装37池取消专属效果并从十类随机4条",len(red_pool_refs)==37 and not red_pool_errors,{"pool_count":len(red_pool_refs),"errors":red_pool_errors},{"pool_count":37,"errors":[]})

    quality_actual={row[quality_col["Quality"]]:(row[quality_col["EffectNumMin"]],row[quality_col["EffectNumMax"]],row[quality_col["GradeCoefficient"]],row[quality_col["QualityCoefficient"]]) for row in quality_rows}
    check("品质系数与副属性条数",quality_actual==QUALITY_EXPECTED,quality_actual,QUALITY_EXPECTED)
    sub_actual={row[sub_col["EquipLv"]]:row[sub_col["SubAttrCoef"]] for row in sub_rows}
    check("副属性12档等级倍率",sub_actual==SUB_LEVEL_EXPECTED,sub_actual,SUB_LEVEL_EXPECTED)
    for w in (sw,mw,dw,ew,qw,scw,iw,fiw,fsw): w.close()
    rw=load_workbook(OUT/"PortableRange.xlsx",read_only=True,data_only=True); ranges=[row for row in rw["PortableRange"].iter_rows(min_row=5,values_only=True) if row[1] is not None]
    range_record=expected["PortableRange.xlsx"]
    expected_ranges=[(1,10,10,1190,5),(2,15,20,1191,10),(3,25,30,1192,20),(4,35,40,1193,30),(5,40,45,1194,35),(6,45,50,1195,40),(7,50,55,1196,45),(8,55,60,1197,50)]
    actual_ranges=[(r[1],r[2],r[3],r[4],r[5]) for r in ranges]
    check("PortableRange八档已确认",actual_ranges==expected_ranges,actual_ranges,expected_ranges)
    check("PortableRange投入等级规则",all(r[5]==r[2]-5 for r in ranges),[(r[1],r[2],r[5]) for r in ranges],"EquipMinLv=LevelMin-5")
    check("PortableRange可粘贴状态",range_record["status"].startswith("可粘贴"),range_record["status"],"可粘贴")
    rw.close()

    # 合成/打造：正式Schema、品质概率、等级Generate和消耗引用闭环。
    schema_cases=[
        ("SpecialEquipment.xlsx","SpecialEquipment","SpecialEquipment/SpecialEquipment.xlsx"),
        ("SpecialEquipmentMain.xlsx","SpecialEquipment","SpecialEquipment/SpecialEquipmentMain.xlsx"),
        ("SpecialEquipmentEffect.xlsx","SpecialEquipment","SpecialEquipment/SpecialEquipmentEffect.xlsx"),
        ("SpecialEquipmentEffectDrop.xlsx","SpecialEquipment","SpecialEquipment/SpecialEquipmentEffectDrop.xlsx"),
        ("SpecialEquipmentQuality.xlsx","SpecialEquipment","SpecialEquipment/SpecialEquipmentQuality.xlsx"),
        ("SpecialEquipmentSubAttrCoef.xlsx","SpecialEquipmentSubAttrCoef","SpecialEquipment/SpecialEquipmentSubAttrCoef.xlsx"),
        ("PortableCompoundProb.xlsx","PortableCompoundProb","SpecialEquipment/Portable/PortableCompoundProb.xlsx"),
        ("PortableCraftProb.xlsx","PortableCompoundProb","SpecialEquipment/Portable/PortableCraftProb.xlsx"),
        ("PortableEquipCompound.xlsx","PortableEquipCompound","SpecialEquipment/Portable/PortableEquipCompound.xlsx"),
        ("PortableEquipCraft.xlsx","PortableEquipCraft","SpecialEquipment/Portable/PortableEquipCraft.xlsx"),
        ("SpecialEquipmentGenerate.xlsx","SpecialEquipment","SpecialEquipment/SpecialEquipmentGenerate.xlsx"),
        ("Unlock.xlsx","Unlock","unlock/Unlock.xlsx"),
        ("SystemOpen.xlsx","SystemOpen","system_open/SystemOpen.xlsx"),
        ("UIlevel.xlsx","level","level/UIlevel.xlsx"),
    ]
    for file_name,sheet_name,formal_rel in schema_cases:
        ow=load_workbook(OUT/file_name,read_only=True,data_only=False); osheet=ow[sheet_name]
        fw=load_workbook(CONFIG/formal_rel,read_only=True,data_only=False); fsheet=fw[sheet_name]
        width=max(c for c in range(1,40) if any(fsheet.cell(r,c).value is not None for r in range(1,5)))
        actual=[[osheet.cell(r,c).value for c in range(1,width+1)] for r in range(1,5)]
        expected_schema=[[fsheet.cell(r,c).value for c in range(1,width+1)] for r in range(1,5)]
        check(f"正式Schema:{file_name}",actual==expected_schema,actual,expected_schema)
        ow.close();fw.close()

    def prob_rows(file_name):
        w=load_workbook(OUT/file_name,read_only=True,data_only=True); s=w["PortableCompoundProb"]
        rows=[r for r in s.iter_rows(min_row=5,values_only=True) if r[1] is not None]; w.close(); return rows
    compound_probs=prob_rows("PortableCompoundProb.xlsx"); craft_probs=prob_rows("PortableCraftProb.xlsx")
    def weight_sum(value): return sum(int(x.split(',')[1]) for x in str(value).split(';') if x)
    check("合成品质概率300行",len(compound_probs)==300,len(compound_probs),300)
    check("打造品质概率180行",len(craft_probs)==180,len(craft_probs),180)
    check("合成品质权重和1000",all(weight_sum(r[4])==1000 for r in compound_probs),[(r[1],r[4]) for r in compound_probs if weight_sum(r[4])!=1000],[])
    check("打造品质权重和1000",all(weight_sum(r[4])==1000 for r in craft_probs),[(r[1],r[4]) for r in craft_probs if weight_sum(r[4])!=1000],[])

    gw=load_workbook(OUT/"SpecialEquipmentGenerate.xlsx",read_only=True,data_only=True); gs=gw["SpecialEquipment"]
    generate={r[1]:r for r in gs.iter_rows(min_row=5,values_only=True) if r[1] is not None}
    check("Generate总行数278",len(generate)==278,len(generate),278)
    expected_box_quality_lists={0:"Green,800;Blue,200",1:"Blue,850;Purple,150",2:"Purple,900;Orange,100",3:"Orange,980;Red,20",4:"Red,1000"}
    box_quality_mismatches=[(gid,generate[gid][5],expected_box_quality_lists[(gid-4201)%5]) for gid in range(4201,4351) if generate.get(gid) is None or generate[gid][5]!=expected_box_quality_lists[(gid-4201)%5]]
    check("宝箱开装品质权重",not box_quality_mismatches,box_quality_mismatches,[])
    check("合成Generate八档",set(range(4401,4409))<=set(generate),sorted(set(range(4401,4409))-set(generate)),[])
    check("打造Generate一百二十条",set(range(4501,4621))<=set(generate),len(set(range(4501,4621))-set(generate)),0)
    check("Portable等级池不重复配品质",all(generate[i][5] is None for i in list(range(4401,4409))+list(range(4501,4621))),[(i,generate[i][5]) for i in list(range(4401,4409))+list(range(4501,4621)) if generate[i][5] is not None],[])

    cw=load_workbook(OUT/"PortableEquipCompound.xlsx",read_only=True,data_only=True); cs=cw["PortableEquipCompound"]
    compound_rows=[r for r in cs.iter_rows(min_row=5,values_only=True) if r[1] is not None]
    check("合成配置可随PortableRange原子落表",expected["PortableEquipCompound.xlsx"]["status"].startswith("可粘贴"),expected["PortableEquipCompound.xlsx"]["status"],"可粘贴")
    check("合成40行引用档位Generate",all(r[4]==4400+r[2] for r in compound_rows),[(r[1],r[2],r[4]) for r in compound_rows if r[4]!=4400+r[2]],[])
    check("合成Generate引用闭环",{r[4] for r in compound_rows}<=set(generate),sorted({r[4] for r in compound_rows}-set(generate)),[])
    cw.close()

    uw=load_workbook(OUT/"Unlock.xlsx",read_only=True,data_only=True); us=uw["Unlock"]
    new_unlocks={r[1]:r for r in us.iter_rows(min_row=5,values_only=True) if r[1] is not None}
    fuw=load_workbook(CONFIG/"unlock/Unlock.xlsx",read_only=True,data_only=True); fus=fuw["Unlock"]
    formal_unlock_ids={r[1] for r in fus.iter_rows(min_row=5,values_only=True) if r[1] is not None}
    available_unlock_ids=formal_unlock_ids|set(new_unlocks)
    range_unlock_ids={r[4] for r in ranges}
    check("PortableRange解锁条件引用闭环",range_unlock_ids<=available_unlock_ids,sorted(range_unlock_ids-available_unlock_ids),[])
    expected_compound_unlocks={1190:10,1191:15,1192:25,1193:35,1194:40,1195:45,1196:50,1197:55}
    expected_new_unlocks={**expected_compound_unlocks,1198:250061}
    check("新增映射装备解锁条件完整",{unlock_id:row[4] for unlock_id,row in new_unlocks.items()}==expected_new_unlocks,{unlock_id:row[4] for unlock_id,row in new_unlocks.items()},expected_new_unlocks)
    check("映射装备UnlockId全部新增",set(expected_new_unlocks).isdisjoint(formal_unlock_ids),sorted(set(expected_new_unlocks)&formal_unlock_ids),[])
    check("合成解锁链保持不变",all(new_unlocks[unlock_id][2]=="映射装备合成等级解锁" and new_unlocks[unlock_id][3]=="UnLock_UUT_player_Grade" and new_unlocks[unlock_id][4]==grade for unlock_id,grade in expected_compound_unlocks.items()),[(unlock_id,new_unlocks.get(unlock_id)) for unlock_id in expected_compound_unlocks],"合成1190~1197保持原玩家等级条件")
    craft_unlock=new_unlocks.get(1198)
    check("打造解锁条件正确",craft_unlock is not None and craft_unlock[2]=="映射装备打造通关解锁" and craft_unlock[3]=="UnLock_UUT_Level" and craft_unlock[4]==250061,craft_unlock,"解锁ID1198 / UnLock_UUT_Level / LevelId250061")
    uw.close();fuw.close()

    sow=load_workbook(OUT/"SystemOpen.xlsx",read_only=True,data_only=True); sos=sow["SystemOpen"]
    fsow=load_workbook(CONFIG/"system_open/SystemOpen.xlsx",read_only=True,data_only=True); fsos=fsow["SystemOpen"]
    so_headers=next(sos.iter_rows(min_row=1,max_row=1,values_only=True)); fso_headers=next(fsos.iter_rows(min_row=1,max_row=1,values_only=True))
    so_cols={field:so_headers.index(field) for field in ("Id","##","UnlockId","Name","ParentId","UIViewName","JumpParams")}
    fso_cols={field:fso_headers.index(field) for field in ("Id","##","UnlockId","Name","ParentId","UIViewName","JumpParams")}
    system_rows=[r for r in sos.iter_rows(min_row=5,values_only=True) if r[so_cols["Id"]] is not None]
    formal_system_row=next(r for r in fsos.iter_rows(min_row=5,values_only=True) if r[fso_cols["Id"]]==54)
    system_row=system_rows[0] if len(system_rows)==1 else None
    check("打造SystemOpen单行",system_row is not None and system_row[so_cols["Id"]]==54,system_rows,"仅Id=54")
    check("打造SystemOpen引用专用解锁",system_row is not None and system_row[so_cols["UnlockId"]]==1198 and system_row[so_cols["##"]]=="通关映射装备第6层",None if system_row is None else (system_row[so_cols["UnlockId"]],system_row[so_cols["##"]]),(1198,"通关映射装备第6层"))
    preserved_fields=("Name","ParentId","UIViewName","JumpParams")
    check("打造SystemOpen关键字段保留",system_row is not None and all(system_row[so_cols[field]]==formal_system_row[fso_cols[field]] for field in preserved_fields),None if system_row is None else {field:system_row[so_cols[field]] for field in preserved_fields},{field:formal_system_row[fso_cols[field]] for field in preserved_fields})
    sow.close();fsow.close()

    crw=load_workbook(OUT/"PortableEquipCraft.xlsx",read_only=True,data_only=True); crs=crw["PortableEquipCraft"]
    craft_rows=[r for r in crs.iter_rows(min_row=5,values_only=True) if r[1] is not None]
    oiw=load_workbook(OUT/"Items.xlsx",read_only=True,data_only=True); output_item_ids=set()
    for r in oiw["items"].iter_rows(min_row=5,values_only=True):
        if r[0] is not None: output_item_ids.add(r[0])
    oiw.close()
    available_items=formal_item_ids|output_item_ids
    cost_item_ids=set()
    for r in craft_rows:
        for field in (r[4],r[5]):
            cost_item_ids.update(int(x.split(',')[0]) for x in str(field).split(';') if x)
    check("打造360行完整配置",len(craft_rows)==360,len(craft_rows),360)
    check("打造行无空业务字段",all(all(r[i] is not None for i in range(1,7)) for r in craft_rows),[r[1] for r in craft_rows if any(r[i] is None for i in range(1,7))],[])
    check("打造Generate引用闭环",{r[3] for r in craft_rows}<=set(generate),sorted({r[3] for r in craft_rows}-set(generate)),[])
    check("打造材料Items闭环",cost_item_ids<=available_items,sorted(cost_item_ids-available_items),[])
    old_material_ids=set(range(4101,4173))|set(range(4181,4185))
    check("打造消耗无旧材料ID",not(cost_item_ids&old_material_ids),sorted(cost_item_ids&old_material_ids),[])
    craft_grade_pools={
        1:"10,1000", 2:"15,700;20,300", 3:"25,700;30,300", 4:"35,700;40,300",
        5:"40,650;45,350", 6:"45,650;50,350", 7:"50,650;55,350", 8:"55,650;60,350",
    }
    craft_generate_tier={}
    for tier, ids in enumerate((range(4501,4509), range(4509,4525), range(4525,4541), range(4541,4557), range(4557,4573), range(4573,4589), range(4589,4605), range(4605,4621)), 1):
        craft_generate_tier.update({gid:tier for gid in ids})
    check("打造Generate按档等级概率池",all(str(generate[gid][4])==craft_grade_pools[tier] for gid,tier in craft_generate_tier.items()),[(gid,generate[gid][4],craft_grade_pools[tier]) for gid,tier in craft_generate_tier.items() if str(generate[gid][4])!=craft_grade_pools[tier]],[])
    check("打造Generate等级权重和1000",all(weight_sum(generate[gid][4])==1000 for gid in craft_generate_tier),[(gid,generate[gid][4]) for gid in craft_generate_tier if weight_sum(generate[gid][4])!=1000],[])
    crw.close();gw.close()

    ddw=load_workbook(OUT/"Drop_dungeon.xlsx",read_only=True,data_only=True); dds=ddw["Drop"]
    drop_rows=[row for row in dds.iter_rows(min_row=5,values_only=True) if isinstance(row[1],int)]
    check("Drop镜像678行",len(drop_rows)==678,len(drop_rows),678)
    drop_sort_keys=[(0 if row[1]>=10_000_000 else 1,row[1]) for row in drop_rows]
    check("Drop按一级后二级稳定排序",drop_sort_keys==sorted(drop_sort_keys),[(index+5,drop_rows[index][1]) for index in range(1,len(drop_rows)) if drop_sort_keys[index]<drop_sort_keys[index-1]],"8位一级DropId升序，随后5位二级DropId升序")
    drop_id_positions={}
    for index,row in enumerate(drop_rows): drop_id_positions.setdefault(row[1],[]).append(index)
    split_drop_ids=[drop_id for drop_id,positions in drop_id_positions.items() if positions!=list(range(min(positions),max(positions)+1))]
    check("同一DropId行连续",not split_drop_ids,split_drop_ids,[])
    drop_item_ids={row[9] for row in dds.iter_rows(min_row=5,values_only=True) if isinstance(row[9],int) and row[9]>0}
    check("Drop材料Items闭环",drop_item_ids<=available_items,sorted(drop_item_ids-available_items),[])
    check("Drop无旧材料ID",not(drop_item_ids&old_material_ids),sorted(drop_item_ids&old_material_ids),[])
    early_regular_material_refs=[(row[1],row[4]) for row in drop_rows if 20005101<=row[1]<=20005105 and ((isinstance(row[4],int) and 22301<=row[4]<=22305) or row[4]==22400)]
    early_first_material_items=[(row[1],row[9]) for row in drop_rows if 20000010<=row[1]<=20000014 and isinstance(row[9],int) and 1080<=row[9]<=1155]
    early_main_pools=sorted({row[1] for row in drop_rows if 22301<=row[1]<=22305})
    check("前5层常规无材料分支",not early_regular_material_refs,early_regular_material_refs,[])
    check("前5层首通无材料",not early_first_material_items,early_first_material_items,[])
    check("前5层不生成主材池",not early_main_pools,early_main_pools,[])
    layer6_regular_refs={row[4] for row in drop_rows if row[1]==20005106}
    layer6_first_materials=[row[9] for row in drop_rows if row[1]==20000015 and isinstance(row[9],int) and 1080<=row[9]<=1155]
    layer7_first_materials=[row[9] for row in drop_rows if row[1]==20000016 and isinstance(row[9],int) and 1080<=row[9]<=1155]
    transition_pool_items={drop_id:{row[9] for row in drop_rows if row[1]==drop_id and isinstance(row[9],int) and 1080<=row[9]<=1151} for drop_id in range(22306,22310)}
    check("第6层常规包含主副材分支",{22306,22400}<=layer6_regular_refs,sorted(layer6_regular_refs),"包含22306和22400")
    check("第6～8层常规为10档主材",all(transition_pool_items[drop_id]=={1080,1083,1086} for drop_id in range(22306,22309)),{drop_id:sorted(items) for drop_id,items in transition_pool_items.items() if drop_id<=22308},{22306:[1080,1083,1086],22307:[1080,1083,1086],22308:[1080,1083,1086]})
    check("第9层恢复20档主材",transition_pool_items[22309]=={1089,1092,1095},sorted(transition_pool_items[22309]),[1089,1092,1095])
    check("第6层首通为10档武器紫主材",layer6_first_materials==[1080],layer6_first_materials,[1080])
    check("第7层首通为10档防具紫主材",layer7_first_materials==[1083],layer7_first_materials,[1083])
    box_pool_items={drop_id:{row[9] for row in drop_rows if row[1]==drop_id and isinstance(row[9],int) and 4201<=row[9]<=4350} for drop_id in range(22101,22131)}
    box_pool_mismatches=[]
    for layer in range(1,31):
        expected={4201+(layer-1)*5+quality_index for quality_index in range(5)}
        if box_pool_items[22100+layer]!=expected:
            box_pool_mismatches.append((layer,sorted(box_pool_items[22100+layer]),sorted(expected)))
    check("每层宝箱品质池包含绿蓝紫金红",not box_pool_mismatches,box_pool_mismatches,[])
    layer1_box_weights={row[9]:row[12] for row in drop_rows if row[1]==22101 and isinstance(row[9],int)}
    check("第1层箱品质权重",layer1_box_weights=={4201:5150,4202:3000,4203:1850,4204:0,4205:0},layer1_box_weights,{4201:5150,4202:3000,4203:1850,4204:0,4205:0})
    layer10_box_weights={row[9]:row[12] for row in drop_rows if row[1]==22110 and isinstance(row[9],int)}
    check("第10层金箱首次出现",layer10_box_weights.get(4249)==100 and layer10_box_weights.get(4250)==0,layer10_box_weights,{4249:100,4250:0})
    layer20_box_weights={row[9]:row[12] for row in drop_rows if row[1]==22120 and isinstance(row[9],int)}
    check("第20层红箱首次出现",layer20_box_weights.get(4299)==1000 and layer20_box_weights.get(4300)==50,layer20_box_weights,{4299:1000,4300:50})
    layer30_box_weights={row[9]:row[12] for row in drop_rows if row[1]==22130 and isinstance(row[9],int)}
    check("第30层箱品质权重",layer30_box_weights=={4346:1200,4347:2800,4348:3700,4349:1900,4350:400},layer30_box_weights,{4346:1200,4347:2800,4348:3700,4349:1900,4350:400})
    ddw.close()

    uiw=load_workbook(OUT/"UIlevel.xlsx",read_only=True,data_only=True); uis=uiw["level"]
    ui_headers=next(uis.iter_rows(min_row=1,max_row=1,values_only=True))
    ui_cols={field:ui_headers.index(field) for field in ("LevelId","LevelFPDrop","LevelPPDrop","Reward","FirstReward","FirstRewardDrop","FallRewardDrop")}
    ui_rows=[row for row in uis.iter_rows(min_row=5,values_only=True) if row[ui_cols["LevelId"]] is not None]
    ui_item_ids=set()
    for row in ui_rows:
        for field in (row[ui_cols["Reward"]],row[ui_cols["FirstReward"]]):
            ui_item_ids.update(int(x.split(',')[0]) for x in str(field).split(';') if x and x.split(',')[0].isdigit())
    check("UI预览Items正式闭环",ui_item_ids<=formal_item_ids,sorted(ui_item_ids-formal_item_ids),[])
    check("UI预览无旧材料ID",not(ui_item_ids&old_material_ids),sorted(ui_item_ids&old_material_ids),[])
    check("映射UIlevel150行",len(ui_rows)==150,len(ui_rows),150)
    craft_tier_by_layer={**{layer:10 for layer in range(6,9)},9:20,**{layer:30 for layer in range(10,16)},**{layer:40 for layer in range(16,22)},**{layer:45 for layer in range(22,25)},**{layer:50 for layer in range(25,27)},**{layer:55 for layer in range(27,29)},**{layer:60 for layer in range(29,31)}}
    preview_grade_box_ids={5:4351,10:4352,15:4353,20:4354,25:4365,30:4366,35:4367,40:4368,45:4369,50:4370,55:4371,60:4372}
    grade_pools={
        **{layer:[5,10] for layer in range(1,5)},
        **{layer:[5,10,15] for layer in range(5,8)},
        **{layer:[10,15,20] for layer in range(8,10)},
        **{layer:[15,20,25] for layer in range(10,12)},
        **{layer:[20,25,30] for layer in range(12,14)},
        **{layer:[25,30,35] for layer in range(14,16)},
        **{layer:[30,35,40] for layer in range(16,18)},
        **{layer:[35,40,45] for layer in range(18,20)},
        **{layer:[40,45,50] for layer in range(20,22)},
        **{layer:[45,50,55] for layer in range(22,24)},
        **{layer:[50,55,60] for layer in range(24,27)},
        **{layer:[55,60] for layer in range(27,31)},
    }
    preview_main_ids={(10,"Purple"):4355,(20,"Purple"):4356,(30,"Purple"):4357,(30,"Orange"):4358,(40,"Orange"):4359,(45,"Orange"):4360,(50,"Red"):4361,(55,"Red"):4362,(60,"Red"):4363}
    preview_mismatches=[]
    for row in ui_rows:
        level_id=row[ui_cols["LevelId"]]
        layer=(level_id-250000)//10
        actual_items=[int(part.split(",")[0]) for part in str(row[ui_cols["Reward"]] or "").split(";") if part]
        expected_items=[preview_grade_box_ids[grade] for grade in grade_pools[layer]]
        if layer>=6:
            material_quality="Purple" if layer<=10 else "Orange" if layer<=24 else "Red"
            expected_items.extend([preview_main_ids[(craft_tier_by_layer[layer],material_quality)],4364])
        if actual_items!=expected_items:
            preview_mismatches.append((level_id,actual_items,expected_items))
    check("UI普通奖励按装备等级池逐级展示",not preview_mismatches,preview_mismatches,[])
    concrete_reward_proxies=[]
    for row in ui_rows:
        for part in str(row[ui_cols["Reward"]] or "").split(";"):
            if not part:
                continue
            item_id=int(part.split(",")[0])
            if 1080<=item_id<=1155 or 4201<=item_id<=4350:
                concrete_reward_proxies.append((row[ui_cols["LevelId"]],item_id))
    check("UI普通奖励不使用具体掉落代理",not concrete_reward_proxies,concrete_reward_proxies,[])
    first_box_layers=set(range(1,31))-{7,11,16,21,26}
    grade_box_item_ids=set(preview_grade_box_ids.values())
    first_grade_mismatches=[]
    first_concrete_box_proxies=[]
    for row in ui_rows:
        level_id=row[ui_cols["LevelId"]]
        layer=(level_id-250000)//10
        first_parts=[part for part in str(row[ui_cols["FirstReward"]] or "").split(";") if part]
        first_items=[int(part.split(",")[0]) for part in first_parts]
        actual_grade_items=[item_id for item_id in first_items if item_id in grade_box_item_ids]
        expected_grade_items=[preview_grade_box_ids[grade] for grade in grade_pools[layer]] if layer in first_box_layers else []
        if actual_grade_items!=expected_grade_items:
            first_grade_mismatches.append((level_id,actual_grade_items,expected_grade_items))
        first_concrete_box_proxies.extend((level_id,item_id) for item_id in first_items if 4201<=item_id<=4350)
    check("UI首通宝箱按装备等级池逐级展示",not first_grade_mismatches,first_grade_mismatches,[])
    check("UI首通不再使用具体品质箱代理",not first_concrete_box_proxies,first_concrete_box_proxies,[])
    early_ui_material_items=[]
    for row in ui_rows:
        level_id=row[ui_cols["LevelId"]]
        if not isinstance(level_id,int) or not (1<=((level_id-250000)//10)<=5):
            continue
        for field in ("Reward","FirstReward"):
            for part in str(row[ui_cols[field]] or "").split(";"):
                if part and part.split(",")[0].isdigit():
                    item_id=int(part.split(",")[0])
                    if 1080<=item_id<=1155:
                        early_ui_material_items.append((level_id,field,item_id))
    check("前5层UI预览无材料",not early_ui_material_items,early_ui_material_items,[])
    layer6_ui_material_items=[]
    for row in ui_rows:
        level_id=row[ui_cols["LevelId"]]
        if not isinstance(level_id,int) or ((level_id-250000)//10)!=6:
            continue
        for field in ("Reward","FirstReward"):
            for part in str(row[ui_cols[field]] or "").split(";"):
                if part and part.split(",")[0].isdigit():
                    item_id=int(part.split(",")[0])
                    if 1080<=item_id<=1155:
                        layer6_ui_material_items.append((level_id,field,item_id))
    check("第6层UI预览包含材料",bool(layer6_ui_material_items),layer6_ui_material_items,"常规主副材+首通主材")
    preview_quantity_errors=[]
    for row in ui_rows:
        for field in ("Reward","FirstReward"):
            for part in str(row[ui_cols[field]] or "").split(";"):
                if not part:
                    continue
                item_id,quantity=(int(x) for x in part.split(","))
                if quantity < 0 and quantity != -1:
                    preview_quantity_errors.append((row[ui_cols["LevelId"]],field,item_id,quantity))
    check("UI无数量展示统一为负一",not preview_quantity_errors,preview_quantity_errors,[])
    fuiw=load_workbook(CONFIG/"level/UIlevel.xlsx",read_only=True,data_only=True); fuis=fuiw["level"]
    formal_ui_headers=next(fuis.iter_rows(min_row=1,max_row=1,values_only=True)); formal_ui_cols={field:formal_ui_headers.index(field) for field in ("LevelId","LevelFPDrop","LevelPPDrop")}
    formal_display_by_level={row[formal_ui_cols["LevelId"]]:(row[formal_ui_cols["LevelFPDrop"]],row[formal_ui_cols["LevelPPDrop"]]) for row in fuis.iter_rows(min_row=5,values_only=True) if row[formal_ui_cols["LevelId"]] is not None}
    display_mismatches=[(row[ui_cols["LevelId"]],(row[ui_cols["LevelFPDrop"]],row[ui_cols["LevelPPDrop"]]),formal_display_by_level.get(row[ui_cols["LevelId"]])) for row in ui_rows if (row[ui_cols["LevelFPDrop"]],row[ui_cols["LevelPPDrop"]])!=formal_display_by_level.get(row[ui_cols["LevelId"]])]
    check("LevelFPDrop和LevelPPDrop保持正式源表原值",not display_mismatches,display_mismatches,[])
    fuiw.close()
    uiw.close()
    check("正式Datas零写入",generation["formal_datas_unchanged"] is True,generation["formal_datas_unchanged"],True)
    result={"summary":{"checks":len(checks),"pass":sum(c["result"]=="PASS" for c in checks),"fail":sum(c["result"]=="FAIL" for c in checks)},"checks":checks}
    REPORT.write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding="utf-8")
    print(json.dumps(result["summary"],ensure_ascii=False))
    raise SystemExit(1 if result["summary"]["fail"] else 0)

if __name__=="__main__": main()
