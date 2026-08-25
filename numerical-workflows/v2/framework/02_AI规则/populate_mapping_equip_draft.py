#!/usr/bin/env python3
"""生成首个映射装备配置草案批次，但绝不写入 Vgame Datas。

批次：MAPEQUIP-20260710-D01

本脚本只修改 ``00_总览/Vgame配置映射总表.xlsx`` 的人工维护区：

1. 更新映射装备相关 ``模块映射``；
2. 幂等写入字段级 ``配置草案``；
3. 更新工作簿阶段与门禁提示；
4. 备份、临时文件重开、原子替换并输出机器报告。

正式配置、导出物、审批记录和落地回执始终保持只读/空白。
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import shutil
import sys
import uuid
from collections import Counter
from copy import copy
from datetime import datetime
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


BATCH_ID = "MAPEQUIP-20260710-D01"
BATCH_VERSION = "1.0"
GENERATION_MARKER = "vgame-config-mapping/v1"

FRAME_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_WORKBOOK = FRAME_ROOT / "00_总览" / "Vgame配置映射总表.xlsx"
DEFAULT_CONFIG_ROOT = Path(
    os.environ.get("VGAME_CONFIG_DATAS", "")
)
DEFAULT_REPORT = FRAME_ROOT / "05_验算核算" / "mapping-equip-draft-D01-report.json"

MD_ATTR = "03_沉淀MD/养成/映射装备养成数值沉淀-2026-07-07.md"
MD_ECON = "03_沉淀MD/经济/21_养成_映射装备经济消耗沉淀-2026-07-09.md"
MD_GEN = "03_沉淀MD/养成/映射装备数值总表生成参数-2026-07-10.md"
MD_TIER = "03_沉淀MD/养成/映射装备合成制作8档调整-2026-07-06.md"

DRAFT_HEADERS = [
    "批次ID", "草案ID", "批次版本", "模块", "参数ID", "来源MD", "来源章节",
    "逻辑表", "源工作簿", "Sheet", "主键字段", "主键值", "目标字段", "RawType",
    "当前值", "建议值", "变换公式", "单位与舍入", "Bean或Enum约束", "引用影响",
    "公式单元格", "写入方式", "风险等级", "校验方法", "草案状态", "创建时间", "备注",
]

INPUT_FONT = Font(name="微软雅黑", size=10, color="0000FF")
PASS_FILL = PatternFill("solid", fgColor="E2F0D9")
WARN_FILL = PatternFill("solid", fgColor="FFF2CC")
INFO_FILL = PatternFill("solid", fgColor="DDEBF7")
FAIL_FILL = PatternFill("solid", fgColor="F4CCCC")


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def display_current(value: Any, *, missing: bool = False) -> Any:
    if missing:
        return "<NOT_EXIST>"
    if value is None or value == "":
        return "<BLANK>"
    return value


def round10(value: Decimal) -> int:
    return int((value / Decimal(10)).quantize(Decimal("1"), rounding=ROUND_HALF_UP) * 10)


FORGE_LEVEL_COSTS = (
    10, 20, 30, 40, 50, 60, 70, 100, 120,
    130, 140, 160, 180, 190, 210, 220, 240, 250, 280,
    350, 400, 450, 500, 550, 650, 750, 850, 950, 1050,
    1130, 1190, 1250, 1310, 1370, 1430, 1490, 1550, 1610, 1670,
    2000, 2100, 2200, 2300, 2400, 2500, 2600, 2700, 2800, 2900,
    3800, 4000, 4200, 4400, 4600, 4800, 5000, 5200, 5400, 5600,
)


def forge_level_cost(target_level: int) -> int:
    if not 2 <= target_level <= 60:
        raise ValueError(f"锻造台目标等级超出范围: {target_level}")
    return FORGE_LEVEL_COSTS[target_level - 2]


def read_config_rows(config_root: Path, relative_path: str, sheet_name: str | None = None) -> list[dict[str, Any]]:
    """只读解析 Luban 横向 Schema 表的业务行。"""

    path = config_root / relative_path
    if not path.is_file():
        raise FileNotFoundError(path)
    workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        worksheet = workbook[sheet_name] if sheet_name else workbook[workbook.sheetnames[0]]
        header_row = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True))
        max_col = max(
            (index for index, value in enumerate(header_row, start=1) if value not in (None, "")),
            default=0,
        )
        headers = list(header_row[:max_col])
        rows: list[dict[str, Any]] = []
        for values in worksheet.iter_rows(min_row=5, max_col=max_col, values_only=True):
            record = {
                str(header): values[index]
                for index, header in enumerate(headers)
                if header not in (None, "", "##var")
            }
            if any(value not in (None, "") for value in record.values()):
                rows.append(record)
        return rows
    finally:
        workbook.close()


def add_draft(
    rows: list[list[Any]],
    *,
    draft_id: str,
    parameter_id: str,
    source_md: str,
    source_section: str,
    logical_table: str,
    source_workbook: str,
    sheet: str,
    key_field: str,
    key_value: Any,
    target_field: str,
    raw_type: str,
    current_value: Any,
    suggested_value: Any,
    transform: str = "N/A",
    unit_rounding: str = "N/A",
    constraint: str = "N/A",
    impact: str = "仅影响映射装备；正式落表前复查运行时消费与引用。",
    formula_cell: str = "否",
    write_mode: str = "CANDIDATE_SAFE_RECHECK",
    risk: str = "MEDIUM",
    validation: str = "复查主键唯一、RawType、引用闭合与Luban导出差异。",
    status: str = "READY_FOR_REVIEW",
    created_at: str,
    note: str = "直接配置候选；本批次仅草案，未写正式配置。",
) -> None:
    record = {
        "批次ID": BATCH_ID,
        "草案ID": draft_id,
        "批次版本": BATCH_VERSION,
        "模块": "映射装备",
        "参数ID": parameter_id,
        "来源MD": source_md,
        "来源章节": source_section,
        "逻辑表": logical_table,
        "源工作簿": source_workbook,
        "Sheet": sheet,
        "主键字段": key_field,
        "主键值": key_value,
        "目标字段": target_field,
        "RawType": raw_type,
        "当前值": current_value,
        "建议值": suggested_value,
        "变换公式": transform,
        "单位与舍入": unit_rounding,
        "Bean或Enum约束": constraint,
        "引用影响": impact,
        "公式单元格": formula_cell,
        "写入方式": write_mode,
        "风险等级": risk,
        "校验方法": validation,
        "草案状态": status,
        "创建时间": created_at,
        "备注": note,
    }
    rows.append([record[header] for header in DRAFT_HEADERS])


def build_draft_rows(config_root: Path, created_at: str) -> list[list[Any]]:
    """从权威 MD 口径与只读正式配置快照生成 D01。"""

    rows: list[list[Any]] = []

    # 1) PortableExp：当前唯一可完整按现有 60 行粒度重写的目标表。
    portable_exp_path = "SpecialEquipment/Portable/PortableExp.xlsx"
    portable_exp = read_config_rows(config_root, portable_exp_path, "PortableExp")
    exp_current = {(row["EquipLv"], row["Quality"]): row for row in portable_exp}
    qualities = ["Green", "Blue", "Purple", "Orange", "Red"]
    equip_exp = {
        5: [1, 2, 5, 12, 30], 10: [1, 2, 5, 12, 30],
        15: [2, 4, 10, 24, 60], 20: [2, 4, 10, 24, 60],
        25: [3, 6, 15, 36, 90], 30: [3, 6, 15, 36, 90],
        35: [4, 8, 20, 48, 120], 40: [4, 8, 20, 48, 120],
        45: [5, 10, 25, 60, 150], 50: [5, 10, 25, 60, 150],
        55: [6, 12, 30, 72, 180], 60: [6, 12, 30, 72, 180],
    }
    for equip_level, target_values in equip_exp.items():
        for quality, target_value in zip(qualities, target_values):
            current = exp_current.get((equip_level, quality))
            if current is None:
                raise RuntimeError(f"PortableExp 缺少 ({equip_level}, {quality})")
            row_id = int(current["Id"])
            add_draft(
                rows,
                draft_id=f"D01-PX-{row_id:03d}",
                parameter_id="ME.FORGE.EQUIP_XP",
                source_md=MD_ECON,
                source_section="锻造台 / 单件装备提供锻造台经验",
                logical_table="PortableExp",
                source_workbook=portable_exp_path,
                sheet="PortableExp",
                key_field="Id",
                key_value=row_id,
                target_field="Exp",
                raw_type="int",
                current_value=display_current(current.get("Exp")),
                suggested_value=target_value,
                transform="按装备等级段与品质查权威MD经验表直接写入",
                unit_rounding="整数经验；不舍入",
                constraint="Quality={Green,Blue,Purple,Orange,Red}",
                impact="分解与9合1向锻造台提供经验；打造不给经验。Portable运行消费链仍需正式落地前确认。",
                validation="60个(Id/EquipLv/Quality)唯一；目标范围1~180；同等级随品质递增；与MD矩阵逐项一致。",
                created_at=created_at,
            )

    # 2) 副属性等级倍率：保留现有 Id，不为插入 Lv10 破坏已有引用。
    sub_path = "SpecialEquipment/SpecialEquipmentSubAttrCoef.xlsx"
    sub_rows = read_config_rows(config_root, sub_path, "SpecialEquipmentSubAttrCoef")
    sub_by_id = {int(row["Id"]): row for row in sub_rows}
    sub_targets = {
        1: (5, Decimal("0.200")),
        2: (15, Decimal("0.345")),
        3: (20, Decimal("0.418")),
        4: (25, Decimal("0.491")),
        5: (30, Decimal("0.564")),
        6: (35, Decimal("0.636")),
        7: (40, Decimal("0.709")),
        8: (45, Decimal("0.782")),
        9: (50, Decimal("0.855")),
        10: (55, Decimal("0.927")),
        11: (60, Decimal("1.000")),
        12: (10, Decimal("0.273")),
    }
    for row_id, (equip_level, target_coef) in sub_targets.items():
        current = sub_by_id.get(row_id)
        if current is None:
            add_draft(
                rows,
                draft_id=f"D01-SC-{row_id:02d}-LV",
                parameter_id="ME.ATTR.SUB_LEVEL_CURVE",
                source_md=MD_ATTR,
                source_section="副属性曲线",
                logical_table="SpecialEquipmentSubAttrCoef",
                source_workbook=sub_path,
                sheet="SpecialEquipmentSubAttrCoef",
                key_field="Id",
                key_value=row_id,
                target_field="EquipLv",
                raw_type="int",
                current_value="<NOT_EXIST>",
                suggested_value=equip_level,
                write_mode="NEW_ROW_MANUAL",
                risk="HIGH",
                status="MANUAL_REQUIRED",
                validation="新增 Id=12 后，装备等级集合必须严格为5/10/.../60且主键唯一。",
                created_at=created_at,
                note="新增Lv10行；禁止自动插行，获批后人工追加并复查引用。",
            )
            add_draft(
                rows,
                draft_id=f"D01-SC-{row_id:02d}-COEF",
                parameter_id="ME.ATTR.SUB_LEVEL_CURVE",
                source_md=MD_ATTR,
                source_section="副属性曲线",
                logical_table="SpecialEquipmentSubAttrCoef",
                source_workbook=sub_path,
                sheet="SpecialEquipmentSubAttrCoef",
                key_field="Id",
                key_value=row_id,
                target_field="SubAttrCoef",
                raw_type="float",
                current_value="<NOT_EXIST>",
                suggested_value=float(target_coef),
                transform="0.20 + 0.80 * ((x - 5) / 55)，按当前权威表值保留3位小数",
                unit_rounding="倍率；3位小数",
                write_mode="NEW_ROW_MANUAL",
                risk="HIGH",
                status="MANUAL_REQUIRED",
                validation="Lv10目标0.273；12档端点Lv5=0.200、Lv60=1.000；单调不降。",
                created_at=created_at,
                note="与同主键EquipLv字段一起新增；本批次不操作正式表行。",
            )
        else:
            if int(current["EquipLv"]) != equip_level:
                raise RuntimeError(
                    f"SubAttrCoef Id={row_id} 当前EquipLv漂移：{current['EquipLv']} != {equip_level}"
                )
            add_draft(
                rows,
                draft_id=f"D01-SC-{row_id:02d}-COEF",
                parameter_id="ME.ATTR.SUB_LEVEL_CURVE",
                source_md=MD_ATTR,
                source_section="副属性曲线",
                logical_table="SpecialEquipmentSubAttrCoef",
                source_workbook=sub_path,
                sheet="SpecialEquipmentSubAttrCoef",
                key_field="Id",
                key_value=row_id,
                target_field="SubAttrCoef",
                raw_type="float",
                current_value=display_current(current.get("SubAttrCoef")),
                suggested_value=float(target_coef),
                transform="0.20 + 0.80 * ((x - 5) / 55)，按当前权威表值保留3位小数",
                unit_rounding="倍率；3位小数",
                validation="12档端点Lv5=0.200、Lv60=1.000；按装备等级单调不降。",
                created_at=created_at,
            )

    # 3) 品质系数、红装分解额外道具及旧字段退出。
    quality_path = "SpecialEquipment/SpecialEquipmentQuality.xlsx"
    quality_rows = read_config_rows(config_root, quality_path, "SpecialEquipment")
    quality_by_name = {str(row["Quality"]): row for row in quality_rows}
    main_quality = {"Green": .20, "Blue": .30, "Purple": .45, "Orange": .70, "Red": 1.00}
    sub_quality = {"Green": .30, "Blue": .45, "Purple": .60, "Orange": .80, "Red": 1.00}
    for quality in qualities:
        current = quality_by_name[quality]
        row_id = int(current["Id"])
        for suffix, parameter_id, field, target_value, section in [
            ("MAIN", "ME.ATTR.QUALITY_MAIN_COEF", "GradeCoefficient", main_quality[quality], "品质系数 / 主属性品质系数"),
            ("SUB", "ME.ATTR.QUALITY_SUB_COEF", "QualityCoefficient", sub_quality[quality], "品质系数 / 副属性品质系数"),
        ]:
            add_draft(
                rows,
                draft_id=f"D01-QC-{row_id:02d}-{suffix}",
                parameter_id=parameter_id,
                source_md=MD_ATTR,
                source_section=section,
                logical_table="SpecialEquipmentQuality",
                source_workbook=quality_path,
                sheet="SpecialEquipment",
                key_field="Id",
                key_value=row_id,
                target_field=field,
                raw_type="float",
                current_value=display_current(current.get(field)),
                suggested_value=target_value,
                unit_rounding="倍率；最多2位小数",
                constraint="Quality枚举；Orange对应中文金品质",
                impact="字段可承载，但旧运行时语义需确认：GradeCoefficient=主属性品质系数，QualityCoefficient=副属性品质系数。",
                write_mode="MANUAL_REQUIRED",
                risk="HIGH",
                validation="五品质系数与MD逐项一致；Green<Blue<Purple<Orange<Red；端点Red=1。",
                status="MANUAL_REQUIRED",
                created_at=created_at,
                note="语义确认后才可写正式配置；当前只形成字段级草案。",
            )

    extra_targets = {"Purple": "<BLANK>", "Orange": "<BLANK>", "Red": "23,1"}
    for quality, target_value in extra_targets.items():
        current = quality_by_name[quality]
        row_id = int(current["Id"])
        add_draft(
            rows,
            draft_id=f"D01-QX-{row_id:02d}",
            parameter_id="ME.DISMANTLE.EXTRA_ITEM",
            source_md=MD_ECON,
            source_section="分解回收 / 红装分解额外产出",
            logical_table="SpecialEquipmentQuality",
            source_workbook=quality_path,
            sheet="SpecialEquipment",
            key_field="Id",
            key_value=row_id,
            target_field="ExtraItemList",
            raw_type="(array#sep=;),ExtraItem",
            current_value=display_current(current.get("ExtraItemList")),
            suggested_value=target_value,
            constraint="ExtraItem(Id:int,Num:int)，Bean分隔符逗号、数组分隔符分号",
            impact="分解UI仍读取ExtraItemList；必须与代码/UI迁移同步，不能单独清空。",
            write_mode="MANUAL_REQUIRED",
            risk="HIGH",
            validation="仅Red配置23,1；Green/Blue/Purple/Orange为空；Item(23)存在。",
            status="MANUAL_REQUIRED",
            created_at=created_at,
            note="旧品质额外产出迁移；正式写入前需确认分解消费链。",
        )

    raw_types = {
        "GradeLimit": "int", "Basic": "int", "HoleMax": "int", "HoleMin": "int",
        "ExtraHoleNum": "int", "InscriptionProbability": "float",
        "IdentifyCost": "(array#sep=;),PropType", "IdentifyCostCurrency": "PropType?",
        "HoleCost": "(array#sep=;),PropType", "HoleCostCurrency": "PropType?",
    }
    for index, field in enumerate(raw_types, start=1):
        current_values = ";".join(
            f"{quality}={display_current(quality_by_name[quality].get(field))}" for quality in qualities
        )
        parameter_id = "ME.DISMANTLE.GOLD" if field == "Basic" else (
            "ME.OLD.UPGRADE_EXIT" if field == "GradeLimit" else "ME.OLD.HOLE_IDENTIFY_EXIT"
        )
        suggested = (
            "<MIGRATE_TO PortableExp.ResolveGold>" if field == "Basic" else
            "<DEPRECATED_AFTER_CODE_MIGRATION>"
        )
        add_draft(
            rows,
            draft_id=f"D01-QOLD-{index:02d}",
            parameter_id=parameter_id,
            source_md=MD_ECON,
            source_section="当前 v2 模块表观察 / 旧升级、打孔、鉴定废弃",
            logical_table="SpecialEquipmentQuality",
            source_workbook=quality_path,
            sheet="SpecialEquipment",
            key_field="Id",
            key_value="1..5",
            target_field=field,
            raw_type=raw_types[field],
            current_value=current_values,
            suggested_value=suggested,
            impact="现有客户端仍消费部分旧字段；必须先完成代码/UI迁移和回归，再清空或移除。",
            write_mode="MANUAL_REQUIRED",
            risk="HIGH",
            validation="代码与UI零引用后再处理；不得只改Excel制造运行时空引用。",
            status="MANUAL_REQUIRED",
            created_at=created_at,
            note="旧系统退出项；不是本批次可直接写入的数值。",
        )

    # 4) 合成金币：40个现存单元格有精确目标，但Range/Generate引用未闭合。
    compound_path = "SpecialEquipment/Portable/PortableEquipCompound.xlsx"
    compound_rows = read_config_rows(config_root, compound_path, "PortableEquipCompound")
    compound_by_id = {int(row["Id"]): row for row in compound_rows}
    compound_gold = {
        1: [100, 140, 200, 280, 400],
        2: [350, 490, 700, 980, 1400],
        3: [1000, 1400, 2000, 2800, 4000],
        4: [2500, 3500, 5000, 7000, 10000],
        5: [3800, 5320, 7600, 10640, 15200],
        6: [5500, 7700, 11000, 15400, 22000],
        7: [7500, 10500, 15000, 21000, 30000],
        8: [10000, 14000, 20000, 28000, 40000],
    }
    for tier, target_values in compound_gold.items():
        for quality_index, (quality, target_value) in enumerate(zip(qualities, target_values), start=1):
            row_id = (tier - 1) * 5 + quality_index
            current = compound_by_id.get(row_id)
            if current is None:
                raise RuntimeError(f"PortableEquipCompound 缺少 Id={row_id}")
            if int(current["RangeId"]) != tier or current["Quality"] != quality:
                raise RuntimeError(f"PortableEquipCompound Id={row_id} 主键邻域漂移")
            add_draft(
                rows,
                draft_id=f"D01-CG-{row_id:03d}",
                parameter_id="ME.COMPOUND.GOLD",
                source_md=MD_ECON,
                source_section="合成消耗 / 合成金币矩阵",
                logical_table="PortableEquipCompound",
                source_workbook=compound_path,
                sheet="PortableEquipCompound",
                key_field="Id",
                key_value=row_id,
                target_field="CostCurrency",
                raw_type="int",
                current_value=display_current(current.get("CostCurrency")),
                suggested_value=target_value,
                transform="基础金币 * 品质金币系数",
                unit_rounding="整数金币；矩阵已给定",
                constraint="Quality枚举；RangeId应引用PortableRange.Id",
                impact="RangeId 3~8当前悬空，EquipResult仍有第二套品质池冲突；引用修复前禁止写正式表。",
                write_mode="MANUAL_REQUIRED",
                risk="HIGH",
                validation="8档×5品质=40行；每档按0.5/0.7/1/1.4/2系数；RangeId引用全部存在。",
                status="MANUAL_REQUIRED",
                created_at=created_at,
                note="金币值已精确映射；跨表引用未闭合，暂不可正式应用。",
            )

    # 5) 锻造台升级：现表结构可写，语义与运行时尚未迁移；新行必须人工追加。
    level_path = "SpecialEquipment/Portable/PortableLvUp.xlsx"
    level_rows = read_config_rows(config_root, level_path, "PortableLvUp")
    level_by_id = {int(row["Id"]): row for row in level_rows}
    for target_level in range(2, 61):
        row_id = target_level - 1
        current = level_by_id.get(row_id)
        is_new = current is None
        if current is not None and int(current["TargetLv"]) != target_level:
            raise RuntimeError(f"PortableLvUp Id={row_id} TargetLv漂移")
        if is_new:
            add_draft(
                rows,
                draft_id=f"D01-FL-{row_id:03d}-LV",
                parameter_id="ME.FORGE.LEVEL_UP_XP",
                source_md=MD_ECON,
                source_section="锻造台 / 升级经验公式",
                logical_table="PortableLvUp",
                source_workbook=level_path,
                sheet="PortableLvUp",
                key_field="Id",
                key_value=row_id,
                target_field="TargetLv",
                raw_type="int",
                current_value="<NOT_EXIST>",
                suggested_value=target_level,
                write_mode="NEW_ROW_MANUAL",
                risk="HIGH",
                status="MANUAL_REQUIRED",
                impact="表需从Lv2~20扩为Lv2~60；Portable运行时消费链尚未接入。",
                validation="新增Id20~59连续且唯一；TargetLv=Id+1；最终覆盖Lv2~60。",
                created_at=created_at,
                note="新增锻造台等级行；禁止自动插行。",
            )
        add_draft(
            rows,
            draft_id=f"D01-FL-{row_id:03d}-XP",
            parameter_id="ME.FORGE.LEVEL_UP_XP",
            source_md=MD_ECON,
            source_section="锻造台 / 升级经验公式",
            logical_table="PortableLvUp",
            source_workbook=level_path,
            sheet="PortableLvUp",
            key_field="Id",
            key_value=row_id,
            target_field="CostExp",
            raw_type="int",
            current_value=display_current(current.get("CostExp") if current else None, missing=is_new),
            suggested_value=forge_level_cost(target_level),
            transform="按2026-07-24确认的59级离散曲线取值",
            unit_rounding="经验；整数",
            impact="锻造台升级成本；现表旧注释/运行语义需从装备升级迁移为锻造台升级。",
            write_mode="NEW_ROW_MANUAL" if is_new else "MANUAL_REQUIRED",
            risk="HIGH",
            validation="TargetLv2成本10；TargetLv60成本5600；Lv1~60累计94500；逐级严格递增。",
            status="MANUAL_REQUIRED",
            created_at=created_at,
            note="数值已精确生成；语义迁移与运行时接入完成前不可正式应用。",
        )

    # 6) 结构、代码和跨表阻断项。每行均指向明确字段或Schema决策。
    compact_rows = [
        dict(id="D01-CFG-001", pid="ME.ATTR.SINGLE_MAIN_RULE", md=MD_ATTR, sec="核心结论 / 部位主属性", table="SpecialEquipment", book="SpecialEquipment/SpecialEquipment.xlsx", sheet="SpecialEquipment", key="Id", kval="80001..80190", field="Attribute2List", raw="(array#sep=;),AttributeMain", cur="190/190行非空", sug="<BLANK_AFTER_CODE_MIGRATION>", mode="MANUAL_REQUIRED", status="MANUAL_REQUIRED", note="旧第二主属性退出；先迁移代码/展示与主属性引用。"),
        dict(id="D01-CFG-002", pid="ME.ATTR.QUALITY_MAIN_COEF", md=MD_ATTR, sec="装备品质 / 绿色品质承接", table="SpecialEquipment", book="SpecialEquipment/SpecialEquipment.xlsx", sheet="SpecialEquipment", key="Id", kval="<PENDING_ID>", field="Quality", raw="Quality", cur="Green装备行=0", sug="新增Green装备承接；Id/Name/Items引用待规划", mode="NEW_ROW_MANUAL", status="DRAFT", note="不得虚构装备Id或ItemId。"),
        dict(id="D01-CFG-003", pid="ME.ATTR.MAIN_LEVEL_CURVE", md=MD_ATTR, sec="主属性曲线", table="SpecialEquipmentMain", book="SpecialEquipment/SpecialEquipmentMain.xlsx", sheet="SpecialEquipment", key="<SCHEMA>", kval="<SCHEMA>", field="<SCHEMA:TargetValue60/CurveId>", raw="<NEW_SCHEMA>", cur="<NOT_EXIST>", sug="新增60红目标与非线性等级曲线引用，或新增独立主属性等级系数表", mode="SCHEMA_CHANGE", status="MANUAL_REQUIRED", note="现有Value+Grow线性结构无法表达1.35次幂曲线。"),
        dict(id="D01-CFG-004", pid="ME.ATTR.MAIN_LEVEL_CURVE", md=MD_ATTR, sec="主属性曲线 / 旧框架切片处理", table="SpecialEquipmentMain", book="SpecialEquipment/SpecialEquipmentMain.xlsx", sheet="SpecialEquipment", key="Id", kval="800011..801902", field="Grow", raw="float", cur="380行旧线性成长", sug="<DEPRECATED_AFTER_MAIN_CURVE_MIGRATION>", mode="MANUAL_REQUIRED", status="MANUAL_REQUIRED", note="与第二主属性行一起退出。"),
        dict(id="D01-CFG-005", pid="ME.ATTR.MAIN_LEVEL_CURVE", md=MD_ATTR, sec="主属性曲线", table="SpecialEquipmentMain", book="SpecialEquipment/SpecialEquipmentMain.xlsx", sheet="SpecialEquipment", key="<DECISION>", kval="曲线精度", field="Value/Grow或新曲线表", raw="float", cur="04终稿按3位小数表值计算；MD同时给出全精度公式", sug="<PENDING_DECISION:全精度公式或3位表值>", mode="MANUAL_REQUIRED", status="DRAFT", note="正式属性落表前必须裁决精度，不能静默选择。"),
        dict(id="D01-CFG-006", pid="ME.ATTR.SUB_TARGET_L60_RED", md=MD_ATTR, sec="副属性曲线 / 60级红装单条目标值", table="SpecialEquipmentEffect", book="SpecialEquipment/SpecialEquipmentEffect.xlsx", sheet="SpecialEquipment", key="Id", kval="<PENDING_EFFECT_ID_MAP>", field="ValueMin/ValueMax", raw="float?", cur="246行旧范围，尚未建立12类属性ID对应", sug="按12类60红目标与品质/等级系数重建上下限", mode="MANUAL_REQUIRED", status="DRAFT", note="先完成EffectId到12类属性的字段级映射。"),
        dict(id="D01-CFG-007", pid="ME.ATTR.NO_REPEAT_SUB", md=MD_ATTR, sec="副属性随机边界", table="SpecialEquipmentEffectDrop", book="SpecialEquipment/SpecialEquipmentEffectDrop.xlsx", sheet="SpecialEquipment", key="<SCHEMA/CODE>", kval="全局", field="<NO_DUPLICATE_RULE>", raw="<NEW_SCHEMA_OR_CODE>", cur="现表仅表达候选池和权重", sug="初始生成与洗练均按属性去重抽取", mode="SCHEMA_CHANGE", status="MANUAL_REQUIRED", note="仅改EffectList不能保证不重复。"),
        dict(id="D01-CFG-008", pid="ME.COMPOUND.TIER", md=MD_ECON, sec="合成消耗 / 合成档位", table="SpecialEquipmentGenerate", book="SpecialEquipment/SpecialEquipmentGenerate.xlsx", sheet="SpecialEquipment", key="Id", kval="1001..及副本池", field="GradeList", raw="(array#sep=;),GradeList", cur="当前通用/固定池仍主要为5/10/15", sug="按12个装备等级及8档/副本来源重建等级池", mode="MANUAL_REQUIRED", status="MANUAL_REQUIRED", note="受宝箱来源档位方案阻断。"),
        dict(id="D01-CFG-009", pid="ME.CONFIG.CHEST_PROVENANCE", md=MD_GEN, sec="配置落地追踪 ID", table="SpecialEquipmentGenerate", book="SpecialEquipment/SpecialEquipmentGenerate.xlsx", sheet="SpecialEquipment", key="<CROSS_TABLE>", kval="Items4001..4005", field="GradeList/QualityList", raw="arrays", cur="5种可堆叠箱只固定引用一个GenerateId，开启时丢失来源层", sug="三选一：按档拆箱；掉落时直生装备；箱实例/请求携带来源档位", mode="SCHEMA_CHANGE", status="DRAFT", note="未裁决前不得重写箱池。"),
        dict(id="D01-CFG-010", pid="ME.OLD.HOLE_IDENTIFY_EXIT", md=MD_ECON, sec="当前 v2 模块表观察 / 鉴定废弃", table="SpecialEquipmentGenerate", book="SpecialEquipment/SpecialEquipmentGenerate.xlsx", sheet="SpecialEquipment", key="Id", kval="全部", field="UnidentifiedList", raw="(array#sep=;),Unidentified", cur="字段仍存在", sug="<DEPRECATED_AFTER_CODE_MIGRATION>", mode="MANUAL_REQUIRED", status="MANUAL_REQUIRED", note="鉴定退出后再清理。"),
        dict(id="D01-CFG-011", pid="ME.COMPOUND.TIER", md=MD_ECON, sec="合成消耗 / 合成档位", table="PortableRange", book="SpecialEquipment/Portable/PortableRange.xlsx", sheet="PortableRange", key="Id", kval="1..8", field="LevelMin", raw="int", cur="Id1=1;Id2=5;Id3..8缺行", sug="10;15;25;35;40;45;50;55", mode="NEW_ROW_MANUAL", status="MANUAL_REQUIRED", note="需先确认LevelMin/Max确为产物等级端点语义。"),
        dict(id="D01-CFG-012", pid="ME.COMPOUND.TIER", md=MD_ECON, sec="合成消耗 / 合成档位", table="PortableRange", book="SpecialEquipment/Portable/PortableRange.xlsx", sheet="PortableRange", key="Id", kval="1..8", field="LevelMax", raw="float", cur="Id1=15;Id2=20;Id3..8缺行", sug="10;20;30;40;45;50;55;60", mode="NEW_ROW_MANUAL", status="MANUAL_REQUIRED", note="字段RawType当前为float，目标等级应裁决是否改int。"),
        dict(id="D01-CFG-013", pid="ME.COMPOUND.TIER", md=MD_ECON, sec="合成消耗 / 合成档位", table="PortableRange", book="SpecialEquipment/Portable/PortableRange.xlsx", sheet="PortableRange", key="Id", kval="1..8", field="EquipMinLv", raw="int", cur="Id1=1;Id2=10;Id3..8缺行", sug="10;10;20;30;40;45;50;55", mode="NEW_ROW_MANUAL", status="MANUAL_REQUIRED", note="对应8档投入装备等级下限。"),
        dict(id="D01-CFG-014", pid="ME.COMPOUND.TIER", md=MD_ECON, sec="合成消耗 / 合成档位", table="PortableRange", book="SpecialEquipment/Portable/PortableRange.xlsx", sheet="PortableRange", key="<SCHEMA>", kval="1..8", field="<SCHEMA:ResultLevelList>", raw="(array#sep=;),LevelWeight", cur="<NOT_EXIST>", sug="10:100%；15/20~35/40:70/30；40/45~55/60:65/35", mode="SCHEMA_CHANGE", status="MANUAL_REQUIRED", note="现有Min/Max不能表达离散权重。"),
        dict(id="D01-CFG-015", pid="ME.CONFIG.UNLOCK_CHAIN", md=MD_GEN, sec="配置落地追踪 ID", table="PortableRange", book="SpecialEquipment/Portable/PortableRange.xlsx", sheet="PortableRange", key="Id", kval="1..8", field="UnlockId", raw="int?", cur="Id1/2均空；Id3..8缺行", sug="<PENDING_UNLOCK_IDS:角色Lv10/20/30/40/45/50/55/60>", mode="NEW_ROW_MANUAL", status="DRAFT", note="UnlockId分配与Unlock表引用待补。"),
        dict(id="D01-CFG-016", pid="ME.COMPOUND.QUALITY_PROB", md=MD_ECON, sec="合成消耗 / 合成品质池", table="PortableEquipCompound", book="SpecialEquipment/Portable/PortableEquipCompound.xlsx", sheet="PortableEquipCompound", key="Id", kval="1..40", field="EquipResult", raw="int", cur="当前指1001..1004固定品质Generate池", sug="只负责部位/等级生成；品质唯一由PortableCompoundProb决定", mode="MANUAL_REQUIRED", status="MANUAL_REQUIRED", note="必须消除双重判品质。"),
        dict(id="D01-CFG-017", pid="ME.CRAFT.TIER", md=MD_ECON, sec="打造消耗 / 打造档", table="PortableEquipCraft", book="SpecialEquipment/Portable/PortableEquipCraft.xlsx", sheet="PortableEquipCraft", key="<SCHEMA>", kval="<SCHEMA>", field="<SCHEMA:CraftTier/RangeId>", raw="int", cur="<NOT_EXIST>", sug="8档：10/20/30/40/45/50/55/60", mode="SCHEMA_CHANGE", status="MANUAL_REQUIRED", note="现24行无档位字段。"),
        dict(id="D01-CFG-018", pid="ME.CRAFT.TIER", md=MD_ECON, sec="打造消耗 / 打造档", table="PortableEquipCraft", book="SpecialEquipment/Portable/PortableEquipCraft.xlsx", sheet="PortableEquipCraft", key="<SCHEMA>", kval="<SCHEMA>", field="<SCHEMA:AllowedEquipLevelList>", raw="(array#sep=;),int", cur="<NOT_EXIST>", sug="10；15/20；25/30；35/40；40/45；45/50；50/55；55/60", mode="SCHEMA_CHANGE", status="MANUAL_REQUIRED", note="支持玩家在档内指定等级。"),
        dict(id="D01-CFG-019", pid="ME.CRAFT.MATERIAL", md=MD_ECON, sec="打造消耗 / 主材和副材对应", table="PortableEquipCraft", book="SpecialEquipment/Portable/PortableEquipCraft.xlsx", sheet="PortableEquipCraft", key="<SCHEMA>", kval="<SCHEMA>", field="<SCHEMA:MainMaterialQuality>", raw="Quality", cur="<NOT_EXIST>", sug="Purple/Orange/Red", mode="SCHEMA_CHANGE", status="MANUAL_REQUIRED", note="主材品质决定成品品质池。"),
        dict(id="D01-CFG-020", pid="ME.CRAFT.MATERIAL", md=MD_ECON, sec="打造消耗 / 主材和副材对应", table="PortableEquipCraft", book="SpecialEquipment/Portable/PortableEquipCraft.xlsx", sheet="PortableEquipCraft", key="Id", kval="<NEW_ROWS>", field="CostItemMain", raw="(array#sep=;),CostItem", cur="24行占位2,1", sug="8档×8部位×3主材品质，主材x1；对应72个主材ItemId", mode="NEW_ROW_MANUAL", status="DRAFT", note="72个ItemId待规划，不能沿用旧24胚料。"),
        dict(id="D01-CFG-021", pid="ME.CRAFT.MATERIAL", md=MD_ECON, sec="打造消耗 / 主材和副材对应", table="PortableEquipCraft", book="SpecialEquipment/Portable/PortableEquipCraft.xlsx", sheet="PortableEquipCraft", key="Id", kval="<NEW_ROWS>", field="CostItemAssist", raw="(array#sep=;),CostItem", cur="24行占位2,2;3,5", sug="武器副1+2；防具副2+3；首饰副3+4；数量按8档×主材品质矩阵", mode="NEW_ROW_MANUAL", status="DRAFT", note="4个副材ItemId待规划。"),
        dict(id="D01-CFG-022", pid="ME.CRAFT.GOLD", md=MD_ECON, sec="打造消耗 / 打造档", table="PortableEquipCraft", book="SpecialEquipment/Portable/PortableEquipCraft.xlsx", sheet="PortableEquipCraft", key="CraftTier", kval="1..8", field="CostCurrency", raw="int", cur="24行均1000", sug="500;1200;2800;6200;9000;13000;18000;24000", mode="MANUAL_REQUIRED", status="DRAFT", note="MD标为暂沿用旧档，正式落表前仍需用户确认。"),
        dict(id="D01-CFG-023", pid="ME.COMPOUND.QUALITY_PROB", md=MD_ECON, sec="合成消耗 / 合成概率逐级表", table="PortableCompoundProb", book="SpecialEquipment/Portable/PortableCompoundProb.xlsx", sheet="PortableCompoundProb", key="Id", kval="<EXPAND>", field="Grade/Quality/QualityList", raw="int/Quality/QualityWeight[]", cur="Lv1~20×5=100行；全部同一占位权重", sug="扩为Lv1~60×5=300行", mode="NEW_ROW_MANUAL", status="MANUAL_REQUIRED", note="概率曲线明确，权重整数化尺度待确认。"),
        dict(id="D01-CFG-024", pid="ME.COMPOUND.QUALITY_PROB", md=MD_ECON, sec="合成消耗 / 合成最高品质概率公式", table="PortableCompoundProb", book="SpecialEquipment/Portable/PortableCompoundProb.xlsx", sheet="PortableCompoundProb", key="Grade+Quality", kval="Lv1..60+Green", field="QualityList", raw="QualityWeight[]", cur="占位Blue39;Purple610;Orange340;Red11", sug="Blue=1-p;Purple=p，p=8%+17%*(Lv-1)/59", mode="MANUAL_REQUIRED", status="MANUAL_REQUIRED", note="绿合蓝/紫。"),
        dict(id="D01-CFG-025", pid="ME.COMPOUND.QUALITY_PROB", md=MD_ECON, sec="合成消耗 / 合成最高品质概率公式", table="PortableCompoundProb", book="SpecialEquipment/Portable/PortableCompoundProb.xlsx", sheet="PortableCompoundProb", key="Grade+Quality", kval="Lv1..60+Blue", field="QualityList", raw="QualityWeight[]", cur="同一占位", sug="Purple=1-p;Orange=p，p=5%+15%*(Lv-1)/59", mode="MANUAL_REQUIRED", status="MANUAL_REQUIRED", note="蓝合紫/金。"),
        dict(id="D01-CFG-026", pid="ME.COMPOUND.QUALITY_PROB", md=MD_ECON, sec="合成消耗 / 合成最高品质概率公式", table="PortableCompoundProb", book="SpecialEquipment/Portable/PortableCompoundProb.xlsx", sheet="PortableCompoundProb", key="Grade+Quality", kval="Lv1..60+Purple", field="QualityList", raw="QualityWeight[]", cur="同一占位", sug="Orange=30%+30%*(Lv-1)/59;Red=1%+2%*(Lv-1)/59;Purple=补集", mode="MANUAL_REQUIRED", status="MANUAL_REQUIRED", note="紫合紫/金/红。"),
        dict(id="D01-CFG-027", pid="ME.COMPOUND.QUALITY_PROB", md=MD_ECON, sec="合成消耗 / 合成最高品质概率公式", table="PortableCompoundProb", book="SpecialEquipment/Portable/PortableCompoundProb.xlsx", sheet="PortableCompoundProb", key="Grade+Quality", kval="Lv1..60+Orange", field="QualityList", raw="QualityWeight[]", cur="同一占位", sug="Orange=1-p;Red=p，p=20%+20%*(Lv-1)/59", mode="MANUAL_REQUIRED", status="MANUAL_REQUIRED", note="金合金/红。"),
        dict(id="D01-CFG-028", pid="ME.COMPOUND.QUALITY_PROB", md=MD_ECON, sec="合成消耗 / 合成最高品质概率公式", table="PortableCompoundProb", book="SpecialEquipment/Portable/PortableCompoundProb.xlsx", sheet="PortableCompoundProb", key="Grade+Quality", kval="Lv1..60+Red", field="QualityList", raw="QualityWeight[]", cur="同一占位", sug="Red=100%", mode="MANUAL_REQUIRED", status="MANUAL_REQUIRED", note="红合红。"),
        dict(id="D01-CFG-029", pid="ME.COMPOUND.QUALITY_PROB", md=MD_ECON, sec="合成消耗 / 合成概率逐级表", table="PortableCompoundProb", book="SpecialEquipment/Portable/PortableCompoundProb.xlsx", sheet="PortableCompoundProb", key="<DECISION>", kval="权重尺度", field="QualityList.Weight", raw="int", cur="现表每行总权重1000", sug="<PENDING_DECISION:总权重1000或10000；最大余数法保证和固定>", mode="MANUAL_REQUIRED", status="DRAFT", note="MD显示到0.01%，1000权重无法逐级完全复现。"),
        dict(id="D01-CFG-030", pid="ME.CRAFT.QUALITY_PROB", md=MD_ECON, sec="打造消耗 / 打造概率逐级表", table="PortableCraftProb", book="SpecialEquipment/Portable/PortableCraftProb.xlsx", sheet="PortableCompoundProb", key="Id", kval="<EXPAND>", field="PortableLv/Quality/QualityList", raw="int/Quality/QualityWeight[]", cur="Lv1~20×3=60行；全部同一占位权重", sug="扩为Lv1~60×3=180行", mode="NEW_ROW_MANUAL", status="MANUAL_REQUIRED", note="实际Sheet名仍误为PortableCompoundProb。"),
        dict(id="D01-CFG-031", pid="ME.CRAFT.QUALITY_PROB", md=MD_ECON, sec="打造消耗 / 打造最高品质概率公式", table="PortableCraftProb", book="SpecialEquipment/Portable/PortableCraftProb.xlsx", sheet="PortableCompoundProb", key="PortableLv+Quality", kval="Lv1..60+Purple", field="QualityList", raw="QualityWeight[]", cur="占位", sug="Purple=1-p;Orange=p，p=10%+10%*(Lv-1)/59", mode="MANUAL_REQUIRED", status="MANUAL_REQUIRED", note="紫主材出紫/金。"),
        dict(id="D01-CFG-032", pid="ME.CRAFT.QUALITY_PROB", md=MD_ECON, sec="打造消耗 / 打造最高品质概率公式", table="PortableCraftProb", book="SpecialEquipment/Portable/PortableCraftProb.xlsx", sheet="PortableCompoundProb", key="PortableLv+Quality", kval="Lv1..60+Orange", field="QualityList", raw="QualityWeight[]", cur="占位", sug="Orange=1-p;Red=p，p=2%+3%*(Lv-1)/59", mode="MANUAL_REQUIRED", status="MANUAL_REQUIRED", note="金主材出金/红。"),
        dict(id="D01-CFG-033", pid="ME.CRAFT.QUALITY_PROB", md=MD_ECON, sec="打造消耗 / 打造最高品质概率公式", table="PortableCraftProb", book="SpecialEquipment/Portable/PortableCraftProb.xlsx", sheet="PortableCompoundProb", key="PortableLv+Quality", kval="Lv1..60+Red", field="QualityList", raw="QualityWeight[]", cur="占位", sug="Red=100%", mode="MANUAL_REQUIRED", status="MANUAL_REQUIRED", note="红主材必出红。"),
        dict(id="D01-CFG-034", pid="ME.CRAFT.QUALITY_PROB", md=MD_ECON, sec="打造消耗", table="PortableCraftProb", book="SpecialEquipment/Portable/PortableCraftProb.xlsx", sheet="PortableCompoundProb", key="<SCHEMA>", kval="Sheet", field="<SCHEMA:SHEET_NAME>", raw="worksheet", cur="PortableCompoundProb", sug="PortableCraftProb", mode="SCHEMA_CHANGE", status="MANUAL_REQUIRED", note="需同步__tables__与生成代码核对，不能只重命名Sheet。"),
        dict(id="D01-CFG-035", pid="ME.COMPOUND.COUNT", md=MD_ECON, sec="合成消耗 / 基础规则", table="globals_conf", book="globals_conf/globals_conf.xlsx", sheet="globals_conf", key="<SINGLETON>", kval="globals_conf", field="DefaultSpecialEquipmentCompoundNum", raw="int", cur=9, sug=9, mode="CANDIDATE_SAFE_RECHECK", status="READY_FOR_REVIEW", note="核对项：现值已等于9，无实际写入差异。"),
        dict(id="D01-CFG-036", pid="ME.DISMANTLE.GOLD", md=MD_ECON, sec="分解回收 / 分解金币矩阵", table="PortableExp", book="SpecialEquipment/Portable/PortableExp.xlsx", sheet="PortableExp", key="<SCHEMA>", kval="EquipLv+Quality", field="<SCHEMA:ResolveGold>", raw="int", cur="<NOT_EXIST>", sug="12等级×5品质分解金币矩阵；基础金币×品质系数", mode="SCHEMA_CHANGE", status="MANUAL_REQUIRED", note="Quality.Basic无法表达等级维度。"),
        dict(id="D01-CFG-037", pid="ME.CONFIG.ITEM_IDS", md=MD_GEN, sec="配置落地追踪 ID", table="Item", book="items/Items.xlsx", sheet="items", key="Id", kval="<PENDING_ID>", field="Id/Name/Quality/Type/ParaList", raw="mixed", cur="<NOT_EXIST>", sug="72种主材+4种副材；与背包类型/图标/堆叠上限一起规划", mode="NEW_ROW_MANUAL", status="DRAFT", note="不得用高阶零件23/低阶零件24替代。"),
        dict(id="D01-CFG-038", pid="ME.CONFIG.CHEST_PROVENANCE", md=MD_GEN, sec="配置落地追踪 ID", table="Item", book="items/Items.xlsx", sheet="items", key="Id", kval="4001..4005", field="ParaList", raw="array", cur="4001和4002均指向Generate1001", sug="按宝箱来源档位裁决结果重连Generate", mode="MANUAL_REQUIRED", status="DRAFT", note="当前1001为Blue100%，与绿箱目标不符。"),
        dict(id="D01-CFG-039", pid="ME.CONFIG.DROP_CHAIN", md=MD_GEN, sec="配置落地追踪 ID", table="Drop", book="Drop/Drop_dungeon.xlsx", sheet="Drop", key="Id", kval="<PENDING_DROP_IDS>", field="DropId/DropId2及奖励字段", raw="mixed", cur="目标行含公式/引用，未建立本批次ID映射", sug="映射副本常驻装备/主材/副材掉落与首通奖励", mode="MANUAL_REQUIRED", status="DRAFT", note="该表含公式，禁止自动插删行或覆盖公式。"),
        dict(id="D01-CFG-040", pid="ME.CONFIG.DROP_CHAIN", md=MD_GEN, sec="配置落地追踪 ID", table="UILevel", book="level/UIlevel.xlsx", sheet="level", key="Id", kval="<PENDING_LEVEL_IDS>", field="奖励预览字段", raw="mixed", cur="<PENDING_AUDIT>", sug="与DropId/DropId2同步映射装备、主材、副材预览", mode="MANUAL_REQUIRED", status="DRAFT", note="正式奖励链落地时按Drop/UIlevel同步流程处理。"),
        dict(id="D01-CFG-041", pid="ME.CONFIG.UNLOCK_CHAIN", md=MD_GEN, sec="配置落地追踪 ID", table="Unlock", book="unlock/Unlock.xlsx", sheet="Unlock", key="Id", kval="<PENDING_UNLOCK_IDS>", field="开放条件", raw="mixed", cur="<NOT_MAPPED>", sug="角色Lv10/20/30/40/45/50/55/60对应8档开放", mode="NEW_ROW_MANUAL", status="DRAFT", note="ID与客户端开放逻辑需一起确认。"),
        dict(id="D01-CFG-042", pid="ME.CONFIG.RUNTIME_PORTABLE", md=MD_GEN, sec="配置落地追踪 ID", table="PortableRange", book="SpecialEquipment/Portable/PortableRange.xlsx", sheet="PortableRange", key="<CODE>", kval="Portable 7表", field="<RUNTIME_CONSUMER>", raw="code", cur="仅观察到注册/导出，未发现非生成客户端消费", sug="补齐客户端/服务端合成、打造、锻造台、分解消费链", mode="MANUAL_REQUIRED", status="MANUAL_REQUIRED", note="只填Excel不会自动形成可运行功能。"),
    ]

    deprecated = [
        ("043", "SpecialEquipmentLvUp", "SpecialEquipment/SpecialEquipmentLvUp.xlsx", "SpecialEquipment", "ME.OLD.UPGRADE_EXIT", "236行旧装备升级消耗", "旧装备经验升级"),
        ("044", "SpecialEquipmentLvLimit", "SpecialEquipment/SpecialEquipmentLvLimit.xlsx", "SpecialEquipmentLvLimit", "ME.OLD.UPGRADE_EXIT", "5行旧装备等级上限", "旧装备升级上限"),
        ("045", "SpecialEquipmentGem", "SpecialEquipment/SpecialEquipmentGem.xlsx", "SpecialEquipment", "ME.OLD.HOLE_IDENTIFY_EXIT", "56行宝石配置；辅助Sheet有公式", "宝石/镶嵌"),
        ("046", "SpecialEquipmentColor", "SpecialEquipment/SpecialEquipmentColor.xlsx", "SpecialEquipmentColor", "ME.OLD.HOLE_IDENTIFY_EXIT", "10行宝石颜色", "宝石颜色"),
        ("047", "SpecialEquipmentPartFun", "SpecialEquipment/SpecialEquipmentPartFun.xlsx", "SpecialEquipmentPartFun", "ME.OLD.HOLE_IDENTIFY_EXIT", "8行孔位/颜色绑定", "孔位/镶嵌"),
        ("048", "SpecialEquipmentChip", "SpecialEquipment/SpecialEquipmentChip.xlsx", "SpecialEquipment", "ME.OLD.HOLE_IDENTIFY_EXIT", "45行芯片配置", "芯片范围未裁决，先冻结"),
    ]
    for number, table, book, sheet, parameter_id, current, label in deprecated:
        compact_rows.append(dict(
            id=f"D01-CFG-{number}", pid=parameter_id, md=MD_ECON,
            sec="当前 v2 模块表观察 / 旧系统退出", table=table, book=book, sheet=sheet,
            key="<TABLE>", kval="全部", field="<TABLE_STATE>", raw="workflow", cur=current,
            sug="<FREEZE_NO_WRITE_UNTIL_CODE_MIGRATION>", mode="MANUAL_REQUIRED",
            status="MANUAL_REQUIRED", note=f"{label}：当前版本不写；不能先删表。",
        ))

    for item in compact_rows:
        mode = item["mode"]
        status = item["status"]
        parameter_id = item["pid"]
        source_md = item["md"]
        add_draft(
            rows,
            draft_id=item["id"],
            parameter_id=parameter_id,
            source_md=source_md,
            source_section=item["sec"],
            logical_table=item["table"],
            source_workbook=item["book"],
            sheet=item["sheet"],
            key_field=item["key"],
            key_value=item["kval"],
            target_field=item["field"],
            raw_type=item["raw"],
            current_value=item["cur"],
            suggested_value=item["sug"],
            transform=item.get("transform", "按来源MD章节与配置裁决执行；不得从草案反向覆盖MD"),
            unit_rounding=item.get("unit", "按来源MD；Schema/代码项为N/A"),
            constraint=item.get("constraint", "Schema/Bean/Enum/引用在迁移方案中复查"),
            impact=item.get("impact", "涉及Schema、代码或跨表引用；正式落地前必须单独迁移并回归。"),
            formula_cell=item.get(
                "formula",
                "否" if mode == "CANDIDATE_SAFE_RECHECK" else (
                    "待复查" if mode == "MANUAL_REQUIRED" else "N/A"
                ),
            ),
            write_mode=mode,
            risk="HIGH",
            validation=item.get("validation", "完成Schema/代码/引用迁移后，执行主键、外键、概率、Luban与运行时回归。"),
            status=status,
            created_at=created_at,
            note=item["note"] + " 本批次未写正式配置。",
        )

    draft_ids = [row[1] for row in rows]
    duplicates = [item for item, count in Counter(draft_ids).items() if count > 1]
    if duplicates:
        raise RuntimeError(f"草案ID重复: {duplicates}")
    return rows


def module_update_values() -> dict[str, dict[str, str]]:
    source_attr = MD_ATTR
    source_econ = MD_ECON
    source_all = f"{MD_ATTR}; {MD_ECON}; {MD_GEN}"
    updates: dict[str, dict[str, str]] = {}

    def put(path: str, submodule: str, source: str, relation: str, status: str, note: str) -> None:
        updates[path] = {
            "业务模块": "映射装备",
            "数值子模块": submodule,
            "来源MD": source,
            "关系类型": relation,
            "映射状态": status,
            "负责人": "待指定",
            "备注": note,
        }

    put("SpecialEquipment/SpecialEquipment.xlsx", "装备实体/单主属性", source_attr, "DIRECT", "设计与旧配置冲突", "旧双主属性、无Green装备；需实体/Items一起迁移。")
    put("SpecialEquipment/SpecialEquipmentEffect.xlsx", "副属性值域", source_attr, "DIRECT", "待建立字段级映射", "结构可用，EffectId到12类属性待映射。")
    put("SpecialEquipment/SpecialEquipmentEffectDrop.xlsx", "副属性词条池", source_attr, "DIRECT", "已映射", "沿用池与权重；不重复规则需代码/Schema保证。")
    put("SpecialEquipment/SpecialEquipmentGem.xlsx", "旧宝石/镶嵌", source_econ, "DEPRECATED", "设计与旧配置冲突", "当前版本冻结，代码/UI退出前不可删。")
    put("SpecialEquipment/SpecialEquipmentGenerate.xlsx", "装备生成/宝箱", source_all, "DIRECT", "设计与旧配置冲突", "等级池、品质池可承接；宝箱来源档位阻断。")
    put("SpecialEquipment/SpecialEquipmentLvUp.xlsx", "旧装备升级", source_econ, "DEPRECATED", "设计与旧配置冲突", "236行旧装备升级，当前版本冻结。")
    put("SpecialEquipment/SpecialEquipmentMain.xlsx", "主属性曲线", source_attr, "DIRECT", "缺正式承载", "旧线性Value/Grow无法表达非线性曲线。")
    put("SpecialEquipment/SpecialEquipmentQuality.xlsx", "品质系数/旧字段", source_all, "DIRECT", "设计与旧配置冲突", "新系数可映射；分解与旧打孔/鉴定字段需拆分。")
    put("SpecialEquipment/SpecialEquipmentChip.xlsx", "芯片（范围外冻结）", source_econ, "DEPRECATED", "待确认", "沉淀未明确芯片去留，本轮不写。")
    put("SpecialEquipment/SpecialEquipmentColor.xlsx", "旧宝石颜色", source_econ, "DEPRECATED", "设计与旧配置冲突", "随宝石/镶嵌退出，当前冻结。")
    put("SpecialEquipment/SpecialEquipmentPartFun.xlsx", "旧孔位/镶嵌", source_econ, "DEPRECATED", "设计与旧配置冲突", "客户端仍读取，当前冻结。")
    put("SpecialEquipment/SpecialEquipmentLvLimit.xlsx", "旧装备升级上限", source_econ, "DEPRECATED", "设计与旧配置冲突", "新版开放进入PortableRange/Unlock。")
    put("SpecialEquipment/SpecialEquipmentSubAttrCoef.xlsx", "副属性等级倍率", source_attr, "DIRECT", "已映射", "11行改12档，新增Lv10。")
    put("SpecialEquipment/Portable/PortableRange.xlsx", "合成8档范围", source_econ, "DIRECT", "缺正式承载", "RangeId3~8悬空，且缺等级权重字段。")
    put("SpecialEquipment/Portable/PortableEquipCompound.xlsx", "合成金币/结果", source_econ, "DIRECT", "已映射", "40行金币已映射；Range/Generate引用待修。")
    put("SpecialEquipment/Portable/PortableEquipCraft.xlsx", "打造配方", source_econ, "DIRECT", "缺正式承载", "缺档位、目标等级、主材品质字段。")
    put("SpecialEquipment/Portable/PortableLvUp.xlsx", "锻造台升级经验", source_econ, "DIRECT", "设计与旧配置冲突", "结构可扩至60级，但需语义与运行时迁移。")
    put("SpecialEquipment/Portable/PortableCompoundProb.xlsx", "合成品质概率", source_econ, "DIRECT", "已映射", "公式已映射；需扩300行并裁决权重尺度。")
    put("SpecialEquipment/Portable/PortableCraftProb.xlsx", "打造品质概率", source_econ, "DIRECT", "设计与旧配置冲突", "需扩180行，且当前Sheet名错误。")
    put("SpecialEquipment/Portable/PortableExp.xlsx", "装备锻造经验", source_econ, "DIRECT", "已映射", "60行可直接形成字段差异。")
    put("globals_conf/globals_conf.xlsx", "合成全局参数", source_econ, "INDIRECT", "已映射", "DefaultSpecialEquipmentCompoundNum=9已对齐；其他旧参数另行退出。")
    put("items/Items.xlsx", "宝箱/打造材料Item", source_all, "INDIRECT", "缺正式承载", "72主材+4副材及箱引用待规划。")
    put("Drop/Drop_dungeon.xlsx", "副本奖励/材料掉落", source_all, "INDIRECT", "待确认", "仅映射Drop_dungeon输入；公式字段禁止自动写。")
    put("unlock/Unlock.xlsx", "8档开放条件", source_all, "INDIRECT", "待确认", "UnlockId待分配。")
    put("level/UIlevel.xlsx", "奖励预览", source_all, "INDIRECT", "待确认", "与DropId/DropId2同步。")
    put("level/level.xlsx", "映射副本层级", source_all, "INDIRECT", "待确认", "副本层与箱/掉落档位链待闭合。")
    put("Rubik/Craft.xlsx", "旧宝石合成（防误用）", source_econ, "DEPRECATED", "不需要数值落表", "Rubik宝石3合1不是映射装备9合1承载表。")
    return updates


def status_fill(token: str) -> PatternFill | None:
    upper = token.upper()
    if upper in {"PASS", "APPROVED", "VERIFIED", "APPLIED", "EXPORTED", "READY_FOR_REVIEW"}:
        return PASS_FILL
    if upper in {"WARN", "REVIEW", "NEEDS_REVISION", "MANUAL_REQUIRED", "PASS_WITH_WARNINGS"}:
        return WARN_FILL
    if upper in {"FAIL", "BLOCK", "REJECTED"}:
        return FAIL_FILL
    if upper in {"INFO", "NOT_STARTED", "DRAFT", "UNCLASSIFIED"}:
        return INFO_FILL
    return None


def update_workbook(workbook_path: Path, config_root: Path, report_path: Path) -> dict[str, Any]:
    if not workbook_path.is_file():
        raise FileNotFoundError(workbook_path)
    if not config_root.is_dir():
        raise NotADirectoryError(config_root)

    created_at = datetime.now().astimezone().isoformat(timespec="seconds")
    draft_rows = build_draft_rows(config_root, created_at)
    workbook_hash_before = sha256_file(workbook_path)
    workbook = load_workbook(workbook_path, data_only=False, keep_links=True)
    try:
        if "__generation_meta__" not in workbook.sheetnames:
            raise RuntimeError("缺少__generation_meta__")
        meta = workbook["__generation_meta__"]
        metadata = {text(meta.cell(row, 1).value): meta.cell(row, 2).value for row in range(1, meta.max_row + 1)}
        if metadata.get("generation_marker") != GENERATION_MARKER:
            raise RuntimeError(f"生成标记不匹配: {metadata.get('generation_marker')!r}")
        if metadata.get("official_config_write") not in (False, "False", "false", 0):
            raise RuntimeError("工作簿元数据声称曾写正式配置，拒绝继续")

        # 更新模块映射人工列，定位只依赖生成列中的相对路径。
        module_sheet = workbook["模块映射"]
        module_headers = {text(cell.value): cell.column for cell in module_sheet[1] if text(cell.value)}
        manual_headers = ["业务模块", "数值子模块", "来源MD", "关系类型", "映射状态", "负责人", "备注"]
        updates = module_update_values()
        updated_paths: list[str] = []
        source_column = module_headers["相对路径"]
        for row_number in range(2, module_sheet.max_row + 1):
            relative_path = text(module_sheet.cell(row_number, source_column).value).replace("\\", "/")
            values = updates.get(relative_path)
            if not values:
                continue
            for header in manual_headers:
                module_sheet.cell(row_number, module_headers[header], values[header])
            updated_paths.append(relative_path)
        missing_paths = sorted(set(updates) - set(updated_paths))
        if missing_paths:
            raise RuntimeError(f"模块映射中未找到目标输入: {missing_paths}")

        # 幂等替换本批次，保留未来可能存在的其他批次。
        draft_sheet = workbook["配置草案"]
        actual_headers = [text(cell.value) for cell in draft_sheet[1][: len(DRAFT_HEADERS)]]
        if actual_headers != DRAFT_HEADERS:
            raise RuntimeError("配置草案表头漂移")
        preserved_rows: list[list[Any]] = []
        for row in draft_sheet.iter_rows(min_row=2, max_col=len(DRAFT_HEADERS), values_only=True):
            if not any(value not in (None, "") for value in row):
                continue
            if text(row[0]) != BATCH_ID:
                preserved_rows.append(list(row))
        if "TblDrafts" in draft_sheet.tables:
            del draft_sheet.tables["TblDrafts"]
        if draft_sheet.max_row > 1:
            draft_sheet.delete_rows(2, draft_sheet.max_row - 1)
        all_rows = preserved_rows + draft_rows
        for row in all_rows:
            draft_sheet.append(row)

        template_style = copy(module_sheet.cell(2, 1)._style)
        for row_number in range(2, len(all_rows) + 2):
            for column in range(1, len(DRAFT_HEADERS) + 1):
                cell = draft_sheet.cell(row_number, column)
                cell._style = copy(template_style)
                cell.font = copy(INPUT_FONT)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            status_cell = draft_sheet.cell(row_number, DRAFT_HEADERS.index("草案状态") + 1)
            fill = status_fill(text(status_cell.value))
            if fill:
                status_cell.fill = copy(fill)
        table = Table(displayName="TblDrafts", ref=f"A1:AA{len(all_rows) + 1}")
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        draft_sheet.add_table(table)
        # Excel 12 会移除与工作表级 AutoFilter 重叠的 Table。
        # 有业务行时由 TblDrafts 自带筛选；5000行输入范围继续由数据验证维护。
        draft_sheet.auto_filter.ref = None

        # 工作簿阶段与门禁只反映草案已生成；绝不伪造批准或落地状态。
        summary = workbook["使用说明"]
        summary["A2"] = "本工作簿已生成映射装备配置草案批次；审批记录与落地回执仍为空。用户批准当前批次并另行授权SVN写入前，Vgame Datas保持只读。"
        summary["B12"] = f"S2 配置草案 {BATCH_ID} 已生成，等待用户审阅"
        gate_sheet = workbook["工作流门禁"]
        gate_status = {"G1": "PASS", "G2": "PASS_WITH_WARNINGS", "G3": "PASS_WITH_WARNINGS"}
        for row_number in range(2, gate_sheet.max_row + 1):
            gate = text(gate_sheet.cell(row_number, 1).value)
            if gate not in gate_status:
                continue
            gate_sheet.cell(row_number, 3, gate_status[gate])
            gate_sheet.cell(row_number, 6, f"{BATCH_ID} / 配置草案与验证报告")
            gate_sheet.cell(row_number, 3).fill = copy(status_fill(gate_status[gate]) or INFO_FILL)

        meta_values = {
            "draft_batch_id": BATCH_ID,
            "draft_batch_version": BATCH_VERSION,
            "draft_generated_at": created_at,
            "official_config_write": False,
        }
        meta_index = {text(meta.cell(row, 1).value): row for row in range(1, meta.max_row + 1)}
        for key, value in meta_values.items():
            row_number = meta_index.get(key)
            if row_number is None:
                row_number = meta.max_row + 1
                meta.cell(row_number, 1, key)
            meta.cell(row_number, 2, value)

        try:
            workbook.calculation.calcMode = "auto"
            workbook.calculation.fullCalcOnLoad = True
            workbook.calculation.forceFullCalc = True
        except Exception:
            pass

        backup_dir = workbook_path.parent / "_backups"
        backup_dir.mkdir(parents=True, exist_ok=True)
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        backup_path = backup_dir / f"{workbook_path.stem}_{timestamp}{workbook_path.suffix}"
        shutil.copy2(workbook_path, backup_path)

        temporary_path = workbook_path.with_name(
            f".{workbook_path.stem}.{uuid.uuid4().hex}.tmp{workbook_path.suffix}"
        )
        workbook.save(temporary_path)
    finally:
        workbook.close()

    check = None
    try:
        check = load_workbook(temporary_path, read_only=True, data_only=False, keep_links=True)
        if "配置草案" not in check.sheetnames or "__generation_meta__" not in check.sheetnames:
            raise RuntimeError("临时工作簿重开后缺Sheet")
    finally:
        if check is not None:
            check.close()
    os.replace(temporary_path, workbook_path)

    scope_rows = []
    for row in draft_rows:
        normalized = list(row)
        normalized[DRAFT_HEADERS.index("创建时间")] = "<TIMESTAMP>"
        scope_rows.append(normalized)
    scope_hash = hashlib.sha256(
        json.dumps(scope_rows, ensure_ascii=False, separators=(",", ":"), default=str).encode("utf-8")
    ).hexdigest()
    status_counter = Counter(text(row[DRAFT_HEADERS.index("草案状态")]) for row in draft_rows)
    mode_counter = Counter(text(row[DRAFT_HEADERS.index("写入方式")]) for row in draft_rows)
    table_counter = Counter(text(row[DRAFT_HEADERS.index("逻辑表")]) for row in draft_rows)
    report = {
        "batch_id": BATCH_ID,
        "batch_version": BATCH_VERSION,
        "created_at": created_at,
        "workbook": str(workbook_path),
        "workbook_backup": str(backup_path),
        "workbook_sha256_before": workbook_hash_before,
        "workbook_sha256_after": sha256_file(workbook_path),
        "batch_scope_sha256": scope_hash,
        "draft_row_count": len(draft_rows),
        "preserved_other_batch_rows": len(preserved_rows),
        "module_mapping_count": len(updated_paths),
        "module_mapping_paths": sorted(updated_paths),
        "counts_by_status": dict(sorted(status_counter.items())),
        "counts_by_write_mode": dict(sorted(mode_counter.items())),
        "counts_by_logical_table": dict(sorted(table_counter.items())),
        "source_snapshot_hash": metadata.get("source_snapshot_hash"),
        "config_root": str(config_root),
        "official_config_write": False,
        "approval_rows_written": 0,
        "receipt_rows_written": 0,
        "luban_executed": False,
    }
    report_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    return report


def main() -> int:
    parser = argparse.ArgumentParser(description="生成映射装备配置草案 D01（不写正式配置）")
    parser.add_argument("--workbook", default=str(DEFAULT_WORKBOOK))
    parser.add_argument("--config-root", default=str(DEFAULT_CONFIG_ROOT))
    parser.add_argument("--report", default=str(DEFAULT_REPORT))
    args = parser.parse_args()

    try:
        report = update_workbook(
            Path(args.workbook).resolve(),
            Path(args.config_root).resolve(),
            Path(args.report).resolve(),
        )
    except Exception as exc:
        print(f"[FAIL] {type(exc).__name__}: {exc}", file=sys.stderr)
        return 1
    print("[OK] 映射装备配置草案已生成；正式Datas未写入")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
