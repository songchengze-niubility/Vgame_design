# -*- coding: utf-8 -*-
"""CHG-026 重建：删除112个任务Drop包的全部旧行，按v15权威值重新生成224行。"""
import sys
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
CFG = HERE.parents[1] / "06_最终配置" / "任务"
V15 = Path("D:/Vgame_design/output/14日任务总表-评审v15-2026-08-31.xlsx")

# v15 权威奖励
wb = load_workbook(V15, read_only=True, data_only=True)
rows = list(wb["14日任务总表"].iter_rows(values_only=True))[1:]
auth = {}
for r in rows:
    tid = r[2]
    if isinstance(tid, int) and 5101 <= tid <= 6510:
        auth[tid] = (int(r[10]), int(r[11]))
wb.close()

HUNDRED = {58: 8, 59: 9, 61: 10, 62: 11, 63: 12, 64: 13, 65: 14}
DAY_BASE = {8: 91000102, 9: 91000108, 10: 91000114, 11: 91000120,
            12: 91000126, 13: 91000132, 14: 91000138}
def pack_of(tid):
    if tid <= 5710:
        day = (tid - 5101) // 100 + 1
        seq = (tid - 5101) % 100 + 1
        return 91000000 + (day - 1) * 10 + seq
    day = HUNDRED[tid // 100]
    return DAY_BASE[day] + (tid % 100) - 1

src = CFG / "Drop_TimeLimitedActivity.xlsx"
wb = load_workbook(src)
ws = wb["Drop"]
head = [c.value for c in next(ws.iter_rows(max_row=1))]
cID = head.index("ID") + 1
cDrop = head.index("DropId") + 1
cName = head.index("DropName") + 1
cId2 = head.index("DropId2") + 1
cType = head.index("DropType") + 1
cMinN = head.index("DropMinNum") + 1
cMaxN = head.index("DropMaxNum") + 1
cItem = head.index("DropItem_id") + 1
cMin = head.index("DropItemNumMin") + 1
cMax = head.index("DropItemNumMax") + 1
cW = head.index("DropWeight") + 1
lbl = next((c for c in range(1, ws.max_column + 1) if head[c-1] is None), None)

def main():
    # 1) 删除区间内所有旧行（自底向上）
    del_rows = []
    for r in range(5, ws.max_row + 1):
        did = ws.cell(row=r, column=cDrop).value
        rid = ws.cell(row=r, column=cID).value
        if isinstance(did, int) and (91000000 <= did <= 91000070 or 91000102 <= did <= 91000143):
            del_rows.append(r)
        elif isinstance(did, int) and did == 91000101 and isinstance(rid, int) and rid >= 400461:
            del_rows.append(r)
    for r in sorted(del_rows, reverse=True):
        ws.delete_rows(r)
    print("删除旧行:", len(del_rows))

    # 2) 按 v15 重建 224 行
    ITEM = {1:"猎鹰金币",6:"沿途随笔",19:"勇士勋章",25:"映射之匙",28:"校准元件",33:"漫游邮戳",
            100:"低阶记录盘",102:"高阶记录盘",106:"旅券",107:"收容凭证",108:"泛黄档案",
            109:"录像带",111:"限时旅券",301:"商品筹码",1074:"房车零件",1156:"高级指南材料",
            3010:"1级宝石掉落箱",3011:"2级宝石掉落箱",3015:"普通能效自选箱",3027:"指南材料自选箱",
            51001:"虚妄之镜",311029:"凯兰"}
    next_id = 400461
    for tid in sorted(auth):
        if tid <= 5710:
            day = (tid - 5101) // 100 + 1
            seq = (tid - 5101) % 100 + 1
            did = 91000000 + (day - 1) * 10 + seq
            name = f"开服挑战任务{seq:02d}"
        else:
            day = HUNDRED[tid // 100]
            seq = tid % 100
            did = DAY_BASE[day] + seq - 1
            name = f"开服挑战第{day}天任务{seq:02d}"
        item, num = auth[tid]
        for it, nm in [(32, 10), (item, num)]:
            row = [None] * ws.max_column
            row[cID - 1] = next_id
            row[cDrop - 1] = did
            if lbl:
                row[lbl - 1] = f"{name}-{it}"
            row[cName - 1] = name
            row[cId2 - 1] = 0
            row[cType - 1] = "fixed_"
            row[cMinN - 1] = 1
            row[cMaxN - 1] = 1
            row[cItem - 1] = it
            row[cMin - 1] = nm
            row[cMax - 1] = nm
            row[cW - 1] = 10000
            ws.append(row)
            next_id += 1
    wb.save(src)
    print("重建224行, ID至", next_id - 1)

    # 3) 复核
    wb2 = load_workbook(src, read_only=True, data_only=True)
    groups = {}
    from collections import defaultdict
    g = defaultdict(list)
    for r in wb2["Drop"].iter_rows(min_row=5, values_only=True):
        if r and len(r) > 2 and isinstance(r[2], int) and (91000001 <= r[2] <= 91000070 or 91000102 <= r[2] <= 91000143):
            g[r[2]].append((r[10], r[11]))
    ok = all(sorted(g[pack_of(t)]) == sorted([(32, 10), auth[t]]) for t in auth)
    print("112包全部与v15一致:", ok)
    print("一致性复核完成")


if __name__ == '__main__':
    main()
