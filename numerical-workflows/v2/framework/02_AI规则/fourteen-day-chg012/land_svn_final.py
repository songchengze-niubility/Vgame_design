# -*- coding: utf-8 -*-
"""CHG-026 终态落SVN：06最终配置5张表 -> SVN工作副本，并做SVN侧独立复验。

只改工作副本，不执行 svn commit。
"""
import json
import shutil
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
FRAME = HERE.parents[1]
CFG = FRAME / "06_最终配置" / "任务"
RUN = FRAME / "05_验算核算" / "十四日活动CHG012_20260831" / "SVN落表-20260903"
SVN = Path("D:/Vgame/Config/GameConfig/Datas")

FILES = [
    (CFG / "Task.xlsx", SVN / "task" / "Task.xlsx", "task/Task.xlsx"),
    (CFG / "TaskTarget.xlsx", SVN / "task" / "TaskTarget.xlsx", "task/TaskTarget.xlsx"),
    (CFG / "FourteenDayActivity.xlsx", SVN / "Activity" / "FourteenDayActivity.xlsx", "Activity/FourteenDayActivity.xlsx"),
    (CFG / "ActiveReward.xlsx", SVN / "Activity" / "ActiveReward.xlsx", "Activity/ActiveReward.xlsx"),
    (CFG / "Drop_TimeLimitedActivity.xlsx", SVN / "Drop" / "Drop_TimeLimitedActivity.xlsx", "Drop/Drop_TimeLimitedActivity.xlsx"),
]

V15 = Path("D:/Vgame_design/output/14日任务总表-评审v15-2026-08-31.xlsx")
ITEM = {1: "猎鹰金币", 6: "沿途随笔", 19: "勇士勋章", 25: "映射之匙", 28: "校准元件",
        33: "漫游邮戳", 100: "低阶记录盘", 102: "高阶记录盘", 106: "旅券", 107: "收容凭证",
        108: "泛黄档案", 109: "录像带", 111: "限时旅券", 301: "商品筹码", 1074: "房车零件",
        1156: "高级指南材料", 3010: "1级宝石掉落箱", 3011: "2级宝石掉落箱",
        3015: "普通能效自选箱", 3027: "指南材料自选箱", 51001: "虚妄之镜", 311029: "凯兰"}
DAY_BASE = {8: 91000102, 9: 91000108, 10: 91000114, 11: 91000120,
            12: 91000126, 13: 91000132, 14: 91000138}
HUNDRED = {58: 8, 59: 9, 61: 10, 62: 11, 63: 12, 64: 13, 65: 14}
T0 = {8: 5801, 9: 5901, 10: 6101, 11: 6201, 12: 6301, 13: 6401, 14: 6501}


def pack_of(tid):
    if tid <= 5710:
        day = (tid - 5101) // 100 + 1
        return 91000000 + (day - 1) * 10 + (tid - 5101) % 100 + 1
    return DAY_BASE[HUNDRED[tid // 100]] + (tid % 100) - 1


def col_index(header, name):
    return head_map[id(header)][name]


head_map = {}


def reg_head(header):
    head_map[id(header)] = {h: i for i, h in enumerate(header) if h}


def main():
    RUN.mkdir(parents=True, exist_ok=True)
    (RUN / "backup-svn").mkdir(parents=True, exist_ok=True)

    # 1) 备份 + 复制
    for src, dst, rel in FILES:
        bdir = RUN / "backup-svn" / Path(rel).parent
        bdir.mkdir(parents=True, exist_ok=True)
        shutil.copy2(dst, bdir / Path(rel).name)
        shutil.copy2(src, dst)
        print("落表:", rel)
    (RUN / "copied.json").write_text(
        json.dumps({"batch": "SVN-LAND-CHG026-FINAL", "created": datetime.now().isoformat(),
                    "files": [rel for _, _, rel in FILES]}, ensure_ascii=False, indent=2),
        encoding="utf-8")

    # 2) SVN 侧独立复验
    errors = []
    wb = load_workbook(SVN / "Drop" / "Drop_TimeLimitedActivity.xlsx", read_only=True, data_only=True)
    groups = defaultdict(list)
    for r in wb["Drop"].iter_rows(min_row=5, values_only=True):
        if r and len(r) > 2 and isinstance(r[2], int):
            if 91000001 <= r[2] <= 91000070 or 91000102 <= r[2] <= 91000143 or 91001001 <= r[2] <= 91001011:
                groups[r[2]].append((r[10], r[11]))
    wb.close()

    wbv = load_workbook(V15, read_only=True, data_only=True)
    auth = {}
    for r in list(wbv["14日任务总表"].iter_rows(values_only=True))[1:]:
        tid = r[2]
        if isinstance(tid, int) and 5101 <= tid <= 6510:
            auth[tid] = (int(r[10]), int(r[11]))
    wbv.close()
    for tid in sorted(auth):
        want = sorted([(32, 10), auth[tid]])
        got = sorted(groups.get(pack_of(tid), []))
        if want != got:
            errors.append(f"Drop{pack_of(tid)} 与v15不一致: {got}")

    # 任务/目标链路
    wt = load_workbook(SVN / "task" / "Task.xlsx", read_only=True, data_only=True)
    trows = list(wt["Task"].iter_rows(values_only=True))
    wt.close()
    th = trows[0]
    m = {h: i for i, h in enumerate(th) if h}
    n7 = n14 = 0
    for r in trows[4:]:
        if r and isinstance(r[m["Id"]], int):
            v = r[m["Id"]]
            if 5101 <= v <= 5710:
                n7 += 1
            elif (5801 <= v <= 5806 or 5901 <= v <= 5906 or 6101 <= v <= 6106
                  or 6201 <= v <= 6206 or 6301 <= v <= 6306 or 6401 <= v <= 6406
                  or 6501 <= v <= 6506):
                n14 += 1
    if n7 != 70 or n14 != 42:
        errors.append(f"SVN任务数: 前7天{n7}(应70) 后7天{n14}(应42)")

    wf = load_workbook(SVN / "Activity" / "FourteenDayActivity.xlsx", read_only=True, data_only=True)
    fda = list(wf.worksheets[0].iter_rows(values_only=True))
    wf.close()
    days = sorted(r[3] for r in fda[4:] if len(r) > 4 and r[2] == 1 and r[3] is not None)
    if days != list(range(1, 15)):
        errors.append(f"SVN FourteenDay OpenDay: {days}")

    wa = load_workbook(SVN / "Activity" / "ActiveReward.xlsx", read_only=True, data_only=True)
    ar = list(wa["ActiveReward"].iter_rows(values_only=True))
    wa.close()
    a_head = ar[0]
    pts = sorted(r[a_head.index("ActivePoint")] for r in ar[4:]
                 if r and len(r) > 4 and r[a_head.index("ActivityId")] == 1 and r[a_head.index("ActivePoint")] is not None)
    if pts != list(range(100, 1101, 100)):
        errors.append(f"SVN ActiveReward 档位: {pts}")

    result = {
        "batch": "SVN-LAND-CHG026-FINAL",
        "generated_at": datetime.now().isoformat(),
        "files": [rel for _, _, rel in FILES],
        "errors": errors,
        "status": "PASS" if not errors else "FAIL",
    }
    (RUN / "validation-svn.json").write_text(
        json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))
    if result["status"] != "PASS":
        raise SystemExit(1)


if __name__ == "__main__":
    main()
