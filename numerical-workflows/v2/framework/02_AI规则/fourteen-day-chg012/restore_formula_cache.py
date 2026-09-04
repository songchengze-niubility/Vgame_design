# -*- coding: utf-8 -*-
"""CHG-012 修复：openpyxl 保存丢失公式缓存值，从备份把缓存 <v> 回填进 Drop 表 XML。"""
import re
import shutil
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, column_index_from_string

HERE = Path(__file__).resolve().parent
RUN = HERE.parents[1] / "05_验算核算" / "十四日活动CHG012_20260831"
BACKUP = RUN / "backup" / "任务" / "Drop_TimeLimitedActivity.xlsx"
CURRENT = HERE.parents[1] / "06_最终配置" / "任务" / "Drop_TimeLimitedActivity.xlsx"
NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"


def cell_ref(row, col):
    return f"{get_column_letter(col)}{row}"


# 1) 备份的缓存值矩阵 + 当前的公式单元格清单
wb_v = load_workbook(BACKUP, read_only=True, data_only=True)["Drop"]
cache = {}
for ri, r in enumerate(wb_v.iter_rows(values_only=True), 1):
    for ci, v in enumerate(r, 1):
        if v is not None:
            cache[(ri, ci)] = v

wb_f = load_workbook(CURRENT, read_only=True)["Drop"]
formula_cells = set()
for ri, r in enumerate(wb_f.iter_rows(values_only=True), 1):
    for ci, v in enumerate(r, 1):
        if isinstance(v, str) and v.startswith("="):
            formula_cells.add((ri, ci))
print("formula cells:", len(formula_cells))

# 2) 找到 Drop 表对应的 sheet XML
tmp = CURRENT.with_suffix(".patched.xlsx")
with zipfile.ZipFile(CURRENT) as zin:
    wb_xml = ET.fromstring(zin.read("xl/workbook.xml"))
    rels_xml = ET.fromstring(zin.read("xl/_rels/workbook.xml.rels"))
    rid_map = {r.get("Id"): r.get("Target") for r in rels_xml}
    ns_r = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
    sheet_target = None
    for sh in wb_xml.iter(f"{{{NS}}}sheet"):
        if sh.get("name") == "Drop":
            target = rid_map[sh.get(ns_r)]
            sheet_target = "xl/" + target.lstrip("/").replace("xl/", "", 1) if target.startswith("/xl/") else ("xl/" + target if not target.startswith("xl/") else target)
    if sheet_target is None:
        raise SystemExit("未找到 Drop sheet")
    print("sheet xml:", sheet_target)

    patched_cache_missing = []
    out = {}
    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename == sheet_target:
            root = ET.fromstring(data)
            for c in root.iter(f"{{{NS}}}c"):
                f_node = c.find(f"{{{NS}}}f")
                if f_node is None:
                    continue
                ref = c.get("r")
                m = re.match(r"([A-Z]+)(\d+)", ref)
                key = (int(m.group(2)), column_index_from_string(m.group(1)))
                if key not in cache:
                    patched_cache_missing.append(ref)
                    continue
                v_node = c.find(f"{{{NS}}}v")
                if v_node is None:
                    v_node = ET.SubElement(c, f"{{{NS}}}v")
                v_node.text = repr(cache[key]) if not float(cache[key]).is_integer() else str(int(cache[key]))
            ET.register_namespace("", NS)
            data = ET.tostring(root, encoding="UTF-8", xml_declaration=True)
        out[item.filename] = data
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            zout.writestr(item, out[item.filename])

if patched_cache_missing:
    raise SystemExit(f"缓存缺失: {patched_cache_missing[:5]}")

shutil.move(tmp, CURRENT)
print("patched ->", CURRENT)

# 3) 复核：data_only 读当前文件，公式单元格缓存应齐全
wb_check = load_workbook(CURRENT, read_only=True, data_only=True)["Drop"]
missing = 0
for ri, r in enumerate(wb_check.iter_rows(values_only=True), 1):
    for ci, v in enumerate(r, 1):
        if (ri, ci) in formula_cells and v is None:
            missing += 1
print("cached missing after patch:", missing)
assert missing == 0
