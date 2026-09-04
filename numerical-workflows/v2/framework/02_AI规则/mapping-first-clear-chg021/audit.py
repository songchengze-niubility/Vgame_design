"""CHG-021 source extraction and independent verification; never saves XLSX."""
import argparse
import hashlib
import importlib
import json
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path
from openpyxl import load_workbook

HERE=Path(__file__).resolve().parent
ROOT=HERE.parent.parent
OUT=ROOT/'06_最终配置/映射装备'
MD=ROOT/'03_沉淀MD/经济/05_玩法_映射副本经济沉淀-2026-07-08.md'
FORMAL=Path('D:/Vgame/Config/GameConfig/Datas')
FILES=['Drop_dungeon.xlsx','UIlevel.xlsx','Items.xlsx','00_配置清单.xlsx']
SHA=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()


def read(p,sheet):
    w=load_workbook(p,read_only=True,data_only=False)
    rs=[list(r) for r in w[sheet].iter_rows(values_only=True)]
    w.close()
    return rs


def tidy(r):
    r=list(r)
    while r and r[-1] is None:r.pop()
    return r


def groups(rs,col):
    d=defaultdict(list)
    for r in rs[4:]:
        if len(r)>col and r[col] is not None:d[r[col]].append(tidy(r))
    return d


def rules():
    colors={'绿':0,'蓝':1,'紫':2,'金':3,'红':4}
    result={}
    for line in MD.read_text(encoding='utf-8-sig').splitlines():
        m=re.match(r'^\| (\d+) \| 装备副本钥匙 x1 \| ([绿蓝紫金红])宝箱 x1(.*?) \|',line)
        if m:
            l=int(m[1]);result[l]={'layer':l,'box':4201+(l-1)*5+colors[m[2]],'quality':m[2],
                                 'preview':4372+l,'first_archive':line.split('|')[3].strip()}
    assert set(result)==set(range(1,31))
    return result


parser=argparse.ArgumentParser(description=__doc__)
parser.add_argument('mode',choices=['prepare','verify','published'])
args=parser.parse_args()
expected=rules()
if args.mode=='prepare':
    drops=groups(read(OUT/FILES[0],'Drop'),1)
    first=[]
    for l,e in expected.items():
        rs=drops[20000009+l]
        key=[r for r in rs if r[9]==25]
        assert len(key)==1 and key[0][10]==key[0][11]==1
        box=list(next((r for r in rs if 4201<=r[9]<=4350),key[0]))
        box[2]='首通'+e['quality']+'箱';box[4]=0;box[5]='fixed_';box[6]=box[7]=1
        box[9]=e['box'];box[10]=box[11]=1;box[12]=10000;box[15]='VALUE_READY'
        materials=[r for r in rs if r[9]!=25 and not 4201<=r[9]<=4350]
        assert all(1080<=r[9]<=1155 and r[10]==r[11]==1 for r in materials)
        first.extend(key+[box]+materials)
        e['materials']=[r[9] for r in materials]
        e['firstPreview']=';'.join(['25,1',str(e['preview'])+',1']+[str(r[9])+',1' for r in materials])
    assert len(first)==75
    olditems=groups(read(OUT/'Items.xlsx','items'),0)
    fi=read(FORMAL/'items/Items.xlsx','items')
    newitems=[]
    for r in fi[4:]:
        if r[0]!='##' and r[1] in {e['preview'] for e in expected.values()} and r[1] not in olditems:
            newitems.append([r[1],r[2],r[3],r[4],r[5],r[6],r[7],r[8],
                             {1:'Green',2:'Blue',3:'Purple',4:'Orange',5:'Red'}[r[9]],r[10],r[11],
                             r[12],r[13],r[14],r[15],'true','CHG-021：正式已有首通占位引用'])
    assert len(newitems)==27
    # Existing formal display text may itself be stale. Explicit quality words must follow the MD.
    for r in newitems:
        e=expected[r[0]-4372]
        r[9]=re.sub(r'[绿蓝紫金红]色映射装备',e['quality']+'色映射装备',r[9])
    snapshot={'revision':'CHG-021','rules':expected,'firstRows':first,'newItems':newitems,
              'mdHash':SHA(MD),'candidateHashes':{p.name:SHA(p) for p in OUT.glob('*.xlsx') if not p.name.startswith('~$')},
              'formalHashes':{str(FORMAL/p):SHA(FORMAL/p) for p in ['Drop/Drop_dungeon.xlsx','level/UIlevel.xlsx','items/Items.xlsx']}}
    (HERE/'source.json').write_text(json.dumps(snapshot,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({'firstRows':75,'newItems':27,'expectedKeys':30,'expectedBoxes':30}))
    raise SystemExit(0)

meta_path=ROOT/'05_验算核算/mapping-equip-first-clear-CHG021.json' if args.mode=='published' else HERE/'repair.json'
meta=json.loads(meta_path.read_text(encoding='utf-8'));source=meta['source'];backup=Path(meta['backup'])
output=OUT if args.mode=='published' else Path(meta['staging'])
checks=[]
def check(n,ok):checks.append({'name':n,'result':'PASS' if ok else 'FAIL'})
before=groups(read(backup/'Drop_dungeon.xlsx','Drop'),1)
after=groups(read(output/'Drop_dungeon.xlsx','Drop'),1)
check('DropId集合保持',set(before)==set(after))
check('Drop有效行695',sum(map(len,after.values()))==695)
for id,rs in after.items():
    if not 20000010<=id<=20000039:check(f'常规及其他Drop{id}逐字段不变',rs==before[id])
    check(f'Drop{id}引用完整',all(not r[4] or r[4] in after for r in rs))
totalkeys=totalboxes=0
for l,e in expected.items():
    rs=after[20000009+l];old=before[20000009+l]
    keys=[r for r in rs if r[9]==25];boxes=[r for r in rs if 4201<=r[9]<=4350]
    check(f'L{l}真实首通钥匙固定1',len(keys)==1 and keys[0][10]==keys[0][11]==1)
    check(f'L{l}真实首通宝箱品质数量匹配MD',len(boxes)==1 and boxes[0][9]==e['box'] and boxes[0][10]==boxes[0][11]==1)
    check(f'L{l}首通无随机和子池',all(r[4]==0 and r[5]=='fixed_' and r[6]==r[7]==1 and r[12]==10000 for r in rs))
    nonbox=lambda rows:[r for r in rows if not 4201<=r[9]<=4350]
    check(f'L{l}钥匙与材料保持原值',nonbox(rs)==nonbox(old))
    totalkeys+=sum(r[10] for r in keys);totalboxes+=sum(r[10] for r in boxes)
check('30层钥匙总量30',totalkeys==30);check('30层首通箱总量30',totalboxes==30)
check('箱品质分布2绿7蓝10紫8金3红',Counter((r[9]-4201)%5 for l in expected for r in after[20000009+l] if 4201<=r[9]<=4350)==Counter({0:2,1:7,2:10,3:8,4:3}))
old_ui=groups(read(backup/'UIlevel.xlsx','level'),1);ui=groups(read(output/'UIlevel.xlsx','level'),1)
check('150条UI身份保留',set(old_ui)==set(ui) and len(ui)==150)
for id,rs in ui.items():
    l=(id-250000)//10;e=source['rules'][str(l)];old=[r.copy() for r in old_ui[id]]
    old[0][17]=e['firstPreview']
    check(f'UI{id}仅FirstReward变化且不乘品质变体',rs==old)
    check(f'UI{id}钥匙箱材料都是固定1',all(int(v.split(',')[1])==1 for v in rs[0][17].split(';')))
olditems=groups(read(backup/'Items.xlsx','items'),0);items=groups(read(output/'Items.xlsx','items'),0)
check('Items新增27条首通引用',set(items)-set(olditems)=={r[0] for r in source['newItems']})
for r in source['newItems']:check(f'新增Item{r[0]}字段与校正后来源一致',items[r[0]]==[tidy(r)])
for id,rs in olditems.items():check(f'Item{id}既有字段保留',items[id]==rs)
for l,e in expected.items():
    item=items[e['preview']][0]
    check(f'L{l}首通展示品质正确',item[8]=={'绿':'Green','蓝':'Blue','紫':'Purple','金':'Orange','红':'Red'}[e['quality']])
    explicit=re.findall(r'([绿蓝紫金红])色映射装备',item[9])
    check(f'L{l}首通描述不误报品质',not explicit or explicit==[e['quality']])
    check(f'L{l}展示道具没有发放引用',not any(r[9]==e['preview'] for rs in after.values() for r in rs))
for file in FILES:
    oldw=load_workbook(backup/file,data_only=False);neww=load_workbook(output/file,data_only=False)
    fm=lambda w:{(s.title,c.coordinate):c.value for s in w for row in s for c in row if c.data_type=='f'}
    check(file+'所有公式不变',fm(oldw)==fm(neww));oldw.close();neww.close()
    w=load_workbook(output/file,data_only=True)
    check(file+'无公式错误',not any(c.data_type=='e' for s in w for row in s for c in row));w.close()
for p,d in source['formalHashes'].items():check('SVN未改 '+Path(p).name,SHA(Path(p))==d)
for f,d in source['candidateHashes'].items():
    if f not in FILES:check('其他候选未改 '+f,SHA(OUT/f)==d)
check('文档数值依据未改变',SHA(MD)==source['mdHash'])

# Execute only pure generator functions. Do not run workbook writers or full production.
sys.path.insert(0,str(HERE.parent))
gen=importlib.import_module('generate_mapping_equip_full_package')
layers=[]
for l,e in expected.items():
    mats=source['rules'][str(l)]['materials']
    tier=gen.CRAFT_TIERS[(mats[0]-1080)//9] if mats else 10
    layer={'layer':l,'first_archive':e['first_archive'],'craft_tier':tier}
    layers.append(layer)
    for variant in range(1,6):
        check(f'生成器L{l}变体{variant}首通不乘倍率',gen.first_preview_items(layer,variant)==source['rules'][str(l)]['firstPreview'])
rows,_=gen.first_clear_rows({'layers':layers})
check('生成器首通75条无旧随机池',len(rows)==75 and all(20000010<=r[1]<=20000039 and r[4]==0 for r in rows))
for l,e in expected.items():
    check(f'生成器L{l}宝箱按文档',sum(r[1]==20000009+l and r[9]==e['box'] and r[10]==r[11]==1 for r in rows)==1)
summary={'checks':len(checks),'pass':sum(c['result']=='PASS' for c in checks),'fail':sum(c['result']=='FAIL' for c in checks)}
report={'revision':'CHG-021','output':str(output),'backup':str(backup),'summary':summary,'checks':checks,
        'files':{f:SHA(output/f) for f in FILES},'keyTotal':totalkeys,'boxTotal':totalboxes}
(ROOT/'05_验算核算/mapping-equip-first-clear-CHG021-validation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(summary))
for c in checks:
    if c['result']=='FAIL':print(c)
raise SystemExit(bool(summary['fail']))
