# -*- coding: utf-8 -*-
"""从本次生成物组装一个全新的 ``04_数值文档`` staging。

这个脚本有意不读取、复制或增量更新现有 ``04_数值文档``。只有在外层流程
完成动态验算后，staging 才能被原子发布为正式目录。
"""

from __future__ import annotations

import argparse
import hashlib
import json
import shutil
from datetime import datetime
from pathlib import Path


V2_DIR = Path(__file__).resolve().parent.parent
FINAL_DIR = V2_DIR / "04_数值文档"

# 最终交付面向 QA 测试对照：只发布可直接验证的数值表；经济产销核算类产物
# （模块产销表、全资源产销总表、资源养成阅读版）保留在工作流内部，不再进入 04。
FIXED_MAPPINGS = [
    ("养成/生成_养成总览.xlsx", "养成/养成总览_数值文档.xlsx"),
    ("养成/生成_角色升级.xlsx", "养成/角色升级_数值文档.xlsx"),
    ("养成/生成_角色星盘.xlsx", "养成/角色星盘_数值文档.xlsx"),
    ("养成/生成_角色升星.xlsx", "养成/角色升星_数值文档.xlsx"),
    ("养成/生成_玩家战力曲线.xlsx", "养成/玩家战力曲线_数值文档.xlsx"),
    ("养成/生成_旅券属性.xlsx", "养成/旅券属性_数值文档.xlsx"),
    ("养成/生成_专武.xlsx", "养成/专武_数值文档.xlsx"),
    ("养成/生成_炮台.xlsx", "养成/炮台_数值文档.xlsx"),
    ("养成/生成_房车.xlsx", "养成/房车养成_数值文档.xlsx"),
    ("养成/生成_映射装备合成制作8档.xlsx", "养成/映射装备合成制作8档_数值文档.xlsx"),
    ("养成/生成_映射装备数值总表.xlsx", "大模块/映射装备数值总表_数值文档.xlsx"),
    ("战斗/生成_怪物标准模型.xlsx", "战斗/怪物标准模型_数值文档.xlsx"),
    ("战斗/镜像_战斗知识工具页.xlsx", "战斗/战斗知识工具页_镜像.xlsx"),
    ("经济/生成_角色升级成本.xlsx", "经济/角色升级成本_数值文档.xlsx"),
    ("经济/镜像_经济知识工具页.xlsx", "经济/经济知识工具页_镜像.xlsx"),
    ("经济/生成_爬塔首通挂机奖励.xlsx", "玩法/爬塔首通挂机奖励_数值文档.xlsx"),
    ("节奏/生成_开启节奏总览.xlsx", "节奏/开放节奏总览_数值文档.xlsx"),
    ("节奏/生成_主角等级单源.xlsx", "节奏/主角等级_数值文档.xlsx"),
    ("节奏/全玩法难度线.xlsx", "节奏/全玩法难度线_数值文档.xlsx"),
    ("节奏/生成_无限挑战难度线.xlsx", "节奏/无限挑战难度线_数值文档.xlsx"),
    ("节奏/生成_爬塔难度线_天数层数.xlsx", "节奏/爬塔难度线_数值文档.xlsx"),
]

EXPECTED_MODULE_WORKBOOKS = 22
EXPECTED_FINAL_WORKBOOKS = len(FIXED_MAPPINGS) + EXPECTED_MODULE_WORKBOOKS

# 系统工作簿按 QA 口径发布：去掉产销核算/汇总页签，只保留数据页。
MODULE_ACCOUNTING_SHEETS = ("阅读摘要", "产销阅读", "资源汇总", "机器明细")
MODULE_FALLBACK_SHEET = "产销阅读"
MODULE_FALLBACK_TITLE = "奖励与消耗"


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def write_qa_module_copy(source: Path, destination: Path) -> None:
    """把经济模块工作簿复制为 QA 版：删除核算页签；若只剩空壳则保留逐条件清单页。"""

    import openpyxl

    wb = openpyxl.load_workbook(source)
    only_accounting = set(wb.sheetnames) <= set(MODULE_ACCOUNTING_SHEETS)
    for name in [n for n in wb.sheetnames if n in MODULE_ACCOUNTING_SHEETS]:
        if only_accounting and name == MODULE_FALLBACK_SHEET:
            continue
        del wb[name]
    if not wb.sheetnames:
        raise RuntimeError("模块工作簿没有可保留的数据页签: %s" % source)
    if set(wb.sheetnames) == {MODULE_FALLBACK_SHEET}:
        wb[MODULE_FALLBACK_SHEET].title = MODULE_FALLBACK_TITLE
    destination.parent.mkdir(parents=True, exist_ok=True)
    wb.save(destination)
    wb.close()


def mappings() -> list[tuple[Path, Path]]:
    result = [(V2_DIR / src, Path(dst)) for src, dst in FIXED_MAPPINGS]
    module_dir = V2_DIR / "经济" / "模块产销表"
    module_files = sorted(module_dir.glob("*.xlsx"))
    if len(module_files) != EXPECTED_MODULE_WORKBOOKS:
        raise RuntimeError(
            "经济模块产销表预期 %d 个，实际 %d 个"
            % (EXPECTED_MODULE_WORKBOOKS, len(module_files))
        )
    for source in module_files:
        result.append((source, Path("系统") / source.name))
    return result


def expected_relatives() -> list[str]:
    """返回正式 04 应包含的全部 xlsx 相对路径，供验算脚本复用。"""

    return sorted(relative.as_posix() for _source, relative in mappings())


def write_readme(output: Path) -> None:
    text = """# 04 数值文档（QA 交付版）

本目录由 `02_AI规则/stage_final_documents_from_md.py` 从本次生成物全量组装，
不是在旧数值文档上增量修改。读者是 QA：内容为可直接对照测试的数值表
（逐级曲线、逐层奖励、节点消耗、难度线、逐系统奖励/消耗清单）。

- 权威输入：`01_源表` 的全局事实，以及 `03_沉淀MD` 的系统参数与裁决。
- 正式位置：`numerical-workflows/v2/framework/04_数值文档/`，由 Git 随工作流管理。
- 工作区与验算报告：仅保存在 `output/numerical-production-workflow-v2/`。
- 参数 Excel：不再作为生成依赖。
- 发布数量：43 个 Excel（养成 10 + 大模块 1 + 战斗 2 + 经济 2 + 玩法 1 + 节奏 5
  + 系统 22）。
- `系统/` 22 份为每个玩法/外循环/养成系统的逐条件数值表：玩法类为「玩法产出」
  逐层/逐条件表，养成类为升级/升阶/升星/制作等阶段表，抽卡为卡池概率表；
  只有账本页的系统保留「奖励与消耗」逐条件清单页。
- 产销核算不发布：系统工作簿已剔除「阅读摘要/产销阅读/资源汇总/机器明细」
  核算页签；全资源产销总账与资源养成阅读版是内部经济分析产物，保留在
  工作流目录，不进入本交付。

每次升级必须先在新 staging 中生成并通过 `05_验算核算` 动态门禁，再整体替换
本目录。禁止从旧 04 回拷缺失文件。
"""
    (output / "README.md").write_text(text, encoding="utf-8")


def refresh_manifest(output: Path, pairs: list[tuple[Path, Path]]) -> None:
    """按 staging 当前字节刷新清单（用于 Excel 完整重算之后）。"""

    manifest_path = output / "generation-manifest.json"
    manifest = {
        "schema": 1,
        "built_at": datetime.now().astimezone().isoformat(timespec="seconds"),
        "architecture": "01_source + 03_md -> generated artifacts -> clean 04",
    }
    if manifest_path.exists():
        try:
            existing = json.loads(manifest_path.read_text(encoding="utf-8"))
            if isinstance(existing, dict):
                manifest.update(existing)
        except (json.JSONDecodeError, OSError):
            pass
    manifest["finalized_at"] = datetime.now().astimezone().isoformat(timespec="seconds")
    manifest["xlsx_count"] = len(pairs)
    manifest["artifacts"] = []
    for source, relative in pairs:
        destination = output / relative
        if not destination.exists():
            raise FileNotFoundError(destination)
        manifest["artifacts"].append(
            {
                "path": relative.as_posix(),
                "generator_output": source.relative_to(V2_DIR).as_posix(),
                "sha256": sha256(destination),
                "bytes": destination.stat().st_size,
            }
        )
    manifest_path.write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("output", help="新的 staging 目录；必须不存在")
    parser.add_argument(
        "--refresh-manifest",
        action="store_true",
        help="Excel 重算后只刷新现有 staging 的哈希清单，不复制文件",
    )
    args = parser.parse_args()
    output = Path(args.output).resolve()
    root = V2_DIR.resolve()
    if output == FINAL_DIR.resolve() or root not in output.parents:
        raise ValueError("staging 必须位于框架v2内且不能是正式04目录: %s" % output)
    if args.refresh_manifest:
        if not output.is_dir():
            raise FileNotFoundError("staging 不存在: %s" % output)
        pairs = mappings()
        actual = sorted(path.relative_to(output).as_posix() for path in output.rglob("*.xlsx"))
        if actual != expected_relatives():
            raise RuntimeError("staging 文件集合与期望清单不一致，拒绝刷新清单")
        refresh_manifest(output, pairs)
        print("[OK] manifest refreshed: %s" % (output / "generation-manifest.json"))
        return 0
    if output.exists():
        raise FileExistsError("staging 已存在，拒绝覆盖: %s" % output)

    pairs = mappings()
    if len(pairs) != EXPECTED_FINAL_WORKBOOKS:
        raise RuntimeError(
            "预期 %d 个最终映射，实际 %d 个" % (EXPECTED_FINAL_WORKBOOKS, len(pairs))
        )
    missing = [str(source) for source, _ in pairs if not source.exists()]
    if missing:
        raise FileNotFoundError("缺少本次生成物:\n" + "\n".join(missing))
    forbidden = [
        str(source)
        for source, _ in pairs
        if source.resolve() == FINAL_DIR.resolve() or FINAL_DIR.resolve() in source.resolve().parents
    ]
    if forbidden:
        raise RuntimeError("映射源不得来自旧 04:\n" + "\n".join(forbidden))

    output.mkdir(parents=True)
    manifest = {
        "schema": 1,
        "built_at": datetime.now().astimezone().isoformat(timespec="seconds"),
        "architecture": "01_source + 03_md -> generated artifacts -> clean 04",
        "xlsx_count": len(pairs),
        "artifacts": [],
    }
    for source, relative in pairs:
        destination = output / relative
        destination.parent.mkdir(parents=True, exist_ok=True)
        if relative.parts[0] == "系统":
            write_qa_module_copy(source, destination)
        else:
            shutil.copy2(source, destination)
        destination_hash = sha256(destination)
        manifest["artifacts"].append(
            {
                "path": relative.as_posix(),
                "generator_output": source.relative_to(V2_DIR).as_posix(),
                "sha256": destination_hash,
                "bytes": destination.stat().st_size,
            }
        )

    write_readme(output)
    (output / "generation-manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )

    actual = sorted(path.relative_to(output).as_posix() for path in output.rglob("*.xlsx"))
    expected = expected_relatives()
    if actual != expected:
        raise RuntimeError("staging 文件集合与期望清单不一致")

    print("[OK] clean staging=%s" % output)
    print("[OK] final_xlsx=%d reused_from_old_04=0" % len(actual))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
