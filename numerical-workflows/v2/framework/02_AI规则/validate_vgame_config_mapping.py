# -*- coding: utf-8 -*-
"""只读验证 ``Vgame配置映射总表.xlsx`` 及其生成报告。

本脚本不会保存工作簿，也不会修改 ``--config-root`` 下的任何文件。报告中的
计数和源快照是期望值；验证器不把 293/300 等项目状态硬编码为判定依据。
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
import zipfile
from pathlib import Path
from typing import Any, Iterable
from xml.etree import ElementTree

from openpyxl import load_workbook

from vgame_config_schema_extractor import extract_config_schema


REQUIRED_SHEETS = (
    "使用说明",
    "逻辑表清单",
    "表输入关系",
    "逻辑字段字典",
    "物理字段来源",
    "Bean与字段",
    "Enum与枚举项",
    "公式字段摘要",
    "未注册工作簿",
    "Schema异常与门禁",
    "模块映射",
    "配置草案",
    "审批记录",
    "落地回执",
    "工作流门禁",
)
META_SHEET = "__generation_meta__"
DEFAULT_MARKER = "vgame-config-mapping/v1"
SCHEMA_FILES = ("__tables__.xlsx", "__beans__.xlsx", "__enums__.xlsx")

ERROR_VALUES = {
    "#NULL!",
    "#DIV/0!",
    "#VALUE!",
    "#REF!",
    "#NAME?",
    "#NUM!",
    "#N/A",
    "#GETTING_DATA",
    "#SPILL!",
    "#CALC!",
    "#FIELD!",
    "#BLOCKED!",
    "#UNKNOWN!",
    "#CONNECT!",
    "#BUSY!",
}
EXTERNAL_FORMULA_RE = re.compile(
    r"\[(?:\d+|[^\[\]]+\.(?:xlsx?|xlsm|xlsb|xltx|xltm|csv|ods))\]"
    r"[^!\r\n]{0,128}!",
    re.IGNORECASE,
)

COUNT_ALIASES = {
    "logical_tables": (
        "logical_tables",
        "logical_table_count",
        "registered_logical_tables",
        "逻辑表数",
        "逻辑表数量",
    ),
    "registered_inputs": (
        "registered_inputs",
        "registered_input_count",
        "input_relation_count",
        "physical_inputs",
        "注册输入数",
        "输入关系数",
    ),
}

SHEET_HEADER_TOKENS = {
    "逻辑表清单": {"逻辑表", "逻辑表名", "full_name", "全名"},
    "表输入关系": {"逻辑表", "逻辑表名", "输入文件", "物理工作簿", "input"},
    "配置草案": {"批次ID", "草案ID", "batch_id", "draft_id"},
    "审批记录": {"审批ID", "批次ID", "approval_id", "batch_id"},
    "落地回执": {"回执ID", "批次ID", "receipt_id", "batch_id"},
}

REPORT_COLLECTION_SHEETS = {
    "logical_tables": "逻辑表清单",
    "table_inputs": "表输入关系",
    "logical_fields": "逻辑字段字典",
    "physical_fields": "物理字段来源",
    "bean_fields": "Bean与字段",
    "enum_items": "Enum与枚举项",
    "formula_fields": "公式字段摘要",
    "unregistered_workbooks": "未注册工作簿",
    "anomalies": "Schema异常与门禁",
}


def normalize_key(value: Any) -> str:
    return re.sub(r"[\s_\-./]+", "", str(value)).casefold()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def iter_dicts(value: Any) -> Iterable[dict]:
    if isinstance(value, dict):
        yield value
        for child in value.values():
            yield from iter_dicts(child)
    elif isinstance(value, list):
        for child in value:
            yield from iter_dicts(child)


def dict_value(mapping: dict, aliases: Iterable[str]) -> Any:
    normalized = {normalize_key(key): value for key, value in mapping.items()}
    for alias in aliases:
        key = normalize_key(alias)
        if key in normalized:
            return normalized[key]
    return None


def recursive_value(root: Any, aliases: Iterable[str]) -> Any:
    for mapping in iter_dicts(root):
        value = dict_value(mapping, aliases)
        if value is not None:
            return value
    return None


def as_count(value: Any) -> int | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if isinstance(value, (list, tuple, set, dict)):
        return len(value)
    try:
        text = str(value).strip()
        return int(text) if text else None
    except ValueError:
        return None


def expected_count(report: dict, name: str) -> int | None:
    counts = report.get("counts")
    if isinstance(counts, dict):
        value = dict_value(counts, COUNT_ALIASES[name])
        count = as_count(value)
        if count is not None:
            return count
    return as_count(recursive_value(report, COUNT_ALIASES[name]))


def cell_text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def find_header_row(worksheet, tokens: set[str]) -> int | None:
    normalized_tokens = {normalize_key(token) for token in tokens}
    best_row: int | None = None
    best_score = 0
    for row_number, row in enumerate(
        worksheet.iter_rows(min_row=1, max_row=min(worksheet.max_row, 30), values_only=True),
        start=1,
    ):
        labels = {normalize_key(value) for value in row if cell_text(value)}
        score = len(labels & normalized_tokens)
        if score > best_score:
            best_row = row_number
            best_score = score
    return best_row if best_score else None


def business_row_count(worksheet, header_tokens: set[str]) -> tuple[int | None, int]:
    header_row = find_header_row(worksheet, header_tokens)
    if header_row is None:
        return None, 0
    count = 0
    for row in worksheet.iter_rows(min_row=header_row + 1, values_only=True):
        if any(cell_text(value) for value in row):
            count += 1
    return header_row, count


def dense_business_row_count(worksheet) -> tuple[int | None, int]:
    """用前 30 行中最早的最密集行识别普通数据表头。"""

    best_row: int | None = None
    best_score = 0
    for row_number, row in enumerate(
        worksheet.iter_rows(min_row=1, max_row=min(worksheet.max_row, 30), values_only=True),
        start=1,
    ):
        score = sum(bool(cell_text(value)) for value in row)
        if score > best_score:
            best_row = row_number
            best_score = score
    if best_row is None or best_score < 2:
        return None, 0
    count = 0
    for row in worksheet.iter_rows(min_row=best_row + 1, values_only=True):
        if any(cell_text(value) for value in row):
            count += 1
    return best_row, count


def report_sheet_counts(report: dict) -> dict[str, int]:
    aliases = ("sheet_counts", "sheet_row_counts", "business_rows_by_sheet", "工作表计数")
    raw = recursive_value(report, aliases)
    if not isinstance(raw, dict):
        return {}
    result: dict[str, int] = {}
    for sheet_name, value in raw.items():
        count = as_count(value)
        if count is not None:
            result[str(sheet_name)] = count
    return result


def inspect_xlsx_package(path: Path, errors: list[str]) -> None:
    try:
        with zipfile.ZipFile(path) as archive:
            corrupt = archive.testzip()
            if corrupt:
                errors.append(f"工作簿 ZIP CRC 失败: {corrupt}")
            names = archive.namelist()
            external_parts = [name for name in names if name.startswith("xl/externalLinks/")]
            if external_parts:
                errors.append("工作簿包含 externalLinks 包: " + ", ".join(external_parts[:10]))
            for name in names:
                if not name.endswith(".rels"):
                    continue
                try:
                    root = ElementTree.fromstring(archive.read(name))
                except Exception as exc:
                    errors.append(f"关系 XML 无法解析 {name}: {type(exc).__name__}: {exc}")
                    continue
                for relationship in root:
                    if str(relationship.attrib.get("TargetMode", "")).casefold() == "external":
                        relation_type = str(relationship.attrib.get("Type", "")).rsplit("/", 1)[-1]
                        # 普通说明链接不是公式依赖；这里只阻断会参与工作簿计算或
                        # 数据装载的外部关系。
                        if relation_type.casefold() not in {
                            "externallink",
                            "externallinkpath",
                            "connection",
                            "oleobject",
                        }:
                            continue
                        target = relationship.attrib.get("Target", "")
                        errors.append(f"存在外部关系 {name}:{relation_type}: {target}")
    except (OSError, zipfile.BadZipFile) as exc:
        errors.append(f"工作簿不是有效 xlsx: {type(exc).__name__}: {exc}")


def defined_name_formulas(workbook) -> list[tuple[str, str]]:
    container = workbook.defined_names
    if callable(getattr(container, "values", None)):
        names = list(container.values())
    else:
        names = list(getattr(container, "definedName", []))
    result = []
    for item in names:
        formula = getattr(item, "attr_text", None)
        if isinstance(formula, str):
            result.append((str(getattr(item, "name", "")), formula))
    return result


def inspect_workbook(path: Path, report: dict, errors: list[str]) -> dict[str, int]:
    inspect_xlsx_package(path, errors)
    sheet_counts: dict[str, int] = {}
    workbook = None
    try:
        workbook = load_workbook(path, read_only=True, data_only=False, keep_links=True)
        missing = [name for name in REQUIRED_SHEETS if name not in workbook.sheetnames]
        if missing:
            errors.append("缺少必需 Sheet: " + "、".join(missing))

        expected_marker = recursive_value(
            report, ("generation_marker", "generator_marker", "生成标记")
        )
        expected_marker = cell_text(expected_marker) or DEFAULT_MARKER
        if META_SHEET not in workbook.sheetnames:
            errors.append(f"缺少隐藏生成标记 Sheet: {META_SHEET}")
        else:
            meta = workbook[META_SHEET]
            if meta.sheet_state not in {"hidden", "veryHidden"}:
                errors.append(f"生成标记 Sheet 未隐藏: {META_SHEET}")
            actual_key = cell_text(meta["A1"].value)
            actual_marker = cell_text(meta["B1"].value)
            if actual_key != "generation_marker" or actual_marker != expected_marker:
                errors.append(
                    "隐藏生成标记不匹配: "
                    f"A1={actual_key!r}, B1={actual_marker!r}, expected={expected_marker!r}"
                )

        links = list(getattr(workbook, "_external_links", []) or [])
        if links:
            errors.append(f"openpyxl 检测到 {len(links)} 个外链对象")

        for name, formula in defined_name_formulas(workbook):
            if EXTERNAL_FORMULA_RE.search(formula):
                errors.append(f"定义名称含外链公式: {name}={formula}")
            upper = formula.upper()
            if any(token in upper for token in ERROR_VALUES):
                errors.append(f"定义名称含显式错误: {name}={formula}")

        formula_issues = 0
        explicit_errors = 0
        for worksheet in workbook.worksheets:
            for row in worksheet.iter_rows():
                for cell in row:
                    value = cell.value
                    text = cell_text(value)
                    upper = text.upper()
                    if cell.data_type == "e" or upper in ERROR_VALUES:
                        explicit_errors += 1
                        if explicit_errors <= 20:
                            errors.append(
                                f"显式 Excel 错误: {worksheet.title}!{cell.coordinate}={text}"
                            )
                    if cell.data_type == "f" or text.startswith("="):
                        if EXTERNAL_FORMULA_RE.search(text):
                            formula_issues += 1
                            if formula_issues <= 20:
                                errors.append(
                                    f"外链公式: {worksheet.title}!{cell.coordinate}={text}"
                                )
                        if any(token in upper for token in ERROR_VALUES):
                            formula_issues += 1
                            if formula_issues <= 20:
                                errors.append(
                                    f"公式含显式错误: {worksheet.title}!{cell.coordinate}={text}"
                                )

        for sheet_name, tokens in SHEET_HEADER_TOKENS.items():
            if sheet_name not in workbook.sheetnames:
                continue
            header_row, count = business_row_count(workbook[sheet_name], tokens)
            if header_row is None:
                errors.append(f"{sheet_name} 未找到可识别表头")
            sheet_counts[sheet_name] = count

        for collection_name, sheet_name in REPORT_COLLECTION_SHEETS.items():
            if sheet_name not in workbook.sheetnames or sheet_name in sheet_counts:
                continue
            header_row, count = dense_business_row_count(workbook[sheet_name])
            if header_row is None:
                errors.append(f"{sheet_name} 未找到可识别表头")
            sheet_counts[sheet_name] = count

        for sheet_name in ("配置草案", "审批记录", "落地回执"):
            if sheet_name in sheet_counts and sheet_counts[sheet_name] != 0:
                errors.append(
                    f"{sheet_name} 初始应无业务行，实际 {sheet_counts[sheet_name]} 行"
                )

        expected_sheet_counts = report_sheet_counts(report)
        for sheet_name, expected in expected_sheet_counts.items():
            if sheet_name not in workbook.sheetnames:
                continue
            if sheet_name in sheet_counts:
                actual = sheet_counts[sheet_name]
            else:
                tokens = {cell_text(value) for value in next(
                    workbook[sheet_name].iter_rows(min_row=1, max_row=1, values_only=True),
                    (),
                ) if cell_text(value)}
                _header, actual = business_row_count(workbook[sheet_name], tokens)
            if actual != expected:
                errors.append(
                    f"Sheet 计数与报告不一致: {sheet_name} workbook={actual}, report={expected}"
                )
    except Exception as exc:
        errors.append(f"工作簿公式视图无法打开: {type(exc).__name__}: {exc}")
    finally:
        if workbook is not None:
            workbook.close()

    cached = None
    try:
        cached = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        cached_errors = 0
        for worksheet in cached.worksheets:
            for row in worksheet.iter_rows():
                for cell in row:
                    text = cell_text(cell.value)
                    if cell.data_type == "e" or text.upper() in ERROR_VALUES:
                        cached_errors += 1
                        if cached_errors <= 20:
                            errors.append(
                                f"缓存值含显式 Excel 错误: {worksheet.title}!{cell.coordinate}={text}"
                            )
    except Exception as exc:
        errors.append(f"工作簿缓存视图无法重开: {type(exc).__name__}: {exc}")
    finally:
        if cached is not None:
            cached.close()

    reopened = None
    try:
        reopened = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        _ = reopened.sheetnames
    except Exception as exc:
        errors.append(f"工作簿二次重开失败: {type(exc).__name__}: {exc}")
    finally:
        if reopened is not None:
            reopened.close()
    return sheet_counts


def validate_report_collection_counts(
    report: dict, sheet_counts: dict[str, int], errors: list[str]
) -> None:
    """报告公开的明细集合应与工作簿对应 Sheet 行数一致。"""

    counts = report.get("counts") if isinstance(report.get("counts"), dict) else {}
    for collection_name, sheet_name in REPORT_COLLECTION_SHEETS.items():
        collection = report.get(collection_name)
        if isinstance(collection, list):
            expected = len(collection)
        else:
            expected = as_count(counts.get(collection_name))
        if expected is None or sheet_name not in sheet_counts:
            continue
        actual = sheet_counts[sheet_name]
        if actual != expected:
            errors.append(
                f"{sheet_name} 行数与 JSON {collection_name} 不一致: "
                f"workbook={actual}, report={expected}"
            )


def registry_state(config_root: Path, errors: list[str]) -> dict:
    registry = config_root / "__tables__.xlsx"
    if not registry.is_file():
        errors.append(f"缺少注册表: {registry}")
        return {
            "logical_tables": 0,
            "registered_inputs": [],
            "empty_registered_inputs": [],
            "source_paths": [],
        }

    path_index: dict[str, list[Path]] = {}
    for path in config_root.rglob("*.xlsx"):
        if not path.name.startswith("~$"):
            path_index.setdefault(path.name.casefold(), []).append(path)

    workbook = None
    logical_tables = 0
    registered_inputs: list[str] = []
    empty_registered_inputs: list[str] = []
    try:
        workbook = load_workbook(registry, read_only=True, data_only=False, keep_links=False)
        worksheet = workbook.active
        headers = {
            cell_text(cell.value): index
            for index, cell in enumerate(worksheet[1])
            if cell_text(cell.value)
        }
        full_index = headers.get("full_name")
        input_index = headers.get("input")
        if full_index is None or input_index is None:
            errors.append("__tables__.xlsx 缺少 full_name/input 表头")
        else:
            for row in worksheet.iter_rows(min_row=4, values_only=True):
                full_name = cell_text(row[full_index]) if full_index < len(row) else ""
                input_value = cell_text(row[input_index]) if input_index < len(row) else ""
                if not full_name:
                    continue
                logical_tables += 1
                if not input_value:
                    empty_registered_inputs.append(full_name)
                    continue
                registered_inputs.extend(
                    part.strip()
                    for part in re.split(r"[,;]", input_value)
                    if part.strip()
                )
    except Exception as exc:
        errors.append(f"注册表无法读取: {type(exc).__name__}: {exc}")
    finally:
        if workbook is not None:
            workbook.close()

    resolved_sources: list[Path] = []
    for filename in registered_inputs:
        matches = path_index.get(Path(filename).name.casefold(), [])
        if not matches:
            errors.append(f"注册输入不存在: {filename}")
        elif len(matches) > 1:
            joined = "、".join(path.relative_to(config_root).as_posix() for path in matches)
            errors.append(f"注册输入存在同名歧义: {filename} => {joined}")
        else:
            resolved_sources.append(matches[0])

    for filename in SCHEMA_FILES:
        path = config_root / filename
        if not path.is_file():
            errors.append(f"缺少 Schema 文件: {path}")
        else:
            resolved_sources.append(path)

    unique_sources = sorted(
        set(path.resolve() for path in resolved_sources),
        key=lambda item: item.relative_to(config_root).as_posix(),
    )
    return {
        "logical_tables": logical_tables,
        "registered_inputs": registered_inputs,
        "empty_registered_inputs": sorted(empty_registered_inputs),
        "source_paths": unique_sources,
    }


def snapshot_records(config_root: Path, paths: Iterable[Path]) -> list[dict]:
    records = []
    for path in paths:
        records.append(
            {
                "path": path.relative_to(config_root).as_posix(),
                "bytes": path.stat().st_size,
                "sha256": sha256_file(path),
            }
        )
    return sorted(records, key=lambda item: item["path"])


def aggregate_snapshot(records: list[dict]) -> str:
    digest = hashlib.sha256()
    for item in sorted(records, key=lambda row: row["path"]):
        digest.update(
            (
                f"{item['path']}\0{item['bytes']}\0{item['sha256']}\n"
            ).encode("utf-8")
        )
    return digest.hexdigest()


def report_snapshot(report: dict) -> tuple[str, str, int | None, list[dict]]:
    node = report.get("source_snapshot")
    if not isinstance(node, dict):
        node = {}
    algorithm = cell_text(node.get("algorithm")) or "sha256-path-size-content-v1"
    expected_hash = cell_text(report.get("source_snapshot_hash"))
    if not expected_hash:
        expected_hash = cell_text(
            dict_value(node, ("sha256", "hash", "snapshot_sha256", "source_snapshot_hash"))
        )
    file_count = as_count(dict_value(node, ("file_count", "count", "source_file_count")))
    raw_files = dict_value(node, ("files", "entries", "sources", "items"))
    files = [item for item in raw_files if isinstance(item, dict)] if isinstance(raw_files, list) else []
    return algorithm, expected_hash, file_count, files


def validate_snapshot(
    config_root: Path, report: dict, state: dict, errors: list[str]
) -> tuple[int, str]:
    algorithm, expected_hash, expected_file_count, report_files = report_snapshot(report)
    snapshot_node = report.get("source_snapshot")
    nested_hash = (
        cell_text(dict_value(snapshot_node, ("sha256", "hash")))
        if isinstance(snapshot_node, dict)
        else ""
    )
    top_hash = cell_text(report.get("source_snapshot_hash"))
    if nested_hash and top_hash and nested_hash != top_hash:
        errors.append(
            f"JSON 内部源快照 hash 不一致: top={top_hash}, nested={nested_hash}"
        )
    if algorithm != "sha256-path-size-content-v1":
        errors.append(f"不支持的源快照算法: {algorithm!r}")
    if not expected_hash:
        errors.append("JSON 报告缺少 source_snapshot_hash")

    live_records = snapshot_records(config_root, state["source_paths"])
    live_hash = aggregate_snapshot(live_records)
    if expected_file_count is not None and len(live_records) != expected_file_count:
        errors.append(
            f"源快照文件数不一致: live={len(live_records)}, report={expected_file_count}"
        )
    if expected_hash and live_hash != expected_hash:
        errors.append(
            f"Datas 注册源快照已变化: live={live_hash}, report={expected_hash}"
        )

    if report_files:
        live_by_path = {item["path"]: item for item in live_records}
        report_by_path: dict[str, dict] = {}
        for item in report_files:
            path = cell_text(dict_value(item, ("path", "relative_path", "file"))).replace("\\", "/")
            if path:
                report_by_path[path] = item
        missing = sorted(set(report_by_path) - set(live_by_path))
        extra = sorted(set(live_by_path) - set(report_by_path))
        if missing or extra:
            errors.append(
                "源快照文件集合不一致: "
                f"missing_live={missing[:20]}, extra_live={extra[:20]}"
            )
        for path in sorted(set(report_by_path) & set(live_by_path)):
            expected_item_hash = cell_text(dict_value(report_by_path[path], ("sha256", "hash")))
            expected_size = as_count(dict_value(report_by_path[path], ("bytes", "size")))
            live_item = live_by_path[path]
            if expected_item_hash and live_item["sha256"] != expected_item_hash:
                errors.append(f"源文件 hash 变化: {path}")
            if expected_size is not None and live_item["bytes"] != expected_size:
                errors.append(
                    f"源文件大小变化: {path} live={live_item['bytes']} report={expected_size}"
                )
    return len(live_records), live_hash


def schema_row_has_business_data(row: tuple[Any, ...]) -> bool:
    if not any(cell_text(value) for value in row):
        return False
    return not row or cell_text(row[0]) != "##"


def validate_empty_registered_inputs(
    config_root: Path, report: dict, errors: list[str]
) -> tuple[int | None, list[str]]:
    """复查已注册输入中没有业务数据行的 Schema Sheet。"""

    counts = report.get("counts") if isinstance(report.get("counts"), dict) else {}
    expected_count = as_count(counts.get("empty_registered_inputs"))
    physical_inputs = report.get("physical_inputs")
    if expected_count is None:
        errors.append("JSON 报告缺少 counts.empty_registered_inputs")
        return None, []
    if not isinstance(physical_inputs, list):
        errors.append("JSON 报告缺少 physical_inputs 明细，无法复查空注册输入")
        return expected_count, []

    expected_empty: set[str] = set()
    live_empty: set[str] = set()
    for item in physical_inputs:
        if not isinstance(item, dict):
            continue
        source_path = cell_text(item.get("source_path")).replace("\\", "/")
        schema = item.get("schema") if isinstance(item.get("schema"), dict) else {}
        sheet_name = cell_text(schema.get("sheet_name"))
        data_start_row = as_count(schema.get("data_start_row")) or 5
        reported_rows = as_count(schema.get("data_row_count"))
        identity = f"{source_path} / {sheet_name}"
        if reported_rows == 0:
            expected_empty.add(identity)
        if not source_path or not sheet_name:
            errors.append(f"physical_inputs 缺少 source_path/sheet_name: {item!r}")

    try:
        live_schema = extract_config_schema(config_root)
        for item in live_schema.get("physical_inputs", []):
            schema = item.get("schema") if isinstance(item.get("schema"), dict) else {}
            if as_count(schema.get("data_row_count")) == 0:
                live_empty.add(
                    f"{cell_text(item.get('source_path')).replace(chr(92), '/')} / "
                    f"{cell_text(schema.get('sheet_name'))}"
                )
    except Exception as exc:
        errors.append(f"当前注册输入无法按真实XML Schema复查: {type(exc).__name__}: {exc}")

    if len(expected_empty) != expected_count:
        errors.append(
            "JSON 空注册输入明细与计数不一致: "
            f"physical_inputs={len(expected_empty)}, counts={expected_count}"
        )
    if len(live_empty) != expected_count or live_empty != expected_empty:
        errors.append(
            "当前空注册输入与报告不一致: "
            f"live={sorted(live_empty)}, report={sorted(expected_empty)}"
        )
    return expected_count, sorted(live_empty)


def validate(args) -> tuple[list[str], dict]:
    errors: list[str] = []
    workbook_path = Path(args.workbook).resolve()
    report_path = Path(args.report).resolve()
    config_root = Path(args.config_root).resolve()

    if not workbook_path.is_file():
        errors.append(f"工作簿不存在: {workbook_path}")
    if not report_path.is_file():
        errors.append(f"JSON 报告不存在: {report_path}")
    if not config_root.is_dir():
        errors.append(f"配置根目录不存在: {config_root}")
    if errors:
        return errors, {}

    try:
        report = json.loads(report_path.read_text(encoding="utf-8"))
        if not isinstance(report, dict):
            raise TypeError("报告根节点必须是 object")
    except Exception as exc:
        return [f"JSON 报告无法读取: {type(exc).__name__}: {exc}"], {}

    logical_expected = expected_count(report, "logical_tables")
    inputs_expected = expected_count(report, "registered_inputs")
    if logical_expected is None:
        errors.append("JSON 报告缺少逻辑表期望计数")
    if inputs_expected is None:
        errors.append("JSON 报告缺少注册输入期望计数")

    sheet_counts = inspect_workbook(workbook_path, report, errors)
    validate_report_collection_counts(report, sheet_counts, errors)
    state = registry_state(config_root, errors)
    live_logical = state["logical_tables"]
    live_inputs = len(state["registered_inputs"])
    live_unique_inputs = len({name.casefold() for name in state["registered_inputs"]})

    if logical_expected is not None:
        workbook_count = sheet_counts.get("逻辑表清单")
        if workbook_count != logical_expected:
            errors.append(
                f"逻辑表清单与报告不一致: workbook={workbook_count}, report={logical_expected}"
            )
        if live_logical != logical_expected:
            errors.append(
                f"当前 __tables__ 与报告逻辑表数不一致: live={live_logical}, report={logical_expected}"
            )
    if inputs_expected is not None:
        workbook_count = sheet_counts.get("表输入关系")
        if workbook_count != inputs_expected:
            errors.append(
                f"表输入关系与报告不一致: workbook={workbook_count}, report={inputs_expected}"
            )
        if live_inputs != inputs_expected:
            errors.append(
                f"当前 __tables__ 与报告注册输入数不一致: live={live_inputs}, report={inputs_expected}"
            )

    report_counts = report.get("counts") if isinstance(report.get("counts"), dict) else {}
    unique_inputs_expected = as_count(report_counts.get("registered_unique_inputs"))
    if unique_inputs_expected is not None and live_unique_inputs != unique_inputs_expected:
        errors.append(
            "当前 __tables__ 与报告唯一输入数不一致: "
            f"live={live_unique_inputs}, report={unique_inputs_expected}"
        )

    expected_empty = report.get("empty_registered_inputs")
    if isinstance(expected_empty, list):
        expected_empty_values = sorted(cell_text(item) for item in expected_empty if cell_text(item))
        if state["empty_registered_inputs"] != expected_empty_values:
            errors.append(
                "空 input 逻辑表与报告不一致: "
                f"live={state['empty_registered_inputs']}, report={expected_empty_values}"
            )

    source_count, source_hash = validate_snapshot(config_root, report, state, errors)
    empty_input_count, empty_inputs = validate_empty_registered_inputs(
        config_root, report, errors
    )

    workbook_hash_expected = cell_text(
        recursive_value(report, ("workbook_sha256", "mapping_workbook_sha256"))
    )
    if workbook_hash_expected:
        workbook_hash = sha256_file(workbook_path)
        if workbook_hash != workbook_hash_expected:
            errors.append(
                f"工作簿 hash 与报告不一致: live={workbook_hash}, report={workbook_hash_expected}"
            )

    summary = {
        "workbook": str(workbook_path),
        "report": str(report_path),
        "config_root": str(config_root),
        "logical_tables": live_logical,
        "registered_inputs": live_inputs,
        "registered_unique_inputs": live_unique_inputs,
        "registered_source_files": source_count,
        "source_snapshot_hash": source_hash,
        "empty_registered_input_count": empty_input_count,
        "empty_registered_inputs": empty_inputs,
        "empty_business_rows": {
            name: sheet_counts.get(name) for name in ("配置草案", "审批记录", "落地回执")
        },
    }
    return errors, summary


def main() -> int:
    parser = argparse.ArgumentParser(description="只读验证 Vgame 配置映射总表")
    parser.add_argument("--workbook", required=True, help="Vgame配置映射总表.xlsx")
    parser.add_argument("--report", required=True, help="生成器输出的 JSON 报告")
    parser.add_argument("--config-root", required=True, help="Vgame Config/GameConfig/Datas")
    args = parser.parse_args()

    errors, summary = validate(args)
    if errors:
        print("[FAIL] Vgame 配置映射总表验证失败", file=sys.stderr)
        for error in errors:
            print(f"- {error}", file=sys.stderr)
        if summary:
            print(json.dumps(summary, ensure_ascii=False, indent=2))
        return 1

    print("[OK] Vgame 配置映射总表验证通过")
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
