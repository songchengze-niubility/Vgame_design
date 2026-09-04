# -*- coding: utf-8 -*-
"""CHG-011：七日竞技场历史最高分任务描述由分数改为段位名称（对齐 ArenaRank.xlsx）。

只改 06_最终配置/任务 的 Task.xlsx（TaskName/TaskDesc）与 TaskTarget.xlsx（TargetDesc），
ParamList 分数参数保持段位最低分不变。备份、plan.json 与逐格校验落在
05_验算核算/七日竞技场段位描述CHG011_20260831/。
"""
import hashlib
import json
import shutil
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
RUN = ROOT / "05_验算核算" / "七日竞技场段位描述CHG011_20260831"

# task_id -> (段位描述, 段位最低分)
CHANGES = {
    5105: ("竞技场历史最高达到黄金Ⅱ", 1100),
    5205: ("竞技场历史最高达到黄金Ⅰ", 1200),
    5305: ("竞技场历史最高达到铂金Ⅲ", 1300),
    5405: ("竞技场历史最高达到铂金Ⅱ", 1400),
    5505: ("竞技场历史最高达到铂金Ⅰ", 1500),
    5605: ("竞技场历史最高达到钻石Ⅲ", 1600),
    5705: ("竞技场历史最高达到钻石Ⅱ", 1700),
}


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main():
    RUN.mkdir(parents=True, exist_ok=True)
    (RUN / "backup" / "任务").mkdir(parents=True, exist_ok=True)

    targets = {
        "任务/Task.xlsx": {
            "sheet": "Task",
            "id_col": 2,
            "edits": {tid: {3: name, 4: name} for tid, (name, _) in CHANGES.items()},
        },
        "任务/TaskTarget.xlsx": {
            "sheet": "TaskTarget",
            "id_col": 2,
            "edits": {tid: {3: name} for tid, (name, _) in CHANGES.items()},
        },
    }

    plan = {
        "batch": "ARENA-RANK-DESC-CHG011",
        "created": datetime.now().isoformat(),
        "root": str(ROOT),
        "run": str(RUN),
        "files": [],
    }

    for rel, spec in targets.items():
        src = ROOT / "06_最终配置" / rel
        backup = RUN / "backup" / rel
        shutil.copy2(src, backup)

        wb = load_workbook(src)
        ws = wb[spec["sheet"]]
        row_by_id = {}
        for row in range(5, ws.max_row + 1):
            value = ws.cell(row=row, column=spec["id_col"]).value
            if value in CHANGES:
                row_by_id[value] = row
        missing = sorted(set(spec["edits"]) - set(row_by_id))
        if missing:
            raise SystemExit(f"{rel}: 未找到任务行 {missing}")

        file_plan = {
            "file": rel,
            "source": str(src),
            "backup": str(backup),
            "sha256_before": sha256(src),
            "edits": [],
        }
        for tid in sorted(spec["edits"]):
            for col, new in spec["edits"][tid].items():
                cell = ws.cell(row=row_by_id[tid], column=col)
                old = cell.value
                if old == new:
                    continue
                file_plan["edits"].append(
                    {
                        "sheet": spec["sheet"],
                        "row": row_by_id[tid],
                        "col": col,
                        "task_id": tid,
                        "old": old,
                        "new": new,
                    }
                )
                cell.value = new

        plan["files"].append(file_plan)
        wb.save(src)
        file_plan["sha256_after"] = sha256(src)

    (RUN / "plan.json").write_text(
        json.dumps(plan, ensure_ascii=False, indent=2), encoding="utf-8"
    )

    # 逐格校验：备份 vs 现场，仅允许 plan 内单元格变化。
    expected = defaultdict(dict)
    for f in plan["files"]:
        for e in f["edits"]:
            expected[(f["file"], e["sheet"])][(e["row"], e["col"])] = e["new"]

    errors = []
    for rel, spec in targets.items():
        old_wb = load_workbook(RUN / "backup" / rel, read_only=True)
        new_wb = load_workbook(ROOT / "06_最终配置" / rel, read_only=True)
        old_rows = list(old_wb[spec["sheet"]].iter_rows(values_only=True))
        new_rows = list(new_wb[spec["sheet"]].iter_rows(values_only=True))
        if len(old_rows) != len(new_rows):
            errors.append(f"{rel}: 行数变化 {len(old_rows)}->{len(new_rows)}")
            continue
        allow = expected.get((rel, spec["sheet"]), {})
        for ri, (old_r, new_r) in enumerate(zip(old_rows, new_rows), 1):
            for ci, (old_v, new_v) in enumerate(zip(old_r, new_r), 1):
                if old_v != new_v and allow.get((ri, ci)) != new_v:
                    errors.append(
                        f"{rel}!{spec['sheet']}!{get_column_letter(ci)}{ri}: "
                        f"{old_v!r}->{new_v!r} 不在计划内"
                    )
        old_wb.close()
        new_wb.close()

    # 业务校验：描述含段位名、ParamList 仍为段位最低分、奖励与解锁不变。
    check = {}
    wb_task = load_workbook(ROOT / "06_最终配置" / "任务" / "Task.xlsx", read_only=True)
    wb_tt = load_workbook(
        ROOT / "06_最终配置" / "任务" / "TaskTarget.xlsx", read_only=True
    )
    param = {}
    for row in wb_tt["TaskTarget"].iter_rows(min_row=5, values_only=True):
        if row[1] is not None:
            param[row[1]] = (row[2], row[4])
    days, expected_drop = {}, {1: 91000005, 2: 91000015, 3: 91000025,
                               4: 91000035, 5: 91000045, 6: 91000055, 7: 91000065}
    for row in wb_task["Task"].iter_rows(min_row=5, values_only=True):
        tid = row[1]
        if isinstance(tid, int) and 5105 <= tid <= 5705 and tid % 100 == 5:
            days[tid // 100 - 50] = {
                "name": row[2],
                "target_desc": param.get(tid, (None, None))[0],
                "param": param.get(tid, (None, None))[1],
                "drop": row[14],
                "unlock": row[10],
            }
    check["七日段位描述"] = {
        f"第{d}天": [days[d]["name"], days[d]["param"]] for d in sorted(days)
    }
    desc_ok = all(
        days[d]["name"] == CHANGES[tid][0]
        and days[d]["target_desc"] == CHANGES[tid][0]
        for d, tid in zip(sorted(days), sorted(CHANGES))
    )
    param_ok = all(
        days[d]["param"] == f"0,{CHANGES[tid][1]}"
        for d, tid in zip(sorted(days), sorted(CHANGES))
    )
    meta_ok = all(
        days[d]["unlock"] == 4112 and days[d]["drop"] == expected_drop[d]
        for d in days
    )
    wb_task.close()
    wb_tt.close()

    result = {
        "batch": "ARENA-RANK-DESC-CHG011",
        "generated_at": datetime.now().isoformat(),
        "cell_diff_outside_plan": errors,
        "business_checks": {
            "描述含段位名": desc_ok,
            "ParamList为段位最低分": param_ok,
            "奖励与解锁不变": meta_ok,
        },
        "status": "PASS" if not errors and desc_ok and param_ok and meta_ok else "FAIL",
    }
    (RUN / "validation.json").write_text(
        json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(json.dumps(result, ensure_ascii=False, indent=2))
    if result["status"] != "PASS":
        raise SystemExit(1)


if __name__ == "__main__":
    main()
