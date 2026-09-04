# -*- coding: utf-8 -*-
"""CHG-011 SVN 落表：把 06 已验证的段位描述改动同步到 ${VGAME_CONFIG_DATAS}/task。

只改 Task.xlsx（TaskName/TaskDesc）与 TaskTarget.xlsx（TargetDesc）共 7 条 21 格，
ParamList 分数参数已在 CHG-010 落表，不动。列位置按表头名定位。
执行前请确认用户已授权；只改工作副本，不执行 svn commit。
"""
import hashlib
import json
import os
import shutil
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
RUN = ROOT / "05_验算核算" / "七日竞技场段位描述CHG011_20260831"
SVN_TASK_DIR = Path(os.environ.get("VGAME_CONFIG_DATAS") or str(
        Path(__file__).resolve().parents[5].parent / "Vgame" / "Config" / "GameConfig" / "Datas"
    )) / "task"

CHANGES = {
    5105: ("竞技场历史最高达到黄金Ⅱ", 1100),
    5205: ("竞技场历史最高达到黄金Ⅰ", 1200),
    5305: ("竞技场历史最高达到铂金Ⅲ", 1300),
    5405: ("竞技场历史最高达到铂金Ⅱ", 1400),
    5505: ("竞技场历史最高达到铂金Ⅰ", 1500),
    5605: ("竞技场历史最高达到钻石Ⅲ", 1600),
    5705: ("竞技场历史最高达到钻石Ⅱ", 1700),
}


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def col_of(header, name):
    for idx, h in enumerate(header, 1):
        if h == name:
            return idx
    raise SystemExit(f"表头缺少列 {name}")


def main():
    (RUN / "backup-svn" / "task").mkdir(parents=True, exist_ok=True)

    specs = [
        ("Task.xlsx", "Task", ("TaskName", "TaskDesc")),
        ("TaskTarget.xlsx", "TaskTarget", ("TargetDesc",)),
    ]

    plan = {
        "batch": "ARENA-RANK-DESC-CHG011-SVN",
        "created": datetime.now().isoformat(),
        "svn_dir": str(SVN_TASK_DIR),
        "files": [],
    }

    for filename, sheet, fields in specs:
        src = SVN_TASK_DIR / filename
        backup = RUN / "backup-svn" / "task" / filename
        shutil.copy2(src, backup)

        wb = load_workbook(src)
        ws = wb[sheet]
        rows = list(ws.iter_rows(values_only=True))
        header = rows[0]
        id_col = col_of(header, "Id")
        row_by_id = {}
        for ri, r in enumerate(rows[4:], 5):
            if r and r[id_col - 1] in CHANGES:
                row_by_id[r[id_col - 1]] = ri
        missing = sorted(set(CHANGES) - set(row_by_id))
        if missing:
            raise SystemExit(f"{filename}: 未找到目标行 {missing}")

        file_plan = {
            "file": f"task/{filename}",
            "source": str(src),
            "backup": str(backup),
            "sha256_before": sha256(src),
            "edits": [],
        }
        for tid in sorted(CHANGES):
            name, _ = CHANGES[tid]
            for field in fields:
                ci = col_of(header, field)
                cell = ws.cell(row=row_by_id[tid], column=ci)
                old = cell.value
                new = name
                if old == new:
                    continue
                file_plan["edits"].append(
                    {
                        "sheet": sheet,
                        "row": row_by_id[tid],
                        "col": ci,
                        "field": field,
                        "task_id": tid,
                        "old": old,
                        "new": new,
                    }
                )
                cell.value = new

        plan["files"].append(file_plan)
        wb.save(src)
        file_plan["sha256_after"] = sha256(src)

    (RUN / "plan-svn.json").write_text(
        json.dumps(plan, ensure_ascii=False, indent=2), encoding="utf-8"
    )

    # 逐格校验：备份 vs 现场，仅允许 plan 内单元格变化。
    expected = defaultdict(dict)
    for f in plan["files"]:
        for e in f["edits"]:
            expected[(f["file"], e["sheet"])][(e["row"], e["col"])] = e["new"]

    errors = []
    for filename, sheet, _ in specs:
        rel = f"task/{filename}"
        old_wb = load_workbook(RUN / "backup-svn" / "task" / filename, read_only=True)
        new_wb = load_workbook(SVN_TASK_DIR / filename, read_only=True)
        old_rows = list(old_wb[sheet].iter_rows(values_only=True))
        new_rows = list(new_wb[sheet].iter_rows(values_only=True))
        if len(old_rows) != len(new_rows):
            errors.append(f"{rel}: 行数变化 {len(old_rows)}->{len(new_rows)}")
            continue
        allow = expected.get((rel, sheet), {})
        for ri, (old_r, new_r) in enumerate(zip(old_rows, new_rows), 1):
            for ci, (old_v, new_v) in enumerate(zip(old_r, new_r), 1):
                if old_v != new_v and allow.get((ri, ci)) != new_v:
                    errors.append(
                        f"{rel}!{sheet}!{get_column_letter(ci)}{ri}: "
                        f"{old_v!r}->{new_v!r} 不在计划内"
                    )
        old_wb.close()
        new_wb.close()

    # 业务校验：段位描述 + ParamList 仍为段位最低分 + 解锁与奖励不变。
    wb_task = load_workbook(SVN_TASK_DIR / "Task.xlsx", read_only=True, data_only=True)
    wb_tt = load_workbook(SVN_TASK_DIR / "TaskTarget.xlsx", read_only=True, data_only=True)
    t_rows = list(wb_task["Task"].iter_rows(values_only=True))
    tt_rows = list(wb_tt["TaskTarget"].iter_rows(values_only=True))
    th = t_rows[0]
    unlock_c, drop_c = col_of(th, "UnlockId"), col_of(th, "DropId")
    param = {r[1]: (r[2], r[4]) for r in tt_rows[4:] if r and r[1] is not None}
    days, expected_drop = {}, {1: "91000005", 2: "91000015", 3: "91000025",
                               4: "91000035", 5: "91000045", 6: "91000055", 7: "91000065"}
    for r in t_rows[4:]:
        tid = r[1]
        if isinstance(tid, int) and 5105 <= tid <= 5705 and tid % 100 == 5:
            days[tid // 100 - 50] = {
                "name": r[2],
                "desc": param.get(tid, (None, None))[0],
                "param": param.get(tid, (None, None))[1],
                "unlock": r[unlock_c - 1],
                "drop": str(r[drop_c - 1]),
            }
    desc_ok = all(
        days[d]["name"] == CHANGES[tid][0] and days[d]["desc"] == CHANGES[tid][0]
        for d, tid in zip(sorted(days), sorted(CHANGES))
    )
    param_ok = all(
        days[d]["param"] == f"0,{CHANGES[tid][1]}"
        for d, tid in zip(sorted(days), sorted(CHANGES))
    )
    meta_ok = all(
        days[d]["unlock"] == 4112 and days[d]["drop"] == expected_drop[d] for d in days
    )
    wb_task.close()
    wb_tt.close()

    result = {
        "batch": "ARENA-RANK-DESC-CHG011-SVN",
        "generated_at": datetime.now().isoformat(),
        "cell_diff_outside_plan": errors,
        "business_checks": {
            "描述含段位名": desc_ok,
            "ParamList为段位最低分": param_ok,
            "解锁与奖励不变": meta_ok,
        },
        "status": "PASS" if not errors and desc_ok and param_ok and meta_ok else "FAIL",
    }
    (RUN / "validation-svn.json").write_text(
        json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(json.dumps(result, ensure_ascii=False, indent=2))
    if result["status"] != "PASS":
        raise SystemExit(1)


if __name__ == "__main__":
    main()
