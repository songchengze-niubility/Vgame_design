# -*- coding: utf-8 -*-
"""生成 Vgame 全量配置 Schema 映射与配置草案工作簿。"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import shutil
import sys
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

from vgame_config_schema_extractor import extract_config_schema


V2_DIR = Path(__file__).resolve().parents[1]
DEFAULT_CONFIG_ROOT = Path(
    os.environ.get("VGAME_CONFIG_DATAS", "")
)
DEFAULT_WORKBOOK = V2_DIR / "00_总览" / "Vgame配置映射总表.xlsx"
DEFAULT_REPORT = V2_DIR / "05_验算核算" / "vgame-config-mapping-report.json"
DEFAULT_MD_REPORT = V2_DIR / "_reports" / "P6_Vgame配置映射总表_生成与审计.md"

MARKER = "vgame-config-mapping/v1"
META_SHEET = "__generation_meta__"
FONT_NAME = "微软雅黑"
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
SUBHEADER_FILL = PatternFill("solid", fgColor="D9EAF7")
TITLE_FILL = PatternFill("solid", fgColor="17365D")
PASS_FILL = PatternFill("solid", fgColor="E2F0D9")
WARN_FILL = PatternFill("solid", fgColor="FFF2CC")
FAIL_FILL = PatternFill("solid", fgColor="FCE4D6")
INFO_FILL = PatternFill("solid", fgColor="DDEBF7")
GRAY_FILL = PatternFill("solid", fgColor="F2F2F2")
WHITE_BOLD = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
BODY_FONT = Font(name=FONT_NAME, size=10, color="000000")
INPUT_FONT = Font(name=FONT_NAME, size=10, color="0000FF")
THIN_GRAY = Side(style="thin", color="D9E1F2")
BODY_BORDER = Border(bottom=THIN_GRAY)

STAGING_SHEETS = ("配置草案", "审批记录", "落地回执")
MODULE_HEADERS = [
    "逻辑表",
    "注册表行",
    "输入序号",
    "输入文件",
    "相对路径",
    "Schema Sheet",
    "Vgame目录模块",
    "建议数值域",
    "建议置信度",
    "业务模块",
    "数值子模块",
    "来源MD",
    "关系类型",
    "映射状态",
    "负责人",
    "备注",
]


class Formula(str):
    pass


def now_text() -> str:
    return datetime.now().astimezone().isoformat(timespec="seconds")


def json_text(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, separators=(",", ":"))


def display(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, Formula):
        return str(value)
    if isinstance(value, bool):
        return "是" if value else "否"
    if isinstance(value, (list, dict, tuple)):
        return json_text(value)
    if isinstance(value, str) and value.startswith("="):
        return "'" + value
    return value


def compact(value: Any, limit: int = 2000) -> str:
    text = "" if value is None else str(value)
    text = " ".join(text.split())
    return text if len(text) <= limit else text[: limit - 1] + "…"


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def parse_index(raw: Any) -> list[str]:
    if raw is None:
        return []
    return [part.strip() for part in re.split(r"[,+]", str(raw)) if part.strip()]


def example_cells(value: Any) -> str:
    if not isinstance(value, list):
        return ""
    cells = []
    for item in value:
        if isinstance(item, dict) and item.get("cell"):
            cells.append(str(item["cell"]))
    return ";".join(cells[:5])


def classify_domain(source_path: str) -> tuple[str, str]:
    top = source_path.replace("\\", "/").split("/", 1)[0].casefold()
    combat = {
        "monster", "skill", "buff", "bullet", "fightitem", "camera", "battle",
        "spskill", "ai", "behaviortree", "summon",
    }
    growth = {
        "hero", "exclusiveweapon", "strategy", "specialequipment", "item_craft",
        "portableequipment", "player", "weapon",
    }
    economy = {
        "drop", "shop", "activity", "task", "battlepass", "signinreward", "monthcard",
        "mail", "items", "guild", "drawcard",
    }
    pacing = {"level", "unlock", "system_open", "guide", "mainlevel", "chapter"}
    if top in combat:
        return "战斗", "AUTO_MEDIUM"
    if top in growth:
        return "养成", "AUTO_MEDIUM"
    if top in economy:
        return "经济", "AUTO_MEDIUM"
    if top in pacing:
        return "节奏", "AUTO_MEDIUM"
    return "综合/待校准", "AUTO_LOW"


def read_preserved(path: Path, force: bool) -> dict[str, Any]:
    if not path.exists():
        return {"staging": {}, "module_manual": {}}
    workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        marked = (
            META_SHEET in workbook.sheetnames
            and str(workbook[META_SHEET]["A1"].value or "") == "generation_marker"
            and str(workbook[META_SHEET]["B1"].value or "") == MARKER
        )
        if not marked and not force:
            raise RuntimeError(
                f"目标文件已存在但没有生成标记，拒绝覆盖: {path}；如确认覆盖请使用 --force"
            )
        staging: dict[str, list[list[Any]]] = {}
        for name in STAGING_SHEETS:
            rows: list[list[Any]] = []
            if name in workbook.sheetnames:
                worksheet = workbook[name]
                for row in worksheet.iter_rows(min_row=2, values_only=True):
                    values = list(row)
                    if any(value not in (None, "") for value in values):
                        rows.append(values)
            staging[name] = rows

        module_manual: dict[tuple[str, str, str], dict[str, Any]] = {}
        if "模块映射" in workbook.sheetnames:
            worksheet = workbook["模块映射"]
            headers = [str(cell.value or "") for cell in worksheet[1]]
            positions = {name: index for index, name in enumerate(headers)}
            manual_headers = MODULE_HEADERS[9:]
            for row in worksheet.iter_rows(min_row=2, values_only=True):
                values = list(row)
                logical = values[positions["逻辑表"]] if "逻辑表" in positions else None
                source = values[positions["相对路径"]] if "相对路径" in positions else None
                sheet = values[positions["Schema Sheet"]] if "Schema Sheet" in positions else None
                if not logical or not source:
                    continue
                module_manual[(str(logical), str(source), str(sheet or ""))] = {
                    header: values[positions[header]]
                    for header in manual_headers
                    if header in positions and positions[header] < len(values)
                }
        return {"staging": staging, "module_manual": module_manual}
    finally:
        workbook.close()


def add_table(worksheet, name: str, row_count: int, column_count: int) -> None:
    if row_count <= 0 or column_count <= 0:
        worksheet.auto_filter.ref = f"A1:{get_column_letter(column_count)}1"
        return
    table = Table(
        displayName=name,
        ref=f"A1:{get_column_letter(column_count)}{row_count + 1}",
    )
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    worksheet.add_table(table)


def style_data_sheet(worksheet, headers: list[str], rows: list[list[Any]], table_name: str) -> None:
    worksheet.sheet_view.showGridLines = False
    worksheet.freeze_panes = "A2"
    worksheet.row_dimensions[1].height = 30
    for cell in worksheet[1]:
        cell.fill = HEADER_FILL
        cell.font = WHITE_BOLD
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = BODY_FONT
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = BODY_BORDER
            token = str(cell.value or "").upper()
            if token in {"PASS", "APPROVED", "VERIFIED", "APPLIED", "EXPORTED"}:
                cell.fill = PASS_FILL
            elif token in {"WARN", "REVIEW", "READY_FOR_REVIEW", "NEEDS_REVISION", "MANUAL_REQUIRED"}:
                cell.fill = WARN_FILL
            elif token in {"FAIL", "BLOCK", "REJECTED"}:
                cell.fill = FAIL_FILL
            elif token in {"INFO", "NOT_STARTED", "DRAFT", "UNCLASSIFIED"}:
                cell.fill = INFO_FILL
    add_table(worksheet, table_name, len(rows), len(headers))
    for index, header in enumerate(headers, start=1):
        values = [header]
        for row in rows[:1000]:
            if index - 1 < len(row):
                values.append(compact(row[index - 1], 120))
        width = max((len(str(value)) for value in values), default=10) + 2
        worksheet.column_dimensions[get_column_letter(index)].width = min(max(width, 10), 48)


def add_data_sheet(
    workbook: Workbook,
    name: str,
    headers: list[str],
    rows: list[list[Any]],
    table_name: str,
):
    worksheet = workbook.create_sheet(name)
    worksheet.append(headers)
    for row in rows:
        worksheet.append([display(value) for value in row])
    style_data_sheet(worksheet, headers, rows, table_name)
    return worksheet


def add_list_validation(worksheet, column: int, values: list[str], end_row: int = 5000) -> None:
    validation = DataValidation(
        type="list",
        formula1='"' + ",".join(values) + '"',
        allow_blank=True,
    )
    validation.error = "请选择预设状态，避免产生不可识别的批次状态。"
    validation.errorTitle = "无效状态"
    validation.prompt = "请选择列表中的值。"
    validation.promptTitle = "配置批次门禁"
    worksheet.add_data_validation(validation)
    validation.add(f"{get_column_letter(column)}2:{get_column_letter(column)}{end_row}")


def create_staging_sheet(
    workbook: Workbook,
    name: str,
    headers: list[str],
    preserved_rows: list[list[Any]],
    table_name: str,
):
    worksheet = add_data_sheet(workbook, name, headers, preserved_rows, table_name)
    # 空白模板没有 Table，使用工作表级筛选；有保留业务行时由 Table 自带筛选。
    # 两者重叠会被旧版 Excel（如本机 Excel 12）判为冲突并移除 Table。
    if preserved_rows:
        worksheet.auto_filter.ref = None
    else:
        worksheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}5000"
    for index, header in enumerate(headers, start=1):
        worksheet.cell(1, index).comment = Comment(
            "人工维护区。Schema 刷新会保留已有业务行；任何正式写入仍需当前批次人工批准。",
            "Vgame数值工作流",
        )
        worksheet.column_dimensions[get_column_letter(index)].width = min(
            max(12, len(header) * 2 + 2), 34
        )
    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = INPUT_FONT
    return worksheet


def build_workbook(data: dict[str, Any], preserved: dict[str, Any], report_path: Path) -> tuple[Workbook, dict[str, int]]:
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.creator = "Vgame 数值框架 v2"
    workbook.properties.title = "Vgame配置映射总表"
    workbook.properties.subject = "Schema 镜像、配置草案、人工审批与正式落表门禁"
    workbook.properties.description = MARKER
    try:
        workbook.calculation.calcMode = "auto"
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
    except Exception:
        pass

    logical_conflicts: dict[str, int] = {}
    for field in data["logical_fields"]:
        if field.get("schema_conflict"):
            logical_conflicts[field["logical_table"]] = logical_conflicts.get(field["logical_table"], 0) + 1

    logical_headers = [
        "逻辑表", "注册表行", "ValueType", "从Excel读SchemaRaw", "从Excel读Schema",
        "注册输入数", "注册输入列表", "生成C#Raw", "生成C#", "显式主键",
        "模式", "分组", "输出", "注释", "Tags", "Schema冲突字段数", "状态",
    ]
    logical_rows = []
    for item in data["logical_tables"]:
        logical_rows.append([
            item["logical_table"], item["registry_row"], item.get("value_type"),
            item.get("read_schema_from_file_raw"), item.get("read_schema_from_file"),
            len(item.get("inputs") or []), ";".join(item.get("inputs") or []),
            item.get("gen_csharp_class_raw"), item.get("gen_csharp_class"), item.get("index_raw"),
            item.get("mode"), item.get("group"), item.get("output"), item.get("comment"),
            item.get("tags"), logical_conflicts.get(item["logical_table"], 0), "PASS",
        ])
    add_data_sheet(workbook, "逻辑表清单", logical_headers, logical_rows, "TblLogicalTables")

    physical_input_map = {item["source_path"]: item for item in data["physical_inputs"]}
    snapshot_map = {item["path"]: item for item in data["source_snapshot"]["files"]}
    input_headers = [
        "逻辑表", "注册表行", "输入序号", "输入文件", "相对路径", "解析状态",
        "Schema Sheet", "Schema方向", "表头布局", "数据起始行", "导出字段数",
        "注释字段数", "工作表数", "辅助Sheet", "公式总数", "Schema公式数",
        "辅助页公式数", "外链公式数", "文件字节", "SHA256", "状态",
    ]
    input_rows = []
    for item in data["table_inputs"]:
        physical = physical_input_map.get(item.get("source_path"), {})
        schema = physical.get("schema") or {}
        snapshot = snapshot_map.get(item.get("source_path"), {})
        schema_names = set(physical.get("schema_sheet_names") or [])
        auxiliary = [name for name in physical.get("sheet_names") or [] if name not in schema_names]
        status = "PASS" if item.get("resolution_status") == "resolved" and physical.get("schema_sheet_count") == 1 else "BLOCK"
        input_rows.append([
            item["logical_table"], item["registry_row"], item["input_index"], item["input_name"],
            item.get("source_path"), item.get("resolution_status"), item.get("schema_sheet"),
            schema.get("orientation"), schema.get("header_layout"), schema.get("data_start_row"),
            item.get("exported_field_count"), item.get("comment_field_count"), physical.get("sheet_count"),
            ";".join(auxiliary), physical.get("formula_count"), physical.get("schema_formula_count"),
            physical.get("auxiliary_formula_count"), physical.get("external_formula_count"),
            snapshot.get("bytes"), snapshot.get("sha256"), status,
        ])
    add_data_sheet(workbook, "表输入关系", input_headers, input_rows, "TblTableInputs")

    table_map = {item["logical_table"]: item for item in data["logical_tables"]}
    key_fields: dict[str, set[str]] = {}
    first_fields: dict[str, str] = {}
    for field in data["logical_fields"]:
        first_fields.setdefault(field["logical_table"], field["field_name"])
    for logical, table in table_map.items():
        explicit = parse_index(table.get("index_raw"))
        key_fields[logical] = set(explicit or ([first_fields[logical]] if logical in first_fields else []))

    logical_field_headers = [
        "逻辑表", "字段序号", "字段名", "规范RawType", "BaseType", "类型形态",
        "可空", "集合类型", "分隔符", "默认值Raw", "字段Group", "是否主键",
        "来源文件数", "缺失来源", "来源变体", "Schema冲突", "状态",
    ]
    logical_field_rows = []
    for item in data["logical_fields"]:
        variants = item.get("source_variants") or []
        logical_field_rows.append([
            item["logical_table"], item["field_ordinal"], item["field_name"],
            item.get("canonical_raw_type"), item.get("base_type"), item.get("collection_kind") or "scalar",
            item.get("is_optional"), item.get("collection_kind"), item.get("separator"), item.get("default_raw"),
            item.get("canonical_group"), item["field_name"] in key_fields.get(item["logical_table"], set()),
            len(variants), ";".join(item.get("missing_sources") or []), variants,
            item.get("schema_conflict"), "REVIEW" if item.get("schema_conflict") else "PASS",
        ])
    add_data_sheet(workbook, "逻辑字段字典", logical_field_headers, logical_field_rows, "TblLogicalFields")

    bean_names = {item["full_name"] for item in data["beans"]}
    enum_names = {item["full_name"] for item in data["enums"]}
    conflict_keys = {
        (item["logical_table"], item["field_name"])
        for item in data["logical_fields"] if item.get("schema_conflict")
    }
    dimension_keys = {
        (item.get("source_path"), item.get("sheet_name"))
        for item in data["anomalies"] if str(item.get("code", "")).startswith("worksheet_dimension")
    }
    physical_headers = [
        "逻辑表", "相对路径", "Schema Sheet", "Schema方向", "字段序号", "字段轴",
        "轴序号", "Excel轴", "表头单元格", "字段名", "RawType", "BaseType",
        "类型形态", "可空", "集合类型", "分隔符", "默认值Raw", "字段Group",
        "字段注释", "注释字段", "是否导出", "是否主键", "Bean或Enum来源",
        "数据起始行", "当前公式单元格数", "当前常量单元格数", "公式占比",
        "外链公式数", "公式常量混合", "公式示例地址", "自动写入策略",
        "Schema冲突", "维度异常", "备注",
    ]
    physical_rows = []
    for row_number, item in enumerate(data["physical_fields"], start=2):
        logicals = item.get("logical_tables") or []
        logical_text = ";".join(logicals)
        base_type = item.get("base_type")
        type_ref = base_type if base_type in bean_names or base_type in enum_names else ""
        formula_count = int(item.get("formula_cell_count") or 0)
        comment_field = bool(item.get("is_comment_field"))
        if comment_field:
            policy = "禁止自动写入（非导出注释字段）"
        elif formula_count:
            policy = "MANUAL_REQUIRED"
        else:
            policy = "CANDIDATE_SAFE_RECHECK"
        is_key = any(item["field_name"] in key_fields.get(logical, set()) for logical in logicals)
        conflict = any((logical, item["field_name"]) in conflict_keys for logical in logicals)
        physical_input = physical_input_map.get(item.get("source_path"), {})
        data_start = (physical_input.get("schema") or {}).get("data_start_row")
        dimension_anomaly = (item.get("source_path"), item.get("schema_sheet")) in dimension_keys
        physical_rows.append([
            logical_text, item.get("source_path"), item.get("schema_sheet"), item.get("orientation"),
            item.get("source_ordinal"), item.get("axis"), item.get("axis_index"), item.get("excel_axis"),
            item.get("header_cell"), item.get("field_name"), item.get("raw_type"), base_type,
            item.get("collection_kind") or "scalar", item.get("is_optional"), item.get("collection_kind"),
            item.get("separator"), item.get("default_raw"), item.get("group"), item.get("comment"),
            comment_field, item.get("is_exported"), is_key, type_ref, data_start,
            formula_count, item.get("constant_cell_count"), Formula(f"=IFERROR(Y{row_number}/(Y{row_number}+Z{row_number}),0)"),
            item.get("external_formula_cell_count"), item.get("formula_constant_mixed"),
            example_cells(item.get("formula_examples")), policy, conflict, dimension_anomaly, "",
        ])
    physical_sheet = add_data_sheet(
        workbook, "物理字段来源", physical_headers, physical_rows, "TblPhysicalFields"
    )
    physical_sheet.column_dimensions["AA"].width = 12
    for cell in physical_sheet["AA"][1:]:
        cell.number_format = "0.0%"

    bean_map = {item["full_name"]: item for item in data["beans"]}
    bean_headers = [
        "Bean", "Bean定义行", "Parent", "ValueType", "Bean分隔符", "Alias", "Bean Group",
        "字段定义行", "字段名", "RawType", "字段Group", "字段注释", "Variants", "Tags",
    ]
    bean_rows = []
    for item in data["bean_fields"]:
        bean = bean_map.get(item["bean"], {})
        bean_rows.append([
            item["bean"], bean.get("definition_row"), bean.get("parent"), bean.get("value_type"),
            bean.get("separator"), bean.get("alias"), bean.get("group"), item.get("definition_row"),
            item.get("field_name"), item.get("raw_type"), item.get("group"), item.get("comment"),
            item.get("variants"), item.get("tags"),
        ])
    add_data_sheet(workbook, "Bean与字段", bean_headers, bean_rows, "TblBeanFields")

    enum_map = {item["full_name"]: item for item in data["enums"]}
    enum_headers = [
        "Enum", "Enum定义行", "FlagsRaw", "Flags", "UniqueRaw", "Unique", "Enum Group",
        "枚举项定义行", "枚举项", "Alias", "Value", "隐式值", "注释", "Tags",
    ]
    enum_rows = []
    for item in data["enum_items"]:
        enum = enum_map.get(item["enum"], {})
        enum_rows.append([
            item["enum"], enum.get("definition_row"), enum.get("flags_raw"), enum.get("flags"),
            enum.get("unique_raw"), enum.get("unique"), enum.get("group"), item.get("definition_row"),
            item.get("item_name"), item.get("alias"), item.get("value"), item.get("value_is_implicit"),
            item.get("comment"), item.get("tags"),
        ])
    add_data_sheet(workbook, "Enum与枚举项", enum_headers, enum_rows, "TblEnumItems")

    formula_headers = [
        "逻辑表", "相对路径", "Schema Sheet", "字段名", "RawType", "是否导出", "注释字段",
        "字段轴", "轴序号", "表头单元格", "公式单元格数", "常量单元格数", "外链公式数",
        "公式常量混合", "公式示例地址", "自动写入策略",
    ]
    formula_rows = []
    for item in data["formula_fields"]:
        policy = "禁止自动写入（非导出注释字段）" if item.get("is_comment_field") else "MANUAL_REQUIRED"
        formula_rows.append([
            ";".join(item.get("logical_tables") or []), item.get("source_path"), item.get("schema_sheet"),
            item.get("field_name"), item.get("raw_type"), item.get("is_exported"),
            item.get("is_comment_field"), item.get("axis"), item.get("axis_index"), item.get("header_cell"),
            item.get("formula_cell_count"), item.get("constant_cell_count"),
            item.get("external_formula_cell_count"), item.get("formula_constant_mixed"),
            example_cells(item.get("formula_examples")), policy,
        ])
    add_data_sheet(workbook, "公式字段摘要", formula_headers, formula_rows, "TblFormulaFields")

    unregistered_headers = [
        "相对路径", "文件名", "分类建议", "解析状态", "具有Schema外观", "Schema Sheet",
        "全部Sheet", "导出字段数", "公式数", "会复制到TempDatas", "被__tables__注册",
        "运行时被Luban消费", "门禁", "处理建议",
    ]
    unregistered_rows = []
    for item in data["unregistered_workbooks"]:
        hint = item.get("classification_hint")
        gate = "WARN" if hint == "backup" else "REVIEW"
        unregistered_rows.append([
            item.get("source_path"), item.get("basename"), hint, item.get("parse_status"),
            item.get("schema_like"), ";".join(item.get("schema_sheet_names") or []),
            ";".join(item.get("sheet_names") or []), item.get("exported_field_count"),
            item.get("formula_count"), True, False, False, gate,
            "确认其为工具/生成/废弃/漏注册；未确认前不得视为正式运行时配置。",
        ])
    add_data_sheet(
        workbook, "未注册工作簿", unregistered_headers, unregistered_rows, "TblUnregistered"
    )

    anomaly_headers = [
        "严重度", "代码", "来源路径", "Sheet", "详情", "上报范围", "实际最大行",
        "实际最大列", "门禁结果", "处理建议",
    ]
    anomaly_rows = []
    for item in data["anomalies"]:
        severity = str(item.get("severity") or "info").lower()
        gate = "BLOCK" if severity in {"error", "fatal"} else ("REVIEW" if severity == "warn" else "INFO")
        anomaly_rows.append([
            severity.upper(), item.get("code"), item.get("source_path"), item.get("sheet_name"),
            item.get("detail"), item.get("reported_ref"), item.get("actual_max_row"),
            item.get("actual_max_col"), gate,
            "按代码复核；不得以 worksheet.max_row/max_column 或 Sheet 名代替真实 Schema 解析。",
        ])
    add_data_sheet(
        workbook, "Schema异常与门禁", anomaly_headers, anomaly_rows, "TblSchemaAnomalies"
    )

    manual = preserved.get("module_manual", {})
    module_rows = []
    for item in data["table_inputs"]:
        source = str(item.get("source_path") or "")
        schema_sheet = str(item.get("schema_sheet") or "")
        key = (str(item["logical_table"]), source, schema_sheet)
        domain, confidence = classify_domain(source)
        top = source.replace("\\", "/").split("/", 1)[0] if source else ""
        manual_values = manual.get(key, {})
        module_rows.append([
            item["logical_table"], item["registry_row"], item["input_index"], item["input_name"],
            source, schema_sheet, top, domain, confidence,
            *[manual_values.get(header, "") for header in MODULE_HEADERS[9:]],
        ])
    module_sheet = add_data_sheet(workbook, "模块映射", MODULE_HEADERS, module_rows, "TblModuleMap")
    add_list_validation(module_sheet, MODULE_HEADERS.index("关系类型") + 1, [
        "UNCLASSIFIED", "DIRECT", "INDIRECT", "NO_NUMERIC", "CODE", "SCENE", "NEW_TABLE", "DEPRECATED"
    ])
    add_list_validation(module_sheet, MODULE_HEADERS.index("映射状态") + 1, [
        "待建立字段级映射", "待确认", "已映射", "缺正式承载", "设计与旧配置冲突", "不需要数值落表"
    ])

    draft_headers = [
        "批次ID", "草案ID", "批次版本", "模块", "参数ID", "来源MD", "来源章节",
        "逻辑表", "源工作簿", "Sheet", "主键字段", "主键值", "目标字段", "RawType",
        "当前值", "建议值", "变换公式", "单位与舍入", "Bean或Enum约束", "引用影响",
        "公式单元格", "写入方式", "风险等级", "校验方法", "草案状态", "创建时间", "备注",
    ]
    draft_sheet = create_staging_sheet(
        workbook, "配置草案", draft_headers, preserved.get("staging", {}).get("配置草案", []), "TblDrafts"
    )
    add_list_validation(draft_sheet, draft_headers.index("写入方式") + 1, [
        "CANDIDATE_SAFE_RECHECK", "MANUAL_REQUIRED", "SCHEMA_CHANGE", "NEW_ROW_MANUAL"
    ])
    add_list_validation(draft_sheet, draft_headers.index("风险等级") + 1, ["LOW", "MEDIUM", "HIGH"])
    add_list_validation(draft_sheet, draft_headers.index("草案状态") + 1, [
        "DRAFT", "READY_FOR_REVIEW", "APPROVED", "REJECTED", "NEEDS_REVISION", "SUPERSEDED",
        "MANUAL_REQUIRED", "APPLIED", "EXPORTED", "VERIFIED"
    ])

    approval_headers = [
        "审批ID", "批次ID", "批次版本", "草案范围哈希", "审批决定", "批准范围", "例外项",
        "审批人", "审批时间", "审批备注",
    ]
    approval_sheet = create_staging_sheet(
        workbook, "审批记录", approval_headers, preserved.get("staging", {}).get("审批记录", []), "TblApprovals"
    )
    add_list_validation(approval_sheet, approval_headers.index("审批决定") + 1, [
        "APPROVED", "REJECTED", "NEEDS_REVISION", "SUPERSEDED"
    ])

    receipt_headers = [
        "回执ID", "批次ID", "批次版本", "执行时间", "执行人", "备份路径", "源工作簿", "Sheet",
        "主键范围", "字段范围", "写入结果", "Luban命令", "Luban结果", "规则校验", "生成物差异",
        "回滚路径", "回执状态", "备注",
    ]
    receipt_sheet = create_staging_sheet(
        workbook, "落地回执", receipt_headers, preserved.get("staging", {}).get("落地回执", []), "TblReceipts"
    )
    add_list_validation(receipt_sheet, receipt_headers.index("回执状态") + 1, [
        "APPLIED", "EXPORTED", "VERIFIED", "FAILED", "ROLLED_BACK"
    ])

    gate_headers = ["门禁", "阶段", "当前状态", "通过条件", "禁止事项", "责任证据"]
    gate_rows = [
        ["G0", "Schema镜像", "PASS_WITH_WARNINGS", "注册输入全部解析；异常与未注册项已显式列出", "不得复制业务数据", "本工作簿+JSON报告"],
        ["G1", "MD追踪", "NOT_STARTED", "每个配置数值有参数ID与来源MD", "不得在草案中另立参数权威", "来源MD/参数ID"],
        ["G2", "配置草案", "NOT_STARTED", "字段级旧值/建议值/转换/影响完整", "不得直接写Datas", "配置草案批次"],
        ["G3", "预检", "NOT_STARTED", "类型/Bean/Enum/引用/公式/风险通过", "不得跳过公式复查", "预检结果"],
        ["G4", "人工批准", "NOT_STARTED", "用户明确批准当前批次版本与范围", "不得复用模糊同意或旧版本批准", "审批记录"],
        ["G5", "正式写入授权", "NOT_STARTED", "用户另行授权修改SVN正式配置", "未授权时Datas只读", "授权记录"],
        ["G6", "备份与漂移复查", "NOT_STARTED", "备份完成且目标单元格未漂移/非公式", "不得插删行列或覆写公式", "备份路径+复查"],
        ["G7", "精确写入", "NOT_STARTED", "只修改获批安全单元格", "不得修改server_json/TempDatas", "写入差异"],
        ["G8", "Luban导出", "NOT_STARTED", "导出成功且规则检查通过", "不得把生成物当源表", "命令与日志"],
        ["G9", "验证与回执", "NOT_STARTED", "生成物差异符合批次并可回滚", "不得自动代替用户最终验收", "落地回执"],
    ]
    add_data_sheet(workbook, "工作流门禁", gate_headers, gate_rows, "TblWorkflowGates")

    summary = workbook.create_sheet("使用说明", 0)
    summary.sheet_view.showGridLines = False
    summary.merge_cells("A1:H1")
    summary["A1"] = "Vgame 配置映射总表（Schema 镜像 / 配置草案人工门禁）"
    summary["A1"].fill = TITLE_FILL
    summary["A1"].font = Font(name=FONT_NAME, size=16, bold=True, color="FFFFFF")
    summary["A1"].alignment = Alignment(horizontal="center", vertical="center")
    summary.row_dimensions[1].height = 34
    summary.merge_cells("A2:H2")
    summary["A2"] = "本工作簿不复制正式配置业务数据；配置草案、审批记录和落地回执初始为空。用户批准具体批次前，Vgame Datas 保持只读。"
    summary["A2"].fill = WARN_FILL
    summary["A2"].font = Font(name=FONT_NAME, size=10, bold=True, color="9C0006")
    summary["A2"].alignment = Alignment(wrap_text=True, vertical="center")
    summary.row_dimensions[2].height = 36

    summary_rows = [
        ["生成标记", MARKER, "Schema版本", data.get("schema_version")],
        ["生成时间", data.get("generated_at"), "正式配置根目录", data.get("config_root")],
        ["源快照SHA256", data.get("source_snapshot_hash"), "快照文件数", data["counts"].get("snapshot_files")],
        ["逻辑表", Formula(f"=COUNTA('逻辑表清单'!A2:A{len(logical_rows)+1})"), "注册输入", Formula(f"=COUNTA('表输入关系'!A2:A{len(input_rows)+1})")],
        ["逻辑字段", Formula(f"=COUNTA('逻辑字段字典'!A2:A{len(logical_field_rows)+1})"), "物理字段（含注释）", Formula(f"=COUNTA('物理字段来源'!A2:A{len(physical_rows)+1})")],
        ["Bean字段", Formula(f"=COUNTA('Bean与字段'!A2:A{len(bean_rows)+1})"), "Enum枚举项", Formula(f"=COUNTA('Enum与枚举项'!A2:A{len(enum_rows)+1})")],
        ["公式字段", Formula(f"=COUNTA('公式字段摘要'!A2:A{len(formula_rows)+1})"), "未注册工作簿", Formula(f"=COUNTA('未注册工作簿'!A2:A{len(unregistered_rows)+1})")],
        ["Schema异常/提示", Formula(f"=COUNTA('Schema异常与门禁'!A2:A{len(anomaly_rows)+1})"), "当前草案业务行", Formula("=COUNTA('配置草案'!A2:A5000)")],
        ["当前阶段", "S1 Schema镜像已生成，等待用户审阅", "正式配置写入", "未授权 / 未执行"],
    ]
    start = 4
    for r_index, row in enumerate(summary_rows, start=start):
        for c_index, value in enumerate(row, start=1):
            cell = summary.cell(r_index, c_index, display(value))
            cell.font = BODY_FONT
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = BODY_BORDER
            if c_index in {1, 3}:
                cell.fill = SUBHEADER_FILL
                cell.font = Font(name=FONT_NAME, size=10, bold=True)
    summary.column_dimensions["A"].width = 22
    summary.column_dimensions["B"].width = 48
    summary.column_dimensions["C"].width = 22
    summary.column_dimensions["D"].width = 54
    for column in "EFGH":
        summary.column_dimensions[column].width = 12
    summary.freeze_panes = "A4"

    note_row = start + len(summary_rows) + 2
    summary.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=8)
    summary.cell(note_row, 1, "标准链路：01源表+03沉淀MD → 04数值文档/05验算 → 配置草案 → 用户批准 → Datas → Luban导出与回执")
    summary.cell(note_row, 1).fill = INFO_FILL
    summary.cell(note_row, 1).font = Font(name=FONT_NAME, size=10, bold=True)
    summary.cell(note_row, 1).alignment = Alignment(wrap_text=True)

    meta = workbook.create_sheet(META_SHEET)
    meta.append(["generation_marker", MARKER])
    meta.append(["schema_version", data.get("schema_version")])
    meta.append(["generated_at", data.get("generated_at")])
    meta.append(["source_snapshot_hash", data.get("source_snapshot_hash")])
    meta.append(["config_root", data.get("config_root")])
    meta.append(["report_path", str(report_path)])
    meta.append(["official_config_write", False])
    meta.sheet_state = "veryHidden"

    sheet_counts = {
        "逻辑表清单": len(logical_rows),
        "表输入关系": len(input_rows),
        "逻辑字段字典": len(logical_field_rows),
        "物理字段来源": len(physical_rows),
        "Bean与字段": len(bean_rows),
        "Enum与枚举项": len(enum_rows),
        "公式字段摘要": len(formula_rows),
        "未注册工作簿": len(unregistered_rows),
        "Schema异常与门禁": len(anomaly_rows),
        "模块映射": len(module_rows),
        "配置草案": len(preserved.get("staging", {}).get("配置草案", [])),
        "审批记录": len(preserved.get("staging", {}).get("审批记录", [])),
        "落地回执": len(preserved.get("staging", {}).get("落地回执", [])),
        "工作流门禁": len(gate_rows),
    }
    return workbook, sheet_counts


def write_json_atomic(value: dict[str, Any], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(path.name + ".tmp")
    temporary.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")
    temporary.replace(path)


def write_text_atomic(text: str, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(path.name + ".tmp")
    temporary.write_text(text, encoding="utf-8")
    temporary.replace(path)


def save_workbook_atomic(workbook: Workbook, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(path.stem + ".tmp" + path.suffix)
    temporary.unlink(missing_ok=True)
    workbook.save(temporary)
    workbook.close()
    check = load_workbook(temporary, read_only=True, data_only=False, keep_links=False)
    check.close()
    temporary.replace(path)


def build_md_report(report: dict[str, Any], workbook_path: Path, report_path: Path) -> str:
    counts = report["counts"]
    anomalies = report.get("anomaly_counts") or {}
    return f"""# P6 Vgame 配置映射总表：生成与审计

生成时间：`{report['generated_at']}`

## 产物

- 工作簿：`{workbook_path}`
- 机器报告：`{report_path}`
- Schema 来源：`{report['config_root']}`
- 源快照：`{report['source_snapshot_hash']}`（`{counts.get('snapshot_files')}` 个文件）

## 覆盖

| 项目 | 数量 |
|---|---:|
| 逻辑表 | {counts.get('logical_tables')} |
| 注册输入 | {counts.get('registered_inputs')} |
| Schema Sheet | {counts.get('schema_worksheets')} |
| 辅助 Sheet | {counts.get('auxiliary_worksheets')} |
| 物理导出字段 | {counts.get('physical_exported_fields')} |
| 逻辑导出字段 | {counts.get('logical_exported_fields')} |
| Bean / 字段 | {counts.get('beans')} / {counts.get('bean_fields')} |
| Enum / Item | {counts.get('enums')} / {counts.get('enum_items')} |
| Schema 公式单元格 | {counts.get('schema_formula_cells')} |
| 未注册工作簿 | {counts.get('unregistered_workbooks')} |

## 已显式保留的边界

- 3 个纵向 Schema、4 个两行表头文件、105 个辅助 Sheet 均按真实结构解析。
- `Drop/DropName` 多输入类型冲突未被压平，保留到逻辑/物理两层字段字典。
- 公式只记录字段、计数、示例地址和自动写入风险；未复制正式配置业务值。
- `配置草案`、`审批记录`、`落地回执` 初始无业务行。
- 未注册工作簿只列为审计对象，不视为正式运行时配置。
- 本次只读正式 Datas；未执行写入、Luban 导出或生成物修改。

## 异常分类

```json
{json.dumps(anomalies, ensure_ascii=False, indent=2)}
```

## 下一门禁

由用户先审阅工作簿结构。之后选择一个明确模块，按 MD 参数 ID 生成首个 `配置草案` 批次；在用户批准该批次且另行授权 SVN 写入前，正式配置继续保持只读。
"""


def resolve_path(value: str | None, default: Path) -> Path:
    if not value:
        return default.resolve()
    path = Path(value)
    return path.resolve() if path.is_absolute() else (V2_DIR / path).resolve()


def main() -> int:
    parser = argparse.ArgumentParser(description="生成 Vgame 全量配置 Schema 映射总表")
    parser.add_argument("--config-root", default=str(DEFAULT_CONFIG_ROOT))
    parser.add_argument("--workbook", default=str(DEFAULT_WORKBOOK))
    parser.add_argument("--report", default=str(DEFAULT_REPORT))
    parser.add_argument("--md-report", default=str(DEFAULT_MD_REPORT))
    parser.add_argument("--force", action="store_true", help="允许覆盖没有生成标记的目标文件")
    parser.add_argument("--no-backup", action="store_true", help="刷新已生成工作簿时不创建备份")
    args = parser.parse_args()

    config_root = resolve_path(args.config_root, DEFAULT_CONFIG_ROOT)
    workbook_path = resolve_path(args.workbook, DEFAULT_WORKBOOK)
    report_path = resolve_path(args.report, DEFAULT_REPORT)
    md_report_path = resolve_path(args.md_report, DEFAULT_MD_REPORT)
    if not config_root.is_dir():
        parser.error(f"配置根目录不存在: {config_root}")

    preserved = read_preserved(workbook_path, args.force)
    data = extract_config_schema(config_root)
    if workbook_path.exists() and not args.no_backup:
        backup_dir = workbook_path.parent / "_backups"
        backup_dir.mkdir(parents=True, exist_ok=True)
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        shutil.copy2(workbook_path, backup_dir / f"{workbook_path.stem}_{stamp}{workbook_path.suffix}")

    workbook, sheet_counts = build_workbook(data, preserved, report_path)
    save_workbook_atomic(workbook, workbook_path)

    report = dict(data)
    report.update({
        "generation_marker": MARKER,
        "mapping_workbook": str(workbook_path),
        "sheet_counts": sheet_counts,
        "official_config_write": False,
        "business_values_copied": False,
        "staging_rows": {name: sheet_counts[name] for name in STAGING_SHEETS},
    })
    write_json_atomic(report, report_path)
    write_text_atomic(build_md_report(report, workbook_path, report_path), md_report_path)

    print(f"[PASS] workbook={workbook_path}")
    print(f"[PASS] report={report_path}")
    print(f"[PASS] md_report={md_report_path}")
    print(f"[INFO] source_snapshot_hash={data['source_snapshot_hash']}")
    print(f"[INFO] sheet_counts={json.dumps(sheet_counts, ensure_ascii=False)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
