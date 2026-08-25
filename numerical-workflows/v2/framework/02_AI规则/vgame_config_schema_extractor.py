from __future__ import annotations

import hashlib
import posixpath
import re
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable
from xml.etree import ElementTree as ET
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.utils.cell import column_index_from_string, get_column_letter, range_boundaries


SCHEMA_VERSION = "vgame-config-schema-extractor/v1"
_SCHEMA_FILES = {"__tables__.xlsx", "__beans__.xlsx", "__enums__.xlsx"}
_SCHEMA_MARKERS = {"##var", "##column#var"}
_CELL_REF_RE = re.compile(r"^([A-Z]+)([0-9]+)$")
_MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_DOC_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"


def _q(namespace: str, name: str) -> str:
    return f"{{{namespace}}}{name}"


def _json_value(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return str(value)


def _normalized_bool(value: Any) -> bool | None:
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)) and value in (0, 1):
        return bool(value)
    if isinstance(value, str):
        text = value.strip().upper().replace(" ", "")
        if text in {"TRUE", "=TRUE()", "1"}:
            return True
        if text in {"FALSE", "=FALSE()", "0"}:
            return False
    return None


def _iter_rows_until_blank(
    worksheet: Any,
    *,
    start_row: int,
    max_col: int,
    blank_limit: int = 100,
    hard_limit: int = 100_000,
) -> Iterable[tuple[int, tuple[Any, ...]]]:
    blank_run = 0
    for row_number, values in enumerate(
        worksheet.iter_rows(
            min_row=start_row,
            max_row=hard_limit,
            min_col=1,
            max_col=max_col,
            values_only=True,
        ),
        start_row,
    ):
        if all(value is None for value in values):
            blank_run += 1
            if blank_run >= blank_limit:
                break
            continue
        blank_run = 0
        yield row_number, values


def _read_table_registry(path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        worksheet = workbook.active
        tables: list[dict[str, Any]] = []
        for row_number, values in _iter_rows_until_blank(
            worksheet, start_row=4, max_col=12
        ):
            full_name = values[1]
            if full_name is None or not str(full_name).strip():
                continue
            raw_input = values[4]
            inputs = (
                [item.strip() for item in str(raw_input).split(",") if item.strip()]
                if raw_input is not None
                else []
            )
            tables.append(
                {
                    "registry_row": row_number,
                    "logical_table": str(full_name).strip(),
                    "value_type": _json_value(values[2]),
                    "read_schema_from_file_raw": _json_value(values[3]),
                    "read_schema_from_file": _normalized_bool(values[3]),
                    "inputs": inputs,
                    "gen_csharp_class_raw": _json_value(values[5]),
                    "gen_csharp_class": _normalized_bool(values[5]),
                    "index_raw": _json_value(values[6]),
                    "mode": _json_value(values[7]),
                    "group": _json_value(values[8]),
                    "comment": _json_value(values[9]),
                    "tags": _json_value(values[10]),
                    "output": _json_value(values[11]),
                }
            )
        return tables
    finally:
        workbook.close()


def _read_beans(path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        worksheet = workbook.active
        beans: list[dict[str, Any]] = []
        current: dict[str, Any] | None = None
        for row_number, values in _iter_rows_until_blank(
            worksheet, start_row=4, max_col=17
        ):
            full_name = values[1]
            if full_name is not None and str(full_name).strip():
                current = {
                    "definition_row": row_number,
                    "full_name": str(full_name).strip(),
                    "parent": _json_value(values[2]),
                    "value_type": _json_value(values[3]),
                    "separator": _json_value(values[4]),
                    "alias": _json_value(values[5]),
                    "comment": _json_value(values[6]),
                    "group": _json_value(values[7]),
                    "tags": _json_value(values[8]),
                    "fields": [],
                }
                beans.append(current)
            field_name = values[9]
            if field_name is not None and str(field_name).strip() and current is not None:
                current["fields"].append(
                    {
                        "definition_row": row_number,
                        "field_name": str(field_name).strip(),
                        "raw_type": _json_value(values[10]),
                        "group": _json_value(values[11]),
                        "comment": _json_value(values[12]),
                        "tags": _json_value(values[13]),
                        "variants": _json_value(values[14]),
                    }
                )
        return beans
    finally:
        workbook.close()


def _read_enums(path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        worksheet = workbook.active
        enums: list[dict[str, Any]] = []
        current: dict[str, Any] | None = None
        for row_number, values in _iter_rows_until_blank(
            worksheet, start_row=4, max_col=15
        ):
            full_name = values[1]
            if full_name is not None and str(full_name).strip():
                current = {
                    "definition_row": row_number,
                    "full_name": str(full_name).strip(),
                    "flags_raw": _json_value(values[2]),
                    "flags": _normalized_bool(values[2]),
                    "unique_raw": _json_value(values[3]),
                    "unique": _normalized_bool(values[3]),
                    "group": _json_value(values[4]),
                    "comment": _json_value(values[5]),
                    "tags": _json_value(values[6]),
                    "gen_csharp_class_raw": _json_value(values[12]),
                    "gen_csharp_class": _normalized_bool(values[12]),
                    "items": [],
                }
                enums.append(current)
            item_name = values[7]
            item_value = values[9]
            if (item_name is not None or item_value is not None) and current is not None:
                current["items"].append(
                    {
                        "definition_row": row_number,
                        "item_name": _json_value(item_name),
                        "alias": _json_value(values[8]),
                        "value": _json_value(item_value),
                        "value_is_implicit": item_value is None,
                        "comment": _json_value(values[10]),
                        "tags": _json_value(values[11]),
                    }
                )
        return enums
    finally:
        workbook.close()


def _read_shared_strings(archive: ZipFile) -> list[str]:
    name = "xl/sharedStrings.xml"
    if name not in archive.namelist():
        return []
    root = ET.fromstring(archive.read(name))
    return [
        "".join(node.text or "" for node in item.iter(_q(_MAIN_NS, "t")))
        for item in root.findall(_q(_MAIN_NS, "si"))
    ]


def _decode_cell(cell: ET.Element, shared_strings: list[str]) -> Any:
    cell_type = cell.attrib.get("t")
    if cell_type == "inlineStr":
        inline = cell.find(_q(_MAIN_NS, "is"))
        if inline is None:
            return ""
        return "".join(node.text or "" for node in inline.iter(_q(_MAIN_NS, "t")))
    value_node = cell.find(_q(_MAIN_NS, "v"))
    if value_node is None:
        return None
    text = value_node.text or ""
    if cell_type == "s":
        try:
            return shared_strings[int(text)]
        except (IndexError, TypeError, ValueError):
            return text
    if cell_type in {"str", "e"}:
        return text
    if cell_type == "b":
        return text == "1"
    try:
        number = float(text)
        return int(number) if number.is_integer() else number
    except ValueError:
        return text


def _sheet_target(target: str) -> str:
    if target.startswith("/"):
        return target.lstrip("/")
    return posixpath.normpath(posixpath.join("xl", target))


def _parse_dimension(reference: str | None) -> dict[str, Any]:
    result = {
        "reported_ref": reference,
        "reported_min_row": None,
        "reported_min_col": None,
        "reported_max_row": None,
        "reported_max_col": None,
    }
    if not reference:
        return result
    try:
        min_col, min_row, max_col, max_row = range_boundaries(reference)
    except (TypeError, ValueError):
        return result
    result.update(
        {
            "reported_min_row": min_row,
            "reported_min_col": min_col,
            "reported_max_row": max_row,
            "reported_max_col": max_col,
        }
    )
    return result


def _cell_ref(reference: str) -> tuple[int, int] | None:
    match = _CELL_REF_RE.match(reference)
    if not match:
        return None
    return int(match.group(2)), column_index_from_string(match.group(1))


def _is_external_formula(formula: str | None) -> bool:
    if not formula:
        return False
    lowered = formula.casefold()
    return "[" in formula or ".xlsx" in lowered or ".xlsm" in lowered


def _parse_workbook_xml(path: Path) -> dict[str, Any]:
    with ZipFile(path) as archive:
        shared_strings = _read_shared_strings(archive)
        workbook_root = ET.fromstring(archive.read("xl/workbook.xml"))
        relation_root = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        relations = {
            node.attrib["Id"]: node.attrib["Target"]
            for node in relation_root.findall(_q(_PKG_REL_NS, "Relationship"))
            if "Id" in node.attrib and "Target" in node.attrib
        }
        sheets: list[dict[str, Any]] = []
        sheets_node = workbook_root.find(_q(_MAIN_NS, "sheets"))
        if sheets_node is None:
            raise ValueError(f"Workbook has no sheets node: {path}")
        for sheet_node in sheets_node.findall(_q(_MAIN_NS, "sheet")):
            relationship_id = sheet_node.attrib.get(_q(_DOC_REL_NS, "id"))
            if relationship_id not in relations:
                continue
            target = _sheet_target(relations[relationship_id])
            if target not in archive.namelist():
                continue
            worksheet_root = ET.fromstring(archive.read(target))
            dimension_node = worksheet_root.find(_q(_MAIN_NS, "dimension"))
            dimension = _parse_dimension(
                dimension_node.attrib.get("ref") if dimension_node is not None else None
            )
            cells: dict[tuple[int, int], dict[str, Any]] = {}
            actual_max_row = 0
            actual_max_col = 0
            formula_count = 0
            external_formula_count = 0
            formula_examples: list[dict[str, Any]] = []
            for cell in worksheet_root.findall(f".//{_q(_MAIN_NS, 'c')}"):
                reference = cell.attrib.get("r")
                if not reference:
                    continue
                position = _cell_ref(reference)
                if position is None:
                    continue
                formula_node = cell.find(_q(_MAIN_NS, "f"))
                formula = None
                if formula_node is not None:
                    formula = "=" + (formula_node.text or "")
                    formula_count += 1
                    if _is_external_formula(formula):
                        external_formula_count += 1
                    if len(formula_examples) < 5:
                        formula_examples.append({"cell": reference, "formula": formula})
                value = _decode_cell(cell, shared_strings)
                has_value = value is not None and value != ""
                if not has_value and formula_node is None:
                    continue
                row_number, column_number = position
                actual_max_row = max(actual_max_row, row_number)
                actual_max_col = max(actual_max_col, column_number)
                cells[position] = {
                    "cell": reference,
                    "value": _json_value(value),
                    "formula": formula,
                    "is_formula": formula_node is not None,
                    "is_external_formula": _is_external_formula(formula),
                }
            dimension.update(
                {"actual_max_row": actual_max_row, "actual_max_col": actual_max_col}
            )
            sheets.append(
                {
                    "name": sheet_node.attrib.get("name", ""),
                    "state": sheet_node.attrib.get("state", "visible"),
                    "xml_path": target,
                    "dimension": dimension,
                    "formula_count": formula_count,
                    "external_formula_count": external_formula_count,
                    "formula_examples": formula_examples,
                    "cells": cells,
                }
            )
        external_parts = sorted(
            name for name in archive.namelist() if name.startswith("xl/externalLinks/")
        )
        return {"sheets": sheets, "external_link_parts": external_parts}


def _cell_value(sheet: dict[str, Any], row: int, column: int) -> Any:
    cell = sheet["cells"].get((row, column))
    return cell["value"] if cell is not None else None


def _formula_profile(cells: Iterable[dict[str, Any]]) -> dict[str, Any]:
    formula_count = 0
    constant_count = 0
    external_count = 0
    formula_examples: list[dict[str, Any]] = []
    for cell in cells:
        if cell["is_formula"]:
            formula_count += 1
            external_count += int(cell["is_external_formula"])
            if len(formula_examples) < 3:
                formula_examples.append(
                    {"cell": cell["cell"], "formula": cell["formula"]}
                )
        else:
            constant_count += 1
    return {
        "formula_cell_count": formula_count,
        "constant_cell_count": constant_count,
        "external_formula_cell_count": external_count,
        "formula_examples": formula_examples,
        "formula_constant_mixed": formula_count > 0 and constant_count > 0,
    }


def _horizontal_schema(sheet: dict[str, Any]) -> dict[str, Any]:
    group_row = 3 if _cell_value(sheet, 3, 1) == "##group" else None
    comment_row = None
    if _cell_value(sheet, 3, 1) == "##":
        comment_row = 3
    elif _cell_value(sheet, 4, 1) == "##":
        comment_row = 4
    metadata_rows = [1, 2]
    if group_row is not None:
        metadata_rows.append(group_row)
    if comment_row is not None:
        metadata_rows.append(comment_row)
    data_start_row = max(metadata_rows) + 1
    if group_row == 3 and comment_row == 4:
        header_layout = "4-row"
    elif group_row is None and comment_row is None:
        header_layout = "2-row"
    else:
        header_layout = "custom"
    header_columns = sorted(
        column
        for (row, column), cell in sheet["cells"].items()
        if row == 1 and column > 1 and cell["value"] not in (None, "")
    )
    fields: list[dict[str, Any]] = []
    data_rows: set[int] = set()
    for ordinal, column in enumerate(header_columns, 1):
        name = str(_cell_value(sheet, 1, column)).strip()
        is_comment = name.startswith("#")
        field_cells = [
            cell
            for (row, cell_column), cell in sheet["cells"].items()
            if cell_column == column and row >= data_start_row
        ]
        data_rows.update(
            row
            for (row, cell_column), cell in sheet["cells"].items()
            if cell_column == column and row >= data_start_row
        )
        field = {
            "source_ordinal": ordinal,
            "axis": "column",
            "axis_index": column,
            "excel_axis": get_column_letter(column),
            "header_cell": f"{get_column_letter(column)}1",
            "field_name": name,
            "raw_type": _json_value(_cell_value(sheet, 2, column)),
            "group": _json_value(
                _cell_value(sheet, group_row, column) if group_row is not None else None
            ),
            "comment": _json_value(
                _cell_value(sheet, comment_row, column)
                if comment_row is not None
                else None
            ),
            "is_comment_field": is_comment,
            "is_exported": not is_comment,
        }
        field.update(_formula_profile(field_cells))
        fields.append(field)
    return {
        "orientation": "horizontal",
        "header_layout": header_layout,
        "group_row": group_row,
        "comment_row": comment_row,
        "data_start_row": data_start_row,
        "data_row_count": len(data_rows),
        "fields": fields,
    }


def _vertical_schema(sheet: dict[str, Any]) -> dict[str, Any]:
    field_rows = sorted(
        row
        for (row, column), cell in sheet["cells"].items()
        if column == 1
        and row >= 4
        and cell["value"] not in (None, "")
        and not str(cell["value"]).strip().startswith("#")
    )
    max_column = sheet["dimension"]["actual_max_col"] or 2
    comment_columns = {
        column
        for column in range(3, max_column + 1)
        if str(_cell_value(sheet, 1, column) or "").strip().startswith("#")
    }
    value_columns = {
        column
        for row in field_rows
        for column in range(3, max_column + 1)
        if column not in comment_columns and (row, column) in sheet["cells"]
    }
    fields: list[dict[str, Any]] = []
    for ordinal, row in enumerate(field_rows, 1):
        name = str(_cell_value(sheet, row, 1)).strip()
        comments = [
            str(_cell_value(sheet, row, column))
            for column in sorted(comment_columns)
            if _cell_value(sheet, row, column) not in (None, "")
        ]
        value_cells = [
            sheet["cells"][(row, column)]
            for column in sorted(value_columns)
            if (row, column) in sheet["cells"]
        ]
        field = {
            "source_ordinal": ordinal,
            "axis": "row",
            "axis_index": row,
            "excel_axis": str(row),
            "header_cell": f"A{row}",
            "field_name": name,
            "raw_type": _json_value(_cell_value(sheet, row, 2)),
            "group": None,
            "comment": " | ".join(comments) if comments else None,
            "is_comment_field": False,
            "is_exported": True,
        }
        field.update(_formula_profile(value_cells))
        fields.append(field)
    return {
        "orientation": "vertical",
        "header_layout": "column-var",
        "group_row": None,
        "comment_row": None,
        "data_start_row": 4,
        "data_row_count": len(value_columns),
        "value_columns": [get_column_letter(column) for column in sorted(value_columns)],
        "comment_columns": [
            get_column_letter(column) for column in sorted(comment_columns)
        ],
        "fields": fields,
    }


def _parse_config_workbook(path: Path, relative_path: str) -> dict[str, Any]:
    parsed = _parse_workbook_xml(path)
    schema_sheets = [
        sheet
        for sheet in parsed["sheets"]
        if _cell_value(sheet, 1, 1) in _SCHEMA_MARKERS
    ]
    schema_sheet = schema_sheets[0] if len(schema_sheets) == 1 else None
    schema: dict[str, Any] | None = None
    if schema_sheet is not None:
        marker = _cell_value(schema_sheet, 1, 1)
        schema = (
            _horizontal_schema(schema_sheet)
            if marker == "##var"
            else _vertical_schema(schema_sheet)
        )
        schema.update(
            {
                "sheet_name": schema_sheet["name"],
                "marker": marker,
                "dimension": schema_sheet["dimension"],
            }
        )
    formula_total = sum(sheet["formula_count"] for sheet in parsed["sheets"])
    external_formula_total = sum(
        sheet["external_formula_count"] for sheet in parsed["sheets"]
    )
    schema_formula_count = schema_sheet["formula_count"] if schema_sheet else 0
    return {
        "source_path": relative_path,
        "basename": path.name,
        "sheet_count": len(parsed["sheets"]),
        "sheet_names": [sheet["name"] for sheet in parsed["sheets"]],
        "schema_sheet_names": [sheet["name"] for sheet in schema_sheets],
        "schema_sheet_count": len(schema_sheets),
        "schema": schema,
        "sheets": [
            {
                "sheet_name": sheet["name"],
                "state": sheet["state"],
                "dimension": sheet["dimension"],
                "formula_count": sheet["formula_count"],
                "external_formula_count": sheet["external_formula_count"],
                "formula_examples": sheet["formula_examples"],
                "is_schema_sheet": sheet in schema_sheets,
            }
            for sheet in parsed["sheets"]
        ],
        "formula_count": formula_total,
        "schema_formula_count": schema_formula_count,
        "auxiliary_formula_count": formula_total - schema_formula_count,
        "external_formula_count": external_formula_total,
        "external_link_parts": parsed["external_link_parts"],
    }


def _type_parts(raw_type: Any) -> dict[str, Any]:
    if raw_type is None:
        return {
            "base_type": None,
            "is_optional": False,
            "collection_kind": None,
            "separator": None,
            "default_raw": None,
        }
    text = str(raw_type).strip()
    collection_kind = None
    separator = None
    remainder = text
    collection_match = re.match(r"^\((array|list)(?:#sep=([^)]*))?\),(.*)$", text)
    if collection_match:
        collection_kind = collection_match.group(1)
        separator = collection_match.group(2)
        remainder = collection_match.group(3).strip()
    is_optional = remainder.endswith("?")
    if is_optional:
        remainder = remainder[:-1]
    default_raw = None
    default_match = re.search(r"#(?:default|defalut|defualt)=(.*)$", remainder)
    if default_match:
        default_raw = default_match.group(1)
        remainder = remainder[: default_match.start()]
    return {
        "base_type": remainder.strip() or None,
        "is_optional": is_optional,
        "collection_kind": collection_kind,
        "separator": separator,
        "default_raw": default_raw,
    }


def _snapshot_hash(root: Path, paths: Iterable[Path]) -> dict[str, Any]:
    digest = hashlib.sha256()
    files: list[dict[str, Any]] = []
    for path in sorted(set(paths), key=lambda item: item.relative_to(root).as_posix()):
        relative = path.relative_to(root).as_posix()
        file_digest = hashlib.sha256()
        with path.open("rb") as stream:
            for chunk in iter(lambda: stream.read(1024 * 1024), b""):
                file_digest.update(chunk)
        size = path.stat().st_size
        sha256 = file_digest.hexdigest()
        digest.update(relative.encode("utf-8"))
        digest.update(b"\0")
        digest.update(str(size).encode("ascii"))
        digest.update(b"\0")
        digest.update(sha256.encode("ascii"))
        digest.update(b"\n")
        files.append({"path": relative, "bytes": size, "sha256": sha256})
    return {
        "algorithm": "sha256-path-size-content-v1",
        "sha256": digest.hexdigest(),
        "files": files,
    }


def _add_anomaly(
    anomalies: list[dict[str, Any]],
    code: str,
    severity: str,
    detail: str,
    **context: Any,
) -> None:
    item = {"code": code, "severity": severity, "detail": detail}
    item.update({key: _json_value(value) for key, value in context.items()})
    anomalies.append(item)


def extract_config_schema(config_root: Path) -> dict[str, Any]:
    """Extract the Vgame Luban schema without writing to the config repository."""

    root = Path(config_root).expanduser().resolve()
    if not root.is_dir():
        raise FileNotFoundError(f"Config root does not exist: {root}")
    required = {name: root / name for name in _SCHEMA_FILES}
    missing_schema = [name for name, path in required.items() if not path.is_file()]
    if missing_schema:
        raise FileNotFoundError(
            f"Missing schema workbooks under {root}: {', '.join(sorted(missing_schema))}"
        )

    anomalies: list[dict[str, Any]] = []
    tables = _read_table_registry(required["__tables__.xlsx"])
    beans = _read_beans(required["__beans__.xlsx"])
    enums = _read_enums(required["__enums__.xlsx"])

    spreadsheets = sorted(
        [*root.rglob("*.xlsx"), *root.rglob("*.xlsm")],
        key=lambda path: path.relative_to(root).as_posix().casefold(),
    )
    lock_files = [path for path in spreadsheets if path.name.startswith("~$")]
    non_lock_files = [path for path in spreadsheets if not path.name.startswith("~$")]
    basename_index: dict[str, list[Path]] = defaultdict(list)
    for path in non_lock_files:
        basename_index[path.name.casefold()].append(path)
    for basename, paths in sorted(basename_index.items()):
        if len(paths) > 1:
            _add_anomaly(
                anomalies,
                "duplicate_flattened_basename",
                "block",
                "Export flattens Datas paths, so duplicate basenames are ambiguous.",
                basename=basename,
                source_paths=";".join(
                    path.relative_to(root).as_posix() for path in paths
                ),
            )

    input_references: list[dict[str, Any]] = []
    registered_paths: set[Path] = set()
    for table in tables:
        for input_index, input_name in enumerate(table["inputs"], 1):
            candidates = basename_index.get(Path(input_name).name.casefold(), [])
            reference = {
                "logical_table": table["logical_table"],
                "registry_row": table["registry_row"],
                "input_index": input_index,
                "input_name": input_name,
                "source_path": None,
                "resolution_status": "missing",
                "schema_sheet": None,
                "schema_sheet_name_matches_logical": None,
                "exported_field_count": None,
                "comment_field_count": None,
            }
            if len(candidates) == 1:
                path = candidates[0]
                registered_paths.add(path)
                reference["source_path"] = path.relative_to(root).as_posix()
                reference["resolution_status"] = "resolved"
            elif len(candidates) > 1:
                reference["resolution_status"] = "ambiguous"
                _add_anomaly(
                    anomalies,
                    "ambiguous_registered_input",
                    "block",
                    "Registered input resolves to multiple Datas files.",
                    logical_table=table["logical_table"],
                    input_name=input_name,
                    source_paths=";".join(
                        path.relative_to(root).as_posix() for path in candidates
                    ),
                )
            else:
                _add_anomaly(
                    anomalies,
                    "missing_registered_input",
                    "block",
                    "Registered input cannot be found under Datas.",
                    logical_table=table["logical_table"],
                    input_name=input_name,
                )
            input_references.append(reference)

    physical_inputs: list[dict[str, Any]] = []
    physical_by_path: dict[str, dict[str, Any]] = {}
    for path in sorted(
        registered_paths, key=lambda item: item.relative_to(root).as_posix().casefold()
    ):
        relative = path.relative_to(root).as_posix()
        try:
            physical = _parse_config_workbook(path, relative)
        except (BadZipFile, ET.ParseError, KeyError, ValueError) as error:
            _add_anomaly(
                anomalies,
                "registered_workbook_parse_error",
                "block",
                str(error),
                source_path=relative,
            )
            continue
        physical_inputs.append(physical)
        physical_by_path[relative] = physical
        if physical["schema_sheet_count"] == 0:
            _add_anomaly(
                anomalies,
                "missing_schema_sheet",
                "block",
                "No sheet starts with ##var or ##column#var.",
                source_path=relative,
            )
        elif physical["schema_sheet_count"] > 1:
            _add_anomaly(
                anomalies,
                "multiple_schema_sheets",
                "block",
                "More than one schema marker sheet exists in the workbook.",
                source_path=relative,
                sheet_names=";".join(physical["schema_sheet_names"]),
            )
        schema = physical["schema"]
        if schema is not None:
            fields = schema["fields"]
            exported_names = [
                field["field_name"].casefold()
                for field in fields
                if field["is_exported"]
            ]
            duplicates = sorted(
                name for name, count in Counter(exported_names).items() if count > 1
            )
            if duplicates:
                _add_anomaly(
                    anomalies,
                    "duplicate_exported_field",
                    "block",
                    "Schema contains duplicate exported field names.",
                    source_path=relative,
                    sheet_name=schema["sheet_name"],
                    fields=";".join(duplicates),
                )
            missing_types = [
                field["field_name"]
                for field in fields
                if field["is_exported"] and field["raw_type"] in (None, "")
            ]
            if missing_types:
                _add_anomaly(
                    anomalies,
                    "missing_exported_field_type",
                    "block",
                    "Exported schema fields must declare a raw type.",
                    source_path=relative,
                    sheet_name=schema["sheet_name"],
                    fields=";".join(missing_types),
                )
            if schema["header_layout"] == "2-row":
                _add_anomaly(
                    anomalies,
                    "two_row_header_layout",
                    "info",
                    "Data begins immediately after ##var and ##type.",
                    source_path=relative,
                    sheet_name=schema["sheet_name"],
                    data_start_row=schema["data_start_row"],
                )
            if schema["data_row_count"] == 0:
                _add_anomaly(
                    anomalies,
                    "empty_registered_input",
                    "info",
                    "Registered schema sheet has no populated data records.",
                    source_path=relative,
                    sheet_name=schema["sheet_name"],
                )
            dimension = schema["dimension"]
            reported_row = dimension["reported_max_row"]
            reported_col = dimension["reported_max_col"]
            actual_row = dimension["actual_max_row"]
            actual_col = dimension["actual_max_col"]
            if (
                reported_row is not None
                and reported_col is not None
                and (actual_row > reported_row or actual_col > reported_col)
            ):
                _add_anomaly(
                    anomalies,
                    "worksheet_dimension_underreported",
                    "warn",
                    "Worksheet dimension is smaller than populated XML cells.",
                    source_path=relative,
                    sheet_name=schema["sheet_name"],
                    reported_ref=dimension["reported_ref"],
                    actual_max_row=actual_row,
                    actual_max_col=actual_col,
                )
            if (reported_row or 0) >= 1_048_576 or (reported_col or 0) >= 16_000:
                _add_anomaly(
                    anomalies,
                    "worksheet_dimension_inflated",
                    "warn",
                    "Worksheet dimension is formatting-inflated; do not iterate to max_row/max_column.",
                    source_path=relative,
                    sheet_name=schema["sheet_name"],
                    reported_ref=dimension["reported_ref"],
                    actual_max_row=actual_row,
                    actual_max_col=actual_col,
                )
        if physical["external_formula_count"] or physical["external_link_parts"]:
            _add_anomaly(
                anomalies,
                "external_formula_or_link",
                "warn",
                "Workbook contains an external-style formula or OOXML external link part.",
                source_path=relative,
                external_formula_count=physical["external_formula_count"],
                external_link_part_count=len(physical["external_link_parts"]),
            )

    for reference in input_references:
        source_path = reference["source_path"]
        if source_path is None or source_path not in physical_by_path:
            continue
        physical = physical_by_path[source_path]
        schema = physical["schema"]
        if schema is None:
            continue
        reference["schema_sheet"] = schema["sheet_name"]
        reference["schema_sheet_name_matches_logical"] = (
            schema["sheet_name"].casefold() == reference["logical_table"].casefold()
        )
        reference["exported_field_count"] = sum(
            field["is_exported"] for field in schema["fields"]
        )
        reference["comment_field_count"] = sum(
            field["is_comment_field"] for field in schema["fields"]
        )
        if not reference["schema_sheet_name_matches_logical"]:
            _add_anomaly(
                anomalies,
                "schema_sheet_name_differs_from_logical_table",
                "info",
                "Schema sheet selection must use the marker, not name equality.",
                logical_table=reference["logical_table"],
                source_path=source_path,
                sheet_name=schema["sheet_name"],
            )

    references_by_table: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for reference in input_references:
        references_by_table[reference["logical_table"]].append(reference)
    logical_fields: list[dict[str, Any]] = []
    for table in tables:
        references = [
            reference
            for reference in references_by_table[table["logical_table"]]
            if reference["source_path"] in physical_by_path
            and physical_by_path[reference["source_path"]]["schema"] is not None
        ]
        ordered_names: list[str] = []
        seen_names: set[str] = set()
        exported_by_source: dict[str, dict[str, dict[str, Any]]] = {}
        for reference in references:
            source_path = reference["source_path"]
            schema = physical_by_path[source_path]["schema"]
            source_fields: dict[str, dict[str, Any]] = {}
            for field in schema["fields"]:
                if not field["is_exported"]:
                    continue
                key = field["field_name"].casefold()
                source_fields[key] = field
                if key not in seen_names:
                    seen_names.add(key)
                    ordered_names.append(key)
            exported_by_source[source_path] = source_fields
        for field_ordinal, field_key in enumerate(ordered_names, 1):
            variants: list[dict[str, Any]] = []
            missing_sources: list[str] = []
            for reference in references:
                source_path = reference["source_path"]
                field = exported_by_source[source_path].get(field_key)
                if field is None:
                    missing_sources.append(source_path)
                    continue
                variants.append(
                    {
                        "source_path": source_path,
                        "field_name": field["field_name"],
                        "raw_type": field["raw_type"],
                        "group": field["group"],
                        "comment": field["comment"],
                        "axis": field["axis"],
                        "axis_index": field["axis_index"],
                        "header_cell": field["header_cell"],
                    }
                )
            first_variant = variants[0]
            signatures = {
                (
                    variant["field_name"],
                    str(variant["raw_type"]),
                    str(variant["group"]),
                )
                for variant in variants
            }
            conflict = bool(missing_sources) or len(signatures) > 1
            logical_field = {
                "logical_table": table["logical_table"],
                "field_ordinal": field_ordinal,
                "field_name": first_variant["field_name"],
                "canonical_raw_type": (
                    first_variant["raw_type"] if not conflict else None
                ),
                "canonical_group": first_variant["group"] if not conflict else None,
                "schema_conflict": conflict,
                "missing_sources": missing_sources,
                "source_variants": variants,
            }
            logical_field.update(_type_parts(first_variant["raw_type"]))
            logical_fields.append(logical_field)
            if conflict:
                _add_anomaly(
                    anomalies,
                    "multi_input_schema_conflict",
                    "block",
                    "A multi-input logical table has incompatible field definitions.",
                    logical_table=table["logical_table"],
                    field_name=first_variant["field_name"],
                    variants=";".join(
                        f"{variant['source_path']}={variant['raw_type']}"
                        for variant in variants
                    ),
                    missing_sources=";".join(missing_sources),
                )

    schema_paths = {path.resolve() for path in required.values()}
    unregistered_paths = [
        path
        for path in non_lock_files
        if path.resolve() not in schema_paths and path not in registered_paths
    ]
    unregistered_workbooks: list[dict[str, Any]] = []
    for path in unregistered_paths:
        relative = path.relative_to(root).as_posix()
        item: dict[str, Any] = {
            "source_path": relative,
            "basename": path.name,
            "classification_hint": (
                "backup" if ".bak_" in path.name.casefold() else "unclassified"
            ),
            "parse_status": "ok",
            "schema_like": False,
            "schema_sheet_names": [],
            "sheet_names": [],
            "exported_field_count": 0,
            "formula_count": 0,
        }
        try:
            parsed = _parse_config_workbook(path, relative)
            item.update(
                {
                    "schema_like": parsed["schema_sheet_count"] > 0,
                    "schema_sheet_names": parsed["schema_sheet_names"],
                    "sheet_names": parsed["sheet_names"],
                    "exported_field_count": (
                        sum(
                            field["is_exported"]
                            for field in parsed["schema"]["fields"]
                        )
                        if parsed["schema"] is not None
                        else 0
                    ),
                    "formula_count": parsed["formula_count"],
                }
            )
        except (BadZipFile, ET.ParseError, KeyError, ValueError) as error:
            item["parse_status"] = "error"
            item["parse_error"] = str(error)
        unregistered_workbooks.append(item)
        _add_anomaly(
            anomalies,
            (
                "unregistered_schema_workbook"
                if item["schema_like"]
                else "unregistered_workbook"
            ),
            "warn",
            "Workbook exists under Datas but is not referenced by __tables__.xlsx.",
            source_path=relative,
            classification_hint=item["classification_hint"],
        )

    for enum in enums:
        for item in enum["items"]:
            if item["value_is_implicit"]:
                _add_anomaly(
                    anomalies,
                    "enum_implicit_value",
                    "info",
                    "Enum item has no explicit value; preserve Luban implicit-value semantics.",
                    enum=enum["full_name"],
                    item=item["item_name"],
                    definition_row=item["definition_row"],
                )

    physical_fields: list[dict[str, Any]] = []
    logical_tables_by_source: dict[str, list[str]] = defaultdict(list)
    for reference in input_references:
        if reference["source_path"] is not None:
            logical_tables_by_source[reference["source_path"]].append(
                reference["logical_table"]
            )
    for physical in physical_inputs:
        schema = physical["schema"]
        if schema is None:
            continue
        for field in schema["fields"]:
            item = {
                "logical_tables": sorted(
                    set(logical_tables_by_source[physical["source_path"]])
                ),
                "source_path": physical["source_path"],
                "schema_sheet": schema["sheet_name"],
                "orientation": schema["orientation"],
                **field,
            }
            item.update(_type_parts(field["raw_type"]))
            physical_fields.append(item)
    bean_fields = [
        {"bean": bean["full_name"], **field}
        for bean in beans
        for field in bean["fields"]
    ]
    enum_items = [
        {"enum": enum["full_name"], **item}
        for enum in enums
        for item in enum["items"]
    ]
    formula_fields = [
        {
            "logical_tables": field["logical_tables"],
            "source_path": field["source_path"],
            "schema_sheet": field["schema_sheet"],
            "field_name": field["field_name"],
            "raw_type": field["raw_type"],
            "is_exported": field["is_exported"],
            "is_comment_field": field["is_comment_field"],
            "axis": field["axis"],
            "axis_index": field["axis_index"],
            "header_cell": field["header_cell"],
            "formula_cell_count": field["formula_cell_count"],
            "constant_cell_count": field["constant_cell_count"],
            "external_formula_cell_count": field["external_formula_cell_count"],
            "formula_constant_mixed": field["formula_constant_mixed"],
            "formula_examples": field["formula_examples"],
        }
        for field in physical_fields
        if field["formula_cell_count"] > 0
    ]

    physical_exported_fields = sum(
        sum(field["is_exported"] for field in physical["schema"]["fields"])
        for physical in physical_inputs
        if physical["schema"] is not None
    )
    physical_comment_fields = sum(
        sum(field["is_comment_field"] for field in physical["schema"]["fields"])
        for physical in physical_inputs
        if physical["schema"] is not None
    )
    schema_formula_count = sum(
        physical["schema_formula_count"] for physical in physical_inputs
    )
    auxiliary_formula_count = sum(
        physical["auxiliary_formula_count"] for physical in physical_inputs
    )
    exported_formula_count = sum(
        field["formula_cell_count"]
        for physical in physical_inputs
        if physical["schema"] is not None
        for field in physical["schema"]["fields"]
        if field["is_exported"]
    )
    comment_formula_count = sum(
        field["formula_cell_count"]
        for physical in physical_inputs
        if physical["schema"] is not None
        for field in physical["schema"]["fields"]
        if field["is_comment_field"]
    )
    snapshot_paths = set(required.values()) | registered_paths
    source_snapshot = _snapshot_hash(root, snapshot_paths)
    snapshot_file_count = len(source_snapshot["files"])
    anomaly_counts = dict(sorted(Counter(item["code"] for item in anomalies).items()))
    empty_registered_inputs = sum(
        physical["schema"] is not None
        and physical["schema"]["data_row_count"] == 0
        for physical in physical_inputs
    )
    counts = {
        "spreadsheet_files": len(spreadsheets),
        "xlsx_files": sum(path.suffix.casefold() == ".xlsx" for path in spreadsheets),
        "xlsm_files": sum(path.suffix.casefold() == ".xlsm" for path in spreadsheets),
        "lock_files": len(lock_files),
        "schema_workbooks": len(required),
        "logical_tables": len(tables),
        "registered_inputs": len(input_references),
        "registered_input_references": len(input_references),
        "registered_unique_inputs": len(registered_paths),
        "parsed_registered_inputs": len(physical_inputs),
        "registered_worksheets": sum(
            physical["sheet_count"] for physical in physical_inputs
        ),
        "schema_worksheets": sum(
            physical["schema_sheet_count"] for physical in physical_inputs
        ),
        "auxiliary_worksheets": sum(
            physical["sheet_count"] - physical["schema_sheet_count"]
            for physical in physical_inputs
        ),
        "multi_sheet_registered_inputs": sum(
            physical["sheet_count"] > 1 for physical in physical_inputs
        ),
        "horizontal_schema_inputs": sum(
            physical["schema"] is not None
            and physical["schema"]["orientation"] == "horizontal"
            for physical in physical_inputs
        ),
        "vertical_schema_inputs": sum(
            physical["schema"] is not None
            and physical["schema"]["orientation"] == "vertical"
            for physical in physical_inputs
        ),
        "physical_exported_fields": physical_exported_fields,
        "logical_exported_fields": len(logical_fields),
        "physical_comment_fields": physical_comment_fields,
        "beans": len(beans),
        "bean_fields": sum(len(bean["fields"]) for bean in beans),
        "enums": len(enums),
        "enum_items": sum(len(enum["items"]) for enum in enums),
        "registered_formula_cells": schema_formula_count + auxiliary_formula_count,
        "schema_formula_cells": schema_formula_count,
        "auxiliary_formula_cells": auxiliary_formula_count,
        "exported_field_formula_cells": exported_formula_count,
        "comment_field_formula_cells": comment_formula_count,
        "external_formula_cells": sum(
            physical["external_formula_count"] for physical in physical_inputs
        ),
        "unregistered_workbooks": len(unregistered_workbooks),
        "empty_registered_inputs": empty_registered_inputs,
        "anomalies": len(anomalies),
        "snapshot_files": snapshot_file_count,
    }
    return {
        "schema_version": SCHEMA_VERSION,
        "generated_at": datetime.now().astimezone().isoformat(timespec="seconds"),
        "config_root": str(root),
        "counts": counts,
        "anomaly_counts": anomaly_counts,
        "anomalies": anomalies,
        "source_snapshot_hash": source_snapshot["sha256"],
        "source_snapshot": source_snapshot,
        "logical_tables": tables,
        "table_inputs": input_references,
        "physical_inputs": physical_inputs,
        "logical_fields": logical_fields,
        "physical_fields": physical_fields,
        "beans": beans,
        "bean_fields": bean_fields,
        "enums": enums,
        "enum_items": enum_items,
        "formula_fields": formula_fields,
        "unregistered_workbooks": unregistered_workbooks,
        "lock_files": [path.relative_to(root).as_posix() for path in lock_files],
    }


__all__ = ["extract_config_schema"]
