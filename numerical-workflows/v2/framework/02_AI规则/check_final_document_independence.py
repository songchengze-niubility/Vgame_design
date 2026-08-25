# -*- coding: utf-8 -*-
"""检查 04_数值文档 是否可脱离参数工作簿独立打开和修改。"""

from __future__ import annotations

import argparse
import re
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook


V2_DIR = Path(__file__).resolve().parent.parent
DEFAULT_ROOT = V2_DIR / "04_数值文档"
PARAMETER_MARKERS = ("参数.xlsx", "共享参数.xlsx", "参数总表.xlsx")
EXCEL_ERRORS = {"#REF!", "#DIV/0!", "#VALUE!", "#N/A", "#NAME?", "#NUM!", "#NULL!"}
SHEET_REF = re.compile(r"(?:'((?:[^']|'')+)'|([A-Za-z0-9_\u3400-\u9fff]+))!")
EXTERNAL_REF = re.compile(r"\[[^\]]+\.(?:xlsx|xlsm|xlsb|xls)\]", re.IGNORECASE)


def formula_issues(formula: str, sheetnames: set[str]) -> list[str]:
    issues = []
    if EXTERNAL_REF.search(formula) or any(marker in formula for marker in PARAMETER_MARKERS):
        issues.append("外部工作簿或参数表引用")
    if any(error in formula.upper() for error in EXCEL_ERRORS):
        issues.append("公式内含 Excel 错误")
    if "×" in formula or "÷" in formula:
        issues.append("疑似把人读说明误写为公式")
    for match in SHEET_REF.finditer(formula):
        target = (match.group(1) or match.group(2)).replace("''", "'")
        if target not in sheetnames:
            issues.append("引用不存在页签: %s" % target)
    return sorted(set(issues))


def inspect(path: Path) -> tuple[int, list[str]]:
    problems = []
    try:
        with ZipFile(path) as archive:
            names = archive.namelist()
    except BadZipFile:
        return 0, ["文件不是有效 xlsx"]
    if any(name.startswith("xl/externalLinks/") for name in names):
        problems.append("包含 externalLinks")
    if "xl/connections.xml" in names or any(name.startswith("xl/queryTables/") for name in names):
        problems.append("包含数据连接或 QueryTable")

    wb = load_workbook(path, data_only=False, read_only=False, keep_links=True)
    values = load_workbook(path, data_only=True, read_only=False, keep_links=True)
    sheetnames = set(wb.sheetnames)
    formulas = 0
    for ws in wb.worksheets:
        value_ws = values[ws.title]
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                if cell.data_type == "f" or (isinstance(value, str) and value.startswith("=")):
                    formulas += 1
                    for issue in formula_issues(str(value), sheetnames):
                        problems.append("%s!%s %s: %s" % (ws.title, cell.coordinate, issue, value))
                cached = value_ws[cell.coordinate].value
                if cached in EXCEL_ERRORS:
                    problems.append("%s!%s 缓存错误: %s" % (ws.title, cell.coordinate, cached))
    wb.close()
    values.close()
    return formulas, sorted(set(problems))


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("root", nargs="?", default=str(DEFAULT_ROOT))
    args = parser.parse_args()
    root = Path(args.root)
    files = sorted(root.rglob("*.xlsx"))
    formula_files = total_formulas = 0
    failed = []
    for path in files:
        formulas, problems = inspect(path)
        if formulas:
            formula_files += 1
            total_formulas += formulas
        if problems:
            failed.append((path.relative_to(root), problems))
    print("最终文档: %d；含内部公式: %d；公式单元格: %d；问题文件: %d" % (
        len(files), formula_files, total_formulas, len(failed)
    ))
    for path, problems in failed:
        print("[FAIL] %s" % path)
        for problem in problems:
            print("  - %s" % problem)
    if not failed:
        print("[OK] 未发现外链、数据连接、参数表引用、失效页签引用或明显公式错误。")
    return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(main())

