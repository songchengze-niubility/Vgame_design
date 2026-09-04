"""Read-only extraction of the three approved first-clear packages; never writes XLSX."""
import hashlib
import json
from pathlib import Path
from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent.parent
FORMAL = Path('D:/Vgame/Config/GameConfig/Datas')
CANDIDATE = ROOT / '06_最终配置/映射装备'
TARGETS = {20000036: 4334, 20000038: 4345, 20000039: 4350}
LAYERS = {27: 4399, 29: 4401, 30: 4402}


def rows(path, sheet):
    with path.open('rb') as stream:
        wb = load_workbook(stream, read_only=True, data_only=True)
        result = list(wb[sheet].iter_rows(values_only=True))
        wb.close()
    return result


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


drop = rows(FORMAL / 'Drop/Drop_dungeon.xlsx', 'Drop')
packages = {}
for target, box in TARGETS.items():
    rs = [r for r in drop[4:] if r[0] != '##' and r[2] == target]
    assert all(r[5] in (None, 0) and r[6] == 'fixed_' for r in rs)
    assert sum(r[10] == box and r[11] == r[12] == 1 for r in rs) == 1
    packages[target] = [list(r[:15]) for r in rs]
ui = rows(FORMAL / 'level/UIlevel.xlsx', 'level')
previews = {r[1]: r[17] for r in ui[4:] if r[0] != '##' and (r[1] - 250000) // 10 in LAYERS}
assert len(previews) == 15
items = rows(FORMAL / 'items/Items.xlsx', 'items')
item_rows = []
for r in items[4:]:
    if r[0] != '##' and r[1] in LAYERS.values():
        item_rows.append([r[1], r[2], r[3], r[4], r[5], r[6], r[7], r[8],
                          {4: 'Orange', 5: 'Red'}[r[9]], r[10], r[11], r[12],
                          r[13], r[14], r[15], 'true', 'CHG-020：正式已有UI占位，仅补齐候选'])
assert len(item_rows) == 3
formal_paths = [FORMAL / f for f in ['Drop/Drop_dungeon.xlsx', 'level/UIlevel.xlsx', 'items/Items.xlsx']]
source = {'revision': 'CHG-020', 'packages': packages, 'previews': previews, 'items': item_rows,
          'formalHashes': {str(p): digest(p) for p in formal_paths},
          'candidateHashes': {p.name: digest(p) for p in CANDIDATE.glob('*.xlsx')}}
(HERE / 'source.json').write_text(json.dumps(source, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps({'packages': len(packages), 'previews': len(previews), 'items': len(item_rows)}))
