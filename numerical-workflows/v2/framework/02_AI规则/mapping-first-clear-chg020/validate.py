"""Validate staged/published CHG-020 against its immutable CHG-019 backup. No XLSX writes."""
import argparse
import hashlib
import json
import subprocess
import sys
from collections import defaultdict
from pathlib import Path
from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent.parent
parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument('--published', action='store_true')
args = parser.parse_args()
metadata = ROOT/'05_验算核算/mapping-equip-first-clear-CHG020.json' if args.published else HERE/'repair.json'
repair = json.loads(metadata.read_text(encoding='utf-8'))
source = repair['source']
backup = Path(repair['backup'])
out = Path(repair['out'] if args.published else repair['staging'])
checks = []


def check(name, ok, detail=None):
    checks.append({'name': name, 'result': 'PASS' if ok else 'FAIL', 'detail': detail})


def load(directory, file, sheet):
    wb = load_workbook(directory/file, data_only=False)
    rows = [list(r) for r in wb[sheet].iter_rows(values_only=True)]
    wb.close()
    return rows


def trim(row):
    row = list(row)
    while row and row[-1] is None:
        row.pop()
    return row


def grouped(rows, col):
    result = defaultdict(list)
    for r in rows[4:]:
        if len(r) > col and r[col] is not None:
            result[r[col]].append(trim(r))
    return result


def sha(p):
    return hashlib.sha256(p.read_bytes()).hexdigest()


baseline_report = ROOT/'05_验算核算/mapping-equip-CHG019-before-CHG020-validation.json'
baseline = subprocess.run([sys.executable,str(HERE.parent/'mapping-drop-chg019/validate.py'),
                           '--out',str(backup),'--report',str(baseline_report)],capture_output=True,text=True,encoding='utf-8')
check('CHG-019历史基线1697项验证通过', baseline.returncode == 0, baseline.stdout[-500:])

old = grouped(load(backup,'Drop_dungeon.xlsx','Drop'),1)
new = grouped(load(out,'Drop_dungeon.xlsx','Drop'),1)
targets={20000036:4334,20000038:4345,20000039:4350}
retired={22451,22452,22453}
check('仅移除三个旧子池ID',set(new)==set(old)-retired)
check('清理6条记录后有效行数690',sum(map(len,new.values()))==690)
for id,rs in new.items():
    if id not in targets:
        check(f'Drop{id}全部字段保持不变',rs==old[id])
    check(f'Drop{id}二级引用完整',all(not r[4] or r[4] in new for r in rs))
    check(f'Drop{id}数量与权重合法',all(r[6]<=r[7] and r[10]<=r[11] and isinstance(r[12],int) and r[12]>=0 for r in rs))
for id,box in targets.items():
    rs=new[id]
    expected=source['packages'][str(id)]
    # Compare actual reward semantics across the different source/candidate layouts.
    actual=sorted(tuple(r[k] for k in [4,5,6,7,9,10,11,12]) for r in rs)
    formal=sorted(tuple(r[k] for k in [5,6,7,8,10,11,12,13]) for r in expected)
    check(f'{id}与正式工作副本奖励逐字段相同',actual==formal)
    check(f'{id}仅固定箱一件且无子掉落',all(r[4]==0 and r[5]=='fixed_' for r in rs) and sum(r[9]==box and r[10]==r[11]==1 for r in rs)==1)
    old_nonbox=[r for r in old[id] if r[4]==0]
    check(f'{id}钥匙与材料原值不变',[r for r in rs if r[9]!=box]==old_nonbox)

old_ui=grouped(load(backup,'UIlevel.xlsx','level'),1)
new_ui=grouped(load(out,'UIlevel.xlsx','level'),1)
check('150个关卡ID不变',set(old_ui)==set(new_ui) and len(new_ui)==150)
for id,rs in new_ui.items():
    expected=[list(r) for r in old_ui[id]]
    if str(id) in source['previews']:
        expected[0][17]=source['previews'][str(id)]
    check(f'UI{id}仅指定FirstReward变化',rs==expected)

old_items=grouped(load(backup,'Items.xlsx','items'),0)
new_items=grouped(load(out,'Items.xlsx','items'),0)
check('仅补入正式已有的三个展示道具',set(new_items)-set(old_items)=={4399,4401,4402} and set(old_items)<=set(new_items))
for id,rs in old_items.items():
    check(f'既有Item{id}逐字段不变',new_items[id]==rs)
for r in source['items']:
    check(f'新Item{r[0]}与源映射一致',new_items[r[0]]==[trim(r)])
    check(f'新Item{r[0]}不可开箱或发放',r[2] is None and r[4] is None and not any(x[9]==r[0] for rs in new.values() for x in rs))
for id,value in source['previews'].items():
    pairs=[tuple(map(int,x.split(','))) for x in value.split(';')]
    check(f'UI{id}首通占位唯一且数量1',sum(i in {4399,4401,4402} and n==1 for i,n in pairs)==1)
    check(f'UI{id}展示引用存在',all(i in new_items or i in {25,1143} for i,_ in pairs))

for file in repair['files']:
    before=load_workbook(backup/file,data_only=False)
    after=load_workbook(out/file,data_only=False)
    check(file+'页签集合不变',before.sheetnames==after.sheetnames)
    def formulas(w):
        return {(s.title,c.coordinate):c.value for s in w for row in s for c in row if c.data_type=='f'}
    check(file+'所有公式逐格不变',formulas(before)==formulas(after))
    before.close();after.close()
    w=load_workbook(out/file,data_only=True)
    errors=[(s.title,c.coordinate,c.value) for s in w for row in s for c in row if c.data_type=='e']
    check(file+'无公式错误',not errors,errors)
    w.close()
for p,digest in source['formalHashes'].items():
    check('正式文件未修改：'+Path(p).name,sha(Path(p))==digest)
for file,digest in source['candidateHashes'].items():
    if file not in repair['files']:
        check('其他候选未修改：'+file,sha(Path(repair['out'])/file)==digest)

report={'revision':'CHG-020','scope':'仅06候选27/29/30层首通与关联展示；其他首通差异保留',
        'output':str(out),'backup':str(backup),'sourceHashes':source['formalHashes'],
        'files':{file:sha(out/file) for file in repair['files']},
        'summary':{'checks':len(checks),'pass':sum(c['result']=='PASS' for c in checks),'fail':sum(c['result']=='FAIL' for c in checks)},
        'checks':checks}
report_path=ROOT/'05_验算核算/mapping-equip-first-clear-CHG020-validation.json'
report_path.write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(report['summary']))
for c in checks:
    if c['result']=='FAIL':print(json.dumps(c,ensure_ascii=False))
raise SystemExit(bool(report['summary']['fail']))
