#!/usr/bin/env python3
"""只读校验映射装备正式配置粘贴包 P01。"""

from __future__ import annotations

import os

import argparse
import hashlib
import json
import math
import sys
from fractions import Fraction
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook


FRAME_ROOT = Path(os.environ.get("VGAME_NUMERICAL_RUNTIME_ROOT", str(Path(__file__).resolve().parents[1])))
DEFAULT_WORKBOOK = (
    FRAME_ROOT / "04_数值文档" / "配置粘贴包" / "映射装备正式配置粘贴包_P01.xlsx"
)
DEFAULT_CONFIG_ROOT = Path(os.environ.get("VGAME_CONFIG_DATAS", ""))
DEFAULT_OUTPUT = (
    FRAME_ROOT / "05_验算核算" / "mapping-equip-paste-package-P01-validation.json"
)
PACKAGE_ID = "MAPEQUIP-20260710-P01"
WEIGHT_SCALE = 1000
ERROR_VALUES = {
    "#DIV/0!", "#N/A", "#NAME?", "#NULL!", "#NUM!", "#REF!", "#VALUE!",
}
QUALITIES = ["Green", "Blue", "Purple", "Orange", "Red"]
EXPECTED_TARGETS = {
    "SpecialEquipment/Portable/PortableExp.xlsx",
    "SpecialEquipment/SpecialEquipmentSubAttrCoef.xlsx",
    "SpecialEquipment/Portable/PortableLvUp.xlsx",
    "SpecialEquipment/Portable/PortableCompoundProb.xlsx",
    "SpecialEquipment/Portable/PortableCraftProb.xlsx",
    "SpecialEquipment/Portable/PortableEquipCompound.xlsx",
    "SpecialEquipment/SpecialEquipmentQuality.xlsx",
}


def clean(value: Any) -> str:
    return "" if value is None else str(value).strip()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def equal_value(actual: Any, expected: Any) -> bool:
    if expected is None:
        return actual in (None, "")
    if isinstance(expected, float):
        try:
            return math.isclose(float(actual), expected, rel_tol=0, abs_tol=1e-12)
        except (TypeError, ValueError):
            return False
    if isinstance(expected, int) and not isinstance(expected, bool):
        try:
            return float(actual).is_integer() and int(actual) == expected
        except (TypeError, ValueError):
            return False
    return actual == expected


def largest_remainder(items: list[tuple[str, Fraction]]) -> list[tuple[str, int]]:
    exact = [(name, probability * WEIGHT_SCALE) for name, probability in items]
    weights = [value.numerator // value.denominator for _, value in exact]
    remaining = WEIGHT_SCALE - sum(weights)
    order = sorted(
        range(len(exact)),
        key=lambda index: (
            -(exact[index][1] - weights[index]),
            index,
        ),
    )
    for index in order[:remaining]:
        weights[index] += 1
    return [(items[index][0], weights[index]) for index in range(len(items))]


def quality_list(items: list[tuple[str, Fraction]]) -> str:
    weights = largest_remainder(items)
    return ";".join(f"{quality},{weight}" for quality, weight in weights if weight > 0)


def build_portable_exp() -> list[list[Any]]:
    bands = {
        5: [1, 2, 5, 12, 30],
        10: [1, 2, 5, 12, 30],
        15: [2, 4, 10, 24, 60],
        20: [2, 4, 10, 24, 60],
        25: [3, 6, 15, 36, 90],
        30: [3, 6, 15, 36, 90],
        35: [4, 8, 20, 48, 120],
        40: [4, 8, 20, 48, 120],
        45: [5, 10, 25, 60, 150],
        50: [5, 10, 25, 60, 150],
        55: [6, 12, 30, 72, 180],
        60: [6, 12, 30, 72, 180],
    }
    rows: list[list[Any]] = []
    row_id = 1
    for level in range(5, 61, 5):
        for quality, exp in zip(QUALITIES, bands[level], strict=True):
            rows.append([row_id, level, quality, exp])
            row_id += 1
    return rows


def build_sub_attr_coef() -> list[list[Any]]:
    # 保留正式表既有 Id 1~11，仅追加缺失的 10 级为 Id 12，避免既有引用整体漂移。
    return [
        [1, 5, 0.200],
        [2, 15, 0.345],
        [3, 20, 0.418],
        [4, 25, 0.491],
        [5, 30, 0.564],
        [6, 35, 0.636],
        [7, 40, 0.709],
        [8, 45, 0.782],
        [9, 50, 0.855],
        [10, 55, 0.927],
        [11, 60, 1.000],
        [12, 10, 0.273],
    ]


LV_UP_COSTS = [
    10, 20, 30, 40, 50, 60, 70, 100, 120,
    130, 140, 160, 180, 190, 210, 220, 240, 250, 280,
    350, 400, 450, 500, 550, 650, 750, 850, 950, 1050,
    1130, 1190, 1250, 1310, 1370, 1430, 1490, 1550, 1610, 1670,
    2000, 2100, 2200, 2300, 2400, 2500, 2600, 2700, 2800, 2900,
    3800, 4000, 4200, 4400, 4600, 4800, 5000, 5200, 5400, 5600,
]


def build_portable_lv_up() -> list[list[Any]]:
    return [[index, index + 1, cost] for index, cost in enumerate(LV_UP_COSTS, start=1)]


def compound_probabilities(level: int, quality: str) -> list[tuple[str, Fraction]]:
    progress = Fraction(level - 1, 59)
    if quality == "Green":
        high = Fraction(8, 100) + Fraction(17, 100) * progress
        return [("Blue", 1 - high), ("Purple", high)]
    if quality == "Blue":
        high = Fraction(5, 100) + Fraction(15, 100) * progress
        return [("Purple", 1 - high), ("Orange", high)]
    if quality == "Purple":
        orange = Fraction(30, 100) + Fraction(30, 100) * progress
        red = Fraction(1, 100) + Fraction(2, 100) * progress
        return [("Purple", 1 - orange - red), ("Orange", orange), ("Red", red)]
    if quality == "Orange":
        red = Fraction(20, 100) + Fraction(20, 100) * progress
        return [("Orange", 1 - red), ("Red", red)]
    if quality == "Red":
        return [("Red", Fraction(1, 1))]
    raise ValueError(quality)


def build_compound_prob() -> list[list[Any]]:
    rows: list[list[Any]] = []
    row_id = 1
    for level in range(1, 61):
        for quality in QUALITIES:
            rows.append([row_id, level, quality, quality_list(compound_probabilities(level, quality))])
            row_id += 1
    return rows


def craft_probabilities(level: int, quality: str) -> list[tuple[str, Fraction]]:
    progress = Fraction(level - 1, 59)
    if quality == "Purple":
        high = Fraction(10, 100) + Fraction(10, 100) * progress
        return [("Purple", 1 - high), ("Orange", high)]
    if quality == "Orange":
        high = Fraction(2, 100) + Fraction(3, 100) * progress
        return [("Orange", 1 - high), ("Red", high)]
    if quality == "Red":
        return [("Red", Fraction(1, 1))]
    raise ValueError(quality)


def build_craft_prob() -> list[list[Any]]:
    rows: list[list[Any]] = []
    row_id = 1
    for level in range(1, 61):
        for quality in ("Purple", "Orange", "Red"):
            rows.append([row_id, level, quality, quality_list(craft_probabilities(level, quality))])
            row_id += 1
    return rows


def build_compound_cost() -> list[list[Any]]:
    matrix = [
        [100, 140, 200, 280, 400],
        [350, 490, 700, 980, 1400],
        [1000, 1400, 2000, 2800, 4000],
        [2500, 3500, 5000, 7000, 10000],
        [3800, 5320, 7600, 10640, 15200],
        [5500, 7700, 11000, 15400, 22000],
        [7500, 10500, 15000, 21000, 30000],
        [10000, 14000, 20000, 28000, 40000],
    ]
    rows: list[list[Any]] = []
    row_id = 1
    for costs in matrix:
        for cost in costs:
            rows.append([row_id, cost])
            row_id += 1
    return rows


def build_quality_patch() -> list[list[Any]]:
    main = [0.20, 0.30, 0.45, 0.70, 1.00]
    sub = [0.30, 0.45, 0.60, 0.80, 1.00]
    extra = [None, None, None, None, "23,1"]
    return [
        [index, quality, extra[index - 1], main[index - 1], sub[index - 1]]
        for index, quality in enumerate(QUALITIES, start=1)
    ]


EXPECTED_SHEETS = {
    "01_PortableExp": {
        "headers": ["Id", "EquipLv", "Quality", "Exp"],
        "start": 5,
        "end": 64,
        "rows": build_portable_exp(),
    },
    "02_SubAttrCoef": {
        "headers": ["Id", "EquipLv", "SubAttrCoef"],
        "start": 5,
        "end": 16,
        "rows": build_sub_attr_coef(),
    },
    "03_PortableLvUp": {
        "headers": ["Id", "TargetLv", "CostExp"],
        "start": 5,
        "end": 63,
        "rows": build_portable_lv_up(),
    },
    "04_CompoundProb": {
        "headers": ["Id", "Grade", "Quality", "QualityList"],
        "start": 5,
        "end": 304,
        "rows": build_compound_prob(),
    },
    "05_CraftProb": {
        "headers": ["Id", "PortableLv", "Quality", "QualityList"],
        "start": 5,
        "end": 184,
        "rows": build_craft_prob(),
    },
    "10_CompoundCost": {
        "headers": ["Id", "CostCurrency"],
        "start": 5,
        "end": 44,
        "rows": build_compound_cost(),
    },
    "11_QualityPatch": {
        "headers": ["Id", "Quality", "ExtraItemList", "GradeCoefficient", "QualityCoefficient"],
        "start": 5,
        "end": 9,
        "rows": build_quality_patch(),
    },
}


def parse_quality_list(value: Any) -> list[tuple[str, int]]:
    raw = clean(value)
    if not raw:
        raise ValueError("空 QualityList")
    result: list[tuple[str, int]] = []
    for item in raw.split(";"):
        fields = [part.strip() for part in item.split(",")]
        if len(fields) != 2 or not fields[0]:
            raise ValueError(f"非法 QualityList 项: {item!r}")
        try:
            weight = int(fields[1])
        except ValueError as exc:
            raise ValueError(f"非法权重: {item!r}") from exc
        result.append((fields[0], weight))
    return result


def allowed_outputs(source: str, craft: bool) -> set[str]:
    if craft:
        return {
            "Purple": {"Purple", "Orange"},
            "Orange": {"Orange", "Red"},
            "Red": {"Red"},
        }[source]
    return {
        "Green": {"Blue", "Purple"},
        "Blue": {"Purple", "Orange"},
        "Purple": {"Purple", "Orange", "Red"},
        "Orange": {"Orange", "Red"},
        "Red": {"Red"},
    }[source]


def high_quality(source: str, craft: bool) -> str:
    if craft:
        return {"Purple": "Orange", "Orange": "Red", "Red": "Red"}[source]
    return {
        "Green": "Purple",
        "Blue": "Orange",
        "Purple": "Red",
        "Orange": "Red",
        "Red": "Red",
    }[source]


def check_rows(workbook, errors: list[str]) -> dict[str, int]:
    row_counts: dict[str, int] = {}
    for sheet_name, spec in EXPECTED_SHEETS.items():
        worksheet = workbook[sheet_name]
        headers = [clean(worksheet.cell(4, column).value) for column in range(1, len(spec["headers"]) + 1)]
        if headers != spec["headers"]:
            errors.append(f"{sheet_name} 第4行字段顺序漂移: {headers} != {spec['headers']}")
        expected_rows = spec["rows"]
        if len(expected_rows) != spec["end"] - spec["start"] + 1:
            errors.append(f"校验器内部行数定义错误: {sheet_name}")
        actual_ids: list[Any] = []
        for offset, expected in enumerate(expected_rows):
            row_number = spec["start"] + offset
            actual = [worksheet.cell(row_number, column).value for column in range(1, len(expected) + 1)]
            actual_ids.append(actual[0])
            for column, (actual_value, expected_value) in enumerate(zip(actual, expected, strict=True), start=1):
                if not equal_value(actual_value, expected_value):
                    errors.append(
                        f"{sheet_name}!{worksheet.cell(row_number, column).coordinate} 数据不符: "
                        f"actual={actual_value!r}, expected={expected_value!r}"
                    )
                    break
                cell = worksheet.cell(row_number, column)
                if cell.data_type == "f" or (isinstance(cell.value, str) and cell.value.startswith("=")):
                    errors.append(f"复制区禁止公式: {sheet_name}!{cell.coordinate}")
        normalized_ids = [clean(value) for value in actual_ids]
        if any(not value for value in normalized_ids):
            errors.append(f"{sheet_name} 存在空 Id")
        if len(set(normalized_ids)) != len(normalized_ids):
            errors.append(f"{sheet_name} Id 不唯一")
        for row_number in range(spec["end"] + 1, worksheet.max_row + 1):
            if any(
                worksheet.cell(row_number, column).value not in (None, "")
                for column in range(1, len(spec["headers"]) + 1)
            ):
                errors.append(f"{sheet_name} 复制列在目标数据区之后仍有数据: 第{row_number}行")
                break
        row_counts[sheet_name] = len(expected_rows)
    return row_counts


def check_probabilities(workbook, errors: list[str]) -> dict[str, Any]:
    totals: list[int] = []
    high_weights: dict[tuple[str, str], list[tuple[int, int]]] = {}
    for sheet_name, craft, level_col in [
        ("04_CompoundProb", False, 2),
        ("05_CraftProb", True, 2),
    ]:
        spec = EXPECTED_SHEETS[sheet_name]
        worksheet = workbook[sheet_name]
        for row_number in range(spec["start"], spec["end"] + 1):
            level = int(worksheet.cell(row_number, level_col).value)
            source = clean(worksheet.cell(row_number, 3).value)
            coordinate = f"{sheet_name}!D{row_number}"
            try:
                items = parse_quality_list(worksheet.cell(row_number, 4).value)
            except ValueError as exc:
                errors.append(f"{coordinate} {exc}")
                continue
            names = [name for name, _ in items]
            weights = [weight for _, weight in items]
            if len(set(names)) != len(names):
                errors.append(f"{coordinate} 输出品质重复: {names}")
            if any(weight <= 0 for weight in weights):
                errors.append(f"{coordinate} 权重必须为正整数: {weights}")
            if set(names) != allowed_outputs(source, craft):
                errors.append(
                    f"{coordinate} 输出品质池非法: actual={set(names)}, "
                    f"expected={allowed_outputs(source, craft)}"
                )
            total = sum(weights)
            totals.append(total)
            if total != WEIGHT_SCALE:
                errors.append(f"{coordinate} 权重和 {total} != {WEIGHT_SCALE}")
            weight_map = dict(items)
            high = high_quality(source, craft)
            high_weights.setdefault((sheet_name, source), []).append((level, weight_map.get(high, 0)))
    for key, values in high_weights.items():
        values.sort()
        for previous, current in zip(values, values[1:]):
            if current[1] < previous[1]:
                errors.append(
                    f"{key[0]} {key[1]} 最高品质权重非单调: "
                    f"Lv{previous[0]}={previous[1]}, Lv{current[0]}={current[1]}"
                )
                break
    return {
        "probability_row_count": len(totals),
        "probability_weight_min": min(totals) if totals else None,
        "probability_weight_max": max(totals) if totals else None,
        "monotonic_series_count": len(high_weights),
    }


def read_metadata(worksheet) -> dict[str, Any]:
    return {
        clean(worksheet.cell(row, 1).value): worksheet.cell(row, 2).value
        for row in range(1, worksheet.max_row + 1)
        if clean(worksheet.cell(row, 1).value)
    }


def false_value(value: Any) -> bool:
    return value is False or value == 0 or clean(value).lower() in {"false", "0", "no"}


def normalized_relative_path(value: Any, config_root: Path) -> str:
    raw = clean(value).replace("\\", "/")
    root = str(config_root.resolve()).replace("\\", "/").rstrip("/")
    if raw.lower().startswith(root.lower() + "/"):
        raw = raw[len(root) + 1:]
    return raw.lstrip("/")


def check_target_hashes(
    metadata: dict[str, Any], config_root: Path, errors: list[str]
) -> tuple[dict[str, dict[str, Any]], dict[str, str]]:
    raw = metadata.get("official_target_hashes_json")
    if not clean(raw):
        errors.append("__meta__ 缺少 official_target_hashes_json")
        return {}, {}
    try:
        decoded = json.loads(str(raw)) if isinstance(raw, str) else raw
    except json.JSONDecodeError as exc:
        errors.append(f"official_target_hashes_json 不是合法 JSON: {exc}")
        return {}, {}
    if not isinstance(decoded, list):
        errors.append("official_target_hashes_json 必须是数组")
        return {}, {}
    recorded: dict[str, dict[str, Any]] = {}
    for item in decoded:
        if not isinstance(item, dict):
            errors.append(f"正式目标哈希记录不是对象: {item!r}")
            continue
        relative = normalized_relative_path(item.get("relative_path"), config_root)
        if not relative:
            errors.append(f"正式目标哈希记录缺 relative_path: {item!r}")
            continue
        if relative in recorded:
            errors.append(f"正式目标哈希记录重复: {relative}")
        recorded[relative] = item
    missing = sorted(EXPECTED_TARGETS - set(recorded))
    if missing:
        errors.append(f"__meta__ 缺少正式目标文件哈希: {missing}")
    live_hashes: dict[str, str] = {}
    for relative in sorted(EXPECTED_TARGETS):
        path = config_root / Path(relative)
        if not path.is_file():
            errors.append(f"正式目标文件不存在: {relative}")
            continue
        live_hash = sha256_file(path)
        live_hashes[relative] = live_hash
        item = recorded.get(relative)
        if item is None:
            continue
        expected_hash = clean(item.get("sha256")).lower()
        if len(expected_hash) != 64:
            errors.append(f"正式目标文件哈希记录非法: {relative}={expected_hash!r}")
        elif live_hash.lower() != expected_hash:
            errors.append(
                f"正式目标文件已变化: {relative} live={live_hash}, recorded={expected_hash}"
            )
        if item.get("bytes") not in (None, ""):
            try:
                expected_bytes = int(item["bytes"])
            except (TypeError, ValueError):
                errors.append(f"正式目标文件 bytes 记录非法: {relative}={item.get('bytes')!r}")
            else:
                if path.stat().st_size != expected_bytes:
                    errors.append(
                        f"正式目标文件大小已变化: {relative} "
                        f"live={path.stat().st_size}, recorded={expected_bytes}"
                    )
    return recorded, live_hashes


def iter_cells(workbook) -> Iterable[tuple[str, Any]]:
    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows():
            for cell in row:
                yield worksheet.title, cell


def check_formula_health(workbook, cached, errors: list[str]) -> dict[str, int]:
    formula_count = 0
    explicit_errors = 0
    cached_errors = 0
    for sheet_name, cell in iter_cells(workbook):
        value = cell.value
        if cell.data_type == "f" or (isinstance(value, str) and value.startswith("=")):
            formula_count += 1
        if cell.data_type == "e" or clean(value).upper() in ERROR_VALUES:
            explicit_errors += 1
            if explicit_errors <= 20:
                errors.append(f"Excel 显式错误: {sheet_name}!{cell.coordinate}={value}")
    for sheet_name, cell in iter_cells(cached):
        value = cell.value
        if cell.data_type == "e" or clean(value).upper() in ERROR_VALUES:
            cached_errors += 1
            if cached_errors <= 20:
                errors.append(f"Excel 缓存错误: {sheet_name}!{cell.coordinate}={value}")
    return {
        "formula_count": formula_count,
        "explicit_formula_error_count": explicit_errors,
        "cached_formula_error_count": cached_errors,
    }


def check_audit_sheet(workbook, cached, errors: list[str]) -> dict[str, int]:
    formula_sheet = workbook["99_验算"]
    cached_sheet = cached["99_验算"]
    status_coordinates: list[str] = []
    for row in formula_sheet.iter_rows():
        for cell in row:
            formula = clean(cell.value).upper()
            if (
                (cell.data_type == "f" or formula.startswith("="))
                and '"PASS"' in formula
                and '"FAIL"' in formula
            ):
                status_coordinates.append(cell.coordinate)
    if not status_coordinates:
        errors.append("99_验算 未找到 PASS/FAIL 状态公式")
    pass_count = 0
    for coordinate in status_coordinates:
        value = clean(cached_sheet[coordinate].value).upper()
        if value == "PASS":
            pass_count += 1
        else:
            errors.append(f"99_验算 未通过: {coordinate}={value or '<EMPTY>'}")
    literal_fail_count = sum(
        1
        for row in cached_sheet.iter_rows()
        for cell in row
        if clean(cell.value).upper() == "FAIL"
    )
    if literal_fail_count:
        errors.append(f"99_验算 存在 {literal_fail_count} 个 FAIL")
    return {
        "audit_status_formula_count": len(status_coordinates),
        "audit_pass_count": pass_count,
        "audit_fail_count": literal_fail_count,
    }


def validate(args: argparse.Namespace) -> tuple[list[str], dict[str, Any]]:
    errors: list[str] = []
    workbook_path = Path(args.workbook).resolve()
    config_root = Path(args.config_root).resolve()
    output_path = Path(args.output).resolve()
    summary: dict[str, Any] = {
        "package_id": PACKAGE_ID,
        "workbook": str(workbook_path),
        "config_root": str(config_root),
        "official_config_write": False,
        "weight_scale": WEIGHT_SCALE,
    }
    workbook = None
    cached = None
    try:
        if not workbook_path.is_file():
            errors.append(f"粘贴包不存在: {workbook_path}")
            return errors, summary
        if not config_root.is_dir():
            errors.append(f"正式 Datas 不存在: {config_root}")
            return errors, summary

        workbook = load_workbook(workbook_path, data_only=False, keep_links=True)
        cached = load_workbook(workbook_path, data_only=True, keep_links=False)
        required = {"00_说明", "90_阻断清单", "99_验算", "__meta__", *EXPECTED_SHEETS}
        missing = sorted(required - set(workbook.sheetnames))
        if missing:
            errors.append(f"粘贴包缺 Sheet: {missing}")
            return errors, summary

        meta_sheet = workbook["__meta__"]
        metadata = read_metadata(meta_sheet)
        if clean(metadata.get("package_id")) != PACKAGE_ID:
            errors.append(
                f"package_id 不匹配: {metadata.get('package_id')!r} != {PACKAGE_ID!r}"
            )
        if not false_value(metadata.get("official_config_write")):
            errors.append("__meta__.official_config_write 必须为 False")
        try:
            meta_weight_scale = int(metadata.get("weight_scale"))
        except (TypeError, ValueError):
            meta_weight_scale = -1
        if meta_weight_scale != WEIGHT_SCALE:
            errors.append(
                f"__meta__.weight_scale 必须为 {WEIGHT_SCALE}: {metadata.get('weight_scale')!r}"
            )
        if meta_sheet.sheet_state != "veryHidden":
            errors.append("__meta__ 必须为 veryHidden")
        meta_root = clean(metadata.get("official_config_root"))
        if not meta_root:
            errors.append("__meta__ 缺少 official_config_root")
        elif Path(meta_root).resolve() != config_root:
            errors.append(
                f"__meta__.official_config_root 不匹配: {Path(meta_root).resolve()} != {config_root}"
            )

        row_counts = check_rows(workbook, errors)
        probability_summary = check_probabilities(workbook, errors)
        recorded_hashes, live_hashes = check_target_hashes(metadata, config_root, errors)
        formula_summary = check_formula_health(workbook, cached, errors)
        audit_summary = check_audit_sheet(workbook, cached, errors)
        external_links = len(getattr(workbook, "_external_links", []))
        if external_links:
            errors.append(f"粘贴包存在外链: {external_links}")

        summary.update(
            {
                "workbook_sha256": sha256_file(workbook_path),
                "sheet_row_counts": row_counts,
                "paste_row_count": sum(row_counts.values()),
                "direct_full_table_row_count": sum(
                    row_counts[name]
                    for name in ("01_PortableExp", "02_SubAttrCoef", "03_PortableLvUp")
                ),
                "conditional_probability_row_count": sum(
                    row_counts[name] for name in ("04_CompoundProb", "05_CraftProb")
                ),
                "conditional_patch_row_count": sum(
                    row_counts[name] for name in ("10_CompoundCost", "11_QualityPatch")
                ),
                "portable_lv_up_total_exp": sum(LV_UP_COSTS),
                "compound_cost_total": sum(row[1] for row in build_compound_cost()),
                "official_target_hash_record_count": len(recorded_hashes),
                "official_target_live_hashes": live_hashes,
                "external_link_count": external_links,
                **probability_summary,
                **formula_summary,
                **audit_summary,
            }
        )
    finally:
        if workbook is not None:
            workbook.close()
        if cached is not None:
            cached.close()
        summary["errors"] = errors
        summary["validation_status"] = "PASS" if not errors else "FAIL"
        output_path.parent.mkdir(parents=True, exist_ok=True)
        output_path.write_text(
            json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
        )
    return errors, summary


def main() -> int:
    parser = argparse.ArgumentParser(description="校验映射装备正式配置粘贴包 P01")
    parser.add_argument("--workbook", default=str(DEFAULT_WORKBOOK))
    parser.add_argument("--config-root", default=str(DEFAULT_CONFIG_ROOT))
    parser.add_argument("--output", default=str(DEFAULT_OUTPUT))
    args = parser.parse_args()

    errors, summary = validate(args)
    if errors:
        print("[FAIL] 映射装备正式配置粘贴包 P01 校验失败", file=sys.stderr)
        for error in errors:
            print(f"- {error}", file=sys.stderr)
        print(json.dumps(summary, ensure_ascii=False, indent=2))
        return 1
    print("[OK] 映射装备正式配置粘贴包 P01 校验通过；正式目标文件哈希未变化")
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
