#!/usr/bin/env python3
"""将已确认的映射装备/旅券副本 UIlevel 奖励预览安全写入正式配置。"""
from __future__ import annotations

import hashlib
import json
import os
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(os.environ.get("VGAME_NUMERICAL_RUNTIME_ROOT", str(Path(__file__).resolve().parents[1])))
CONFIG = Path(os.environ.get("VGAME_CONFIG_DATAS", ""))
FORMAL = CONFIG / "level" / "UIlevel.xlsx"
SOURCES = {
    "mapping": (ROOT / "06_最终配置" / "映射装备" / "UIlevel.xlsx", 150),
    "travel": (ROOT / "06_最终配置" / "旅券副本" / "UIlevel.xlsx", 40),
}
FIELDS = ("Reward", "FirstReward")
REPORT = ROOT / "05_验算核算" / "uilevel-formal-preview-apply-F01.json"


def sha(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def header_columns(ws, fields: tuple[str, ...]) -> dict[str, int]:
    headers = [ws.cell(1, col).value for col in range(1, ws.max_column + 1)]
    return {field: headers.index(field) + 1 for field in fields}


def formal_item_ids() -> set[int]:
    wb = load_workbook(CONFIG / "items" / "Items.xlsx", read_only=False, data_only=True)
    ws = wb["items"]
    item_col = header_columns(ws, ("Id",))["Id"]
    ids = {cell.value for (row, col), cell in ws._cells.items() if row >= 5 and col == item_col and isinstance(cell.value, int)}
    wb.close()
    return ids


def validate_preview(value, item_ids: set[int], context: str) -> None:
    for part in str(value or "").split(";"):
        if not part:
            continue
        try:
            item_id, quantity = (int(x) for x in part.split(","))
        except ValueError as exc:
            raise ValueError(f"{context}: invalid PropType {part!r}") from exc
        if item_id not in item_ids:
            raise ValueError(f"{context}: ItemId={item_id} does not exist in formal Items.xlsx")
        if quantity < 0 and quantity != -1:
            raise ValueError(f"{context}: no-count preview must use -1, got {quantity}")


def source_previews(item_ids: set[int]) -> tuple[dict[int, dict[str, object]], dict[str, int]]:
    previews: dict[int, dict[str, object]] = {}
    counts: dict[str, int] = {}
    for name, (path, expected_count) in SOURCES.items():
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb["level"]
        cols = header_columns(ws, ("LevelId", *FIELDS))
        count = 0
        for row in range(5, ws.max_row + 1):
            level_id = ws.cell(row, cols["LevelId"]).value
            if level_id is None:
                continue
            if level_id in previews:
                raise ValueError(f"duplicate LevelId across delivery files: {level_id}")
            previews[level_id] = {field: ws.cell(row, cols[field]).value for field in FIELDS}
            for field in FIELDS:
                validate_preview(previews[level_id][field], item_ids, f"{name}[{level_id}].{field}")
            count += 1
        wb.close()
        if count != expected_count:
            raise ValueError(f"{name}: expected {expected_count} rows, got {count}")
        counts[name] = count
    return previews, counts


def main() -> None:
    if os.environ.get("VGAME_ALLOW_FORMAL_CONFIG_WRITE") != "USER_APPROVED":
        raise RuntimeError("formal config write denied: explicit user authorization required")
    backup_value = os.environ.get("UILEVEL_FORMAL_BACKUP")
    if not backup_value:
        raise RuntimeError("UILEVEL_FORMAL_BACKUP is required")
    backup = Path(backup_value)
    if not backup.exists() or sha(backup) != sha(FORMAL):
        raise RuntimeError("formal UIlevel changed after backup; stop and create a fresh backup")

    item_ids = formal_item_ids()
    previews, source_counts = source_previews(item_ids)
    wb = load_workbook(FORMAL, data_only=False)
    ws = wb["level"]
    cols = header_columns(ws, ("LevelId", *FIELDS))
    formal_rows = {cell.value: row for (row, col), cell in ws._cells.items() if row >= 5 and col == cols["LevelId"] and cell.value is not None}
    missing = sorted(set(previews) - set(formal_rows))
    if missing:
        raise ValueError(f"formal UIlevel missing target LevelId: {missing}")

    before_values = {(row, col): cell.value for (row, col), cell in ws._cells.items()}
    changes = []
    intended_cells = set()
    for level_id, values in previews.items():
        row = formal_rows[level_id]
        for field, new_value in values.items():
            cell = ws.cell(row, cols[field])
            if isinstance(cell.value, str) and cell.value.startswith("="):
                raise ValueError(f"refuse to overwrite formula: LevelId={level_id}, field={field}")
            if cell.value != new_value:
                changes.append({"LevelId": level_id, "field": field, "old": cell.value, "new": new_value})
                intended_cells.add((row, cols[field]))
                cell.value = new_value

    wb.save(FORMAL)
    wb.close()

    verify_wb = load_workbook(FORMAL, read_only=False, data_only=False)
    verify_ws = verify_wb["level"]
    after_values = {(row, col): cell.value for (row, col), cell in verify_ws._cells.items()}
    unexpected = [
        {"row": row, "col": col, "old": before_values.get((row, col)), "new": value}
        for (row, col), value in after_values.items()
        if before_values.get((row, col)) != value and (row, col) not in intended_cells
    ]
    removed = [
        {"row": row, "col": col, "old": value}
        for (row, col), value in before_values.items()
        if value is not None and (row, col) not in after_values and (row, col) not in intended_cells
    ]
    verify_cols = header_columns(verify_ws, ("LevelId", *FIELDS))
    verify_rows = {cell.value: row for (row, col), cell in verify_ws._cells.items() if row >= 5 and col == verify_cols["LevelId"] and cell.value is not None}
    mismatches = [
        {"LevelId": level_id, "field": field, "actual": verify_ws.cell(verify_rows[level_id], verify_cols[field]).value, "expected": value}
        for level_id, fields in previews.items()
        for field, value in fields.items()
        if verify_ws.cell(verify_rows[level_id], verify_cols[field]).value != value
    ]
    formula_errors = [
        f"{cell.coordinate}:{cell.value}"
        for cell in verify_ws._cells.values()
        if isinstance(cell.value, str) and cell.value.startswith(("#REF!", "#VALUE!", "#DIV/0!", "#NAME?", "#N/A"))
    ]
    verify_wb.close()
    if unexpected or removed or mismatches or formula_errors:
        raise RuntimeError(json.dumps({"unexpected": unexpected, "removed": removed, "mismatches": mismatches, "formula_errors": formula_errors}, ensure_ascii=False))

    report = {
        "applied_at": datetime.now().astimezone().isoformat(),
        "formal": str(FORMAL),
        "backup": str(backup),
        "backup_sha256": sha(backup),
        "formal_sha256_after": sha(FORMAL),
        "source_rows": source_counts,
        "fields": list(FIELDS),
        "changes": changes,
        "checks": {
            "target_levels": len(previews),
            "changed_cells": len(changes),
            "formal_item_id_closure": True,
            "negative_preview_normalized": True,
            "unrelated_value_changes": 0,
            "formula_errors": 0,
        },
    }
    REPORT.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report["checks"], ensure_ascii=False))


if __name__ == "__main__":
    main()
