"""Create a scoped human-readable landing checklist; no config writes."""
import hashlib,json,subprocess
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
import compare_svn_scope as audit

OUT=audit.OUT
raw=json.loads((OUT/'raw-comparison.json').read_text(encoding='utf-8'))
C={c['section']:c for c in raw['comparisons']}
def digest(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def celltext(x):
    if x is None:return '空'
    return str(x).replace('\n','<br>').replace('|','\\|')
def table(headers,rows):return '\n'.join(['| '+' | '.join(headers)+' |','|'+'|'.join('---' for h in headers)+'|']+['| '+' | '.join(celltext(v) for v in row)+' |' for row in rows])
def package(rs):return '；'.join(f"{r['DropItem_id']}×{r['DropItemNumMin']}" for r in rs) or '无'

for p,h in raw['source_hashes'].items():assert digest(p)==h,p
flat_sections=['章节挂载','章节任务本体','章节任务目标']
actions=[];new_rows=[];disable_rows=[]
for section in flat_sections:
    c=C[section]
    for r in c['changed']:
        for e in r['edits']:
            actions.append({'section':section,'file':c['svn'],'sheet':c['sheet'],'id':r['id'],**e})

for section in ['章节单项奖励','无限之旅钥匙','第二章36星钥匙','映射首通奖励']:
    c=C[section];s=audit.table(audit.SVN/c['svn'],'DropId')
    for change in c['changed']:
        old,new=change['old_only'],change['candidate_only']
        assert len(new)==1,(section,change['drop_id'])
        if not old:
            new_rows.append({'section':section,'file':c['svn'],'sheet':c['sheet'],'drop_id':change['drop_id'],'values':{'DropId':change['drop_id'],**{f:new[0][f] for f in audit.DF}},'candidate_row':new[0]['row']})
            continue
        # Replace one physical row, disable surplus material rows instead of deleting formula-bearing rows.
        for f in audit.DF:
            if audit.norm(old[0][f])!=audit.norm(new[0][f]):
                actions.append({'section':section,'file':c['svn'],'sheet':c['sheet'],'id':change['drop_id'],'field':f,'old':old[0][f],'new':new[0][f],'svn_cell':get_column_letter(s['header'][f]+1)+str(old[0]['row'])})
        for r in old[1:]:disable_rows.append({'section':section,'file':c['svn'],'sheet':c['sheet'],'id':change['drop_id'],'svn_row':r['row'],'svn_cell':f"A{r['row']}",'new':'##','old_item':r['DropItem_id']})

# Item naming, icons, material subtype placeholders and enum encodings are not authorized copy targets.
item=C['映射箱子与预览道具']
special=next(r for r in item['changed'] if r['id']==4375)
specific=next(e for e in special['edits'] if e['field']=='Description')
actions.append({'section':'4375实际箱描述','file':item['svn'],'sheet':item['sheet'],'id':4375,**specific})
assert len({a['file'] for a in actions}|{r['file'] for r in new_rows}|{r['file'] for r in disable_rows})==7
assert len(new_rows)==6 and len(disable_rows)==15

# Independent physical-cell validation catches shifted candidate/formal layouts.
books={};checked=0
for a in actions:
    p=audit.SVN/a['file']
    if p not in books:books[p]=load_workbook(p,read_only=False,data_only=False)
    cell=books[p][a['sheet']][a['svn_cell']]
    assert cell.data_type!='f',(a['file'],a['svn_cell'],'formula: manual review required')
    assert audit.norm(cell.value)==audit.norm(a['old']),(a,cell.value)
    checked+=1
for a in disable_rows:
    assert books[audit.SVN/a['file']][a['sheet']][a['svn_cell']].value!='##'
for w in books.values():w.close()

# Verify the new chapter package does not collide with another logical Drop source.
registered=next(r for r in raw['metadata']['__tables__.xlsx'] if r['values'][1]=='Drop')['values'][4].split(',')
collision=[]
for filename in registered:
    p=audit.SVN/'Drop'/filename
    _,packs=audit.packages(p)
    if 1030205 in packs:collision.append(str(p))
assert not collision,collision
assert raw['task_drop_refs_outside_chapter_scope']=={}

# Check current 06 numerical semantics rather than relying on an old run's PASS label.
from sync_chapter_md import source,parse_table,MD,rewards_text
chapters,tasks,single,complete=source();body=MD.read_text(encoding='utf-8')
details=parse_table(body,'挂载与目标明细');assert len(details)==69
for row,t in zip(details,tasks):
    assert row==list(map(str,[t['chapter'],t['group'],t['task'],t['target'],t['desc'],t['type'],t['param'],t['drop'],rewards_text(t['rewards'])]))
_,first=audit.packages(audit.FINAL/'映射装备/Drop_dungeon.xlsx')
gen=audit.index(audit.table(audit.FINAL/'映射装备/SpecialEquipmentGenerate.xlsx','Id'),'Id')
for did in range(20000010,20000040):
    rs=first[did];assert all(r['DropId2']==0 and r['DropType']=='fixed_' and r['DropItemNumMin']==r['DropItemNumMax']==1 for r in rs)
    assert sum(r['DropItem_id']==25 for r in rs)==1
    assert sum(4201<=r['DropItem_id']<=4350 or r['DropItem_id']==4375 for r in rs)==1
assert {r['DropItem_id'] for r in first[20000012]}=={25,4375}
assert gen[4375]['GradeList']=='10,1000' and gen[4375]['QualityList']=='Orange,1000'
assert gen[4375]['GenerateMin']==gen[4375]['GenerateMax']==1

svn='C:/Program Files/TortoiseSVN/bin/svn.exe'
for command,name in [('info','svn-info.xml'),('status','svn-status.xml')]:
    result=subprocess.run([svn,command,'--xml',str(audit.SVN)],capture_output=True,check=True)
    (OUT/name).write_bytes(result.stdout)

files=['task/Jottings.xlsx','task/Task.xlsx','task/TaskTarget.xlsx','Drop/Drop_other.xlsx','Drop/Drop_dungeon.xlsx','Drop/Drop_mainlevel.xlsx','items/Items.xlsx']
parts=['# 06候选 → SVN工作副本落表差异清单',
 '核对日期：2026-08-27。基线：本机 `D:/Vgame/Config/GameConfig/Datas`，目录修订r40473；未查询远端HEAD。受控配置没有本地未提交修改；仅有一个未纳入版本控制的Items副本，本次未读取或处理。',
 '## 结论与范围',
 '用户确认的章节任务06数值与当前沉淀MD一致。按本次对话范围落SVN，需要修改7张真实配置表；若只落章节任务，则只需前4张。不能把整个06目录直接覆盖SVN：它含分模块候选、说明页、不同的列布局和行ID。',
 table(['SVN相对路径','需要同步的内容'],[
  [files[0],'第1、2章挂载列表，共2行'],[files[1],'11条章节任务文案；其中1315/1327/1328另补DropId'],[files[2],'11条目标记录：文案、6条参数、2条目标类型及2条跳转入口'],[files[3],'8个章节单项奖励包：新增1包，改7包；清掉旧副材的有效发放'],[files[4],'19个映射首通包的箱子ID/数量，外加无限之旅23000004钥匙1→2'],[files[5],'11000012的钥匙行ShowItemID 101→25，实际5把已正确'],[files[6],'仅4375.Description改成固定1件10级金装；开箱规则已正确']]),
 '本次只生成核对清单，未写任何06或SVN工作簿，未提交SVN，未执行Luban或客户端/服务端验收。',
 '## 一、章节挂载',
 table(['章节ID','字段','SVN当前','06目标','SVN单元格'],[[r['id'],e['field'],e['old'],e['new'],e['svn_cell']] for r in C['章节挂载']['changed'] for e in r['edits']]),
 'JottingsTask的56个组与TaskID列表已一致，不新增或重编号。第1章补组24，第2章补组35、36。',
 '## 二、Task本体与TaskTarget目标',
 'TaskId和TargetId均复用现有记录；不要新增重复任务。下表只列有差异的字段，未列出的字段保持SVN原值。日周任务及那9条曾被否决的UnlockId不在修改范围。']
for section in ['章节任务本体','章节任务目标']:
    c=C[section];parts.extend(['### '+c['svn'],table(['ID','字段','SVN当前','06目标','SVN单元格'],[[r['id'],e['field'],e['old'],e['new'],e['svn_cell']] for r in c['changed'] for e in r['edits']])])
parts.extend(['## 三、Drop_other章节奖励',table(['DropId','SVN当前奖励','06目标奖励','处理'],[[r['drop_id'],package(r['old_full']),package(r['new_full']),'新增1行' if r['svn_package_missing'] else '替换奖励，停用多余旧行' if len(r['old_only'])>1 else '修改现有行'] for r in C['章节单项奖励']['changed']]),
 '所有新目标行均为fixed_，DropId2=0，包次数1/1，物品Min/Max一致，权重10000。六条映射任务奖励均为25×3，不保留副材；仅追加钥匙而不处理旧行，会同时发放副材。章节完成奖7个包及日周活跃钥匙3个包已经一致，不重复改。',
 '多余15条副材行可在首列标记`##`停用，避免删除行导致ID公式移位：',table(['DropId','SVN行','原道具','首列目标'],[[r['id'],r['svn_row'],r['old_item'],r['new']] for r in disable_rows]),
 '## 四、Drop_dungeon映射首通与无限之旅',
 '以下19个映射首通包只调整箱子；各包钥匙及材料保持现有值。第一、二层与后续层箱子按06逐包同步；第3层20000012已正确发25×1和4375×1，不再改。',
 table(['层数','DropId','SVN有差异的箱子','06目标箱子','SVN行/操作'],[[r['drop_id']-20000009,r['drop_id'],package(r['old_only']),package(r['candidate_only']),'、'.join(str(x['row']) for x in r['old_only']) or '新增箱子行'] for r in C['映射首通奖励']['changed']]),
 '无限之旅：普通难度4000分对应DropId23000004，Item25在SVN第2735行的L/M列均由1改2；其余19个钥匙节点已各2把，不动。四难度每档5个钥匙节点，因此补齐后每难度10把/周，日周任务31把/周，共41把/周；实际领取仍以四难度互斥规则为前提。',
 '## 五、主线36星与4375说明',
 'Drop_mainlevel：11000012（第二章36星）已经固定发25×5，不要再追加钥匙。仅把该钥匙行O976（ShowItemID）从101改25。',
 f"Items：4375的{specific['svn_cell']}（Description）改为“{specific['new']}”。当前旧描述还写5级或10级随机；实际Generate4375已是10级、Orange、固定1件，生成表无需再改。",
 '## 六、所有修改单元格',
 '下面是现有纯数据格的精确新旧值，已独立复读其SVN地址；未发现目标单元格为公式。行号只适用于本次r40473核对快照，操作前先按ID重新定位。',
 table(['SVN文件','业务ID','字段','单元格','当前值','目标值'],[[a['file'],a['id'],a['field'],a['svn_cell'],a['old'],a['new']] for a in actions]),
 '## 七、需要手工新增的6行',
 '5行补缺失首通箱，1行新增第一章舱室奖励包。不要直接复制候选B列行ID，不要自动插入或删除公式所在行；在正式表尾部准备新行后，按正式表规则维护ID及DropId公式。',
 table(['SVN文件','DropId','道具','数量Min/Max','DropId2','DropType','包次数Min/Max','权重','ShowItemID'],[[r['file'],r['drop_id'],r['values']['DropItem_id'],f"{r['values']['DropItemNumMin']}/{r['values']['DropItemNumMax']}",0,'fixed_','1/1',10000,r['values']['ShowItemID']] for r in new_rows]),
 '## 八、不需要同步或不能整表覆盖的内容',
 '- 玩家等级10/20/25/30/33/36/38已在SVN，其他未列出的章节目标不改。JottingsTask共56组已一致。\n- 每日80活跃3把、每周20/40活跃4/6把、月签5/12/19/26日各2把已正确；DailyActiveReward、WeeklyReward、DailySign无需再改。\n- 映射UIlevel的150行FirstReward/Reward/FirstRewardDrop/FallRewardDrop已一致；116个常规及子掉落包、151条开箱规则也一致，不重刷这些表。\n- NR_HighestScoreReward按“难度+分数”对齐后40个节点均一致。SVN采用1～10、101～110、201～210、301～310；06采用1～40，不能为了追求表面一致改正式ID。\n- Items候选的品质名Green/Blue/Purple/Orange/Red与SVN数值1～5是编码表示差异。候选还有通用素材名称/图标/描述、空ItemSubType等范围外差异，保留SVN现有正式字段；本次Items只列4375.Description。\n- 06中的00_配置清单、中文说明、校验JSON是交付辅助内容，不复制到Datas业务表；Task的候选和SVN列位置并不完全相同，应按字段名贴值。',
 '## 九、落表顺序与验收',
 '1. 备份这7张SVN源表，先按ID准备6条新增Drop行；更新奖励并停用15条旧副材行。\n2. 同步TaskTarget、Task和Jottings，最后改4375说明与主线钥匙预览。\n3. 复查实际挂载7章/56组/69条，第二至七章映射任务各25×3且无副材；第一、二章舱室各3028×2/×3。\n4. 复查30层首通匙1/箱1，第三层4375固定10级金装；普通难度4000分钥匙2把；36星只发5把。\n5. 使用正式工程Luban导出并检查引用；联调任务持有量初始化、进度、章节奖励领取与周奖励互斥。静态相等不能替代这些运行验收。',
 '## 复核证据',
 '原始字段对比见同目录raw-comparison.json；精确修改与新增/停用计划见landing-plan.json；独立单元格核对与只读保护见validation.json。没有执行落表。'])

plan={'baseline_revision':'40473','scope':raw['scope'],'files':files,'cell_edits':actions,'manual_new_rows':new_rows,'disable_rows':disable_rows,'excluded_item_differences':len(item['changed'])-1,'source_hashes':raw['source_hashes']}
(OUT/'landing-plan.json').write_text(json.dumps(plan,ensure_ascii=False,indent=2),encoding='utf-8')
(OUT/'SVN配置修改清单.md').write_text('\n\n'.join(parts)+'\n',encoding='utf-8')
for p,h in raw['source_hashes'].items():assert digest(p)==h,p
for p,h in audit.hashes.items():assert digest(p)==h,p
validation={'status':'PASS_READ_ONLY_SCOPED_COMPARISON','files_to_change':7,'existing_pure_data_cell_edits':len(actions),'manual_new_rows':len(new_rows),'old_material_rows_to_disable':len(disable_rows),'physical_cells_rechecked':checked,'chapter_drop_collision':collision,'chapter_task_refs_outside_scope':raw['task_drop_refs_outside_chapter_scope'],'candidate_chapter_md_match':True,'candidate_mapping_first_clear_30_layers':True,'third_floor_fixed_10_orange':True,'all_read_sources_unchanged':True,'svn_written':False,'remote_head_checked':False,'runtime_verified':False,'additional_source_hashes':audit.hashes}
# JSON keys require serializable paths.
validation['additional_source_hashes']={str(k):v for k,v in audit.hashes.items()}
(OUT/'validation.json').write_text(json.dumps(validation,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({k:v for k,v in validation.items() if k!='additional_source_hashes'},ensure_ascii=False))
