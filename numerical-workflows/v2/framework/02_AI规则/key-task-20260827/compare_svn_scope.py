"""Read-only comparison of approved chapter/key/mapping scope, never writes XLSX."""
import hashlib,json,re,subprocess
from collections import Counter,defaultdict
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

ROOT=Path(__file__).resolve().parents[2]
FINAL=ROOT/'06_最终配置'
SVN=Path('D:/Vgame/Config/GameConfig/Datas')
OUT=ROOT/'05_验算核算/06与SVN落表差异_20260827'
OUT.mkdir(parents=True,exist_ok=True)
cache={};hashes={};comparisons=[]
QUALITY={'Green':1,'Blue':2,'Purple':3,'Orange':4,'Red':5}

def read(path):
    path=path.resolve()
    if path in cache:return cache[path]
    hashes[str(path)]=hashlib.sha256(path.read_bytes()).hexdigest()
    w=load_workbook(path,read_only=True,data_only=True);result=[]
    for s in w:
        # Mapping candidates use a three-line title before their field header.
        iterator=iter(s.values);top=[]
        for _ in range(4):top.append(next(iterator,()))
        hr=0 if top[0] and top[0][0]=='##var' else 3
        header={str(v).split('\n')[-1]:i for i,v in enumerate(top[hr]) if v is not None and not str(v).startswith('##')}
        records=[]
        for ri,r in enumerate(iterator,5):
            if not r or r[0]=='##':continue
            row={k:(r[i] if i<len(r) else None) for k,i in header.items()}
            if not any(v is not None for v in row.values()):continue
            row['_row']=ri;records.append(row)
        result.append({'sheet':s.title,'header':header,'rows':records})
    w.close();cache[path]=result;return result

def table(path,key):
    matches=[s for s in read(path) if key in s['header']]
    assert len(matches)==1,(str(path),key,[s['sheet'] for s in matches])
    return matches[0]

def index(t,key):
    d={}
    for r in t['rows']:
        if r.get(key) is not None:
            assert r[key] not in d,(key,r[key])
            d[r[key]]=r
    return d

def norm(x):
    if x is None or x=='':return None
    if isinstance(x,bool):return str(int(x))
    if isinstance(x,float) and x.is_integer():x=int(x)
    return str(x)

def flat(section,cfile,sfile,key,ids=None,fields=None):
    ct=table(FINAL/cfile,key);st=table(SVN/sfile,key)
    ci=index(ct,key);si=index(st,key)
    keys=list(ci) if ids is None else ids
    fs=fields or [f for f in ct['header'] if f in st['header'] and f!=key and not f.startswith('_')]
    changed=[];missing=[];equal=[]
    for k in keys:
        assert k in ci,(cfile,key,k)
        if k not in si:
            missing.append({'id':k,'candidate_row':ci[k]['_row'],'values':{f:ci[k].get(f) for f in fs}});continue
        edits=[]
        for f in fs:
            a,b=si[k].get(f),ci[k].get(f)
            equivalent=norm(a)==norm(b) or (f=='ItemQuality' and a==QUALITY.get(b))
            if not equivalent:edits.append({'field':f,'old':a,'new':b,'svn_cell':get_column_letter(st['header'][f]+1)+str(si[k]['_row']),'candidate_cell':get_column_letter(ct['header'][f]+1)+str(ci[k]['_row'])})
        if edits:changed.append({'id':k,'svn_row':si[k]['_row'],'candidate_row':ci[k]['_row'],'edits':edits})
        else:equal.append(k)
    item={'section':section,'candidate':cfile,'svn':sfile,'sheet':st['sheet'],'key':key,'fields':fs,'changed':changed,'missing':missing,'equal':equal,'checked':len(keys)}
    comparisons.append(item);return item

def score_nodes(cfile,sfile):
    ct=table(FINAL/cfile,'Id');st=table(SVN/sfile,'Id')
    key=lambda r:(r['Difficult'],r['Score'])
    ci={key(r):r for r in ct['rows'] if r.get('Id') is not None};si={key(r):r for r in st['rows'] if r.get('Id') is not None}
    assert len(ci)==40 and len(si)==40
    changed=[];equal=[];id_differences=[]
    for k,c in ci.items():
        assert k in si
        s=si[k]
        if c['DropId']!=s['DropId']:changed.append({'difficulty':k[0],'score':k[1],'old':s['DropId'],'new':c['DropId']})
        else:equal.append(list(k))
        if c['Id']!=s['Id']:id_differences.append({'difficulty':k[0],'score':k[1],'svn_id':s['Id'],'candidate_id':c['Id']})
    comparisons.append({'section':'无限之旅奖励挂载','candidate':cfile,'svn':sfile,'sheet':st['sheet'],'key':'Difficult+Score','checked':40,'changed':changed,'equal':equal,'id_differences_not_to_copy':id_differences})

DF=['DropId2','DropType','DropMinNum','DropMaxNum','DropItem_id','DropItemNumMin','DropItemNumMax','DropWeight','ShowItemID']
def packages(path):
    t=table(path,'DropId');d=defaultdict(list)
    for r in t['rows']:
        if isinstance(r.get('DropId'),int):d[r['DropId']].append(r)
    return t,d

def drops(section,cfile,sfile,ids,item_filter=None):
    ct,cp=packages(FINAL/cfile);st,sp=packages(SVN/sfile)
    changes=[];equal=[]
    sig=lambda r:tuple(norm(r.get(f)) for f in DF)
    clean=lambda r:{f:r.get(f) for f in DF}|{'row':r['_row']}
    for did in sorted(ids):
        assert did in cp,(cfile,did)
        cr=[r for r in cp[did] if item_filter is None or r['DropItem_id']==item_filter]
        sr=[r for r in sp.get(did,[]) if item_filter is None or r['DropItem_id']==item_filter]
        cs,ss=Counter(map(sig,cr)),Counter(map(sig,sr))
        if cs==ss:equal.append(did);continue
        removed=[];added=[];extra=ss-cs
        for r in sr:
            if extra[sig(r)]:removed.append(clean(r));extra[sig(r)]-=1
        extra=cs-ss
        for r in cr:
            if extra[sig(r)]:added.append(clean(r));extra[sig(r)]-=1
        changes.append({'drop_id':did,'svn_package_missing':did not in sp,'old_only':removed,'candidate_only':added,'old_full':[clean(r) for r in sr],'new_full':[clean(r) for r in cr]})
    item={'section':section,'candidate':cfile,'svn':sfile,'sheet':st['sheet'],'key':'DropId','item_filter':item_filter,'changed':changes,'equal':equal,'checked':len(ids)}
    comparisons.append(item);return item

def main():
    # The seven-chapter candidate explicitly defines the scope. No unused old task or daily unlock row is copied.
    j=index(table(FINAL/'任务/Jottings.xlsx','ID'),'ID');g=index(table(FINAL/'任务/JottingsTask.xlsx','ID'),'ID')
    tids=[];gids=[]
    for c in j.values():
        for gid in map(int,str(c['JottingsTasks']).split(';')):
            gids.append(gid);tids.extend(map(int,str(g[gid]['TaskID']).split(';')))
    assert len(tids)==len(set(tids))==69 and len(gids)==56
    tc=index(table(FINAL/'任务/Task.xlsx','Id'),'Id')
    targetids=sorted({int(x) for tid in tids for x in str(tc[tid]['TaskTarget']).split(';')})
    taskdrops={tc[tid]['DropId'] for tid in tids};chapterdrops={r['DropID'] for r in j.values()}
    flat('章节挂载','任务/Jottings.xlsx','task/Jottings.xlsx','ID')
    flat('章节任务组','任务/JottingsTask.xlsx','task/JottingsTask.xlsx','ID',gids)
    flat('章节任务本体','任务/Task.xlsx','task/Task.xlsx','Id',tids)
    flat('章节任务目标','任务/TaskTarget.xlsx','task/TaskTarget.xlsx','Id',targetids)
    drops('章节单项奖励','任务/Drop_other.xlsx','Drop/Drop_other.xlsx',taskdrops)
    drops('章节完成奖励','任务/Drop_other.xlsx','Drop/Drop_other.xlsx',chapterdrops)
    drops('日周活跃钥匙','任务/Drop_other.xlsx','Drop/Drop_other.xlsx',{1010204,1020202,1020204},25)
    flat('日活奖励挂载','任务/DailyActiveReward.xlsx','task/DailyActiveReward.xlsx','ID')
    # Locate known registered filenames without assuming directory names.
    wk=list(SVN.rglob('WeeklyReward.xlsx'));assert len(wk)==1
    flat('周活奖励挂载','任务/WeeklyReward.xlsx',str(wk[0].relative_to(SVN)),'ID')
    sign=table(FINAL/'签到/DailySign.xlsx','id');ids=[r['id'] for r in sign['rows'] if str(r.get('day')) in {'5','12','19','26'}]
    flat('每月签到钥匙','签到/DailySign.xlsx','SignInReward/DailySign.xlsx','id',ids,['rewardItems'])
    drops('无限之旅钥匙','无限之旅/Drop_dungeon.xlsx','Drop/Drop_dungeon.xlsx',set(range(23000002,23000041,2)),25)
    nr=list(SVN.rglob('NR_HighestScoreReward.xlsx'));assert len(nr)==1
    score_nodes('无限之旅/NR_HighestScoreReward.xlsx',str(nr[0].relative_to(SVN)))
    drops('第二章36星钥匙','主线高难/Drop_mainlevel.xlsx','Drop/Drop_mainlevel.xlsx',{11000012},25)
    star=index(table(FINAL/'主线高难/ChapterStarReward.xlsx','Id'),'Id')
    starids=[k for k,r in star.items() if r.get('ChapterId')==10002 and r.get('Star')==36 and not r.get('IsDifficult')]
    flat('第二章36星挂载','主线高难/ChapterStarReward.xlsx','level/ChapterStarReward.xlsx','Id',starids)
    # Mapping reward fields only: preserve level costs, progression, combat and enum fields.
    ui=table(FINAL/'映射装备/UIlevel.xlsx','LevelId');uids=[r['LevelId'] for r in ui['rows'] if r.get('LevelType')=='LevelType_SpecialEquipment']
    assert len(uids)==150
    flat('映射首通和常规预览','映射装备/UIlevel.xlsx','level/UIlevel.xlsx','LevelId',uids,['FirstReward','Reward','FirstRewardDrop','FallRewardDrop'])
    _,mapping=packages(FINAL/'映射装备/Drop_dungeon.xlsx')
    first=set(range(20000010,20000040))
    drops('映射首通奖励','映射装备/Drop_dungeon.xlsx','Drop/Drop_dungeon.xlsx',first)
    # CHG019 ordinary rewards and their subtree definitions, scoped to candidate mapping packages only.
    drops('映射常规及子掉落','映射装备/Drop_dungeon.xlsx','Drop/Drop_dungeon.xlsx',set(mapping)-first)
    # Only item IDs referenced by the actual scoped drops or their UI previews.
    im=index(table(FINAL/'映射装备/Items.xlsx','Id'),'Id')
    referenced={r['DropItem_id'] for rr in mapping.values() for r in rr if r.get('DropItem_id')}
    for u in ui['rows']:
        if u.get('LevelId') not in uids:continue
        for f in ['FirstReward','Reward']:
            referenced.update(int(p.split(',')[0]) for p in str(u.get(f) or '').split(';') if p)
    itemids=sorted(referenced&set(im))
    flat('映射箱子与预览道具','映射装备/Items.xlsx','items/Items.xlsx','Id',itemids,['Name','ParaList','ItemType','ItemSubType','Icon','ItemQuality','Description'])
    genids=sorted({int(im[i]['ParaList']) for i in itemids if im[i].get('ItemSubType')=='ItemSubType_SpecialEquip_Box' and im[i].get('ParaList')})
    flat('映射开箱规则','映射装备/SpecialEquipmentGenerate.xlsx','SpecialEquipment/SpecialEquipmentGenerate.xlsx','Id',genids,['GradeList','QualityList','UnidentifiedList','PartList','SpecialEquipmentId','GenerateMax','GenerateMin'])
    # Explicitly inspect formal schema and enum evidence; it is not a copy target.
    metadata={}
    for filename in ['__tables__.xlsx','__enums__.xlsx']:
        p=SVN/filename;hashes[str(p)]=hashlib.sha256(p.read_bytes()).hexdigest()
        w=load_workbook(p,read_only=True,data_only=True)
        wanted=['Jottings','TaskTarget','Drop_dungeon','DailySign','UIlevel','NR_HighestScoreReward'] if filename=='__tables__.xlsx' else ['Target_TTT_RvCabin_LevelUp_Num','Target_TTT_SpecialEquipmentLevelUp_Num']
        found=[]
        for s in w:
            for ri,r in enumerate(s.values,1):
                if any(v and any(q in str(v) for q in wanted) for v in r) or (filename=='__enums__.xlsx' and len(r)>9 and r[7] in QUALITY):found.append({'sheet':s.title,'row':ri,'values':list(r)})
        w.close();metadata[filename]=found
    for name,value in QUALITY.items():
        assert any(r['values'][7]==name and r['values'][9]==value for r in metadata['__enums__.xlsx'])
    shared_refs={str(did):[r['Id'] for r in table(SVN/'task/Task.xlsx','Id')['rows'] if r.get('DropId')==did and r.get('Id') not in tids] for did in taskdrops}
    protected_ok=all(hashlib.sha256(Path(p).read_bytes()).hexdigest()==h for p,h in hashes.items())
    assert protected_ok
    report={'scope':'用户确认章节任务、映射之匙产出、映射常规/首通及预览；只读本机SVN工作副本，未查询远端HEAD','comparisons':comparisons,'source_hashes':hashes,'sources_unchanged':protected_ok,'metadata':metadata,'task_drop_refs_outside_chapter_scope':{k:v for k,v in shared_refs.items() if v}}
    (OUT/'raw-comparison.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    for c in comparisons:print(json.dumps({'section':c['section'],'file':c['svn'],'checked':c['checked'],'changed':len(c['changed']),'missing':len(c.get('missing',[])),'equal':len(c['equal'])},ensure_ascii=False))

if __name__=='__main__':main()
