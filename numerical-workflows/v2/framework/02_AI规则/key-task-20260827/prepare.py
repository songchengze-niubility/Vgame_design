import json, hashlib, shutil
from pathlib import Path
from datetime import datetime
from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
if (ROOT/'05_验算核算/章节任务纠偏_20260827/plan.json').exists():
    raise SystemExit('Historical batch superseded: rejected daily/weekly unlock edits must not be reapplied.')
FINAL = ROOT/'06_最终配置'
RUN = ROOT/'05_验算核算'/'映射之匙任务_20260827'
RUN.mkdir(parents=True, exist_ok=True)
S = json.loads((HERE/'snapshot.json').read_text(encoding='utf-8'))
changes = {}
def put(file,sheet,row,col,value):
    rows=S[file][sheet]
    old=rows[row-1][col-1] if row<=len(rows) and col<=len(rows[row-1]) else None
    if old!=value:
        changes.setdefault(file,[]).append({'sheet':sheet,'row':row,'col':col,'old':old,'new':value})
def replace(file,sheet,old,new):
    for ri,row in enumerate(S[file][sheet],1):
        for ci,v in enumerate(row,1):
            if isinstance(v,str) and old in v: put(file,sheet,ri,ci,v.replace(old,new))

f='任务/Drop_other.xlsx'
for ri,r in enumerate(S[f]['Drop'],1):
    if len(r)>12 and r[10]==25 and r[2] in {1010204,1020202,1020204}:
        qty={1010204:3,1020202:4,1020204:6}[r[2]]
        for c in [12,13]: put(f,'Drop',ri,c,qty)
f='任务/Task.xlsx'
for ri,r in enumerate(S[f]['Task'],1):
    if len(r)<11: continue
    if r[1] in [1108,12051,12052,12053]:
        assert r[10]==4112
        put(f,'Task',ri,11,123)
        put(f,'Task',ri,12,'主线2-8')
    if r[1] in [1109,1110,12061,12062,12063]:
        assert r[10]==4314
        put(f,'Task',ri,11,140)
        put(f,'Task',ri,12,'主线3-10')
f='签到/DailySign.xlsx'
for ri,r in enumerate(S[f]['Sheet1'],1):
    if len(r)>4 and r[4]=='25,1':
        assert int(r[3]) in [5,12,19,26]
        put(f,'Sheet1',ri,5,'25,2')
assert len(changes[f])==48
f='无限之旅/Drop_dungeon.xlsx'
for ri,r in enumerate(S[f]['Drop'],1):
    if len(r)>12 and r[10]==25:
        assert r[11:13]==[1,1]
        for c in [12,13]: put(f,'Drop',ri,c,2)
notes=S[f]['中文说明']
put(f,'中文说明',3,2,'KEY-TASK-20260827 / CHG-010')
put(f,'中文说明',4,2,'PASS_STATIC_WITH_RUNTIME_GATES_CHG010')
put(f,'中文说明',6,2,'06_玩法_无限挑战经济沉淀CHG-010；沿用DropId23000001～23000040')
put(f,'中文说明',8,2,'保留CHG-009的可见DropId字体和B列ID公式；本轮只增加20项钥匙数量，其他奖励不变。')
put(f,'中文说明',len(notes)+1,1,'CHG-010 钥匙')
put(f,'中文说明',len(notes)+1,2,'2000/4000/6000/8000/10000分每档映射之匙×2，单难度10把/周；四难度互斥，不可合计40把。其余奖励及B列ID公式不变。')
f='主线高难/Drop_mainlevel.xlsx'
rows=S[f]['Drop']; appendrow=len(rows)+1
new=[None,None,11000012,'第二章36星追加映射之匙','映射之匙',0,'fixed_',1,1,None,25,5,5,10000,25,None]
for ci,v in enumerate(new,1): put(f,'Drop',appendrow,ci,v)
put(f,'中文说明',len(S[f]['中文说明'])+1,1,'CHG-004 钥匙')
put(f,'中文说明',len(S[f]['中文说明'])+1,2,'第二章36星实际引用11000012，追加25×5，保留原奖励；11000011是28星，不改。候选追加1行；正式表需人工预建行并下拉ID公式。')

f='任务/00_配置清单.xlsx'
replace(f,'任务总览','映射之匙1、','映射之匙3、')
replace(f,'任务总览','映射之匙5','映射之匙10')
put(f,'数值校验',10,1,'常驻任务映射之匙/周')
put(f,'数值校验',10,2,"='钥匙来源核对'!D2+'钥匙来源核对'!D3")
put(f,'数值校验',10,3,31)
put(f,'落地说明',5,3,'按2026-08-26任务MD：竞技场任务Unlock=123（2-8），映射任务Unlock=140（3-10）；其余保留。次日/次周刷新入池需联调。')
for row,v in enumerate([
 ['统计联调','资源副本日周任务1107/12041~12043','现有Target_TTT_Type_LevelClearNum参数0,N；未找到参数0仅统计资源副本的运行证据。需验证资源类型过滤、挑战/通关差别，不能算作语义已通过。'],
 ['注册核对','DailyTask.xlsx','当前__tables__.xlsx无DailyTask输入，可能由TaskType枚举入池；交接时确认任务池消费入口，不擅自注册新表。'],
 ['本次核对','新增任务已存在06','日常10、周常33、章节68、七日70；只补钥匙和9条日周任务开放条件，章节及七日数据保留。'],
 ['本次来源','映射之匙与新增任务落06-2026-08-27','参见03_沉淀MD/经济/映射之匙与任务落06-2026-08-27.md；只更新候选，未导出SVN。'],
],7):
    for c,value in enumerate(v,1): put(f,'落地说明',row,c,value)
put(f,'配置清单',4,3,'保留四套任务本体及引用；按每日每周任务MD修正9条任务UnlockId为123/140')
put(f,'配置清单',13,3,'日活80分钥匙3；周活20/40分钥匙4/6；章节任务奖励保留')

f='签到/00_配置清单.xlsx'
put(f,'配置清单',2,8,'每月5/12/19/26日映射之匙各2把；6/13/20/27日副材依次1152/1153/1154/1155×2；其余不变')
replace(f,'验收结果','25: 48','25: 96')
put(f,'验收结果',6,3,'每月映射之匙8把，四种副材各×2；28/30/31天均不额外追加钥匙')
f='无限之旅/00_配置清单.xlsx'
for ri,r in enumerate(S[f]['周奖励预览'],1):
    if len(r)>10 and r[8]==25: put(f,'周奖励预览',ri,11,2)
for ri in range(2,6): put(f,'周奖励汇总',ri,6,f"=SUMIFS('周奖励预览'!K2:K41,'周奖励预览'!B2:B41,A{ri},'周奖励预览'!I2:I41,25)")
put(f,'配置清单',7,4,'VALUE_READY_CHG010_KEYS_10_PER_WEEK')
put(f,'配置清单',7,6,'40包×3项=120行；每难度5个钥匙节点各×2；ID公式和其他奖励保持。周奖励四难度互斥，程序待联调。')
put(f,'配置清单',9,4,'PASS_STATIC_WITH_RUNTIME_GATES_CHG010')
put(f,'配置清单',9,6,'本轮重新核验钥匙、40包奖励、ID公式和可见字体；历史CHG009不作为当前校验结果。')
f='主线高难/00_配置清单.xlsx'
put(f,'配置清单',8,3,999)
put(f,'配置清单',8,5,'第二章36星11000012追加25×5；保留原998行。正式表先预建新增1行并下拉ID公式，再按DropId粘贴。')
put(f,'验收结果',6,3,'37节点；第二章7/14/21/28/36星；36星引用11000012，原奖励外追加映射之匙×5')

hashof=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
guards={}
for folder in ['任务','签到','无限之旅','主线高难','映射装备']:
    for p in (FINAL/folder).glob('*.xlsx'):
        if not p.name.startswith('~$'): guards[str(p)]=hashof(p)
formal=Path('D:/Vgame/Config/GameConfig/Datas')
for p in formal.rglob('*.xlsx'):
    if not p.name.startswith('~$'): guards[str(p)]=hashof(p)
files=[]
for file,edits in changes.items():
    p=FINAL/file
    wb=load_workbook(p,read_only=True,data_only=False)
    for e in edits:
        assert wb[e['sheet']].cell(e['row'],e['col']).value==e['old'],(file,e)
    wb.close()
    backup=RUN/'backup'/file; backup.parent.mkdir(parents=True,exist_ok=True)
    if backup.exists(): assert hashof(backup)==hashof(p),'Backup already differs; do not overwrite'
    else: shutil.copy2(p,backup)
    files.append({'file':file,'source':str(p),'backup':str(backup),'sha256':hashof(p),'edits':edits})
plan={'batch':'KEY-TASK-20260827','created':datetime.now().isoformat(),'root':str(ROOT),'run':str(RUN),'files':files,'guards':guards,'append_main_row':appendrow}
(RUN/'plan.json').write_text(json.dumps(plan,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'files':[(f['file'],len(f['edits'])) for f in files],'guards':len(guards),'main_append_row':appendrow},ensure_ascii=False))
