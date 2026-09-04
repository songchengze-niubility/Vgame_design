import json, hashlib, sys
from pathlib import Path
from collections import Counter, defaultdict
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

HERE=Path(__file__).resolve().parent
ROOT=HERE.parents[1]
RUN=ROOT/'05_验算核算'/'映射之匙任务_20260827'
PLAN=json.loads((RUN/'plan.json').read_text(encoding='utf-8'))
PUBLISHED='--published' in sys.argv
DEST=ROOT/'06_最终配置' if PUBLISHED else RUN/'stage'
checks=[]; cache={}; differences=[]
def check(name,passed,detail=None):
    checks.append({'check':name,'pass':bool(passed),'detail':detail})
def read(p,data=False):
    w=load_workbook(p,read_only=True,data_only=data)
    rows={s.title:[list(r) for r in s.values] for s in w}
    w.close(); return rows
def book(file):
    if file not in cache:
        p=DEST/file
        cache[file]=read(p if p.exists() else ROOT/'06_最终配置'/file)
    return cache[file]
def table(file,sheet):
    rows=book(file)[sheet]; head=rows[0]
    return [dict(zip(head,r)) for r in rows[4:] if len(r)>1 and r[0]!='##' and r[1] is not None]
def drop(file):
    # Candidate mainline uses intentionally blank ID cells; identify by DropId.
    rows=book(file)['Drop']; head=rows[0]; result=defaultdict(list)
    p=DEST/file
    cached=read(p if p.exists() else ROOT/'06_最终配置'/file,True)['Drop']
    for ri,r in enumerate(rows[4:],4):
        if len(r)>2 and r[0]!='##' and r[2] is not None:
            resolved=list(r)
            for ci,v in enumerate(r):
                if isinstance(v,str) and v.startswith('='):
                    resolved[ci]=cached[ri][ci]
                    check('公式缓存存在:'+file+f'!{get_column_letter(ci+1)}{ri+1}',resolved[ci] is not None)
            result[resolved[2]].append(dict(zip(head,resolved)))
    return result
def reward(rows):
    c=Counter()
    for r in rows:
        check(f"固定掉落:{r['DropId']}/{r['DropItem_id']}",r['DropType']=='fixed_' and r['DropId2']==0 and r['DropMinNum']==r['DropMaxNum']==1 and r['DropItemNumMin']==r['DropItemNumMax'])
        c[r['DropItem_id']]+=r['DropItemNumMin']
    return c
for f in PLAN['files']:
    file=f['file']; old=read(Path(f['backup'])); new=book(file)
    expected={(e['sheet'],e['row'],e['col']):e['new'] for e in f['edits']}
    errors=[]; formula_errors=[]
    for sheet in old:
        a=old[sheet]; b=new[sheet]
        for ri in range(max(len(a),len(b))):
            ar=a[ri] if ri<len(a) else []; br=b[ri] if ri<len(b) else []
            for ci in range(max(len(ar),len(br))):
                av=ar[ci] if ci<len(ar) else None; bv=br[ci] if ci<len(br) else None
                ev=expected.get((sheet,ri+1,ci+1),av)
                if bv!=ev: errors.append([sheet,ri+1,ci+1,ev,bv])
                if av!=bv: differences.append({'file':file,'sheet':sheet,'cell':f'{get_column_letter(ci+1)}{ri+1}','old':av,'new':bv})
                if isinstance(bv,str) and bv in ['#REF!','#VALUE!','#DIV/0!','#N/A','#NAME?','#NUM!']: formula_errors.append([sheet,ri+1,ci+1,bv])
    check('逐单元格改动白名单:'+file,not errors,errors[:10])
    check('公式错误:'+file,not formula_errors,formula_errors)
    check('页签完整:'+file,set(new)==set(old)|({'钥匙来源核对'} if file=='任务/00_配置清单.xlsx' else set()))

daily=table('任务/DailyTask.xlsx','Sheet1'); weekly=table('任务/WeeklyTask.xlsx','Sheet1')
tasks={r['Id']:r for r in table('任务/Task.xlsx','Task')}
targets={r['Id']:r for r in table('任务/TaskTarget.xlsx','TaskTarget')}
dd=drop('任务/Drop_other.xlsx')
dailyids=[1101,1103,1102,1104,1106,1107,1108,1109,1110,1111]
weekids=[[12000+g*10+t for t in [1,2,3]] for g in range(1,12)]
check('每日10条索引及顺序',[r['Task'] for r in daily]==dailyids)
check('每周11组三级索引',[[int(x) for x in str(r['TaskID']).split(';')] for r in weekly]==weekids)
daily_points=[reward(dd[tasks[i]['DropId']])[16] for i in dailyids]
check('日活10/20分档及130总量',daily_points==[10,20,10,10,10,20,20,10,10,10] and sum(daily_points)==130)
check('前7条100日活',sum(daily_points[:7])==100)
for group in weekids:check('周常1/2/4活跃:'+str(group),[reward(dd[tasks[i]['DropId']])[17] for i in group]==[1,2,4])
check('周活上限与基础可达',sum([1,2,4])*11==77 and 6*7+7*3==63 and 77+42==119)
daily_expected={20:{1:20000,4:300,17:1},40:{102:5,4:300,17:1},60:{2:50,4:300,17:1},80:{25:3,4:300,17:1},100:{2:50,4:300,17:2,26:800}}
week_expected={10:{2:50,110:5},20:{2:50,25:4},30:{2:100,102:5},40:{2:100,25:6},50:{2:150,3027:5},60:{2:150,3016:2}}
totalday=Counter(); totalweek=Counter()
for file,sheet,expect,total in [('任务/DailyActiveReward.xlsx','BattlePassDailyActiveReward',daily_expected,totalday),('任务/WeeklyReward.xlsx','Sheet1',week_expected,totalweek)]:
    ts=table(file,sheet);check('活跃节点:'+file,[r['ActivePoint'] for r in ts]==list(expect))
    for r in ts:
        got=reward(dd[r['RewardId']]);total.update(got)
        check('活跃真实包:'+str(r['ActivePoint'])+file,got==Counter(expect[r['ActivePoint']]),dict(got))
check('任务钥匙31/周',totalday[25]*7+totalweek[25]==31)
expectedtypes=[('Target_TTT_Load','0,1'),('Target_TTT_Item_Consume','3,160'),('Target_TTT_Type_LevelClearNum','1,3'),('Target_TTT_Sweep_Times','0,3'),('Target_TTT_TypeMonsterNum','3,50'),('Target_TTT_Type_LevelClearNum','0,1'),('Target_TTT_Arena_Challenge_Times','0,1'),('Target_TTT_SpecialEquipDungeon_Challenge_Times','0,1'),('Target_TTT_BreakdowntSpecialEquipment_Num','0,1'),('Target_TTT_TravelDungeon_Challenge_Times','0,1')]
for i,expected in zip(dailyids,expectedtypes):
    t=tasks[i]; target=targets[int(t['TaskTarget'])]
    check('日常目标配置:'+str(i),(target['TargetType'],target['ParamList'])==expected and t['TaskType']=='Task_TTT_Daily')
thresholds=[[2,4,6],[300,700,1200],[10,25,40],[3,7,12],[3,8,15],[3,8,15],[1000,3000,5000],[2,4,7],[5,10,15],[1000,3000,5000],[5,10,15]]
weektypes=['Login_Total','Item_Consume','DailyTask_Finish','Type_LevelClearNum','Arena_Challenge_Times','SpecialEquipDungeon_Challenge_Times','InfiniteChallenge_Score','Trade_CompleteTimes','Type_LevelClearNum','Roguelike_AccumulateScore','TravelDungeon_Challenge_Times']
for g,(group,nums,suffix) in enumerate(zip(weekids,thresholds,weektypes),1):
    for i,n in zip(group,nums):
        t=tasks[i]; target=targets[int(t['TaskTarget'])]; param=f'{3 if g==2 else 37 if g==9 else 0},{n}'
        check('周常目标配置:'+str(i),(target['TargetType'],target['ParamList'])==('Target_TTT_'+suffix,param) and t['TaskType']=='Task_TTT_Weekly')
formal=json.loads((HERE/'formal.json').read_text(encoding='utf-8'))
unlock={r[1]:(r[3],r[4]) for r in formal['unlock/Unlock.xlsx']['Unlock'][4:] if len(r)>4 and isinstance(r[1],int)}
for i in [1108,12051,12052,12053]:check('竞技场任务2-8:'+str(i),unlock[tasks[i]['UnlockId']]==('UnLock_UUT_Level',100208))
for i in [1109,1110,12061,12062,12063]:check('映射任务3-10:'+str(i),unlock[tasks[i]['UnlockId']]==('UnLock_UUT_Level',100310))
enum_names={str(r[7]) for rows in formal['__enums__.xlsx'].values() for r in rows if len(r)>7}
check('日周任务类型枚举已注册',all(targets[int(tasks[i]['TaskTarget'])]['TargetType'] in enum_names for i in dailyids+sum(weekids,[])))

nr=drop('无限之旅/Drop_dungeon.xlsx')
score=table('无限之旅/NR_HighestScoreReward.xlsx','WeekScoreReward')
check('神经漫游40包120项',len(nr)==40 and sum(map(len,nr.values()))==120)
nr_totals=[]
for difficulty in range(1,5):
    total=Counter()
    for node in range(1,11):
        did=23000000+(difficulty-1)*10+node; got=reward(nr[did]); total.update(got)
        check('神经漫游钥匙节点:'+str(did),got[25]==(2 if node%2==0 else 0))
    check('单难度10钥匙:'+str(difficulty),total[25]==10)
    nr_totals.append(dict(total))
nrrows=book('无限之旅/Drop_dungeon.xlsx')['Drop']
check('神经漫游120条ID公式完整',all(r[1]=='=ROW()+200000-4' for r in nrrows[4:] if len(r)>2 and r[2] is not None))
nw=load_workbook(DEST/'无限之旅/Drop_dungeon.xlsx',read_only=True,data_only=False)
badfont=[]
for ri in range(5,125):
    color=nw['Drop'].cell(ri,3).font.color
    if color and color.type=='rgb' and color.rgb[-6:].upper()=='FFFFFF':badfont.append(ri)
check('神经漫游DropId不为白色字体',not badfont,badfont);nw.close()
iw=load_workbook(Path('D:/Vgame/Config/GameConfig/Datas/items/Items.xlsx'),read_only=True,data_only=False)
itemids={r[1] for r in iw['items'].values if len(r)>1 and r[0]!='##'};iw.close()
check('钥匙与日周活跃道具存在',set([25,16,17])<=itemids)
signrows=book('签到/DailySign.xlsx')['Sheet1']
months=defaultdict(Counter);dates=[]
for r in signrows:
    if len(r)>4 and r[0] not in ['##','##var','##type','##group'] and str(r[1]).isdigit() and str(r[2]).isdigit() and str(r[3]).isdigit():
        dates.append((int(r[2]),int(r[3])))
        for pair in str(r[4]).split(';'):
            item,qty=map(int,pair.split(','));months[int(r[2])][item]+=qty
            if item==25:check('月签钥匙日:'+str(r[1]),int(r[3]) in [5,12,19,26] and qty==2)
check('365天唯一日期',len(dates)==len(set(dates))==365)
check('每月8全年96',len(months)==12 and all(m[25]==8 for m in months.values()) and sum(m[25] for m in months.values())==96)
star=[r for r in table('主线高难/ChapterStarReward.xlsx','ChapterStarReward') if r['ChapterId']==10002 and r['Star']==36 and not r['IsDifficult']]
check('第二章36星实际Drop11000012',len(star)==1 and star[0]['DropId']==11000012)
main=drop('主线高难/Drop_mainlevel.xlsx')
check('第二章36星原奖励保留追加5把',reward(main[11000012])==Counter({2:150,101:12,111:3,25:5}))
check('第二章28星不发钥匙',reward(main[11000011])[25]==0)
check('主线候选999行',sum(map(len,main.values()))==999)
activity=[r for r in table('任务/FourteenDayActivity.xlsx','FourteenDayActivity') if r['ActivityId']==1]
sevenids=[int(v) for r in activity for v in str(r['TaskList']).split(';')]
sevendrop=drop('任务/Drop_TimeLimitedActivity.xlsx'); seven=Counter()
for i in sevenids: seven.update(reward(sevendrop[tasks[i]['DropId']]))
check('七日70任务8钥匙保持',len(sevenids)==70 and seven[25]==8,{'tasks':len(sevenids),'keys':seven[25]})
check('每周41与每月8分开计算',totalday[25]*7+totalweek[25]+nr_totals[0][25]==41 and months[1][25]==8)
summaries=read(DEST/'任务/00_配置清单.xlsx',True)
check('清单公式缓存31及41',summaries['数值校验'][9][1]==31 and summaries['钥匙来源核对'][7][3]==41)
ns=read(DEST/'无限之旅/00_配置清单.xlsx',True)
check('神经漫游清单四难度各10',[r[5] for r in ns['周奖励汇总'][1:]]==[10]*4)
edited={str(Path(f['source'])) for f in PLAN['files']}
drift=[]
for name,h in PLAN['guards'].items():
    if PUBLISHED and name in edited:continue
    if hashlib.sha256(Path(name).read_bytes()).hexdigest()!=h: drift.append(name)
check('SVN与本轮未改候选hash保护',not drift,drift)
issues=[
 '资源副本目标1107/12041~12043复用Type_LevelClearNum参数0,N；仅资源类型过滤和挑战/通关语义未验证。',
 '无限之旅周最高分暂复用InfiniteChallenge_Score；需验证最高值/非累计、周清零和四难度奖励互斥。',
 '未解锁任务隐藏及解锁后下一日/周入池需服务端联调。',
 '当前__tables__未注册DailyTask输入；需确认TaskType与任务池配置的消费入口。',
 '本轮未执行Luban或服务端运行验证；SVN未改。'
]
report={'batch':PLAN['batch'],'phase':'published' if PUBLISHED else 'staged','status':'PASS_STATIC_WITH_RUNTIME_GATES' if all(c['pass'] for c in checks) else 'FAIL','passed':sum(c['pass'] for c in checks),'failed':[c for c in checks if not c['pass']],'checks':checks,'weekly_keys':41,'monthly_keys':8,'nr_totals':nr_totals,'runtime_gates':issues,'differences':differences}
name='published-validation.json' if PUBLISHED else 'staged-validation.json'
(RUN/name).write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({k:report[k] for k in ['phase','status','passed','failed','weekly_keys','monthly_keys']},ensure_ascii=False))
sys.exit(0 if report['status'].startswith('PASS') else 1)
