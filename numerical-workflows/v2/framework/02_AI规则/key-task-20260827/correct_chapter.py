"""Read-only audit / JSON staging for the user-requested chapter correction.

The Task workbook is restored byte-for-byte only after confirming that its sole
data differences from the pre-error backup are the rejected 18 cells.
Spreadsheet authoring is handled by correct_chapter.mjs (artifact-tool).
"""
import hashlib
import json
import shutil
import sys
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
FINAL = ROOT / '06_最终配置'
RUN = ROOT / '05_验算核算' / '章节任务纠偏_20260827'
OLD = ROOT / '05_验算核算' / '映射之匙任务_20260827'
ARENA = {1108, 12051, 12052, 12053}
MAPPING = {1109, 1110, 12061, 12062, 12063}

def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()

def workbook(path, data_only=False):
    w = load_workbook(path, read_only=True, data_only=data_only)
    result = {s.title: [list(r) for r in s.values] for s in w}
    w.close()
    return result

def records(path, sheet=None):
    w = workbook(path)
    rows = w[sheet] if sheet else next(iter(w.values()))
    return {r[1]: (i, r) for i, r in enumerate(rows, 1) if i > 4 and len(r)>1 and isinstance(r[1], int) and r[0] != '##'}

def cellmap(path):
    return {(s, ri, ci): v for s, rows in workbook(path).items()
            for ri, row in enumerate(rows, 1) for ci, v in enumerate(row, 1) if v is not None}

def compare(a, b):
    x, y = cellmap(a), cellmap(b)
    return [{'sheet':k[0], 'row':k[1], 'col':k[2], 'old':x.get(k), 'new':y.get(k)}
            for k in sorted(x.keys() | y.keys()) if x.get(k) != y.get(k)]

def audit_chapters(task_path):
    task = records(task_path, 'Task')
    target = records(FINAL/'任务/TaskTarget.xlsx', 'TaskTarget')
    groups = records(FINAL/'任务/JottingsTask.xlsx')
    chapters = records(FINAL/'任务/Jottings.xlsx')
    drop = workbook(FINAL/'任务/Drop_other.xlsx')['Drop']
    packs = {}
    for r in drop[4:]:
        if len(r)>12 and r[0]!='##' and isinstance(r[2], int): packs.setdefault(r[2], []).append(r)
    result, all_tasks, sizes = [], [], []
    for ch, (_, row) in chapters.items():
        gids = [int(x) for x in str(row[3]).split(';')]
        sizes.append(len(gids))
        for gid in gids:
            assert gid in groups
            for tid in map(int, str(groups[gid][1][2]).split(';')):
                all_tasks.append(tid)
                assert tid in task
                tr, t = task[tid]
                assert t[6]=='Task_TTT_Jottings' and t[7] in target
                rr, tar = target[t[7]]
                assert t[14] in packs
                reward_rows = packs[t[14]]
                data = dict(chapter=ch, group_id=gid, task_id=tid, target_id=t[7],
                            task_row=tr, target_row=rr, description=t[2], target_type=tar[3],
                            param=tar[4], drop_id=t[14],
                            rewards=[{'item':r[10], 'min':r[11], 'max':r[12]} for r in reward_rows])
                if tar[3] == 'Target_TTT_SpecialEquipmentLevelUp_Num':
                    data['category']='映射装备'
                    assert len(reward_rows)==1
                    r=reward_rows[0]
                    assert r[5]==0 and r[6]=='fixed_' and r[7:9]==[1,1] and r[10:14]==[25,3,3,10000]
                elif tar[3] == 'Target_TTT_RvCabin_LevelUp_Num': data['category']='舱室'
                elif tar[3] == 'Target_TTT_RoleLevel': data['category']='玩家等级'
                else: continue
                assert t[2]==t[3]==tar[2]
                result.append(data)
    assert sizes==[5,6,7,8,10,10,10]
    assert len(all_tasks)==len(set(all_tasks))==69
    mapping=[(x['chapter'],x['group_id'],x['task_id'],x['param'],x['drop_id']) for x in result if x['category']=='映射装备']
    assert mapping==[(2,35,1327,'5,4',1030307),(3,47,1339,'10,4',1030407),(4,57,1349,'15,8',1030507),(5,67,1381,'25,8',1030607),(6,77,1384,'35,8',1030707),(7,87,1388,'45,8',1030807)]
    cabin=[(x['chapter'],x['task_id'],x['param']) for x in result if x['category']=='舱室']
    assert cabin==[(1,1315,'2,1'),(2,1328,'2,2'),(3,1338,'3,1'),(4,1348,'3,2'),(5,1401,'4,1')]
    levels=[x['param'] for x in result if x['category']=='玩家等级']
    assert levels==['10,1','20,1','25,1','30,1','33,1','36,1','38,1']
    return {'group_counts':sizes, 'group_total':sum(sizes), 'task_count':len(all_tasks),
            'mapping_keys_one_time':18, 'verified_existing_chapter_rows':result,
            'chapter_config_changes_this_batch':0}

def prepare():
    assert not (RUN/'plan.json').exists(), 'Existing correction plan; do not overwrite baseline'
    RUN.mkdir(parents=True, exist_ok=True)
    current=FINAL/'任务/Task.xlsx'; old=OLD/'backup/任务/Task.xlsx'
    differences=compare(current,old)
    tasks=records(current,'Task')
    expected=[]
    for tid in sorted(ARENA|MAPPING):
        row=tasks[tid][0]; arena=tid in ARENA
        expected.extend([{'sheet':'Task','row':row,'col':11,'old':123 if arena else 140,'new':4112 if arena else 4314},
                         {'sheet':'Task','row':row,'col':12,'old':'主线2-8' if arena else '主线3-10','new':None}])
    assert sorted(differences,key=lambda x:(x['row'],x['col']))==sorted(expected,key=lambda x:(x['row'],x['col']))
    chapter=audit_chapters(current)
    changes={
      ('配置清单','C3'):'每日10、每周33、章节当前挂载69条任务目标、七日70；章节玩家等级10/20/25/30/33/36/38。未挂载旧记录不计入章节数量。',
      ('配置清单','C4'):'保留四套任务本体及引用；撤回误改的9条日周UnlockId及备注。章节任务实际落点见“章节任务落点”。',
      ('配置清单','C10'):'7章共56组目标；主线分步展开后69条任务。第二章组35已挂载1327映射装备任务。',
      ('落地说明','C5'):'撤回本轮擅自同步：竞技场任务UnlockId恢复4112，映射日周任务恢复4314，9条新增备注清空。其余解锁不动。',
      ('落地说明','C9'):'已从Jottings逐层核对56组、69条章节任务：前两章舱室及第二章映射任务均已在06；本轮未新增或覆盖章节任务。',
      ('落地说明','C10'):'章节要求按8月26日用户确认；本轮纠偏见03_沉淀MD/经济/映射之匙与任务落06-2026-08-27.md。仅06，SVN未改。',
      ('任务总览','C4'):'5/6/7/8/10/10/10组；玩家等级10/20/25/30/33/36/38。舱室按1个Lv2、2个Lv2、1个Lv3、2个Lv3、1个Lv4递进。第二至七章映射任务各奖3把钥匙。',
      ('任务总览','D4'):'主线每分步元晶30；章节映射任务共18把钥匙（一次性）。其余任务奖励及章节完成奖励保留。',
      ('数值校验','B6'):56, ('数值校验','C6'):56,
      ('数值校验','B7'):69, ('数值校验','C7'):69,
    }
    files=[]
    for relative in ['任务/Task.xlsx','任务/00_配置清单.xlsx','任务/任务配置校验.json']:
        source=FINAL/relative; backup=RUN/'backup'/relative; backup.parent.mkdir(parents=True,exist_ok=True)
        shutil.copy2(source,backup)
        files.append({'file':relative,'source':str(source),'backup':str(backup),'sha256':digest(source)})
    mf=files[1]; w=load_workbook(FINAL/mf['file'],read_only=True)
    mf['edits']=[{'sheet':s,'cell':c,'old':w[s][c].value,'new':v} for (s,c),v in changes.items()];w.close()
    plan={'batch':'CHAPTER-CORRECTION-20260827','files':files,'task_restore_from':str(old),'task_restore_sha256':digest(old),
          'task_edits':differences,'chapter_audit':chapter,
          'preserved_workbooks':{str(p):digest(p) for p in FINAL.rglob('*.xlsx') if p not in [FINAL/f['file'] for f in files]},
          'formal_protected':{str(p):digest(p) for rel in ['task/Task.xlsx','task/TaskTarget.xlsx','task/Jottings.xlsx','task/JottingsTask.xlsx','Drop/Drop_other.xlsx'] if (p:=Path('D:/Vgame/Config/GameConfig/Datas')/rel).exists()}}
    (RUN/'plan.json').write_text(json.dumps(plan,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({'prepared':len(files),'rollback_cells':len(differences),'chapter_groups':chapter['group_counts'],'chapter_tasks':69,'existing_mapping_keys':18},ensure_ascii=False))

def stage_restore_and_report():
    plan=json.loads((RUN/'plan.json').read_text(encoding='utf-8'))
    old=Path(plan['task_restore_from']);assert digest(old)==plan['task_restore_sha256']
    dest=RUN/'stage/任务/Task.xlsx';dest.parent.mkdir(parents=True,exist_ok=True);shutil.copy2(old,dest)
    f=plan['files'][2];report=json.loads(Path(f['source']).read_text(encoding='utf-8'))
    report.update(batch=plan['batch'],status='PASS_STATIC_WITH_RUNTIME_GATES',current_validation='../../05_验算核算/章节任务纠偏_20260827/published-validation.json',
                  task_unlocks={'arena':4112,'mapping':4314,'basis':'撤回用户否决的9条日周解锁改动，恢复原06值；不是新增设计确认'},
                  chapter_audit=plan['chapter_audit'],previous_validation='../../05_验算核算/映射之匙任务_20260827/published-validation.json',
                  previous_validation_notice='旧报告中9条UnlockId为123/140的规则被用户否决；钥匙数值修改保留。')
    (RUN/'stage'/f['file']).write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    for f in plan['files']:f['stage_sha256']=digest(RUN/'stage'/f['file'])
    (RUN/'plan.json').write_text(json.dumps(plan,ensure_ascii=False,indent=2),encoding='utf-8')

def validate(published):
    plan=json.loads((RUN/'plan.json').read_text(encoding='utf-8'))
    base=FINAL if published else RUN/'stage'
    task=base/'任务/Task.xlsx'
    assert digest(task)==plan['task_restore_sha256'], 'Task is not exact pre-error backup'
    chapter=audit_chapters(task)
    assert chapter==plan['chapter_audit']
    for tid in ARENA|MAPPING:
        r=records(task,'Task')[tid][1]
        assert r[10]==(4112 if tid in ARENA else 4314) and r[11] is None
    mf=plan['files'][1]
    diff=compare(Path(mf['backup']),base/mf['file'])
    w=load_workbook(base/mf['file'],read_only=True,data_only=False)
    allowed={(e['sheet'],w[e['sheet']][e['cell']].row,w[e['sheet']][e['cell']].column) for e in mf['edits']}
    for e in mf['edits']:assert w[e['sheet']][e['cell']].value==e['new']
    for d in diff:
        assert (d['sheet'],d['row'],d['col']) in allowed or d['sheet']=='章节任务落点' or (d['sheet']=='钥匙来源核对' and d['row']==9),d
    assert w['钥匙来源核对']['D8'].value=='=SUM(D2:D4)'
    assert w['钥匙来源核对']['D9'].value=='=B9*C9'
    assert w['章节任务落点']['D2'].value==1327
    w.close()
    caches=workbook(base/mf['file'],True)
    assert caches['钥匙来源核对'][7][3]==41
    assert caches['钥匙来源核对'][8][3]==18
    for values in caches.values():
        for row in values:
            assert not any(isinstance(v,str) and v in {'#REF!','#VALUE!','#DIV/0!','#NAME?','#NUM!','#N/A'} for v in row)
    for p,sha in (plan['preserved_workbooks']|plan['formal_protected']).items():assert digest(Path(p))==sha,p
    for f in plan['files']:assert digest(base/f['file'])==f['stage_sha256']
    output={'batch':plan['batch'],'status':'PASS_STATIC_WITH_RUNTIME_GATES','stage':'published' if published else 'staged',
            'task_restored_byte_identical':True,'restored_cells':plan['task_edits'], 'chapter_audit':chapter,
            'preserved_workbooks':len(plan['preserved_workbooks']),'formal_protected':plan['formal_protected'],
            'mapping_first_clear_4375_unchanged':True,'weekly_keys':41,'monthly_keys':8,
            'manifest_diff':diff,'runtime_gates':['仅06静态验收；未合SVN、未执行Luban或客户端/服务端运行验收。']}
    (RUN/('published-validation.json' if published else 'staged-validation.json')).write_text(json.dumps(output,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({'status':output['status'],'stage':output['stage'],'chapter_tasks':69,'chapter_changes_this_batch':0,'restored_cells':18,'preserved_workbooks':len(plan['preserved_workbooks'])},ensure_ascii=False))

if __name__=='__main__':
    if sys.argv[1]=='prepare':prepare()
    elif sys.argv[1]=='stage':stage_restore_and_report()
    elif sys.argv[1]=='validate':validate('--published' in sys.argv)
    else:raise ValueError(sys.argv[1])
