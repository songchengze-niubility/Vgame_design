#!/usr/bin/env python3
"""只读验证映射装备配置草案 D01 与正式配置零写入证据。"""

from __future__ import annotations

import argparse
import hashlib
import json
import subprocess
import sys
from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

import populate_mapping_equip_draft as draft_generator


ERROR_VALUES = {"#DIV/0!", "#N/A", "#NAME?", "#NULL!", "#NUM!", "#REF!", "#VALUE!"}
EXPECTED_FORMULAS = 2781
DEFAULT_SCHEMA_REPORT = (
    draft_generator.FRAME_ROOT / "05_验算核算" / "vgame-config-mapping-report.json"
)
DEFAULT_OUTPUT = (
    draft_generator.FRAME_ROOT / "05_验算核算" / "mapping-equip-draft-D01-validation.json"
)


def text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def business_rows(worksheet) -> list[list[Any]]:
    rows: list[list[Any]] = []
    for row in worksheet.iter_rows(min_row=2, values_only=True):
        values = list(row)
        if any(value not in (None, "") for value in values):
            rows.append(values)
    return rows


def normalize(value: Any) -> Any:
    if isinstance(value, float):
        return round(value, 12)
    return value


def aggregate_snapshot(records: list[dict[str, Any]]) -> str:
    digest = hashlib.sha256()
    for item in sorted(records, key=lambda row: row["path"]):
        digest.update(
            f"{item['path']}\0{item['bytes']}\0{item['sha256']}\n".encode("utf-8")
        )
    return digest.hexdigest()


def validate_source_snapshot(
    config_root: Path,
    schema_report_path: Path,
    errors: list[str],
) -> tuple[int, str, str]:
    report = json.loads(schema_report_path.read_text(encoding="utf-8"))
    source_snapshot = report.get("source_snapshot") or {}
    expected_records = source_snapshot.get("files") or []
    expected_hash = text(report.get("source_snapshot_hash") or source_snapshot.get("sha256"))
    live_records: list[dict[str, Any]] = []
    for item in expected_records:
        relative_path = text(item.get("path")).replace("\\", "/")
        path = config_root / relative_path
        if not path.is_file():
            errors.append(f"正式配置快照文件缺失: {relative_path}")
            continue
        live_size = path.stat().st_size
        live_sha = sha256_file(path)
        if live_size != int(item.get("bytes") or -1):
            errors.append(f"正式配置文件大小变化: {relative_path}")
        if live_sha != text(item.get("sha256")):
            errors.append(f"正式配置文件hash变化: {relative_path}")
        live_records.append({"path": relative_path, "bytes": live_size, "sha256": live_sha})
    live_hash = aggregate_snapshot(live_records)
    if len(live_records) != len(expected_records):
        errors.append(
            f"正式配置快照文件数变化: live={len(live_records)}, expected={len(expected_records)}"
        )
    if expected_hash and live_hash != expected_hash:
        errors.append(f"正式Datas注册源快照变化: live={live_hash}, expected={expected_hash}")
    return len(live_records), live_hash, expected_hash


def font_is_blue(cell) -> bool:
    color = cell.font.color
    if color is None:
        return False
    if color.type == "rgb":
        return text(color.rgb).upper().endswith("0000FF")
    return False


def validate(args) -> tuple[list[str], dict[str, Any]]:
    errors: list[str] = []
    workbook_path = Path(args.workbook).resolve()
    config_root = Path(args.config_root).resolve()
    schema_report_path = Path(args.schema_report).resolve()
    output_path = Path(args.output).resolve()

    for path, label in [
        (workbook_path, "映射总表"),
        (schema_report_path, "Schema报告"),
    ]:
        if not path.is_file():
            errors.append(f"{label}不存在: {path}")
    if not config_root.is_dir():
        errors.append(f"正式配置根目录不存在: {config_root}")
    if errors:
        return errors, {}

    workbook = None
    formula_count = 0
    formula_errors = 0
    cached_errors = 0
    external_links = 0
    batch_rows: list[list[Any]] = []
    module_paths: set[str] = set()
    source_count = 0
    source_hash = ""
    expected_source_hash = ""
    try:
        workbook = load_workbook(workbook_path, data_only=False, keep_links=True)
        required_sheets = {
            "使用说明", "物理字段来源", "模块映射", "配置草案", "审批记录",
            "落地回执", "工作流门禁", "__generation_meta__",
        }
        missing_sheets = sorted(required_sheets - set(workbook.sheetnames))
        if missing_sheets:
            errors.append(f"工作簿缺少Sheet: {missing_sheets}")
            return errors, {}

        meta_sheet = workbook["__generation_meta__"]
        metadata = {
            text(meta_sheet.cell(row, 1).value): meta_sheet.cell(row, 2).value
            for row in range(1, meta_sheet.max_row + 1)
        }
        if metadata.get("generation_marker") != draft_generator.GENERATION_MARKER:
            errors.append("生成标记不匹配")
        if metadata.get("draft_batch_id") != draft_generator.BATCH_ID:
            errors.append("元数据批次ID不匹配")
        if text(metadata.get("draft_batch_version")) != draft_generator.BATCH_VERSION:
            errors.append("元数据批次版本不匹配")
        if metadata.get("official_config_write") not in (False, "False", "false", 0):
            errors.append("元数据official_config_write必须为False")
        if meta_sheet.sheet_state != "veryHidden":
            errors.append("__generation_meta__必须保持veryHidden")

        draft_sheet = workbook["配置草案"]
        actual_headers = [text(cell.value) for cell in draft_sheet[1][: len(draft_generator.DRAFT_HEADERS)]]
        if actual_headers != draft_generator.DRAFT_HEADERS:
            errors.append("配置草案表头漂移")
        all_draft_rows = business_rows(draft_sheet)
        batch_rows = [row[: len(draft_generator.DRAFT_HEADERS)] for row in all_draft_rows if text(row[0]) == draft_generator.BATCH_ID]
        other_batch_ids = sorted({text(row[0]) for row in all_draft_rows if text(row[0]) != draft_generator.BATCH_ID})
        if other_batch_ids:
            errors.append(f"D01首批验证发现其他批次: {other_batch_ids}")

        expected_rows = draft_generator.build_draft_rows(config_root, "<TIMESTAMP>")
        expected_by_id = {text(row[1]): row for row in expected_rows}
        actual_by_id: dict[str, list[Any]] = {}
        for row_number, row in enumerate(batch_rows, start=2):
            draft_id = text(row[1])
            if not draft_id:
                errors.append(f"配置草案第{row_number}行缺草案ID")
                continue
            if draft_id in actual_by_id:
                errors.append(f"草案ID重复: {draft_id}")
            actual_by_id[draft_id] = row
            for column, header in enumerate(draft_generator.DRAFT_HEADERS):
                if row[column] in (None, ""):
                    errors.append(f"{draft_id} 必填列为空: {header}")
            if text(row[2]) != draft_generator.BATCH_VERSION:
                errors.append(f"{draft_id} 批次版本错误")
            if text(row[21]) not in {
                "CANDIDATE_SAFE_RECHECK", "MANUAL_REQUIRED", "SCHEMA_CHANGE", "NEW_ROW_MANUAL"
            }:
                errors.append(f"{draft_id} 写入方式非法: {row[21]}")
            if text(row[22]) not in {"LOW", "MEDIUM", "HIGH"}:
                errors.append(f"{draft_id} 风险等级非法: {row[22]}")
            if text(row[24]) not in {"DRAFT", "READY_FOR_REVIEW", "MANUAL_REQUIRED"}:
                errors.append(f"{draft_id} 出现不允许的首批状态: {row[24]}")
            if text(row[24]) in {"APPROVED", "APPLIED", "EXPORTED", "VERIFIED"}:
                errors.append(f"{draft_id} 非法越过人工门禁")
            source_md = draft_generator.FRAME_ROOT / text(row[5])
            if not source_md.is_file():
                errors.append(f"{draft_id} 来源MD不存在: {row[5]}")
            else:
                md_text = source_md.read_text(encoding="utf-8")
                if text(row[4]) not in md_text:
                    errors.append(f"{draft_id} 参数ID未登记在来源MD: {row[4]}")
            transform = text(row[16])
            if transform.startswith("="):
                errors.append(f"{draft_id} 变换公式被误写成Excel公式")

        if set(actual_by_id) != set(expected_by_id):
            errors.append(
                "草案ID集合与生成器不一致: "
                f"missing={sorted(set(expected_by_id)-set(actual_by_id))[:20]}, "
                f"extra={sorted(set(actual_by_id)-set(expected_by_id))[:20]}"
            )
        compare_columns = [index for index in range(len(draft_generator.DRAFT_HEADERS)) if index != 25]
        for draft_id in sorted(set(actual_by_id) & set(expected_by_id)):
            actual = actual_by_id[draft_id]
            expected = expected_by_id[draft_id]
            for index in compare_columns:
                if normalize(actual[index]) != normalize(expected[index]):
                    errors.append(
                        f"{draft_id} 与权威生成结果不一致: "
                        f"{draft_generator.DRAFT_HEADERS[index]} actual={actual[index]!r} expected={expected[index]!r}"
                    )
                    break

        table = draft_sheet.tables.get("TblDrafts")
        expected_table_ref = f"A1:AA{len(batch_rows) + 1}"
        if table is None:
            errors.append("配置草案缺少TblDrafts")
        elif table.ref != expected_table_ref:
            errors.append(f"TblDrafts范围错误: {table.ref} != {expected_table_ref}")
        if draft_sheet.auto_filter.ref is not None:
            errors.append(
                f"配置草案不应同时存在工作表级AutoFilter与TblDrafts: {draft_sheet.auto_filter.ref}"
            )
        validation_ranges = {str(item.sqref) for item in draft_sheet.data_validations.dataValidation}
        for expected_range in {"V2:V5000", "W2:W5000", "Y2:Y5000"}:
            if expected_range not in validation_ranges:
                errors.append(f"配置草案数据验证缺失: {expected_range}")
        for row_number in range(2, len(batch_rows) + 2):
            for column in range(1, len(draft_generator.DRAFT_HEADERS) + 1):
                if not font_is_blue(draft_sheet.cell(row_number, column)):
                    errors.append(f"配置草案人工输入字体不是蓝色: {draft_sheet.cell(row_number,column).coordinate}")
                    break

        # 所有安全候选必须回查到真实物理字段，且该字段当前没有公式。
        field_sheet = workbook["物理字段来源"]
        field_headers = {text(cell.value): cell.column - 1 for cell in field_sheet[1] if text(cell.value)}
        schema_index: dict[tuple[str, str, str, str], dict[str, Any]] = {}
        for row in field_sheet.iter_rows(min_row=2, values_only=True):
            logicals = text(row[field_headers["逻辑表"]]).split(";")
            relative_path = text(row[field_headers["相对路径"]]).replace("\\", "/")
            sheet_name = text(row[field_headers["Schema Sheet"]])
            field_name = text(row[field_headers["字段名"]])
            for logical in logicals:
                schema_index[(logical, relative_path, sheet_name, field_name)] = {
                    "formula_count": int(row[field_headers["当前公式单元格数"]] or 0),
                    "raw_type": text(row[field_headers["RawType"]]),
                }
        for row in batch_rows:
            if text(row[21]) != "CANDIDATE_SAFE_RECHECK":
                continue
            key = (text(row[7]), text(row[8]).replace("\\", "/"), text(row[9]), text(row[12]))
            item = schema_index.get(key)
            if item is None:
                errors.append(f"安全候选无法回查Schema: {row[1]} {key}")
                continue
            if item["formula_count"]:
                errors.append(f"安全候选指向公式字段: {row[1]} formula_count={item['formula_count']}")
            if text(row[20]) != "否":
                errors.append(f"安全候选公式单元格标记错误: {row[1]}")

        # 模块映射只检查本批次预期路径，不干扰其他人工映射。
        module_sheet = workbook["模块映射"]
        module_headers = {text(cell.value): cell.column - 1 for cell in module_sheet[1] if text(cell.value)}
        expected_module_paths = set(draft_generator.module_update_values())
        for row in module_sheet.iter_rows(min_row=2, values_only=True):
            relative_path = text(row[module_headers["相对路径"]]).replace("\\", "/")
            if relative_path not in expected_module_paths:
                continue
            module_paths.add(relative_path)
            if text(row[module_headers["业务模块"]]) != "映射装备":
                errors.append(f"模块映射业务模块错误: {relative_path}")
            if not text(row[module_headers["来源MD"]]):
                errors.append(f"模块映射缺来源MD: {relative_path}")
            if text(row[module_headers["负责人"]]) != "待指定":
                errors.append(f"模块映射负责人不应虚构: {relative_path}")
        if module_paths != expected_module_paths:
            errors.append(
                f"模块映射路径不完整: missing={sorted(expected_module_paths-module_paths)}"
            )

        if business_rows(workbook["审批记录"]):
            errors.append("审批记录必须为空，不能替用户批准")
        if business_rows(workbook["落地回执"]):
            errors.append("落地回执必须为空，尚未写正式配置")
        if draft_generator.BATCH_ID not in text(workbook["使用说明"]["B12"].value):
            errors.append("使用说明当前阶段未指向D01")
        gate_rows = {text(row[0]): text(row[2]) for row in workbook["工作流门禁"].iter_rows(min_row=2, values_only=True)}
        if gate_rows.get("G1") != "PASS":
            errors.append("G1 MD追踪应为PASS")
        for gate in ("G2", "G3"):
            if gate_rows.get(gate) != "PASS_WITH_WARNINGS":
                errors.append(f"{gate}应为PASS_WITH_WARNINGS")
        for gate in ("G4", "G5", "G6", "G7", "G8", "G9"):
            if gate_rows.get(gate) != "NOT_STARTED":
                errors.append(f"{gate}不应提前推进: {gate_rows.get(gate)}")

        for worksheet in workbook.worksheets:
            for row in worksheet.iter_rows():
                for cell in row:
                    value = cell.value
                    if cell.data_type == "f" or (isinstance(value, str) and value.startswith("=")):
                        formula_count += 1
                    if cell.data_type == "e" or text(value).upper() in ERROR_VALUES:
                        formula_errors += 1
                        if formula_errors <= 20:
                            errors.append(f"显式Excel错误: {worksheet.title}!{cell.coordinate}={value}")
        external_links = len(getattr(workbook, "_external_links", []))
        if formula_count != int(args.expected_formulas):
            errors.append(f"公式数量变化: {formula_count} != {args.expected_formulas}")
        if external_links:
            errors.append(f"出现外链: {external_links}")
    finally:
        if workbook is not None:
            workbook.close()

    cached = None
    try:
        cached = load_workbook(workbook_path, read_only=True, data_only=True, keep_links=False)
        for worksheet in cached.worksheets:
            for row in worksheet.iter_rows():
                for cell in row:
                    value = cell.value
                    if cell.data_type == "e" or text(value).upper() in ERROR_VALUES:
                        cached_errors += 1
                        if cached_errors <= 20:
                            errors.append(f"缓存Excel错误: {worksheet.title}!{cell.coordinate}={value}")
    finally:
        if cached is not None:
            cached.close()

    source_count, source_hash, expected_source_hash = validate_source_snapshot(
        config_root, schema_report_path, errors
    )

    try:
        completed = subprocess.run(
            ["svn", "status", str(config_root)],
            check=False,
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
        )
        svn_status = [line for line in completed.stdout.splitlines() if line.strip()]
    except Exception as exc:
        svn_status = [f"<svn status unavailable: {type(exc).__name__}: {exc}>"]

    status_counter = Counter(text(row[24]) for row in batch_rows)
    mode_counter = Counter(text(row[21]) for row in batch_rows)
    table_counter = Counter(text(row[7]) for row in batch_rows)
    summary = {
        "batch_id": draft_generator.BATCH_ID,
        "batch_version": draft_generator.BATCH_VERSION,
        "workbook": str(workbook_path),
        "workbook_sha256": sha256_file(workbook_path),
        "draft_row_count": len(batch_rows),
        "draft_id_unique_count": len({text(row[1]) for row in batch_rows}),
        "counts_by_status": dict(sorted(status_counter.items())),
        "counts_by_write_mode": dict(sorted(mode_counter.items())),
        "counts_by_logical_table": dict(sorted(table_counter.items())),
        "module_mapping_count": len(module_paths),
        "formula_count": formula_count,
        "formula_error_count": formula_errors,
        "cached_error_count": cached_errors,
        "external_link_count": external_links,
        "registered_source_file_count": source_count,
        "source_snapshot_hash": source_hash,
        "expected_source_snapshot_hash": expected_source_hash,
        "svn_status": svn_status,
        "approval_row_count": 0,
        "receipt_row_count": 0,
        "official_config_write": False,
        "luban_executed": False,
        "validation_status": "PASS" if not errors else "FAIL",
    }
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    return errors, summary


def main() -> int:
    parser = argparse.ArgumentParser(description="验证映射装备配置草案D01")
    parser.add_argument("--workbook", default=str(draft_generator.DEFAULT_WORKBOOK))
    parser.add_argument("--config-root", default=str(draft_generator.DEFAULT_CONFIG_ROOT))
    parser.add_argument("--schema-report", default=str(DEFAULT_SCHEMA_REPORT))
    parser.add_argument("--output", default=str(DEFAULT_OUTPUT))
    parser.add_argument("--expected-formulas", type=int, default=EXPECTED_FORMULAS)
    args = parser.parse_args()

    errors, summary = validate(args)
    if errors:
        print("[FAIL] 映射装备配置草案验证失败", file=sys.stderr)
        for error in errors:
            print(f"- {error}", file=sys.stderr)
        print(json.dumps(summary, ensure_ascii=False, indent=2))
        return 1
    print("[OK] 映射装备配置草案验证通过；正式Datas快照未变化")
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
