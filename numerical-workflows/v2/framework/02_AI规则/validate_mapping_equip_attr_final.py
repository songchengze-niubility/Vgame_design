#!/usr/bin/env python3
"""只读校验映射装备属性最终配置镜像，不检查掉落、解锁等非属性表。"""
from __future__ import annotations

import json
import os
from pathlib import Path

from openpyxl import load_workbook
from mapping_rare_rules import EFFECT_IDS as RARE_EFFECT_IDS, POOL_IDS as RARE_POOL_IDS, rare_pool_id
from mapping_named_rules import EFFECT_IDS as NAMED_EFFECT_IDS, POOL_IDS as NAMED_POOL_IDS, named_effect_id

ROOT = Path(os.environ.get("VGAME_NUMERICAL_RUNTIME_ROOT", str(Path(__file__).resolve().parents[1])))
OUT = Path(os.environ.get("MAPPING_EQUIP_FINAL_OUT", str(ROOT / "06_最终配置" / "映射装备")))
CONFIG = Path(os.environ.get("VGAME_CONFIG_DATAS", ""))
REPORT = Path(os.environ.get("MAPPING_EQUIP_ATTR_REPORT", str(ROOT / "05_验算核算" / "mapping-equip-attr-final-validation.json")))

FILES = {
    "SpecialEquipment.xlsx": ("SpecialEquipment", "SpecialEquipment/SpecialEquipment.xlsx"),
    "SpecialEquipmentMain.xlsx": ("SpecialEquipment", "SpecialEquipment/SpecialEquipmentMain.xlsx"),
    "SpecialEquipmentEffect.xlsx": ("SpecialEquipment", "SpecialEquipment/SpecialEquipmentEffect.xlsx"),
    "SpecialEquipmentEffectDrop.xlsx": ("SpecialEquipment", "SpecialEquipment/SpecialEquipmentEffectDrop.xlsx"),
    "SpecialEquipmentQuality.xlsx": ("SpecialEquipment", "SpecialEquipment/SpecialEquipmentQuality.xlsx"),
    "SpecialEquipmentSubAttrCoef.xlsx": ("SpecialEquipmentSubAttrCoef", "SpecialEquipment/SpecialEquipmentSubAttrCoef.xlsx"),
}
MAIN_TARGET = {
    "SEPart_Weapon": ("Attack", 11250), "SEPart_Head": ("Hp", 18750),
    "SEPart_Armour": ("Hp", 18750), "SEPart_Glove": ("Armor", 2500),
    "SEPart_Belt": ("Armor", 2500), "SEPart_Shoe": ("Armor", 2500),
    "SEPart_Device": ("CritDmgBonusLevel", 8800), "SEPart_Decoration": ("CritLevel", 7200),
}
PERCENT_SUB_ATTRS = {"AttackBonus", "ArmorBonus", "HpBonus"}
FIXED_SUB_ATTRS = {"Attack", "Armor", "Hp", "CritLevel", "CritDmgBonusLevel", "EnergyPreAddBonusLevel", "CoolDownReduceLevel"}
SUB_CURRENT_CAP = {
    "AttackBonus": .329, "ArmorBonus": .329, "HpBonus": .329,
    "Attack": 3701, "Armor": 2468, "Hp": 12338,
    "CritLevel": 2961, "CritDmgBonusLevel": 3619,
    "EnergyPreAddBonusLevel": 2961, "CoolDownReduceLevel": 4277,
}
SUB_STANDARD = {attr: cap * .50 for attr, cap in SUB_CURRENT_CAP.items()}
SUB_RANDOM_MULTIPLIER = {
    "Green": (1.00, 1.00), "Blue": (1.00, 1.00), "Purple": (1.00, 1.00),
    "Orange": (1.00, 1.20), "Red": (1.20, 1.40),
}
SUB_RANDOM_STEP = {attr: (.001 if attr in PERCENT_SUB_ATTRS else 1) for attr in SUB_CURRENT_CAP}
SUB_PERCENT_LEVEL_COEF = {level: .45 + .55 * ((level - 5) / 55) for level in range(5, 61, 5)}
SUB_FIXED_LEVEL_COEF = {level: .03 + .97 * (((level - 5) / 55) ** 1.35) for level in range(5, 61, 5)}
QUALITY_LEVEL_COMBOS = tuple(
    (quality, level)
    for quality in ("Green", "Blue", "Purple", "Orange", "Red")
    for level in range(5, 61, 5)
)
SUB_POOL_ID_BY_COMBO = {combo: 5001 + index for index, combo in enumerate(QUALITY_LEVEL_COMBOS)}
SUB_EFFECT_IDS_BY_COMBO = {
    combo: tuple(50001 + index * 10 + slot for slot in range(10))
    for index, combo in enumerate(QUALITY_LEVEL_COMBOS)
}
NEW_SUB_POOL_IDS = set(SUB_POOL_ID_BY_COMBO.values())
NEW_SUB_EFFECT_IDS = {effect_id for ids in SUB_EFFECT_IDS_BY_COMBO.values() for effect_id in ids}
SUB_ATTR_ORDER = (
    "AttackBonus", "ArmorBonus", "HpBonus", "CoolDownReduceLevel", "EnergyPreAddBonusLevel",
    "CritLevel", "CritDmgBonusLevel", "Attack", "Armor", "Hp",
)
QUALITY = {
    "Green": (2, 2, .20, .35), "Blue": (2, 3, .30, .50),
    "Purple": (4, 4, .45, .65), "Orange": (4, 4, .70, .85),
    "Red": (4, 4, 1.00, 1.00),
}
SUB_LEVEL = {level: 1.0 for level in range(5, 61, 5)}


def headers(ws):
    return {value: index for index, value in enumerate(next(ws.iter_rows(min_row=1, max_row=1, values_only=True))) if value is not None}


def rows(ws, id_field="Id"):
    cols = headers(ws)
    return cols, [row for row in ws.iter_rows(min_row=5, max_col=max(cols.values())+1, values_only=True) if row[cols[id_field]] is not None]


def schema(ws):
    top = list(ws.iter_rows(min_row=1, max_row=4, max_col=min(ws.max_column or 100, 100), values_only=True))
    width = max(c+1 for row in top for c,value in enumerate(row) if value is not None)
    return [list(row[:width]) for row in top]


def level_coef(attr, level):
    return SUB_PERCENT_LEVEL_COEF[level] if attr in PERCENT_SUB_ATTRS else SUB_FIXED_LEVEL_COEF[level]


def expected_effect_range(attr, quality, level):
    random_min, random_max = SUB_RANDOM_MULTIPLIER[quality]
    digits = 9 if attr in PERCENT_SUB_ATTRS else 6
    base = SUB_STANDARD[attr] * level_coef(attr, level) * QUALITY[quality][3]
    return round(base * random_max, digits), round(base * random_min, digits), SUB_RANDOM_STEP[attr]


def close_tuple(actual, expected, tolerance=1e-8):
    return all(abs(float(a) - float(e)) <= tolerance for a, e in zip(actual, expected))


def main():
    checks = []

    def check(name, ok, actual=None, expected=True):
        checks.append({"name": name, "result": "PASS" if ok else "FAIL", "actual": actual, "expected": expected})

    opened = {}
    for file_name, (sheet_name, formal_rel) in FILES.items():
        path = OUT / file_name
        check(f"存在:{file_name}", path.exists(), path.exists(), True)
        if not path.exists():
            continue
        wb = load_workbook(path, data_only=False, read_only=True)
        formal_wb = load_workbook(CONFIG / formal_rel, data_only=False, read_only=True)
        ws, formal_ws = wb[sheet_name], formal_wb[sheet_name]
        check(f"正式Schema:{file_name}", schema(ws) == schema(formal_ws), schema(ws), schema(formal_ws))
        formulas = errors = 0
        for row in ws.iter_rows(values_only=True):
            for value in row:
                formulas += isinstance(value, str) and value.startswith("=")
                errors += isinstance(value, str) and value.startswith(("#REF!", "#DIV/0!", "#VALUE!", "#N/A", "#NAME?"))
        check(f"公式与错误值:{file_name}", formulas == 0 and errors == 0, {"formulas": formulas, "errors": errors}, {"formulas": 0, "errors": 0})
        opened[file_name] = (wb, ws, formal_wb, formal_ws)

    if len(opened) != len(FILES):
        result = {"status": "FAIL", "checks": len(checks), "pass": sum(c["result"] == "PASS" for c in checks), "fail": sum(c["result"] == "FAIL" for c in checks), "details": checks}
        REPORT.parent.mkdir(parents=True, exist_ok=True)
        REPORT.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
        print(json.dumps({k: result[k] for k in ("status", "checks", "pass", "fail")}, ensure_ascii=False))
        raise SystemExit(1)

    _, equip_ws, _, formal_equip_ws = opened["SpecialEquipment.xlsx"]
    ec, equips = rows(equip_ws); fec, formal_equip_rows = rows(formal_equip_ws)
    formal_equips = {row[fec["Id"]]: row for row in formal_equip_rows}
    preserve = [field for field in fec if field not in {"Attribute2List", "Effect", "RareEffect", "SpecialEffect", "FixEffect"}]
    equip_errors = []; combos = set()
    for row in equips:
        equip_id = row[ec["Id"]]; formal = formal_equips.get(equip_id)
        combo = (row[ec["Quality"]], row[ec["Grade"]]); combos.add(combo)
        expected_pool = SUB_POOL_ID_BY_COMBO.get(combo)
        expected_rare = rare_pool_id(*combo)
        expected_special = named_effect_id(combo[0], row[ec["Name"]])
        expected_fixed = formal[fec["FixEffect"]] if formal else None
        if expected_fixed in NAMED_EFFECT_IDS: expected_fixed = None
        if (formal is None or row[ec["Attribute2List"]] is not None or row[ec["Effect"]] != expected_pool
                or row[ec["RareEffect"]] != expected_rare
                or row[ec["SpecialEffect"]] != expected_special
                or row[ec["FixEffect"]] != expected_fixed
                or any(row[ec[field]] != formal[fec[field]] for field in preserve)):
            equip_errors.append((equip_id, combo, row[ec["Effect"]], expected_pool))
    expected_equip_count = len(formal_equip_rows)
    check(f"{expected_equip_count}件装备主属性2、普通副属性池、金红稀有池及红装具名绑定", len(equips) == expected_equip_count and not equip_errors, {"rows": len(equips), "errors": equip_errors}, {"rows": expected_equip_count, "errors": []})
    check("装备品质等级组合60组", combos == set(QUALITY_LEVEL_COMBOS), sorted(combos), list(QUALITY_LEVEL_COMBOS))

    _, main_ws, _, formal_main_ws = opened["SpecialEquipmentMain.xlsx"]
    mc, main_rows = rows(main_ws); main_by_id = {row[mc["Id"]]: row for row in main_rows}
    fmc, formal_main_rows = rows(formal_main_ws)
    formal_main_ids = {row[fmc["Id"]] for row in formal_main_rows}
    main_errors = []; lv5_errors = []; lv60_errors = []
    for equip in equips:
        main_id = int(str(equip[ec["Attribute1List"]]).split(",")[0]); row = main_by_id.get(main_id)
        attr, target = MAIN_TARGET[equip[ec["Part"]]]
        if row is None or row[mc["Attribute"]] != attr:
            main_errors.append((equip[ec["Id"]], main_id)); continue
        value, grow, param = row[mc["Value"]], row[mc["Grow"]], row[mc["Param"]]
        lv5_errors.append(abs(value * grow + param - target * .05)); lv60_errors.append(abs(value * grow ** 12 + param - target))
    main_ids_match = set(main_by_id) == formal_main_ids
    check(f"{len(formal_main_rows)}条Main兼容定义及在用部位映射", len(main_rows) == len(formal_main_rows) and main_ids_match and not main_errors, {"rows": len(main_rows), "id_set_match": main_ids_match, "errors": main_errors}, {"rows": len(formal_main_rows), "id_set_match": True, "errors": []})
    check("主属性Lv5与Lv60锚点", max(lv5_errors, default=999) < 1e-5 and max(lv60_errors, default=999) < 1e-5, {"lv5_max_error": max(lv5_errors, default=None), "lv60_max_error": max(lv60_errors, default=None)}, "均<1e-5")

    _, effect_ws, _, formal_effect_ws = opened["SpecialEquipmentEffect.xlsx"]
    fc, effect_rows = rows(effect_ws); ffc, formal_effect_rows = rows(formal_effect_ws)
    effect_by_id = {row[fc["Id"]]: row for row in effect_rows}
    formal_effect_by_id = {row[ffc["Id"]]: row for row in formal_effect_rows if row[ffc["Id"]] not in NEW_SUB_EFFECT_IDS}
    legacy_errors = []
    for effect_id, formal_row in formal_effect_by_id.items():
        output_row = effect_by_id.get(effect_id)
        if output_row is None or any(output_row[fc[field]] != formal_row[ffc[field]] for field in ffc):
            legacy_errors.append(effect_id)
    check(f"{len(formal_effect_by_id)}条正式历史Effect逐行保留", not legacy_errors, legacy_errors, [])
    new_effect_errors = []
    for combo in QUALITY_LEVEL_COMBOS:
        quality, level = combo
        for attr, effect_id in zip(SUB_ATTR_ORDER, SUB_EFFECT_IDS_BY_COMBO[combo]):
            row = effect_by_id.get(effect_id)
            if row is None:
                new_effect_errors.append((effect_id, "missing")); continue
            actual = (row[fc["ValueMax"]], row[fc["ValueMin"]], row[fc["RandomMin"]]); expected = expected_effect_range(attr, quality, level)
            if row[fc["Attribute"]] != attr or not close_tuple(actual, expected):
                new_effect_errors.append((effect_id, combo, attr, row[fc["Attribute"]], actual, expected))
    expected_effect_rows = len(formal_effect_by_id) + 600 + len((RARE_EFFECT_IDS | NAMED_EFFECT_IDS)-set(formal_effect_by_id))
    check("600条双曲线新Effect值", len(effect_rows) == expected_effect_rows and not new_effect_errors, {"rows": len(effect_rows), "errors": new_effect_errors}, {"rows": expected_effect_rows, "errors": []})
    fixed_random_errors = []
    for combo in QUALITY_LEVEL_COMBOS:
        quality, _ = combo
        for effect_id in SUB_EFFECT_IDS_BY_COMBO[combo]:
            row = effect_by_id[effect_id]; value_max, value_min = row[fc["ValueMax"]], row[fc["ValueMin"]]
            should_fixed = quality in {"Green", "Blue", "Purple"}
            if should_fixed != (abs(float(value_max) - float(value_min)) <= 1e-9):
                fixed_random_errors.append((effect_id, quality, value_min, value_max))
    check("绿蓝紫定值且金红随机", not fixed_random_errors, fixed_random_errors, [])

    _, pool_ws, _, formal_pool_ws = opened["SpecialEquipmentEffectDrop.xlsx"]
    pc, pool_rows = rows(pool_ws); fpc, formal_pool_rows = rows(formal_pool_ws)
    formal_legacy_pool_rows = [row for row in formal_pool_rows if row[fpc["Id"]] not in NEW_SUB_POOL_IDS]
    pools = {row[pc["Id"]]: row[pc["EffectList"]] for row in pool_rows}
    expected_pool_rows = len(formal_legacy_pool_rows) + 60 + len((RARE_POOL_IDS | NAMED_POOL_IDS)-{r[fpc["Id"]] for r in formal_legacy_pool_rows})
    check("新增60个品质等级池", len(pool_rows) == expected_pool_rows and NEW_SUB_POOL_IDS.issubset(pools), len(pool_rows), expected_pool_rows)
    check("72个稀有Effect绑定同号Skill", all(i in effect_by_id and str(effect_by_id[i][fc["SkillList"]]) == str(i) for i in RARE_EFFECT_IDS))
    check("21个具名Effect绑定同号Skill", all(i in effect_by_id and str(effect_by_id[i][fc["SkillList"]]) == str(i) for i in NAMED_EFFECT_IDS))
    check("21个具名池各含唯一同号Effect", all(pools.get(i) == f"{i},100" for i in NAMED_POOL_IDS))
    check("绿蓝紫没有RareEffect", all(r[ec["RareEffect"]] is None for r in equips if r[ec["Quality"]] in {"Green","Blue","Purple"}))
    check("非红没有SpecialEffect", all(r[ec["SpecialEffect"]] is None for r in equips if r[ec["Quality"]] != "Red"))
    check("撤回具名FixEffect", all(r[ec["FixEffect"]] not in NAMED_EFFECT_IDS for r in equips))
    check("6个稀有池仅含同档12条等权技能", all(pools.get(6000+s) == ";".join(f"{base*100+s},100" for base in range(311,323)) for s in range(1,7)))
    active_pools = {row[ec["Effect"]] for row in equips}; pool_errors = []
    for pool_id in sorted(active_pools):
        ids = [int(token.split(",")[0]) for token in str(pools.get(pool_id, "")).split(";") if token]
        attrs = [effect_by_id[i][fc["Attribute"]] for i in ids if i in effect_by_id]
        if len(ids) != 10 or set(attrs) != set(SUB_ATTR_ORDER) or len(attrs) != len(set(attrs)):
            pool_errors.append((pool_id, ids, attrs))
    check("全部在用池均为十类无重复词条", not pool_errors, pool_errors, [])
    combo_pool_errors = []
    for combo, pool_id in SUB_POOL_ID_BY_COMBO.items():
        actual = tuple(int(token.split(",")[0]) for token in str(pools.get(pool_id, "")).split(";") if token)
        if actual != SUB_EFFECT_IDS_BY_COMBO[combo]: combo_pool_errors.append((combo, pool_id, actual))
    check("60个品质等级池引用正确Effect", not combo_pool_errors, combo_pool_errors, [])
    legacy_active = active_pools.intersection(set(range(1, 38)) | {101, 201, 301})
    check("旧池退出新装备引用", not legacy_active, sorted(legacy_active), [])

    _, quality_ws, _, _ = opened["SpecialEquipmentQuality.xlsx"]
    qc, quality_rows = rows(quality_ws)
    quality_actual = {row[qc["Quality"]]: (row[qc["EffectNumMin"]], row[qc["EffectNumMax"]], row[qc["GradeCoefficient"]], row[qc["QualityCoefficient"]]) for row in quality_rows}
    check("五品质系数与副属性条数", quality_actual == QUALITY, quality_actual, QUALITY)
    green_attack = effect_by_id[SUB_EFFECT_IDS_BY_COMBO[("Green", 5)][0]][fc["ValueMax"]]
    red_attack = effect_by_id[SUB_EFFECT_IDS_BY_COMBO[("Red", 5)][0]][fc["ValueMax"]]
    red_green_ratio = red_attack / green_attack
    check("红色上限为绿色定值4倍", abs(red_green_ratio - 4.0) < 1e-9, red_green_ratio, 4.0)

    _, sub_ws, _, _ = opened["SpecialEquipmentSubAttrCoef.xlsx"]
    sc, sub_rows = rows(sub_ws); sub_actual = {row[sc["EquipLv"]]: row[sc["SubAttrCoef"]] for row in sub_rows}
    check("SubAttrCoef统一为1避免双重缩放", sub_actual == SUB_LEVEL, sub_actual, SUB_LEVEL)
    check("百分比曲线Lv5/Lv60锚点", abs(SUB_PERCENT_LEVEL_COEF[5] - .45) < 1e-12 and abs(SUB_PERCENT_LEVEL_COEF[60] - 1) < 1e-12, {5: SUB_PERCENT_LEVEL_COEF[5], 60: SUB_PERCENT_LEVEL_COEF[60]}, {5: .45, 60: 1.0})
    check("固定值曲线Lv5/Lv60锚点", abs(SUB_FIXED_LEVEL_COEF[5] - .03) < 1e-12 and abs(SUB_FIXED_LEVEL_COEF[60] - 1) < 1e-12, {5: SUB_FIXED_LEVEL_COEF[5], 60: SUB_FIXED_LEVEL_COEF[60]}, {5: .03, 60: 1.0})

    for wb, _, formal_wb, _ in opened.values(): wb.close(); formal_wb.close()
    failed = [item for item in checks if item["result"] == "FAIL"]
    result = {"status": "PASS" if not failed else "FAIL", "checks": len(checks), "pass": len(checks) - len(failed), "fail": len(failed), "output": str(OUT), "details": checks}
    REPORT.parent.mkdir(parents=True, exist_ok=True)
    REPORT.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({k: result[k] for k in ("status", "checks", "pass", "fail")}, ensure_ascii=False))
    if failed: raise SystemExit(1)


if __name__ == "__main__":
    main()
