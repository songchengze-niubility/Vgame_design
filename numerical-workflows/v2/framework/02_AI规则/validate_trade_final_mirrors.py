from __future__ import annotations

import os

import argparse
import json
import zipfile
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_CONFIG = Path(os.environ.get("VGAME_CONFIG_DATAS", ""))
DEFAULT_OUT = ROOT / "06_最终配置" / "跑商"


def same_schema(source: Path, mirror: Path, rows: int) -> bool:
    src = load_workbook(source, data_only=False).active
    dst = load_workbook(mirror, data_only=False).active
    # 正式表可能带有仅含样式的尾随空列；Schema 只比较表头区域内实际有值的列。
    last_col = max(
        c
        for r in range(1, rows + 1)
        for c in range(1, src.max_column + 1)
        if src.cell(r, c).value is not None
    )
    for r in range(1, rows + 1):
        for c in range(1, last_col + 1):
            if src.cell(r, c).value != dst.cell(r, c).value:
                return False
    return True


def fields(ws) -> dict[str, int]:
    return {str(ws.cell(1, c).value): c for c in range(1, ws.max_column + 1) if ws.cell(1, c).value is not None}


def column_values(ws, field: str, start: int = 5) -> list:
    col = fields(ws)[field]
    return [ws.cell(r, col).value for r in range(start, ws.max_row + 1)]


def same_data_rows(source: Path, mirror: Path, key: str) -> bool:
    src = load_workbook(source, data_only=False).active
    dst = load_workbook(mirror, data_only=False).active
    src_fields = fields(src)
    dst_fields = fields(dst)
    names = [name for name in src_fields if name not in {"##var"}]
    if any(name not in dst_fields for name in names):
        return False
    src_rows = [
        tuple(src.cell(r, src_fields[name]).value for name in names)
        for r in range(5, src.max_row + 1)
        if src.cell(r, src_fields[key]).value is not None
    ]
    dst_rows = [
        tuple(dst.cell(r, dst_fields[name]).value for name in names)
        for r in range(5, dst.max_row + 1)
        if dst.cell(r, dst_fields[key]).value is not None
    ]
    return src_rows == dst_rows


def workbook_safety(path: Path) -> dict:
    wb = load_workbook(path, data_only=False)
    formulas = []
    formula_errors = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.data_type == "f" or (isinstance(cell.value, str) and cell.value.startswith("=")):
                    formulas.append(f"{ws.title}!{cell.coordinate}")
                if cell.data_type == "e":
                    formula_errors.append(f"{ws.title}!{cell.coordinate}:{cell.value}")
    with zipfile.ZipFile(path) as zf:
        external_links = [name for name in zf.namelist() if name.startswith("xl/externalLinks/")]
    return {
        "formulas": formulas,
        "formula_errors": formula_errors,
        "external_links": external_links,
        "sheets": wb.sheetnames,
    }


def pending_blank_with_comment(ws, row: int, field: str) -> bool:
    cell = ws.cell(row, fields(ws)[field])
    return cell.value in (None, "") and cell.comment is not None and bool(cell.comment.text.strip())


def validate(config: Path, out: Path) -> dict:
    expected = [
        "00_配置清单.xlsx",
        "globals_conf.xlsx",
        "TradeGoodsMarket.xlsx",
        "TradeLvUp.xlsx",
        "TradeLocation.xlsx",
        "TradeEvent.xlsx",
        "Drop_other.xlsx",
        "Unlock.xlsx",
        "Items.xlsx",
        "Shop.xlsx",
        "ShopGoods.xlsx",
    ]
    checks: list[dict] = []

    def check(name: str, ok: bool, detail: str = "") -> None:
        checks.append({"name": name, "ok": bool(ok), "detail": detail})

    actual = sorted(p.name for p in out.glob("*.xlsx")) if out.exists() else []
    check("文件清单", sorted(expected) == actual, f"actual={actual}")

    for name in expected:
        path = out / name
        if not path.exists():
            continue
        safety = workbook_safety(path)
        expected_sheets = 1 if name == "00_配置清单.xlsx" else 2
        check(f"{name}:工作表结构", len(safety["sheets"]) == expected_sheets and (expected_sheets == 1 or safety["sheets"][1] == "中文说明"), str(safety["sheets"]))
        check(f"{name}:无公式", not safety["formulas"], str(safety["formulas"][:5]))
        check(f"{name}:无公式错误", not safety["formula_errors"], str(safety["formula_errors"][:5]))
        check(f"{name}:无外链", not safety["external_links"], str(safety["external_links"]))

    schema_cases = [
        (config / "globals_conf" / "globals_conf.xlsx", out / "globals_conf.xlsx", 3),
        (config / "Trade" / "TradeGoodsMarket.xlsx", out / "TradeGoodsMarket.xlsx", 4),
        (config / "Trade" / "TradeLvUp.xlsx", out / "TradeLvUp.xlsx", 4),
        (config / "Trade" / "TradeLocation.xlsx", out / "TradeLocation.xlsx", 4),
        (config / "Trade" / "TradeEvent.xlsx", out / "TradeEvent.xlsx", 4),
        (config / "Drop" / "Drop_other.xlsx", out / "Drop_other.xlsx", 4),
        (config / "unlock" / "Unlock.xlsx", out / "Unlock.xlsx", 4),
        (config / "items" / "Items.xlsx", out / "Items.xlsx", 4),
        (config / "Shop" / "Shop.xlsx", out / "Shop.xlsx", 4),
        (config / "Shop" / "ShopGoods.xlsx", out / "ShopGoods.xlsx", 4),
    ]
    for source, mirror, rows in schema_cases:
        if source.exists() and mirror.exists():
            check(f"{mirror.name}:Schema一致", same_schema(source, mirror, rows), str(source))

    globals_ws = load_workbook(out / "globals_conf.xlsx", data_only=False).active
    globals_values = {globals_ws.cell(r, 1).value: globals_ws.cell(r, 4).value for r in range(4, globals_ws.max_row + 1)}
    confirmed_globals = {
        "TradeEventExistHours": 18,
        "TradeEventMaxExistCount": 2,
        "TradeRunEventRefreshMin": 1,
        "TradeRunEventRefreshMax": 1,
        "TradeDefaultTicketItemId": 301,
        "TradeDefaultTradeLevel": 1,
        "TradeDefaultTicketCount": 30000,
        "TradeRecycleParam": 0.7,
        "DefaultCargoCapacity": 30,
        "TradeTransportHours": 20,
    }
    check("Globals确认值", all(globals_values.get(k) == v for k, v in confirmed_globals.items()), str({k: globals_values.get(k) for k in confirmed_globals}))
    check("Globals全部确认", all(globals_values.get(k) is not None for k in confirmed_globals), str(confirmed_globals))

    goods_ws = load_workbook(out / "TradeGoodsMarket.xlsx", data_only=False).active
    check("货物基础价", column_values(goods_ws, "BasicPrice") == [900, 1000, 1200, 1300, 1500, 1700, 1600, 2000, 2100, 1100, 1900, 2300, 1400, 1800, 2100], str(column_values(goods_ws, "BasicPrice")))
    rate_expected = {
        "BuyRateMin": 0.80, "BuyRateMax": 1.05,
        "NormalSellRateMin": 0.90, "NormalSellRateMax": 1.20,
        "UrgentSellRateMin": 1.60, "UrgentSellRateMax": 2.00,
    }
    check("货物价格倍率", all(set(column_values(goods_ws, field)) == {value} for field, value in rate_expected.items()), str(rate_expected))
    check("货物基础每日限购", column_values(goods_ws, "DailyBuyLimit") == [25, 20, 15] * 5, str(column_values(goods_ws, "DailyBuyLimit")))
    goods_ids = list(range(910001, 910016))
    check("货物预留 ID", column_values(goods_ws, "Id") == goods_ids, str(column_values(goods_ws, "Id")))
    check("货物跑商等级解锁", column_values(goods_ws, "UnlockId") == [910101] * 3 + [910102] * 3 + [910103] * 3 + [910104] * 3 + [910105] * 3, str(column_values(goods_ws, "UnlockId")))

    level_ws = load_workbook(out / "TradeLvUp.xlsx", data_only=False).active
    thresholds = [0, 15000, 45000, 75000, 105000, 185000, 270000, 355000, 445000, 590000, 740000, 995000, 1290000, 1685000, 2130000, 2620000, 3190000, 3800000, 4455000, 5115000]
    stage_exp = [thresholds[i] - thresholds[i - 1] for i in range(1, len(thresholds))]
    tickets = [13000, 13000, 13000, 16000, 16000, 16000, 20000, 20000, 20000, 20000, 20000, 20000, 25000, 25000, 25000, 25000, 32000, 32000, 40000, 40000]
    check("等级累计阈值", column_values(level_ws, "CostExp") == thresholds, str(column_values(level_ws, "CostExp")))
    check("单级经验非递减", all(stage_exp[i] >= stage_exp[i - 1] for i in range(1, len(stage_exp))), str(stage_exp))
    check("每日商券", column_values(level_ws, "DailyTicketCount") == tickets, str(column_values(level_ws, "DailyTicketCount")))
    check("等级预留 ID", column_values(level_ws, "Id") == list(range(940001, 940021)), str(column_values(level_ws, "Id")))
    level_drop_ids = [None] + list(range(69090101, 69090120))
    check("等级奖励 DropId", column_values(level_ws, "DropId") == level_drop_ids, str(column_values(level_ws, "DropId")))
    check("Lv.1 初始商券不重复配置奖励包", pending_blank_with_comment(level_ws, 5, "DropId"))

    location_ws = load_workbook(out / "TradeLocation.xlsx", data_only=False).active
    check("站点预留 ID", column_values(location_ws, "Id") == list(range(920001, 920006)), str(column_values(location_ws, "Id")))
    check("站点急需货物引用", column_values(location_ws, "DemandGoodsList") == [
        "910001;910002;910003", "910004;910005;910006", "910007;910008;910009", "910010;910011;910012", "910013;910014;910015",
    ], str(column_values(location_ws, "DemandGoodsList")))
    event_pools = ["1,250;2,250;3,250;4,150;5,100", "6,250;7,250;8,250;9,150;10,100", "11,250;12,250;13,250;14,150;15,100", "16,250;17,250;18,250;19,150;20,100", "21,250;22,250;23,250;24,150;25,100"]
    check("站点 V22 事件概率池", column_values(location_ws, "TradeEventList") == event_pools, str(column_values(location_ws, "TradeEventList")))
    check("站点跑商等级解锁", column_values(location_ws, "UnlockId") == [910101, 910102, 910103, 910104, 910105], str(column_values(location_ws, "UnlockId")))

    event_ws = load_workbook(out / "TradeEvent.xlsx", data_only=False).active
    check("正式事件 ID", column_values(event_ws, "Id") == list(range(1, 26)), str(column_values(event_ws, "Id")))
    check("正式事件类型", set(column_values(event_ws, "TradeEventType")) == {"Reward"}, str(column_values(event_ws, "TradeEventType")))
    check("正式事件选项1完整", all(column_values(event_ws, "Option1Title")), str(column_values(event_ws, "Option1Title")))
    check("正式事件选项2完整", all(column_values(event_ws, "Option2Title")), str(column_values(event_ws, "Option2Title")))
    check("正式事件 DropId1", column_values(event_ws, "DropId1") == list(range(69090201, 69090250, 2)), str(column_values(event_ws, "DropId1")))
    check("正式事件 DropId2", column_values(event_ws, "DropId2") == list(range(69090202, 69090251, 2)), str(column_values(event_ws, "DropId2")))
    expected_event_unlocks = [None, None, 11001, None, None, None, 11001, None, None, None, 11001, None, None, None, None, None, 11001, None, None, None, 11001, None, None, None, 11001]
    check("正式事件 UnlockId", column_values(event_ws, "UnlockId") == expected_event_unlocks, str(column_values(event_ws, "UnlockId")))

    drop_ws = load_workbook(out / "Drop_other.xlsx", data_only=False).active
    level_drop_ids = list(range(69090101, 69090120))
    all_drop_ids = column_values(drop_ws, "DropId")
    all_drop_row_ids = column_values(drop_ws, "ID")
    check("跑商等级奖励包 DropId", all_drop_ids[:19] == level_drop_ids, str(all_drop_ids[:19]))
    check("跑商等级奖励包行 ID", all_drop_row_ids[:19] == list(range(950001, 950020)), str(all_drop_row_ids[:19]))
    check("跑商事件奖励包行 ID", all_drop_row_ids[19:] == list(range(950101, 950152)), str(all_drop_row_ids[19:]))
    expected_event_drop_ids = list(range(69090201, 69090250)) + [69090250, 69090250]
    check("跑商事件奖励包 DropId", all_drop_ids[19:] == expected_event_drop_ids, str(all_drop_ids[19:]))
    level_ticket_expected = {
        69090101: 3000, 69090102: 4000, 69090105: 5000, 69090108: 6000,
        69090110: 7000, 69090112: 8000, 69090115: 9000, 69090118: 10000,
    }
    level_ticket_actual = {
        drop_ws.cell(row, fields(drop_ws)["DropId"]).value: drop_ws.cell(row, fields(drop_ws)["DropItemNumMin"]).value
        for row in range(5, drop_ws.max_row + 1)
        if drop_ws.cell(row, fields(drop_ws)["DropItem_id"]).value == 301
        and drop_ws.cell(row, fields(drop_ws)["DropId"]).value in level_ticket_expected
    }
    check("等级商券奖励", level_ticket_actual == level_ticket_expected, str(level_ticket_actual))
    check("旧事件奖励包已移除", not any(drop_id in set(range(69090001, 69090007)) for drop_id in column_values(drop_ws, "DropId")))
    drop_fields = fields(drop_ws)
    event_drop_actual: dict[int, list[tuple]] = {}
    for row in range(24, drop_ws.max_row + 1):
        drop_id = drop_ws.cell(row, drop_fields["DropId"]).value
        event_drop_actual.setdefault(drop_id, []).append((
            drop_ws.cell(row, drop_fields["DropItem_id"]).value,
            drop_ws.cell(row, drop_fields["DropItemNumMin"]).value,
            drop_ws.cell(row, drop_fields["DropWeight"]).value,
            drop_ws.cell(row, drop_fields["DropType"]).value,
        ))
    reward1_qty = [6500, 7000, 7500, 8500, 9000, 7000, 7500, 8000, 9000, 9500, 7500, 8000, 8500, 9500, 10000, 8000, 8500, 9000, 10000, 10500, 8500, 9000, 9500, 10500, 11000]
    reward2 = [(1, 3000), (100, 4), (1, 4000), (100, 8), (2, 10), (1, 5000), (100, 6), (101, 1), (101, 2), (2, 15), (1, 8000), (101, 2), (1, 10000), (102, 1), (2, 20), (1, 12000), (101, 3), (1, 15000), (102, 1), (2, 25), (1, 18000), (101, 5), (1, 20000), (102, 2)]
    expected_event_drop = {}
    for event_id, qty in enumerate(reward1_qty, 1):
        expected_event_drop[69090200 + (event_id - 1) * 2 + 1] = [(301, qty, 10000, "fixed_")]
        if event_id < 25:
            item_id, item_count = reward2[event_id - 1]
            expected_event_drop[69090200 + (event_id - 1) * 2 + 2] = [(item_id, item_count, 10000, "fixed_")]
    expected_event_drop[69090250] = [(2, 30, 9000, "random"), (106, 1, 1000, "random")]
    check("事件奖励内容", event_drop_actual == expected_event_drop, f"packages={len(event_drop_actual)}")
    check("事件奖励道具引用", {item for rows in event_drop_actual.values() for item, _, _, _ in rows} <= {1, 2, 100, 101, 102, 106, 301}, str(event_drop_actual))

    unlock_ws = load_workbook(out / "Unlock.xlsx", data_only=False).active
    unlock_ids = [910101, 910102, 910106, 910103, 910107, 910104, 910108, 910105, 910109, 910110, 910111, 910112, 910113]
    unlock_levels = [1, 2, 3, 5, 6, 8, 9, 10, 14, 15, 17, 18, 20]
    check("跑商解锁预留 ID", column_values(unlock_ws, "Id") == unlock_ids, str(column_values(unlock_ws, "Id")))
    check("跑商解锁类型", set(column_values(unlock_ws, "UnLockType")) == {"UnLock_UUT_Trade_Grade"}, str(column_values(unlock_ws, "UnLockType")))
    check("跑商解锁等级", column_values(unlock_ws, "UnLockParams") == unlock_levels, str(column_values(unlock_ws, "UnLockParams")))

    items_ws = load_workbook(out / "Items.xlsx", data_only=False).active
    check("便签正式 ItemId", column_values(items_ws, "Id") == [1075, 1076, 1077, 1078, 1079], str(column_values(items_ws, "Id")))
    check("便签正式名称", column_values(items_ws, "Name") == ["行程便签", "安途清单", "探路标记", "见闻手札", "调律灯片"], str(column_values(items_ws, "Name")))
    check("便签道具类型", set(column_values(items_ws, "ItemType")) == {"ItemType_Prop"}, str(column_values(items_ws, "ItemType")))

    shop_ws = load_workbook(out / "Shop.xlsx", data_only=False).active
    check("复用行商商店", column_values(shop_ws, "Id") == [1110] and column_values(shop_ws, "Name") == ["行商商店"], str(column_values(shop_ws, "Id")))
    check("行商商店货币", column_values(shop_ws, "CurrencyShow") == [301], str(column_values(shop_ws, "CurrencyShow")))
    check("行商商店入口解锁", column_values(shop_ws, "Unlock") == [4314], str(column_values(shop_ws, "Unlock")))
    check("行商商店版本", column_values(shop_ws, "version") == [2], str(column_values(shop_ws, "version")))

    shop_goods_ws = load_workbook(out / "ShopGoods.xlsx", data_only=False).active
    expected_shop_item_ids = [1075, 1076, 1077, 1078, 1079, 106, 107, 106, 107, 106, 107, 34, 106, 107, 35, 106, 107, 36]
    expected_shop_prices = ["301,10000"] * 5 + ["301,300000"] * 6 + ["301,100000", "301,300000", "301,300000", "301,150000", "301,300000", "301,300000", "301,300000"]
    expected_limit_types = ["BuyLimit_Week"] * 5 + ["BuyLimit_Ever"] * 6 + ["BuyLimit_Week", "BuyLimit_Ever", "BuyLimit_Ever", "BuyLimit_Week", "BuyLimit_Ever", "BuyLimit_Ever", "BuyLimit_Month"]
    expected_limit_nums = [5] * 5 + [1] * 13
    expected_shop_unlocks = [910101] * 5 + [910102, 910106, 910103, 910107, 910104, 910108, 910105, 910109, 910110, 910110, 910111, 910112, 910113]
    check("商店商品预留 ID", column_values(shop_goods_ws, "Id") == list(range(1110001, 1110019)), str(column_values(shop_goods_ws, "Id")))
    check("商店商品归属", column_values(shop_goods_ws, "ShopId") == [1110] * 18, str(column_values(shop_goods_ws, "ShopId")))
    check("商店商品 ItemId", column_values(shop_goods_ws, "ItemId") == expected_shop_item_ids, str(column_values(shop_goods_ws, "ItemId")))
    check("商店商品价格", column_values(shop_goods_ws, "Price") == expected_shop_prices, str(column_values(shop_goods_ws, "Price")))
    check("商店商品限购类型", column_values(shop_goods_ws, "LimitType") == expected_limit_types, str(column_values(shop_goods_ws, "LimitType")))
    check("商店商品限购数量", column_values(shop_goods_ws, "LimitNum") == expected_limit_nums, str(column_values(shop_goods_ws, "LimitNum")))
    check("商店商品等级解锁", column_values(shop_goods_ws, "Unlock") == expected_shop_unlocks, str(column_values(shop_goods_ws, "Unlock")))

    manifest = load_workbook(out / "00_配置清单.xlsx", data_only=False).active
    manifest_status = [str(manifest.cell(r, 4).value or "") for r in range(2, manifest.max_row + 1)]
    check(
        "清单保留未决项",
        sum("未生成" in x for x in manifest_status) == 0
        and sum("部分可粘贴" in x for x in manifest_status) == 7
        and sum("只读核对" in x for x in manifest_status) == 1
        and len(manifest_status) == 10,
        str(manifest_status),
    )

    failed = [item for item in checks if not item["ok"]]
    return {"status": "success" if not failed else "failed", "checks": len(checks), "failed": failed, "output": str(out)}


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--config-root", type=Path, default=DEFAULT_CONFIG)
    parser.add_argument("--out", type=Path, default=DEFAULT_OUT)
    args = parser.parse_args()
    result = validate(args.config_root, args.out)
    print(json.dumps(result, ensure_ascii=False, indent=2))
    raise SystemExit(0 if result["status"] == "success" else 1)


if __name__ == "__main__":
    main()
