#!/usr/bin/env python3
"""独立验证映射装备全链路配置包 P02。

验证器不导入生成器，避免实现错误在生成/验证两侧共享。
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import os
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path
from typing import Sequence

from openpyxl import load_workbook


FRAME_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_WORKBOOK = FRAME_ROOT / "04_数值文档" / "配置粘贴包" / "映射装备全链路配置包_P02.xlsx"
DEFAULT_GENERATION_REPORT = FRAME_ROOT / "05_验算核算" / "mapping-equip-full-package-P02-generation.json"
DEFAULT_OUTPUT = FRAME_ROOT / "05_验算核算" / "mapping-equip-full-package-P02-validation.json"
DEFAULT_CONFIG_ROOT = Path(os.environ.get("VGAME_CONFIG_DATAS", ""))


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def header_map(ws) -> dict[str, int]:
    result: dict[str, int] = {}
    for cell in ws[4]:
        if cell.value is None:
            continue
        value = str(cell.value)
        result[value] = cell.column
        # P02 可见表头为“中文\n正式英文字段”；验证器仍用正式字段名取列。
        for line in value.splitlines():
            if line.strip():
                result[line.strip()] = cell.column
    return result


def data_rows(ws) -> list[list]:
    headers = header_map(ws)
    max_col = max(headers.values()) if headers else ws.max_column
    return [[ws.cell(row, col).value for col in range(1, max_col + 1)] for row in range(5, ws.max_row + 1)]


def sheet_ids(wb, sheet: str, field: str = "Id") -> list[int]:
    ws = wb[sheet]
    headers = header_map(ws)
    column = headers[field]
    return [ws.cell(row, column).value for row in range(5, ws.max_row + 1) if isinstance(ws.cell(row, column).value, int)]


def official_hashes(config_root: Path, expected: list[dict]) -> tuple[bool, list[str]]:
    changed = []
    for record in expected:
        path = config_root / record["relative_path"]
        if not path.is_file() or path.stat().st_size != record["bytes"] or sha256_file(path) != record["sha256"]:
            changed.append(record["relative_path"])
    return not changed, changed


def add_check(checks: list[dict], name: str, actual, expected, detail: str = "") -> None:
    checks.append({"name": name, "actual": actual, "expected": expected, "status": "PASS" if actual == expected else "FAIL", "detail": detail})


def validate(workbook_path: Path, generation_report_path: Path, config_root: Path) -> dict:
    report = json.loads(generation_report_path.read_text(encoding="utf-8"))
    # 这里需要对同一工作表做大量随机访问；read_only模式的cell()
    # 会每次重新从XML开头扫描，反而成为O(n²)。正常模式一次解析更快。
    wb = load_workbook(workbook_path, read_only=False, data_only=False)
    checks: list[dict] = []

    required_sheets = [
        "00_P02总览", "01_全链状态", "02_ID分配", "10_主属性480", "11_副属性600", "12_属性落表门禁",
        "20_副本30层", "21_层级宝箱Items150", "22_Generate150", "23_主材Items72", "24_副材Items4", "24A_预览占位Items22",
        "25_Drop常规根110", "26_Drop宝箱池150", "27_Drop额外箱池", "28_Drop主材池", "29_Drop副材池4",
        "30_首通提案", "30A_首通Drop纯值", "31_EquipmentLevel150", "32_UIlevel150", "33_关卡入口链", "34_ItemAccess",
        "35_合成打造覆盖", "40_Drop粘贴规程", "41_边界排除", "90_P02阻断", "99_P02验算",
        "01_PortableExp", "02_SubAttrCoef", "03_PortableLvUp", "04_CompoundProb", "05_CraftProb", "10_CompoundCost", "11_QualityPatch",
    ]
    add_check(checks, "required_sheets", sorted(set(required_sheets) - set(wb.sheetnames)), [], "缺失Sheet")
    if any(check["status"] == "FAIL" for check in checks):
        wb.close()
        return {"status": "FAIL", "checks": checks}

    layer_ws = wb["20_副本30层"]
    lh = header_map(layer_ws)
    expected_box_pool_rows = 0
    expected_extra_pool_rows = 0
    expected_main_pool_rows = 0
    chest_probs_by_layer = {}
    for row in range(5, layer_ws.max_row + 1):
        layer_no = int(layer_ws.cell(row, lh["Layer"]).value)
        chest_probs_by_layer[layer_no] = {
            quality: float(layer_ws.cell(row, lh[field]).value)
            for quality, field in [("Green", "GreenBox"), ("Blue", "BlueBox"), ("Purple", "PurpleBox"), ("Orange", "OrangeBox"), ("Red", "RedBox")]
        }
        extra_probability = float(layer_ws.cell(row, lh["ExtraBoxProb"]).value)
        positive_material_qualities = sum(1 for field in ["PurpleMat", "OrangeMat", "RedMat"] if float(layer_ws.cell(row, lh[field]).value) > 0)
        expected_box_pool_rows += 5
        expected_extra_pool_rows += (1 if extra_probability < 1 else 0) + (5 if extra_probability > 0 else 0)
        main_probability = float(layer_ws.cell(row, lh["MainMatProb"]).value)
        if main_probability > 0:
            expected_main_pool_rows += 1 + 3 * positive_material_qualities  # 空掉落 + 3大类×正权重品质

    add_check(checks, "purple_box_start_layer", min(layer for layer, probs in chest_probs_by_layer.items() if probs["Purple"] > 0), 1)
    add_check(checks, "orange_box_start_layer", min(layer for layer, probs in chest_probs_by_layer.items() if probs["Orange"] > 0), 10)
    add_check(checks, "red_box_start_layer", min(layer for layer, probs in chest_probs_by_layer.items() if probs["Red"] > 0), 20)
    add_check(checks, "layer1_box_probs", [chest_probs_by_layer[1][quality] for quality in ["Green", "Blue", "Purple", "Orange", "Red"]], [0.515, 0.30, 0.185, 0.0, 0.0])
    add_check(checks, "layer10_high_box_probs", [chest_probs_by_layer[10][quality] for quality in ["Orange", "Red"]], [0.01, 0.0])
    add_check(checks, "layer20_high_box_probs", [chest_probs_by_layer[20][quality] for quality in ["Orange", "Red"]], [0.10, 0.005])
    add_check(checks, "layer30_high_box_probs", [chest_probs_by_layer[30][quality] for quality in ["Purple", "Orange", "Red"]], [0.37, 0.19, 0.04])
    actual_gold = {layer: probs["Purple"] * 0.10 + probs["Orange"] * 0.98 for layer, probs in chest_probs_by_layer.items()}
    actual_red = {layer: probs["Orange"] * 0.02 + probs["Red"] for layer, probs in chest_probs_by_layer.items()}
    add_check(checks, "actual_gold_strictly_increasing", all(actual_gold[layer] < actual_gold[layer + 1] for layer in range(1, 30)), True)
    add_check(checks, "actual_red_zero_before_layer10", all(math.isclose(actual_red[layer], 0, abs_tol=1e-12) for layer in range(1, 10)), True)
    add_check(checks, "actual_red_strictly_increasing_from_layer10", all(actual_red[layer] < actual_red[layer + 1] for layer in range(10, 30)), True)

    expected_rows = {
        "10_主属性480": 480,
        "11_副属性600": 600,
        "20_副本30层": 30,
        "21_层级宝箱Items150": 150,
        "22_Generate150": 150,
        "23_主材Items72": 72,
        "24_副材Items4": 4,
        "24A_预览占位Items22": 22,
        "25_Drop常规根110": 110,
        "26_Drop宝箱池150": expected_box_pool_rows,
        "27_Drop额外箱池": expected_extra_pool_rows,
        "28_Drop主材池": expected_main_pool_rows,
        "29_Drop副材池4": 4,
        "30_首通提案": 30,
        "30A_首通Drop纯值": 76,
        "31_EquipmentLevel150": 150,
        "32_UIlevel150": 150,
    }
    for sheet, expected in expected_rows.items():
        add_check(checks, f"rows:{sheet}", wb[sheet].max_row - 4, expected)

    bilingual_sheets = [
        "10_主属性480", "11_副属性600", "20_副本30层", "21_层级宝箱Items150", "22_Generate150",
        "23_主材Items72", "24_副材Items4", "24A_预览占位Items22", "25_Drop常规根110", "26_Drop宝箱池150", "27_Drop额外箱池",
        "28_Drop主材池", "29_Drop副材池4", "30A_首通Drop纯值", "31_EquipmentLevel150", "32_UIlevel150",
        "34_ItemAccess", "01_PortableExp", "02_SubAttrCoef", "03_PortableLvUp", "04_CompoundProb", "05_CraftProb",
        "10_CompoundCost", "11_QualityPatch",
    ]
    english_only_headers = []
    malformed_bilingual = []
    bilingual_count = 0
    for sheet in bilingual_sheets:
        ws = wb[sheet]
        for cell in ws[4]:
            if not isinstance(cell.value, str):
                continue
            value = cell.value
            if "\n" in value:
                first, *rest = value.splitlines()
                english = rest[-1] if rest else ""
                bilingual_count += 1
                if not re.search(r"[\u4e00-\u9fff]", first) or not (re.search(r"[A-Za-z]", english) or english == "##"):
                    malformed_bilingual.append((sheet, cell.coordinate, value))
            elif re.fullmatch(r"[A-Za-z][A-Za-z0-9_ ()/.-]*", value) or value in {"##Name"}:
                english_only_headers.append((sheet, cell.coordinate, value))
        if ws.row_dimensions[4].height is None or ws.row_dimensions[4].height < 40:
            malformed_bilingual.append((sheet, "row4", f"height={ws.row_dimensions[4].height}"))
    add_check(checks, "bilingual_header_count_min", bilingual_count >= 150, True, str(bilingual_count))
    add_check(checks, "english_only_headers", english_only_headers, [])
    add_check(checks, "bilingual_header_format", malformed_bilingual, [])

    # Items / Generate 一对一、ID不重复。
    box_ids = sheet_ids(wb, "21_层级宝箱Items150")
    main_item_ids = sheet_ids(wb, "23_主材Items72")
    sub_item_ids = sheet_ids(wb, "24_副材Items4")
    preview_item_ids = sheet_ids(wb, "24A_预览占位Items22")
    all_item_ids = box_ids + main_item_ids + sub_item_ids + preview_item_ids
    add_check(checks, "candidate_item_id_unique", len(all_item_ids), len(set(all_item_ids)))
    add_check(checks, "box_item_id_range", (min(box_ids), max(box_ids)), (4201, 4350))
    add_check(checks, "main_item_id_range", (min(main_item_ids), max(main_item_ids)), (1080, 1151))
    add_check(checks, "sub_item_id_range", (min(sub_item_ids), max(sub_item_ids)), (1152, 1155))
    add_check(checks, "preview_item_id_range", (min(preview_item_ids), max(preview_item_ids)), (4351, 4372))
    add_check(checks, "preview_item_id_contiguous", sorted(preview_item_ids), list(range(4351, 4373)))

    generate_ids = sheet_ids(wb, "22_Generate150")
    box_ws = wb["21_层级宝箱Items150"]
    box_headers = header_map(box_ws)
    box_para = [box_ws.cell(row, box_headers["ParaList"]).value for row in range(5, box_ws.max_row + 1)]
    add_check(checks, "generate_id_unique", len(generate_ids), len(set(generate_ids)))
    add_check(checks, "box_para_matches_generate", sorted(box_para), sorted(generate_ids))

    # 属性矩阵唯一性和公式独立复算。
    main_ws = wb["10_主属性480"]
    mh = header_map(main_ws)
    main_keys = []
    main_mismatches = []
    for row in range(5, main_ws.max_row + 1):
        part = main_ws.cell(row, mh["Part"]).value
        quality = main_ws.cell(row, mh["Quality"]).value
        level = main_ws.cell(row, mh["EquipLv"]).value
        target = float(main_ws.cell(row, mh["Red60Target"]).value)
        qcoef = float(main_ws.cell(row, mh["QualityCoef"]).value)
        value = float(main_ws.cell(row, mh["TargetValue"]).value)
        formula_coef = 0.05 + 0.95 * (((level - 5) / 55) ** 1.35)
        expected = target * qcoef * formula_coef
        main_keys.append((part, quality, level))
        if not math.isclose(value, expected, rel_tol=1e-8, abs_tol=1e-8):
            main_mismatches.append((row, value, expected))
    add_check(checks, "main_key_unique", len(main_keys), len(set(main_keys)))
    add_check(checks, "main_formula_mismatches", len(main_mismatches), 0, str(main_mismatches[:3]))

    sub_ws = wb["11_副属性600"]
    sh = header_map(sub_ws)
    sub_keys = []
    sub_mismatches = []
    for row in range(5, sub_ws.max_row + 1):
        attr = sub_ws.cell(row, sh["Attribute"]).value
        quality = sub_ws.cell(row, sh["Quality"]).value
        level = sub_ws.cell(row, sh["EquipLv"]).value
        target = float(sub_ws.cell(row, sh["Red60Target"]).value)
        qcoef = float(sub_ws.cell(row, sh["QualityCoef"]).value)
        value = float(sub_ws.cell(row, sh["TargetValue"]).value)
        formula_coef = 0.40 + 0.60 * ((level - 5) / 55)
        expected = target * qcoef * formula_coef
        sub_keys.append((attr, quality, level))
        if not math.isclose(value, expected, rel_tol=1e-8, abs_tol=1e-8):
            sub_mismatches.append((row, value, expected))
    add_check(checks, "sub_key_unique", len(sub_keys), len(set(sub_keys)))
    add_check(checks, "sub_formula_mismatches", len(sub_mismatches), 0, str(sub_mismatches[:3]))

    # Drop 权重、根包分支与外键。
    pool_ids = set()
    for sheet in ["26_Drop宝箱池150", "27_Drop额外箱池", "28_Drop主材池", "29_Drop副材池4"]:
        ws = wb[sheet]
        h = header_map(ws)
        grouped = defaultdict(int)
        non_positive = []
        for row in range(5, ws.max_row + 1):
            drop_id = ws.cell(row, h["DropId"]).value
            weight = ws.cell(row, h["DropWeight"]).value
            grouped[drop_id] += weight
            pool_ids.add(drop_id)
            # 宝箱品质池保留全品质稳定行结构；尚未开放的品质以 0 权重占位。
            allow_zero = sheet in ["26_Drop宝箱池150", "27_Drop额外箱池"]
            if not isinstance(weight, int) or weight < 0 or (weight == 0 and not allow_zero):
                non_positive.append((row, weight))
        bad = {drop_id: total for drop_id, total in grouped.items() if total != 10000}
        add_check(checks, f"weight_sum:{sheet}", bad, {})
        add_check(checks, f"positive_weight:{sheet}", non_positive, [])

    root_ws = wb["25_Drop常规根110"]
    rh = header_map(root_ws)
    root_groups = defaultdict(list)
    missing_pool_refs = []
    for row in range(5, root_ws.max_row + 1):
        drop_id = root_ws.cell(row, rh["DropId"]).value
        drop_id2 = root_ws.cell(row, rh["DropId2"]).value
        root_groups[drop_id].append(drop_id2)
        if drop_id2 not in pool_ids:
            missing_pool_refs.append((drop_id, drop_id2))
    add_check(checks, "regular_root_count", len(root_groups), 30)
    add_check(checks, "regular_root_branch_distribution", Counter(len(values) for values in root_groups.values()), Counter({2: 5, 4: 25}))
    early_material_refs = [(drop_id, drop_id2) for drop_id, values in root_groups.items() if 20005101 <= drop_id <= 20005105 for drop_id2 in values if 22301 <= drop_id2 <= 22305 or drop_id2 == 22400]
    add_check(checks, "early_regular_material_refs", early_material_refs, [])
    add_check(checks, "layer6_regular_material_refs", {22306, 22400}.issubset(set(root_groups[20005106])), True)
    add_check(checks, "regular_missing_pool_refs", missing_pool_refs, [])

    main_ws = wb["28_Drop主材池"]
    mh = header_map(main_ws)
    transition_pool_items = defaultdict(set)
    for row in range(5, main_ws.max_row + 1):
        drop_id = main_ws.cell(row, mh["DropId"]).value
        item_id = main_ws.cell(row, mh["DropItem_id"]).value
        if 22306 <= drop_id <= 22309 and isinstance(item_id, int) and 1080 <= item_id <= 1151:
            transition_pool_items[drop_id].add(item_id)
    expected_10_tier_items = [1080, 1083, 1086]
    add_check(checks, "layer6_8_regular_10_tier", {drop_id: sorted(transition_pool_items[drop_id]) for drop_id in range(22306, 22309)}, {drop_id: expected_10_tier_items for drop_id in range(22306, 22309)})
    add_check(checks, "layer9_regular_20_tier", sorted(transition_pool_items[22309]), [1089, 1092, 1095])

    first_ws = wb["30A_首通Drop纯值"]
    fh = header_map(first_ws)
    first_root_ids = set()
    first_pool_weights = defaultdict(int)
    first_fixed_bad = []
    first_nested_bad = []
    early_first_material_items = []
    for row in range(5, first_ws.max_row + 1):
        drop_id = first_ws.cell(row, fh["DropId"]).value
        weight = first_ws.cell(row, fh["DropWeight"]).value
        drop_type = first_ws.cell(row, fh["DropType"]).value
        drop_min = first_ws.cell(row, fh["DropMinNum"]).value
        drop_max = first_ws.cell(row, fh["DropMaxNum"]).value
        if drop_type == "fixed_" and (drop_min != 1 or drop_max != 1 or weight != 10000):
            first_fixed_bad.append((row, drop_id, drop_min, drop_max, weight))
        if drop_type == "fixed__random" and weight != 0:
            first_nested_bad.append((row, drop_id, weight))
        if 20000010 <= drop_id <= 20000039:
            first_root_ids.add(drop_id)
            item_id = first_ws.cell(row, fh["DropItem_id"]).value
            if drop_id <= 20000014 and isinstance(item_id, int) and 1080 <= item_id <= 1155:
                early_first_material_items.append((drop_id, item_id))
        elif 22451 <= drop_id <= 22453:
            first_pool_weights[drop_id] += weight
    add_check(checks, "first_root_coverage", sorted(first_root_ids), list(range(20000010, 20000040)))
    add_check(checks, "first_random_pool_weights", dict(first_pool_weights), {22451: 10000, 22452: 10000, 22453: 10000})
    add_check(checks, "first_fixed_semantics", first_fixed_bad, [])
    add_check(checks, "first_nested_semantics", first_nested_bad, [])
    add_check(checks, "early_first_clear_material_items", early_first_material_items, [])
    layer6_first_material_items=[]
    for row in range(5, first_ws.max_row + 1):
        if first_ws.cell(row, fh["DropId"]).value == 20000015:
            item_id=first_ws.cell(row, fh["DropItem_id"]).value
            if isinstance(item_id, int) and 1080 <= item_id <= 1155:
                layer6_first_material_items.append(item_id)
    add_check(checks, "layer6_first_clear_material_items", layer6_first_material_items, [1080])
    layer7_first_material_items=[]
    for row in range(5, first_ws.max_row + 1):
        if first_ws.cell(row, fh["DropId"]).value == 20000016:
            item_id=first_ws.cell(row, fh["DropItem_id"]).value
            if isinstance(item_id, int) and 1080 <= item_id <= 1155:
                layer7_first_material_items.append(item_id)
    add_check(checks, "layer7_first_clear_material_items", layer7_first_material_items, [1083])

    # EquipmentLevel 五品质映射。
    eq_ws = wb["31_EquipmentLevel150"]
    eh = header_map(eq_ws)
    eq_groups = defaultdict(dict)
    for row in range(5, eq_ws.max_row + 1):
        ui_id = eq_ws.cell(row, eh["UiLevelId"]).value
        quality = eq_ws.cell(row, eh["EquipmentQuality"]).value
        level_id = eq_ws.cell(row, eh["LevelId"]).value
        eq_groups[ui_id][quality] = level_id
    equipment_bad = []
    quality_order = ["Green", "Blue", "Purple", "Orange", "Red"]
    for ui_id, mapping in eq_groups.items():
        for suffix, quality in enumerate(quality_order, start=1):
            if mapping.get(quality) != ui_id - 1 + suffix:
                equipment_bad.append((ui_id, quality, mapping.get(quality)))
    add_check(checks, "equipment_level_group_count", len(eq_groups), 30)
    add_check(checks, "equipment_level_mapping", equipment_bad, [])

    # UIlevel 基准行外键和非基准行空引用。
    ui_ws = wb["32_UIlevel150"]
    uh = header_map(ui_ws)
    ui_bad = []
    base_count = 0
    for row in range(5, ui_ws.max_row + 1):
        layer = ui_ws.cell(row, uh["TierNum"]).value
        variant = ui_ws.cell(row, uh["Variant"]).value
        first_id = ui_ws.cell(row, uh["FirstRewardDrop"]).value
        fall_id = ui_ws.cell(row, uh["FallRewardDrop"]).value
        if variant == 1:
            base_count += 1
            if first_id != 20000009 + layer or fall_id != 20005100 + layer:
                ui_bad.append((layer, variant, first_id, fall_id))
        elif first_id is not None or fall_id is not None:
            ui_bad.append((layer, variant, first_id, fall_id))
    add_check(checks, "ui_base_count", base_count, 30)
    add_check(checks, "ui_drop_refs", ui_bad, [])

    craft_tier_by_layer = {
        **{layer: 10 for layer in range(6, 9)},
        9: 20,
        **{layer: 30 for layer in range(10, 16)},
        **{layer: 40 for layer in range(16, 22)},
        **{layer: 45 for layer in range(22, 25)},
        **{layer: 50 for layer in range(25, 27)},
        **{layer: 55 for layer in range(27, 29)},
        **{layer: 60 for layer in range(29, 31)},
    }
    preview_grade_box_ids = {5: 4351, 10: 4352, 15: 4353, 20: 4354, 25: 4365, 30: 4366, 35: 4367, 40: 4368, 45: 4369, 50: 4370, 55: 4371, 60: 4372}
    grade_pools = {
        **{layer: [5, 10] for layer in range(1, 5)},
        **{layer: [5, 10, 15] for layer in range(5, 8)},
        **{layer: [10, 15, 20] for layer in range(8, 10)},
        **{layer: [15, 20, 25] for layer in range(10, 12)},
        **{layer: [20, 25, 30] for layer in range(12, 14)},
        **{layer: [25, 30, 35] for layer in range(14, 16)},
        **{layer: [30, 35, 40] for layer in range(16, 18)},
        **{layer: [35, 40, 45] for layer in range(18, 20)},
        **{layer: [40, 45, 50] for layer in range(20, 22)},
        **{layer: [45, 50, 55] for layer in range(22, 24)},
        **{layer: [50, 55, 60] for layer in range(24, 27)},
        **{layer: [55, 60] for layer in range(27, 31)},
    }
    preview_main_ids = {
        (10, "Purple"): 4355, (20, "Purple"): 4356, (30, "Purple"): 4357,
        (30, "Orange"): 4358, (40, "Orange"): 4359, (45, "Orange"): 4360,
        (50, "Red"): 4361, (55, "Red"): 4362, (60, "Red"): 4363,
    }
    preview_mismatches = []
    preview_quantity_errors = []
    for row in range(5, ui_ws.max_row + 1):
        layer = int(ui_ws.cell(row, uh["TierNum"]).value)
        reward = str(ui_ws.cell(row, uh["Reward"]).value or "")
        actual_items = [int(part.split(",")[0]) for part in reward.split(";") if part]
        actual_quantities = [int(part.split(",")[1]) for part in reward.split(";") if part]
        if any(quantity != -1 for quantity in actual_quantities):
            preview_quantity_errors.append((layer, ui_ws.cell(row, uh["Variant"]).value, actual_quantities))
        expected_items = [preview_grade_box_ids[grade] for grade in grade_pools[layer]]
        if layer >= 6:
            material_quality = "Purple" if layer <= 10 else "Orange" if layer <= 24 else "Red"
            expected_items.extend([preview_main_ids[(craft_tier_by_layer[layer], material_quality)], 4364])
        if actual_items != expected_items:
            preview_mismatches.append((layer, ui_ws.cell(row, uh["Variant"]).value, actual_items, expected_items))
    add_check(checks, "ui_grade_box_preview", preview_mismatches, [])
    add_check(checks, "ui_preview_quantity_negative_one", preview_quantity_errors, [])

    # P01品质补丁必须在P02中降级为原子条件块。
    quality_warning = str(wb["11_QualityPatch"]["B3"].value)
    add_check(checks, "quality_patch_atomic_warning", all(token in quality_warning for token in ["不得单贴", "SpecialEquipmentMain", "EffectDrop"]), True)

    # 公式文本和已缓存错误值扫描。
    formula_errors = []
    error_tokens = ["#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A"]
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                if isinstance(value, str) and value.startswith("=") and any(token in value for token in error_tokens):
                    formula_errors.append((ws.title, cell.coordinate, value))
    add_check(checks, "formula_error_tokens", formula_errors, [])
    wb.close()

    wb_values = load_workbook(workbook_path, read_only=False, data_only=True)
    cached_errors = []
    for ws in wb_values.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.data_type == "e":
                    cached_errors.append((ws.title, cell.coordinate, cell.value))
    wb_values.close()
    add_check(checks, "cached_formula_errors", cached_errors, [])

    unchanged, changed = official_hashes(config_root, report["official_hashes_after"])
    add_check(checks, "official_datas_unchanged", unchanged, True, str(changed))
    add_check(checks, "generation_report_official_write", report.get("official_config_write"), False)
    add_check(checks, "generation_report_luban", report.get("luban_executed"), False)
    add_check(checks, "workbook_sha256", sha256_file(workbook_path), report["workbook_sha256"])

    failures = [check for check in checks if check["status"] != "PASS"]
    return {
        "status": "PASS" if not failures else "FAIL",
        "package_id": report.get("package_id"),
        "workbook": str(workbook_path),
        "workbook_sha256": sha256_file(workbook_path),
        "official_datas_unchanged": unchanged,
        "formula_error_count": len(formula_errors) + len(cached_errors),
        "check_count": len(checks),
        "failure_count": len(failures),
        "failures": failures,
        "checks": checks,
    }


def parse_args(argv: Sequence[str]) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="验证映射装备全链路配置包P02")
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--generation-report", type=Path, default=DEFAULT_GENERATION_REPORT)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--config-root", type=Path, default=DEFAULT_CONFIG_ROOT)
    return parser.parse_args(argv)


def main(argv: Sequence[str] | None = None) -> int:
    args = parse_args(sys.argv[1:] if argv is None else argv)
    result = validate(args.workbook.resolve(), args.generation_report.resolve(), args.config_root.resolve())
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({key: result[key] for key in ["status", "package_id", "workbook", "check_count", "failure_count", "formula_error_count", "official_datas_unchanged"]}, ensure_ascii=False, indent=2))
    return 0 if result["status"] == "PASS" else 1


if __name__ == "__main__":
    raise SystemExit(main())
