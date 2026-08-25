from __future__ import annotations

import json
import os
import zipfile
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
CONFIG = Path(os.environ.get("VGAME_CONFIG_DATAS", ""))
OUT = Path(os.environ.get("RV_FINAL_OUT", str(ROOT / "06_最终配置" / "房车")))
EXPECTED = {
    "00_配置清单.xlsx",
    "RvAttribute.xlsx", "RvGuide.xlsx", "RvGuideNode.xlsx", "RvGuideNodeLv.xlsx",
    "RvLvUp.xlsx", "RvStage.xlsx", "RvStrengthenResult.xlsx", "Unlock.xlsx",
    "ShopGoods.xlsx", "Drop_dungeon.xlsx", "UIlevel.xlsx",
}
SOURCES = {
    "RvAttribute.xlsx": CONFIG / "Car" / "RvAttribute.xlsx",
    "RvGuide.xlsx": CONFIG / "Car" / "RvGuide.xlsx",
    "RvGuideNode.xlsx": CONFIG / "Car" / "RvGuideNode.xlsx",
    "RvGuideNodeLv.xlsx": CONFIG / "Car" / "RvGuideNodeLv.xlsx",
    "RvLvUp.xlsx": CONFIG / "Car" / "RvLvUp.xlsx",
    "RvStage.xlsx": CONFIG / "Car" / "RvStage.xlsx",
    "RvStrengthenResult.xlsx": CONFIG / "Car" / "RvStrengthenResult.xlsx",
    "Unlock.xlsx": CONFIG / "unlock" / "Unlock.xlsx",
    "ShopGoods.xlsx": CONFIG / "Shop" / "ShopGoods.xlsx",
    "Drop_dungeon.xlsx": CONFIG / "Drop" / "Drop_dungeon.xlsx",
    "UIlevel.xlsx": CONFIG / "level" / "UIlevel.xlsx",
}


def fields(ws) -> dict[str, int]:
    return {str(c.value): c.column for c in ws[1] if c.value is not None}


def rows(ws, key: str) -> list[dict]:
    h = fields(ws)
    result = []
    for r in range(5, ws.max_row + 1):
        if ws.cell(r, h[key]).value in (None, ""):
            continue
        result.append({name: ws.cell(r, col).value for name, col in h.items()})
    return result


def same_schema(source: Path, mirror: Path) -> bool:
    a = load_workbook(source, data_only=False, read_only=True).active
    b = load_workbook(mirror, data_only=False, read_only=True).active
    last = max(c for r in range(1, 5) for c in range(1, a.max_column + 1) if a.cell(r, c).value is not None)
    return all(a.cell(r, c).value == b.cell(r, c).value for r in range(1, 5) for c in range(1, last + 1))


def safety(path: Path) -> tuple[list[str], list[str], list[str]]:
    wb = load_workbook(path, data_only=False, read_only=True)
    formulas, errors = [], []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.data_type == "f" or (isinstance(cell.value, str) and cell.value.startswith("=")):
                    formulas.append(f"{ws.title}!{cell.coordinate}")
                if cell.data_type == "e":
                    errors.append(f"{ws.title}!{cell.coordinate}:{cell.value}")
    with zipfile.ZipFile(path) as zf:
        external = [x for x in zf.namelist() if x.startswith("xl/externalLinks/")]
    return formulas, errors, external


def validate() -> dict:
    checks = []

    def check(name: str, ok: bool, detail=""):
        checks.append({"name": name, "ok": bool(ok), "detail": str(detail)})

    actual = {p.name for p in OUT.glob("*.xlsx")}
    check("文件清单", actual == EXPECTED, sorted(actual))
    for name in sorted(EXPECTED):
        path = OUT / name
        if not path.exists():
            continue
        wb = load_workbook(path, data_only=False, read_only=True)
        expected_sheets = 1 if name == "00_配置清单.xlsx" else 2
        check(f"{name}:工作表", len(wb.sheetnames) == expected_sheets and (expected_sheets == 1 or wb.sheetnames[1] == "中文说明"), wb.sheetnames)
        formulas, errors, external = safety(path)
        check(f"{name}:无公式", not formulas, formulas[:5])
        check(f"{name}:无公式错误", not errors, errors[:5])
        check(f"{name}:无外链", not external, external)
        if name in SOURCES:
            check(f"{name}:Schema一致", same_schema(SOURCES[name], path), SOURCES[name])

    attr = rows(load_workbook(OUT / "RvAttribute.xlsx", data_only=False).active, "Id")
    check("房车属性3条", len(attr) == 3, len(attr))
    check("房车满级属性", {r["Attribute"]: round(r["Grow"] * 99 + r["Value"]) for r in attr} == {"Hp": 25000, "Attack": 7500, "Armor": 5000})

    lv = rows(load_workbook(OUT / "RvLvUp.xlsx", data_only=False).active, "Id")
    check("房车升级99档", len(lv) == 99 and {r["TargetLv"] for r in lv} == set(range(2, 101)), len(lv))
    check("房车升级消耗", sum(r["CostItemNum"] for r in lv) == 15457 and sum(r["CostCurrencyNum"] for r in lv) == 2164000)

    result = rows(load_workbook(OUT / "RvStrengthenResult.xlsx", data_only=False).active, "Id")
    check("强化结果", [(r["AddProgress"], r["Weight"]) for r in result] == [(30, 5000), (60, 3000), (100, 2000)])

    guide = rows(load_workbook(OUT / "RvGuide.xlsx", data_only=False).active, "Id")
    stage = rows(load_workbook(OUT / "RvStage.xlsx", data_only=False).active, "Id")
    node = rows(load_workbook(OUT / "RvGuideNode.xlsx", data_only=False).active, "Id")
    node_lv = rows(load_workbook(OUT / "RvGuideNodeLv.xlsx", data_only=False).active, "Id")
    check("指南结构数量", (len(guide), len(stage), len(node), len(node_lv)) == (5, 20, 310, 310), (len(guide), len(stage), len(node), len(node_lv)))
    check("节点引用", {r["StageId"] for r in node} <= {r["Id"] for r in stage} and {r["NodeId"] for r in node_lv} <= {r["Id"] for r in node})
    node_by_id = {r["Id"]: r for r in node}
    topology_widths = [
        [1, 2, 1, 1, 2, 1],
        [1, 2, 1, 1, 2, 2, 1, 1, 1, 2, 1],
        [1, 2, 1, 1, 2, 2, 1, 1, 1, 2, 1, 1, 2, 1],
        [1, 2, 2, 1, 1, 1, 2, 1, 1, 2, 1, 1, 1, 2, 1],
    ]
    stage_spans = [(1, 8), (9, 23), (24, 42), (43, 62)]
    topology_ok = True
    topology_detail = {}
    expected_parents = defaultdict(list)
    for guide_id in range(1, 6):
        for stage_index, (start, end) in enumerate(stage_spans):
            ids = [guide_id * 100 + i for i in range(start, end + 1)]
            widths = topology_widths[stage_index]
            layers = []
            cursor = 0
            for width in widths:
                layers.append(ids[cursor:cursor + width])
                cursor += width
            expected = []
            for layer_index, current in enumerate(layers):
                for lane, node_id in enumerate(current):
                    if layer_index + 1 < len(layers):
                        following = layers[layer_index + 1]
                        next_ids = [following[lane]] if len(current) == len(following) == 2 else following
                    elif end < 62:
                        next_ids = [guide_id * 100 + end + 1]
                    else:
                        next_ids = []
                    for next_id in next_ids:
                        expected_parents[next_id].append(node_id)
                    expected.append((layer_index + 1, ";".join(map(str, next_ids)) or None))
            actual = [(node_by_id[i]["ColIdx"], node_by_id[i]["NextNodeIdList"]) for i in ids]
            topology_detail[(guide_id, stage_index + 1)] = actual
            topology_ok &= actual == expected
    check("四阶段新拓扑无2-1-2", topology_ok, topology_detail)
    by_item = Counter()
    by_gold = Counter()
    for r in node_lv:
        item_id, num = map(int, str(r["CostItemList"]).split(","))
        by_item[item_id] += num
        by_gold[(r["NodeId"] // 100)] += r["CostCurrencyNum"]
    check("五职业材料各5000", by_item == Counter({1075: 5000, 1076: 5000, 1077: 5000, 1078: 5000, 1079: 5000}), by_item)
    check("五职业金币各1364500", all(by_gold[i] == 1364500 for i in range(1, 6)), by_gold)

    unlock = rows(load_workbook(OUT / "Unlock.xlsx", data_only=False).active, "Id")
    check("房车解锁314条", len(unlock) == 314, len(unlock))
    check("阶段等级门槛", {r["Id"]: r["UnLockParams"] for r in unlock if r["Id"] in range(12001, 12005)} == {12001: 1, 12002: 15, 12003: 35, 12004: 50})
    unlock_by_node = {
        int(str(r["UnLockParams"]).split(";")[0]): r["Id"]
        for r in unlock
        if r["UnLockType"] == "UnLock_UUT_Car_Attribute_Grade"
    }
    node_lv_by_node = {r["NodeId"]: r for r in node_lv}
    preposition_ok = True
    for node_id, row in node_lv_by_node.items():
        actual = {int(x) for x in str(row["PrepositionGradeList"] or "").split(";") if x}
        expected = {unlock_by_node[parent] for parent in expected_parents[node_id]}
        preposition_ok &= actual == expected
    check("节点前置链匹配新拓扑", preposition_ok)

    shop = rows(load_workbook(OUT / "ShopGoods.xlsx", data_only=False).active, "Id")
    mapping = {1109011: 1075, 1109012: 1076, 1109013: 1077, 1109014: 1078, 1109015: 1079}
    check("行商商品映射", {r["Id"]: r["ItemId"] for r in shop} == mapping)
    check("行商房车材料20个×10次×4000", all(r["ShopId"] == 1110 and r["Num"] == 20 and r["Price"] == "301,4000" and r["LimitType"] == "BuyLimit_Week" and r["LimitNum"] == 10 for r in shop))

    drop = rows(load_workbook(OUT / "Drop_dungeon.xlsx", data_only=False).active, "DropId")
    totals = Counter()
    for r in drop:
        totals[r["DropItem_id"]] += r["DropItemNumMin"]
    check("首通100包245条", len(drop) == 245 and len({r["DropId"] for r in drop}) == 100, (len(drop), len({r["DropId"] for r in drop})))
    check("首通核心总量", all(totals[k] == v for k, v in {1074: 4000, 106: 40, 107: 40, 111: 19, 112: 19, 2: 4000}.items()), totals)

    ui = rows(load_workbook(OUT / "UIlevel.xlsx", data_only=False).active, "LevelId")
    check("旧重复奖励清理400条", len(ui) == 400 and all(r["Reward"] in (None, "") and r["FallRewardDrop"] in (None, "") for r in ui), len(ui))

    failed = [x for x in checks if not x["ok"]]
    return {"status": "pass" if not failed else "fail", "checks": len(checks), "passed": len(checks) - len(failed), "failed": failed}


if __name__ == "__main__":
    report = validate()
    print(json.dumps(report, ensure_ascii=False, indent=2))
    raise SystemExit(0 if report["status"] == "pass" else 1)
