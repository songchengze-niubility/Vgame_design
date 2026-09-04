"""Read-only Excel extraction and validation plan; does not author workbooks."""
import hashlib
import json
import os
import sys
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
FRAMEWORK = HERE.parent.parent
OUT = FRAMEWORK / "06_最终配置/映射装备"
CONFIG = Path(os.environ["VGAME_CONFIG_DATAS"])
SOURCE = Path(sys.argv[1])
RULES = json.loads((HERE / "rules.json").read_text(encoding="utf-8"))


def sha(p):
    return hashlib.sha256(p.read_bytes()).hexdigest()


def rows(p, width, sheet=None, data_only=False):
    w = openpyxl.load_workbook(p, read_only=True, data_only=data_only)
    s = w[sheet] if sheet else w.worksheets[0]
    s.reset_dimensions()
    result = [list(r) for r in s.iter_rows(max_col=width, values_only=True)]
    w.close()
    return result


def snapshot(p):
    w = openpyxl.load_workbook(p, read_only=True, data_only=False)
    result = {}
    for s in w:
        s.reset_dimensions()
        rs = [list(r) for r in s.iter_rows(values_only=True)]
        result[s.title] = rs
        if s.title == "SpecialEquipment":
            assert not any(isinstance(v, str) and v.startswith("=") for r in rs for v in r), p
    w.close()
    return result


ids = set(RULES["skills"].values())
source_rows = rows(SOURCE, 9, RULES["sourceSheet"])
source = {r[3]: {"id": r[2], "row": i+1, "effect": r[4]} for i,r in enumerate(source_rows) if i >= 98 and isinstance(r[2], int)}
assert {k:v["id"] for k,v in source.items()} == RULES["skills"]
skill_rows = {r[1]: r for r in rows(CONFIG / "Skill/Skill.xlsx", 14, data_only=True)[4:] if r[0] != "##" and r[1] in ids}
assert set(skill_rows) == ids
assert all(r[9] == "SkillsSpecialEquipment" and r[11] == f"MappingEquip/ME_skill_{i}" for i,r in skill_rows.items())
registry = rows(CONFIG / "__tables__.xlsx", 20)
for logical in ["SpecialEquipment", "SpecialEquipmentEffect", "Skill"]:
    assert any(logical in r and logical+".xlsx" in r for r in registry), logical
files = ["SpecialEquipment.xlsx", "SpecialEquipmentEffect.xlsx", "00_配置清单.xlsx"]
before = {f:snapshot(OUT/f) for f in files}
equip = before[files[0]]["SpecialEquipment"]
effects = before[files[1]]["SpecialEquipment"]
formal = {r[1]:r for r in rows(CONFIG / "SpecialEquipment/SpecialEquipment.xlsx", 13)[4:] if isinstance(r[1],int) and r[0] != "##"}
changes = []
bindings = []
for n,r in enumerate(equip[4:],5):
    if r[5] != "Red":
        continue
    f = formal[r[1]]
    assert (f[3],f[4],f[5]) == (r[3],r[4],r[5])
    if r[2] != f[2]:
        assert r[2] == "待占位", (r[1],r[2],f[2])
        changes.append({"row":n,"col":2,"old":r[2],"new":f[2]})
    name = RULES["aliases"].get(f[2], f[2])
    skill = source.get(name)
    if skill:
        assert r[11] is None
        changes.append({"row":n,"col":11,"old":None,"new":skill["id"]})
        bindings.append({"equipmentId":r[1],"name":f[2],"grade":r[3],"part":r[4],"skillName":name,"skillId":skill["id"],"sourceRow":skill["row"]})
assert sorted(b["equipmentId"] for b in bindings) == RULES["expectedEquipmentIds"]
assert len([c for c in changes if c["col"] == 2]) == 10
assert len(bindings) == 22 and {b["skillId"] for b in bindings} == ids
for p in [OUT / files[1], CONFIG / "SpecialEquipment" / files[1]]:
    assert not ids.intersection(r[1] for r in rows(p, 9)[4:]), p
tracked = [CONFIG/"SpecialEquipment"/f for f in [files[0],files[1],"SpecialEquipmentEffectDrop.xlsx","SpecialEquipmentQuality.xlsx"]]
tracked += [CONFIG/"Skill/Skill.xlsx", CONFIG/"__tables__.xlsx", SOURCE]
result = {"source":str(SOURCE),"files":files,"before":before,"sourceSkills":source,"bindings":bindings,"changes":changes,
          "formalNames":{str(i):r[2] for i,r in formal.items()},"skillRows":skill_rows,
          "formalHashes":{str(p):sha(p) for p in tracked},"candidateHashes":{p.name:sha(p) for p in OUT.glob("*.xlsx")}}
(HERE / "preflight.json").write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding="utf-8")
print(json.dumps({"bindings":len(bindings),"skills":len(ids),"nameUpdates":10,"changes":changes},ensure_ascii=False))
