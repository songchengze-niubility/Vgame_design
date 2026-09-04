# -*- coding: utf-8 -*-
"""CHG-010：七日活动竞技场历史最高分目标调整为1100/1200/1300/1400/1500/1600/1700。

只改 06_最终配置/任务 的 Task.xlsx、TaskTarget.xlsx 各4行，
备份、plan.json 与逐格校验落在 05_验算核算/七日竞技场分数CHG010_20260831/。
不改 SVN 工作副本，不改奖励与解锁。
"""
import hashlib
import json
import shutil
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
RUN = ROOT / "05_验算核算" / "七日竞技场分数CHG010_20260831"

# task_id -> (task_name/desc, param_score)
CHANGES = {
    5105: ("竞技场历史最高达到1100分", 1100),
    5205: ("竞技场历史最高达到1200分", 1200),
    5605: ("竞技场历史最高达到1600分", 1600),
    5705: ("竞技场历史最高达到1700分", 1700),
}
UNCHANGED = {5305: 1300, 5405: 1400, 5505: 1500}


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def find_rows(ws, id_col, ids):
    hits = {}
    for row in range(5, ws.max_row + 1):
        value = ws.cell(row=row, column=id_col).value
        if value in ids:
            hits[value] = row
    return hits


def main():
    RUN.mkdir(parents=True, exist_ok=True)
    (RUN / "backup" / "任务").mkdir(parents=True, exist_ok=True)

    targets = {
        "任务/Task.xlsx": {
            "sheet": "Task",
            "id_col": 2,
            "edits": {tid: {3: name, 4: name} for tid, (name, _) in CHANGES.items()},
        },
        "任务/TaskTarget.xlsx": {
            "sheet": "TaskTarget",
            "id_col": 2,
            "edits": {
                tid: {3: name, 5: f"0,{score}"}
                for tid, (name, score) in CHANGES.items()
            },
        },
    }

    plan = {
        "batch": "ARENA-SCORE-CHG010",
        "created": datetime.now().isoformat(),
        "root": str(ROOT),
        "run": str(RUN),
        "files": [],
    }
    applied = []

    for rel, spec in targets.items():
        src = ROOT / "06_最终配置" / rel
        backup = RUN / "backup" / rel
        shutil.copy2(src, backup)

        wb = load_workbook(src)
        ws = wb[spec["sheet"]]
        rows = find_rows(ws, spec["id_col"], set(spec["edits"]))
        missing = sorted(set(spec["edits"]) - set(rows))
        if missing:
            raise SystemExit(f"{rel}: 未找到任务行 {missing}")

        file_plan = {
            "file": rel,
            "source": str(src),
            "backup": str(backup),
            "sha256_before": sha256(src),
            "edits": [],
        }
        for tid in sorted(spec["edits"]):
            for col, new in spec["edits"][tid].items():
                cell = ws.cell(row=rows[tid], column=col)
                old = cell.value
                if old == new:
                    continue
                file_plan["edits"].append(
                    {
                        "sheet": spec["sheet"],
                        "row": rows[tid],
                        "col": col,
                        "task_id": tid,
                        "old": old,
                        "new": new,
                    }
                )
                cell.value = new
                applied.append((rel, spec["sheet"], rows[tid], col, old, new))

        plan["files"].append(file_plan)
        wb.save(src)
        file_plan["sha256_after"] = sha256(src)

    (RUN / "plan.json").write_text(
        json.dumps(plan, ensure_ascii=False, indent=2), encoding="utf-8"
    )

    # 校验：备份 vs 现场逐格对比，仅允许 plan 内的单元格发生变化。
    expected = defaultdict(dict)
    for f in plan["files"]:
        for e in f["edits"]:
            expected[(f["file"], e["sheet"])][(e["row"], e["col"])] = e["new"]

    errors = []
    for rel, spec in targets.items():
        old_wb = load_workbook(RUN / "backup" / rel, read_only=True)
        new_wb = load_workbook(ROOT / "06_最终配置" / rel, read_only=True)
        for sheet in spec["sheet"] and [spec["sheet"]]:
            old_rows = list(old_wb[sheet].iter_rows(values_only=True))
            new_rows = list(new_wb[sheet].iter_rows(values_only=True))
            if len(old_rows) != len(new_rows):
                errors.append(f"{rel}!{sheet}: 行数变化 {len(old_rows)}->{len(new_rows)}")
                continue
            allow = expected.get((rel, sheet), {})
            for ri, (old_r, new_r) in enumerate(zip(old_rows, new_rows), 1):
                for ci, (old_v, new_v) in enumerate(zip(old_r, new_r), 1):
                    if old_v != new_v and allow.get((ri, ci)) != new_v:
                        errors.append(
                            f"{rel}!{sheet}!{get_column_letter(ci)}{ri}: "
                            f"{old_v!r}->{new_v!r} 不在计划内"
                        )
        old_wb.close()
        new_wb.close()

    # 业务校验：7天分数线与奖励/解锁不变。
    check = {}
    wb_task = load_workbook(ROOT / "06_最终配置" / "任务" / "Task.xlsx", read_only=True)
    wb_tt = load_workbook(
        ROOT / "06_最终配置" / "任务" / "TaskTarget.xlsx", read_only=True
    )
    tt_param = {}
    for row in wb_tt["TaskTarget"].iter_rows(min_row=5, values_only=True):
        if row[1] is not None:
            tt_param[row[1]] = row[4]
    days = {}
    for row in wb_task["Task"].iter_rows(min_row=5, values_only=True):
        tid = row[1]
        if isinstance(tid, int) and 5105 <= tid <= 5705 and tid % 100 == 5:
            day = tid // 100 - 50
            days[day] = {
                "task_id": tid,
                "name": row[2],
                "param": tt_param.get(tid),
                "drop": row[14],
                "unlock": row[10],
            }
    check["七日竞技场分数线"] = {
        f"第{d}天": days[d]["param"] for d in sorted(days)
    }
    check["分数线符合1100~1700"] = [
        days[d]["param"] == f"0,{1100 + (d - 1) * 100}" for d in sorted(days)
    ] == [True] * 7
    check["奖励与解锁不变"] = all(
        days[d]["unlock"] == 4112
        and days[d]["drop"] == {1: 91000005, 2: 91000015, 3: 91000025,
                                4: 91000035, 5: 91000045, 6: 91000055,
                                7: 91000065}[d]
        for d in days
    )
    wb_task.close()
    wb_tt.close()

    result = {
        "batch": "ARENA-SCORE-CHG010",
        "generated_at": datetime.now().isoformat(),
        "applied_cells": len(applied),
        "cell_diff_outside_plan": errors,
        "business_checks": check,
        "status": "PASS" if not errors and check.get("分数线符合1100~1700")
        and check.get("奖励与解锁不变") else "FAIL",
    }
    (RUN / "validation.json").write_text(
        json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(json.dumps(result, ensure_ascii=False, indent=2))
    if result["status"] != "PASS":
        raise SystemExit(1)


if __name__ == "__main__":
    main()
