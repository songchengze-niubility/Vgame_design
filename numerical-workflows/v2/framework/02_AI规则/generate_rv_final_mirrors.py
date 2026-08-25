from __future__ import annotations

import hashlib
import os
import shutil
import subprocess
import sys
import tempfile
from copy import copy
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill


ROOT = Path(__file__).resolve().parents[1]
CONFIG = Path(os.environ.get("VGAME_CONFIG_DATAS", ""))
DESIGN_ROOT = Path(os.environ.get("VGAME_DESIGN_ROOT", ""))
APPLY_SCRIPT = DESIGN_ROOT / "output" / "rv_growth" / "apply_final_config.py"
OUT = Path(os.environ.get("RV_FINAL_OUT", str(ROOT / "06_最终配置" / "房车")))
BATCH = "RV-20260723-F03"
SOURCE_NOTE = "Source: 03_沉淀MD/参数基线/养成参数基线.md；CHG-028（用户确认 2026-07-23）"
NAVY = "17365D"
WHITE = "FFFFFF"
BLUE = "0000FF"
YELLOW = "FFF2CC"


SPECS = {
    "RvAttribute.xlsx": ("Car/RvAttribute.xlsx", "RvAttribute", "Id", "全表替换", "VALUE_READY"),
    "RvGuide.xlsx": ("Car/RvGuide.xlsx", "RvGuide", "Id", "全表替换", "VALUE_READY"),
    "RvGuideNode.xlsx": ("Car/RvGuideNode.xlsx", "RvGuideNode", "Id", "全表替换", "VALUE_READY"),
    "RvGuideNodeLv.xlsx": ("Car/RvGuideNodeLv.xlsx", "RvGuideNodeLv", "Id", "全表替换", "VALUE_READY"),
    "RvLvUp.xlsx": ("Car/RvLvUp.xlsx", "RvLvUp", "Id", "全表替换", "VALUE_READY"),
    "RvStage.xlsx": ("Car/RvStage.xlsx", "RvStage", "Id", "全表替换", "VALUE_READY"),
    "RvStrengthenResult.xlsx": ("Car/RvStrengthenResult.xlsx", "RvStrengthenResult", "Id", "全表替换", "VALUE_READY"),
    "Unlock.xlsx": ("unlock/Unlock.xlsx", "Unlock", "Id", "按 Id 覆盖 12001-12004、12101-12410", "VALUE_READY"),
    "ShopGoods.xlsx": ("Shop/ShopGoods.xlsx", "Goods", "Id", "按 Id 覆盖 1109011-1109015", "VALUE_READY"),
    "Drop_dungeon.xlsx": ("Drop/Drop_dungeon.xlsx", "Drop", "DropId", "复用旧爬塔奖励块；ID 公式列不粘贴", "VALUE_READY_FORMULA_MANUAL"),
    "UIlevel.xlsx": ("level/UIlevel.xlsx", "level", "LevelId", "按 LevelId 覆盖旧重复奖励字段", "VALUE_READY"),
}


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def copy_cell(src, dst, value=True) -> None:
    if value:
        dst.value = src.value
    if getattr(src, "has_style", False):
        dst.font = copy(src.font)
        dst.fill = copy(src.fill)
        dst.border = copy(src.border)
        dst.number_format = src.number_format
        dst.alignment = copy(src.alignment)
        dst.protection = copy(src.protection)


def header_map(ws) -> dict[str, int]:
    return {str(c.value): c.column for c in ws[1] if c.value is not None}


def copy_schema(src, dst) -> None:
    schema_rows = list(src.iter_rows(min_row=1, max_row=4))
    last_col = max(c.column for row in schema_rows for c in row if c.value is not None)
    for r, row in enumerate(schema_rows, 1):
        for c, cell in enumerate(row[:last_col], 1):
            copy_cell(cell, dst.cell(r, c))
    for c in range(1, last_col + 1):
        dst.column_dimensions[dst.cell(1, c).column_letter].width = 16


def add_cn_info(wb, file_name: str, formal_path: Path, rows: int, usage: str, status: str) -> None:
    ws = wb.create_sheet("中文说明")
    warning = "这是 06_最终配置评审镜像，不是正式 Datas；未经用户再次批准不得同步 SVN。"
    if file_name == "Drop_dungeon.xlsx":
        warning += " 正式表 ID 列为公式列，镜像中保持空白；正式落表时必须保留原公式，仅粘贴纯值列。"
    data = [
        ["项目", "内容"],
        ["配置镜像文件", file_name],
        ["交付批次", BATCH],
        ["落表状态", status],
        ["正式配置位置", str(formal_path)],
        ["数值来源", SOURCE_NOTE],
        ["有效数据行", rows],
        ["使用方式", usage],
        ["重要提醒", warning],
    ]
    for r, row in enumerate(data, 1):
        for c, value in enumerate(row, 1):
            cell = ws.cell(r, c, value)
            cell.font = Font(name="Microsoft YaHei", size=10, bold=r == 1, color=WHITE if r == 1 else "222222")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if r == 1:
                cell.fill = PatternFill("solid", fgColor=NAVY)
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 110
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False


def active_rows(ws, key: str) -> list[int]:
    col = header_map(ws)[key]
    return [r for r, row in enumerate(ws.iter_rows(min_row=5), 5) if row[0].value != "##" and row[col - 1].value not in (None, "")]


def write_mirror(name: str, source: Path, staged: Path, selected_rows: list[int]) -> None:
    src_wb = load_workbook(source, data_only=False, read_only=True)
    stage_wb = load_workbook(staged, data_only=False, read_only=True)
    src = src_wb.active
    stage = stage_wb.active
    wb = Workbook()
    ws = wb.active
    ws.title = stage.title
    copy_schema(src, ws)
    headers = header_map(stage)
    formula_blank_cols = {headers["ID"]} if name == "Drop_dungeon.xlsx" else set()
    selected = set(selected_rows)
    out_r = 4
    for stage_r, row in enumerate(stage.iter_rows(min_row=5), 5):
        if stage_r not in selected:
            continue
        out_r += 1
        for c, source_cell in enumerate(row, 1):
            target_cell = ws.cell(out_r, c)
            copy_cell(source_cell, target_cell, value=False)
            value = source_cell.value
            if isinstance(value, str) and value.startswith("="):
                if c not in formula_blank_cols:
                    raise RuntimeError(f"镜像中发现未处理公式: {name}!{source_cell.coordinate}")
                target_cell.value = None
                target_cell.fill = PatternFill("solid", fgColor=YELLOW)
                target_cell.comment = Comment("正式表公式列，保持原公式，不从镜像粘贴。", "Codex")
            else:
                target_cell.value = value
                if value not in (None, ""):
                    target_cell.font = copy(source_cell.font)
                    target_cell.font = Font(
                        name=target_cell.font.name or "Microsoft YaHei",
                        size=target_cell.font.sz or 10,
                        bold=target_cell.font.bold,
                        italic=target_cell.font.italic,
                        color=BLUE,
                    )
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:{ws.cell(max(ws.max_row, 4), ws.max_column).coordinate}"
    add_cn_info(wb, name, source, len(selected_rows), SPECS[name][3], SPECS[name][4])
    wb.save(OUT / name)
    src_wb.close()
    stage_wb.close()


def make_manifest(rows: list[tuple[str, str, int, str, str]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "配置清单"
    headers = ["文件", "正式目标", "有效数据行", "状态", "落表方式", "说明"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(name="Microsoft YaHei", size=10, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for name, target, count, status, usage in rows:
        note = "房车最终配置评审镜像；未修改正式 Datas。"
        if name == "Drop_dungeon.xlsx":
            note += " ID 公式列留空。"
        ws.append([name, target, count, status, usage, note])
    ws.append(["依赖（不生成）", "items/Items.xlsx", 6, "NO_CHANGE", "复用既有 ItemId 1074-1079", "正式 Items 已存在，不需要配置变更。"])
    ws.append(["依赖（不生成）", "Tower/TowerFirstReward.xlsx", 100, "NO_CHANGE", "继续引用 DropId 24000001-24000100", "首通引用链不需要配置变更。"])
    for row in ws.iter_rows():
        for cell in row:
            cell.font = cell.font.copy(name="Microsoft YaHei", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for col, width in zip("ABCDEF", [28, 48, 14, 30, 38, 64]):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(OUT / "00_配置清单.xlsx")


def main() -> None:
    if not APPLY_SCRIPT.exists():
        raise FileNotFoundError(APPLY_SCRIPT)
    inputs = [CONFIG / rel for rel, *_ in SPECS.values()] + [CONFIG / "items" / "Items.xlsx"]
    before = {path: sha256(path) for path in inputs}
    OUT.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="rv-final-mirror-") as temp:
        stage_root = Path(temp)
        for source in inputs:
            rel = source.relative_to(CONFIG)
            if source.name == "Items.xlsx":
                rel = Path("Items") / source.name
            target = stage_root / rel
            target.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(source, target)
        env = dict(os.environ)
        env["RV_FINAL_DATAS"] = str(stage_root)
        subprocess.run([sys.executable, str(APPLY_SCRIPT)], check=True, env=env, capture_output=True, text=True, encoding="utf-8")

        manifest_rows = []
        for name, (rel, _sheet, key, usage, status) in SPECS.items():
            source = CONFIG / rel
            staged = stage_root / rel
            stage_book = load_workbook(staged, data_only=False, read_only=True)
            stage_ws = stage_book.active
            original_book = None
            if name == "Unlock.xlsx":
                h = header_map(stage_ws)
                ids = set(range(12001, 12005)) | set(range(12101, 12411))
                selected = [r for r, row in enumerate(stage_ws.iter_rows(min_row=5), 5) if row[h["Id"] - 1].value in ids]
            elif name == "ShopGoods.xlsx":
                h = header_map(stage_ws)
                ids = set(range(1109011, 1109016))
                selected = [r for r, row in enumerate(stage_ws.iter_rows(min_row=5), 5) if row[h["Id"] - 1].value in ids]
            elif name == "Drop_dungeon.xlsx":
                h = header_map(stage_ws)
                selected = [r for r, row in enumerate(stage_ws.iter_rows(min_row=5), 5) if isinstance(row[h["DropId"] - 1].value, int) and 24000001 <= row[h["DropId"] - 1].value <= 24000100 and row[0].value != "##"]
            elif name == "UIlevel.xlsx":
                original_book = load_workbook(source, data_only=False, read_only=True)
                original_ws = original_book.active
                h = header_map(original_ws)
                selected = [r for r, row in enumerate(original_ws.iter_rows(min_row=5), 5) if isinstance(row[h["FallRewardDrop"] - 1].value, int) and 24001001 <= row[h["FallRewardDrop"] - 1].value <= 24001100]
            else:
                selected = active_rows(stage_ws, key)
            stage_book.close()
            if original_book is not None:
                original_book.close()
            write_mirror(name, source, staged, selected)
            manifest_rows.append((name, rel, len(selected), status, usage))
        make_manifest(manifest_rows)

    after = {path: sha256(path) for path in inputs}
    if before != after:
        changed = [str(path) for path in inputs if before[path] != after[path]]
        raise RuntimeError(f"正式 Datas 被意外修改: {changed}")
    print(f"[OK] 房车最终配置镜像已生成: {OUT}")
    print("[OK] 正式 Datas 输入哈希保持不变")


if __name__ == "__main__":
    main()
