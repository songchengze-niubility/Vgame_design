#!/usr/bin/env python3
"""更新映射装备06配置清单中的CHG-017正式模板重基记录。"""
from __future__ import annotations

import os
from pathlib import Path

from openpyxl import load_workbook


OUT = Path(os.environ.get("MAPPING_EQUIP_FINAL_OUT", ""))
INDEX = OUT / "00_配置清单.xlsx"
UPDATES = {
    "SpecialEquipment.xlsx": (
        "SpecialEquipment", "可粘贴：480条正式装备属性改造", 480,
        "CHG-017：按当前正式80001~80480重基；Effect改指向5001~5060品质等级池；Attribute2List清空。",
    ),
    "SpecialEquipmentMain.xlsx": (
        "SpecialEquipment", "可粘贴：190条主属性参数", 190,
        "主属性数值不变；480件装备引用171个在用Main ID，完整保留190条Main兼容定义。",
    ),
    "SpecialEquipmentEffect.xlsx": (
        "SpecialEquipment", "可粘贴：双曲线分品质等级副属性", 846,
        "CHG-018：保留246条历史定义；600条新Effect已预乘等级与品质系数，直接存最终副属性值。",
    ),
    "SpecialEquipmentEffectDrop.xlsx": (
        "SpecialEquipment", "可粘贴：60个品质等级十类副属性池", 253,
        "保留193条历史池，新增5001~5060；旧池退出新装备引用。",
    ),
    "SpecialEquipmentQuality.xlsx": (
        "SpecialEquipment", "可粘贴：品质系数与副属性条数", 5,
        "副属性品质系数0.35/0.50/0.65/0.85/1.00；已预乘进新Effect，字段保留作参数门槛/评分。",
    ),
    "SpecialEquipmentSubAttrCoef.xlsx": (
        "SpecialEquipmentSubAttrCoef", "可粘贴：12档副属性公共倍率", 12,
        "双曲线已预乘进Effect，12档SubAttrCoef统一1.00。",
    ),
}


def main():
    if not INDEX.exists():
        raise FileNotFoundError(INDEX)
    wb = load_workbook(INDEX)
    ws = wb[wb.sheetnames[0]]
    row_by_file = {ws.cell(row, 2).value: row for row in range(2, ws.max_row + 1)}
    missing = sorted(set(UPDATES) - set(row_by_file))
    if missing:
        raise ValueError(f"配置清单缺少文件：{missing}")
    for file_name, (sheet, status, count, note) in UPDATES.items():
        row = row_by_file[file_name]
        ws.cell(row, 3).value = sheet
        ws.cell(row, 4).value = status
        ws.cell(row, 5).value = count
        ws.cell(row, 7).value = note
    wb.save(INDEX)
    print({"index": str(INDEX), "updated": len(UPDATES), "rows": ws.max_row - 1})


if __name__ == "__main__":
    main()
