from __future__ import annotations

import argparse
import json
import os
import sys
from copy import copy
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
ECON_SCRIPTS = ROOT / "经济" / "scripts"
sys.path.insert(0, str(ECON_SCRIPTS))
from generate_trade_workbook import integer, optional_integer, parse_md, percent, percent_range, reward_specs, row_map  # noqa: E402


MD = ROOT / "03_沉淀MD" / "经济" / "跑商玩法数值设计文档-2026-07-06.md"
CONFIG_ROOT = Path(os.environ.get("VGAME_CONFIG_DATAS", ""))
OUT = Path(os.environ.get("TRADE_FINAL_OUT", str(ROOT / "06_最终配置" / "跑商")))
BATCH = "TRADE-20260716-F15"
YELLOW = "FFF2CC"
BLUE = "0000FF"
NAVY = "17365D"
WHITE = "FFFFFF"
SOURCE_NOTE = "Source: 03_沉淀MD/经济/跑商玩法数值设计文档-2026-07-06.md（V22）"
FORMAL_EVENT_SOURCE_NOTE = "Source: 正式 Datas/Trade/TradeEvent.xlsx（用户确认 2026-07-15）"
GOODS_ID_START = 910001
LOCATION_ID_START = 920001
LEVEL_ROW_ID_START = 940001
DROP_ROW_ID_START = 950001
EVENT_DROP_ROW_ID_START = 950101
LEVEL_REWARD_DROP_ID_START = 69090101
EVENT_REWARD_DROP_ID_START = 69090201
LEVEL_REWARD_SPECS = {
    2: (301, 3000), 3: (301, 4000), 4: (2, 50), 5: (106, 1), 6: (301, 5000),
    7: (2, 80), 8: (106, 1), 9: (301, 6000), 10: (106, 1), 11: (301, 7000),
    12: (2, 100), 13: (301, 8000), 14: (106, 1), 15: (2, 150), 16: (301, 9000),
    17: (106, 1), 18: (2, 200), 19: (301, 10000), 20: (106, 2),
}
TRADE_UNLOCK_IDS = {
    1: 910101, 2: 910102, 5: 910103, 8: 910104, 10: 910105,
    3: 910106, 6: 910107, 9: 910108, 14: 910109, 15: 910110,
    17: 910111, 18: 910112, 20: 910113,
}
SHOP_ID = 1110
SHOP_GOODS_ID_START = 1110001
NOTE_ITEM_IDS = {
    1075: "行程便签", 1076: "安途清单", 1077: "探路标记", 1078: "见闻手札", 1079: "调律灯片",
}
PROVISIONAL_ID_NOTE = "预留 ID：按用户授权临时生成，可在写入 SVN 前替换；不得与正式 ID 分配混用。"


def copy_cell(src, dst) -> None:
    dst.value = src.value
    if src.has_style:
        dst._style = copy(src._style)
    if src.number_format:
        dst.number_format = src.number_format
    dst.alignment = copy(src.alignment)
    dst.protection = copy(src.protection)


def copy_schema_rows(src_ws, dst_ws, rows: int = 4) -> None:
    for r in range(1, rows + 1):
        for c in range(1, src_ws.max_column + 1):
            copy_cell(src_ws.cell(r, c), dst_ws.cell(r, c))
    for key, dim in src_ws.column_dimensions.items():
        dst_ws.column_dimensions[key].width = dim.width
    for r in range(1, rows + 1):
        dst_ws.row_dimensions[r].height = src_ws.row_dimensions[r].height


def header_map(ws) -> dict[str, int]:
    return {str(ws.cell(1, c).value): c for c in range(2, ws.max_column + 1) if ws.cell(1, c).value}


def style_data_cell(cell, from_md: bool = True, pending: str | None = None, source_note: str = SOURCE_NOTE) -> None:
    cell.alignment = Alignment(vertical="top", wrap_text=True)
    if from_md and cell.value not in (None, ""):
        cell.font = Font(name="Microsoft YaHei", size=10, color=BLUE)
        cell.comment = Comment(source_note, "Codex")
    if pending:
        cell.fill = PatternFill("solid", fgColor=YELLOW)
        cell.comment = Comment(pending, "Codex")


def add_cn_info(wb: Workbook, file_name: str, status: str, formal_path: Path, rows: int, usage: str, blockers: str, reminder: str, data_source: str | None = None) -> None:
    ws = wb.create_sheet("中文说明")
    data = [
        ["项目", "内容"],
        ["配置镜像文件", file_name],
        ["交付批次", BATCH],
        ["落表状态", status],
        ["正式配置位置", str(formal_path)],
        ["数据来源", data_source or "03_沉淀MD/经济/跑商玩法数值设计文档-2026-07-06.md；正式 Datas 仅用于 Schema"],
        ["有效数据行", rows],
        ["使用方式", usage],
        ["重要提醒", reminder],
        ["阻断项", blockers],
    ]
    for r, row in enumerate(data, 1):
        for c, value in enumerate(row, 1):
            cell = ws.cell(r, c, value)
            cell.font = Font(name="Microsoft YaHei", size=10, bold=r == 1, color=WHITE if r == 1 else "222222")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if r == 1:
                cell.fill = PatternFill("solid", fgColor=NAVY)
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 105
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False
    for r in range(1, len(data) + 1):
        ws.row_dimensions[r].height = 30 if r == 1 else 40


def create_schema_mirror(source: Path, out: Path, rows: list[dict], status: str, blockers: str, usage: str, reminder: str, pending: dict[int, dict[str, str]] | None = None, source_note: str = SOURCE_NOTE, data_source: str | None = None) -> dict:
    src_wb = load_workbook(source, data_only=False)
    src_ws = src_wb.active
    wb = Workbook()
    ws = wb.active
    ws.title = src_ws.title
    copy_schema_rows(src_ws, ws)
    fields = header_map(src_ws)
    pending = pending or {}
    for r_idx, values in enumerate(rows, 5):
        ws.cell(r_idx, 1, "")
        for field, value in values.items():
            if field not in fields:
                raise ValueError(f"{source.name} 缺少字段 {field}")
            cell = ws.cell(r_idx, fields[field], value)
            style_data_cell(cell, source_note=source_note)
        for field, reason in pending.get(r_idx - 5, {}).items():
            if field not in fields:
                raise ValueError(f"{source.name} 缺少待确认字段 {field}")
            style_data_cell(ws.cell(r_idx, fields[field]), from_md=False, pending=reason)
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:{get_column_letter(src_ws.max_column)}{4 + len(rows)}"
    ws.sheet_view.showGridLines = False
    add_cn_info(wb, out.name, status, source, len(rows), usage, blockers, reminder, data_source)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return {"file": out.name, "sheet": ws.title, "rows": len(rows), "status": status, "formal": str(source)}


def create_globals(md_data: dict, out: Path) -> dict:
    source = CONFIG_ROOT / "globals_conf" / "globals_conf.xlsx"
    src_wb = load_workbook(source, data_only=False)
    src_ws = src_wb.active
    dispatch = row_map(*md_data["dispatch"], "项")
    currency = row_map(*md_data["currency"], "项")
    event_trigger = row_map(*md_data["event_trigger"], "项")
    event_refresh = [integer(part) for part in event_trigger["单次刷新数量"].split("、")[:2]]
    confirmed = {
        "TradeEventExistHours": integer(event_trigger["事件存在时间"]),
        "TradeEventMaxExistCount": integer(event_trigger["事件最多存在数量"]),
        "TradeRunEventRefreshMin": event_refresh[0],
        "TradeRunEventRefreshMax": event_refresh[1],
        "TradeDefaultTicketItemId": integer(currency["商券道具 ID"]),
        "TradeDefaultTradeLevel": 1,
        "TradeDefaultTicketCount": integer(dispatch["初始商券"]),
        "TradeRecycleParam": percent(currency["货物回收"]),
        "DefaultCargoCapacity": integer(dispatch["单次派遣载货量"]),
        "TradeTransportHours": integer(dispatch["常规派遣时间"]),
    }
    wanted = [
        "TradeEventExistHours", "TradeEventMaxExistCount", "TradeRunEventRefreshMin", "TradeRunEventRefreshMax",
        "TradeDefaultTicketItemId", "TradeDefaultTradeLevel", "TradeDefaultTicketCount", "TradeRecycleParam",
        "DefaultCargoCapacity", "TradeTransportHours",
    ]
    source_rows = {src_ws.cell(r, 1).value: r for r in range(1, src_ws.max_row + 1)}
    wb = Workbook()
    ws = wb.active
    ws.title = src_ws.title
    copy_schema_rows(src_ws, ws, 3)
    for idx, key in enumerate(wanted, 4):
        src_row = source_rows[key]
        for col in range(1, src_ws.max_column + 1):
            copy_cell(src_ws.cell(src_row, col), ws.cell(idx, col))
        # 未确认项必须显式清空，不能残留正式表中的临时测试值。
        ws.cell(idx, 4).value = confirmed.get(key)
        if key in confirmed:
            style_data_cell(ws.cell(idx, 4))
        else:
            style_data_cell(ws.cell(idx, 4), from_md=False, pending="待确认：沉淀 MD 未提供该参数")
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:E{3 + len(wanted)}"
    ws.sheet_view.showGridLines = False
    add_cn_info(
        wb, out.name, "可按参数名替换：10 项已确认", source, len(wanted),
        "只按 A 列参数名定位并替换 D 列确认值；不得整表覆盖正式 globals_conf。",
        "新手首单 10 分钟和每日 4:00 刷新仍缺少专项字段映射；持久事件的领取、过期和列表满额行为待程序/UI 确认。",
        "不得整表覆盖 globals_conf；本镜像只包含跑商相关参数切片。",
    )
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return {"file": out.name, "sheet": ws.title, "rows": len(wanted), "status": "可按参数名替换：10 项已确认", "formal": str(source)}


def create_trade_event_drops(md_data: dict, out: Path) -> dict:
    source = CONFIG_ROOT / "Drop" / "Drop_other.xlsx"
    rows = []
    for level, (item_id, item_count) in LEVEL_REWARD_SPECS.items():
        rows.append({
            "ID": DROP_ROW_ID_START + len(rows), "DropId": LEVEL_REWARD_DROP_ID_START + level - 2,
            "DropName": f"跑商-商会等级{level}级奖励", "DropId2": 0, "DropType": "fixed_",
            "DropMinNum": 1, "DropMaxNum": 1, "DropItem_id": item_id,
            "DropItemNumMin": item_count, "DropItemNumMax": item_count, "DropWeight": 10000,
            "ShowItemID": str(item_id),
        })
    event_headers, event_rows = md_data["events"]
    event_idx = {name: idx for idx, name in enumerate(event_headers)}
    next_event_row_id = EVENT_DROP_ROW_ID_START
    for event in event_rows:
        event_id = integer(event[event_idx["事件ID"]])
        for option in (1, 2):
            drop_id = integer(event[event_idx[f"DropId{option}"]])
            expected_drop_id = EVENT_REWARD_DROP_ID_START + (event_id - 1) * 2 + option - 1
            if drop_id != expected_drop_id:
                raise ValueError(f"事件 {event_id} 选项 {option} DropId 异常：{drop_id}")
            specs = reward_specs(event[event_idx[f"奖励{option}"]])
            drop_type = "random" if len(specs) > 1 else "fixed_"
            if len(specs) > 1 and sum(weight for _, _, weight in specs) != 10000:
                raise ValueError(f"事件 {event_id} 选项 {option} 随机权重不等于 10000")
            for item_id, item_count, weight in specs:
                rows.append({
                    "ID": next_event_row_id, "DropId": drop_id,
                    "DropName": f"跑商-事件{event_id}-选项{option}", "DropId2": 0, "DropType": drop_type,
                    "DropMinNum": 1, "DropMaxNum": 1, "DropItem_id": item_id,
                    "DropItemNumMin": item_count, "DropItemNumMax": item_count, "DropWeight": weight,
                    "ShowItemID": str(item_id),
                })
                next_event_row_id += 1
    pending = {index: {"ID": PROVISIONAL_ID_NOTE, "DropId": PROVISIONAL_ID_NOTE} for index in range(len(rows))}
    return create_schema_mirror(
        source, out,
        rows,
        "部分可粘贴：19 个等级奖励包与 50 个事件选项奖励包已输出",
        "预留行 ID、等级奖励 DropId 与事件奖励 DropId 写入 SVN 前均需做冲突检查并确认。",
        "按完整 DropId 奖励包复制；不得用本镜像覆盖正式 Drop_other。",
        "事件奖励使用专属 DropId 69090201-69090250；最后一包为 90% 元晶30 / 10% 旅券1，其余均为固定单项。",
        pending,
    )


def create_formal_trade_events(md_data: dict, out: Path) -> dict:
    source = CONFIG_ROOT / "Trade" / "TradeEvent.xlsx"
    headers, event_rows = md_data["events"]
    idx = {name: col for col, name in enumerate(headers)}
    rows = []
    for event in event_rows:
        rows.append({
            "Id": integer(event[idx["事件ID"]]), "Name": event[idx["事件名"]], "Desc": event[idx["事件描述"]],
            "Option1Title": event[idx["选项1"]], "Option2Title": event[idx["选项2"]],
            "UnlockId": optional_integer(event[idx["UnlockId"]]), "TradeEventType": event[idx["类型"]],
            "DropId1": integer(event[idx["DropId1"]]), "DropId2": integer(event[idx["DropId2"]]),
        })
    pending = {index: {"DropId1": PROVISIONAL_ID_NOTE, "DropId2": PROVISIONAL_ID_NOTE} for index in range(len(rows))}
    return create_schema_mirror(
        source, out, rows,
        "部分可粘贴：25 条正式事件内容与 V22 双选奖励引用已输出",
        "事件奖励 DropId 69090201-69090250 为预留值，写入 SVN 前需做冲突检查并确认。",
        "保留正式事件 Id/名称/描述/UnlockId；替换两个选项标题与 DropId 时须与 Drop_other 奖励包同步。",
        "每个事件固定两个互斥选项，只结算玩家所选 Drop；不再复用主线 DropId 10001。",
        pending,
        data_source="正式 Datas/Trade/TradeEvent.xlsx 提供事件内容；沉淀 MD V22 提供选项与奖励配置",
    )


def create_trade_unlocks(out: Path) -> dict:
    source = CONFIG_ROOT / "unlock" / "Unlock.xlsx"
    rows = [
        {"Id": unlock_id, "UnLockDesc": f"跑商等级达到 Lv.{level}", "UnLockType": "UnLock_UUT_Trade_Grade", "UnLockParams": level}
        for level, unlock_id in sorted(TRADE_UNLOCK_IDS.items())
    ]
    pending = {index: {"Id": PROVISIONAL_ID_NOTE} for index in range(len(rows))}
    return create_schema_mirror(
        source, out,
        rows,
        "部分可粘贴：跑商等级与商店商品解锁条件、预留 UnlockId 已输出",
        "预留 UnlockId 写入 SVN 前需替换；跑商系统入口的 SystemOpen 关联尚未纳入本期镜像。",
        "按 Id 整行复制；TradeGoodsMarket、TradeLocation、ShopGoods 已引用这些 UnlockId。正式 TradeEvent 使用既有 UnlockId 11001。不得用本镜像覆盖正式 Unlock。",
        "解锁类型使用项目既有枚举 UnLock_UUT_Trade_Grade，参数覆盖跑商商品所需等级 1/2/3/5/6/8/9/10/14/15/17/18/20。",
        pending,
    )


def create_trade_shop_items(out: Path) -> dict:
    source = CONFIG_ROOT / "items" / "Items.xlsx"
    src_wb = load_workbook(source, data_only=False)
    src_ws = src_wb.active
    fields = header_map(src_ws)
    source_rows = {
        src_ws.cell(row, fields["Id"]).value: row
        for row in range(5, src_ws.max_row + 1)
        if src_ws.cell(row, fields["Id"]).value in NOTE_ITEM_IDS
    }
    if set(source_rows) != set(NOTE_ITEM_IDS):
        raise ValueError(f"正式 Items 缺少便签 ItemId：{sorted(set(NOTE_ITEM_IDS) - set(source_rows))}")
    rows = []
    for item_id, name in NOTE_ITEM_IDS.items():
        source_row = source_rows[item_id]
        formal_name = src_ws.cell(source_row, fields["Name"]).value
        if formal_name != name:
            raise ValueError(f"便签 ItemId {item_id} 正式名称异常：{formal_name} != {name}")
        rows.append({
            field: src_ws.cell(source_row, col).value
            for field, col in fields.items()
            if src_ws.cell(source_row, col).value not in (None, "")
        })
    return create_schema_mirror(
        source, out, rows,
        "只读核对：复用正式 Items 已有 ItemId 1075-1079",
        "无；五个 ItemId 与正式名称均已核对存在。",
        "本文件仅作为正式 Items 道具切片核对，不新增、不覆盖正式 Items；ShopGoods 直接引用这些既有 ItemId。",
        "正式名称依次为行程便签、安途清单、探路标记、见闻手札、调律灯片。",
        data_source="正式 Datas/items/Items.xlsx；沉淀 MD V22 确认跑商商店引用关系",
    )


def create_trade_shop(out: Path) -> dict:
    source = CONFIG_ROOT / "Shop" / "Shop.xlsx"
    rows = [{
        "Id": SHOP_ID,
        "Name": "行商商店",
        "IsTab": None,
        "##": "行商商店",
        "IconPath": "Icon_tab_zahuopu_jingjichang.png",
        "RelateTab": 1101,
        "Sort": 3,
        "ViewPanel": "UIGroceriesShopPanel",
        "RefreshType": "SRT_None",
        "CurrencyShow": 301,
        "HiddenWhenSoldOut": False,
        "Unlock": 4314,
        "version": 2,
        "ShopType": "ShopType_None",
    }]
    return create_schema_mirror(
        source, out, rows,
        "可按 Id=1110 更新：行商商店入口与版本号已输出",
        "商店入口 UnlockId=4314 沿用正式表既有值；商品级跑商等级解锁由 ShopGoods.Unlock 承载。",
        "按 Id=1110 更新该行，不得新增第二个行商商店；商品种类更新后版本号由 1 调整为 2。",
        "复用正式配置已有行商商店，顶部货币为商券 ItemId=301；商店本体不统一刷新。",
    )


def create_trade_shop_goods(md_data: dict, out: Path) -> dict:
    source = CONFIG_ROOT / "Shop" / "ShopGoods.xlsx"
    limit_types = {"每周": "BuyLimit_Week", "每月": "BuyLimit_Month", "永久一次": "BuyLimit_Ever"}
    rows = []
    pending = {}
    for index, row in enumerate(md_data["shop"][1]):
        seq = integer(row[0])
        item_id = integer(row[2])
        unlock_level = integer(row[7])
        rows.append({
            "Id": SHOP_GOODS_ID_START + seq - 1,
            "Name": row[1],
            "ShopId": SHOP_ID,
            "##": "行商商店",
            "ItemId": item_id,
            "Num": integer(row[3]),
            "PriceType": "SGPT_Item",
            "Price": f"301,{integer(row[4])}",
            "LimitType": limit_types[row[5]],
            "LimitNum": integer(row[6]),
            "Sort": seq,
            "Hide": False,
            "Unlock": TRADE_UNLOCK_IDS[unlock_level],
            "HideRewardPop": False,
        })
        pending_fields = {
            "Id": PROVISIONAL_ID_NOTE,
            "Unlock": PROVISIONAL_ID_NOTE,
        }
        pending[index] = pending_fields
    return create_schema_mirror(
        source, out, rows,
        "部分可粘贴：18 条行商商店商品与预留商品 ID/UnlockId 已输出",
        "商品 ID 与跑商等级 UnlockId 为预留值；便签 ItemId 1075-1079 已确认复用正式 Items；后期第 4 个重复商品只保留预算，不生成配置行。",
        "正式落表时先删除或停用现有 ShopId=1110 临时测试商品，再写入本镜像 18 行；不得与旧测试行并存。",
        "价格统一消耗商券 ItemId=301；五种房车天赋材料引用正式 ItemId 1075-1079；旅券/收容凭证仅永久一次，攻略笔记Ⅰ/Ⅱ为周限，攻略笔记Ⅲ为月限。",
        pending,
    )


def build() -> dict:
    md_data = parse_md(MD)
    dispatch = row_map(*md_data["dispatch"], "项")
    currency = row_map(*md_data["currency"], "项")
    price = row_map(*md_data["price"], "项")
    buy_min, buy_max = percent_range(price["今日购入价"])
    normal_min, normal_max = percent_range(price["普通收购价"])
    urgent_min, urgent_max = percent_range(price["急需收购价"])

    OUT.mkdir(parents=True, exist_ok=True)
    results = []
    results.append(create_globals(md_data, OUT / "globals_conf.xlsx"))

    goods_rows = []
    goods_pending = {}
    goods_by_location: dict[str, list[str]] = {}
    goods_id_by_name: dict[str, int] = {}
    location_levels: dict[str, int] = {}
    for row in md_data["locations"][1]:
        lv = integer(row[0])
        location_levels[row[1]] = lv
        goods_by_location[row[1]] = [row[2], row[5], row[8]]
        for name_col, price_col, limit_col in ((2, 3, 4), (5, 6, 7), (8, 9, 10)):
            goods_id = GOODS_ID_START + len(goods_rows)
            goods_id_by_name[row[name_col]] = goods_id
            goods_rows.append({
                "Id": goods_id,
                "Name": row[name_col],
                "UnlockId": TRADE_UNLOCK_IDS[lv],
                "BasicPrice": integer(row[price_col]),
                "DailyBuyLimit": integer(row[limit_col]),
                "BuyRateMin": buy_min,
                "BuyRateMax": buy_max,
                "NormalSellRateMin": normal_min,
                "NormalSellRateMax": normal_max,
                "UrgentSellRateMin": urgent_min,
                "UrgentSellRateMax": urgent_max,
            })
            goods_pending[len(goods_rows) - 1] = {
                "Id": PROVISIONAL_ID_NOTE,
            }
    results.append(create_schema_mirror(
        CONFIG_ROOT / "Trade" / "TradeGoodsMarket.xlsx", OUT / "TradeGoodsMarket.xlsx", goods_rows,
        "部分可粘贴：预留 ID、UnlockId、基础价、基础限购和价格倍率已输出", "预留 ID 写入 SVN 前需替换",
        "只复制已确认的 Name/UnlockId/BasicPrice/DailyBuyLimit/六个价格倍率字段；预留 ID 确认前不可整行粘贴。",
        "基础 DailyBuyLimit 按同站低/中/高价货物配置为 25/20/15；实际可购买数量还受行商等级上浮影响。", goods_pending,
    ))

    benefit_delta = {
        3: ("DailyBuyLimitAddRate", 0.10), 6: ("SellPriceChangeRate", 0.04), 9: ("BuyPriceChangeRate", -0.03),
        11: ("DailyBuyLimitAddRate", 0.10), 12: ("SellPriceChangeRate", 0.04), 14: ("BuyPriceChangeRate", -0.03),
        15: ("SellPriceChangeRate", 0.04), 16: ("DailyBuyLimitAddRate", 0.10), 18: ("BuyPriceChangeRate", -0.04),
        20: ("SellPriceChangeRate", 0.03),
    }
    level_rows = []
    level_pending = {}
    for row in md_data["levels"][1]:
        lv = integer(row[0])
        benefit, param = benefit_delta.get(lv, (None, None))
        level_rows.append({
            "Id": LEVEL_ROW_ID_START + lv - 1, "TargetLv": lv, "Name": None, "Icon": None, "Description": row[7],
            "CostExp": integer(row[1]), "TradeBenefitType": benefit, "param": param,
            "DropId": None if lv == 1 else LEVEL_REWARD_DROP_ID_START + lv - 2, "DailyTicketCount": integer(row[4]),
        })
        pending = {
            "Id": PROVISIONAL_ID_NOTE,
        }
        if lv == 1:
            pending["DropId"] = "Lv.1 初始商券由 globals_conf.TradeDefaultTicketCount=30000 发放，不重复配置升级奖励包。"
        if benefit == "BuyPriceChangeRate":
            pending["param"] = "设计值按负增量输出；正式写入前必须确认运行时对 BuyPriceChangeRate 的符号处理"
        level_pending[len(level_rows) - 1] = pending
    results.append(create_schema_mirror(
        CONFIG_ROOT / "Trade" / "TradeLvUp.xlsx", OUT / "TradeLvUp.xlsx", level_rows,
        "部分可粘贴：预留 ID、等级阈值、每日商券、等级奖励、设计增益已输出", "购入优惠 param 的运行时符号需确认；预留 ID 写入 SVN 前需替换",
        "可核对 TargetLv/CostExp/DailyTicketCount；增益字段需程序确认符号后再粘贴；黄色字段不得写入。",
        "CostExp 按累计正利润阈值输出；商会补给写入 DailyTicketCount；每级只保留一个永久提升。", level_pending,
    ))

    location_rows = []
    location_pending = {}
    event_pools = [row[2] for row in md_data["event_pools"][1]]
    if len(event_pools) != 5:
        raise ValueError(f"V22 站点事件池数量异常：{len(event_pools)}")
    for offset, (name, goods) in enumerate(goods_by_location.items()):
        location_rows.append({
            "Id": LOCATION_ID_START + offset,
            "Name": name,
            "UnlockId": TRADE_UNLOCK_IDS[location_levels[name]],
            "DemandGoodsList": ";".join(str(goods_id_by_name[good]) for good in goods),
            "TradeEventList": event_pools[offset],
        })
        location_pending[len(location_rows) - 1] = {
            "Id": PROVISIONAL_ID_NOTE,
            "DemandGoodsList": f"预留货物 ID 映射：{'；'.join(goods)}；正式 ID 分配后需同步替换",
        }
    results.append(create_schema_mirror(
        CONFIG_ROOT / "Trade" / "TradeLocation.xlsx", OUT / "TradeLocation.xlsx", location_rows,
        "部分可粘贴：预留站点/货物引用与 V22 事件概率池已输出", "预留站点与货物 ID 写入 SVN 前需确认；事件池权重需整体替换",
        "可核对站点名称、UnlockId、急需池；TradeEventList 按沉淀 MD V22 写入。",
        "每站急需 3 个货物；事件池每站 5 条，权重 250/250/250/150/100。", location_pending,
    ))

    results.append(create_formal_trade_events(md_data, OUT / "TradeEvent.xlsx"))
    results.append(create_trade_event_drops(md_data, OUT / "Drop_other.xlsx"))
    results.append(create_trade_unlocks(OUT / "Unlock.xlsx"))
    results.append(create_trade_shop_items(OUT / "Items.xlsx"))
    results.append(create_trade_shop(OUT / "Shop.xlsx"))
    results.append(create_trade_shop_goods(md_data, OUT / "ShopGoods.xlsx"))

    manifest_rows = [
        [1, "globals_conf.xlsx", "globals_conf", results[0]["status"], results[0]["rows"], results[0]["formal"], "按参数名替换"],
        [2, "TradeGoodsMarket.xlsx", "TradeGoodsMarket", results[1]["status"], results[1]["rows"], results[1]["formal"], "预留ID与UnlockId已写入；当前 Schema 无 Icon 字段"],
        [3, "TradeLvUp.xlsx", "TradeLevel", results[2]["status"], results[2]["rows"], results[2]["formal"], "预留ID与等级奖励已写入；购入优惠符号待确认"],
        [4, "TradeLocation.xlsx", "TradeLocation", results[3]["status"], results[3]["rows"], results[3]["formal"], "TradeEventList 已按 V22 的 75%/15%/10% 品级概率同步"],
        [5, "TradeEvent.xlsx", "TradeEvent", results[4]["status"], results[4]["rows"], results[4]["formal"], "25 条事件全部双选互斥；两个选项均配置专属 DropId"],
        [6, "Drop_other.xlsx", "Drop", results[5]["status"], results[5]["rows"], results[5]["formal"], "19 个等级奖励包 + 50 个事件选项奖励包，共 70 行"],
        [7, "Unlock.xlsx", "Unlock", results[6]["status"], results[6]["rows"], results[6]["formal"], "预留解锁供货物、站点和商店商品使用；正式事件使用既有 UnlockId 11001"],
        [8, "Items.xlsx", "items", results[7]["status"], results[7]["rows"], results[7]["formal"], "复用正式 Items 已有 ItemId 1075-1079；本镜像仅作核对切片"],
        [9, "Shop.xlsx", "Shop", results[8]["status"], results[8]["rows"], results[8]["formal"], "复用并更新既有行商商店 Id=1110，版本号调整为 2"],
        [10, "ShopGoods.xlsx", "Goods", results[9]["status"], results[9]["rows"], results[9]["formal"], "18 条有效商品；旧 ShopId=1110 临时测试行须先删除或停用"],
    ]
    wb = Workbook()
    ws = wb.active
    ws.title = "配置清单"
    headers = ["序号", "镜像文件", "配置页签", "状态", "数据行数", "正式配置路径", "说明"]
    for c, value in enumerate(headers, 1):
        cell = ws.cell(1, c, value)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name="Microsoft YaHei", size=10, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for r, row in enumerate(manifest_rows, 2):
        for c, value in enumerate(row, 1):
            cell = ws.cell(r, c, value)
            cell.font = Font(name="Microsoft YaHei", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if "阻塞" in str(row[3]) or "未生成" in str(row[3]):
                cell.fill = PatternFill("solid", fgColor=YELLOW)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:G{1 + len(manifest_rows)}"
    ws.sheet_view.showGridLines = False
    for col, width in enumerate([8, 28, 22, 52, 12, 70, 44], 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    wb.save(OUT / "00_配置清单.xlsx")

    return {"status": "success", "output": str(OUT), "files": 11, "generated_mirrors": 10, "manifest_rows": len(manifest_rows), "batch": BATCH}


def main() -> None:
    global CONFIG_ROOT, OUT
    parser = argparse.ArgumentParser()
    parser.add_argument("--config-root", type=Path, default=CONFIG_ROOT)
    parser.add_argument("--out", type=Path, default=OUT)
    args = parser.parse_args()
    CONFIG_ROOT = args.config_root
    OUT = args.out
    print(json.dumps(build(), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
