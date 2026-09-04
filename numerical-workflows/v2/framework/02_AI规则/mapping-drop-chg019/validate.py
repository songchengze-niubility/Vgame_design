"""Read-only XLSX validation of CHG-019; never saves a workbook."""
import hashlib
import argparse
import json
from collections import defaultdict
from fractions import Fraction as F
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / '06_最终配置/映射装备'
REPORT = ROOT / '05_验算核算/mapping-equip-drop-preview-CHG019.json'
parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument('--out', type=Path, default=OUT, help='Candidate or archived CHG-019 snapshot')
parser.add_argument('--report', type=Path, default=ROOT/'05_验算核算/mapping-equip-drop-preview-CHG019-validation.json')
args = parser.parse_args()
OUT = args.out
generation = json.loads(REPORT.read_text(encoding='utf-8'))
BACKUP = Path(generation['backup'])
checks = []


def check(name, ok, detail=None):
    checks.append({'name': name, 'result': 'PASS' if ok else 'FAIL', 'detail': detail})


def rows(directory, file):
    ws = load_workbook(directory / file, data_only=False).worksheets[0]
    return [list(r) for r in ws.iter_rows(values_only=True)]


def group(data):
    result = defaultdict(list)
    for r in data[4:]:
        if r[1] is not None:
            result[r[1]].append(r)
    return result


before = group(rows(BACKUP, 'Drop_dungeon.xlsx'))
after = group(rows(OUT, 'Drop_dungeon.xlsx'))
check('Drop业务ID集合未变', set(before) == set(after))
check('Drop总行数696', sum(map(len, after.values())) == 696)
new_records = {(d, r[9]) for d, rs in after.items() for r in rs} - {(d, r[9]) for d, rs in before.items() for r in rs}
check('本次补18行或幂等重跑0行', len(new_records) in (0, 18), sorted(new_records))
for d, rs in after.items():
    if not 22101 <= d <= 22230 and not 22310 <= d <= 22330:
        check(f'非本次DropId{d}逐值不变', rs == before[d])
    check(f'DropId{d}数量范围合法', all(r[6] <= r[7] and r[10] <= r[11] for r in rs))
    check(f'DropId{d}权重为非负int32', all(isinstance(r[12], int) and 0 <= r[12] <= 2147483647 for r in rs))

items_before = {r[0]: r for r in rows(BACKUP, 'Items.xlsx')[4:] if r[0] is not None}
items = {r[0]: r for r in rows(OUT, 'Items.xlsx')[4:] if r[0] is not None}
for i, r in items_before.items():
    check(f'既有Item{i}全部字段未变', r == items[i])
new_ids = set(range(4403, 4448))
check('45个新UI Item全部存在', new_ids <= items.keys())
check('UI占位无发放引用', not any(r[9] in new_ids for rs in after.values() for r in rs))
check('UI占位不调用开箱', all(items[i][2] is None and items[i][4] is None for i in new_ids))
formal_items_path = (
    Path(__file__).resolve().parents[5].parent / "Vgame" / "Config" / "GameConfig" / "Datas" / "items" / "Items.xlsx"
)
formal_items = {r[1] for r in load_workbook(formal_items_path, data_only=True)['items'].iter_rows(min_row=5, values_only=True) if r[0] != '##'}
check('新增候选Item未占正式ID', not (new_ids & formal_items))
check('真实掉落Item引用完整', all(r[9] == 0 or r[9] in items or r[9] in formal_items for rs in after.values() for r in rs))
gen = {r[1]: r for r in rows(OUT, 'SpecialEquipmentGenerate.xlsx')[4:] if r[1] is not None}
ui_before = {r[1]: r for r in rows(BACKUP, 'UIlevel.xlsx')[4:] if r[1] is not None}
ui_rows = rows(OUT, 'UIlevel.xlsx')
ui = {r[1]: r for r in ui_rows[4:] if r[1] is not None}
check('UI表头不变', ui_rows[:4] == rows(BACKUP, 'UIlevel.xlsx')[:4])
check('150个UI LevelId完全保留', set(ui) == set(ui_before) and len(ui) == 150)
for l in range(1, 31):
    q = {((r[9] - 4201) % 5): F(r[12], sum(x[12] for x in after[22100+l])) for r in after[22100+l]}
    check(f'L{l}金箱曲线', q[3] == (F(0) if l < 10 else F(1, 100) + F(l-10, 400)))
    check(f'L{l}红箱曲线', q[4] == (F(0) if l < 20 else F(3, 1000) + F(3*(l-20), 2500)))
    oldq = {((r[9]-4201)%5): F(r[12], sum(x[12] for x in before[22100+l])) for r in before[22100+l]}
    check(f'L{l}绿蓝与原值相同', all(q[k] == oldq[k] for k in (0,1)))
    check(f'L{l}紫承接剩余', q[2] == 1-q[0]-q[1]-q[3]-q[4])
    ext, oldext = after[22200+l], before[22200+l]
    total, oldtotal = sum(r[12] for r in ext), sum(r[12] for r in oldext)
    miss = F(next(r[12] for r in ext if r[9] == 0), total)
    oldmiss = F(next(r[12] for r in oldext if r[9] == 0), oldtotal)
    check(f'L{l}额外箱触发率保持', miss == oldmiss)
    if l >= 10 and miss < 1:
        check(f'L{l}额外箱内品质精确匹配', all(F(r[12], total)/(1-miss) == q[(r[9]-4201)%5] for r in ext if r[9] != 0))
    if l >= 6:
        mat, oldmat = after[22300+l], before[22300+l]
        total = sum(r[12] for r in mat)
        hit = sum(r[12] for r in mat if r[9])
        oldtotal = sum(r[12] for r in oldmat)
        oldhit = sum(r[12] for r in oldmat if r[9])
        check(f'L{l}主材总掉率保持', F(hit,total) == F(oldhit,oldtotal))
        golden = F(0) if l < 10 else F(3,200) + F(13*(l-10),4000)
        red = F(0) if l < 20 else F(1,200)+F(l-20,1000)
        expected = [1-golden-red,golden,red]
        for k in range(3):
            check(f'L{l}主材品质{k}精确', F(sum(r[12] for r in mat if r[9] and (r[9]-1080)%3 == k), hit) == expected[k])
        for k, weight in enumerate([F(1,5),F(11,20),F(1,4)]):
            check(f'L{l}主材大类{k}比例', F(sum(r[12] for r in mat if r[9] and ((r[9]-1080)%9)//3 == k),hit)==weight)
    levels = sorted([int(s.split(',')[0]) for s in gen[4201+(l-1)*5][4].split(';') if int(s.split(',')[1]) > 0], reverse=True)
    qs = [k for k in range(4,-1,-1) if q[k] > 0]
    base, remain = divmod(len(qs), len(levels))
    allotted = [level for j,level in enumerate(levels) for _ in range(base+(j>=len(levels)-remain))]
    expected = list(zip(allotted, qs))
    for variant in range(1,6):
        lid = 250000+l*10+variant
        r, old = ui[lid], ui_before[lid]
        check(f'UI{lid}只改Des/Reward', all(a==b for c,(a,b) in enumerate(zip(r,old)) if c not in (15,16)))
        refs = [tuple(map(int,s.split(','))) for s in r[16].split(';')]
        check(f'UI{lid}全部引用存在且数量-1', all(i in items and qty==-1 for i,qty in refs))
        selected = [items[i] for i,_ in refs[:len(qs)]]
        actual = [(int(i[1].split('级')[0]),['Green','Blue','Purple','Orange','Red'].index(i[8])) for i in selected]
        check(f'UI{lid}品质等级顺序匹配', actual==expected)
        check(f'UI{lid}等级文案完整', all(str(g)+'级' in r[15].split('\n')[0] for g in levels))
        if l>=6:
            check(f'UI{lid}材料最高品质匹配', items[refs[-2][0]][8] == ('Red' if l>=20 else 'Orange' if l>=10 else 'Purple') and refs[-1][0]==4364)

for file in ['Drop_dungeon.xlsx','Items.xlsx','UIlevel.xlsx','00_配置清单.xlsx']:
    wb = load_workbook(OUT/file, data_only=True)
    errors = [(s.title,c.coordinate,c.value) for s in wb for row in s for c in row if c.data_type=='e']
    check(file+'缓存无公式错误', not errors, errors)
summary = load_workbook(OUT/'00_配置清单.xlsx',data_only=True)['CHG019概率展示']
check('核对页第30层终点', all(abs(summary[cell].value-value)<1e-12 for cell,value in [('E38',.06),('F38',.015),('I38',.08),('J38',.015)]))
report = {'revision':'CHG-019','scope':'06候选配置；非正式Datas导出', 'summary':{'checks':len(checks),'pass':sum(c['result']=='PASS' for c in checks),'fail':sum(c['result']=='FAIL' for c in checks)}, 'checks':checks, 'knownLimits':['未运行正式Datas导出或客户端实机验收。','首通保留本次前06镜像；与正式工作副本已有差异未纳入本次变更。','旧第1～9层额外箱池保留原整数取整；第10层起精确归一化。'], 'formal_items_sha256':hashlib.sha256(formal_items_path.read_bytes()).hexdigest()}
args.report.write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(report['summary'],ensure_ascii=False))
for c in checks:
    if c['result']=='FAIL': print(c)
raise SystemExit(bool(report['summary']['fail']))
