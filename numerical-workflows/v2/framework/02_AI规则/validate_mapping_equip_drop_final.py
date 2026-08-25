#!/usr/bin/env python3
from __future__ import annotations

import json
import os
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
OUT = Path(os.environ.get("MAPPING_EQUIP_FINAL_OUT", str(ROOT / "06_最终配置" / "映射装备")))
P02 = ROOT / "04_数值文档" / "配置粘贴包" / "映射装备全链路配置包_P02.xlsx"
CONFIG = Path(os.environ.get("VGAME_CONFIG_DATAS", ""))
REPORT = ROOT / "05_验算核算" / "mapping-equip-drop-final-CHG008-validation.json"


def values(ws, row: int, width: int) -> list:
    return [ws.cell(row, col).value for col in range(1, width + 1)]


def header_map(ws, row: int) -> dict[str, int]:
    result = {}
    for cell in ws[row]:
        if cell.value is None:
            continue
        for name in str(cell.value).splitlines():
            result[name.strip()] = cell.column
    return result


def var_map(ws) -> dict[str, int]:
    return {cell.value: cell.column for cell in ws[1] if cell.value}


def normalize_preview(value):
    if value in (None, ""):
        return value
    result = []
    for part in str(value).split(";"):
        item_id, quantity = (int(item) for item in part.split(","))
        result.append(f"{item_id},{-1 if quantity < 0 else quantity}")
    return ";".join(result)


def rows(ws, width: int) -> list[list]:
    return [
        values(ws, row, width)
        for row in range(5, ws.max_row + 1)
        if any(ws.cell(row, col).value is not None for col in range(1, width + 1))
    ]


def main() -> None:
    checks = []

    def check(name: str, ok: bool, actual=None, expected=None) -> None:
        checks.append({"name": name, "status": "PASS" if ok else "FAIL", "actual": actual, "expected": expected})

    p02 = load_workbook(P02, data_only=False)

    drop_path = OUT / "Drop_dungeon.xlsx"
    drop = load_workbook(drop_path, data_only=False)
    ds = drop["Drop"]
    drop_sources = ["25_Drop常规根110", "26_Drop宝箱池150", "27_Drop额外箱池", "28_Drop主材池", "29_Drop副材池4", "30A_首通Drop纯值"]
    width = p02[drop_sources[0]].max_column
    expected_drop = []
    sequence = 0
    for sheet_name in drop_sources:
        source = p02[sheet_name]
        for source_row in rows(source, width):
            drop_id = source_row[1]
            expected_drop.append(((0 if drop_id >= 10_000_000 else 1, drop_id, sequence), source_row))
            sequence += 1
    expected_drop = [row for _, row in sorted(expected_drop, key=lambda item: item[0])]
    actual_drop = rows(ds, width)
    check("Drop行数678", len(actual_drop) == 678, len(actual_drop), 678)
    check("Drop逐格等于P02", actual_drop == expected_drop, len(actual_drop), len(expected_drop))

    grouped = {}
    for row in actual_drop:
        grouped.setdefault(row[1], []).append(row)
    bad_weights = {
        drop_id: sum(row[12] for row in group)
        for drop_id, group in grouped.items()
        if 22101 <= drop_id <= 22453 and drop_id not in range(22301, 22306)
        and sum(row[12] for row in group) != 10000
    }
    check("随机池权重和10000", not bad_weights, bad_weights, {})
    check("固定箱池每层五品质", all(len(grouped.get(22100 + layer, [])) == 5 for layer in range(1, 31)))

    generate_path = OUT / "SpecialEquipmentGenerate.xlsx"
    generate = load_workbook(generate_path, data_only=False)
    gs = generate["SpecialEquipment"]
    formal_generate = load_workbook(CONFIG / "SpecialEquipment" / "SpecialEquipmentGenerate.xlsx", data_only=False)
    fgs = formal_generate["SpecialEquipment"]
    gwidth = max(col for col in range(1, fgs.max_column + 1) if any(fgs.cell(row, col).value is not None for row in range(1, 5)))
    check("Generate正式Schema", [values(gs, row, gwidth) for row in range(1, 5)] == [values(fgs, row, gwidth) for row in range(1, 5)])
    ph = header_map(p02["22_Generate150"], 4)
    p_generate = {
        p02["22_Generate150"].cell(row, ph["Id"]).value: [
            None,
            p02["22_Generate150"].cell(row, ph["Id"]).value,
            p02["22_Generate150"].cell(row, ph["##Name"]).value,
            None,
            p02["22_Generate150"].cell(row, ph["GradeList"]).value,
            p02["22_Generate150"].cell(row, ph["QualityList"]).value,
            p02["22_Generate150"].cell(row, ph["UnidentifiedList"]).value,
            p02["22_Generate150"].cell(row, ph["PartList"]).value,
            p02["22_Generate150"].cell(row, ph["SpecialEquipmentId"]).value,
            p02["22_Generate150"].cell(row, ph["GenerateMax"]).value,
            p02["22_Generate150"].cell(row, ph["GenerateMin"]).value,
        ]
        for row in range(5, p02["22_Generate150"].max_row + 1)
    }
    actual_generate = {row[1]: list(row) for row in gs.iter_rows(min_row=5, values_only=True) if isinstance(row[1], int)}
    generate_mismatch = [gid for gid in range(4201, 4351) if actual_generate.get(gid) != p_generate.get(gid)]
    check("150个宝箱Generate逐格等于P02", not generate_mismatch, generate_mismatch, [])

    ui_path = OUT / "UIlevel.xlsx"
    ui = load_workbook(ui_path, data_only=False)
    us = ui["level"]
    formal_ui = load_workbook(CONFIG / "level" / "UIlevel.xlsx", data_only=False)
    fus = formal_ui["level"]
    uwidth = max(col for col in range(1, fus.max_column + 1) if any(fus.cell(row, col).value is not None for row in range(1, 5)))
    check("UIlevel正式Schema", [values(us, row, uwidth) for row in range(1, 5)] == [values(fus, row, uwidth) for row in range(1, 5)])
    uc = var_map(us)
    fuc = var_map(fus)
    pc = header_map(p02["32_UIlevel150"], 4)
    proposal = {
        p02["32_UIlevel150"].cell(row, pc["LevelId"]).value: {
            field: (
                normalize_preview(p02["32_UIlevel150"].cell(row, pc[field]).value)
                if field in ("Reward", "FirstReward")
                else p02["32_UIlevel150"].cell(row, pc[field]).value
            )
            for field in ("Reward", "FirstReward", "FirstRewardDrop", "FallRewardDrop", "MultiSettlementTimes", "Vitality")
        }
        for row in range(5, p02["32_UIlevel150"].max_row + 1)
    }
    formal_rows = {fus.cell(row, fuc["LevelId"]).value: row for row in range(5, fus.max_row + 1)}
    ui_rows = {us.cell(row, uc["LevelId"]).value: row for row in range(5, us.max_row + 1) if us.cell(row, uc["LevelId"]).value is not None}
    target_fields = set(next(iter(proposal.values())))
    ui_target_bad = []
    ui_unrelated_bad = []
    for level_id, row in ui_rows.items():
        for field in target_fields:
            if us.cell(row, uc[field]).value != proposal[level_id][field]:
                ui_target_bad.append((level_id, field))
        formal_row = formal_rows[level_id]
        for field, col in uc.items():
            if field not in target_fields and us.cell(row, col).value != fus.cell(formal_row, fuc[field]).value:
                ui_unrelated_bad.append((level_id, field))
    check("UIlevel150行", len(ui_rows) == 150, len(ui_rows), 150)
    check("UIlevel目标字段等于P02", not ui_target_bad, ui_target_bad, [])
    check("UIlevel无关字段逐格保留", not ui_unrelated_bad, ui_unrelated_bad, [])
    check("每层最高装备箱预览为红色", all(str(us.cell(row, uc["Reward"]).value).split(";")[0] == "4354,-1" for row in ui_rows.values()))

    formal_items = load_workbook(CONFIG / "items" / "Items.xlsx", data_only=True, read_only=True)
    fis = formal_items["items"]
    item_ids = {row[1] for row in fis.iter_rows(min_row=5, values_only=True) if isinstance(row[1], int)}
    preview_ids = set()
    for row in ui_rows.values():
        for field in ("Reward", "FirstReward"):
            preview_ids.update(int(part.split(",")[0]) for part in str(us.cell(row, uc[field]).value or "").split(";") if part)
    check("UI预览Item正式闭环", preview_ids <= item_ids, sorted(preview_ids - item_ids), [])

    error_tokens = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A")
    for name, workbook in [("Drop", drop), ("Generate", generate), ("UIlevel", ui)]:
        errors = [
            (ws.title, cell.coordinate, cell.value)
            for ws in workbook.worksheets
            for row in ws.iter_rows()
            for cell in row
            if isinstance(cell.value, str) and cell.value.startswith(error_tokens)
        ]
        check(f"{name}无公式错误文本", not errors, errors, [])

    for workbook in (p02, drop, generate, formal_generate, ui, formal_ui, formal_items):
        workbook.close()

    result = {
        "output": str(OUT),
        "summary": {
            "checks": len(checks),
            "pass": sum(check["status"] == "PASS" for check in checks),
            "fail": sum(check["status"] == "FAIL" for check in checks),
        },
        "checks": checks,
    }
    REPORT.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(result["summary"], ensure_ascii=False))
    raise SystemExit(1 if result["summary"]["fail"] else 0)


if __name__ == "__main__":
    main()
