from __future__ import annotations

import argparse
import json
import zipfile
from pathlib import Path

from openpyxl import load_workbook


EXPECTED_SHEETS = [
    "说明", "全局参数", "基础规则", "价格规则", "站点货物", "货物明细", "行商等级", "等级解锁明细",
    "商品购买数量规则", "升级一次性奖励", "商会补给", "旅途事件", "商券商店", "商店预算测算",
    "收益与产销测算", "升级时间测算", "公式测算", "待确认项",
]


def validate(path: Path) -> dict:
    formulas_wb = load_workbook(path, data_only=False, read_only=False)
    values_wb = load_workbook(path, data_only=True, read_only=False)
    errors: list[str] = []
    formulas = []
    blank_cache = []
    formula_errors = []

    if formulas_wb.sheetnames != EXPECTED_SHEETS:
        errors.append("页签清单或顺序不正确")

    for ws in formulas_wb.worksheets:
        values_ws = values_wb[ws.title]
        for row in ws.iter_rows():
            for cell in row:
                if cell.data_type != "f":
                    continue
                formulas.append(f"{ws.title}!{cell.coordinate}")
                if "[" in str(cell.value):
                    errors.append(f"发现外部工作簿引用：{ws.title}!{cell.coordinate}")
                cached = values_ws[cell.coordinate].value
                if cached is None:
                    blank_cache.append(f"{ws.title}!{cell.coordinate}")
                if isinstance(cached, str) and cached.startswith("#"):
                    formula_errors.append(f"{ws.title}!{cell.coordinate}={cached}")

    if len(formulas) < 150:
        errors.append(f"公式数量过少：{len(formulas)}")
    if blank_cache:
        errors.append(f"存在未重算公式：{len(blank_cache)}")
    if formula_errors:
        errors.append(f"存在公式错误：{len(formula_errors)}")

    event = values_wb["旅途事件"]
    checks = values_wb["公式测算"]
    expected = {
        "正式旅途事件数量": (event["B38"].value, 25),
        "旅途事件商券期望": (event["B39"].value, 5055),
        "Lv1旅途事件商券期望": (event["B40"].value, 4455),
        "Lv20旅途事件商券期望": (event["B41"].value, 5655),
        "Lv20普通均利率": (checks["B7"].value, 0.375),
        "Lv20急需均利率": (checks["B8"].value, 1.2375),
        "Lv20毕业天数": (checks["B10"].value, 89),
        "Lv1五种便签周消耗": (checks["B11"].value, 250000),
        "Lv20当前重复兑换周消耗": (checks["B12"].value, 575000),
        "Lv20含预留周消耗": (checks["B13"].value, 725000),
        "Lv20含预留周结余": (checks["B14"].value, 86006),
        "一次性礼包库存总价": (checks["B15"].value, 3000000),
        "便签90天规划完成度": (checks["B16"].value, 0.65),
    }
    for name, (actual, target) in expected.items():
        if actual is None or abs(actual - target) > 1e-9:
            errors.append(f"{name}异常：{actual} != {target}")

    check_results = [checks[f"D{row}"].value for row in range(2, 17)]
    if any(value != "通过" for value in check_results):
        errors.append(f"公式测算未全部通过：{check_results}")

    event_ids = [event.cell(row, 1).value for row in range(2, 27)]
    if event_ids != list(range(1, 26)):
        errors.append(f"正式事件 ID 异常：{event_ids}")
    event_weights = [event.cell(row, 12).value for row in range(2, 27)]
    if event_weights != [250, 250, 250, 150, 100] * 5:
        errors.append(f"正式事件权重异常：{event_weights}")
    event_probs = [event.cell(row, 13).value for row in range(2, 27)]
    if event_probs != [0.25, 0.25, 0.25, 0.15, 0.10] * 5:
        errors.append(f"正式事件概率异常：{event_probs}")
    if [event.cell(row, 10).value for row in range(2, 27)] != list(range(69090201, 69090250, 2)):
        errors.append("正式事件 DropId1 异常")
    if [event.cell(row, 11).value for row in range(2, 27)] != list(range(69090202, 69090251, 2)):
        errors.append("正式事件 DropId2 异常")
    if any(not event.cell(row, 7).value for row in range(2, 27)):
        errors.append("正式事件存在空选项2")
    if [event.cell(row, 14).value for row in range(2, 27)] != [301] * 25:
        errors.append("正式事件选项1未全部配置商券")
    expected_reward1_qty = [6500, 7000, 7500, 8500, 9000, 7000, 7500, 8000, 9000, 9500, 7500, 8000, 8500, 9500, 10000, 8000, 8500, 9000, 10000, 10500, 8500, 9000, 9500, 10500, 11000]
    if [event.cell(row, 15).value for row in range(2, 27)] != expected_reward1_qty:
        errors.append("正式事件选项1商券数量异常")
    if event["P26"].value != "2,30,9000;106,1,1000":
        errors.append(f"海港高级事件随机奖励异常：{event['P26'].value}")
    if [event.cell(row, 13).value for row in range(31, 36)] != [1000] * 5:
        errors.append("正式站点事件池权重合计异常")
    if [event.cell(row, 14).value for row in range(31, 36)] != [0.75] * 5:
        errors.append("正式站点普通事件概率异常")
    if [event.cell(row, 15).value for row in range(31, 36)] != [0.15] * 5:
        errors.append("正式站点中级事件概率异常")
    if [event.cell(row, 16).value for row in range(31, 36)] != [0.10] * 5:
        errors.append("正式站点高级事件概率异常")
    if [event.cell(row, 17).value for row in range(31, 36)] != [7425, 7925, 8425, 8925, 9425]:
        errors.append("站点选项1商券加权均值异常")
    if [event.cell(row, 18).value for row in range(31, 36)] != [0.6] * 5:
        errors.append("商券选项测算选择率异常")
    if [event.cell(row, 19).value for row in range(31, 36)] != [4455, 4755, 5055, 5355, 5655]:
        errors.append("站点事件商券期望异常")

    global_ws = values_wb["全局参数"]
    global_values = {global_ws.cell(row, 1).value: global_ws.cell(row, 3).value for row in range(2, global_ws.max_row + 1)}
    confirmed_globals = {
        "TR-N01": 30000,
        "TR-N16": 30,
        "TR-N21": 301,
        "TR-N22": 18,
        "TR-N23": 2,
        "TR-N24": 1,
        "TR-N25": 1,
    }
    for param_id, target in confirmed_globals.items():
        if global_values.get(param_id) != target:
            errors.append(f"{param_id}异常：{global_values.get(param_id)} != {target}")

    goods_ws = values_wb["站点货物"]
    goods_prices = [goods_ws.cell(row, col).value for row in range(2, 7) for col in (4, 7, 10)]
    expected_goods_prices = [900, 1000, 1200, 1300, 1500, 1700, 1600, 2000, 2100, 1100, 1900, 2300, 1400, 1800, 2100]
    if goods_prices != expected_goods_prices:
        errors.append(f"货物基础价异常：{goods_prices}")

    daily_tickets = [values_wb["行商等级"].cell(row, 5).value for row in range(2, 22)]
    expected_daily_tickets = [13000, 13000, 13000, 16000, 16000, 16000, 20000, 20000, 20000, 20000, 20000, 20000, 25000, 25000, 25000, 25000, 32000, 32000, 40000, 40000]
    if daily_tickets != expected_daily_tickets:
        errors.append(f"商会补给异常：{daily_tickets}")

    reward_tickets = [values_wb["升级一次性奖励"].cell(row, 4).value for row in range(2, 22)]
    expected_reward_tickets = [30000, 3000, 4000, 0, 0, 5000, 0, 0, 6000, 0, 7000, 0, 8000, 0, 0, 9000, 0, 0, 10000, 0]
    if reward_tickets != expected_reward_tickets:
        errors.append(f"升级商券奖励异常：{reward_tickets}")

    buy_rules = values_wb["商品购买数量规则"]
    for label, base_col, formula_col, md_col, base, expected_values in [
        ("低价", "C", "D", "E", 25, [25, 28, 30, 33]),
        ("中价", "F", "G", "H", 20, [20, 22, 24, 26]),
        ("高价", "I", "J", "K", 15, [15, 17, 18, 20]),
    ]:
        bases = [buy_rules[f"{base_col}{row}"].value for row in range(2, 6)]
        formula_values = [buy_rules[f"{formula_col}{row}"].value for row in range(2, 6)]
        md_values = [buy_rules[f"{md_col}{row}"].value for row in range(2, 6)]
        if bases != [base] * 4:
            errors.append(f"{label}货物基础购买值异常：{bases}")
        if formula_values != expected_values or md_values != expected_values:
            errors.append(f"{label}货物实际购买数量异常：公式={formula_values}，MD={md_values}")

    shop_ws = values_wb["商券商店"]
    shop_item_ids = [shop_ws.cell(row, 3).value for row in range(2, 20)]
    expected_shop_item_ids = [1075, 1076, 1077, 1078, 1079, 106, 107, 106, 107, 106, 107, 34, 106, 107, 35, 106, 107, 36]
    if shop_item_ids != expected_shop_item_ids:
        errors.append(f"商券商店 ItemId 异常：{shop_item_ids}")
    shop_prices = [shop_ws.cell(row, 5).value for row in range(2, 20)]
    expected_shop_prices = [10000] * 5 + [300000] * 6 + [100000, 300000, 300000, 150000, 300000, 300000, 300000]
    if shop_prices != expected_shop_prices:
        errors.append(f"商券商店价格异常：{shop_prices}")
    shop_limits = [(shop_ws.cell(row, 6).value, shop_ws.cell(row, 7).value) for row in range(2, 20)]
    expected_shop_limits = [("每周", 5)] * 5 + [("永久一次", 1)] * 6 + [("每周", 1), ("永久一次", 1), ("永久一次", 1), ("每周", 1), ("永久一次", 1), ("永久一次", 1), ("每月", 1)]
    if shop_limits != expected_shop_limits:
        errors.append(f"商券商店限购异常：{shop_limits}")
    shop_unlocks = [shop_ws.cell(row, 8).value for row in range(2, 20)]
    expected_shop_unlocks = [1, 1, 1, 1, 1, 2, 3, 5, 6, 8, 9, 10, 14, 15, 15, 17, 18, 20]
    if shop_unlocks != expected_shop_unlocks:
        errors.append(f"商券商店解锁异常：{shop_unlocks}")

    shop_budget = values_wb["商店预算测算"]
    recurring = [shop_budget.cell(row, 4).value for row in range(2, 6)]
    if recurring != [250000, 350000, 500000, 575000]:
        errors.append(f"商店重复兑换周消耗异常：{recurring}")
    if shop_budget["E2"].value != 3337:
        errors.append(f"Lv1 便签周预算结余异常：{shop_budget['E2'].value}")
    if shop_budget["H5"].value != 725000 or shop_budget["I5"].value != 86006:
        errors.append(f"商店预留预算异常：H5={shop_budget['H5'].value}, I5={shop_budget['I5'].value}")
    if shop_budget["B17"].value != 0.65:
        errors.append(f"便签90天规划完成度异常：{shop_budget['B17'].value}")

    thresholds = [0, 15000, 45000, 75000, 105000, 185000, 270000, 355000, 445000, 590000, 740000, 995000, 1290000, 1685000, 2130000, 2620000, 3190000, 3800000, 4455000, 5115000]
    stage_exp = [thresholds[i] - thresholds[i - 1] for i in range(1, len(thresholds))]
    level_thresholds = [values_wb["行商等级"].cell(row, 2).value for row in range(2, 22)]
    timeline_thresholds = [values_wb["升级时间测算"].cell(row, 2).value for row in range(2, 21)]
    if level_thresholds != thresholds or timeline_thresholds != thresholds[1:]:
        errors.append(f"升级阈值异常：等级={level_thresholds}，时间={timeline_thresholds}")
    if any(stage_exp[i] < stage_exp[i - 1] for i in range(1, len(stage_exp))):
        errors.append(f"本级所需正利润存在倒挂：{stage_exp}")

    target_days = [1, 2, 3, 4, 6, 8, 10, 12, 15, 18, 23, 28, 35, 43, 51, 60, 69, 79, 89]
    timeline_days = [values_wb["升级时间测算"].cell(row, 3).value for row in range(2, 21)]
    if timeline_days != target_days:
        errors.append(f"升级目标日异常：{timeline_days} != {target_days}")

    info = values_wb["说明"]
    if info["B3"].value != "V22":
        errors.append(f"说明页版本异常：{info['B3'].value} != V22")
    if "03_沉淀MD" not in str(info["B6"].value):
        errors.append("说明页未指向正式沉淀 MD")
    if "04_数值文档" not in str(info["B7"].value):
        errors.append("说明页未指向正式发布路径")
    if "货物基础价" not in str(info["B10"].value) or "25/20/15" not in str(info["B10"].value):
        errors.append("说明页未记录 V22 货物基础价与基础购买量口径")

    with zipfile.ZipFile(path) as archive:
        external_parts = [name for name in archive.namelist() if "externalLink" in name]
    if external_parts:
        errors.append(f"发现外部链接部件：{external_parts}")

    return {
        "status": "success" if not errors else "errors_found",
        "file": str(path),
        "sheets": len(formulas_wb.sheetnames),
        "formulas": len(formulas),
        "blank_formula_cache": len(blank_cache),
        "formula_errors": formula_errors,
        "external_links": external_parts,
        "key_values": {name: actual for name, (actual, _) in expected.items()},
        "checks": check_results,
        "errors": errors,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    args = parser.parse_args()
    result = validate(args.workbook)
    print(json.dumps(result, ensure_ascii=False, indent=2))
    raise SystemExit(0 if result["status"] == "success" else 1)


if __name__ == "__main__":
    main()
