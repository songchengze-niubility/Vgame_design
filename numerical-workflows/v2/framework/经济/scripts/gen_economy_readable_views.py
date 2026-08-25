# -*- coding: utf-8 -*-
"""
gen_economy_readable_views.py - P4g 资源养成阅读版

从机器校验用的 `全资源产销总表.xlsx` 派生一份给人读的工作簿：
- 资源卡片：每个资源从哪来、花到哪、是否还有奖励包引用
- 养成消耗：角色/专武/旅券/映射装备/炮台/共鸣等养成阶段视图
- 玩法产出：玩法按周期和条件聚合后的奖励/消耗视图
- 商店兑换：把买到什么、花什么钱、限制写在一行
- 奖励包引用：Drop/Reward 先按系统/分组汇总，不展开随机包
"""
import io
import os
import re
import sys
from collections import Counter, defaultdict

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

HERE = os.path.dirname(os.path.abspath(__file__))
ECONOMY_DIR = os.path.dirname(HERE)
SUMMARY = os.path.join(ECONOMY_DIR, "全资源产销总表.xlsx")
MODULE_DIR = os.path.join(ECONOMY_DIR, "模块产销表")
OUT = os.path.join(ECONOMY_DIR, "资源养成_阅读版.xlsx")
REPORT = os.path.join(ECONOMY_DIR, "_reports", "P4g_资源养成阅读版.md")

LEDGER_SHEET = "追溯_全明细"
GROWTH_SYSTEMS = {"角色养成", "专武养成", "旅券养成", "映射装备养成", "炮台养成", "房车养成", "共鸣装备"}
SHOP_BASIS = "商店兑换商品"
REF_DIR = "引用"

HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
SUB_FILL = PatternFill("solid", fgColor="D9EAF7")
OK_FILL = PatternFill("solid", fgColor="E2F0D9")
WARN_FILL = PatternFill("solid", fgColor="FFF2CC")
THIN = Side(style="thin", color="D9E2F3")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def clean(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def to_num(value):
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return value
    text = clean(value)
    if re.fullmatch(r"-?\d+(\.\d+)?", text):
        num = float(text)
        return int(num) if num.is_integer() else num
    return None


def fmt_num(value):
    num = to_num(value)
    if num is None:
        return clean(value)
    if isinstance(num, float) and not num.is_integer():
        return ("%.2f" % num).rstrip("0").rstrip(".")
    return str(int(num))


def item_text(name, item_id="", qty=""):
    name = clean(name)
    item_id = clean(item_id)
    qty = clean(qty)
    base = name or ("ItemId %s" % item_id if item_id else "")
    if item_id and item_id != "None":
        base = "%s(%s)" % (base, item_id)
    if qty and qty != "None":
        base = "%s x%s" % (base, fmt_num(qty))
    return base


def compact_list(values, limit=8):
    values = [v for v in values if clean(v)]
    if not values:
        return ""
    counts = Counter(values)
    ordered = []
    for value, count in counts.most_common():
        ordered.append("%s%s" % (value, " ×%s" % count if count > 1 else ""))
    if len(ordered) > limit:
        ordered = ordered[:limit] + ["另%s项" % (len(values) - limit)]
    return "；".join(ordered)


def compact_systems(rows, direction=None, limit=6):
    pairs = []
    for row in rows:
        if direction and row.get("方向") != direction:
            continue
        system = clean(row.get("系统"))
        group = clean(row.get("分组"))
        period = clean(row.get("周期/限制"))
        label = system
        if group:
            label += "/" + group
        if period:
            label += "（%s）" % period
        pairs.append(label)
    return compact_list(pairs, limit)


def category_of(row):
    name = clean(row.get("资源/道具"))
    item_id = clean(row.get("道具ID"))
    system = clean(row.get("系统"))
    text = "%s %s %s" % (name, item_id, system)
    if any(k in text for k in ["钻石", "金币", "体力", "储能", "元晶"]):
        return "基础货币/体力"
    if any(k in text for k in ["抽卡", "旅券", "收容凭证", "招募", "票"]):
        return "抽卡/角色获取"
    if any(k in text for k in ["经验", "升级", "材料", "元件", "星盘", "突破", "宝石", "装备", "炮台", "房车", "便签指南", "映射之匙"]):
        return "养成材料"
    if any(k in text for k in ["代币", "勋章", "积分", "贡献", "货币"]):
        return "外循环代币"
    if clean(row.get("方向")) == REF_DIR:
        return "奖励包/配置引用"
    return "其他资源"


def risk_tags(rows):
    text = " ".join(clean(v) for row in rows for v in row.values())
    tags = []
    if any(k in text for k in ["钻石", "抽卡", "元晶", "付费"]):
        tags.append("高敏感")
    if any(k in text for k in ["每日", "每周", "赛季", "BuyLimit", "限购", "可重复"]):
        tags.append("周期/重复")
    if any(clean(row.get("方向")) == REF_DIR for row in rows):
        tags.append("需展开Drop/Reward")
    if any(k in text for k in ["随机", "权重", "Drop", "RewardId"]):
        tags.append("随机/引用")
    return "；".join(tags) if tags else "常规"


def load_ledger():
    wb = openpyxl.load_workbook(SUMMARY, read_only=True, data_only=True)
    ws = wb[LEDGER_SHEET]
    headers = [clean(c.value) for c in ws[1]]
    rows = []
    for values in ws.iter_rows(min_row=2, values_only=True):
        row = {headers[i]: values[i] if i < len(values) else "" for i in range(len(headers))}
        rows.append(row)
    wb.close()
    return rows


def add_sheet(wb, title, headers, rows, widths=None):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    for row in rows:
        ws.append([row.get(h, "") if isinstance(row, dict) else row[i] if i < len(row) else "" for i, h in enumerate(headers)])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = BORDER
    widths = widths or [14] * len(headers)
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    return ws


def build_resource_cards(rows):
    grouped = defaultdict(list)
    for row in rows:
        key = (clean(row.get("资源/道具")), clean(row.get("道具ID")))
        grouped[key].append(row)
    out = []
    for (name, item_id), group_rows in grouped.items():
        if not name:
            continue
        dirs = Counter(clean(r.get("方向")) for r in group_rows)
        qty_by_dir = defaultdict(float)
        qty_state = defaultdict(lambda: True)
        val_by_dir = defaultdict(float)
        for row in group_rows:
            direction = clean(row.get("方向"))
            qty = to_num(row.get("数量"))
            if qty is None:
                qty_state[direction] = False
            else:
                qty_by_dir[direction] += qty
            val = to_num(row.get("价值(钻石)"))
            if val is not None:
                val_by_dir[direction] += val
        def qty_text(direction):
            if dirs[direction] == 0:
                return ""
            if qty_state[direction]:
                return fmt_num(qty_by_dir[direction])
            return "含区间/非数值"
        out.append({
            "资源分类": category_of(group_rows[0]),
            "资源/道具": name,
            "道具ID": item_id,
            "产出记录": dirs["产出"],
            "消耗记录": dirs["消耗"],
            "引用记录": dirs["引用"],
            "产出数量": qty_text("产出"),
            "消耗数量": qty_text("消耗"),
            "产出价值(钻石)": fmt_num(val_by_dir["产出"]) if val_by_dir["产出"] else "",
            "消耗价值(钻石)": fmt_num(val_by_dir["消耗"]) if val_by_dir["消耗"] else "",
            "主要来源": compact_systems(group_rows, "产出"),
            "主要消耗": compact_systems(group_rows, "消耗"),
            "引用位置": compact_systems(group_rows, "引用"),
            "阅读提示": risk_tags(group_rows),
        })
    return sorted(out, key=lambda r: (r["资源分类"], -int(r["产出记录"] or 0) - int(r["消耗记录"] or 0), r["资源/道具"]))


def summarize_items(rows, direction, limit=8):
    grouped = defaultdict(lambda: {"qty": 0.0, "numeric": True, "count": 0, "name": "", "id": ""})
    for row in rows:
        if clean(row.get("方向")) != direction:
            continue
        key = (clean(row.get("资源/道具")), clean(row.get("道具ID")))
        rec = grouped[key]
        rec["name"], rec["id"] = key
        rec["count"] += 1
        qty = to_num(row.get("数量"))
        if qty is None:
            rec["numeric"] = False
        else:
            rec["qty"] += qty
    ordered = sorted(grouped.values(), key=lambda r: (-r["count"], r["name"]))
    parts = []
    for rec in ordered[:limit]:
        label = item_text(rec["name"], rec["id"])
        if rec["numeric"] and rec["qty"]:
            label += " 合计x%s" % fmt_num(rec["qty"])
        else:
            label += " %s条" % rec["count"]
        parts.append(label)
    if len(ordered) > limit:
        parts.append("另%s项" % (len(ordered) - limit))
    return "；".join(parts)


def build_module_overview(rows):
    buckets = defaultdict(list)
    for row in rows:
        buckets[clean(row.get("系统"))].append(row)
    out = []
    for system, group_rows in buckets.items():
        if not system:
            continue
        dirs = Counter(clean(r.get("方向")) for r in group_rows)
        groups = compact_list([clean(r.get("分组")) for r in group_rows if clean(r.get("分组"))], 10)
        periods = compact_list([clean(r.get("周期/限制")) for r in group_rows if clean(r.get("周期/限制"))], 8)
        sources = compact_list([clean(r.get("来源页")) for r in group_rows if clean(r.get("来源页"))], 8)
        out.append({
            "系统": system,
            "阅读定位": "养成消耗" if system in GROWTH_SYSTEMS else ("商店兑换" if system == "商店" else "玩法/外循环"),
            "记录数": len(group_rows),
            "产出": dirs["产出"],
            "消耗": dirs["消耗"],
            "引用": dirs["引用"],
            "主要分组": groups,
            "主要产出": summarize_items(group_rows, "产出"),
            "主要消耗": summarize_items(group_rows, "消耗"),
            "奖励包/引用": summarize_items(group_rows, REF_DIR),
            "周期/限制": periods,
            "来源页": sources,
            "阅读提示": risk_tags(group_rows),
        })
    return sorted(out, key=lambda r: (r["阅读定位"], -int(r["记录数"]), r["系统"]))


def aggregate_by_condition(rows, systems=None, exclude_systems=None):
    buckets = defaultdict(list)
    for row in rows:
        system = clean(row.get("系统"))
        if systems and system not in systems:
            continue
        if exclude_systems and system in exclude_systems:
            continue
        key = (
            system,
            clean(row.get("分组")),
            clean(row.get("周期/限制")),
            clean(row.get("进度/条件")),
            clean(row.get("来源页")),
            clean(row.get("来源行")),
        )
        buckets[key].append(row)
    out = []
    for key, group_rows in buckets.items():
        system, group, period, condition, source, source_row = key
        rewards = [item_text(r.get("资源/道具"), r.get("道具ID"), r.get("数量")) for r in group_rows if clean(r.get("方向")) == "产出"]
        costs = [item_text(r.get("资源/道具"), r.get("道具ID"), r.get("数量")) for r in group_rows if clean(r.get("方向")) == "消耗"]
        refs = [item_text(r.get("资源/道具"), r.get("道具ID"), r.get("数量")) for r in group_rows if clean(r.get("方向")) == REF_DIR]
        notes = [clean(r.get("备注")) for r in group_rows if clean(r.get("备注"))]
        out.append({
            "系统": system,
            "分组": group,
            "周期/限制": period,
            "条件/阶段": condition,
            "产出/获得": compact_list(rewards, 10),
            "消耗/花费": compact_list(costs, 10),
            "奖励包/引用": compact_list(refs, 8),
            "来源": "%s%s" % (source, ":%s" % source_row if source_row else ""),
            "备注": compact_list(notes, 3),
        })
    return sorted(out, key=lambda r: (r["系统"], r["分组"], r["周期/限制"], r["条件/阶段"], r["来源"]))


def build_shop_rows(rows):
    out = []
    for row in rows:
        if clean(row.get("口径")) != SHOP_BASIS or clean(row.get("方向")) != "产出":
            continue
        out.append({
            "系统": clean(row.get("系统")),
            "商店/分组": clean(row.get("分组")),
            "买到": item_text(row.get("资源/道具"), row.get("道具ID"), row.get("数量")),
            "花费": item_text(row.get("货币/成本"), "", row.get("货币数量")),
            "货币价值(钻石)": fmt_num(row.get("货币价值(钻石)")),
            "折扣/倍率": clean(row.get("折扣/倍率")),
            "限制": clean(row.get("周期/限制")),
            "条件": clean(row.get("进度/条件")),
            "来源": "%s:%s" % (clean(row.get("来源页")), clean(row.get("来源行"))),
            "备注": clean(row.get("备注")),
        })
    return sorted(out, key=lambda r: (r["系统"], r["商店/分组"], r["限制"], r["买到"]))


def build_reference_rows(rows):
    buckets = defaultdict(list)
    for row in rows:
        if clean(row.get("方向")) != REF_DIR:
            continue
        key = (clean(row.get("系统")), clean(row.get("分组")), clean(row.get("来源页")))
        buckets[key].append(row)
    out = []
    for (system, group, source), group_rows in buckets.items():
        refs = [item_text(r.get("资源/道具"), r.get("道具ID"), r.get("数量")) for r in group_rows]
        conditions = [clean(r.get("进度/条件")) for r in group_rows if clean(r.get("进度/条件"))]
        notes = [clean(r.get("备注")) for r in group_rows if clean(r.get("备注"))]
        out.append({
            "系统": system,
            "分组": group,
            "来源": source,
            "引用数量": len(group_rows),
            "引用示例": compact_list(refs, 10),
            "条件示例": compact_list(conditions, 6),
            "处理建议": "后续用 Drop/Reward 展开脚本拆包；本页只做人读索引。",
            "备注": compact_list(notes, 4),
        })
    return sorted(out, key=lambda r: (-int(r["引用数量"]), r["系统"], r["分组"]))


def build_nav(rows, module_rows, resource_cards, growth_rows, gameplay_rows, shop_rows, ref_rows):
    dirs = Counter(clean(r.get("方向")) for r in rows)
    systems = Counter(clean(r.get("系统")) for r in rows)
    return [
        ["产物", OUT],
        ["来源", SUMMARY],
        ["用途", "给策划/数值阅读；机器校验仍以 全资源产销总表.xlsx 为准"],
        ["总明细", "%s 行" % len(rows)],
        ["模块总览", "%s 个系统" % len(module_rows)],
        ["资源卡片", "%s 个资源/道具" % len(resource_cards)],
        ["养成视图", "%s 行" % len(growth_rows)],
        ["玩法视图", "%s 行" % len(gameplay_rows)],
        ["商店视图", "%s 行" % len(shop_rows)],
        ["奖励引用汇总", "%s 组" % len(ref_rows)],
        ["方向分布", "；".join("%s:%s" % (k, v) for k, v in dirs.items())],
        ["Top系统", "；".join("%s:%s" % (k, v) for k, v in systems.most_common(8))],
        ["阅读顺序", "先看 模块总览，再看 资源卡片；需要细节时再进 养成消耗/玩法产出/商店兑换"],
    ]


def style_nav(ws):
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 120
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = BORDER
    for cell in ws["A"]:
        cell.font = Font(bold=True)
        cell.fill = SUB_FILL


def write_report(rows, module_rows, resource_cards, growth_rows, gameplay_rows, shop_rows, ref_rows):
    os.makedirs(os.path.dirname(REPORT), exist_ok=True)
    lines = [
        "# P4g · 资源养成阅读版",
        "",
        "> 从 `全资源产销总表.xlsx` 派生，不替代机器校验总账；目标是让策划/数值能快速看资源从哪来、花到哪、哪些还只是奖励包引用。",
        "",
        "## 产物",
        "",
        f"- 阅读版 Excel: `{OUT}`",
        f"- 来源总账: `{SUMMARY}`",
        f"- 模块产销表: `{MODULE_DIR}`（每个模块文件已内置 `阅读摘要`、`产销阅读`、`资源汇总`、`机器明细`）。",
        f"- 总明细: {len(rows)} 行。",
        f"- 模块总览: {len(module_rows)} 个系统。",
        f"- 资源卡片: {len(resource_cards)} 个资源/道具。",
        f"- 养成消耗视图: {len(growth_rows)} 行。",
        f"- 玩法产出视图: {len(gameplay_rows)} 行。",
        f"- 商店兑换视图: {len(shop_rows)} 行。",
        f"- 奖励包引用汇总: {len(ref_rows)} 组。",
        "",
        "## 页签口径",
        "",
        "- `模块总览`：每个系统压成一行，先看主要产出、主要消耗、奖励包引用和风险提示。",
        "- `资源卡片`：按资源/道具聚合产出、消耗、引用、主要系统和风险提示。",
        "- `养成消耗`：角色、专武、旅券、映射装备、炮台、共鸣等养成阶段合并到人读行。",
        "- `玩法产出`：玩法/外循环按系统、分组、周期、条件聚合产出与消耗。",
        "- `商店兑换`：一行同时展示买到什么、花什么、限制和来源。",
        "- `奖励包引用`：DropId/RewardId 暂不展开随机权重，只按系统汇总引用位置。",
        "",
        "## 边界",
        "",
        "- 日常看单个模块时优先打开 `经济/模块产销表/*.xlsx`；本阅读版用于跨模块检索、资源卡片和全局导航。",
        "- 本阅读版不写公式，不参与正式配置校验；校验仍读取 `全资源产销总表.xlsx` 的 `追溯_全明细`。",
        "- Drop/Reward 随机包真实内容仍需后续专项展开。",
    ]
    with open(REPORT, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))


def main():
    if not os.path.exists(SUMMARY):
        raise FileNotFoundError(SUMMARY)
    rows = load_ledger()
    module_rows = build_module_overview(rows)
    resource_cards = build_resource_cards(rows)
    growth_rows = aggregate_by_condition(rows, systems=GROWTH_SYSTEMS)
    gameplay_rows = aggregate_by_condition(rows, exclude_systems=GROWTH_SYSTEMS | {"商店"})
    shop_rows = build_shop_rows(rows)
    ref_rows = build_reference_rows(rows)
    nav_rows = build_nav(rows, module_rows, resource_cards, growth_rows, gameplay_rows, shop_rows, ref_rows)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "阅读导航"
    for row in nav_rows:
        ws.append(row)
    style_nav(ws)

    add_sheet(
        wb,
        "模块总览",
        ["系统", "阅读定位", "记录数", "产出", "消耗", "引用", "主要分组", "主要产出", "主要消耗",
         "奖励包/引用", "周期/限制", "来源页", "阅读提示"],
        module_rows,
        [16, 14, 10, 8, 8, 8, 40, 54, 54, 46, 30, 38, 22],
    )
    add_sheet(
        wb,
        "资源卡片",
        ["资源分类", "资源/道具", "道具ID", "产出记录", "消耗记录", "引用记录", "产出数量", "消耗数量",
         "产出价值(钻石)", "消耗价值(钻石)", "主要来源", "主要消耗", "引用位置", "阅读提示"],
        resource_cards,
        [14, 22, 10, 10, 10, 10, 14, 14, 16, 16, 44, 44, 36, 22],
    )
    add_sheet(
        wb,
        "养成消耗",
        ["系统", "分组", "周期/限制", "条件/阶段", "产出/获得", "消耗/花费", "奖励包/引用", "来源", "备注"],
        growth_rows,
        [14, 18, 16, 28, 42, 52, 34, 28, 48],
    )
    add_sheet(
        wb,
        "玩法产出",
        ["系统", "分组", "周期/限制", "条件/阶段", "产出/获得", "消耗/花费", "奖励包/引用", "来源", "备注"],
        gameplay_rows,
        [16, 20, 16, 34, 52, 36, 42, 30, 48],
    )
    add_sheet(
        wb,
        "商店兑换",
        ["系统", "商店/分组", "买到", "花费", "货币价值(钻石)", "折扣/倍率", "限制", "条件", "来源", "备注"],
        shop_rows,
        [16, 22, 34, 28, 16, 12, 20, 34, 28, 46],
    )
    add_sheet(
        wb,
        "奖励包引用",
        ["系统", "分组", "来源", "引用数量", "引用示例", "条件示例", "处理建议", "备注"],
        ref_rows,
        [16, 22, 34, 10, 54, 44, 42, 50],
    )

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    wb.save(OUT)
    write_report(rows, module_rows, resource_cards, growth_rows, gameplay_rows, shop_rows, ref_rows)
    print("资源养成阅读版生成完成: 模块总览 %s 个，资源卡片 %s 个，养成视图 %s 行，玩法视图 %s 行，商店 %s 行，奖励引用 %s 组" %
          (len(module_rows), len(resource_cards), len(growth_rows), len(gameplay_rows), len(shop_rows), len(ref_rows)))
    print("[阅读版] %s" % OUT)
    print("[报告] %s" % REPORT)


if __name__ == "__main__":
    main()
