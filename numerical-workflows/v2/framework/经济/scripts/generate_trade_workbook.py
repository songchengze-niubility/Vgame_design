from __future__ import annotations

import argparse
import json
import re
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


NAVY = "17365D"
BLUE = "1F4E78"
LIGHT_BLUE = "D9EAF7"
LIGHT_GREEN = "E2F0D9"
LIGHT_YELLOW = "FFF2CC"
LIGHT_GRAY = "E7E6E6"
WHITE = "FFFFFF"
INPUT_BLUE = "0000FF"
LINK_GREEN = "008000"
TEXT = "222222"
THIN = Side(style="thin", color="B7C9D6")
SOURCE = "跑商玩法数值设计文档-2026-07-06.md（V22）"


def clean(value: str) -> str:
    return value.strip().strip("`").replace("<br>", "\n")


def split_row(line: str) -> list[str]:
    return [clean(x) for x in line.strip().strip("|").split("|")]


def is_separator(row: list[str]) -> bool:
    return bool(row) and all(re.fullmatch(r":?-{3,}:?", x.replace(" ", "")) for x in row)


def table_after(lines: list[str], heading: str) -> tuple[list[str], list[list[str]]]:
    try:
        start = next(i for i, line in enumerate(lines) if line.strip() == heading)
    except StopIteration as exc:
        raise ValueError(f"找不到章节：{heading}") from exc
    first = next((i for i in range(start + 1, len(lines)) if lines[i].lstrip().startswith("|")), None)
    if first is None:
        raise ValueError(f"章节下没有表格：{heading}")
    raw = []
    for line in lines[first:]:
        if not line.lstrip().startswith("|"):
            break
        raw.append(split_row(line))
    if len(raw) < 2 or not is_separator(raw[1]):
        raise ValueError(f"Markdown 表头异常：{heading}")
    width = len(raw[0])
    rows = [(row + [""] * width)[:width] for row in raw[2:] if not is_separator(row)]
    return raw[0], rows


def row_map(headers: list[str], rows: list[list[str]], key: str) -> dict[str, str]:
    idx = headers.index(key)
    return {row[idx]: row[1] for row in rows}


def number(value: str) -> float:
    match = re.search(r"-?\d[\d,]*(?:\.\d+)?", value)
    if not match:
        raise ValueError(f"无法读取数值：{value}")
    return float(match.group(0).replace(",", ""))


def integer(value: str) -> int:
    return int(round(number(value)))


def optional_integer(value: str) -> int | None:
    return integer(value) if str(value).strip() else None


def reward_specs(value: str) -> list[tuple[int, int, int]]:
    specs = []
    for part in str(value).split(";"):
        tokens = [int(x.strip()) for x in part.split(",") if x.strip()]
        if len(tokens) == 2:
            specs.append((tokens[0], tokens[1], 10000))
        elif len(tokens) == 3:
            specs.append((tokens[0], tokens[1], tokens[2]))
        else:
            raise ValueError(f"奖励格式异常：{value}")
    return specs


def percent(value: str) -> float:
    match = re.search(r"([+-]?\d+(?:\.\d+)?)%", value)
    if not match:
        raise ValueError(f"无法读取百分比：{value}")
    return float(match.group(1)) / 100


def percent_range(value: str) -> tuple[float, float]:
    nums = [float(x) / 100 for x in re.findall(r"\d+(?:\.\d+)?(?=%)", value)]
    if len(nums) < 2:
        raise ValueError(f"无法读取百分比区间：{value}")
    return nums[0], nums[1]


def numeric_range(value: str) -> tuple[int, int]:
    nums = [int(x.replace(",", "")) for x in re.findall(r"\d[\d,]*", value)]
    if len(nums) < 2:
        raise ValueError(f"无法读取数值区间：{value}")
    return nums[0], nums[1]


def level(value: str) -> int:
    return integer(value)


def source_comment(section: str) -> Comment:
    return Comment(f"Source: {SOURCE}, {section}", "Codex")


def add_table(ws, start_row: int, headers: list[str], rows: list[list], source: str, input_cols: set[int] | None = None) -> tuple[int, int]:
    input_cols = input_cols or set()
    for col, value in enumerate(headers, 1):
        cell = ws.cell(start_row, col, value)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.font = Font(name="Microsoft YaHei", size=10, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=THIN)
    for r_off, row in enumerate(rows, 1):
        for col, value in enumerate(row, 1):
            cell = ws.cell(start_row + r_off, col, value)
            cell.font = Font(name="Microsoft YaHei", size=10, color=INPUT_BLUE if col in input_cols else TEXT)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=Side(style="hair", color="D9E1F2"))
            if col in input_cols and value not in (None, ""):
                cell.comment = source_comment(source)
    return start_row + len(rows), len(headers)


def add_section(ws, row: int, title: str, width: int) -> None:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=width)
    cell = ws.cell(row, 1, title)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.font = Font(name="Microsoft YaHei", size=11, bold=True, color=WHITE)
    cell.alignment = Alignment(vertical="center")


def finish_sheet(ws, filter_ref: str | None = None, freeze: str = "A2", widths: dict[int, float] | None = None) -> None:
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = freeze
    if filter_ref:
        ws.auto_filter.ref = filter_ref
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            if cell.data_type == "f":
                color = LINK_GREEN if "!" in str(cell.value) else TEXT
                cell.font = Font(name="Microsoft YaHei", size=10, color=color)
            elif not cell.font or cell.font.name != "Microsoft YaHei":
                cell.font = Font(name="Microsoft YaHei", size=10, color=cell.font.color if cell.font and cell.font.color else TEXT)
    for col in range(1, ws.max_column + 1):
        width = widths.get(col) if widths else None
        if width is None:
            longest = max((len(str(ws.cell(r, col).value or "")) for r in range(1, min(ws.max_row, 100) + 1)), default=8)
            width = min(max(longest * 1.15 + 2, 10), 36)
        ws.column_dimensions[get_column_letter(col)].width = width
    for row in range(1, ws.max_row + 1):
        ws.row_dimensions[row].height = 30 if row == 1 else 24


def pct_columns(ws, columns: list[int], start: int, end: int) -> None:
    for col in columns:
        for row in range(start, end + 1):
            ws.cell(row, col).number_format = "0.0%;(0.0%);-"


def int_columns(ws, columns: list[int], start: int, end: int) -> None:
    for col in columns:
        for row in range(start, end + 1):
            ws.cell(row, col).number_format = "#,##0;(#,##0);-"


def parse_md(md_path: Path) -> dict:
    lines = md_path.read_text(encoding="utf-8").splitlines()
    result = {"lines": lines}
    for key, heading in {
        "meta": "## 元信息",
        "scope": "## 本期范围",
        "dispatch": "### R-01 基础派遣",
        "currency": "### R-02 商券与货物",
        "price": "### R-03 价格倍率",
        "locations": "## 站点与货物",
        "levels": "### R-05 等级经验",
        "unlocks": "### R-05A 等级解锁明细",
        "buy_rules": "### R-05B 商品购买数量上浮",
        "level_rewards": "### R-05C 升级一次性奖励",
        "supply": "### R-06 商会补给",
        "event_trigger": "### R-07 事件触发",
        "events": "### R-08 正式事件配置",
        "event_pools": "### R-08A 正式站点事件池",
        "shop": "### R-09 商店商品与兑换节奏",
        "ledger": "## 收益与产销测算",
        "timeline": "## 升级时间测算",
        "points": "## 数值点清单",
        "risks": "## 风险与待确认",
    }.items():
        result[key] = table_after(lines, heading)
    return result


def build(md_path: Path, output: Path, published_path: Path | None = None) -> dict:
    data = parse_md(md_path)
    dispatch = row_map(*data["dispatch"], "项")
    currency = row_map(*data["currency"], "项")
    price = row_map(*data["price"], "项")
    event_trigger = row_map(*data["event_trigger"], "项")
    event_refresh = [int(x) for x in re.findall(r"\d+", event_trigger["单次刷新数量"])][:2]
    buy_min, buy_max = percent_range(price["今日购入价"])
    normal_min, normal_max = percent_range(price["普通收购价"])
    urgent_min, urgent_max = percent_range(price["急需收购价"])
    event_headers = data["events"][0]
    event_rows = data["events"][1]
    event_pool_rows = data["event_pools"][1]
    if len(event_rows) != 25 or len(event_pool_rows) != 5:
        raise ValueError(f"正式事件数量异常：events={len(event_rows)}, pools={len(event_pool_rows)}")
    event_idx = {name: idx for idx, name in enumerate(event_headers)}
    event_by_id = {integer(row[event_idx["事件ID"]]): row for row in event_rows}
    choice_rate = percent(event_trigger["商券选项测算选择率"])
    event_ticket_by_location = []
    for pool in event_pool_rows:
        weighted_sum = 0
        total_weight = 0
        for part in pool[2].split(";"):
            event_id, weight = [int(x) for x in part.split(",")]
            reward1 = reward_specs(event_by_id[event_id][event_idx["奖励1"]])
            if len(reward1) != 1 or reward1[0][0] != 301:
                raise ValueError(f"事件 {event_id} 的选项1不是固定商券奖励")
            weighted_sum += reward1[0][1] * weight
            total_weight += weight
        event_ticket_by_location.append(weighted_sum / total_weight * choice_rate)
    event_expected = sum(event_ticket_by_location) / len(event_ticket_by_location)

    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.creator = "Codex"
    wb.properties.title = "跑商玩法数值文档 V22"
    wb.properties.subject = "由正式沉淀 MD 生成的跑商玩法数值阅读文档"
    wb.properties.keywords = "Vgame,跑商,行商委托,数值,V22"
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"

    ws = wb.create_sheet("说明")
    info = [
        ["功能名称", "跑商 / 行商委托"],
        ["当前版本", "V22"],
        ["文档类型", "正式数值阅读文档（派生输出）"],
        ["生成日期", date.today().isoformat()],
        ["唯一参数源", str(md_path)],
        ["输出位置", str(published_path or output)],
        ["当前状态", "InReview"],
        ["风险档位", "MEDIUM"],
        ["配置口径", "正式 TradeEvent 的 Id/名称/描述/UnlockId 为内容源；货物基础价、25/20/15 基础购买量、便签 ItemId 1075-1079、价格倍率、事件双选奖励与事件池概率以沉淀 MD V22 为准"],
        ["阅读顺序", "全局参数 → 站点货物/行商等级 → 商券商店/预算 → 收益与产销测算 → 升级时间测算 → 待确认项"],
        ["维护规则", "参数修改先回到 03_沉淀MD；本 Excel 重新生成，不直接形成第二份权威"],
        ["颜色说明", "蓝字=沉淀 MD 输入；绿字=跨页公式；黑字=本页公式；黄色=待确认"],
    ]
    add_table(ws, 1, ["项", "内容"], info, "元信息", {2})
    ws["B5"].font = Font(name="Microsoft YaHei", size=10, color=LINK_GREEN)
    finish_sheet(ws, "A1:B13", widths={1: 18, 2: 100})

    ws = wb.create_sheet("全局参数")
    params = [
        ["TR-N01", "初始商券", integer(dispatch["初始商券"]), "商券", "解锁跑商时发放", "R-01"],
        ["TR-N02", "常规派遣时间", integer(dispatch["常规派遣时间"]), "小时", "每日约 1 次", "R-01"],
        ["TR-N03", "新手首单", integer(dispatch["新手首单"]), "分钟", "仅 1 次", "R-01"],
        ["TR-N04-A", "买价最低倍率", buy_min, "倍率", "今日购入价下限", "R-03"],
        ["TR-N04-M", "买价均值倍率", "=AVERAGE(C5,C7)", "倍率", "测算口径", "R-03"],
        ["TR-N04-B", "买价最高倍率", buy_max, "倍率", "今日购入价上限", "R-03"],
        ["TR-N05-A", "普通收购最低倍率", normal_min, "倍率", "普通收购价下限", "R-03"],
        ["TR-N05-M", "普通收购均值倍率", "=AVERAGE(C8,C10)", "倍率", "测算口径", "R-03"],
        ["TR-N05-B", "普通收购最高倍率", normal_max, "倍率", "普通收购价上限", "R-03"],
        ["TR-N06-A", "急需收购最低倍率", urgent_min, "倍率", "急需收购价下限", "R-03"],
        ["TR-N06-M", "急需收购均值倍率", "=AVERAGE(C11,C13)", "倍率", "测算口径", "R-03"],
        ["TR-N06-B", "急需收购最高倍率", urgent_max, "倍率", "急需收购价上限", "R-03"],
        ["TR-N19", "货物回收系数", percent(currency["货物回收"]), "倍率", "今日购入价乘该系数", "R-02"],
        ["TR-N16", "固定载货量", integer(dispatch["单次派遣载货量"]), "件", "全等级固定", "R-01"],
        ["TR-N20", "每日刷新时间", dispatch["每日刷新时间"], "时间", "商会补给与日刷新", "R-01/R-06"],
        ["TR-N09", "行商等级上限", 20, "级", "当前 V22 上限", "R-05"],
        ["TR-N12-E", "事件商券期望", "='旅途事件'!B39", "商券/次", "五站均值；按 60% 玩家选择商券测算", "R-08B"],
        ["TR-N21", "默认商券道具 ID", integer(currency["商券道具 ID"]), "ItemId", "globals_conf.TradeDefaultTicketItemId", "R-02"],
        ["TR-N22", "事件存在时间", integer(event_trigger["事件存在时间"]), "小时", "globals_conf.TradeEventExistHours", "R-07"],
        ["TR-N23", "事件最多存在数量", integer(event_trigger["事件最多存在数量"]), "个", "globals_conf.TradeEventMaxExistCount", "R-07"],
        ["TR-N24", "单次跑商事件刷新下限", event_refresh[0], "个", "globals_conf.TradeRunEventRefreshMin", "R-07"],
        ["TR-N25", "单次跑商事件刷新上限", event_refresh[1], "个", "globals_conf.TradeRunEventRefreshMax", "R-07"],
    ]
    add_table(ws, 1, ["参数ID", "参数", "数值", "单位", "说明", "来源"], params, "全局参数", {3})
    pct_columns(ws, [3], 5, 14)
    int_columns(ws, [3], 2, 23)
    for r in [5, 6, 7, 8, 9, 10, 11, 12, 13, 14]:
        ws.cell(r, 3).number_format = "0.0%;(0.0%);-"
    finish_sheet(ws, "A1:F23", widths={1: 16, 2: 26, 3: 16, 4: 14, 5: 42, 6: 14})

    ws = wb.create_sheet("基础规则")
    add_section(ws, 1, "R-01 基础派遣", 3)
    d_headers, d_rows = data["dispatch"]
    end, _ = add_table(ws, 2, d_headers + ["来源"], [r + ["R-01"] for r in d_rows], "R-01", {2})
    add_section(ws, end + 2, "R-02 商券与货物", 3)
    c_headers, c_rows = data["currency"]
    add_table(ws, end + 3, c_headers + ["来源"], [r + ["R-02"] for r in c_rows], "R-02", {2})
    finish_sheet(ws, freeze="A3", widths={1: 24, 2: 90, 3: 14})

    ws = wb.create_sheet("价格规则")
    p_headers, p_rows = data["price"]
    end, _ = add_table(ws, 1, p_headers + ["来源"], [r + ["R-03"] for r in p_rows], "R-03", {2})
    add_section(ws, end + 2, "R-04 结算公式", 3)
    formulas = [
        ["今日购入价", "基础价 * 购入倍率 * (1 + 等级购入价格加成)", "R-04"],
        ["今日收购价", "基础价 * 收购倍率 * (1 + 等级出售价格加成)", "R-04"],
        ["实际可购买数量", "CEIL(商品基础购买值 * (1 + 等级购买数量上浮%))", "R-04"],
        ["本次利润", "Σ((今日收购价 - 今日购入价) * 出售数量)", "R-04"],
        ["等级经验", "MAX(本次利润, 0)", "R-04"],
    ]
    add_table(ws, end + 3, ["项目", "公式", "来源"], formulas, "R-04", {2})
    finish_sheet(ws, freeze="A2", widths={1: 26, 2: 90, 3: 14})

    ws = wb.create_sheet("站点货物")
    loc_headers, loc_rows = data["locations"]
    clean_rows = []
    goods = []
    for row in loc_rows:
        lv = level(row[0])
        clean_rows.append([
            lv, row[1], row[2], integer(row[3]), integer(row[4]), row[5], integer(row[6]), integer(row[7]),
            row[8], integer(row[9]), integer(row[10]),
        ])
        for name_col, price_col, limit_col in ((2, 3, 4), (5, 6, 7), (8, 9, 10)):
            goods.append([lv, row[1], row[name_col], integer(row[price_col]), integer(row[limit_col])])
    add_table(
        ws, 1,
        ["解锁等级", "站点", "低价货物", "基础价", "基础购买量", "中价货物", "基础价", "基础购买量", "高价货物", "基础价", "基础购买量"],
        clean_rows, "站点与货物", set(range(1, 12)),
    )
    int_columns(ws, [1, 4, 5, 7, 8, 10, 11], 2, 6)
    finish_sheet(ws, "A1:K6", widths={1: 12, 2: 18, 3: 18, 4: 12, 5: 14, 6: 18, 7: 12, 8: 14, 9: 18, 10: 12, 11: 14})

    ws = wb.create_sheet("货物明细")
    detail_rows = []
    for idx, row in enumerate(goods, 2):
        detail_rows.append(row + [
            f"=D{idx}*全局参数!$C$5", f"=D{idx}*全局参数!$C$6", f"=D{idx}*全局参数!$C$7",
            f"=D{idx}*全局参数!$C$9", f"=D{idx}*全局参数!$C$12", f"=I{idx}-G{idx}",
            f"=J{idx}-G{idx}", f"=G{idx}*全局参数!$C$14",
        ])
    detail_headers = ["解锁等级", "站点", "货物", "基础价", "基础购买量", "最低买价", "平均买价", "最高买价", "普通平均卖价", "急需平均卖价", "普通单件均利", "急需单件均利", "平均回收价"]
    add_table(ws, 1, detail_headers, detail_rows, "站点与货物", {1, 2, 3, 4})
    int_columns(ws, list(range(1, 14)), 2, 16)
    finish_sheet(ws, "A1:M16", widths={1: 12, 2: 18, 3: 18, 4: 12, 5: 14, 6: 13, 7: 13, 8: 13, 9: 16, 10: 16, 11: 16, 12: 16, 13: 16})

    ws = wb.create_sheet("行商等级")
    lvl_rows = []
    for idx, row in enumerate(data["levels"][1], 2):
        lvl_rows.append([
            level(row[0]), integer(row[1]), integer(row[2]), percent(row[3]), integer(row[4]), percent(row[5]), percent(row[6]),
            f"=B{idx}" if idx == 2 else f"=B{idx}-B{idx-1}", row[7], row[8],
        ])
    headers = ["等级", "累计正利润", "固定载货量", "购买数量上浮", "商会补给", "出售加成", "购入优惠", "本级所需正利润", "本级唯一提升", "升级一次性奖励"]
    add_table(ws, 1, headers, lvl_rows, "R-05", {1, 2, 3, 4, 5, 6, 7, 9, 10})
    pct_columns(ws, [4, 6, 7], 2, 21)
    int_columns(ws, [1, 2, 3, 5, 8], 2, 21)
    finish_sheet(ws, "A1:J21", widths={1: 10, 2: 16, 3: 14, 4: 16, 5: 14, 6: 12, 7: 12, 8: 18, 9: 62, 10: 28})
    chart = LineChart()
    chart.title = "行商等级累计正利润曲线"
    chart.y_axis.title = "累计正利润"
    chart.x_axis.title = "行商等级"
    chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=21), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=21))
    chart.height = 7
    chart.width = 14
    ws.add_chart(chart, "L2")

    ws = wb.create_sheet("等级解锁明细")
    unlock_rows = []
    for row in data["unlocks"][1]:
        unlock_rows.append([level(row[0]), row[1], integer(row[2]), percent(row[3]), integer(row[4]), row[5], row[6]])
    add_table(ws, 1, data["unlocks"][0], unlock_rows, "R-05A", set(range(1, 8)))
    pct_columns(ws, [4], 2, 21)
    int_columns(ws, [1, 3, 5], 2, 21)
    finish_sheet(ws, "A1:G21", widths={1: 10, 2: 24, 3: 14, 4: 18, 5: 14, 6: 24, 7: 68})

    ws = wb.create_sheet("商品购买数量规则")
    buy_rows = []
    for idx, row in enumerate(data["buy_rules"][1], 2):
        rate = percent(row[1])
        buy_rows.append([
            row[0], rate,
            25, f"=ROUNDUP($C{idx}*(1+$B{idx}),0)", integer(row[2]),
            20, f"=ROUNDUP($F{idx}*(1+$B{idx}),0)", integer(row[3]),
            15, f"=ROUNDUP($I{idx}*(1+$B{idx}),0)", integer(row[4]),
            "V22 已确认",
        ])
    add_table(
        ws, 1,
        ["等级段", "购买数量上浮", "低价基础量", "低价公式实际", "低价MD实际", "中价基础量", "中价公式实际", "中价MD实际", "高价基础量", "高价公式实际", "高价MD实际", "状态"],
        buy_rows, "R-05B", {1, 2, 3, 5, 6, 8, 9, 11, 12},
    )
    pct_columns(ws, [2], 2, 5)
    int_columns(ws, list(range(3, 12)), 2, 5)
    finish_sheet(ws, "A1:L5", widths={1: 16, 2: 18, 3: 14, 4: 16, 5: 16, 6: 14, 7: 16, 8: 16, 9: 14, 10: 16, 11: 16, 12: 16})

    ws = wb.create_sheet("升级一次性奖励")
    reward_rows = []
    for idx, row in enumerate(data["level_rewards"][1], 2):
        reward_rows.append([
            level(row[0]), row[1], integer(row[2]), f'=IF($B{idx}="商券",$C{idx},0)', f'=IF($B{idx}="钻石",$C{idx},0)',
            f'=IF($B{idx}="角色常驻抽卡券",$C{idx},0)', row[3], row[4], "待确认" if row[1] != "商券" else "玩法内货币", "待确认" if row[1] != "商券" else "不使用Drop",
        ])
    add_table(ws, 1, ["等级", "奖励类型", "数量", "商券", "钻石", "角色常驻抽卡券", "奖励摘要", "备注", "ItemId", "DropId"], reward_rows, "R-05C", {1, 2, 3, 7, 8, 9, 10})
    total_row = 22
    ws.cell(total_row, 1, "合计")
    for col in (4, 5, 6):
        ws.cell(total_row, col, f"=SUM({get_column_letter(col)}2:{get_column_letter(col)}21)")
    for cell in ws[total_row]:
        cell.fill = PatternFill("solid", fgColor=LIGHT_GREEN)
        cell.font = Font(name="Microsoft YaHei", size=10, bold=True)
    int_columns(ws, [1, 3, 4, 5, 6], 2, total_row)
    for row in range(2, 22):
        if ws.cell(row, 9).value == "待确认":
            ws.cell(row, 9).fill = PatternFill("solid", fgColor=LIGHT_YELLOW)
            ws.cell(row, 10).fill = PatternFill("solid", fgColor=LIGHT_YELLOW)
    finish_sheet(ws, "A1:J21", widths={1: 10, 2: 22, 3: 12, 4: 12, 5: 12, 6: 20, 7: 32, 8: 34, 9: 16, 10: 16})

    ws = wb.create_sheet("商会补给")
    add_table(ws, 1, data["supply"][0], data["supply"][1], "R-06", {2})
    finish_sheet(ws, "A1:B9", widths={1: 24, 2: 90})

    ws = wb.create_sheet("旅途事件")
    evt_rows = []
    for idx, row in enumerate(event_rows, 2):
        reward1 = reward_specs(row[event_idx["奖励1"]])
        if len(reward1) != 1:
            raise ValueError(f"事件选项1必须是固定单项：{row}")
        evt_rows.append([
            integer(row[event_idx["事件ID"]]), integer(row[event_idx["站点序号"]]), row[event_idx["品级"]],
            row[event_idx["事件名"]], row[event_idx["事件描述"]], row[event_idx["选项1"]], row[event_idx["选项2"]],
            optional_integer(row[event_idx["UnlockId"]]), row[event_idx["类型"]], integer(row[event_idx["DropId1"]]),
            integer(row[event_idx["DropId2"]]), integer(row[event_idx["权重"]]), percent(row[event_idx["概率"]]),
            reward1[0][0], reward1[0][1], row[event_idx["奖励2"]], choice_rate,
            f"=O{idx}*M{idx}*Q{idx}",
            f'=IF(AND(F{idx}<>"",G{idx}<>"",J{idx}>0,K{idx}>0),"通过","异常")',
        ])
    sheet_event_headers = [
        "事件ID", "站点序号", "品级", "事件名", "事件描述", "选项1", "选项2", "UnlockId", "类型",
        "DropId1", "DropId2", "权重", "事件概率", "奖励1 ItemId", "奖励1数量", "奖励2配置", "商券选项测算选择率",
        "商券期望贡献", "配置校验",
    ]
    add_table(ws, 1, sheet_event_headers, evt_rows, "R-08 正式 TradeEvent", set(range(1, 18)))
    int_columns(ws, [1, 2, 8, 10, 11, 12, 14, 15, 18], 2, 26)
    pct_columns(ws, [13, 17], 2, 26)

    add_section(ws, 29, "R-08A 正式站点事件池", 19)
    pool_rows = []
    for idx, row in enumerate(event_pool_rows, 31):
        entries = []
        for part in row[2].split(";"):
            event_id, weight = part.split(",")
            entries.extend([int(event_id), int(weight)])
        if len(entries) != 10:
            raise ValueError(f"站点事件池不是 5 条：{row}")
        first_event_row = 2 + (integer(row[0]) - 1) * 5
        pool_rows.append([
            integer(row[0]), row[1], *entries,
            f"=SUM(D{idx},F{idx},H{idx},J{idx},L{idx})",
            f"=SUM(D{idx},F{idx},H{idx})/M{idx}", f"=J{idx}/M{idx}", f"=L{idx}/M{idx}",
            f"=(O{first_event_row}*D{idx}+O{first_event_row+1}*F{idx}+O{first_event_row+2}*H{idx}+O{first_event_row+3}*J{idx}+O{first_event_row+4}*L{idx})/M{idx}",
            choice_rate, f"=Q{idx}*R{idx}",
        ])
    pool_headers = ["站点序号", "正式站点名"]
    for i in range(1, 6):
        pool_headers.extend([f"事件{i}ID", f"权重{i}"])
    pool_headers.extend(["权重合计", "普通合计概率", "中级概率", "高级概率", "选项1商券加权均值", "商券选项测算选择率", "事件商券期望"])
    add_table(ws, 30, pool_headers, pool_rows, "R-08A 正式 TradeLocation.TradeEventList", set(range(1, 13)))
    int_columns(ws, list(range(1, 14)) + [17, 19], 31, 35)
    pct_columns(ws, [14, 15, 16, 18], 31, 35)

    ws["A38"] = "正式事件数量"
    ws["B38"] = "=COUNTA(A2:A26)"
    ws["A39"] = "五站事件商券期望均值"
    ws["B39"] = "=AVERAGE(S31:S35)"
    ws["A40"] = "Lv.1事件商券期望"
    ws["B40"] = "=S31"
    ws["A41"] = "Lv.20事件商券期望"
    ws["B41"] = "=S35"
    ws["A42"] = "奖励结算口径"
    ws["B42"] = "每个事件二选一，只发放所选选项的专属 Drop"
    ws["A43"] = "事件池权重校验"
    ws["B43"] = '=IF(AND(COUNTIF(M31:M35,1000)=5,COUNTIF(N31:N35,75%)=5,COUNTIF(O31:O35,15%)=5,COUNTIF(P31:P35,10%)=5),"通过","异常")'
    for row in (38, 39, 40, 41, 42, 43):
        for cell in ws[row]:
            cell.fill = PatternFill("solid", fgColor=LIGHT_GREEN)
            cell.font = Font(name="Microsoft YaHei", size=10, bold=True)

    add_section(ws, 45, "R-07 触发规则", 19)
    trigger_rows = [r + ["R-07"] for r in data["event_trigger"][1]]
    add_table(ws, 46, data["event_trigger"][0] + ["来源"], trigger_rows, "R-07", {2})
    finish_sheet(ws, freeze="A2", widths={1: 11, 2: 11, 3: 10, 4: 18, 5: 25, 6: 18, 7: 18, 8: 12, 9: 10, 10: 13, 11: 13, 12: 10, 13: 12, 14: 13, 15: 13, 16: 30, 17: 18, 18: 14, 19: 12})

    ws = wb.create_sheet("商券商店")
    shop_rows = []
    for idx, row in enumerate(data["shop"][1], 2):
        limit_type = row[5]
        shop_rows.append([
            integer(row[0]), row[1], integer(row[2]), integer(row[3]), integer(row[4]), limit_type,
            integer(row[6]), level(row[7]),
            f'=IF(F{idx}="每周",G{idx},IF(F{idx}="每月",G{idx}/4,0))',
            f"=E{idx}*I{idx}",
            f'=IF(F{idx}="永久一次",E{idx}*G{idx},0)',
            row[8],
        ])
    shop_headers = [
        "商品序号", "商品名称", "ItemId", "单次获得数量", "单价（商券）", "限购类型", "限购次数", "解锁等级",
        "周等价兑换次数", "周等价消耗", "一次性库存总价", "说明",
    ]
    add_table(ws, 1, shop_headers, shop_rows, "R-09", set(range(1, 9)) | {12})
    total_row = 2 + len(shop_rows)
    ws.cell(total_row, 1, "合计")
    ws.cell(total_row, 10, f"=SUM(J2:J{total_row-1})")
    ws.cell(total_row, 11, f"=SUM(K2:K{total_row-1})")
    for cell in ws[total_row]:
        cell.fill = PatternFill("solid", fgColor=LIGHT_GREEN)
        cell.font = Font(name="Microsoft YaHei", size=10, bold=True)
    int_columns(ws, [1, 3, 4, 5, 7, 8, 10, 11], 2, total_row)
    for row in range(2, total_row):
        ws.cell(row, 9).number_format = "0.00"
    finish_sheet(ws, f"A1:L{total_row-1}", widths={1: 11, 2: 28, 3: 11, 4: 16, 5: 16, 6: 14, 7: 12, 8: 12, 9: 18, 10: 16, 11: 18, 12: 50})

    ws = wb.create_sheet("商店预算测算")
    budget_rows = []
    for idx, lv in enumerate((1, 10, 15, 20), 2):
        reserve = 150000 if lv == 20 else 0
        budget_rows.append([
            lv,
            f"=AVERAGE(INDEX('收益与产销测算'!$N$2:$N$21,MATCH(A{idx},'收益与产销测算'!$A$2:$A$21,0)),INDEX('收益与产销测算'!$O$2:$O$21,MATCH(A{idx},'收益与产销测算'!$A$2:$A$21,0)))",
            f"=B{idx}*7",
            f'=SUMIFS(\'商券商店\'!$J$2:$J$19,\'商券商店\'!$H$2:$H$19,"<="&A{idx})',
            f"=C{idx}-D{idx}",
            f"=D{idx}/C{idx}",
            reserve,
            f"=D{idx}+G{idx}",
            f"=C{idx}-H{idx}",
        ])
    add_table(
        ws, 1,
        ["等级", "日净商券中位收入", "周净商券中位收入", "已解锁重复兑换周消耗", "当前周结余", "当前消耗率", "后期商品预留周预算", "含预留周消耗", "含预留周结余"],
        budget_rows, "R-09 商店预算", {1, 7},
    )
    int_columns(ws, [1, 2, 3, 4, 5, 7, 8, 9], 2, 5)
    pct_columns(ws, [6], 2, 5)
    add_section(ws, 8, "房车天赋便签 90 天进度", 5)
    note_rows = [
        ["单种每周兑换次数", "=INDEX('商券商店'!$G$2:$G$19,MATCH(1075,'商券商店'!$C$2:$C$19,0))", 5, "次", "Lv.1 同时开放五种"],
        ["单次获得数量", "=INDEX('商券商店'!$D$2:$D$19,MATCH(1075,'商券商店'!$C$2:$C$19,0))", 5, "个", "五种一致"],
        ["单种每周产出", "=B10*B11", 25, "个", "5 次 × 5 个"],
        ["90 天完整周数", "=ROUNDDOWN(90/7,0)", 12, "周", "只计完整刷新周；自然周边界可能实际覆盖 13 次"],
        ["90 天保守产出", "=B12*B13", 300, "个/种", "不含首尾自然周额外刷新"],
        ["规划口径产出", "=B12*13", 325, "个/种", "按 13 个周刷新周期"],
        ["单职业完整需求", 500, 500, "个", "房车天赋树目标"],
        ["90 天规划完成度", "=B15/B16", 0.65, "比例", "保留活动及其他来源空间"],
    ]
    add_table(ws, 9, ["项目", "公式结果", "MD预期", "单位", "说明"], note_rows, "R-09 商店预算", {3, 4, 5})
    int_columns(ws, [2, 3], 10, 16)
    pct_columns(ws, [2, 3], 17, 17)
    finish_sheet(ws, freeze="A2", widths={1: 24, 2: 22, 3: 16, 4: 14, 5: 54, 6: 14, 7: 20, 8: 18, 9: 18})

    ws = wb.create_sheet("收益与产销测算")
    ledger_rows = []
    for idx, row in enumerate(data["ledger"][1], 2):
        buy_lo, buy_hi = numeric_range(row[6])
        sale_lo, sale_hi = numeric_range(row[7])
        profit_lo, profit_hi = numeric_range(row[8])
        event_for_level = event_ticket_by_location[integer(row[1]) - 1]
        if abs(number(row[9]) - event_for_level) > 0.001:
            raise ValueError(f"收益表事件期望与事件池不一致：{row[0]}")
        net_lo, net_hi = numeric_range(row[10])
        calc_profit_lo, calc_profit_hi = sale_lo - buy_lo, sale_hi - buy_hi
        calc_net_lo = calc_profit_lo + integer(row[5]) + round(event_for_level)
        calc_net_hi = calc_profit_hi + integer(row[5]) + round(event_for_level)
        diffs = [abs(calc_profit_lo - profit_lo), abs(calc_profit_hi - profit_hi), abs(calc_net_lo - net_lo), abs(calc_net_hi - net_hi)]
        if max(diffs) > 2:
            raise ValueError(f"收益表与公式不一致：{row[0]}")
        ledger_rows.append([
            level(row[0]), integer(row[1]), percent(row[2]), percent(row[3]), percent(row[4]), integer(row[5]),
            buy_lo, buy_hi, sale_lo, sale_hi, f"=I{idx}-G{idx}", f"=J{idx}-H{idx}", f"=INDEX('旅途事件'!$S$31:$S$35,B{idx})",
            f"=K{idx}+F{idx}+M{idx}", f"=L{idx}+F{idx}+M{idx}",
        ])
    ledger_headers = ["等级", "开放站点", "购买上浮", "出售加成", "购入优惠", "商会补给", "采购低", "采购高", "回款低", "回款高", "利润低", "利润高", "事件期望", "净流入低", "净流入高"]
    add_table(ws, 1, ledger_headers, ledger_rows, "收益与产销测算", set(range(1, 11)))
    pct_columns(ws, [3, 4, 5], 2, 21)
    int_columns(ws, [1, 2, 6, 7, 8, 9, 10, 11, 12, 13, 14, 15], 2, 21)
    finish_sheet(ws, "A1:O21", widths={1: 9, 2: 12, 3: 12, 4: 12, 5: 12, 6: 13, 7: 12, 8: 12, 9: 12, 10: 12, 11: 12, 12: 12, 13: 13, 14: 13, 15: 13})
    chart = LineChart()
    chart.title = "每日净商券流入区间"
    chart.y_axis.title = "商券/日"
    chart.x_axis.title = "行商等级"
    chart.add_data(Reference(ws, min_col=14, max_col=15, min_row=1, max_row=21), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=21))
    chart.height = 7
    chart.width = 14
    ws.add_chart(chart, "Q2")

    ws = wb.create_sheet("升级时间测算")
    timeline_rows = []
    for idx, row in enumerate(data["timeline"][1], 2):
        day = integer(row[2])
        timeline_rows.append([
            level(row[0]), integer(row[1]), day, f"=C{idx}" if idx == 2 else f"=C{idx}-C{idx-1}",
            f"=B{idx}" if idx == 2 else f"=B{idx}-B{idx-1}", f"=IF(D{idx}=0,0,E{idx}/D{idx})", row[3],
        ])
    add_table(ws, 1, ["达到等级", "累计正利润阈值", "预计天数", "阶段天数", "阶段所需正利润", "阶段日均正利润", "关键提升/解锁"], timeline_rows, "升级时间测算", {1, 2, 3, 7})
    int_columns(ws, [1, 2, 3, 4, 5, 6], 2, 20)
    finish_sheet(ws, "A1:G20", widths={1: 12, 2: 18, 3: 12, 4: 12, 5: 18, 6: 18, 7: 72})

    ws = wb.create_sheet("公式测算")
    formula_rows = [
        ["正式旅途事件数量", "='旅途事件'!B38", 25, "=IF(B2=C2,\"通过\",\"异常\")"],
        ["旅途事件商券期望", "='旅途事件'!B39", round(event_expected), "=IF(ABS(B3-C3)<0.000001,\"通过\",\"异常\")"],
        ["平均买价倍率", "=全局参数!C6", (buy_min + buy_max) / 2, "=IF(ABS(B4-C4)<0.000001,\"通过\",\"异常\")"],
        ["普通单件均利率", "=全局参数!C9-全局参数!C6", (normal_min + normal_max) / 2 - (buy_min + buy_max) / 2, "=IF(ABS(B5-C5)<0.000001,\"通过\",\"异常\")"],
        ["急需单件均利率", "=全局参数!C12-全局参数!C6", (urgent_min + urgent_max) / 2 - (buy_min + buy_max) / 2, "=IF(ABS(B6-C6)<0.000001,\"通过\",\"异常\")"],
        ["Lv20普通单件均利率", "=全局参数!C9*(1+行商等级!F21)-全局参数!C6*(1+行商等级!G21)", 0.375, "=IF(ABS(B7-C7)<0.000001,\"通过\",\"异常\")"],
        ["Lv20急需单件均利率", "=全局参数!C12*(1+行商等级!F21)-全局参数!C6*(1+行商等级!G21)", 1.2375, "=IF(ABS(B8-C8)<0.000001,\"通过\",\"异常\")"],
        ["平均回收价倍率", "=全局参数!C6*全局参数!C14", ((buy_min + buy_max) / 2) * percent(currency["货物回收"]), "=IF(ABS(B9-C9)<0.000001,\"通过\",\"异常\")"],
        ["Lv20毕业天数", "=升级时间测算!C20", 89, "=IF(B10=C10,\"通过\",\"异常\")"],
        ["Lv1五种便签周消耗", "='商店预算测算'!D2", 250000, "=IF(B11=C11,\"通过\",\"异常\")"],
        ["Lv20当前重复兑换周消耗", "='商店预算测算'!D5", 575000, "=IF(B12=C12,\"通过\",\"异常\")"],
        ["Lv20含预留周消耗", "='商店预算测算'!H5", 725000, "=IF(B13=C13,\"通过\",\"异常\")"],
        ["Lv20含预留周结余", "='商店预算测算'!I5", 86006, "=IF(B14=C14,\"通过\",\"异常\")"],
        ["一次性礼包库存总价", "='商券商店'!K20", 3000000, "=IF(B15=C15,\"通过\",\"异常\")"],
        ["便签90天规划完成度", "='商店预算测算'!B17", 0.65, "=IF(ABS(B16-C16)<0.000001,\"通过\",\"异常\")"],
    ]
    add_table(ws, 1, ["项目", "公式结果", "MD预期", "校验"], formula_rows, "公式测算", {3})
    int_columns(ws, [2, 3], 2, 3)
    for row in [4, 5, 6, 7, 8, 9]:
        ws.cell(row, 2).number_format = "0.0%;(0.0%);-"
        ws.cell(row, 3).number_format = "0.0%;(0.0%);-"
    int_columns(ws, [2, 3], 3, 3)
    int_columns(ws, [2, 3], 10, 15)
    pct_columns(ws, [2, 3], 16, 16)
    finish_sheet(ws, "A1:D16", widths={1: 32, 2: 24, 3: 18, 4: 14})

    ws = wb.create_sheet("待确认项")
    add_section(ws, 1, "数值点清单", 4)
    end, _ = add_table(ws, 2, data["points"][0], data["points"][1], "数值点清单", {1, 2, 3, 4})
    add_section(ws, end + 2, "风险与待确认", 4)
    risk_start = end + 3
    add_table(ws, risk_start, data["risks"][0], data["risks"][1], "风险与待确认", {1, 2, 3})
    for row in range(risk_start + 1, risk_start + 1 + len(data["risks"][1])):
        for col in range(1, 4):
            ws.cell(row, col).fill = PatternFill("solid", fgColor=LIGHT_YELLOW)
    finish_sheet(ws, freeze="A3", widths={1: 14, 2: 36, 3: 82, 4: 16})

    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    reopened = load_workbook(output, data_only=False, read_only=False)
    formula_count = sum(1 for ws in reopened.worksheets for row in ws.iter_rows() for cell in row if cell.data_type == "f")
    if reopened.sheetnames != ["说明", "全局参数", "基础规则", "价格规则", "站点货物", "货物明细", "行商等级", "等级解锁明细", "商品购买数量规则", "升级一次性奖励", "商会补给", "旅途事件", "商券商店", "商店预算测算", "收益与产销测算", "升级时间测算", "公式测算", "待确认项"]:
        raise ValueError("页签清单异常")
    if formula_count < 150:
        raise ValueError(f"公式数量异常：{formula_count}")
    return {"status": "success", "output": str(output), "sheets": len(reopened.sheetnames), "formulas": formula_count, "event_ticket_expected": event_expected}


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--md", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--published-path", type=Path)
    args = parser.parse_args()
    print(json.dumps(build(args.md, args.output, args.published_path), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
