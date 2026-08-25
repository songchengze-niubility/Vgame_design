# -*- coding: utf-8 -*-
"""
sync_economy_knowledge_pages.py — 经济知识/工具页选择性镜像 (P4a)

只把不能直接用参数生成替代、且需要保留上下文的经济知识/工具页纳入 v2：
  - 道具定价
  - 角色、专武、旅券掉落池
  - 讨论页-玩法定位和产出

商店、任务、签到、通行证、各玩法奖励等结构化内容走经济总账生成器，
不通过整本镜像替代。
"""
import io
import os
import sys
import warnings

warnings.filterwarnings("ignore")

import openpyxl

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

HERE = os.path.dirname(os.path.abspath(__file__))
ECONOMY_DIR = os.path.dirname(HERE)
V2_DIR = os.path.dirname(ECONOMY_DIR)
NUM_DIR = os.path.dirname(V2_DIR)

SRC = os.path.join(NUM_DIR, "Vgame经济框架.xlsx")
OUT = os.path.join(ECONOMY_DIR, "镜像_经济知识工具页.xlsx")
REPORT = os.path.join(ECONOMY_DIR, "_reports", "P4a_经济知识工具页_1比1镜像.md")

KEEP_SHEETS = [
    ("道具定价", "工具/定价锚点", "全经济账本的钻石/RMB 等价锚点，必须保留原始上下文和备注。"),
    ("角色、专武、旅券掉落池", "工具/概率池", "抽卡、专武、旅券池的概率与价值口径，首版先保留原工具页。"),
    ("讨论页-玩法定位和产出", "知识/定位表", "玩法职责与产出定位讨论，不适合直接参数化生成。"),
]

EXCLUDED_NOTE = [
    "商店",
    "签到",
    "7日任务",
    "任务",
    "通行证",
    "无限挑战",
    "竞技场",
    "爬塔",
    "大秘境",
    "资源、装备副本",
    "旅券副本",
    "映射装备副本",
    "主线副本",
]


def is_formula(value):
    return (
        isinstance(value, str) and value.startswith("=")
    ) or (
        hasattr(value, "text") and hasattr(value, "ref")
    )


def comparable_value(value):
    if hasattr(value, "text") and hasattr(value, "ref"):
        return ("ArrayFormula", value.ref, value.text)
    return value


def sheet_stats(ws):
    nonempty = 0
    formulas = 0
    for row in ws.iter_rows():
        for cell in row:
            value = cell.value
            if value not in (None, ""):
                nonempty += 1
                if is_formula(value):
                    formulas += 1
    return {
        "rows": ws.max_row,
        "cols": ws.max_column,
        "nonempty": nonempty,
        "formulas": formulas,
        "merged": len(ws.merged_cells.ranges),
    }


def compare_sheet(src_ws, out_ws, src_value_ws):
    max_row = max(src_ws.max_row, out_ws.max_row)
    max_col = max(src_ws.max_column, out_ws.max_column)
    checked = 0
    fails = []
    for r in range(1, max_row + 1):
        for c in range(1, max_col + 1):
            checked += 1
            src_value = comparable_value(src_ws.cell(r, c).value)
            out_value = comparable_value(out_ws.cell(r, c).value)
            materialized = is_formula(src_ws.cell(r, c).value) and out_value == comparable_value(src_value_ws.cell(r, c).value)
            if src_value != out_value and not materialized:
                fails.append((src_ws.cell(r, c).coordinate, src_value, out_value))
                if len(fails) >= 20:
                    return checked, fails
    return checked, fails


def write_report(rows, total_checked, total_fails, materialized):
    os.makedirs(os.path.dirname(REPORT), exist_ok=True)
    lines = [
        "# P4a · 经济知识/工具页 1:1 镜像",
        "",
        "> 纠偏记录：经济层不做整本镜像。只镜像不能直接参数量化、且当前需要保留上下文的知识/工具页；其他页签进入 `全资源产销总表.xlsx` 的结构化总账。",
        "",
        "## 结论",
        "",
        f"- 源表: `{SRC}`",
        f"- 产物: `{OUT}`",
        f"- 保留页签: {len(KEEP_SHEETS)} 个。",
        f"- 内容/公式逐格校验: {total_checked - total_fails}/{total_checked} 通过。",
        "- 商店、任务、签到、通行证、玩法奖励、养成消耗等不进镜像，走 P4b 经济总账生成链路。",
        "- 工具页镜像不额外输出索引页；`道具定价` 只保留原工具页，定价锚点由脚本内部读取。",
        f"- 为保证独立打开，{len(materialized)} 个引用未保留页签的公式已固化为源表缓存值。",
        "",
        "## 保留页签",
        "",
        "| 页签 | 类型 | 保留原因 | 行×列 | 非空 | 公式 | 合并格 | 校验 |",
        "|---|---|---|---:|---:|---:|---:|---|",
    ]
    for row in rows:
        stats = row["stats"]
        lines.append(
            "| {sheet} | {kind} | {reason} | {rows}×{cols} | {nonempty} | {formulas} | {merged} | {status} |".format(
                sheet=row["sheet"],
                kind=row["kind"],
                reason=row["reason"],
                rows=stats["rows"],
                cols=stats["cols"],
                nonempty=stats["nonempty"],
                formulas=stats["formulas"],
                merged=stats["merged"],
                status=row["status"],
            )
        )

    lines += [
        "",
        "## 结构化处理页签",
        "",
        "以下页签不通过镜像保留，已纳入或预留给 `gen_economy_ledger.py` 做产销账本追踪：",
        "",
    ]
    for sheet in EXCLUDED_NOTE:
        lines.append(f"- `{sheet}`")

    if total_fails:
        lines += ["", "## 不一致样例", ""]
        for row in rows:
            for cell, src_value, out_value in row["fails"]:
                lines.append(f"- `{row['sheet']}`!`{cell}`: 源=`{src_value}` / 镜像=`{out_value}`")

    with open(REPORT, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))


def main():
    if not os.path.exists(SRC):
        raise FileNotFoundError(SRC)

    src_wb = openpyxl.load_workbook(SRC, data_only=False)
    src_value_wb = openpyxl.load_workbook(SRC, data_only=True)
    keep_names = [item[0] for item in KEEP_SHEETS]
    missing = [name for name in keep_names if name not in src_wb.sheetnames]
    if missing:
        src_wb.close()
        raise RuntimeError("源表缺少必要页签: %s" % ", ".join(missing))

    mirror_wb = openpyxl.load_workbook(SRC, data_only=False)
    for sheet in list(mirror_wb.sheetnames):
        if sheet not in keep_names:
            del mirror_wb[sheet]
    materialized = []
    for sheet in keep_names:
        for row in mirror_wb[sheet].iter_rows():
            for cell in row:
                value = cell.value
                if not is_formula(value):
                    continue
                removed_ref = any((name + "!") in str(value) or ("'%s'!" % name) in str(value) for name in src_wb.sheetnames if name not in keep_names)
                if removed_ref:
                    cached = src_value_wb[sheet][cell.coordinate].value
                    if cached is None:
                        raise RuntimeError("无法固化缺页公式: %s!%s %s" % (sheet, cell.coordinate, value))
                    materialized.append((sheet, cell.coordinate, value, cached))
                    cell.value = cached
    mirror_wb.active = 0
    mirror_wb.save(OUT)
    mirror_wb.close()

    out_wb = openpyxl.load_workbook(OUT, data_only=False)
    rows = []
    total_checked = 0
    total_fails = 0
    for sheet, kind, reason in KEEP_SHEETS:
        stats = sheet_stats(src_wb[sheet])
        checked, fails = compare_sheet(src_wb[sheet], out_wb[sheet], src_value_wb[sheet])
        total_checked += checked
        total_fails += len(fails)
        rows.append({
            "sheet": sheet,
            "kind": kind,
            "reason": reason,
            "stats": stats,
            "fails": fails,
            "status": "✅ 内容/公式一致" if not fails else "🔴 不一致",
        })

    write_report(rows, total_checked, total_fails, materialized)
    out_wb.close()
    src_wb.close()
    src_value_wb.close()

    if total_fails:
        print(f"不一致: {total_fails}")
    else:
        print(f"经济知识/工具页逐格复现: {len(KEEP_SHEETS)}/{len(KEEP_SHEETS)}")
    print(f"[独立化] 固化跨未保留页签公式: {len(materialized)}")
    print(f"[镜像] {OUT}")
    print(f"[报告] {REPORT}")


if __name__ == "__main__":
    main()
