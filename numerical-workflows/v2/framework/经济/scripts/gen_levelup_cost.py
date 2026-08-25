# -*- coding: utf-8 -*-
"""
gen_levelup_cost.py — 角色升级成本曲线 生成器 (框架重搭 · 经济)

读取经济参数基线 MD / 经济参数 + 角色升级的「相对上级提升」输入列 →
重算 每级经验/总经验/每级金币/总金币，与现有 经济框架/角色升级 逐格对照。

公式 (已逆向验证):
  每级所需经验(L) = 首级经验 + Σ 相对上级提升(2..L)
  总经验(L)       = 每级所需经验 累计
  每级所需金币(L) = 每级所需经验(L) × 金币比
  总金币(L)       = 每级所需金币 累计
「相对上级提升」为手调输入(步进), 从源表取。

只读源表。生成表为独立新文件。
用法: python gen_levelup_cost.py [--check]
"""
import sys, io, os, argparse
import warnings
warnings.filterwarnings('ignore')
import openpyxl
from openpyxl.styles import Font, PatternFill

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

HERE = os.path.dirname(os.path.abspath(__file__))
ECONOMY_DIR = os.path.dirname(HERE)  # 经济/
V2_DIR = os.path.dirname(ECONOMY_DIR)  # 框架v2/
NUM_DIR = os.path.dirname(V2_DIR)  # 数值文档/数值文档/
sys.path.insert(0, V2_DIR)
from _shared.md_parameter_reader import load_parameter_workbook, parameter_md_path

ECONOMY_PARAM = parameter_md_path('经济')

OUT     = os.path.join(ECONOMY_DIR, '生成_角色升级成本.xlsx')
REPORT  = os.path.join(ECONOMY_DIR, '_reports', 'P经济_角色升级成本_生成与对照.md')
SRC     = os.path.join(NUM_DIR, 'Vgame经济框架.xlsx')
SRC_SHEET = '角色升级'
DATA_ROW0 = 7   # L1 行
# 列: c2等级 c3相对提升(输入) c4每级经验 c5总经验 c6每级金币 c7总金币
C_LV, C_STEP, C_NEED, C_TOTEXP, C_GOLD, C_TOTGOLD = 2, 3, 4, 5, 6, 7

def read_params():
    wb = load_parameter_workbook(ECONOMY_PARAM, data_only=True); ws = wb['经济参数']
    p={}
    for r in range(1, ws.max_row+1):
        k=ws.cell(r,1).value
        if k in ('角色升级等级上限','角色升级首级经验','角色升级金币比'): p[k]=ws.cell(r,2).value
    wb.close(); return p

def main():
    ap=argparse.ArgumentParser(); ap.add_argument('--check',action='store_true'); args=ap.parse_args()
    p=read_params(); cap=int(p['角色升级等级上限']); first=p['角色升级首级经验']; gratio=p['角色升级金币比']

    wb=openpyxl.load_workbook(SRC, data_only=True); ws=wb[SRC_SHEET]
    steps=[ws.cell(DATA_ROW0+i, C_STEP).value for i in range(cap)]  # 输入: 相对提升

    need=[]; totexp=[]; gold=[]; totgold=[]; ce=0; cg=0
    for i in range(cap):
        nd = first if i==0 else need[-1] + (steps[i] or 0)
        need.append(nd); ce += nd; totexp.append(ce)
        g = nd * gratio; gold.append(g); cg += g; totgold.append(cg)

    # 对照
    e=dict(每级经验=0, 总经验=0, 每级金币=0, 总金币=0)
    for i in range(cap):
        r=DATA_ROW0+i
        for key,col,val in [('每级经验',C_NEED,need[i]),('总经验',C_TOTEXP,totexp[i]),
                            ('每级金币',C_GOLD,gold[i]),('总金币',C_TOTGOLD,totgold[i])]:
            a=ws.cell(r,col).value
            if a is not None: e[key]=max(e[key], abs(val-a))
    wb.close()

    if not args.check:
        owb=openpyxl.Workbook(); bold=Font(bold=True); hf=PatternFill('solid',fgColor='D9E1F2'); der=PatternFill('solid',fgColor='E2EFDA')
        ows=owb.active; ows.title='角色升级成本'
        ows.append(['【自动生成·勿手改】每级经验=首级+累计相对提升; 金币=经验×%g; 总列=累计'%gratio])
        ows.append([]); ows.append(['等级','相对上级提升(输入)','每级所需经验','总经验','每级所需金币','总金币'])
        for c in range(1,7): ows.cell(3,c).font=bold; ows.cell(3,c).fill=hf
        for i in range(cap):
            ows.append([i+1, steps[i], need[i], totexp[i], gold[i], totgold[i]])
            for c in (3,4,5,6): ows.cell(ows.max_row,c).fill=der
        owb.save(OUT)

    md=[]
    def o(s=''):
        print(s); md.append(s)
    o('# 经济 · 角色升级成本 生成与对照'); o()
    o('> 每级经验=首级%g+累计相对提升; 金币=经验×%g; 总列=累计。相对提升为手调输入。源表只读。' % (first, gratio)); o()
    allok = all(v < 1e-6 for v in e.values())
    o('- 对照 %d 级。各派生列最大误差: 每级经验 %.3g, 总经验 %.3g, 每级金币 %.3g, 总金币 %.3g' %
      (cap, e['每级经验'], e['总经验'], e['每级金币'], e['总金币'])); o()
    o('- 末级总经验 %d, 总金币 %d' % (totexp[-1], int(totgold[-1]))); o()
    o('✅ **逐格复现**' if allok else '🔴 存在偏差，需复查')
    os.makedirs(os.path.dirname(REPORT),exist_ok=True)
    with open(REPORT,'w',encoding='utf-8') as f: f.write('\n'.join(md))
    print('\n[报告] %s'%REPORT)
    if not args.check: print('[生成表] %s'%OUT)

if __name__=='__main__':
    main()
